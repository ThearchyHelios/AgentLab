"""Excel / CSV → 可以用 SQL 查的表。

这块存在的理由不是"多支持一种格式"。这个项目有一条不可回退的决策：叙述层
无算术权限，所有算术下沉到 SQL 或口径卡（engine/issuance.py）。表格传进知识库
只能被切块检索，数字就成了模型从片段里"读"出来的——复核层拦得住凭空多出的
数字，拦不住"从一堆片段里读错了一格"。

所以这里最要紧的一条不是"能不能导进去"，是**导进去的数字能不能真的被 SQL 算**：
一个整数列如果落成 TEXT，ORDER BY 就变成字典序（'10' < '9'），而没有任何
地方会报错。
"""

from __future__ import annotations

import io
import sqlite3

import pytest
from httpx import ASGITransport, AsyncClient

from app.data.tabular import UnsupportedTable, infer_type, load_into
from app.main import app


@pytest.fixture
async def client():
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as c:
        yield c


def xlsx(sheets: dict[str, list[list]]) -> bytes:
    from openpyxl import Workbook

    book = Workbook()
    book.remove(book.active)
    for title, rows in sheets.items():
        sheet = book.create_sheet(title)
        for row in rows:
            sheet.append(row)
    buf = io.BytesIO()
    book.save(buf)
    return buf.getvalue()


def query(db, sql):
    conn = sqlite3.connect(str(db))
    try:
        return conn.execute(sql).fetchall()
    finally:
        conn.close()


def col_types(db, table):
    return {r[1]: r[2] for r in query(db, f'PRAGMA table_info("{table}")')}


# --------------------------------------------------------------------------
# 类型推断：这块的全部意义
# --------------------------------------------------------------------------


def test_infer_type_basics():
    assert infer_type([1, 2, 3]) == "INTEGER"
    assert infer_type(["1", "2", "3"]) == "INTEGER"       # CSV 读出来全是字符串
    assert infer_type([1, 2.5]) == "REAL"
    assert infer_type(["1.5", "2"]) == "REAL"
    assert infer_type([1, "两个"]) == "TEXT"
    assert infer_type([]) == "TEXT"                        # 整列空的，别猜
    assert infer_type([None, "  ", None]) == "TEXT"
    # 空单元格不该把一个数字列拖成 TEXT——那正是 SUM() 静默出错的起点
    assert infer_type([1, None, 3]) == "INTEGER"
    assert infer_type([1, "", 3]) == "INTEGER"
    # 真/假不是数量，不该被当成 INTEGER 去求和
    assert infer_type([True, False]) == "TEXT"


def test_numbers_sort_as_numbers_not_as_text(tmp_path):
    """整块设计的落点：ORDER BY 一个数字列要是 9 < 10，不是 '10' < '9'。

    列落成 TEXT 的话这条会反过来，而且不报任何错——只是从此以后每一次
    排序、每一次 MAX()、每一次区间筛选都悄悄错着。
    """
    db = tmp_path / "t.db"
    load_into(str(db), xlsx({"数据": [["名称", "销量"], ["甲", 9], ["乙", 10], ["丙", 100]]}), "a.xlsx")

    table = query(db, "SELECT name FROM sqlite_master WHERE type='table'")[0][0]
    assert col_types(db, table)["销量"] == "INTEGER"
    assert [r[0] for r in query(db, f'SELECT "名称" FROM "{table}" ORDER BY "销量"')] == \
        ["甲", "乙", "丙"]
    assert query(db, f'SELECT SUM("销量") FROM "{table}"')[0][0] == 119


def test_empty_cells_become_null_not_empty_string(tmp_path):
    db = tmp_path / "t.db"
    load_into(str(db), xlsx({"s": [["a", "b"], [1, None], [2, "  "]]}), "a.xlsx")
    table = query(db, "SELECT name FROM sqlite_master WHERE type='table'")[0][0]
    assert query(db, f'SELECT COUNT("b") FROM "{table}"')[0][0] == 0   # COUNT 不数 NULL
    assert query(db, f'SELECT SUM("a") FROM "{table}"')[0][0] == 3


# --------------------------------------------------------------------------
# 表 / 列
# --------------------------------------------------------------------------


def test_multiple_sheets_become_multiple_tables(tmp_path):
    db = tmp_path / "t.db"
    report = load_into(str(db), xlsx({
        "销售明细": [["月份", "金额"], ["1月", 100]],
        "Summary": [["k", "v"], ["总计", 100]],
    }), "a.xlsx")
    assert len(report.tables) == 2
    # 中文 sheet 名原样留着。清洗成 ASCII 的话「明细」「汇总」都会变成
    # sheet、sheet_2——模型和用户都认不出哪张是哪张
    assert {t.name for t in report.tables} == {"销售明细", "Summary"}
    assert query(db, 'SELECT "金额" FROM 销售明细')[0][0] == 100   # 不加引号也查得了


def test_chinese_columns_are_kept_verbatim(tmp_path):
    """列名保持原样。转写是有损的，而 SQLite 认带引号的标识符。"""
    db = tmp_path / "t.db"
    load_into(str(db), xlsx({"s": [["本月销量 (万元)", "x"], [12, "a"]]}), "a.xlsx")
    table = query(db, "SELECT name FROM sqlite_master WHERE type='table'")[0][0]
    assert "本月销量 (万元)" in col_types(db, table)
    assert query(db, f'SELECT "本月销量 (万元)" FROM "{table}"')[0][0] == 12


def test_broken_headers_get_usable_names(tmp_path):
    db = tmp_path / "t.db"
    load_into(str(db), xlsx({"s": [["名称", None, "名称"], ["a", "b", "c"]]}), "a.xlsx")
    table = query(db, "SELECT name FROM sqlite_master WHERE type='table'")[0][0]
    cols = list(col_types(db, table))
    assert cols == ["名称", "col_2", "名称_2"], cols


def test_header_row_can_be_moved(tmp_path):
    """不猜表头行。猜错的表现是列名变成一行数据，而且没有任何征兆。"""
    db = tmp_path / "t.db"
    sheets = {"s": [["2026 年销售报表", None], ["月份", "金额"], ["1月", 100]]}

    load_into(str(db), xlsx(sheets), "a.xlsx", header_row=2)
    table = query(db, "SELECT name FROM sqlite_master WHERE type='table'")[0][0]
    assert list(col_types(db, table)) == ["月份", "金额"]

    # 取默认第 1 行就会把标题当表头——这正是要让用户当场看见的那种错
    load_into(str(db), xlsx(sheets), "a.xlsx")
    assert "2026 年销售报表" in col_types(db, table)


def test_ragged_rows_are_padded(tmp_path):
    db = tmp_path / "t.db"
    load_into(str(db), xlsx({"s": [["a", "b", "c"], [1], [1, 2, 3, 4]]}), "a.xlsx")
    table = query(db, "SELECT name FROM sqlite_master WHERE type='table'")[0][0]
    assert query(db, f'SELECT COUNT(*) FROM "{table}"')[0][0] == 2


# --------------------------------------------------------------------------
# CSV
# --------------------------------------------------------------------------


def test_csv_with_gbk_encoding(tmp_path):
    """中文环境里 GBK 的 CSV 极常见（Excel 另存为 CSV 的默认行为）。

    一律按 UTF-8 读会得到乱码列名，而乱码是不报错的——只是从此搜不到、看不懂。
    """
    db = tmp_path / "t.db"
    raw = "月份,金额\n1月,100\n".encode("gb18030")
    load_into(str(db), raw, "销售.csv")
    table = query(db, "SELECT name FROM sqlite_master WHERE type='table'")[0][0]
    assert "月份" in col_types(db, table)
    assert query(db, f'SELECT "金额" FROM "{table}"')[0][0] == 100


def test_tsv_and_semicolon_csv(tmp_path):
    db = tmp_path / "t.db"
    load_into(str(db), b"a\tb\n1\t2\n", "x.tsv")
    table = query(db, "SELECT name FROM sqlite_master WHERE type='table'")[0][0]
    assert list(col_types(db, table)) == ["a", "b"]


def test_unsupported_formats_say_what_to_do(tmp_path):
    db = tmp_path / "t.db"
    with pytest.raises(UnsupportedTable, match="另存为"):
        load_into(str(db), b"x", "old.xls")
    with pytest.raises(UnsupportedTable, match="只认"):
        load_into(str(db), b"x", "a.pdf")


# --------------------------------------------------------------------------
# 重传
# --------------------------------------------------------------------------


def test_reupload_replaces_instead_of_appending(tmp_path):
    """重传的语义是"这份数据更新了"。追加会让行数悄悄翻倍，没人会发现。"""
    db = tmp_path / "t.db"
    load_into(str(db), xlsx({"s": [["a"], [1], [2]]}), "a.xlsx")
    load_into(str(db), xlsx({"s": [["a"], [9]]}), "a.xlsx")
    table = query(db, "SELECT name FROM sqlite_master WHERE type='table'")[0][0]
    assert [r[0] for r in query(db, f'SELECT "a" FROM "{table}"')] == [9]


# --------------------------------------------------------------------------
# 接口
# --------------------------------------------------------------------------


async def _upload(client, name, content, filename="a.xlsx", **form):
    return await client.post(
        "/api/datasources/upload",
        files={"file": (filename, content,
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        data={"name": name, **form},
    )


@pytest.mark.asyncio
async def test_upload_creates_a_queryable_source(client):
    resp = await _upload(
        client, "tab_sales",
        xlsx({"明细": [["月份", "金额"], ["1月", 100], ["2月", 250]]}),
    )
    assert resp.status_code == 201, resp.text
    body = resp.json()
    assert body["source"]["kind"] == "sqlite"
    assert body["replaced"] is False
    # 列名和类型要原样回显：表头行取错了，用户当场就该看见
    cols = {c["name"]: c["type"] for c in body["tables"][0]["columns"]}
    assert cols == {"月份": "TEXT", "金额": "INTEGER"}
    assert body["tables"][0]["rows"] == 2
    # 建完立刻探查过，否则 db_query 工具的描述里没有表清单
    assert body["source"]["table_count"] >= 1

    # 真能查——而且走的是和别的数据源完全相同的那条路
    from app.data.engine import run_query
    from app.db.base import SessionLocal
    from app.db.models import DataSource
    from sqlalchemy import select as _select

    async with SessionLocal() as session:
        row = (await session.execute(
            _select(DataSource).where(DataSource.name == "tab_sales")
        )).scalar_one()
    result = await run_query(row, 'SELECT SUM("金额") AS s FROM "明细"')
    assert result.rows[0][0] == 350


@pytest.mark.asyncio
async def test_reupload_keeps_the_same_source_and_name(client):
    """重传不能换 id 或名字——名字是工具名（db_query__x），写进了保存过的图。"""
    first = await _upload(client, "tab_keep", xlsx({"s": [["a"], [1]]}))
    assert first.status_code == 201
    src_id = first.json()["source"]["id"]

    again = await _upload(client, "tab_keep", xlsx({"s": [["a"], [7], [8]]}))
    assert again.status_code == 201, again.text
    assert again.json()["replaced"] is True
    assert again.json()["source"]["id"] == src_id
    assert again.json()["source"]["name"] == "tab_keep"
    assert again.json()["tables"][0]["rows"] == 2


@pytest.mark.asyncio
async def test_upload_rejects_a_name_that_cannot_be_a_tool(client):
    resp = await _upload(client, "带中文的名字", xlsx({"s": [["a"], [1]]}))
    assert resp.status_code == 400
    assert "工具名" in resp.json()["detail"]


@pytest.mark.asyncio
async def test_upload_will_not_hijack_a_real_database(client):
    """同名的真数据库不能被一个 Excel 顶掉。"""
    from app.db.base import SessionLocal
    from app.db.models import DataSource

    async with SessionLocal() as session:
        session.add(DataSource(name="tab_real", kind="mysql", host="h", database="d"))
        await session.commit()

    resp = await _upload(client, "tab_real", xlsx({"s": [["a"], [1]]}))
    assert resp.status_code == 409
    assert "换个名字" in resp.json()["detail"]
