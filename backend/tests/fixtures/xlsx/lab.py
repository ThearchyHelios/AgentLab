"""9.3 的实验室难缠用例，按 xlsx-lab/case*.py 的结构重写成合成夹具：地区、产品、人名一律换成假名。

各函数返回 (xlsx 字节, 文件名)，不写磁盘。openpyxl 写不出来的形态（公式缓存值、XlsxWriter 式的「缓存为 0 加
打开时重算」）在压缩包层面补。期望（规则草稿 → 试运行）见 P2-SPEC 9.3，WP-7 验收。
"""
from __future__ import annotations

from openpyxl import Workbook
from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.utils.cell import coordinate_from_string

from . import excel_saved, save

ZONES = ("分区甲", "分区乙", "分区丙")
PRODUCTS = ("产品甲", "产品乙", "产品丙")

C01_VARIANTS = ("literal", "uncached", "cached")
C05_VARIANTS = ("uncached", "zero_full_calc", "cached")
C08_VARIANTS = ("stacked", "gap", "side")
C10_MONTHS = (7, 8, 9, 10)

#: 9.3 表里的叫法 → 这里的 variant
_ALIASES = {"写死": "literal", "公式无缓存": "uncached", "公式有缓存": "cached", "有缓存": "cached",
            "openpyxl": "uncached", "xlsxwriter": "zero_full_calc"}


def _variant(v: str, allowed: tuple[str, ...]) -> str:
    v = _ALIASES.get(v, v)
    if v not in allowed:
        raise ValueError(f"variant 只能是 {allowed}，收到 {v!r}")
    return v


#: c01 挪位（shift）时，表的上方另加的两行说明（不含数字：只让「挪了位置」这一件事不同，P3-SPEC 11.1）
C01_SHIFT_NOTES = ("说明：本表为合成测试数据", "说明：口径与上月相同")


def _shift_ref(ref: str, dr: int, dc: int) -> str:
    """「C5」或「C1:F1」整体平移 dr 行、dc 列。"""
    def one(a: str) -> str:
        col, row = coordinate_from_string(a)
        return f"{get_column_letter(column_index_from_string(col) + dc)}{row + dr}"
    return ":".join(one(a) for a in ref.split(":"))


def c01(variant: str = "literal", *, shift: tuple[int, int] = (0, 0)) -> tuple[bytes, str]:
    """表头在 C5，上方标题、单位、编制行，下方合计行、空行、备注。

    literal：合计写死；uncached：合计是公式、没有缓存值（openpyxl 直接写出的样子）；cached：补上缓存值。
    期望：列表块 + total_row(pick=合计, keep_as)；区域外含数字的文字只有 C1、C3 两格（outside_digits ×2）。

    shift=(行, 列)（期 3，框选回放验收 4.6 第 2 条）：整张表（含 C1:C3 的标题、单位、编制行和下方的备注、
    合计公式的引用）整体往下挪若干行、往右挪若干列；往下挪了 2 行及以上时，第 1、2 行另加两行不含数字的说明
    （C01_SHIFT_NOTES，写在挪位后的首列）。shift=(3, 2) 时表头落在 E8。默认 (0, 0) 与期 2 逐字节相同。
    """
    variant = _variant(variant, C01_VARIANTS)
    dr, dc = shift
    if dr < 0 or dc < 0:
        raise ValueError(f"shift 只能往下、往右挪，收到 {shift!r}")

    def at(ref: str) -> str:
        return _shift_ref(ref, dr, dc)

    book = Workbook()
    ws = book.active
    ws.title = "月报"
    ws[at("C1")] = "2026年8月 销售月报"
    ws.merge_cells(at("C1:F1"))
    ws[at("C2")] = "单位：万元"
    ws[at("C3")] = "编制：财务部    日期：2026-09-05"
    for i, h in enumerate(("地区", "产品", "销量", "金额")):
        ws.cell(5 + dr, 3 + dc + i, h)
    data = [(ZONES[0], PRODUCTS[0], 10, 100), (ZONES[1], PRODUCTS[1], 20, 120), (ZONES[2], PRODUCTS[2], 15, 80)]
    for r, row in enumerate(data, start=6):
        for i, v in enumerate(row):
            ws.cell(r + dr, 3 + dc + i, v)
    ws[at("C9")] = "合计"
    cached: dict[str, object] = {}
    if variant == "literal":
        ws[at("E9")], ws[at("F9")] = 45, 300
    else:
        ws[at("E9")], ws[at("F9")] = f"=SUM({at('E6:E8')})", f"=SUM({at('F6:F8')})"
        if variant == "cached":
            cached = {at("E9"): 45, at("F9"): 300}
    # 备注不含数字：这一格若带数字会多出一条 outside_digits，与 9.3 的期望（C1、C3 两格）不符
    ws[at("C11")] = "注：数据来源于业务系统，金额含税。"
    ws[at("C12")] = "制表人：经办甲"
    if dr >= 2:
        for i, note in enumerate(C01_SHIFT_NOTES):
            ws.cell(1 + i, 3 + dc, note)
    name = f"c01_{variant}.xlsx" if shift == (0, 0) else f"c01_{variant}_shifted.xlsx"
    return excel_saved(save(book), "月报", cached), name


def c04() -> tuple[bytes, str]:
    """合并单元格：A1:F1 标题；两行表头（地区、产品纵向合并，上下半年横向合并、下层销量 / 金额）；地区纵向合并。"""
    book = Workbook()
    ws = book.active
    ws.title = "合并"
    ws["A1"] = "分地区销售（万元）"
    ws.merge_cells("A1:F1")
    ws["A3"], ws["B3"], ws["C3"], ws["E3"] = "地区", "产品", "2026年上半年", "2026年下半年"
    for m in ("A3:A4", "B3:B4", "C3:D3", "E3:F3"):
        ws.merge_cells(m)
    for i, h in enumerate(("销量", "金额", "销量", "金额")):
        ws.cell(4, 3 + i, h)
    data = [(PRODUCTS[0], 1, 10, 2, 20), (PRODUCTS[1], 3, 30, 4, 40), (PRODUCTS[2], 5, 50, 6, 60),
            (PRODUCTS[0], 7, 70, 8, 80), (PRODUCTS[1], 9, 90, 1, 10), (PRODUCTS[2], 2, 20, 3, 30)]
    for r, row in enumerate(data, start=5):
        for i, v in enumerate(row):
            ws.cell(r, 2 + i, v)
    ws["A5"], ws["A8"] = ZONES[0], ZONES[1]
    ws.merge_cells("A5:A7")
    ws.merge_cells("A8:A10")
    return excel_saved(save(book), "合并", {}), "c04_merged.xlsx"


def c05(variant: str = "cached") -> tuple[bytes, str]:
    """公式与缓存值：金额列是公式（单价 × 数量），末行合计。

    uncached：openpyxl 直接写出，公式没有缓存值；zero_full_calc：模拟 XlsxWriter 的默认写法，缓存值全是 0、
    并设打开时重算（fullCalcOnLoad）；cached：补上正确的缓存值（模拟 Excel 保存过）。
    """
    variant = _variant(variant, C05_VARIANTS)
    book = Workbook()
    ws = book.active
    ws.title = "明细"
    rows = [("产品", "单价", "数量", "金额"), (PRODUCTS[0], 2.5, 4), (PRODUCTS[1], 3, 10), (PRODUCTS[2], 10, 1)]
    for row in rows:
        ws.append(list(row))
    for r in (2, 3, 4):
        ws.cell(r, 4, f"=B{r}*C{r}")
    ws["A5"], ws["D5"] = "合计", "=SUM(D2:D4)"
    raw = save(book)
    if variant == "uncached":
        return excel_saved(raw, "明细", {}), "c05_uncached.xlsx"
    if variant == "zero_full_calc":
        return excel_saved(raw, "明细", {"D2": 0, "D3": 0, "D4": 0, "D5": 0}, full_calc=True), "c05_zero_full_calc.xlsx"
    return excel_saved(raw, "明细", {"D2": 10, "D3": 30, "D4": 10, "D5": 50}), "c05_cached.xlsx"


def c06() -> tuple[bytes, str]:
    """隐藏：明细表第 4 行手动隐藏、C 列隐藏、自动筛选（地区 = 分区甲）隐藏其余行；另有 hidden、veryHidden 工作表。"""
    book = Workbook()
    ws = book.active
    ws.title = "明细"
    ws.append(["地区", "产品", "成本", "金额"])
    regions = [ZONES[0], ZONES[1], ZONES[0], ZONES[2], ZONES[0], ZONES[1], ZONES[0], ZONES[2], ZONES[0], ZONES[1]]
    for i, reg in enumerate(regions):
        ws.append([reg, f"型号{'甲乙丙丁戊己庚辛壬癸'[i]}", i, 10 * (i + 1)])
    ws.row_dimensions[4].hidden = True
    ws.column_dimensions["C"].hidden = True
    ws.auto_filter.ref = "A1:D11"
    ws.auto_filter.add_filter_column(0, [ZONES[0]])
    for r in range(2, 12):                  # Excel 保存筛选结果时就是给不满足条件的行打 hidden
        if ws.cell(r, 1).value != ZONES[0]:
            ws.row_dimensions[r].hidden = True
    draft = book.create_sheet("草稿")
    draft.append(["甲", "乙"])
    draft.append([1, 2])
    draft.sheet_state = "hidden"
    conf = book.create_sheet("配置")
    conf.append(["键", "值"])
    conf.append(["参数甲", "示例值"])
    conf.sheet_state = "veryHidden"
    return excel_saved(save(book), "明细", {}), "c06_hidden.xlsx"


def c08(variant: str = "stacked") -> tuple[bytes, str]:
    """同一工作表里的多张表。stacked：上下两张（各有小标题）；gap：一张表中间一行空白；side：左右两张、隔一空列。"""
    variant = _variant(variant, C08_VARIANTS)
    book = Workbook()
    ws = book.active
    if variant == "stacked":
        ws.title = "汇总"
        ws["B1"] = "一、销售"
        for r, row in enumerate([("地区", "销量", "金额"), (ZONES[0], 10, 100), (ZONES[1], 20, 120),
                                 (ZONES[2], 15, 80)], start=2):
            for c, v in enumerate(row, start=2):
                ws.cell(r, c, v)
        ws["B7"] = "二、费用"
        for r, row in enumerate([("部门", "费用"), ("部门甲", 5), ("部门乙", 7), ("部门丙", 8)], start=8):
            for c, v in enumerate(row, start=2):
                ws.cell(r, c, v)
        sheet = "汇总"
    elif variant == "gap":
        ws.title = "明细"
        for row in [("地区", "产品", "金额"), (ZONES[0], PRODUCTS[0], 10), (ZONES[0], PRODUCTS[1], 20),
                    (None, None, None), (ZONES[1], PRODUCTS[0], 30), (ZONES[1], PRODUCTS[1], 40)]:
            ws.append(list(row))
        sheet = "明细"
    else:
        ws.title = "并排"
        for r, (a, b) in enumerate([("地区", "金额"), (ZONES[0], 1), (ZONES[1], 2)], start=1):
            ws.cell(r, 1, a)
            ws.cell(r, 2, b)
        for r, (a, b) in enumerate([("部门", "费用"), ("部门甲", 5), ("部门乙", 7)], start=1):
            ws.cell(r, 4, a)
            ws.cell(r, 5, b)
        sheet = "并排"
    return excel_saved(save(book), sheet, {}), f"c08_{variant}.xlsx"


_C10_DATA = ((ZONES[0], PRODUCTS[0], 10, 100), (ZONES[1], PRODUCTS[1], 20, 120), (ZONES[2], PRODUCTS[2], 15, 80))


def c10(month: int) -> tuple[bytes, str]:
    """版式漂移（以 7 月起草并启用）：

    7 月：表头在 C5：地区 产品 销量 金额，下面合计行；
    8 月：表头挪到 C7（上面多两行说明），产品和销量之间多「单价」列 → column_extra；
    9 月：表头在 B6，外观差异「地区　」「销量\\n」「金 额」→ 通过、差异卡 header_writing；
    10 月：「金额」改名「销售额」→ column_missing 加 column_extra。
    """
    if month not in C10_MONTHS:
        raise ValueError(f"month 只能是 {C10_MONTHS}")
    top, left, headers, extra_rows, with_price = {
        7: (5, 3, ["地区", "产品", "销量", "金额"], 0, False),
        8: (7, 3, ["地区", "产品", "单价", "销量", "金额"], 2, True),
        9: (6, 2, ["地区　", "产品", "销量\n", "金 额"], 1, False),
        10: (5, 3, ["地区", "产品", "销量", "销售额"], 0, False),
    }[month]
    book = Workbook()
    ws = book.active
    ws.title = "月报"
    ws.cell(1, left, "销售月报")
    ws.cell(2, left, "单位：万元")
    for i in range(extra_rows):             # 说明行不含数字：只让表头的变化产生差异
        ws.cell(3 + i, left, f"说明{'甲乙丙'[i]}：口径同上月")
    for j, h in enumerate(headers):
        ws.cell(top, left + j, h)
    for r, (zone, prod, qty, amt) in enumerate(_C10_DATA, start=top + 1):
        vals = [zone, prod] + ([amt / qty] if with_price else []) + [qty, amt]
        for j, v in enumerate(vals):
            ws.cell(r, left + j, v)
    tot = top + 1 + len(_C10_DATA)
    ws.cell(tot, left, "合计")
    ws.cell(tot, left + len(headers) - 2, sum(d[2] for d in _C10_DATA))
    ws.cell(tot, left + len(headers) - 1, sum(d[3] for d in _C10_DATA))
    return excel_saved(save(book), "月报", {}), f"c10_2026-{month:02d}.xlsx"


def kv_form() -> tuple[bytes, str]:
    """键值表单式的版式（规则起草认不出、交给 AI 兜底起草的用例，9.7 第 10 步）：两列「项目：值」，没有表头行。"""
    book = Workbook()
    ws = book.active
    ws.title = "基本情况"
    ws["A1"] = "基本情况登记表"
    items = [("单位名称", "示例单位甲"), ("所在分区", ZONES[0]), ("从业人数", 120), ("营业面积", 860),
             ("年度客流", 35600), ("负责人", "经办甲")]
    for r, (k, v) in enumerate(items, start=3):
        ws.cell(r, 1, k)
        ws.cell(r, 2, v)
    return excel_saved(save(book), "基本情况", {}), "kv_form.xlsx"
