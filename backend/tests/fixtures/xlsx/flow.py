"""「结构仿照客流表」的合成交叉表（P2-SPEC 9.1）。名字全是假名（分区甲 / 分区乙 / 分区丙），数字随机。

坐标与真实报表一致（B2 统计期、第 4 行日期、第 5–7 行指标、第 8 行空、第 9 / 21 行分段标题、第 28–30 行
公式合计），这样执行器、核对和说明对着它跑出来的格子账、行数、核对结果才有可比的期望（FLOW_EXPECT）。
variant 取 9.2 的用例代号，在同一套生成逻辑上改出各种漂移：行的增删挪动都在「行清单」上做，行号和公式
引用跟着重新算，不在 XML 里硬改坐标。
"""
from __future__ import annotations

import datetime as dt
import random
from dataclasses import dataclass
from typing import Any

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from . import excel_saved, save

SHEET = "客流汇总"
TITLE = "客流汇总表"
PLACEHOLDER = "·"
TOTAL_LABEL, ZONE_A_LABEL, ZONE_B_LABEL = "全日客流（人次）", "分区甲（人次）", "分区乙（人次）"
ZONE_C_LABEL = "分区丙（人次）"
DAY_TITLE = "日间时段客流（人次）"
NIGHT_TITLE = "夜间时段客流（人次）"
DAY_HOURS = tuple(range(7, 18))
NIGHT_HOURS = tuple(range(18, 24))
#: 表内合计：(标签, 起始小时, 结束小时)
TOTALS = (("18-22 时合计", 18, 22), ("22-24 时合计", 22, 24), ("18-24 时合计", 18, 24))
LABEL_COL, FIRST_COL = 2, 3     # B 列是标签，C 列起是日期
EXTRA_SHEET = "说明"

#: 只换起止日期、版式与 D00 相同的代号
_PLAIN = {"D00", "D01", "D02", "D03", "D03b"}
#: 9.2 的全部代号，外加 noperiod（把 B2 删掉，9.7 第 6 步的人工录入统计期用）和期 3 的 D24w
#: （时段标签用全角数字和全角减号「７－８」，其余同 D01；P3-SPEC 11.1，验收 A5 用）
VARIANTS = frozenset(_PLAIN | {
    "D04", "D05", "D06", "D07", "D08", "D09", "D10", "D11", "D12", "D13", "D14", "D15", "D16", "D17", "D18",
    "D19", "D20", "D21", "D22", "D23", "D24", "D25", "D26", "D27", "D28", "D28b", "noperiod", "D24w",
})

#: 半角数字 → 全角数字（D24w）
_FULLWIDTH_DIGITS = str.maketrans("0123456789", "０１２３４５６７８９")

#: D00（2026-08-01 起 31 天、参考配方）的期望，照 P2-SPEC 9.1 抄录，WP-7 验收时逐项比
FLOW_EXPECT: dict[str, Any] = {
    "nonempty": 771,
    "bounds": "B2:AG30",
    "merged": 5,
    "formulas": 93,
    #: 格子账去向（契约 Role）：数据值、合计格、合计标签、列表头（日期）、行标签、分段标题、统计期、区域外文字
    "ledger": {"value": 620, "derived_value": 93, "derived_label": 3, "col_header": 31, "row_label": 20,
               "section_title": 2, "context": 1, "outside_text": 1},
    "rows": {"日客流": 31, "时段客流": 527, "时段客流_表内合计": 93},
    "placeholders": {PLACEHOLDER: 31},
    "outside_text": [("客流汇总!B3", TITLE, "text")],
    "period_cells": ["客流汇总!B2"],
    "checks": {"C1": "passed", "C2": "passed", "K1": "passed", "G1": "passed", "R1": "passed", "R2": "info",
               "N1": "passed", "N2": "passed", "N3": "passed", "P1": "passed", "P2": "passed", "P3": "passed"},
    #: R2：31 天中 0 天相等
    "r2_equal_days": 0,
    "canonicalized": 0,
    "confirms": frozenset({
        "placeholder:·", "derived:夜间合计", "year_from:交叉表", "relation:R1", "relation:R2",
        "unit:日客流.全日客流", "unit:日客流.分区甲", "unit:日客流.分区乙", "unit:时段客流.客流",
        "unit:时段客流_表内合计.客流", "mode",
    }),
    "note_has": ("已逐日核对",),
    "note_lacks": ("本期原表附有说明文字",),
}


@dataclass
class _Row:
    #: empty / period / text / axis / measure / blank / section / hour / total / subtotal / note / raw
    kind: str
    label: Any = None
    values: list[Any] | None = None
    merged: bool = False
    #: hour：小时；total / subtotal：(起始小时, 结束小时)
    key: Any = None
    r: int = 0


def flow_filename(start: dt.date, days: int, variant: str = "D00") -> str:
    end = start + dt.timedelta(days=days - 1)
    if variant == "D23":
        return f"{start.month}月客流.xlsx"
    return f"月报导出_{start.isoformat()}_{end.isoformat()}.xlsx"


def _hour_label(h: int, variant: str) -> str:
    if variant == "D11":
        return f"{h:02d}:00-{h + 1:02d}:00"
    if variant == "D24":
        return f"{h}–{h + 1}"          # en dash
    if variant == "D24w":               # 全角数字、全角减号（U+FF0D）：按规范写法存成「7-8」
        return f"{h}－{h + 1}".translate(_FULLWIDTH_DIGITS)
    return f"{h}-{h + 1}"


def _numbers(days: int, seed: int, variant: str):
    rnd = random.Random(seed)
    zone_a = [rnd.randint(3000, 9000) for _ in range(days)]
    zone_b = [rnd.randint(2000, 8000) for _ in range(days)]
    hours: dict[int, list[Any]] = {}
    for h in DAY_HOURS + NIGHT_HOURS:
        # 第 10 行（7-8）整行是占位符：真实报表早上这一档没有数
        hours[h] = [PLACEHOLDER] * days if h == 7 else [rnd.randint(50, 900) for _ in range(days)]
    # 变体专用的数另起一个随机源，基线的数不跟着变
    extra = random.Random(seed * 7919 + 13)
    zone_c = [extra.randint(1000, 5000) for _ in range(days)] if variant == "D13" else None
    total = [zone_a[i] + zone_b[i] + (zone_c[i] if zone_c else 0) for i in range(days)]
    if variant == "D26" and days >= 10:
        total[9] += 7                       # 第 10 个日期：全日 = 甲 + 乙 + 7
    # 保证每个日期「各时段之和 ≠ 全日客流」（R2 的口径不同在每一天都看得出来）
    for i in range(days):
        s = sum(v[i] for h, v in hours.items() if h != 7)
        if s == total[i]:
            hours[8][i] += 1 if hours[8][i] < 900 else -1
    return zone_a, zone_b, zone_c, total, hours


def _period_text(start: dt.date, end: dt.date, variant: str) -> str:
    if variant == "D12":                    # 第二个日期省略年份
        return f"统计时间范围：{start.year}年{start.month}月{start.day}日至{end.month}月{end.day}日"
    return f"统计时间范围：{start.year}年{start.month}月{start.day}日至{end.year}年{end.month}月{end.day}日"


def _title_text(start: dt.date, end: dt.date, variant: str) -> str:
    if variant == "D17":
        return f"{TITLE}（{start.year}年{start.month}月）"
    if variant == "D27":
        return f"{TITLE}（{start.year}年{start.month}月{start.day}日至{end.year}年{end.month}月{end.day}日）"
    return TITLE


def _layout(start: dt.date, end: dt.date, days: int, seed: int, variant: str) -> list[_Row]:
    zone_a, zone_b, zone_c, total, hours = _numbers(days, seed, variant)
    rows: list[_Row] = []
    if variant == "noperiod":
        rows.append(_Row("empty"))
    else:
        rows.append(_Row("period", _period_text(start, end, variant), merged=True))
    rows.append(_Row("text", _title_text(start, end, variant), merged=True))
    rows.append(_Row("axis"))

    measures = [_Row("measure", TOTAL_LABEL, total), _Row("measure", ZONE_A_LABEL, zone_a),
                _Row("measure", ZONE_B_LABEL, zone_b)]
    if zone_c is not None:                  # D13：多一个分区，全日 = 甲 + 乙 + 丙
        measures.append(_Row("measure", ZONE_C_LABEL, zone_c))
    if variant == "D19":                    # 顺序改成 甲、乙、全日
        measures = [measures[1], measures[2], measures[0]]
    blank = [] if variant == "D14" else [_Row("blank", merged=True)]

    titled = variant != "D15"               # D15：分段标题不合并
    day = [_Row("section", "日间分时段客流（人次）" if variant == "D09" else DAY_TITLE, merged=titled)]
    for h in DAY_HOURS:
        if variant == "D04" and h == 7:     # D04：日间少了 7-8
            continue
        day.append(_Row("hour", _hour_label(h, variant), hours[h], key=h))
    if variant == "D18":                    # 日间末尾多一行公式小计
        day.append(_Row("subtotal", "7-18 时合计", key=(7, 18)))
    night = [_Row("section", NIGHT_TITLE, merged=titled)]
    night += [_Row("hour", _hour_label(h, variant), hours[h], key=h) for h in NIGHT_HOURS]
    for label, s, e in TOTALS:
        night.append(_Row("total", f"{s}:00-{e}:00合计" if variant == "D22" else label, key=(s, e)))

    if variant == "D05":                    # 夜间块（标题、时段、合计）移到日间块之前
        rows += measures + blank + night + day
    elif variant == "D06":                  # 日客流三行移到最后，前面加标题
        rows += day + night + [_Row("empty"), _Row("text", "日客流（人次）")] + measures
    else:
        rows += measures + blank + day + night

    if variant == "D21":                    # 表下空一行补录一行数据
        rows += [_Row("empty"), _Row("raw", values=["补录（人次）", "1,234", "2,345"])]
    elif variant == "D28":
        rows += [_Row("empty"), _Row("note", f"注：{start.month}月15日闸机故障，当日客流为估算值")]
    elif variant == "D28b":
        rows += [_Row("empty"), _Row("note", f"注：{start.year}年{start.month}月数据为初步统计，待修订")]
    for i, row in enumerate(rows):
        row.r = 2 + i                       # 第 1 行空着，和真实报表一样从第 2 行开始
    return rows


def flow_workbook(start: dt.date, days: int, *, seed: int = 0, variant: str = "D00") -> tuple[bytes, str]:
    """返回 (xlsx 字节, 文件名)。variant 见 VARIANTS（9.2 的用例代号）。"""
    if variant not in VARIANTS:
        raise ValueError(f"不认识的 variant：{variant}")
    if days < 1:
        raise ValueError("days 至少为 1")
    end = start + dt.timedelta(days=days - 1)
    dates = [start + dt.timedelta(days=i) for i in range(days)]
    rows = _layout(start, end, days, seed, variant)
    last = FIRST_COL + days - 1
    last_letter = get_column_letter(last)
    hour_row = {row.key: row.r for row in rows if row.kind == "hour"}
    hour_vals = {row.key: row.values for row in rows if row.kind == "hour"}

    book = Workbook()
    ws = book.active
    ws.title = SHEET
    cached: dict[str, object] = {}
    computed: dict[int, list[int]] = {}       # 合计行 → 各日期的值（D16 的右侧合计列要用）

    def merge(r: int) -> None:
        ws.merge_cells(start_row=r, start_column=LABEL_COL, end_row=r, end_column=last)

    for row in rows:
        r = row.r
        if row.kind == "empty":
            continue
        if row.kind in ("period", "text", "section", "note"):
            ws.cell(r, LABEL_COL, row.label)
            if row.merged:
                merge(r)
        elif row.kind == "blank":
            if row.merged:
                merge(r)
        elif row.kind == "axis":
            for i, d in enumerate(dates):
                if variant == "D10":        # 表头是日期格，显示成「9月1日」
                    cell = ws.cell(r, FIRST_COL + i, dt.datetime(d.year, d.month, d.day))
                    cell.number_format = 'm"月"d"日"'
                else:
                    ws.cell(r, FIRST_COL + i, f"{d.month}月{d.day}日")
            if variant == "D16":
                ws.cell(r, last + 1, "合计")
        elif row.kind in ("measure", "hour"):
            ws.cell(r, LABEL_COL, row.label)
            for i, v in enumerate(row.values or []):
                ws.cell(r, FIRST_COL + i, v)
        elif row.kind in ("total", "subtotal"):
            ws.cell(r, LABEL_COL, row.label)
            s, e = row.key
            span = [h for h in range(s, e) if h in hour_row]
            r1, r2 = hour_row[span[0]], hour_row[span[-1]]
            assert r2 - r1 == len(span) - 1, "合计覆盖的时段行必须连续"
            vals = []
            for i in range(days):
                v = sum(x for h in span if isinstance(x := hour_vals[h][i], int))
                col = get_column_letter(FIRST_COL + i)
                literal = variant in ("D07", "D08") and row.kind == "total"
                if variant == "D08" and row.key == (18, 22) and i == 4:
                    v += 100                # 写死的「18-22 时合计」第 5 个日期多 100
                vals.append(v)
                if literal:
                    ws.cell(r, FIRST_COL + i, v)
                else:
                    ws.cell(r, FIRST_COL + i, f"=SUM({col}{r1}:{col}{r2})")
                    cached[f"{col}{r}"] = v
            computed[r] = vals
        elif row.kind == "raw":
            for j, v in enumerate(row.values or []):
                ws.cell(r, LABEL_COL + j, v)
        else:  # pragma: no cover - 行清单只会有上面这些
            raise AssertionError(row.kind)

    if variant == "D16":                    # 右侧多一列「合计」：每个数据行一个 SUM 公式，有缓存
        col = get_column_letter(last + 1)
        for row in rows:
            if row.kind in ("measure", "hour"):
                vals = [v for v in row.values or [] if isinstance(v, int)]
            elif row.kind in ("total", "subtotal"):
                vals = computed[row.r]
            else:
                continue
            ws.cell(row.r, last + 1, f"=SUM(C{row.r}:{last_letter}{row.r})")
            cached[f"{col}{row.r}"] = sum(vals)

    if variant == "D20":                    # 多一个有内容的可见工作表
        extra = book.create_sheet(EXTRA_SHEET)
        extra["A1"] = "本工作簿为合成测试数据，仅供测试使用"

    raw = excel_saved(save(book), SHEET, cached, full_calc=variant == "D25")
    return raw, flow_filename(start, days, variant)
