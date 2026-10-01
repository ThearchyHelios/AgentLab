"""配方执行器：列表（P2-SPEC 4.4、4.5、4.6 的列表部分、4.9 的流式路径）。

夹具是测试里现造的销售明细（地区甲/乙/丙、产品甲…，数字手写的假数），不依赖 WP-1 的夹具。
"""
from __future__ import annotations

import copy
import datetime as dt
import io
import os
import sqlite3
from collections.abc import Callable
from typing import Any

import pytest
from openpyxl import Workbook

from app.data import recipe_engine, xlsx_cells, xlsx_scan
from app.data.recipe_types import Extraction, Recipe

SHEET = "销售"
REGIONS = ("地区甲", "地区乙", "地区丙")

LIST = {
    "recipe_format": "agentlab-recipe/2",
    "sheets": [{
        "id": "s1", "match": {"name": SHEET},
        "blocks": [{
            "id": "列表1", "layout": "list", "table": "销售",
            "columns": [{"header": "地区", "name": "地区", "type": "TEXT"},
                        {"header": "产品", "name": "产品", "type": "TEXT"},
                        {"header": "日期", "name": "日期", "type": "DATE"},
                        {"header": "销量", "name": "销量", "type": "INTEGER"},
                        {"header": "金额（万元）", "name": "金额", "type": "REAL"}],
            "rows": {"total_row": {"label_column": "地区", "pick": "合计", "keep_as": "销售_表内合计"}},
        }],
    }],
    "tables": [{"name": "销售", "grain": ["地区", "产品"], "units": {"金额": "万元"}},
               {"name": "销售_表内合计", "grain": ["合计项"], "kind": "reported_total", "units": {"金额": "万元"}}],
}


def sales_book(n: int = 5, *, edit: Callable[[Any], None] | None = None, total: bool = True,
               title: str | None = "销售月报") -> tuple[bytes, str]:
    """A1 标题（合并 A1:E1）、第 3 行表头、第 4 行起 n 行明细、合计行（写死的数）。"""
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET
    if title:
        ws["A1"] = title
        ws.merge_cells("A1:E1")
    for c, h in enumerate(("地区", "产品", "日期", "销量", "金额（万元）"), 1):
        ws.cell(3, c, h)
    for i in range(n):
        r = 4 + i
        ws.cell(r, 1, REGIONS[i % 3])
        ws.cell(r, 2, f"产品{i:03d}")
        ws.cell(r, 3, dt.datetime(2026, 8, 1) + dt.timedelta(days=i))
        ws.cell(r, 4, 10 + i)
        ws.cell(r, 5, 1.5 * (i + 1))
    if total:
        r = 4 + n
        ws.cell(r, 1, "合计")
        ws.cell(r, 4, sum(10 + i for i in range(n)))
        ws.cell(r, 5, sum(1.5 * (i + 1) for i in range(n)))
    if edit is not None:
        edit(ws)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), "销售明细.xlsx"


def list_recipe(**block: Any) -> dict[str, Any]:
    data = copy.deepcopy(LIST)
    data["sheets"][0]["blocks"][0].update(block)
    return data


def run(book: tuple[bytes, str], recipe: dict[str, Any] | None = None, db: str | None = None,
        scan: xlsx_scan.WorkbookScan | None = None, **kw: Any) -> Extraction:
    raw, filename = book
    return recipe_engine.execute(Recipe.model_validate(recipe or LIST), raw, filename,
                                 scan or xlsx_scan.scan(raw), db, **kw)


def codes(ex: Extraction) -> set[str]:
    return {p.code for p in ex.problems}


def problem(ex: Extraction, code: str):
    found = [p for p in ex.problems if p.code == code]
    assert found, f"没有 {code}：{[(p.code, p.message) for p in ex.problems]}"
    return found[0]


def rows_of(db: str, sql: str) -> list[tuple[Any, ...]]:
    conn = sqlite3.connect(db)
    try:
        return conn.execute(sql).fetchall()
    finally:
        conn.close()


# --------------------------------------------------------------------------
# 基线
# --------------------------------------------------------------------------


def test_list_reference(tmp_path):
    db = str(tmp_path / "l.db")
    ex = run(sales_book(), db=db)
    assert ex.ok and not ex.problems
    assert ex.ledger[0].roles == {"col_header": 5, "value": 25, "total_label": 1, "total_value": 2, "outside_text": 1}
    assert ex.ledger[0].unclaimed == 0 and ex.ledger[0].nonempty_read == ex.ledger[0].nonempty_scan == 34
    assert ex.expected_rows == {"销售": 5, "销售_表内合计": 1}
    assert [c.header for c in ex.tables[0].columns] == ["地区", "产品", "日期", "销量", "金额（万元）"]
    assert ex.tables[0].sources == ["s1/列表1"] and ex.block_order == {SHEET: ["列表1"]}
    assert ex.lineage["销售"]["金额"] == [[1, SHEET, "E4", 5, "down"]]
    assert ex.lineage["销售_表内合计"]["销量"] == [[1, SHEET, "D9", 1, "right"]]
    assert [(d.kind, d.cell, d.base_value, d.value, d.rowid_first, d.rowid_last, d.label, d.keep_table)
            for d in ex.derived] == [
        ("column_sum", f"{SHEET}!D9", "销量", 60, 1, 5, "合计", "销售_表内合计"),
        ("column_sum", f"{SHEET}!E9", "金额", 22.5, 1, 5, "合计", "销售_表内合计")]
    assert rows_of(db, "SELECT rowid, * FROM 销售 ORDER BY rowid LIMIT 2") == [
        (1, "地区甲", "产品000", "2026-08-01", 10, 1.5), (2, "地区乙", "产品001", "2026-08-02", 11, 3.0)]
    assert rows_of(db, "SELECT * FROM 销售_表内合计") == [("合计", 60, 22.5)]
    assert [(o.cell, o.kind) for o in ex.outside_text] == [(f"{SHEET}!A1", "text")]


def test_text_column_dates_use_simple_import_writing(tmp_path):
    def stamp(ws):
        ws["C5"] = dt.datetime(2026, 8, 2, 9, 30)
    data = copy.deepcopy(LIST)
    data["sheets"][0]["blocks"][0]["columns"][2]["type"] = "TEXT"
    db = str(tmp_path / "t.db")
    ex = run(sales_book(edit=stamp), data, db=db)
    assert ex.ok
    assert [d for (d,) in rows_of(db, "SELECT 日期 FROM 销售 ORDER BY rowid LIMIT 2")] == [
        "2026-08-01", "2026-08-02 09:30:00"]          # 期 1 的 #23：零点日期不写成「… 00:00:00」


def test_single_integer_primary_key_is_not_a_rowid_alias(tmp_path):
    def codes_col(ws):
        for i, v in enumerate((100, 50, 75, 20, 90)):
            ws.cell(4 + i, 4, v)
    data = copy.deepcopy(LIST)
    data["tables"][0]["grain"] = ["销量"]
    data["sheets"][0]["blocks"][0]["rows"] = {}         # 夹具没有合计行：配方也不要合计行（否则报 total_row_missing）
    db = str(tmp_path / "i.db")
    ex = run(sales_book(edit=codes_col, total=False), data, db=db)
    assert ex.ok
    assert rows_of(db, "SELECT rowid, 销量 FROM 销售 ORDER BY rowid") == [(1, 100), (2, 50), (3, 75), (4, 20), (5, 90)]
    info = rows_of(db, "PRAGMA table_info('销售')")
    assert [name for _, name, _, _, _, pk in info if pk] == ["销量"]


# --------------------------------------------------------------------------
# 表头
# --------------------------------------------------------------------------


def test_header_codes():
    def rename_all(ws):
        for c in range(1, 6):
            ws.cell(3, c).value = f"列{c}"
    assert codes(run(sales_book(edit=rename_all))) == {"header_not_found"}
    def second(ws):
        for c, h in enumerate(("地区", "产品", "日期", "销量", "金额（万元）"), 1):
            ws.cell(20, c, h)
    p = problem(run(sales_book(edit=second)), "header_ambiguous")
    assert p.cells == [f"{SHEET}!A3", f"{SHEET}!A20"]
    ex = run(sales_book(edit=lambda ws: ws.cell(3, 5, "销售额（万元）")))      # c10 第 10 期：改名
    assert codes(ex) == {"column_missing", "column_extra"}
    assert "金额（万元）" in problem(ex, "column_missing").message
    ex = run(sales_book(edit=lambda ws: (ws.cell(3, 6, "单价"), ws.cell(4, 6, 3))))   # c10 第 08 期：多一列
    assert codes(ex) == {"column_extra"} and problem(ex, "column_extra").cells == [f"{SHEET}!F3"]


def test_header_writing_variants_still_match():
    def writing(ws):
        ws["E3"] = "金 额（万元）"
        ws["D3"] = "销量\n"
        ws["A3"] = "地区　"
    ex = run(sales_book(edit=writing))
    assert ex.ok
    # 本期读到的表头原文（首尾空白也保留），差异卡的 header_writing 才看得到这三处写法变化（9.3）
    assert [c.header for c in ex.tables[0].columns] == ["地区　", "产品", "日期", "销量\n", "金 额（万元）"]


def test_header_not_found_names_the_rows_actually_searched():
    def rename_all(ws):
        for c in range(1, 6):
            ws.cell(3, c).value = f"列{c}"
    p = problem(run(sales_book(edit=rename_all)), "header_not_found")
    assert "在第 1–9 行中没有找到表头" in p.message         # 已用区域只有 9 行，不写「第 1–200 行」
    wb = Workbook()
    wb.active.title = SHEET
    wb.active["A1"] = "销售月报"                          # 标题在最后一行：标题之后没有可找的行
    buf = io.BytesIO()
    wb.save(buf)
    p = problem(run((buf.getvalue(), "x.xlsx"), list_recipe(after_title="销售月报")), "header_not_found")
    assert "行中" not in p.message and "没有找到表头" in p.message


def test_extra_columns_ignored():
    def remark(ws):
        ws["F3"] = "备注"
        ws["F4"] = "加急"
    ex = run(sales_book(edit=remark), list_recipe(extra_columns="ignore"))
    assert ex.ok and ex.ignored_columns == {"列表1": ["备注"]}
    assert ex.ledger[0].roles["ignored_column"] == 2


def test_multi_row_header_composes_merged_spans(tmp_path):
    def build() -> tuple[bytes, str]:
        wb = Workbook()
        ws = wb.active
        ws.title = SHEET
        ws["A1"], ws["B1"], ws["C1"] = "地区", "产品", "2026年上半年"
        ws.merge_cells("A1:A2")
        ws.merge_cells("B1:B2")
        ws.merge_cells("C1:D1")
        ws["C2"], ws["D2"] = "销量", "金额"
        for i, region in enumerate(REGIONS):
            ws.cell(3 + i, 1, region)
            ws.cell(3 + i, 2, "产品甲")
            ws.cell(3 + i, 3, 5 + i)
            ws.cell(3 + i, 4, 2.5)
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue(), "上半年.xlsx"
    data = {"recipe_format": "agentlab-recipe/2",
            "sheets": [{"id": "s1", "match": {"name": SHEET}, "blocks": [{
                "id": "列表1", "layout": "list", "table": "销售", "header_rows": 2,
                "columns": [{"header": "地区", "name": "地区", "type": "TEXT"},
                            {"header": "产品", "name": "产品", "type": "TEXT"},
                            {"header": "2026年上半年_销量", "name": "上半年销量", "type": "INTEGER"},
                            {"header": "2026年上半年_金额", "name": "上半年金额", "type": "REAL"}]}]}],
            "tables": [{"name": "销售", "grain": ["地区"]}]}
    ex = run(build(), data, db=str(tmp_path / "h.db"))
    assert ex.ok, [(p.code, p.message) for p in ex.problems]
    assert [c.header for c in ex.tables[0].columns] == ["地区", "产品", "2026年上半年_销量", "2026年上半年_金额"]
    assert ex.ledger[0].roles["col_header"] == 5 and ex.expected_rows["销售"] == 3


def test_after_title_stacked_tables():
    def build(second_title: str = "二、费用", dup: bool = False) -> tuple[bytes, str]:
        wb = Workbook()
        ws = wb.active
        ws.title = SHEET
        ws["B1"] = "一、销售"
        ws["B2"], ws["C2"] = "项目", "金额"
        ws["B3"], ws["C3"] = "项目甲", 10
        ws["B4"], ws["C4"] = "项目乙", 20
        ws["B6"] = second_title
        ws["B7"], ws["C7"] = "项目", "金额"
        ws["B8"], ws["C8"] = "项目丙", 30
        if dup:
            ws["B10"] = "一、销售"
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue(), "两张表.xlsx"
    def block(bid: str, table: str, title: str) -> dict[str, Any]:
        return {"id": bid, "layout": "list", "table": table, "after_title": title,
                "columns": [{"header": "项目", "name": "项目", "type": "TEXT"},
                            {"header": "金额", "name": "金额", "type": "INTEGER"}]}
    data = {"recipe_format": "agentlab-recipe/2",
            "sheets": [{"id": "s1", "match": {"name": SHEET},
                        "blocks": [block("列表1", "销售", "一、销售"), block("列表2", "费用", "二、费用")]}],
            "tables": [{"name": "销售", "grain": ["项目"]}, {"name": "费用", "grain": ["项目"]}]}
    ex = run(build(), data)
    assert ex.ok and ex.expected_rows == {"销售": 2, "费用": 1}
    assert ex.block_order == {SHEET: ["列表1", "列表2"]}
    assert {o.text for o in ex.outside_text} == {"一、销售", "二、费用"}
    assert "title_not_found" in codes(run(build("二、支出"), data))
    p = problem(run(build(dup=True), data), "title_ambiguous")
    assert p.cells == [f"{SHEET}!B1", f"{SHEET}!B10"]


def _two_tables(*, upper_title: bool, same_header: bool) -> tuple[tuple[bytes, str], dict[str, Any]]:
    """上下两张表：上面一张（B1 可选标题「一、销售」），下面一张从 B7 起、没有标题。"""
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET
    if upper_title:
        ws["B1"] = "一、销售"
    ws["B2"], ws["C2"] = "地区", "销量"
    ws["B3"], ws["C3"] = "地区甲", 1
    ws["B4"], ws["C4"] = "地区乙", 2
    lower = ("地区", "销量") if same_header else ("部门", "费用")
    ws["B7"], ws["C7"] = lower
    ws["B8"], ws["C8"] = f"{lower[0]}丙", 3
    ws["B9"], ws["C9"] = f"{lower[0]}丁", 4
    buf = io.BytesIO()
    wb.save(buf)
    def block(bid: str, head: tuple[str, str], title: str | None) -> dict[str, Any]:
        return {"id": bid, "layout": "list", "table": bid, "after_title": title,
                "columns": [{"header": head[0], "name": head[0], "type": "TEXT"},
                            {"header": head[1], "name": head[1], "type": "INTEGER"}]}
    data = {"recipe_format": "agentlab-recipe/2",
            "sheets": [{"id": "s1", "match": {"name": SHEET}, "blocks": [
                block("上表", ("地区", "销量"), "一、销售" if upper_title else None),
                block("下表", lower, None)]}],
            "tables": [{"name": "上表", "grain": ["地区"]}, {"name": "下表", "grain": [lower[0]]}]}
    return (buf.getvalue(), "两张表.xlsx"), data


def test_block_without_title_below_a_titled_block_is_found():
    # 下表没写 after_title：窗口从第 1 行起，第 1 行正是上表的标题。窗口里还没有候选表头时不截断（评审：曾报
    # 「在第 1–0 行中没有找到表头」，这份合法的配方永远过不了）
    book, data = _two_tables(upper_title=True, same_header=False)
    ex = run(book, data)
    assert ex.ok and not ex.problems and ex.expected_rows == {"上表": 2, "下表": 2}
    assert ex.block_order == {SHEET: ["上表", "下表"]}
    flipped = copy.deepcopy(data)
    flipped["sheets"][0]["blocks"].reverse()
    ex = run(book, flipped)
    assert ex.ok and ex.expected_rows == {"下表": 2, "上表": 2}


def test_untitled_block_window_still_stops_at_the_next_title_after_its_header():
    # 上表没写标题、下表写了：两张表头相同。上表的窗口里已经有候选表头之后再碰到下表的标题，窗口到此为止，
    # 下表的同名表头不进上表的窗口，「恰好一处」判得出来
    book, data = _two_tables(upper_title=False, same_header=True)
    raw = book[0]
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(raw))
    wb.active["B6"] = "二、续表"
    buf = io.BytesIO()
    wb.save(buf)
    data["sheets"][0]["blocks"][1]["after_title"] = "二、续表"
    ex = run((buf.getvalue(), book[1]), data)
    assert ex.ok and not ex.problems and ex.expected_rows == {"上表": 2, "下表": 2}
    # 两张表都没写标题、表头又相同：照实报「多处都像表头」，不猜
    data["sheets"][0]["blocks"][1]["after_title"] = None
    assert "header_ambiguous" in codes(run((buf.getvalue(), book[1]), data))


def test_side_by_side_tables_split_by_blank_column():
    def build() -> tuple[bytes, str]:
        wb = Workbook()
        ws = wb.active
        ws.title = SHEET
        ws["A1"], ws["B1"], ws["D1"], ws["E1"] = "项目", "金额", "科目", "费用"
        ws["A2"], ws["B2"], ws["D2"], ws["E2"] = "项目甲", 1, "科目甲", 5
        ws["A3"], ws["B3"], ws["D3"], ws["E3"] = "项目乙", 2, "科目乙", 6
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue(), "左右.xlsx"
    data = {"recipe_format": "agentlab-recipe/2",
            "sheets": [{"id": "s1", "match": {"name": SHEET}, "blocks": [
                {"id": "列表1", "layout": "list", "table": "左表",
                 "columns": [{"header": "项目", "name": "项目", "type": "TEXT"},
                             {"header": "金额", "name": "金额", "type": "INTEGER"}]},
                {"id": "列表2", "layout": "list", "table": "右表",
                 "columns": [{"header": "科目", "name": "科目", "type": "TEXT"},
                             {"header": "费用", "name": "费用", "type": "INTEGER"}]}]}],
            "tables": [{"name": "左表", "grain": ["项目"]}, {"name": "右表", "grain": ["科目"]}]}
    ex = run(build(), data)
    assert ex.ok and ex.expected_rows == {"左表": 2, "右表": 2} and ex.ledger[0].unclaimed == 0
    # 占位符按块声明：左表声明的「—」在右表里不算占位符
    marked = copy.deepcopy(data)
    marked["sheets"][0]["blocks"][0]["values"] = {"placeholders": [{"text": "—"}]}
    def dash(ws):
        ws["E3"] = "—"
    raw, fn = build()
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(raw))
    dash(wb.active)
    buf = io.BytesIO()
    wb.save(buf)
    ex = run((buf.getvalue(), fn), marked)
    p = problem(ex, "value_not_number")
    assert p.cells == [f"{SHEET}!E3"] and ex.placeholders == {}
    one = copy.deepcopy(data)
    one["sheets"][0]["blocks"].pop()
    one["tables"].pop()
    p = problem(run(build(), one), "cell_unclaimed")       # 右边那张表的格在左表的数据行上，没有去处
    # 期 3 起单格没有修复按钮（P3-SPEC 3.1：单格没有能泛化的锚点，只能改文件或框选）
    assert p.fix is None and p.fix_args is None


# --------------------------------------------------------------------------
# 数据行：空行、合计行、停止之后
# --------------------------------------------------------------------------


def test_rows_after_stop_need_confirmation():
    def names(ws):
        ws["A11"], ws["B11"] = "名单甲", "名单乙"
        ws["A12"], ws["B12"] = "名单丙", "名单丁"
        ws["A14"] = "单独一格"                             # 只有一格：不算像数据的行
    ex = run(sales_book(total=False, edit=names), list_recipe(rows={}))
    p = problem(ex, "rows_after_stop")
    assert p.category == "confirm" and "第 11 行起有 2 行" in p.message
    assert p.cells == [f"{SHEET}!A11", f"{SHEET}!B11", f"{SHEET}!A12", f"{SHEET}!B12"]
    assert ex.ok


def test_header_without_data_rows_needs_confirmation():
    """AU-1：只有表头、没有数据行：不拒收（本月可能确实没有记录），但出 confirm 类问题，表里 0 行。"""
    def wipe(ws):
        for row in ws.iter_rows(min_row=4):
            for cell in row:
                cell.value = None
    ex = run(sales_book(total=False, edit=wipe), list_recipe(rows={}))
    p = problem(ex, "list_empty")
    assert p.category == "confirm" and "列表「列表1」" in p.message and "0 行" in p.message
    assert p.cells == [f"{SHEET}!A3"]
    assert ex.ok and next(t for t in ex.tables if t.name == "销售").rows == 0
    # 有数据行时不报
    assert not [x for x in run(sales_book(total=False), list_recipe(rows={})).problems if x.code == "list_empty"]


def test_rows_after_stop_in_a_one_column_list():
    # 只有 1 列的名单：空行之后的名字每行只有 1 格，也要出确认（规格写「2 格及以上」，1 列的块按 1 格算）
    wb = Workbook()
    ws = wb.active
    ws.title = "名单"
    ws["A1"], ws["A2"], ws["A3"], ws["A5"], ws["A6"] = "姓名", "甲某", "乙某", "丙某", "丁某"
    buf = io.BytesIO()
    wb.save(buf)
    data = {"recipe_format": "agentlab-recipe/2",
            "sheets": [{"id": "s1", "match": {"name": "名单"}, "blocks": [
                {"id": "名单", "layout": "list", "table": "名单",
                 "columns": [{"header": "姓名", "name": "姓名", "type": "TEXT"}]}]}],
            "tables": [{"name": "名单", "grain": []}]}
    ex = run((buf.getvalue(), "名单.xlsx"), data)
    p = problem(ex, "rows_after_stop")
    assert "第 5 行起有 2 行" in p.message and p.cells == ["名单!A5", "名单!A6"]
    assert ex.ok and ex.expected_rows == {"名单": 2}


def test_row_with_only_ignored_columns_is_not_a_record(tmp_path):
    # 紧跟数据的下一行只在忽略的「备注」列写了字：不是数据行，按空行处理（默认 stop），那一格进区域外文字。
    # 曾当成数据行写进一条全是空值的记录（没有主键时完全不出声）
    def remark(ws):
        ws["F3"] = "备注"
        ws["F4"] = "加急"
        ws["F9"] = "注：以上为初步统计"
    data = list_recipe(extra_columns="ignore", rows={})   # 夹具没有合计行
    data["tables"][0]["grain"] = []
    db = str(tmp_path / "r.db")
    ex = run(sales_book(edit=remark, total=False), data, db=db)
    assert ex.ok and not ex.problems and ex.expected_rows["销售"] == 5
    assert rows_of(db, "SELECT COUNT(*) FROM 销售 WHERE 地区 IS NULL") == [(0,)]
    assert (f"{SHEET}!F9", "text") in [(o.cell, o.kind) for o in ex.outside_text]
    assert ex.ledger[0].roles["ignored_column"] == 2 and ex.ledger[0].unclaimed == 0
    ex = run(sales_book(edit=remark, total=False), dict(data, sheets=[
        dict(data["sheets"][0], blocks=[dict(data["sheets"][0]["blocks"][0], rows={"blank_rows": "skip"})])]))
    assert ex.ok and ex.expected_rows["销售"] == 5 and ex.blank_rows_skipped == 1


def test_blank_row_stop_and_skip():
    def gap(ws):
        ws.insert_rows(6)
    ex = run(sales_book(edit=gap, total=False), list_recipe(rows={}))
    assert {"rows_after_stop", "outside_number"} <= codes(ex) and not ex.ok
    assert ex.expected_rows == {"销售": 2}
    ex = run(sales_book(edit=gap, total=False), list_recipe(rows={"blank_rows": "skip"}))
    assert ex.ok and ex.blank_rows_skipped == 1 and ex.expected_rows["销售"] == 5
    assert ex.lineage["销售"]["地区"] == [[1, SHEET, "A4", 2, "down"], [3, SHEET, "A7", 3, "down"]]


@pytest.mark.parametrize("stream", [False, True])
def test_note_below_a_skipped_blank_row_is_outside_text_not_a_record(tmp_path, monkeypatch, stream):
    """跳过空行（blank_rows=skip）时，表下隔一个空行、只有首列一格文字的说明行（「注：……」）是表的下边界，不是数据行
    （WP-8 修补）：照起草器 _compatible 的认法交给区域外，文字进回执；那个空行不算「数据中间的空行」，从计数和排除的
    行里撤回。曾静默写进一条首列是「注：……」、其余全为空值的记录，回执里只多一行。"""
    def gap_and_note(ws):
        ws.insert_rows(6)                      # 数据中间一行空白（第 6 行）：明细到第 9 行
        ws["G10"] = "附注"                     # 第 10 行只在块外有字：先按空行跳过，随后要一并撤回
        ws["A11"] = "注：以上为初步统计，以终稿为准"
        ws["A12"] = "制表人：经办甲"
    data = list_recipe(rows={"blank_rows": "skip"})
    if stream:
        monkeypatch.setattr(xlsx_cells, "GRID_MAX_CELLS", 20)
    db = str(tmp_path / "n.db")
    ex = run(sales_book(edit=gap_and_note, total=False), data, db=db)
    assert ex.ok and not ex.problems, [(p.code, p.message) for p in ex.problems]
    assert ex.expected_rows["销售"] == 5 and ex.blank_rows_skipped == 1
    assert rows_of(db, "SELECT COUNT(*) FROM 销售 WHERE 产品 IS NULL") == [(0,)]
    assert [(o.cell, o.kind) for o in ex.outside_text if o.cell.endswith(("11", "12"))] == [
        (f"{SHEET}!A11", "text"), (f"{SHEET}!A12", "text")]
    # 判成表的下边界也看得出来（WP-8 评审意见 4）：说明那一行和其下交给区域外的行（第 12 行的落款）记进排除的行
    # after_stop，锚点是说明那格的文字、格数是本块列范围内的非空格；像说明的不出问题
    assert [(x.reason, x.rows, x.cells, x.anchor, x.block) for x in ex.rows_excluded] == [
        ("blank_skipped", [[6, 6]], 0, None, "列表1"),
        ("after_stop", [[11, 12]], 2, "注：以上为初步统计，以终稿为准", "列表1")]
    # 说明行之后再有数据：照样不读，数字报 outside_number（响亮地拒收，而不是悄悄截断或读进去）
    def note_then_data(ws):
        gap_and_note(ws)
        ws["D12"] = 987654
    ex = run(sales_book(edit=note_then_data, total=False), data)
    assert "outside_number" in codes(ex) and ex.expected_rows["销售"] == 5
    stop = problem(ex, "rows_after_stop")
    assert "表下说明（第 11 行）之后" in stop.message and "跳过继续" not in stop.message, stop.message
    assert [(x.rows, x.cells) for x in ex.rows_excluded if x.reason == "after_stop"] == [([[11, 12]], 3)]


def _gap_and_lone(*tail: tuple[str, Any]) -> Callable[[Any], None]:
    """数据中间一行空白（第 6 行，明细到第 9 行），第 10 行空，第 11 行起按 tail 逐行写（(格, 值)，值为 None 则这一行空着）。"""
    def edit(ws):
        ws.insert_rows(6)
        for i, (cell, value) in enumerate(tail):
            if value is not None:
                ws[f"{cell}{11 + i}"] = value
    return edit


def _seen(ex: Extraction) -> tuple[Any, ...]:
    """两条路径要一致的部分：问题（code、类别、消息、格）、排除的行、区域外文字、各表行数、跳过的空行数。"""
    return ([(p.code, p.category, p.message, p.cells) for p in ex.problems],
            [(x.reason, x.rows, x.cells, x.anchor, x.block) for x in ex.rows_excluded],
            [(o.cell, o.text, o.kind) for o in ex.outside_text], ex.expected_rows, ex.blank_rows_skipped)


def test_a_lone_first_column_row_that_does_not_read_like_a_note_must_be_confirmed(tmp_path, monkeypatch):
    """空行之后只有首列一格文字的行照旧当表的下边界（与起草器一致），但文字不像说明（不以「注」「说明」「备注」……开头）
    时，它也可能是表尾一条只填了首列的数据行：出 confirm 类的 rows_after_stop（消息写明第几行只有首列文字、已当作表下
    说明、没有导入），并记进排除的行（WP-8 评审意见 4）。曾只出现在区域外文字里，确认项、排除的行都没有它。网格路径和
    流式路径结果逐项相同；下面再接一行只有首列的「地区戊」，一并记进排除的行。"""
    data = list_recipe(rows={"blank_rows": "skip"})
    book = sales_book(edit=_gap_and_lone(("A", "地区丁"), ("A", "地区戊")), total=False)
    db = str(tmp_path / "g.db")
    grid = run(book, data, db=db)
    monkeypatch.setattr(xlsx_cells, "GRID_MAX_CELLS", 20)
    stream = run(book, data, db=str(tmp_path / "s.db"))
    assert _seen(grid) == _seen(stream)
    for ex in (grid, stream):
        assert ex.ok and ex.expected_rows["销售"] == 5 and ex.blank_rows_skipped == 1
        assert [p.code for p in ex.problems] == ["rows_after_stop"], [(p.code, p.message) for p in ex.problems]
        p = ex.problems[0]
        assert p.category == "confirm" and p.cells == [f"{SHEET}!A11"]
        assert "第 11 行只有首列文字" in p.message and "已当作表下说明，没有导入" in p.message, p.message
        assert "地区丁" not in p.model_message
        assert [(x.reason, x.rows, x.cells, x.anchor) for x in ex.rows_excluded] == [
            ("blank_skipped", [[6, 6]], 0, None), ("after_stop", [[11, 12]], 2, "地区丁")]
        assert [(o.cell, o.text) for o in ex.outside_text if o.cell != f"{SHEET}!A1"] == [
            (f"{SHEET}!A11", "地区丁"), (f"{SHEET}!A12", "地区戊")]
    assert rows_of(db, "SELECT COUNT(*) FROM 销售 WHERE 地区 IN ('地区丁', '地区戊')") == [(0,)]


@pytest.mark.parametrize("text", ["注：本表为初步统计", "说明：不含退货", "　备注：见附件", "\u200b注：零宽字符开头",
                                  "制表：经办甲", "填表人：经办乙",
                                  "填报单位：合成单位", "来源：合成系统", "数据来源：合成系统", "单位：万元",
                                  "审核：经办丙", "ＮＯＴＥ"])
def test_which_lone_rows_read_like_a_note(text):
    """像说明的开头按 canon 之后比（全角空格、零宽字符开头的照样认）：像说明的只记进排除的行和区域外文字，不出问题；
    不像的（这里用全角字母写的 NOTE）出 rows_after_stop。"""
    ex = run(sales_book(edit=_gap_and_lone(("A", text)), total=False), list_recipe(rows={"blank_rows": "skip"}))
    assert ex.expected_rows["销售"] == 5
    assert [(x.reason, x.rows, x.anchor) for x in ex.rows_excluded if x.reason == "after_stop"] == [
        ("after_stop", [[11, 11]], text.strip())]
    assert (codes(ex) == {"rows_after_stop"}) is (text == "ＮＯＴＥ"), [(p.code, p.message) for p in ex.problems]


def test_lone_first_column_rows_that_are_data_or_total_are_still_read():
    """只认「首列一格文字」：首列是数、日期，或者正是合计行的标签时仍按原样读（缺数的数据行、合计行），
    中间的空行照旧算跳过。"""
    data = list_recipe(rows={"blank_rows": "skip"})
    data["tables"][0]["grain"] = []
    for cell, value in (("C11", dt.datetime(2026, 8, 30)), ("A11", "2026-08-30"), ("A11", "1,234")):
        def gap_and_lone(ws, cell=cell, value=value):
            ws.insert_rows(6)
            ws[cell] = value
        ex = run(sales_book(edit=gap_and_lone, total=False), data)
        assert ex.expected_rows["销售"] == 6 and ex.blank_rows_skipped == 2, (cell, value)
    def gap_and_bare_total(ws):
        ws.insert_rows(6)
        ws.insert_rows(10)                     # 合计行（第 11 行）上方也空一行，而且只写了「合计」两个字
        ws["D11"], ws["E11"] = None, None
    ex = run(sales_book(edit=gap_and_bare_total), list_recipe(rows={"blank_rows": "skip", "total_row": {
        "label_column": "地区", "pick": "合计", "keep_as": "销售_表内合计"}}))
    assert "total_row_missing" not in codes(ex), [(p.code, p.message) for p in ex.problems]
    assert ex.expected_rows == {"销售": 5, "销售_表内合计": 1} and ex.blank_rows_skipped == 2


def test_total_row_formulas_and_hidden_rows():
    def subtotal(ws):
        ws["D9"] = "=SUBTOTAL(9,D4:D8)"
        ws["E9"] = "=SUM(E4:E8)"
    ex = run(sales_book(edit=subtotal))
    assert ex.ok
    d, e = ex.derived
    assert (d.func, d.ref_problem, d.value, d.ref_rowids) == ("SUBTOTAL", "hidden", None, [1, 2, 3, 4, 5])
    assert (e.func, e.ref_problem, e.is_formula) == ("SUM", None, True)
    def hide(ws):
        ws.row_dimensions[5].hidden = True
    ex = run(sales_book(edit=hide), copy.deepcopy(LIST) | {"sheets": [
        dict(LIST["sheets"][0], hidden={"rows": "exclude"})]})
    assert ex.ok and ex.expected_rows["销售"] == 4
    assert all(d.hidden_in_range and (d.rowid_first, d.rowid_last) == (1, 4) for d in ex.derived)
    assert ex.hidden == {SHEET: {"rows": [5], "cols": [], "policy_rows": "exclude", "policy_cols": "reject_if_any"}}
    ex = run(sales_book(edit=hide))
    assert "hidden_rows" in codes(ex)


# --------------------------------------------------------------------------
# 值
# --------------------------------------------------------------------------


def test_placeholder_in_a_number_primary_key_is_rejected():
    # INTEGER 主键列里的占位符按空值存，主键不能为空：曾静默丢掉这一行（ok=True、没有问题、少一行）
    data = copy.deepcopy(LIST)
    data["tables"][0]["grain"] = ["销量"]
    data["sheets"][0]["blocks"][0]["values"] = {"placeholders": [{"text": "·"}]}
    data["sheets"][0]["blocks"][0]["rows"] = {}         # 夹具没有合计行
    ex = run(sales_book(edit=lambda ws: ws.cell(5, 4, "·"), total=False), data)
    assert not ex.ok and codes(ex) == {"value_blank"}
    p = problem(ex, "value_blank")
    assert p.cells == [f"{SHEET}!D5"] and "占位符" in p.message and "·" not in p.model_message
    assert ex.placeholders == {}
    data["tables"][0]["grain"] = []                      # 不是主键：照常存空值、计数
    ex = run(sales_book(edit=lambda ws: ws.cell(5, 4, "·"), total=False), data)
    assert ex.ok and ex.placeholders == {"·": 1} and ex.expected_rows["销售"] == 5


def test_total_row_configured_but_missing_or_blank_is_reported(tmp_path):
    """配方里有合计行：本期没有找到（或合计行的数字格全空）时明确拒收，不能沉默——否则核对里没有 T、说明照旧
    写「原表的合计行未导入本表」，另存的合计表是空的（或多一行空值）。部分干跑合计行可能还在窗口之外，不报。"""
    db = str(tmp_path / "m.db")
    ex = run(sales_book(total=False), LIST, db=db)
    p = problem(ex, "total_row_missing")
    assert not ex.ok and p.category == "structure" and not os.path.exists(db)
    assert "「地区」列以「合计」开头" in p.message and "取消「有合计行」" in p.message and p.cells == [f"{SHEET}!A8"]
    assert not any(d.kind == "column_sum" for d in ex.derived)
    # 合计行在、数字格全空
    def blank(ws):
        ws["D9"] = ws["E9"] = None
    ex = run(sales_book(edit=blank), LIST)
    p = problem(ex, "total_row_blank")
    assert not ex.ok and p.category == "structure" and p.cells == [f"{SHEET}!A9"] and "全部为空" in p.message
    assert "total_row_missing" not in codes(ex)
    # 只空了一列：照常核对另一列（不报）
    def one_blank(ws):
        ws["D9"] = None
    ex = run(sales_book(edit=one_blank), list_recipe())
    assert "total_row_blank" not in codes(ex) and "total_row_missing" not in codes(ex)
    # 部分干跑：只读前几行，合计行在窗口之外，不报
    ex = run(sales_book(n=40), LIST, max_rows=10)
    assert ex.partial and "total_row_missing" not in codes(ex)


def test_blank_number_in_kept_total_primary_key_is_reported():
    # 合计行另存的表把数字列放进主键：合计格是空的，那一行不能静默不存
    data = copy.deepcopy(LIST)
    data["tables"][1]["grain"] = ["合计项", "销量"]
    def blank(ws):
        ws["D9"] = None
    ex = run(sales_book(edit=blank), data)
    assert not ex.ok and codes(ex) == {"value_blank"}
    p = problem(ex, "value_blank")
    assert p.cells == [f"{SHEET}!D9"] and "主键列「销量」" in p.message


@pytest.mark.parametrize("ref,value,block,expected", [
    ("D4", None, {}, None),                                           # 列表的空格默认存空值
    ("D4", None, {"values": {"blank": "reject"}}, "value_blank"),
    ("A4", None, {}, "value_blank"),                                  # 主键列的空格一律拒收
    ("D4", "1,234", {}, "value_text_number"),
    ("D4", "1,234", {"values": {"text_number": "parse_thousands"}}, 1234),
    ("D4", "—", {"values": {"placeholders": [{"text": "—", "meaning": "不适用"}]}}, None),
    ("D4", 2.5, {}, "value_not_integer"),
    ("D4", "#REF!", {}, "value_error"),
    ("D4", "若干", {}, "value_not_number"),
    ("C4", "2026-08-09", {}, "2026-08-09"),
    ("C4", "8月9日", {}, "value_not_date"),                            # 列表的日期必须自带年份
    ("C4", "8/9", {}, "value_not_date"),
    ("C4", dt.datetime(2026, 8, 9, 10, 0), {}, "value_not_date"),
    ("C4", None, {"values": {"blank": "reject"}}, "value_blank"),
])
def test_list_value_conversion(tmp_path, ref, value, block, expected):
    def edit(ws):
        ws[ref] = value
    db = str(tmp_path / "v.db")
    ex = run(sales_book(edit=edit), list_recipe(**block), db=db)
    if isinstance(expected, str) and expected.startswith("value_"):
        assert codes(ex) == {expected} and not ex.ok
        return
    assert ex.ok, [(p.code, p.message) for p in ex.problems]
    column = {"C": "日期", "D": "销量"}[ref[0]]
    (got,) = rows_of(db, f"SELECT {column} FROM 销售 WHERE rowid = 1")[0]
    assert got == expected


def test_list_formulas(tmp_path):
    from tests.test_recipe_engine_crosstab import inject_cache

    def amount(ws):
        for r in range(4, 9):
            ws.cell(r, 5).value = f"=D{r}*0.15"
    raw, fn = sales_book(edit=amount)
    cached = inject_cache(raw, {f"E{r}": (10 + r - 4) * 0.15 for r in range(4, 9)})
    ex = run((cached, fn))
    assert ex.ok and ex.formula_cells_accepted == 5
    assert codes(run((cached, fn), list_recipe(values={"formula": "reject"}))) == {"value_formula"}
    assert codes(run((raw, fn))) == {"formula_uncached"}
    full = inject_cache(raw, {f"E{r}": 0 for r in range(4, 9)}, keep_full_calc=True)
    assert codes(run((full, fn))) == {"formula_full_calc"}


def test_canonical_text_and_canonicalized(tmp_path):
    def writing(ws):
        ws["A4"] = "地区甲 "                     # 首尾空白：as_text 已去掉，不算规范化
        ws["B5"] = "产品００１"                  # 全角数字：主键列按规范写法存
        ws["A6"] = "地区　丙"
    db = str(tmp_path / "c.db")
    ex = run(sales_book(edit=writing), db=db)
    assert ex.ok
    assert [(c.column, c.raw, c.canonical, c.count, c.first_cell) for c in ex.canonicalized] == [
        ("产品", "产品００１", "产品001", 1, f"{SHEET}!B5"), ("地区", "地区　丙", "地区 丙", 1, f"{SHEET}!A6")]
    assert rows_of(db, "SELECT 地区, 产品 FROM 销售 WHERE rowid IN (2, 3) ORDER BY rowid") == [
        ("地区乙", "产品001"), ("地区 丙", "产品002")]
    raw_store = copy.deepcopy(LIST)
    raw_store["sheets"][0]["blocks"][0]["columns"][1]["store"] = "raw"
    ex = run(sales_book(edit=writing), raw_store)
    assert [c.column for c in ex.canonicalized] == ["地区"]


def test_pk_duplicate_lists_both_cells():
    def dup(ws):
        ws["A5"], ws["B5"] = "地区甲", "产品000"
    p = problem(run(sales_book(edit=dup)), "pk_duplicate")
    assert "A4 和 A5" in p.message and p.cells == [f"{SHEET}!A4", f"{SHEET}!A5"]


def test_merged_data_reject_and_fill(tmp_path):
    def merged(ws):
        ws["A5"] = None
        ws.merge_cells("A4:A5")
    ex = run(sales_book(edit=merged))
    assert codes(ex) == {"merged_in_values"}                # 合并区里的空格不再另报为空
    db = str(tmp_path / "f.db")
    ex = run(sales_book(edit=merged), list_recipe(merged_data="fill"), db=db)
    assert ex.ok and rows_of(db, "SELECT 地区 FROM 销售 WHERE rowid = 2") == [("地区甲",)]
    def merged_number(ws):
        ws["D5"] = None
        ws.merge_cells("D4:D5")
    assert codes(run(sales_book(edit=merged_number), list_recipe(merged_data="fill"))) == {"merged_in_values"}


def test_empty_matched_sheet_reports_header_not_found():
    wb = Workbook()
    wb.active.title = SHEET
    wb.create_sheet("别的")["A1"] = "填表说明"
    buf = io.BytesIO()
    wb.save(buf)
    ex = run((buf.getvalue(), "空表.xlsx"))
    assert "header_not_found" in codes(ex) and not ex.ok


# --------------------------------------------------------------------------
# 流式路径、部分干跑
# --------------------------------------------------------------------------


def test_stream_path_matches_grid_path(tmp_path, monkeypatch):
    book = sales_book(n=40)
    grid = run(book, db=str(tmp_path / "g.db"))
    monkeypatch.setattr(xlsx_cells, "GRID_MAX_CELLS", 50)
    stream = run(book, db=str(tmp_path / "s.db"))
    assert grid.ok and stream.ok
    assert stream.table_hashes == grid.table_hashes
    assert stream.ledger == grid.ledger and stream.lineage == grid.lineage
    assert stream.derived == grid.derived and stream.regions == grid.regions
    assert stream.outside_text == grid.outside_text


def test_stream_path_layout_limits(monkeypatch):
    monkeypatch.setattr(xlsx_cells, "GRID_MAX_CELLS", 10)
    two = copy.deepcopy(LIST)
    second = copy.deepcopy(two["sheets"][0]["blocks"][0])
    second.update(id="列表2", table="销售", after_title="不存在")
    two["sheets"][0]["blocks"].append(second)
    assert codes(run(sales_book(), two)) == {"sheet_too_large_for_layout"}


def test_stream_path_problems(tmp_path, monkeypatch):
    monkeypatch.setattr(xlsx_cells, "GRID_MAX_CELLS", 10)
    raw, fn = sales_book()
    sc = xlsx_scan.scan(raw)
    a, b, n = sc.sheets[0].row_runs[-1]
    sc.sheets[0].row_runs[-1] = (a, b, n + 1)
    db = str(tmp_path / "s.db")
    ex = run((raw, fn), scan=sc, db=db)
    assert codes(ex) == {"pass_mismatch"} and not os.path.exists(db)
    ex = run(sales_book(edit=lambda ws: ws.cell(3, 5, "销售额")), db=db)
    assert codes(ex) == {"column_missing", "column_extra"} and not os.path.exists(db)
    def after(ws):
        ws["A12"], ws["B12"] = 987654, "说明"
    ex = run(sales_book(total=False, edit=after), list_recipe(rows={}))
    assert {"rows_after_stop", "outside_number"} <= codes(ex)


@pytest.mark.parametrize("stream", [False, True])
def test_regions_stay_rectangles_when_edge_columns_are_sparse(monkeypatch, stream):
    # 末列（金额）隔行为空、首列偶有空格：区域按块的列范围出矩形，不会每行一个（4.9「区域标记按矩形」）
    def sparse(ws):
        for i in range(400):
            r = 4 + i
            ws.cell(r, 1, REGIONS[i % 3] if i % 7 else None)
            ws.cell(r, 2, f"产品{i:03d}")
            ws.cell(r, 3, dt.datetime(2026, 8, 1))
            ws.cell(r, 4, i)
            ws.cell(r, 5, 1.5 if i % 2 else None)
    data = list_recipe(rows={})
    data["tables"][0]["grain"] = ["产品"]
    if stream:
        monkeypatch.setattr(xlsx_cells, "GRID_MAX_CELLS", 50)
    ex = run(sales_book(n=0, edit=sparse, total=False), data)
    assert ex.ok
    assert [(m.role, m.ref) for m in ex.regions] == [
        ("outside_text", "A1"), ("col_header", "A3:E3"), ("value", "A4:E403")]


def test_partial_dry_run_on_a_long_list():
    ex = run(sales_book(n=60), max_rows=10)
    assert ex.partial and ex.ok and not ex.problems
    assert ex.expected_rows == {"销售": 7, "销售_表内合计": 0}          # 第 1–10 行：表头在第 3 行
    assert ex.table_hashes == {}
