"""按配方导入期 3 的接口（WP-5，P3-SPEC 9.1、12.6）：真管线、经 HTTP。

覆盖按期累积的试运行与提交（追加、重叠拒收、替换该期、复用同一构建）、修复按钮（预览、应用、撤销、过期、
PUT 覆盖）、按规则重新起草与采用、改配方后回答的延续、StagingOut 与 TrialOut 的新字段。版本页的接口在
test_source_versions_api.py，错误码与遗留项在 test_p3_api_errors.py。

夹具一律合成、假名（tests/fixtures/xlsx 的客流表：分区甲、分区乙……），数字是随机数；不调用任何模型。
"""
from __future__ import annotations

import json
import uuid
from pathlib import Path
from typing import Any

import pytest
from httpx import ASGITransport, AsyncClient
from sqlalchemy import select

from app.core import artifact_store
from app.core.config import settings
from app.data import table_versions
from app.db.base import SessionLocal
from app.db.models import ImportStaging, SourceSnapshot, TableImport
from app.main import app
from tests.fixtures.xlsx.drift import AUG, DRIFT_CASES, OCT, SEP
from tests.fixtures.xlsx.flow import flow_workbook

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
API = "/api/datasources"
TABLES = ("日客流", "时段客流", "时段客流_表内合计")


@pytest.fixture(autouse=True)
def store(monkeypatch, tmp_path):
    data = tmp_path / "data"
    (data / "uploads").mkdir(parents=True)
    monkeypatch.setattr(settings, "data_dir", data)
    yield data


@pytest.fixture
async def client():
    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as c:
        yield c


# ==========================================================================
# 小工具（经 HTTP）
# ==========================================================================


def unique(prefix: str = "p3") -> str:
    return f"{prefix}_{uuid.uuid4().hex[:10]}"


def rename_wide(rec: dict[str, Any], old: str, new: str) -> dict[str, Any]:
    """把规则起草的宽表名改成参考配方的名字（写法同 test_recipe_e2e.rename_wide）。"""
    rec = json.loads(json.dumps(rec, ensure_ascii=False))
    for sheet in rec["sheets"]:
        for block in sheet["blocks"]:
            for seg in block.get("segments", []):
                if seg.get("table") == old:
                    seg["table"] = new
                if seg.get("id") == old:
                    seg["id"] = new
    for t in rec["tables"]:
        if t["name"] == old:
            t["name"] = new
    for r in rec["relations"]:
        if r.get("table") == old:
            r["table"] = new
        for side in ("a", "b"):
            if isinstance(r.get(side), dict) and r[side].get("table") == old:
                r[side]["table"] = new
    return rec


async def stage(client, name: str, raw: bytes, filename: str):
    return await client.post(f"{API}/imports/stage", files={"file": (filename, raw, XLSX)}, data={"name": name})


async def reupload(client, source_id: str, raw: bytes, filename: str) -> dict[str, Any]:
    resp = await client.post(f"{API}/{source_id}/reupload", files={"file": (filename, raw, XLSX)})
    assert resp.status_code == 201, resp.text
    return resp.json()


async def trial(client, sid: str, **body: Any) -> dict[str, Any]:
    resp = await client.post(f"{API}/imports/{sid}/trial", json=body)
    assert resp.status_code == 200, resp.text
    return resp.json()


def confirms(st: dict[str, Any]) -> list[str]:
    return [i["id"] for i in (st.get("trial") or {}).get("confirm_items") or []]


async def commit(client, st: dict[str, Any], *, confirmations: list[str] | None = None,
                 acceptances: list[dict[str, str]] | None = None):
    body: dict[str, Any] = {"trial_id": st["trial"]["trial_id"],
                            "confirmations": confirms(st) if confirmations is None else confirmations}
    if acceptances is not None:
        body["acceptances"] = acceptances
    return await client.post(f"{API}/imports/{st['id']}/commit", json=body)


def answers_for(questions: list[dict[str, Any]], mode: str) -> dict[str, dict[str, Any]]:
    out = {}
    for q in questions:
        values = [o["value"] for o in q["options"]]
        if q["id"].startswith("q_relation"):
            out[q["id"]] = {"value": "register"}
        elif q["id"] == "q_mode":
            out[q["id"]] = {"value": mode}
        elif "null" in values:
            out[q["id"]] = {"value": "null"}
        else:
            out[q["id"]] = {"value": values[0]}
    return out


async def first_import(client, name: str, *, seed: int, mode: str = "accumulate", start=AUG, days: int = 31,
                       variant: str = "D00") -> dict[str, Any]:
    """首次导入（回答问题、宽表改名「日客流」、试运行、提交），q_mode 按 mode 答。返回提交响应和试运行。"""
    raw, fn = flow_workbook(start, days, seed=seed, variant=variant)
    resp = await stage(client, name, raw, fn)
    assert resp.status_code == 201, resp.text
    st = resp.json()
    assert st["questions"][-1]["id"] == "q_mode", "q_mode 固定排在问题列表最后"
    resp = await client.post(f"{API}/imports/{st['id']}/answers",
                             json={"answers": answers_for(st["questions"], mode)})
    assert resp.status_code == 200, resp.text
    st = resp.json()
    wide = next(t["name"] for t in st["recipe"]["tables"] if t["name"] not in TABLES)
    resp = await client.put(f"{API}/imports/{st['id']}/recipe", json={"recipe": rename_wide(st["recipe"], wide, "日客流")})
    assert resp.status_code == 200, resp.text
    st = await trial(client, st["id"])
    assert st["trial"]["status"] == "passed", st["trial"]["problems"]
    resp = await commit(client, st)
    assert resp.status_code == 201, resp.text
    out = resp.json()
    return {"commit": out, "trial": st, "source_id": out["source"]["id"], "raw": raw, "file_name": fn}


async def tool_rows(client, source_name: str, sql: str) -> list[list[Any]]:
    resp = await client.post(f"/api/tools/db_query__{source_name}/run", json={"args": {"sql": sql}})
    assert resp.status_code == 200, resp.text
    out = resp.json()
    assert out["ok"], out
    return json.loads(out["result"])["rows"]


async def count(client, name: str, table: str, where: str = "") -> int:
    return int((await tool_rows(client, name, f'SELECT COUNT(*) FROM "{table}"{where}'))[0][0])


async def db_get(model, key):
    async with SessionLocal() as session:
        return await session.get(model, key)


async def source_row(client, source_id: str) -> dict[str, Any]:
    resp = await client.get(API)
    return next(s for s in resp.json() if s["id"] == source_id)


# ==========================================================================
# 按期累积：追加
# ==========================================================================


async def test_accumulate_august_then_september_materializes_and_publishes_the_union(client):
    """8 月首次导入（按期累积）+ 9 月上传新一期：试运行出累积计划、物化并集（U1–U3 通过），提交走 publish_snapshot，
    当前版本两期、库在 snapshots/ 下，三表 61 / 1037 / 183，按月 31 / 30（2.12）。"""
    name = unique()
    base = await first_import(client, name, seed=301)
    source_id = base["source_id"]
    first_snapshot = base["commit"]["snapshot_id"]
    snap = await db_get(SourceSnapshot, first_snapshot)
    assert snap.mode == "accumulate" and Path(snap.db_path).parent.name == "builds", "单期快照库就是那一期的构建库"
    card = await source_row(client, source_id)
    assert card["current_snapshot"]["mode"] == "accumulate" and card["current_snapshot"]["periods"] == 1

    raw, fn = flow_workbook(SEP, 30, seed=302)
    st = await reupload(client, source_id, raw, fn)
    t = st["trial"]
    assert t["status"] == "passed", t["problems"]
    plan = t["accumulate"]
    assert plan["action"] == "append" and plan["mode"] == "accumulate"
    assert [(p["start"], p["new"]) for p in plan["parts"]] == [("2026-08-01", False), ("2026-09-01", True)]
    assert plan["parts"][1]["rows"] == {"日客流": 30, "时段客流": 510, "时段客流_表内合计": 90}
    assert plan["parts"][1]["file_name"] == fn
    assert plan["union"]["rows"] == {"日客流": 61, "时段客流": 1037, "时段客流_表内合计": 183}
    assert [(c["id"], c["status"]) for c in t["union_checks"]] == [("U1", "passed"), ("U2", "passed"),
                                                                    ("U3", "passed")]
    assert t["recipe_compare"] is None, "配方没变：不给对照"
    assert t["prior_acceptances"] == []
    assert "union" not in t and "prev_import_id" not in t, "并集的内部信息不给界面"
    assert t["receipt"]["rows_excluded"] == [], "TrialOut.receipt 带排除的行（期 3 的试运行一律带这个键）"
    union_file = table_versions.union_trial_path(Path((await db_get(ImportStaging, st["id"])).trial_path))
    assert union_file.is_file(), "并集在试运行里物化，和主试运行库放在同一目录"
    assert "本表按统计期累积多期数据" in t["notes"]["日客流"]["comment"], t["notes"]["日客流"]

    resp = await commit(client, st)
    assert resp.status_code == 201, resp.text
    out = resp.json()
    assert out["parts"] == 2 and out["snapshot_reused"] is False and out["build_restored"] is False
    assert not union_file.exists(), "并集的试运行文件被发布消耗"
    snap = await db_get(SourceSnapshot, out["snapshot_id"])
    imps = [await db_get(TableImport, i) for i in snap.imports]
    assert [i.period_start for i in imps] == ["2026-08-01", "2026-09-01"]
    assert [i.status for i in imps] == ["active", "active"], "当前快照里的各期都是 active（2.1）"
    assert snap.mode == "accumulate" and Path(snap.db_path).parent.name == "snapshots"
    assert snap.manifest_artifact and snap.schema_cache["snapshot_manifest"] == snap.manifest_artifact
    for table, n in zip(TABLES, (61, 1037, 183)):
        assert await count(client, name, table) == n
    months = await tool_rows(client, name, 'SELECT substr("日期", 1, 7), COUNT(*) FROM "日客流" GROUP BY 1 ORDER BY 1')
    assert months == [["2026-08", 31], ["2026-09", 30]]
    # 快照清单与证据链（2.9）：union.db_sha256 等于快照库的哈希，各期 rowid 区间相接
    manifest = artifact_store.load(snap.manifest_artifact)
    assert manifest["format"] == "agentlab-snapshot-manifest/1" and manifest["action"] == "append"
    assert manifest["union"]["db_sha256"] == snap.db_sha256
    rows = [p["rows"]["日客流"] for p in manifest["parts"]]
    assert rows[0]["union"] == [1, 31] and rows[1]["union"] == [32, 61] and rows[1]["part"] == [1, 30]
    assert snap.schema_cache["import_manifests"] == [i.manifest_artifact for i in imps]
    # 新一期的导入清单：receipt 带排除的行，edits 是空列表
    new_manifest = artifact_store.load(imps[1].manifest_artifact)
    assert new_manifest["receipt"]["rows_excluded"] == [] and new_manifest["edits"] == []
    # 卡片：按期累积 · 2 期，统计期范围
    card = await source_row(client, source_id)
    cur = card["current_snapshot"]
    assert cur["mode"] == "accumulate" and cur["periods"] == 2
    assert (cur["period_start"], cur["period_end"]) == ("2026-08-01", "2026-09-30") and cur["activated_at"]


async def test_overlapping_period_is_rejected_and_the_current_version_stays(client):
    """统计期与已有的部分重叠：试运行拒收（period_overlap，消息写两段统计期），提交 409 trial_not_passed（A2）。"""
    name = unique()
    base = await first_import(client, name, seed=303)
    source_id = base["source_id"]
    raw, fn = flow_workbook(SEP, 30, seed=304)
    st = await reupload(client, source_id, raw, fn)
    assert (await commit(client, st)).status_code == 201
    current = (await source_row(client, source_id))["current_snapshot"]["id"]
    raw, fn = flow_workbook(AUG.replace(day=15), 31, seed=305)
    st = await reupload(client, source_id, raw, fn)
    t = st["trial"]
    assert t["status"] == "rejected"
    overlap = next(p for p in t["problems"] if p["code"] == "period_overlap")
    assert overlap["category"] == "structure"
    assert "2026-08-01 至 2026-08-31" in overlap["message"] and "2026-09-01 至 2026-09-30" in overlap["message"]
    assert t["accumulate"]["action"] == "rejected" and t["union_checks"] is None
    resp = await commit(client, st)
    assert resp.status_code == 409 and resp.json()["code"] == "trial_not_passed"
    assert (await source_row(client, source_id))["current_snapshot"]["id"] == current


async def test_replace_period_needs_its_confirmation_and_keeps_the_row_counts(client):
    """同一统计期再传（数字不同）：replace_period，必勾 period_replace:<起>~<止>；勾上后新快照是 [8 月, 9 月新]（A3）。"""
    name = unique()
    base = await first_import(client, name, seed=306)
    source_id = base["source_id"]
    raw, fn = flow_workbook(SEP, 30, seed=307)
    st = await reupload(client, source_id, raw, fn)
    assert (await commit(client, st)).status_code == 201
    old_sep = (await tool_rows(client, name, 'SELECT "全日客流" FROM "日客流" WHERE "日期" = \'2026-09-05\''))[0][0]
    raw2, fn2 = flow_workbook(SEP, 30, seed=308)
    st = await reupload(client, source_id, raw2, fn2)
    t = st["trial"]
    assert t["status"] == "passed", t["problems"]
    assert t["accumulate"]["action"] == "replace_period"
    assert t["accumulate"]["replaces"]["start"] == "2026-09-01"
    item = "period_replace:2026-09-01~2026-09-30"
    got = {i["id"]: i for i in t["confirm_items"]}
    assert item in got and got[item]["source"] == "accumulate" and got[item]["required"]
    resp = await commit(client, st, confirmations=[i for i in confirms(st) if i != item])
    assert resp.status_code == 422 and resp.json()["code"] == "confirm_required" and item in resp.json()["detail"]
    resp = await commit(client, st)
    assert resp.status_code == 201, resp.text
    snap = await db_get(SourceSnapshot, resp.json()["snapshot_id"])
    assert len(snap.imports) == 2
    for table, n in zip(TABLES, (61, 1037, 183)):
        assert await count(client, name, table) == n
    new_sep = (await tool_rows(client, name, 'SELECT "全日客流" FROM "日客流" WHERE "日期" = \'2026-09-05\''))[0][0]
    assert new_sep != old_sep


async def test_same_file_again_in_accumulate_mode_is_unchanged(client):
    """累积模式下同一个 9 月文件原样再传：替换该期且构建相同 → 不新建任何记录（same_as_import 的累积扩展）。"""
    name = unique()
    base = await first_import(client, name, seed=309)
    source_id = base["source_id"]
    raw, fn = flow_workbook(SEP, 30, seed=310)
    st = await reupload(client, source_id, raw, fn)
    first = await commit(client, st)
    assert first.status_code == 201
    st = await reupload(client, source_id, raw, fn)
    assert st["trial"]["accumulate"]["action"] == "replace_period"
    assert st["trial"]["same_as_import"]["id"] == first.json()["import_id"]
    resp = await commit(client, st)
    assert resp.status_code == 201, resp.text
    out = resp.json()
    assert out["unchanged"] is True and out["snapshot_id"] == first.json()["snapshot_id"] and out["parts"] == 2
    async with SessionLocal() as session:
        n = len(list((await session.execute(select(TableImport).where(TableImport.source_id == source_id))).scalars()))
    assert n == 2


# ==========================================================================
# 修复按钮：预览、应用、撤销、过期、PUT 覆盖
# ==========================================================================


def problem_with(st: dict[str, Any], code: str) -> dict[str, Any]:
    trial_problems = (st.get("trial") or {}).get("problems") or []
    return next(p for p in [*trial_problems, *st["recipe_problems"], *st["draft_problems"]] if p["code"] == code)


def fix_of(st: dict[str, Any], kind: str) -> dict[str, Any]:
    return next(f for f in st["fixes"] if f["kind"] == kind)


async def preview(client, sid: str, body: dict[str, Any]):
    return await client.post(f"{API}/imports/{sid}/edits/preview", json=body)


async def apply(client, sid: str, body: dict[str, Any], expected: str | None, **extra: Any):
    return await client.post(f"{API}/imports/{sid}/edits/apply", json={**body, "expected_sha256": expected, **extra})


async def fixed(client, st: dict[str, Any], kind: str, option: str, reason: str | None = None) -> dict[str, Any]:
    """按提议预览再应用，返回应用后的 StagingOut。"""
    fx = fix_of(st, kind)
    body = {"fix": {"id": fx["id"], "option": option, **({"reason": reason} if reason else {})}}
    resp = await preview(client, st["id"], body)
    assert resp.status_code == 200, resp.text
    pv = resp.json()
    assert pv["ok"], pv["problems"]
    resp = await apply(client, st["id"], body, pv["recipe_sha256_after"])
    assert resp.status_code == 200, resp.text
    return resp.json()


async def test_d04_remove_label_fix_preview_apply_trial_and_commit(client):
    """D04（日间少了 7-8）：拒收的问题带 fix_ids，提议去掉标签；预览不改暂存区，应用后试运行通过，确认清单的配方类只有
    fix:*（A6、A15），并集 61 / 1007 / 183（3.4），修改记录进导入清单。"""
    name = unique()
    base = await first_import(client, name, seed=311)
    source_id = base["source_id"]
    raw, fn = DRIFT_CASES["D04"].build(311)
    st = await reupload(client, source_id, raw, fn)
    assert st["trial"]["status"] == "rejected"
    missing = problem_with(st, "label_missing")
    fx = fix_of(st, "remove_label")
    assert missing["fix_ids"] == [fx["id"]] and fx["anchor"]["kind"] == "problem"
    assert all("fix_ids" in p for p in st["trial"]["problems"]), "每条问题都带 fix_ids（没有提议时是 []）"
    assert st["recipe_sha256"] and st["edits"] == [] and st["answers_dropped"] == []
    before_sha = st["recipe_sha256"]

    body = {"fix": {"id": fx["id"], "option": "remove"}, "seq": 3}
    resp = await preview(client, st["id"], body)
    assert resp.status_code == 200, resp.text
    pv = resp.json()
    assert pv["ok"] is True and pv["seq"] == 3 and pv["kind"] == "fix"
    assert pv["key"] == "remove_label:日间:7-8" and pv["recipe_sha256_before"] == before_sha
    assert pv["recipe_sha256_after"] and pv["recipe_sha256_after"] != before_sha
    assert pv["breaking"] == {} and isinstance(pv["breaking"], dict), "dimension 段去标签不改表结构"
    assert pv["recipe_problems"] == [] and pv["dry_run"]["problems"] == [] and pv["dry_run"]["marks"]
    assert pv["accumulate_change"] == "compatible"
    seg = next(s for s in pv["compare"]["segments"] if s["id"] == "日间")
    assert seg["labels"]["removed"] == ["7-8"]
    # 预览不改暂存区
    again = (await client.get(f"{API}/imports/{st['id']}")).json()
    assert again["recipe_sha256"] == before_sha and again["edits"] == []

    resp = await apply(client, st["id"], {"fix": body["fix"]}, pv["recipe_sha256_after"], signed_by="合成署名甲")
    assert resp.status_code == 200, resp.text
    st = resp.json()
    assert st["status"] == "drafting" and st["recipe_sha256"] == pv["recipe_sha256_after"]
    [edit] = st["edits"]
    assert edit["key"] == "remove_label:日间:7-8" and edit["undoable"] is True and edit["superseded"] is False
    assert edit["signed_by"] == "合成署名甲" and "before" not in edit and "base_sha256_after" not in edit
    st = await trial(client, st["id"])
    t = st["trial"]
    assert t["status"] == "passed", t["problems"]
    assert t["accumulate"]["action"] == "append"
    assert t["accumulate"]["union"]["rows"] == {"日客流": 61, "时段客流": 1007, "时段客流_表内合计": 183}
    assert t["recipe_compare"] is not None, "工作配方与现行配方不同：给新旧对照"
    items = {i["id"]: i for i in t["confirm_items"]}
    assert items["fix:remove_label:日间:7-8"]["source"] == "edit"
    recipe_class = [i for i, v in items.items() if v.get("source", "recipe") == "recipe"]
    assert recipe_class == [], recipe_class
    resp = await commit(client, st)
    assert resp.status_code == 201, resp.text
    imp = await db_get(TableImport, resp.json()["import_id"])
    manifest = artifact_store.load(imp.manifest_artifact)
    assert [e["key"] for e in manifest["edits"]] == ["remove_label:日间:7-8"]
    assert "before" not in manifest["edits"][0]
    # 3.1：修改记录（谁、何时、补丁）进导入清单——补丁和所选的选项都在，不只剩人话摘要
    assert manifest["edits"][0]["fix"] == {"id": fx["id"], "option": "remove", "reason": None}
    assert {"op": "replace", "path": manifest["edits"][0]["ops"][0]["path"],
            "value": manifest["edits"][0]["ops"][0]["value"]} == manifest["edits"][0]["ops"][0]
    assert manifest["edits"][0]["ops"][0]["path"].endswith("/labels/expect")
    assert "7-8" not in manifest["edits"][0]["ops"][0]["value"]
    assert await count(client, name, "时段客流") == 1007


async def test_undo_restores_the_start_and_put_supersedes_earlier_edits(client):
    """撤销恢复修改前的起点；撤销栈空了回 nothing_to_undo。PUT 之后之前的修改标「已被覆盖」，不能撤销、不出 fix:*（3.1）。"""
    name = unique()
    base = await first_import(client, name, seed=312)
    raw, fn = DRIFT_CASES["D04"].build(312)
    st = await reupload(client, base["source_id"], raw, fn)
    before_sha = st["recipe_sha256"]
    st = await fixed(client, st, "remove_label", "remove")
    resp = await client.post(f"{API}/imports/{st['id']}/edits/undo", json={})
    assert resp.status_code == 200, resp.text
    st = resp.json()
    assert st["recipe_sha256"] == before_sha and st["edits"] == []
    resp = await client.post(f"{API}/imports/{st['id']}/edits/undo", json={})
    assert resp.status_code == 409 and resp.json()["code"] == "nothing_to_undo"

    st = await trial(client, st["id"])
    st = await fixed(client, st, "remove_label", "remove")
    resp = await client.put(f"{API}/imports/{st['id']}/recipe", json={"recipe": st["recipe"]})
    assert resp.status_code == 200, resp.text
    st = resp.json()
    [edit] = st["edits"]
    assert edit["superseded"] is True and edit["undoable"] is False
    resp = await client.post(f"{API}/imports/{st['id']}/edits/undo", json={})
    assert resp.status_code == 409 and resp.json()["code"] == "nothing_to_undo"
    st = await trial(client, st["id"])
    assert st["trial"]["status"] == "passed", st["trial"]["problems"]
    assert not [i for i in confirms(st) if i.startswith("fix:")], "被覆盖的修改不再出确认项"


async def test_needs_reason_fix_refuses_a_changed_reason_and_records_rows_excluded(client):
    """D21（表下补录一行）：按行标签忽略要理由。没理由预览 422 reason_required；预览之后改了理由，应用 409 edit_stale
    （理由进配方哈希）；照原样应用后试运行通过，回执的 rows_excluded 有 ignored_rows（A6b）。提议不在了回 fix_stale。"""
    name = unique()
    base = await first_import(client, name, seed=313)
    raw, fn = DRIFT_CASES["D21"].build(313)
    st = await reupload(client, base["source_id"], raw, fn)
    assert st["trial"]["status"] == "rejected"
    fx = fix_of(st, "ignore_cells")
    assert fx["options"][0]["needs_reason"] is True
    option = fx["options"][0]["value"]
    resp = await preview(client, st["id"], {"fix": {"id": fx["id"], "option": option}})
    assert resp.status_code == 422 and resp.json()["code"] == "reason_required"
    body = {"fix": {"id": fx["id"], "option": option, "reason": "合成理由：补录行另有出处"}}
    pv = (await preview(client, st["id"], body)).json()
    assert pv["ok"], pv["problems"]
    changed = {"fix": {**body["fix"], "reason": "合成理由：改过的理由"}}
    resp = await apply(client, st["id"], changed, pv["recipe_sha256_after"])
    assert resp.status_code == 409 and resp.json()["code"] == "edit_stale"
    resp = await apply(client, st["id"], body, pv["recipe_sha256_after"])
    assert resp.status_code == 200, resp.text
    resp = await preview(client, st["id"], body)
    assert resp.status_code == 409 and resp.json()["code"] == "fix_stale", "应用之后问题没了，提议也不在了"
    st = await trial(client, st["id"])
    t = st["trial"]
    assert t["status"] == "passed", t["problems"]
    reasons = {r["reason"] for r in t["receipt"]["rows_excluded"]}
    assert "ignored_rows" in reasons, t["receipt"]["rows_excluded"]
    resp = await commit(client, st)
    assert resp.status_code == 201, resp.text
    imp = await db_get(TableImport, resp.json()["import_id"])
    manifest = artifact_store.load(imp.manifest_artifact)
    assert {r["reason"] for r in manifest["receipt"]["rows_excluded"]} == reasons


async def test_d13_add_label_then_edit_members_two_steps(client):
    """D13（多一个分区）：加标签之后静态校验报 fact_claim_mismatch，配方问题带 fix_ids，提议「更新关系成员」；两步
    修复后试运行通过，确认项有两条 fix:*，提交后 8 月的分区丙为空值（3.4）。"""
    name = unique()
    base = await first_import(client, name, seed=314)
    raw, fn = DRIFT_CASES["D13"].build(314)
    st = await reupload(client, base["source_id"], raw, fn)
    assert st["trial"]["status"] == "rejected"
    st = await fixed(client, st, "add_label", "add")
    mismatch = problem_with(st, "fact_claim_mismatch")
    fx = fix_of(st, "edit_members")
    assert mismatch["fix_ids"] == [fx["id"]] and fx["anchor"]["kind"] == "recipe_problem"
    st = await fixed(client, st, "edit_members", "update")
    assert st["recipe_problems"] == [], st["recipe_problems"]
    assert [e["kind"] for e in st["edits"]] == ["fix", "fix"] and [e["undoable"] for e in st["edits"]] == [False, True]
    st = await trial(client, st["id"])
    assert st["trial"]["status"] == "passed", st["trial"]["problems"]
    ids = confirms(st)
    assert len([i for i in ids if i.startswith("fix:")]) == 2, ids
    resp = await commit(client, st)
    assert resp.status_code == 201, resp.text
    assert await count(client, name, "日客流", ' WHERE "分区丙" IS NULL') == 31


async def test_d09_rename_title_keeps_the_confirmed_word(client):
    """D09（日间标题改字）：改分段标题、沿用「日间」（3.3 (b)）。传 base 之后静态校验放行，确认项没有 breaking:*。"""
    name = unique()
    base = await first_import(client, name, seed=315)
    raw, fn = DRIFT_CASES["D09"].build(315)
    st = await reupload(client, base["source_id"], raw, fn)
    fx = fix_of(st, "rename_title")
    keep = next(o["value"] for o in fx["options"] if o["value"].endswith("|keep"))
    st = await fixed(client, st, "rename_title", keep)
    assert st["recipe_problems"] == [], st["recipe_problems"]
    st = await trial(client, st["id"])
    assert st["trial"]["status"] == "passed", st["trial"]["problems"]
    assert not [i for i in confirms(st) if i.startswith("breaking:")], confirms(st)
    resp = await commit(client, st)
    assert resp.status_code == 201, resp.text
    kinds = await tool_rows(client, name, 'SELECT DISTINCT "时段类别" FROM "时段客流" ORDER BY 1')
    assert sorted(k[0] for k in kinds) == ["夜间", "日间"]


# ==========================================================================
# 按规则重新起草、改配方后回答的延续
# ==========================================================================


async def test_redraft_rules_then_adopt_lists_every_recipe_item_and_redraft_adopted(client):
    """redraft-rules 不改工作配方，记下对齐后的哈希；PUT（origin=rules_redraft）核对它，对不上 409 redraft_mismatch；
    采用之后确认清单按首次导入列出全部配方类项，另出 redraft_adopted（6.3，评审一-M6）。首次导入不提供。"""
    name = unique()
    base = await first_import(client, name, seed=316, mode="replace")
    # 与现行配方逐字相同的重新起草（同一版式的 9 月）：采用之后没有可确认的变化，只出 redraft_adopted 说明「相同」
    raw, fn = flow_workbook(SEP, 30, seed=317)
    same = await reupload(client, base["source_id"], raw, fn)
    out = (await client.post(f"{API}/imports/{same['id']}/redraft-rules", json={})).json()
    assert out["aligned_recipe"] is not None
    resp = await client.put(f"{API}/imports/{same['id']}/recipe",
                            json={"recipe": out["aligned_recipe"], "origin": "rules_redraft"})
    assert resp.json()["recipe_sha256"] == same["recipe_sha256"], "按来源对齐之后与现行配方逐字相同"
    resp = await client.delete(f"{API}/imports/{same['id']}")
    assert resp.status_code == 204
    # D04（日间少了 7-8）：重新起草出的分段标签与现行不同
    raw, fn = DRIFT_CASES["D04"].build(317)
    st = await reupload(client, base["source_id"], raw, fn)
    before = st["recipe_sha256"]
    resp = await client.post(f"{API}/imports/{st['id']}/redraft-rules", json={})
    assert resp.status_code == 200, resp.text
    out = resp.json()
    assert out["aligned_recipe"] is not None and isinstance(out["alignment"], list)
    assert out["compare"] is not None and "failures_for_model" not in out["draft"]
    assert (await client.get(f"{API}/imports/{st['id']}")).json()["recipe_sha256"] == before, "不改工作配方"
    tampered = json.loads(json.dumps(out["aligned_recipe"], ensure_ascii=False))
    tampered["tables"][0]["note"] = "合成说明"
    resp = await client.put(f"{API}/imports/{st['id']}/recipe", json={"recipe": tampered, "origin": "rules_redraft"})
    assert resp.status_code == 409 and resp.json()["code"] == "redraft_mismatch"
    resp = await client.put(f"{API}/imports/{st['id']}/recipe",
                            json={"recipe": out["aligned_recipe"], "origin": "rules_redraft"})
    assert resp.status_code == 200, resp.text
    st = resp.json()
    assert st["recipe_origin"] == "rules_redraft"
    st = await trial(client, st["id"])
    assert st["trial"]["status"] == "passed", st["trial"]["problems"]
    items = {i["id"]: i for i in st["trial"]["confirm_items"]}
    assert items["redraft_adopted"]["source"] == "edit" and "7-8" in items["redraft_adopted"]["label"]
    assert "mode" in items, "采用重新起草的结果时配方类项全部列出（只改了分段标签也一样）"
    resp = await commit(client, st)
    assert resp.status_code == 201, resp.text

    raw, fn = flow_workbook(AUG, 31, seed=318)
    first = (await stage(client, unique(), raw, fn)).json()
    resp = await client.post(f"{API}/imports/{first['id']}/redraft-rules", json={})
    assert resp.status_code == 409 and resp.json()["code"] == "redraft_not_offered"


async def test_put_keeps_answers_the_new_start_embodies_and_drops_the_rest(client):
    """L2：先回答占位符「·」存为空值，再 PUT 一份体现了这个回答的配方：回答仍在、answers_dropped 为空；PUT 一份去掉
    占位符的配方：answers_dropped 列出它（带问题文字，不只给 id）。"""
    raw, fn = flow_workbook(AUG, 31, seed=319)
    st = (await stage(client, unique(), raw, fn)).json()
    qid = "q_placeholder:·"
    assert qid in {q["id"] for q in st["questions"]}
    resp = await client.post(f"{API}/imports/{st['id']}/answers", json={"answers": {qid: {"value": "null"}}})
    st = resp.json()
    resp = await client.put(f"{API}/imports/{st['id']}/recipe", json={"recipe": st["recipe"]})
    st = resp.json()
    assert st["answers"] == {qid: {"value": "null", "reason": None}} and st["answers_dropped"] == []
    rec = json.loads(json.dumps(st["recipe"], ensure_ascii=False))
    for sheet in rec["sheets"]:
        for block in sheet["blocks"]:
            if isinstance(block.get("values"), dict):
                block["values"]["placeholders"] = []
    resp = await client.put(f"{API}/imports/{st['id']}/recipe", json={"recipe": rec})
    assert resp.status_code == 200, resp.text
    st = resp.json()
    dropped = {d["id"]: d for d in st["answers_dropped"]}
    assert qid in dropped and dropped[qid]["text"] and dropped[qid]["reason"] in ("changed", "gone")
    assert qid not in st["answers"]


# ==========================================================================
# 框选转锚点（4.4、4.6）
# ==========================================================================


async def test_c01_selection_preview_matches_the_rule_draft_and_applies(client):
    """lab.c01 暂存后在「月报!C5:F8」框选为列表：预览 ok，锚点是表头，重放与框一致，换算后的配方与规则草稿逐字相同
    （哈希不变）；框到第 7 行时说不清下边界，照样 200，ok=false、problems 写原因；应用后修改记录是 selection。"""
    from tests.fixtures.xlsx import lab

    raw, fn = lab.c01("literal")
    st = (await stage(client, unique(), raw, fn)).json()
    sel = {"sheet": "月报", "ref": "C5:F8", "as": "list", "options": {"header_rows": 1, "bottom": "box"}}
    resp = await preview(client, st["id"], {"selection": sel, "seq": 1})
    assert resp.status_code == 200, resp.text
    pv = resp.json()
    assert pv["ok"] is True and pv["kind"] == "selection" and pv["seq"] == 1
    assert [a["text"] for a in pv["anchors"] if a["kind"] == "header"] == ["地区", "产品", "销量", "金额"]
    assert pv["replay"]["match"] is True, pv["replay"]
    assert pv["replay"]["expected"] == {"header": "C5:F5", "data": "C6:F8", "total": "C9:F9"}
    assert pv["recipe_sha256_after"] == st["recipe_sha256"], "框选得到的规则与起草器完全一样"
    assert pv["breaking"] == {} and pv["expected"] == pv["replay"]["expected"]
    bad = (await preview(client, st["id"], {"selection": {**sel, "ref": "C5:F7"}})).json()
    assert bad["ok"] is False and [p["code"] for p in bad["problems"]] == ["selection_bottom_not_anchorable"]
    assert bad["breaking"] == {} and bad["recipe_problems"] == [] and bad["dry_run"] is None
    resp = await apply(client, st["id"], {"selection": {**sel, "ref": "C5:F7"}}, None)
    assert resp.status_code == 422 and resp.json()["code"] == "edit_not_applicable"
    resp = await apply(client, st["id"], {"selection": sel}, pv["recipe_sha256_after"])
    assert resp.status_code == 200, resp.text
    st = resp.json()
    [edit] = st["edits"]
    assert edit["kind"] == "selection" and edit["key"].startswith("list:") and edit["undoable"] is True
    # 修改记录记下选区（键是 as，Selection.to_json）；补丁只进导入清单，StagingOut 不给
    assert edit["selection"] == sel and "ops" not in edit
    row = await db_get(ImportStaging, st["id"])
    assert row.edits[0]["ops"] and row.edits[0]["selection"] == sel


# ==========================================================================
# 评审后补：撤销不丢回答、预览的 breaking、采用重新起草之后的来源、重新起草没得到配方、带接受的重传、退役链
# ==========================================================================


async def test_undo_keeps_the_answers_given_before_the_edit(client):
    """先回答（导入模式选每期替换）、再用修复按钮、再撤销：撤销只撤这一次修改，修改前的回答照旧、不报「需要重新
    回答」，工作配方回到修改前（3.1）。原来撤销恢复的是还没叠加回答的起点，有效果的回答一律判 changed 丢掉，导入
    模式悄悄退回草稿默认的按期累积（评审意见）。"""
    raw, fn = DRIFT_CASES["D21"].build(331)
    resp = await stage(client, unique(), raw, fn)
    assert resp.status_code == 201, resp.text
    st = resp.json()
    resp = await client.post(f"{API}/imports/{st['id']}/answers",
                             json={"answers": answers_for(st["questions"], "replace")})
    assert resp.status_code == 200, resp.text
    st = resp.json()
    assert st["recipe"].get("mode", "replace") == "replace"
    before_sha, before_answers = st["recipe_sha256"], st["answers"]
    assert before_answers["q_mode"] == {"value": "replace", "reason": None}
    fx = fix_of(st, "ignore_cells")
    opt = fx["options"][0]
    body = {"fix": {"id": fx["id"], "option": opt["value"],
                    **({"reason": "合成理由：表下另有补录行"} if opt["needs_reason"] else {})}}
    pv = (await preview(client, st["id"], body)).json()
    assert pv["ok"], pv["problems"]
    resp = await apply(client, st["id"], body, pv["recipe_sha256_after"])
    assert resp.status_code == 200, resp.text
    st = resp.json()
    assert st["answers"] == before_answers and st["answers_dropped"] == []
    resp = await client.post(f"{API}/imports/{st['id']}/edits/undo", json={})
    assert resp.status_code == 200, resp.text
    st = resp.json()
    assert st["edits"] == [] and st["recipe_sha256"] == before_sha
    assert st["answers"] == before_answers and st["answers_dropped"] == []
    assert st["recipe"].get("mode", "replace") == "replace"


async def test_preview_breaking_lists_the_table_changes_of_a_breaking_option(client):
    """预览的 breaking 是 table_changes 算出的「表名 → 人话列表」（WP-0 交接）：D09 改分段标题时另选新词（pick，常量
    取值变了），「时段客流」列出取值的变化；沿用原词（keep）为空字典。只断言「是字典」抓不住漏算（评审意见）。"""
    base = await first_import(client, unique(), seed=336)
    raw, fn = DRIFT_CASES["D09"].build(336)
    st = await reupload(client, base["source_id"], raw, fn)
    fx = fix_of(st, "rename_title")
    pick = next(o for o in fx["options"] if "|pick:" in o["value"])
    keep = next(o for o in fx["options"] if o["value"].endswith("|keep"))
    assert pick["breaking"] is True and keep["breaking"] is False
    resp = await preview(client, st["id"], {"fix": {"id": fx["id"], "option": pick["value"]}})
    assert resp.status_code == 200, resp.text
    pv = resp.json()
    assert pv["ok"], pv["problems"]
    word = pick["value"].split("|pick:", 1)[1]
    assert list(pv["breaking"]) == ["时段客流"], pv["breaking"]
    assert any("日间" in line and word in line for line in pv["breaking"]["时段客流"]), pv["breaking"]
    pv = (await preview(client, st["id"], {"fix": {"id": fx["id"], "option": keep["value"]}})).json()
    assert pv["ok"] and pv["breaking"] == {}


async def test_an_adopted_redraft_is_stored_as_rules_and_next_month_does_not_ask_again(client):
    """采用按规则重新起草的配方并提交：配方记录、导入清单记 rules（只取 rules / ai / manual / mixed，界面按
    RECIPE_ORIGIN_LABEL 显示）。下个月照常上传同一版式：暂存区的来源是 rules，不再出 redraft_adopted，也不再列出全部
    配方类项（6.2、6.3：只有在本暂存区采用时才出）。原来 rules_redraft 原样入库，之后每个月都要再勾一次（评审意见）。"""
    name = unique()
    base = await first_import(client, name, seed=337, mode="replace")
    source_id = base["source_id"]
    raw, fn = DRIFT_CASES["D04"].build(337)
    st = await reupload(client, source_id, raw, fn)
    out = (await client.post(f"{API}/imports/{st['id']}/redraft-rules", json={})).json()
    assert out["aligned_recipe"] is not None
    resp = await client.put(f"{API}/imports/{st['id']}/recipe",
                            json={"recipe": out["aligned_recipe"], "origin": "rules_redraft"})
    assert resp.status_code == 200, resp.text
    assert resp.json()["recipe_origin"] == "rules_redraft"
    st = await trial(client, st["id"])
    assert st["trial"]["status"] == "passed", st["trial"]["problems"]
    assert "redraft_adopted" in confirms(st), "本暂存区采用了重新起草：照旧出 redraft_adopted"
    resp = await commit(client, st)
    assert resp.status_code == 201, resp.text
    card = await source_row(client, source_id)
    assert card["current_recipe"]["origin"] == "rules"
    imp = await db_get(TableImport, resp.json()["import_id"])
    assert artifact_store.load(imp.manifest_artifact)["recipe"]["origin"] == "rules"

    raw, fn = DRIFT_CASES["D04"].build(338)
    nxt = await reupload(client, source_id, raw, fn)
    assert nxt["recipe_origin"] == "rules"
    t = nxt["trial"]
    assert t["status"] == "passed", t["problems"]
    ids = confirms(nxt)
    assert "redraft_adopted" not in ids, ids
    assert [i["id"] for i in t["confirm_items"] if i.get("source", "recipe") == "recipe"] == [], ids


async def test_a_stored_rules_redraft_origin_is_read_as_rules_for_the_next_staging(client):
    """防住已经入库的 rules_redraft（修复之前提交的配方记录）：上传新一期、修改配方给暂存区赋来源时同样映射成 rules，
    不出 redraft_adopted。"""
    from sqlalchemy import update

    from app.db.models import TableRecipe

    base = await first_import(client, unique(), seed=344, mode="replace")
    async with SessionLocal() as session:
        await session.execute(update(TableRecipe).where(TableRecipe.id == base["commit"]["recipe_id"])
                              .values(origin="rules_redraft"))
        await session.commit()
    raw, fn = flow_workbook(SEP, 30, seed=345)
    st = await reupload(client, base["source_id"], raw, fn)
    assert st["recipe_origin"] == "rules" and st["trial"]["status"] == "passed", st["trial"]["problems"]
    assert "redraft_adopted" not in confirms(st)
    resp = await client.delete(f"{API}/imports/{st['id']}")
    assert resp.status_code in (200, 204), resp.text
    resp = await client.post(f"{API}/{base['source_id']}/redraft", json={})
    assert resp.status_code == 201, resp.text
    assert resp.json()["recipe_origin"] == "rules"


async def test_redraft_rules_without_a_recipe_leaves_alignment_empty(client):
    """规则起草没得到配方时，失败原因只在 draft.failures；alignment 是「名字对齐」的报告，为空（界面把它放在那个标题下，
    抄一份过去同一句话就出现两次，check-manage 伪造的也是 []，评审意见）。"""
    import io

    import openpyxl

    base = await first_import(client, unique(), seed=339)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "客流汇总"
    ws["A1"] = "说明文字甲"
    ws["A3"] = "说明文字乙"
    buf = io.BytesIO()
    wb.save(buf)
    st = await reupload(client, base["source_id"], buf.getvalue(), "说明.xlsx")
    resp = await client.post(f"{API}/imports/{st['id']}/redraft-rules", json={})
    assert resp.status_code == 200, resp.text
    out = resp.json()
    assert out["draft"]["recipe"] is None and out["draft"]["failures"]
    assert out["aligned_recipe"] is None and out["compare"] is None and out["alignment"] == []


async def test_same_file_again_with_an_acceptance_is_a_new_import(client):
    """带接受的重传照常新建导入记录，接受要重写（第 5 节：unchanged 只适用于没有接受的情况）。回执旁给出上次（没作废的）
    接受理由作参考，不预填。"""
    name = unique()
    base = await first_import(client, name, seed=340)
    source_id = base["source_id"]
    raw, fn = DRIFT_CASES["D26"].build(340)
    st = await reupload(client, source_id, raw, fn)
    t = st["trial"]
    assert t["status"] == "needs_decision", t["problems"]
    first = await commit(client, st, acceptances=[{"check_id": c, "reason": "合成理由甲"} for c in t["acceptable"]])
    assert first.status_code == 201, first.text
    st = await reupload(client, source_id, raw, fn)
    t = st["trial"]
    assert t["accumulate"]["action"] == "replace_period" and t["same_as_import"]["id"] == first.json()["import_id"]
    assert [a["reason"] for a in t["prior_acceptances"]] == ["合成理由甲"] * len(t["acceptable"])
    resp = await commit(client, st, acceptances=[{"check_id": c, "reason": "合成理由乙"} for c in t["acceptable"]])
    assert resp.status_code == 201, resp.text
    out = resp.json()
    assert out["unchanged"] is False and out["import_id"] != first.json()["import_id"] and out["parts"] == 2
    imp = await db_get(TableImport, out["import_id"])
    assert [o["reason"] for o in imp.overrides or []] + [w["reason"] for w in imp.waivers or []] \
        == ["合成理由乙"] * len(t["acceptable"])
    async with SessionLocal() as session:
        n = len(list((await session.execute(select(TableImport).where(TableImport.source_id == source_id))).scalars()))
    assert n == 3


def _drop_total_table(recipe: dict[str, Any]) -> dict[str, Any]:
    """合计行改为「不另存」（keep_as 置空），「时段客流_表内合计」从目标配方里去掉。"""
    rec = json.loads(json.dumps(recipe, ensure_ascii=False))
    n = 0
    for sheet in rec["sheets"]:
        for block in sheet["blocks"]:
            for seg in block.get("segments", []):
                if seg.get("keep_as"):
                    seg["keep_as"] = None
                    n += 1
    rec["tables"] = [t for t in rec["tables"] if t["name"] != "时段客流_表内合计"]
    assert n
    return rec


async def test_a_table_retired_once_is_not_retired_again_next_month(client):
    """退役链（2.5）：9 月去掉合计表（必勾 retire:时段客流_表内合计，早期各期保留原值）；10 月同一配方，这张表已是
    「早已退役」（快照清单 columns 的 status），不再出 retire:*。试运行要把快照清单的退役记录传给分档，漏传就每个月
    都再报一次退役（评审意见：变异「不传 retired_status」没有测试抓到）。"""
    name = unique()
    base = await first_import(client, name, seed=341)
    source_id = base["source_id"]
    raw, fn = flow_workbook(SEP, 30, seed=342)
    st = await reupload(client, source_id, raw, fn)
    resp = await client.put(f"{API}/imports/{st['id']}/recipe", json={"recipe": _drop_total_table(st["recipe"])})
    assert resp.status_code == 200, resp.text
    assert resp.json()["recipe_problems"] == []
    st = await trial(client, st["id"])
    assert st["trial"]["status"] == "passed", st["trial"]["problems"]
    plan = st["trial"]["accumulate"]
    assert plan["change"] == "retire" and plan["retired_new"] == [{"table": "时段客流_表内合计", "column": None}]
    assert "retire:时段客流_表内合计" in confirms(st)
    resp = await commit(client, st)
    assert resp.status_code == 201, resp.text
    snap = await db_get(SourceSnapshot, resp.json()["snapshot_id"])
    columns = artifact_store.load(snap.manifest_artifact)["columns"]["时段客流_表内合计"]
    assert columns and all(c["status"] == "retired" and c["since"] == "2026-09-01~2026-09-30" for c in columns.values())
    assert await count(client, name, "时段客流_表内合计") == 93

    raw, fn = flow_workbook(OCT, 31, seed=343)
    st = await reupload(client, source_id, raw, fn)
    t = st["trial"]
    assert t["status"] == "passed", t["problems"]
    assert t["accumulate"]["retired_new"] == []
    assert t["accumulate"]["retired_existing"] == [{"table": "时段客流_表内合计", "column": None}]
    assert not [i for i in confirms(st) if i.startswith("retire")], confirms(st)
    resp = await commit(client, st)
    assert resp.status_code == 201, resp.text
    assert await count(client, name, "时段客流_表内合计") == 93


async def test_appending_after_removing_the_only_period_with_a_column_asks_no_phantom_retire(client):
    """试运行路径的幽灵退役项（WP-8 修补，WP-5 交付说明「需要你处理」第 2 条）：8 月 R0；9 月 D13 修复加「分区丙」（R1）；
    移除 9 月之后当前版本按 R1 物化成单期并集，「分区丙」是一列全为空值的目标列；10 月修复去掉分区丙（R2）追加。
    结果各期（8 月、10 月）和当前版本的任何一期都没有分区丙的数据，不该出 retire:日客流.分区丙，并集里也不再有这一列。
    改回「把未筛选的当前表结构传给累积计划」，这里会多出一条任何一期都没有数据的必勾退役项。"""
    from tests.test_source_versions_api import _fix_until_passed

    name = unique()
    base = await first_import(client, name, seed=351)
    sid = base["source_id"]
    raw, fn = DRIFT_CASES["D13"].build(351)
    st = await _fix_until_passed(client, await reupload(client, sid, raw, fn))
    assert st["trial"]["status"] == "passed", st["trial"]["problems"]
    resp = await commit(client, st)
    assert resp.status_code == 201, resp.text
    sep, s2 = resp.json()["import_id"], resp.json()["snapshot_id"]
    resp = await client.post(f"{API}/{sid}/imports/{sep}/remove",
                             json={"confirm": True, "expected_current_snapshot_id": s2, "reason": "合成理由"})
    assert resp.status_code == 200, resp.text
    s3 = await db_get(SourceSnapshot, resp.json()["snapshot_id"])
    assert Path(s3.db_path).parent.name == "snapshots", "8 月的配方不是目标配方：按 R1 物化成单期并集"
    assert "分区丙" in [c["name"] for c in s3.schema_cache["tables"]["日客流"]["columns"]], "目标列照旧在，全为空值"

    raw, fn = flow_workbook(OCT, 31, seed=352)
    st = await _fix_until_passed(client, await reupload(client, sid, raw, fn))
    t = st["trial"]
    assert t["status"] == "passed", (t["problems"], st["recipe_problems"])
    assert t["accumulate"]["action"] == "append"
    assert {"table": "日客流", "column": "分区丙"} not in t["accumulate"]["retired_new"], t["accumulate"]
    assert {"table": "日客流", "column": "分区丙"} not in t["accumulate"]["retired_existing"], t["accumulate"]
    assert not [i for i in confirms(st) if i.startswith("retire")], confirms(st)
    resp = await commit(client, st)
    assert resp.status_code == 201, resp.text
    snap = await db_get(SourceSnapshot, resp.json()["snapshot_id"])
    assert "分区丙" not in [c["name"] for c in snap.schema_cache["tables"]["日客流"]["columns"]]
    assert "分区丙" not in artifact_store.load(snap.manifest_artifact)["columns"].get("日客流", {})
    assert await count(client, name, "日客流") == 62



async def test_a_target_column_without_data_is_not_reported_as_added_again(client):
    """WP-8 修补的另一半：筛当前表结构时，目标配方里有的列照旧留着。8 月 R0；9 月 D13 修复加「分区丙」（R1）；移除 9 月之后
    当前版本按 R1 物化，分区丙是全为空值的目标列；10 月的文件又有分区丙、按 R1 直接重放：当前版本本来就有这一列，计划里
    不该再报「新增列」（2.5 第 2 步按当前版本实际的表结构比）。只按各期配方筛、不留目标配方的列，这里会多出一条新增。"""
    from tests.test_source_versions_api import _fix_until_passed

    name = unique()
    base = await first_import(client, name, seed=361)
    sid = base["source_id"]
    raw, fn = DRIFT_CASES["D13"].build(361)
    st = await _fix_until_passed(client, await reupload(client, sid, raw, fn))
    resp = await commit(client, st)
    assert resp.status_code == 201, resp.text
    sep, s2 = resp.json()["import_id"], resp.json()["snapshot_id"]
    resp = await client.post(f"{API}/{sid}/imports/{sep}/remove",
                             json={"confirm": True, "expected_current_snapshot_id": s2, "reason": "合成理由"})
    assert resp.status_code == 200, resp.text
    raw, fn = flow_workbook(OCT, 31, seed=362, variant="D13")
    st = await reupload(client, sid, raw, fn)
    t = st["trial"]
    assert t["status"] == "passed", (t["problems"], st["recipe_problems"])
    assert t["accumulate"]["action"] == "append"
    assert t["accumulate"]["added"] == [] and t["accumulate"]["retired_new"] == [], t["accumulate"]
    assert not [d for d in t["diff"] or [] if d["kind"] == "columns_added"], t["diff"]


async def test_replacing_the_only_period_with_a_column_asks_breaking_not_a_phantom_retire(client):
    """替换该期的幽灵退役项（WP-8 评审意见 2）：8 月 R0；9 月 D13 修复加「分区丙」（R1）后提交；再上传不含分区丙的 9 月
    （修复去掉分区丙后为 R2），替换该期。结果各期（8 月 R0、新 9 月 R2）都没有分区丙：不出 retire:*（文案写「早期各期
    保留原值」，可哪一期都没有这一列），改出必勾的 breaking:日客流（列从现行配方里消失，仍要人确认，不会变成静默）；
    提交后并集里没有分区丙，也没有全为空值的列。改回「把要被替换的那一期也算进结果各期」，这里会出 retire:日客流.分区丙，
    并集多出一列全空的分区丙。"""
    import sqlite3
    from contextlib import closing

    from tests.test_source_versions_api import _fix_until_passed

    name = unique()
    base = await first_import(client, name, seed=361)
    sid = base["source_id"]
    raw, fn = DRIFT_CASES["D13"].build(361)
    st = await _fix_until_passed(client, await reupload(client, sid, raw, fn))
    assert st["trial"]["status"] == "passed", st["trial"]["problems"]
    resp = await commit(client, st)
    assert resp.status_code == 201, resp.text
    sep_r1 = resp.json()["import_id"]

    raw, fn = flow_workbook(SEP, 30, seed=362)
    st = await _fix_until_passed(client, await reupload(client, sid, raw, fn))
    t = st["trial"]
    assert t["status"] == "passed", (t["problems"], st["recipe_problems"])
    assert t["accumulate"]["action"] == "replace_period" and t["accumulate"]["replaces"]["import_id"] == sep_r1
    assert {"table": "日客流", "column": "分区丙"} not in t["accumulate"]["retired_new"], t["accumulate"]
    assert {"table": "日客流", "column": "分区丙"} not in t["accumulate"]["retired_existing"], t["accumulate"]
    ids = confirms(st)
    assert not [i for i in ids if i.startswith("retire")], ids
    assert "breaking:日客流" in ids and "period_replace:2026-09-01~2026-09-30" in ids, ids
    resp = await commit(client, st, confirmations=[i for i in ids if i != "breaking:日客流"])
    assert resp.status_code == 422 and resp.json()["code"] == "confirm_required", resp.text
    assert "breaking:日客流" in resp.json()["detail"]
    resp = await commit(client, st)
    assert resp.status_code == 201, resp.text
    snap = await db_get(SourceSnapshot, resp.json()["snapshot_id"])
    assert "分区丙" not in [c["name"] for c in snap.schema_cache["tables"]["日客流"]["columns"]]
    assert "分区丙" not in artifact_store.load(snap.manifest_artifact)["columns"].get("日客流", {})
    empty = []
    with closing(sqlite3.connect(f"file:{snap.db_path}?mode=ro", uri=True)) as conn:
        for table, info in snap.schema_cache["tables"].items():
            for col in info["columns"]:
                n = conn.execute(f'SELECT COUNT("{col["name"]}") FROM "{table}"').fetchone()[0]
                if n == 0:
                    empty.append(f"{table}.{col['name']}")
    assert empty == [], f"并集里不该有全为空值的列：{empty}"
    assert await count(client, name, "日客流") == 61
