"""期 4「证据下钻」集成验收（P4-SPEC 6.2 的 A1–A26、6.5 的性能，WP-E）。

真管线、真运行：从暂存合成工作簿开始，经接口回答问题、试运行、提交（导入），再用真 runner 跑「工具节点
db_query__<源> × N → 报告节点 → 输出」的图（写作模型是剧本化的 MockChatModel，不调用任何真实模型），运行跑完自动
封存；之后一律经 AsyncClient(transport=ASGITransport(app=app)) 请求片段接口和推断来源接口
`GET /api/runs/{id}/evidence/segments/{seg}/provenance`，不直接调内部函数拼结果。只有接口不提供的观察点
（快照记录、导入记录、内容寻址的工件、数据文件、原件字节）才直接读库、看文件；造「数据文件被改过」「登记哈希被改过」
「封存范围内的事件被改过」这类状态时才直接改库、改文件（6.1 的篡改）。

每份响应都按契约（provenance_types）逐层核对键集合、还原成 ProvenanceOut，再断言 `contract_problems(out) == []`
（WP-0 交接：漏了映射的细分、回查参数对不上都靠这一条暴露）。

**夹具组织**：不改状态的用例共用一个模块级的「世界」：十二个合成的源（每期替换、按期累积、D26 写理由接受（替换与
累积各一）、人工录入统计期、设遮罩（宽表、列表各一）、列表、简单导入、手工源、D28 区域外文字及其隐藏行变体）和四次运行（主运行一次
引用全部用例的格子；期 4 之前形状的文档一次（A21 用）；停在审批上的一次；留给 A23 篡改事件的一次）。世界在单独的
事件循环里建好，建完关掉数据层和数据库的连接，用例各自的事件循环里重新连。改数据文件、改登记哈希、清除原件、作废
接受、改坏清单的用例（A5、A5b、A9、A14、A25、A26），以及运行期间要装桩的 A6，各建自己的源和运行，互不影响。

夹具全部合成、假名（tests/fixtures/xlsx：分区甲、分区乙……数字随机）。断言照规格写，不放宽、不标 xfail。
"""
from __future__ import annotations

import asyncio
import hashlib
import io
import json
import os
import shutil
import sqlite3
import time
import uuid
from contextlib import closing
from dataclasses import dataclass, field, fields, replace
from pathlib import Path
from typing import Any

import pytest
from httpx import ASGITransport, AsyncClient
from langchain_core.messages import AIMessage
from openpyxl import load_workbook
from sqlalchemy import select

from app.core import artifact_store
from app.core.config import settings
from app.data.engine import engines
from app.data.provenance_types import (
    REASON_TEXT,
    AcceptanceView,
    Alert,
    CanonicalView,
    CellRef,
    CellSource,
    CompileFacts,
    FromCell,
    PartView,
    PeriodView,
    ProvenanceOut,
    Reason,
    Recheck,
    RelatedCheck,
    ReportRef,
    StateNote,
    TableRef,
    VersionView,
    YearSource,
    contract_problems,
    refuse,
)
from app.db.base import SessionLocal
from app.db.base import engine as db_engine
from app.db.models import DataSource, Run, RunEvent, SourceSnapshot, TableImport
from app.engine.evidence import iter_segments, iter_units, verify_doc
from app.engine.runner import run_manager
from app.main import app
from tests.fixtures.xlsx import excel_saved, lab, save
from tests.fixtures.xlsx.drift import AUG, SEP
from tests.fixtures.xlsx.flow import DAY_TITLE, SHEET, flow_workbook

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
API = "/api/datasources"
TABLES = ("日客流", "时段客流", "时段客流_表内合计")
REASON_D26 = "合成理由：分区合计口径调整"
SIGNER = "录入员甲"
#: 句子里的标签：不含数字（报告里的裸数字会被标成无出处），一句一个，按片段的引用认片段、不按文字
LABELS = tuple(a + b for a, b in zip("甲乙丙丁戊己庚辛壬癸" * 6, "子丑寅卯辰巳午未申酉戌亥" * 5))
#: D28 夹具在区域外第 32 行写的注释（9 月那一期）
D28_NOTE = "注：9月15日闸机故障，当日客流为估算值"
OUTSIDE_HEAD = "区域外文字（原表区域外的文字原文，未经系统核对；其中的数字不是查询结果，不能用来支持或否定报告中的数字）："
ACCEPT_HEAD = "已接受（用户填写的理由，未经系统核对）："
#: 6.5：超过目标的这个倍数算不通过
FAIL_FACTOR = 1.5


def wide_sql(day: str, column: str = "全日客流") -> str:
    return f'SELECT "日期", "{column}" FROM "日客流" WHERE "日期" = \'{day}\''


def long_sql(day: str, hour: str) -> str:
    return f'SELECT "日期", "时段", "客流" FROM "时段客流" WHERE "日期" = \'{day}\' AND "时段" = \'{hour}\''


#: 实测 B7 的两种写法和 WP-A 补充的 U+FEFF 写法：藏在后面的 UNION ALL 把分区甲的值装进名为「全日客流」的结果列。
#: 第一段一行都选不出来（只有 8 月的数），结果的 r0 实际来自分区甲；夹具让 2026-08-03 的全日客流恰好等于分区甲，
#: 所以只比值的回查分不出来，只能靠 R7 认出集合运算
A24_SQL = {
    "escape": ('SELECT "日期", "全日客流" FROM "日客流" WHERE "日期" LIKE \'2026-09%\' ESCAPE \'\\\' '
               'UNION ALL SELECT "日期", "分区甲" FROM "日客流" WHERE "日期" = \'2026-08-03\''),
    "bracket": ('SELECT "日期", "全日客流" FROM "日客流" AS [t\'] WHERE "日期" = \'2026-09-01\' '
                'UNION ALL SELECT "日期", "分区甲" FROM "日客流" AS [u\'] WHERE "日期" = \'2026-08-03\''),
    "bom": ('SELECT "日期", "全日客流" FROM "日客流" WHERE "日期" = \'2026-09-01\' \ufeffUNION ALL '
            '\ufeffSELECT "日期", "分区甲" \ufeffFROM "日客流" WHERE "日期" = \'2026-08-03\''),
}


# ==========================================================================
# 小工具（经 HTTP 导入）
# ==========================================================================


def unique(prefix: str = "p4e") -> str:
    return f"{prefix}_{uuid.uuid4().hex[:10]}"


def answers_for(questions: list[dict[str, Any]], mode: str) -> dict[str, dict[str, Any]]:
    """系统发现一律「登记」，q_mode 按 mode，其余能存为空值的选「存为空值」（期 2、3 验收的写法）。"""
    out = {}
    for q in questions:
        values = [o["value"] for o in q["options"]]
        if q["id"].startswith("q_relation"):
            out[q["id"]] = {"value": "register"}
        elif q["id"] == "q_mode":
            out[q["id"]] = {"value": mode}
        elif "null" in values:
            out[q["id"]] = {"value": "null"}
        else:
            out[q["id"]] = {"value": values[0]}
    return out


def rename_wide(rec: dict[str, Any], old: str, new: str) -> dict[str, Any]:
    """把规则起草的宽表名改成 new：表名、写入它的分段（id 一起改）、关系里引用的表名（期 2 验收 9.7 第 1 步）。"""
    rec = json.loads(json.dumps(rec, ensure_ascii=False))
    for sheet in rec["sheets"]:
        for block in sheet["blocks"]:
            for seg in block.get("segments", []):
                if seg.get("table") == old:
                    seg["table"] = new
                if seg.get("id") == old:
                    seg["id"] = new
    for t in rec["tables"]:
        if t["name"] == old:
            t["name"] = new
    for r in rec["relations"]:
        if r.get("table") == old:
            r["table"] = new
        for side in ("a", "b"):
            if isinstance(r.get(side), dict) and r[side].get("table") == old:
                r[side]["table"] = new
    return rec


def confirm_ids(st: dict[str, Any]) -> list[str]:
    return [i["id"] for i in (st.get("trial") or {}).get("confirm_items") or []]


async def commit(c: AsyncClient, st: dict[str, Any], acceptances: list[dict[str, str]] | None = None) -> dict[str, Any]:
    body: dict[str, Any] = {"trial_id": st["trial"]["trial_id"], "confirmations": confirm_ids(st)}
    if acceptances is not None:
        body["acceptances"] = acceptances
    r = await c.post(f"{API}/imports/{st['id']}/commit", json=body)
    assert r.status_code == 201, r.text
    return r.json()


@dataclass
class Source:
    """一个合成的源：名字、id、当前版本、各期原件（字节、文件名）、各期导入记录 id（按导入先后）。"""

    name: str
    id: str
    snapshot: str
    raws: list[tuple[bytes, str]]
    imports: list[str]
    #: 每一次提交后的版本 id（按导入先后）；替换模式下当前版本只含最后一期
    snapshots: list[str] = field(default_factory=list)


async def first_import(c: AsyncClient, raw: bytes, fn: str, *, mode: str = "replace", name: str | None = None,
                       rename: bool = True) -> Source:
    """暂存 → 回答问题（q_mode=mode）→ 宽表改名「日客流」（flow 夹具）→ 试运行 → 勾全部确认项提交。"""
    name = name or unique()
    r = await c.post(f"{API}/imports/stage", files={"file": (fn, raw, XLSX)}, data={"name": name})
    assert r.status_code == 201, r.text
    st = r.json()
    if st.get("questions"):
        r = await c.post(f"{API}/imports/{st['id']}/answers", json={"answers": answers_for(st["questions"], mode)})
        assert r.status_code == 200, r.text
        st = r.json()
    if rename:
        wide = next(t["name"] for t in st["recipe"]["tables"] if t["name"] not in TABLES)
        r = await c.put(f"{API}/imports/{st['id']}/recipe", json={"recipe": rename_wide(st["recipe"], wide, "日客流")})
        assert r.status_code == 200, r.text
        st = r.json()
    r = await c.post(f"{API}/imports/{st['id']}/trial", json={})
    assert r.status_code == 200, r.text
    st = r.json()
    assert st["trial"]["status"] == "passed", st["trial"]["problems"]
    out = await commit(c, st)
    return Source(name=name, id=out["source"]["id"], snapshot=out["snapshot_id"], raws=[(raw, fn)],
                  imports=[out["import_id"]], snapshots=[out["snapshot_id"]])


async def add_period(c: AsyncClient, src: Source, raw: bytes, fn: str, *,
                     acceptances: list[dict[str, str]] | None = None, human: tuple[str, str] | None = None) -> None:
    """上传新一期 → 试运行（human=(起, 止) 时按人工录入的统计期、署名「录入员甲」重新试运行）→ 提交。"""
    r = await c.post(f"{API}/{src.id}/reupload", files={"file": (fn, raw, XLSX)})
    assert r.status_code == 201, r.text
    st = r.json()
    if human is not None:
        r = await c.post(f"{API}/imports/{st['id']}/trial", json={
            "context_inputs": {"统计期": {"start": human[0], "end": human[1]}}, "signed_by": SIGNER})
        assert r.status_code == 200, r.text
        st = r.json()
    assert st["trial"]["status"] in ("passed", "needs_decision"), st["trial"]["problems"]
    out = await commit(c, st, acceptances)
    src.snapshot = out["snapshot_id"]
    src.snapshots.append(out["snapshot_id"])
    src.raws.append((raw, fn))
    src.imports.append(out["import_id"])


async def simple_upload(c: AsyncClient, raw: bytes, fn: str, name: str) -> Source:
    """期 1 的简单导入（/upload）：没有导入清单，表结构快照没有 import_mode。"""
    r = await c.post(f"{API}/upload", files={"file": (fn, raw, XLSX)}, data={"name": name})
    assert r.status_code == 201, r.text
    async with SessionLocal() as session:
        row = (await session.execute(select(DataSource).where(DataSource.name == name))).scalar_one()
        return Source(name=name, id=row.id, snapshot=row.current_snapshot_id or "", raws=[(raw, fn)], imports=[])


async def manual_source(name: str, path: Path) -> Source:
    """手工登记的 SQLite 源：查询快照没有 data_version。"""
    db = sqlite3.connect(path)
    db.executescript('CREATE TABLE "订单汇总" ("地区" TEXT PRIMARY KEY, "订单数" INTEGER, "金额" REAL);')
    db.executemany('INSERT INTO "订单汇总" VALUES (?, ?, ?)',
                   [("分区甲", 523, 40211.5), ("分区乙", 466, 35982.0), ("分区丙", 301, 22870.75)])
    db.commit()
    db.close()
    async with SessionLocal() as session:
        row = DataSource(name=name, kind="sqlite", database=str(path), readonly=True, enabled=True, options={})
        session.add(row)
        await session.commit()
        return Source(name=name, id=row.id, snapshot="", raws=[], imports=[])


async def set_masks(source_id: str, masks: str | None) -> None:
    async with SessionLocal() as session:
        row = await session.get(DataSource, source_id)
        options = dict(row.options or {})
        if masks:
            options["mask_columns"] = masks
        else:
            options.pop("mask_columns", None)
        row.options = options
        await session.commit()


def list_book(sheet: str, header: tuple[str, ...], rows: list[tuple[Any, ...]]) -> bytes:
    """合成的简单列表 xlsx：第 1 行表头，之后逐行。"""
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    ws.append(list(header))
    for row in rows:
        ws.append(list(row))
    return save(wb)


def rework(raw: bytes, edit: Any, sheet: str = SHEET) -> bytes:
    """在合成夹具上改几处（edit(工作表)），其余照旧。openpyxl 存盘会丢掉公式的缓存值：先按原文件读出缓存值，存盘后
    用 excel_saved 填回去，表内合计的公式照样有缓存，核对不受影响（期 3 验收 edit_cells 的写法）。"""
    book = load_workbook(io.BytesIO(raw))
    values = load_workbook(io.BytesIO(raw), data_only=True)
    ws, wv = book[sheet], values[sheet]
    cached = {c.coordinate: wv[c.coordinate].value for row in ws.iter_rows() for c in row
              if isinstance(c.value, str) and c.value.startswith("=")}
    edit(ws)
    return excel_saved(save(book), sheet, cached)


def xl(raw: bytes, ref: str, sheet: str = SHEET) -> Any:
    """原件里一格的保存值（公式取缓存值）。只在测试里读、作对照，产品代码不读原件。"""
    return load_workbook(io.BytesIO(raw), data_only=True)[sheet][ref].value


def sha256(raw: bytes) -> str:
    return hashlib.sha256(raw).hexdigest()


# ==========================================================================
# 运行
# ==========================================================================


@dataclass
class Cite:
    """报告里的一处单元格引用：key 认片段；node 相同的几处共用一个工具节点（同一条查询）。"""

    key: str
    source: str
    sql: str
    column: str
    row: int = 0
    node: str | None = None


@dataclass
class RunInfo:
    id: str
    doc_artifact: str
    doc: dict[str, Any]
    #: 用例 key → 片段
    segs: dict[str, dict[str, Any]]
    #: 用例 key → 所在的句子 id
    units: dict[str, str]
    #: 用例 key → 目录里的查询编号
    aliases: dict[str, str]


def node(nid: str, ntype: str, **config: Any) -> dict[str, Any]:
    return {"id": nid, "type": ntype, "position": {"x": 0, "y": 0}, "data": {"label": nid, "config": config}}


def chain(*nodes: dict[str, Any]) -> dict[str, Any]:
    return {"nodes": list(nodes), "edges": [{"source": a["id"], "target": b["id"]} for a, b in zip(nodes, nodes[1:])]}


async def _wait(run_id: str, statuses: tuple[str, ...]) -> Run:
    for _ in range(1200):
        await asyncio.sleep(0.05)
        async with SessionLocal() as session:
            row = await session.get(Run, run_id)
        if row.status in statuses and (row.status == "interrupted" or row.manifest_seq is not None):
            return row
    raise AssertionError(f"run {run_id} 没有结束：{row.status}")


async def run_report(sources: dict[str, Source], cites: list[Cite], *, gate: bool = False,
                     legacy: bool = False) -> RunInfo:
    """跑一次「工具节点 × N → 报告节点 →（审批）→ 输出」。每处引用一句话，句子不含数字；查询按工具节点的先后编号。

    legacy=True：报告节点的 compose_doc 换成去掉 provenance 键的包装（6.1「期 4 之前的文档」），存下的就是一份真正的
    老形状文档。gate=True：报告之后接一个人工审批节点，运行停在审批上（未封存）。"""
    from app.engine.nodes import report as report_node
    from app.providers import mock_model

    tools: list[str] = []
    alias_of: dict[str, str] = {}
    nodes = [node("start", "input")]
    for cite in cites:
        nid = f"q_{cite.node or cite.key}"
        if nid not in tools:
            tools.append(nid)
            nodes.append(node(nid, "tool", tool=f"db_query__{sources[cite.source].name}", args={"sql": cite.sql}))
        alias_of[cite.key] = f"Q{tools.index(nid) + 1}"
    assert len(cites) <= len(LABELS)
    lines = ["## 推断来源验收", ""]
    for label, cite in zip(LABELS, cites):
        lines += [f"{label}的值为 [[v:{alias_of[cite.key]}.r{cite.row}.{cite.column}]]。", ""]
    text = "\n".join(lines)
    nodes.append(node("write", "report", instructions="写推断来源验收报告", on_violation="flag"))
    if gate:
        nodes.append(node("gate", "human", mode="approve", title="确认"))
    nodes.append(node("out", "output", fields=[{"name": "answer", "value": "{{ nodes.write.text }}"}]))

    real_decide, real_compose = mock_model.MockChatModel._decide, report_node.compose_doc

    def legacy_compose(*a: Any, **kw: Any) -> dict[str, Any]:
        doc = real_compose(*a, **kw)
        doc.pop("provenance", None)
        return doc

    mock_model.MockChatModel._decide = lambda self, messages: AIMessage(content=text)
    if legacy:
        report_node.compose_doc = legacy_compose
    try:
        run = await run_manager.start(graph=chain(*nodes), input_payload={})
        row = await _wait(run.id, ("interrupted",) if gate else ("succeeded", "failed"))
    finally:
        mock_model.MockChatModel._decide = real_decide
        report_node.compose_doc = real_compose
    if gate:
        assert row.status == "interrupted", row.error
        async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as c:
            doc_id = (await c.get(f"/api/runs/{row.id}/evidence")).json()["reports"][0]["doc_artifact"]
    else:
        assert row.status == "succeeded", row.error
        doc_id = row.output["_evidence"]["doc_artifact"]
    doc = artifact_store.load(doc_id)
    assert ("provenance" in doc) is not legacy
    catalog = doc["catalog"]
    segs: dict[str, dict[str, Any]] = {}
    units: dict[str, str] = {}
    unit_of = {s["id"]: u["id"] for _, u in iter_units(doc) for s in u.get("segments") or []}
    cells = [s for s in iter_segments(doc) if (s.get("cite") or {}).get("kind") == "cell"]
    for cite in cites:
        alias = alias_of[cite.key]
        assert catalog[alias]["node_id"] == f"q_{cite.node or cite.key}", (cite.key, catalog[alias])
        hit = [s for s in cells if s["cite"].get("alias") == alias
               and (s["cite"].get("locator") or {}).get("row") == cite.row
               and (s["cite"].get("locator") or {}).get("column") == cite.column]
        assert len(hit) == 1, (cite.key, hit)
        segs[cite.key] = hit[0]
        units[cite.key] = unit_of[hit[0]["id"]]
    return RunInfo(id=row.id, doc_artifact=doc_id, doc=doc, segs=segs, units=units, aliases=alias_of)


# ==========================================================================
# 接口与契约
# ==========================================================================


def _dc(cls: type, d: Any, **nested: Any) -> Any:
    """JSON 的一层 → dataclass：键集合与顺序必须和契约的字段逐个相同（多一个、少一个、改名都报）。"""
    assert isinstance(d, dict), (cls.__name__, d)
    names = [f.name for f in fields(cls)]
    assert list(d) == names, f"{cls.__name__} 的键与契约不符：{list(d)} != {names}"
    return cls(**{**d, **nested})


def _opt(cls: type, d: Any, **nested: Any) -> Any:
    return None if d is None else _dc(cls, d, **nested)


def _accept(d: Any) -> AcceptanceView | None:
    return _opt(AcceptanceView, d)


def from_json(body: dict[str, Any]) -> ProvenanceOut:
    """推断来源接口的响应 → ProvenanceOut（逐层核对键集合），给 contract_problems 用。"""
    v = body.get("version")
    version = None if v is None else _dc(
        VersionView, v, tables=[_dc(TableRef, t) for t in v["tables"]],
        parts=[_dc(PartView, p, period=_opt(PeriodView, p["period"]),
                   acceptances=[_accept(a) for a in p["acceptances"]],
                   purged=_opt(StateNote, p["purged"]), revoked=_opt(StateNote, p["revoked"])) for p in v["parts"]])
    cs = body.get("cell_source")
    cell_source = None if cs is None else _dc(
        CellSource, cs, **{"from": [_dc(FromCell, f) for f in cs["from"]]}, year=_opt(YearSource, cs["year"]),
        canonical=_opt(CanonicalView, cs["canonical"]), recheck=_opt(Recheck, cs["recheck"]))
    return _dc(ProvenanceOut, body, report=_dc(ReportRef, body["report"]), cell=_opt(CellRef, body["cell"]),
               reason=_opt(Reason, body["reason"]), alert=_opt(Alert, body["alert"]), version=version,
               cell_source=cell_source,
               checks=[_dc(RelatedCheck, c, acceptance=_accept(c["acceptance"])) for c in body["checks"]])


async def provenance(c: AsyncClient, run: RunInfo, key: str, *, report: str | None = None) -> ProvenanceOut:
    """GET 推断来源：200、键集合与契约相同、contract_problems 为空、报告和片段对得上。"""
    seg_id = run.segs[key]["id"]
    url = f"/api/runs/{run.id}/evidence/segments/{seg_id}/provenance" + (f"?report={report}" if report else "")
    r = await c.get(url)
    assert r.status_code == 200, r.text
    out = from_json(r.json())
    assert contract_problems(out) == [], contract_problems(out)
    assert out.segment == seg_id and out.report == ReportRef(node_id="write", doc_artifact=run.doc_artifact)
    return out


async def query_step(c: AsyncClient, run: RunInfo, key: str) -> dict[str, Any]:
    """片段接口里这一格的查询步骤（单元格片段的出处链只有这一步）。"""
    r = await c.get(f"/api/runs/{run.id}/evidence/segments/{run.segs[key]['id']}")
    assert r.status_code == 200, r.text
    [step] = r.json()["chain"]
    assert step["step"] == "query"
    return step


def result_value(run: RunInfo, key: str) -> Any:
    """查询快照里被引用的那一格（内容寻址的工件，观察点）。"""
    seg = run.segs[key]
    snap = artifact_store.load(run.doc["catalog"][seg["cite"]["alias"]]["artifact"])
    loc = seg["cite"]["locator"]
    return snap["rows"][loc["row"]][snap["columns"].index(loc["column"])]


def query_snapshot(run: RunInfo, key: str) -> dict[str, Any]:
    return artifact_store.load(run.doc["catalog"][run.aliases[key]]["artifact"])


def cell_ref(run: RunInfo, key: str) -> CellRef:
    seg = run.segs[key]
    loc = seg["cite"]["locator"]
    return CellRef(alias=seg["cite"]["alias"], row=loc["row"], column=loc["column"],
                   artifact=run.doc["catalog"][seg["cite"]["alias"]]["artifact"])


def froms(out: ProvenanceOut) -> list[tuple[str, str | None, str, str, str | None, str | None]]:
    return [(f.role, f.column, f.sheet, f.cell, f.text, f.locate_title) for f in out.cell_source.from_]


def check_of(out: ProvenanceOut, cid: str) -> RelatedCheck:
    hit = [x for x in out.checks if x.id == cid]
    assert len(hit) == 1, (cid, [x.id for x in out.checks])
    return hit[0]


def assert_table_only(out: ProvenanceOut, code: str, *, detail: str = "", alert: bool = False,
                      text: str | None = None) -> None:
    """只给表级来历：原因、细分、原文；标红的同时给 alert；不给格子，相关核对为空，数据版本照给（2.3）。"""
    assert out.status == "table_only", (out.status, out.reason)
    assert out.reason is not None and (out.reason.code, out.reason.detail) == (code, detail), out.reason
    assert out.reason.text == (text if text is not None else REASON_TEXT[code])
    if alert:
        assert out.alert == Alert(code=code, text=out.reason.text)
    else:
        assert out.alert is None
    assert out.cell_source is None and out.checks == [] and out.version is not None
    assert not any(p.has_row for p in out.version.parts)


async def import_row(import_id: str) -> TableImport:
    async with SessionLocal() as session:
        return await session.get(TableImport, import_id)


async def snapshot_row(snapshot_id: str) -> SourceSnapshot:
    async with SessionLocal() as session:
        return await session.get(SourceSnapshot, snapshot_id)


async def recipe_seq_of(recipe_id: str | None) -> int | None:
    from app.db.models import TableRecipe

    async with SessionLocal() as session:
        rec = await session.get(TableRecipe, recipe_id) if recipe_id else None
        return rec.seq if rec else None


async def manifest_id_of(import_id: str) -> str:
    """导入记录上的清单 id（观察点：只用来对照，产品代码按表结构快照的哈希链取，不经这个可改的列）。"""
    async with SessionLocal() as session:
        return (await session.get(TableImport, import_id)).manifest_artifact


async def expected_part(src: Source, index: int, *, has_row: bool, period: PeriodView | None, region: str | None,
                        raw_state: str = "kept", purged: StateNote | None = None, revoked: StateNote | None = None,
                        acceptances: list[AcceptanceView] | None = None) -> PartView:
    """数据版本里的一期应有的样子（2.8.1）：文件名、原件哈希取自上传的字节；seq、配方哈希、提交时间、署名取自这一期的
    导入清单（内容寻址的工件，观察点）；配方第几版取自配方记录；原件状态、清除、作废是导入记录的当前状态。"""
    raw, fn = src.raws[index]
    mid = await manifest_id_of(src.imports[index])
    m = artifact_store.load(mid)
    return PartView(import_id=src.imports[index], seq=m["seq"], manifest=mid, period=period, file_name=fn,
                    raw_sha256=sha256(raw), region=region, excluded_rows=0, recipe_sha256=m["recipe"]["sha256"],
                    recipe_seq=await recipe_seq_of(m["recipe"]["id"]), committed_at=m["created_at"],
                    signed_by=(m.get("signed_by") or {}).get("name"), acceptances=acceptances or [],
                    raw_state=raw_state, purged=purged, revoked=revoked, has_row=has_row)


def manifest_acceptance(manifest_id: str, check_id: str) -> AcceptanceView:
    """清单里这条接受（内容寻址，是当时的事实）应有的 AcceptanceView。"""
    m = artifact_store.load(manifest_id)
    [a] = [a for a in m["acceptances"]["overrides"] if a["check_id"] == check_id]
    [c] = [c for c in m["checks"] if c["id"] == check_id]
    return AcceptanceView(check_id=check_id, title=c["title"], kind="override", reason=a["reason"],
                          signed_by=a.get("signed_by"), at=a.get("at"))


def august(source: str = "cells") -> PeriodView:
    return PeriodView(start="2026-08-01", end="2026-08-31", source=source, cells=[f"{SHEET}!B2"], signed_by=None)


def september(source: str = "cells") -> PeriodView:
    if source == "human":
        return PeriodView(start="2026-09-01", end="2026-09-30", source="human", cells=[], signed_by=SIGNER)
    return PeriodView(start="2026-09-01", end="2026-09-30", source="cells", cells=[f"{SHEET}!B2"], signed_by=None)


def axis(col: str, row: int = 4) -> tuple[str, str | None, str, str, str | None, str | None]:
    return ("axis_header", "日期", SHEET, f"{col}{row}", None, None)


def label(ref: str, text: str | None, column: str | None, role: str = "row_label"
          ) -> tuple[str, str | None, str, str, str | None, str | None]:
    return (role, column, SHEET, ref, text, None)


def title(ref: str, column: str = "时段类别") -> tuple[str, str | None, str, str, str | None, str | None]:
    return ("section_title", column, SHEET, ref, None, DAY_TITLE)


# ==========================================================================
# 夹具：模块级的世界（不改状态的用例共用）
# ==========================================================================


@dataclass
class World:
    base: Path
    sources: dict[str, Source]
    main: RunInfo
    legacy: RunInfo
    gate: RunInfo
    broken: RunInfo


#: 主运行引用的格子（每处一句话）。flow 是 8 月 D00（每期替换），把 2026-08-03 的分区乙改成 0、全日客流改成等于
#: 分区甲（A24 的陷阱：两列的值恰好相等）
def _main_cites() -> list[Cite]:
    cites = [
        Cite("a4", "flow", wide_sql("2026-08-05"), "全日客流"),
        Cite("a1", "flow", 'SELECT "日期", "分区甲" AS "分区乙" FROM "日客流" WHERE "日期" = \'2026-08-06\'', "分区乙"),
        Cite("a2_join", "flow", 'SELECT a."日期", b."全日客流" FROM "日客流" a JOIN "日客流" b ON a."日期" = b."日期" '
                                'WHERE a."日期" = \'2026-08-07\'', "全日客流"),
        Cite("a2_comma", "flow", 'SELECT a."日期", b."全日客流" FROM "日客流" a, "日客流" b WHERE a."日期" = b."日期" '
                                 'AND a."日期" = \'2026-08-07\'', "全日客流"),
        Cite("a3", "flow", 'SELECT "日期", "时段", "客流" FROM "时段客流" WHERE "时段" = \'7-8\' '
                           'AND "日期" = \'2026-08-03\'', "客流"),
        Cite("a8", "flow", 'SELECT "日期", "合计项", "客流" FROM "时段客流_表内合计" WHERE "日期" = \'2026-08-10\' '
                           'AND "合计项" = \'18-22时合计\'', "客流"),
        Cite("a11", "flow", 'SELECT SUM("全日客流") AS "合计", COUNT(*) AS "天数" FROM "日客流"', "合计"),
        Cite("a16", "flow", 'SELECT "全日客流" FROM "日客流" WHERE "日期" = \'2026-08-05\'', "全日客流"),
        Cite("long", "flow", long_sql("2026-08-20", "13-14"), "客流"),
        Cite("a7", "acc", long_sql("2026-09-12", "8-9"), "客流"),
        Cite("a18", "acc", long_sql("2026-09-12", "8-9"), "时段", node="a7"),
        Cite("a13_bad", "d26", wide_sql("2026-09-10"), "全日客流"),
        Cite("a13_other", "d26", wide_sql("2026-09-11"), "全日客流"),
        Cite("union_bad", "d26acc", wide_sql("2026-09-10"), "全日客流"),
        Cite("union_other", "d26acc", wide_sql("2026-09-11"), "全日客流"),
        Cite("a15", "masked", wide_sql("2026-08-16", "分区甲"), "分区甲"),
        Cite("a17", "human", wide_sql("2026-09-04"), "全日客流"),
        Cite("a19", "list", 'SELECT "地区", "金额" FROM "月报" WHERE "地区" = \'分区丙\'', "金额"),
        Cite("a15_pk", "list_masked", 'SELECT "地区", "金额" FROM "月报" WHERE "地区" = \'分区丙\'', "金额"),
        Cite("a12", "simple", 'SELECT "地区", "金额" FROM "地区汇总" WHERE "地区" = \'分区乙\'', "金额"),
        Cite("a10", "manual", 'SELECT "地区", "订单数" FROM "订单汇总" WHERE "地区" = \'分区丙\'', "订单数"),
        Cite("d28", "d28", wide_sql("2026-09-06"), "全日客流"),
        Cite("d28h", "d28h", wide_sql("2026-09-06"), "全日客流"),
    ]
    cites += [Cite(f"a24_{k}", "flow", sql, "全日客流") for k, sql in A24_SQL.items()]
    # A20 ⑤：同一数据版本（D28 那个源）5 条查询、15 句话
    for i, day in enumerate(("2026-09-01", "2026-09-02", "2026-09-03", "2026-09-04", "2026-09-05")):
        sql = f'SELECT "日期", "全日客流", "分区甲", "分区乙" FROM "日客流" WHERE "日期" = \'{day}\''
        for col in ("全日客流", "分区甲", "分区乙"):
            cites.append(Cite(f"cost_{i}_{col}", "d28", sql, col, node=f"cost_{i}"))
    return cites


async def _build(base: Path) -> World:
    sources: dict[str, Source] = {}
    await run_manager.setup()
    try:
        async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as c:
            raw, fn = flow_workbook(AUG, 31, seed=71)

            def trap(ws: Any) -> None:
                ws["E7"] = 0                      # 2026-08-03 的分区乙
                ws["E5"] = ws["E6"].value         # 全日客流 = 分区甲 + 0
            sources["flow"] = await first_import(c, rework(raw, trap), fn)
            sources["acc"] = await first_import(c, *flow_workbook(AUG, 31, seed=72), mode="accumulate")
            await add_period(c, sources["acc"], *flow_workbook(SEP, 30, seed=73, variant="D24w"))
            sources["d26"] = await first_import(c, *flow_workbook(AUG, 31, seed=74))
            await add_period(c, sources["d26"], *flow_workbook(SEP, 30, seed=75, variant="D26"),
                             acceptances=[{"check_id": "R1", "reason": REASON_D26}])
            sources["d26acc"] = await first_import(c, *flow_workbook(AUG, 31, seed=89), mode="accumulate")
            await add_period(c, sources["d26acc"], *flow_workbook(SEP, 30, seed=90, variant="D26"),
                             acceptances=[{"check_id": "R1", "reason": REASON_D26}])
            sources["human"] = await first_import(c, *flow_workbook(AUG, 31, seed=76))
            await add_period(c, sources["human"], *flow_workbook(SEP, 30, seed=77, variant="noperiod"),
                             human=("2026-09-01", "2026-09-30"))
            sources["masked"] = await first_import(c, *flow_workbook(AUG, 31, seed=78))
            sources["list"] = await first_import(c, *lab.c01(), rename=False)
            sources["list_masked"] = await first_import(c, *lab.c01(), rename=False)
            book = list_book("地区汇总", ("地区", "金额", "订单数"),
                             [("分区甲", 31847.5, 412), ("分区乙", 27093.25, 377), ("分区丙", 18456.0, 268)])
            sources["simple"] = await simple_upload(c, book, "地区汇总_2026-08.xlsx", unique())
            sources["d28"] = await first_import(c, *flow_workbook(SEP, 30, seed=79, variant="D28"))
            raw, fn = flow_workbook(SEP, 30, seed=80, variant="D28")

            def hide(ws: Any) -> None:
                ws.row_dimensions[32].hidden = True
            sources["d28h"] = await first_import(c, rework(raw, hide), fn)
        sources["manual"] = await manual_source(unique(), base / "orders.db")
        await set_masks(sources["masked"].id, "分区甲")
        await set_masks(sources["list_masked"].id, "地区")

        main = await run_report(sources, _main_cites())
        # list_masked 的遮罩只留在查询快照里（mask_columns）：数据源现在的设置清掉，R10 要靠快照记下的那份拦住
        await set_masks(sources["list_masked"].id, None)
        legacy = await run_report(sources, [Cite("a6", "flow", wide_sql("2026-08-05"), "全日客流")], legacy=True)
        gate = await run_report(sources, [Cite("a22", "flow", wide_sql("2026-08-05"), "全日客流")], gate=True)
        broken = await run_report(sources, [Cite("a23", "flow", wide_sql("2026-08-05"), "全日客流")])
    finally:
        await run_manager.shutdown()
        for src in sources.values():
            await engines.invalidate(src.id)
        # 这个事件循环里建的连接留在连接池里，换了事件循环的用例取到会出错：交还之前全部关掉
        await db_engine.dispose()
    return World(base, sources, main, legacy, gate, broken)


@pytest.fixture(scope="module", autouse=True)
def module_store(tmp_path_factory):
    """整个模块一个数据目录（世界的数据文件、工件、原件存档都在里面）；规则起草不完整时的 AI 可用性换成「不可用」
    的桩，连模型对象都不建；验收必须跑真管线（recipe_imports.pipeline() 接的是各工作包的真函数）。"""
    from app.data import recipe_ai, recipe_engine, recipe_imports, recipe_suggest
    from app.data.recipe_types import AiAvailability

    p = recipe_imports.pipeline()
    assert p.execute is recipe_engine.execute and p.draft is recipe_suggest.draft

    async def no_model(session: Any) -> Any:
        return None, "", AiAvailability(False, reason="验收测试不接模型")

    data = tmp_path_factory.mktemp("p4e") / "data"
    (data / "uploads").mkdir(parents=True)
    mp = pytest.MonkeyPatch()
    mp.setattr(settings, "data_dir", data)
    mp.setattr(recipe_ai, "resolve_draft_model", no_model)
    yield data
    mp.undo()


@pytest.fixture(scope="module")
def world(module_store) -> World:
    loop = asyncio.new_event_loop()
    try:
        return loop.run_until_complete(_build(module_store))
    finally:
        loop.close()


@pytest.fixture
async def client():
    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as c:
        yield c


@pytest.fixture
async def fresh_engines(world):
    """世界的数据文件引擎在用例自己的事件循环里重新建（连接池不跨事件循环），用完关掉。"""
    for src in world.sources.values():
        await engines.invalidate(src.id)
    yield world
    for src in world.sources.values():
        await engines.invalidate(src.id)


@pytest.fixture
async def engine_up():
    await run_manager.setup()
    yield
    await run_manager.shutdown()


@pytest.fixture
async def own(engine_up):
    """改状态的用例自己的源：用完关掉它们的数据层引擎。"""
    made: list[Source] = []
    yield made
    for src in made:
        await engines.invalidate(src.id)


async def own_flow(made: list[Source], *, seed: int, legacy: bool = False) -> tuple[Source, RunInfo]:
    """一个独立的 8 月 D00 每期替换源，跑一次只引用 A4 那一格的运行。"""
    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as c:
        src = await first_import(c, *flow_workbook(AUG, 31, seed=seed))
    made.append(src)
    run = await run_report({"flow": src}, [Cite("a4", "flow", wide_sql("2026-08-05"), "全日客流")], legacy=legacy)
    return src, run




@pytest.fixture
def tripwire(monkeypatch) -> list[str]:
    """把 provenance、direct_select 两个模块的公开函数换成一调用就记下来并抛错的桩（P4-SPEC 3.3）：期 4 之前的文档、
    非上传源的裁判摘录不许碰它们。"""
    from app.data import provenance as provenance_mod
    from app.engine import direct_select as direct_select_mod

    hits: list[str] = []
    for mod in (provenance_mod, direct_select_mod):
        for name, obj in vars(mod).items():
            if name.startswith("_") or not callable(obj) or isinstance(obj, type) \
                    or getattr(obj, "__module__", None) != mod.__name__:
                continue

            def boom(*_a: Any, __name: str = f"{mod.__name__}.{name}", **_kw: Any) -> Any:
                hits.append(__name)
                raise AssertionError(f"不该调用 {__name}")

            monkeypatch.setattr(mod, name, boom)
    return hits


@pytest.fixture
def judge_spy(monkeypatch) -> list[dict[str, str]]:
    """按需裁判的桩：截下交给裁判的摘录，不调用任何模型、不记判定、不计费（outcome 里 calls 为 0，接口不追加事件，
    同一句可以再点一次）。"""
    from app.engine import judge as judging

    seen: list[dict[str, str]] = []

    async def fake_run(request: Any, **_kw: Any) -> dict[str, Any]:
        seen.append(dict(request.excerpts))
        return {"verdicts": {}, "calls": 0, "cost_usd": 0.0, "limits_hit": [], "unjudged": {}, "gaps": [],
                "notes": [], "model": None, "priced": None, "duration_ms": 0}

    monkeypatch.setattr(judging, "run_request", fake_run)
    return seen


def db_rowid(snapshot: SourceSnapshot, table: str, pk: dict[str, Any]) -> int:
    """按主键在快照库上取 rowid（只读、immutable，不改文件）：cell_source.rowid 必须是快照库的 rowid。"""
    where = " AND ".join(f'"{k}" = ?' for k in pk)
    with closing(sqlite3.connect(f"file:{snapshot.db_path}?mode=ro&immutable=1", uri=True)) as conn:
        got = conn.execute(f'SELECT rowid FROM "{table}" WHERE {where}', list(pk.values())).fetchall()
    assert len(got) == 1, (table, pk, got)
    return got[0][0]


async def assert_drilled(out: ProvenanceOut, run: RunInfo, key: str, *, table: str, column: str, role: str,
                         pk: dict[str, Any], sheet: str, cell: str, part_seq: int, kind: str = "data",
                         sealed: bool = True) -> None:
    """inferred 的公共部分：格子、主键、快照库的 rowid、按主键回查的那一句（? 占位、按主键顺序，WP-0 交接）。"""
    assert out.status == "inferred", (out.status, out.reason, out.alert)
    assert out.reason is None and out.alert is None and out.sealed is sealed
    assert out.cell == cell_ref(run, key)
    cs = out.cell_source
    assert (cs.table, cs.column, cs.column_role, cs.kind) == (table, column, role, kind)
    assert cs.pk == pk and list(cs.pk) == list(pk)
    # 键序（回查的 WHERE、params 同序）是冻结表结构的 primary_key，不是 SELECT 里的顺序（WP-0 交接）
    frozen = artifact_store.load(query_snapshot(run, key)["schema_artifact"])["tables"][table]
    assert list(cs.pk) == frozen["primary_key"]
    assert (cs.sheet, cs.cell, cs.part_seq) == (sheet, cell, part_seq)
    snap = await snapshot_row(query_snapshot(run, key)["data_version"])
    assert cs.rowid == db_rowid(snap, table, pk)
    where = " AND ".join(f'"{k}" = ?' for k in pk)
    assert cs.recheck == Recheck(sql=f'SELECT rowid, "{column}" FROM "{table}" WHERE {where}',
                                 params=list(pk.values()), ok=True)
    assert cs.raw_purged is False and cs.merged_fill is False
    hits = [p for p in out.version.parts if p.has_row]
    assert len(hits) == 1 and hits[0].seq == part_seq


# ==========================================================================
# final.md 第 5 节「期 4」的四条（A1–A6）
# ==========================================================================



async def flow_version(world: World, out: ProvenanceOut, tables: list[TableRef]) -> None:
    """flow 源（8 月 D00、每期替换）的数据版本：一期、清单里的事实逐项对得上（表级来历只靠清单，2.6）。"""
    src = world.sources["flow"]
    has_row = out.status == "inferred"
    assert out.version == VersionView(
        source=src.name, snapshot_id=src.snapshot, mode="replace", union=False, tables=tables, manifest_view=True,
        parts=[await expected_part(src, 0, has_row=has_row, period=august(), region=f"{SHEET}!B4:AG30")])


DAILY = [TableRef(name="日客流", kind="data")]


async def test_a1_renamed_column_only_gets_table_level_provenance(world, fresh_engines, client):
    """A1：`SELECT "日期", "分区甲" AS "分区乙"` 不下钻——改名之后无法确认这一格对应哪一列（R7 · alias）。
    数据版本照给；查询步骤带 `provenance: true`（界面因此才会来问）。"""
    w = world
    out = await provenance(client, w.main, "a1")
    assert_table_only(out, "alias", detail="alias")
    assert out.cell == cell_ref(w.main, "a1") and out.sealed is True
    await flow_version(w, out, DAILY)
    assert (await query_step(client, w.main, "a1"))["provenance"] is True


@pytest.mark.parametrize("key,detail", [("a2_join", "join"), ("a2_comma", "comma_join")])
async def test_a2_self_joins_are_not_drilled(world, fresh_engines, client, key, detail):
    """A2：自连接（JOIN 与逗号两种写法）不下钻：R7 · multi_table，细分分别是 join、comma_join。"""
    out = await provenance(client, world.main, key)
    assert_table_only(out, "multi_table", detail=detail)
    await flow_version(world, out, DAILY)


async def test_a3_a_null_cell_is_not_drilled(world, fresh_engines, client):
    """A3：占位符那一行（7-8 时）存为空值：片段本身是 resolved、显示「—」；推断来源 R11 · null_value。"""
    w = world
    seg = w.main.segs["a3"]
    assert seg["state"] == "deterministic" and seg["cite"]["status"] == "resolved" and seg["text"] == "—"
    assert result_value(w.main, "a3") is None
    assert xl(w.sources["flow"].raws[0][0], "E10") == "·"           # 原表里那一格是占位符
    out = await provenance(client, w.main, "a3")
    assert_table_only(out, "null_value")
    await flow_version(w, out, [TableRef(name="时段客流", kind="data")])



WIDE_IDS = ["R1", "R2", "C1", "C2"]


async def test_a4_a_direct_cell_names_its_sheet_and_cell(world, fresh_engines, client):
    """A4：满足条件时给出推断的来源：客流汇总!G5；日期取自表头格 G4、年份取自统计期（B2）、指标名取自行标签 B5；
    按主键回查一致；R1 本期通过、这一行成立；数据版本一期，文件名、原件哈希与上传的文件一致。
    G5 的保存值（openpyxl 读夹具字节，只作对照）等于查询结果那一格。"""
    w = world
    raw, _ = w.sources["flow"].raws[0]
    out = await provenance(client, w.main, "a4")
    await assert_drilled(out, w.main, "a4", table="日客流", column="全日客流", role="measure",
                         pk={"日期": "2026-08-05"}, sheet=SHEET, cell="G5", part_seq=1)
    cs = out.cell_source
    assert cs.part_rowid == cs.rowid                                  # 单期：快照库就是那一期的构建库
    assert (cs.header, cs.unit) == ("全日客流（人次）", "人次")
    assert froms(out) == [axis("G"), label("B5", "全日客流（人次）", None)]
    assert cs.year == YearSource(source="period", cells=[f"{SHEET}!B2"], signed_by=None, mixed=False)
    assert cs.canonical is None
    assert [x.id for x in out.checks] == WIDE_IDS
    r1 = check_of(out, "R1")
    assert (r1.kind, r1.part_status, r1.row_status, r1.cell_status, r1.acceptance) == \
        ("relation_sum_eq", "passed", "passed", None, None)
    r2 = check_of(out, "R2")
    assert (r2.kind, r2.part_status, r2.row_status, r2.cell_status) == ("relation_not_comparable", "info", None, None)
    assert [(x.part_status, x.row_status, x.cell_status) for x in out.checks[2:]] == [("passed", None, None)] * 2
    await flow_version(w, out, DAILY)
    assert out.version.parts[0].raw_sha256 == sha256(raw)
    # 对照原件：G5 的保存值就是报告里那一格；G4 是 8 月 5 日的表头，B5 是指标名
    assert xl(raw, "G5") == result_value(w.main, "a4")
    assert xl(raw, "G4") == "8月5日" and xl(raw, "B5") == "全日客流（人次）"
    assert (await query_step(client, w.main, "a4"))["provenance"] is True
    # 指定报告节点（?report=）选的是同一份报告，答复相同
    assert await provenance(client, w.main, "a4", report="write") == out


async def test_long_table_cell_names_the_section_title_by_its_recipe_text(world, fresh_engines, client):
    """2.5 的推演表「交叉表逆透视后的长表」：V16；日期 V4；行标签 B16「13-14」；分段标题 B9 只给配方里的定位文字
    （text 为 null，locate_title 是配方的标题），时段类别取自它。标签原文与库里的值相同，canonical 为 null。"""
    w = world
    raw, _ = w.sources["flow"].raws[0]
    out = await provenance(client, w.main, "long")
    await assert_drilled(out, w.main, "long", table="时段客流", column="客流", role="value",
                         pk={"日期": "2026-08-20", "时段": "13-14"}, sheet=SHEET, cell="V16", part_seq=1)
    assert froms(out) == [axis("V"), label("B16", "13-14", "时段"), title("B9")]
    assert out.cell_source.canonical is None
    assert xl(raw, "V16") == result_value(w.main, "long") and xl(raw, "B9") == DAY_TITLE
    assert [x.id for x in out.checks] == ["R2", "C1", "C2"]


async def test_a5_a_tampered_data_file_is_flagged_red(own, client):
    """A5：A4 的运行跑完之后，数据文件被改了一个字节（chmod 0644 → 改 → chmod 0444，实测 B5）：下一次取引擎时指纹
    对不上、重新核对整份哈希失败 → table_only · db_tampered，标红；数据版本照给（只靠清单），不给格子。"""
    src, run = await own_flow(own, seed=81)
    first = await provenance(client, run, "a4")
    assert first.status == "inferred"
    path = Path((await snapshot_row(src.snapshot)).db_path)
    data = bytearray(path.read_bytes())
    data[len(data) // 2] ^= 0xFF
    os.chmod(path, 0o644)
    path.write_bytes(bytes(data))
    os.chmod(path, 0o444)
    out = await provenance(client, run, "a4")
    assert_table_only(out, "db_tampered", alert=True)
    assert out.alert.text == REASON_TEXT["db_tampered"] and "可能被修改过" in out.alert.text
    assert [p.import_id for p in out.version.parts] == src.imports


async def test_a5b_registered_hash_that_disagrees_with_the_manifest_is_flagged_red(own, client):
    """A5b：用另一个合法的库替换数据文件，并把登记哈希 source_snapshots.db_sha256 改成新文件的哈希：首次开库校验能
    通过，只有「登记哈希 ≠ 导入清单里的库哈希」（R13）拦得住 → table_only · chain_mismatch，标红，数据版本照给。"""
    src, run = await own_flow(own, seed=82)
    assert (await provenance(client, run, "a4")).status == "inferred"
    snap = await snapshot_row(src.snapshot)
    path = Path(snap.db_path)
    other = path.with_name(path.stem + ".other.db")
    shutil.copyfile(path, other)
    os.chmod(other, 0o644)
    with closing(sqlite3.connect(other)) as conn:
        conn.execute('UPDATE "日客流" SET "分区甲" = "分区甲" + 1 WHERE "日期" = \'2026-08-20\'')
        conn.commit()
    new_sha = sha256(other.read_bytes())
    assert new_sha != snap.db_sha256
    os.chmod(path, 0o644)
    os.replace(other, path)
    os.chmod(path, 0o444)
    async with SessionLocal() as session:
        row = await session.get(SourceSnapshot, src.snapshot)
        row.db_sha256 = new_sha
        await session.commit()
    await engines.invalidate(src.id)
    out = await provenance(client, run, "a4")
    assert_table_only(out, "chain_mismatch", alert=True)
    assert out.version.source == src.name and out.version.snapshot_id == src.snapshot


async def test_a6_a_pre_p4_run_keeps_its_panel_and_review(engine_up, client, tripwire):
    """A6：期 4 之前的文档（报告节点的 compose_doc 换成去掉 provenance 键的包装，存下的就是老形状，6.1）：
    查询步骤没有 provenance 键（界面一个请求都不多发）；推断来源接口回 none · legacy_doc，cell 照给；复核（verify_doc）
    没有 state_mismatch、证据图里报告文档复验通过；裁判摘录在 SQL 与「列：」之间没有新行，也不碰 provenance、
    direct_select 两个模块。三份金样逐字不变见 test_provenance_unchanged.py（6.6）。"""
    from app.engine import judge as judging

    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as c:
        src = await first_import(c, *flow_workbook(AUG, 31, seed=83))
    try:
        run = await run_report({"flow": src}, [Cite("a6", "flow", wide_sql("2026-08-05"), "全日客流")], legacy=True)
        assert "provenance" not in run.doc
        step = await query_step(client, run, "a6")
        assert "provenance" not in step and step["rows"] and step["sealed"] is True
        out = await provenance(client, run, "a6")
        assert (out.status, out.reason.code, out.alert, out.version) == ("none", "legacy_doc", None, None)
        assert out.reason.text == REASON_TEXT["legacy_doc"] and out.cell == cell_ref(run, "a6")
        graph = (await client.get(f"/api/runs/{run.id}/evidence")).json()
        assert graph["reports"][0]["hash_ok"] is True and graph["seal"]["ok"] is True
        checked = verify_doc(run.doc, run.doc["catalog"], loader=artifact_store.load)
        assert checked["ok"] is True and not [v for v in checked["violations"] if v.get("code") == "state_mismatch"]
        text = judging.prepare(run.doc, run.doc["catalog"], loader=artifact_store.load).excerpts[run.aliases["a6"]]
        assert excerpt_new_lines(text) == []
        assert tripwire == []
    finally:
        await engines.invalidate(src.id)


def excerpt_new_lines(text: str) -> list[str]:
    """查询摘录里「SQL：」那一行之后、「列：」那一行之前的行（3.2 来历行的约定位置）。"""
    lines = text.split("\n")
    sql = [i for i, line in enumerate(lines) if line.startswith("SQL：")]
    cols = [i for i, line in enumerate(lines) if line.startswith("列：")]
    assert len(sql) == 1 and len(cols) == 1 and sql[0] < cols[0], text
    return lines[sql[0] + 1:cols[0]]


# ==========================================================================
# 补充用例（A7–A26）
# ==========================================================================


async def acc_version(world: World, out: ProvenanceOut, has_row_seq: int | None) -> None:
    """acc 源（8 月 D00 + 9 月 D24w，按期累积）的数据版本：两期按统计期升序，快照清单的 mode、并集。"""
    src = world.sources["acc"]
    assert out.version == VersionView(
        source=src.name, snapshot_id=src.snapshot, mode="accumulate", union=True,
        tables=[TableRef(name="时段客流", kind="data")], manifest_view=True,
        parts=[await expected_part(src, 0, has_row=has_row_seq == 1, period=august(), region=f"{SHEET}!B4:AG30"),
               await expected_part(src, 1, has_row=has_row_seq == 2, period=september(),
                                   region=f"{SHEET}!B4:AF30")])


async def test_a7_a_union_cell_points_into_the_second_period_file(world, fresh_engines, client):
    """A7：累积并集（8 月加 9 月）里 9 月 12 日 8-9 时那一格：第 2 期，9 月文件的 N11（openpyxl 对照）；
    part_rowid = rowid − 快照清单里这张表这一期的 union 起点 + 1；数据版本两期，只有第 2 期 has_row；
    mode accumulate、union true。9 月那一期的时段标签是全角的「８－９」，库里存规范写法「8-9」：canonical 照给。"""
    w = world
    src = w.sources["acc"]
    sep_raw, _ = src.raws[1]
    out = await provenance(client, w.main, "a7")
    await assert_drilled(out, w.main, "a7", table="时段客流", column="客流", role="value",
                         pk={"日期": "2026-09-12", "时段": "8-9"}, sheet=SHEET, cell="N11", part_seq=2)
    cs = out.cell_source
    schema = artifact_store.load(query_snapshot(w.main, "a7")["schema_artifact"])
    sm = artifact_store.load(schema["snapshot_manifest"])
    [part2] = [p for p in sm["parts"] if p["seq"] == 2]
    assert cs.part_rowid == cs.rowid - part2["rows"]["时段客流"]["union"][0] + 1
    assert part2["rows"]["时段客流"]["part"][0] == 1 and cs.part_rowid != cs.rowid
    assert froms(out) == [axis("N"), label("B11", "８－９", "时段"), title("B9")]
    assert cs.canonical == CanonicalView(raw="８－９", canonical="8-9")
    assert xl(sep_raw, "N11") == result_value(w.main, "a7")
    assert xl(sep_raw, "B11") == "８－９" and xl(sep_raw, "N4") == "9月12日"
    await acc_version(w, out, 2)


async def test_a8_a_reported_total_cell(world, fresh_engines, client):
    """A8：原表写明的合计（时段客流_表内合计）：L28（第 28 行、2026-08-10 那一列，openpyxl 对照）；附带合计标签 B28，
    text 取自 DerivedItem.label_raw「18-22 时合计」，canonical 给出规范写法「18-22时合计」；kind reported_total；
    相关核对 K1 这一格一致，G1 这一格一致（L28 在夹具里是公式格）。"""
    w = world
    raw, _ = w.sources["flow"].raws[0]
    out = await provenance(client, w.main, "a8")
    await assert_drilled(out, w.main, "a8", table="时段客流_表内合计", column="客流", role="value",
                         pk={"日期": "2026-08-10", "合计项": "18-22时合计"}, sheet=SHEET, cell="L28", part_seq=1,
                         kind="reported_total")
    assert froms(out) == [axis("L"), label("B28", "18-22 时合计", "合计项", role="total_label")]
    assert out.cell_source.canonical == CanonicalView(raw="18-22 时合计", canonical="18-22时合计")
    assert xl(raw, "L28") == result_value(w.main, "a8") and xl(raw, "B28") == "18-22 时合计"
    assert str(load_workbook(io.BytesIO(raw))[SHEET]["L28"].value).startswith("=SUM(")   # 公式格
    assert [x.id for x in out.checks] == ["K1", "G1", "C1", "C2"]
    for cid, kind in (("K1", "derived_sum"), ("G1", "formula_refs")):
        x = check_of(out, cid)
        assert (x.kind, x.part_status, x.row_status, x.cell_status, x.acceptance) == (kind, "passed", None, "ok", None)
    await flow_version(w, out, [TableRef(name="时段客流_表内合计", kind="reported_total")])


async def test_a9_purging_the_original_keeps_the_cell(own, client):
    """A9：A4 的源经接口 purge-raw 清除原件后再点同一格：仍是 inferred、格子不变（坐标来自导入清单，不读原件）；
    数据版本里 raw_state purged、purged 带时间、署名、理由（导入记录的当前状态）；cell_source.raw_purged 为真。"""
    src, run = await own_flow(own, seed=84)
    before = await provenance(client, run, "a4")
    assert before.status == "inferred" and before.version.parts[0].raw_state == "kept"
    reason = "合成理由：原件不再保留"
    r = await client.post(f"{API}/{src.id}/imports/{src.imports[0]}/purge-raw",
                          json={"confirm": True, "reason": reason, "signed_by": SIGNER})
    assert r.status_code == 200, r.text
    out = await provenance(client, run, "a4")
    assert out.status == "inferred" and out.cell_source.raw_purged is True
    assert replace(out.cell_source, raw_purged=False) == before.cell_source
    part = out.version.parts[0]
    assert part.raw_state == "purged" and part.has_row is True
    assert part.purged is not None and (part.purged.reason, part.purged.signed_by) == (reason, SIGNER)
    assert isinstance(part.purged.at, str) and part.purged.at
    imp = await import_row(src.imports[0])
    assert imp.raw_state == "purged" and imp.purged["reason"] == reason


async def test_a10_a_manual_source_is_not_an_upload(world, fresh_engines, client, tripwire):
    """A10：手工 SQLite 源上的格子：查询步骤**没有** provenance 键；推断来源接口回 none · not_upload；裁判摘录与期 4
    之前同一个形状（SQL 与「列：」之间没有新行，带不带文档标记逐字相同），也不碰 provenance、direct_select。"""
    from app.engine import judge as judging

    w = world
    step = await query_step(client, w.main, "a10")
    assert "provenance" not in step and step["rows"]
    out = await provenance(client, w.main, "a10")
    assert (out.status, out.reason.code, out.alert, out.version) == ("none", "not_upload", None, None)
    assert out.cell == cell_ref(w.main, "a10")
    alias, unit = w.main.aliases["a10"], w.main.units["a10"]
    marked = judging.prepare(w.main.doc, w.main.doc["catalog"], units=[unit], loader=artifact_store.load).excerpts
    old = {k: v for k, v in w.main.doc.items() if k != "provenance"}
    plain = judging.prepare(old, old["catalog"], units=[unit], loader=artifact_store.load).excerpts
    assert list(marked) == [alias] and marked == plain
    assert excerpt_new_lines(marked[alias]) == []
    assert tripwire == []


async def test_a11_an_aggregate_only_gets_table_level_provenance(world, fresh_engines, client):
    """A11：聚合（SUM、COUNT）在 R7 被拒：table_only · expression；数据版本列出涉及的表 [日客流]。"""
    out = await provenance(client, world.main, "a11")
    assert_table_only(out, "expression", detail="expression")
    assert out.version.tables == DAILY
    await flow_version(world, out, DAILY)


async def test_a12_a_simple_upload_has_no_lineage(world, fresh_engines, client):
    """A12：期 1 的简单导入（/upload）：none · simple_upload，只说一句、不标红、没有数据版本。查询步骤带提示（上传源
    都有 data_version），界面会来问一次，得到这一句（2.8.2）。"""
    w = world
    assert (await query_step(client, w.main, "a12"))["provenance"] is True
    out = await provenance(client, w.main, "a12")
    assert (out.status, out.reason.code, out.alert, out.version, out.cell_source) == \
        ("none", "simple_upload", None, None, None)
    assert out.reason.text == REASON_TEXT["simple_upload"] and out.cell == cell_ref(w.main, "a12")


async def d26_parts(src: Source, *, has_row: bool, revoked: StateNote | None = None) -> list[PartView]:
    mid = await manifest_id_of(src.imports[1])
    return [await expected_part(src, 1, has_row=has_row, period=september(), region=f"{SHEET}!B4:AF30",
                                acceptances=[manifest_acceptance(mid, "R1")], revoked=revoked)]


async def test_a13_an_accepted_mismatch_shows_the_row_and_the_reason(world, fresh_engines, client):
    """A13：每期替换的源，先导入 D00、再导入 D26 并写理由接受 R1（9 月 10 日 全日 ≠ 甲 + 乙）；运行固定在 D26 这一版。
    不成立的那一天：R1 本期不成立、这一行不成立，接受理由等于提交时写的；另一天：这一行成立，接受同样带着（这一期的
    接受）。数据版本只有 D26 这一期（seq 2），它的 acceptances 有这一条。"""
    w = world
    src = w.sources["d26"]
    raw, _ = src.raws[1]
    mid = await manifest_id_of(src.imports[1])
    accepted = manifest_acceptance(mid, "R1")
    assert accepted.reason == REASON_D26 and accepted.signed_by is None and accepted.at
    assert "全日客流 = 分区甲 + 分区乙" in accepted.title
    # 夹具：9 月 10 日（L 列）全日 ≠ 甲 + 乙，9 月 11 日（M 列）相等
    assert xl(raw, "L5") != xl(raw, "L6") + xl(raw, "L7") and xl(raw, "M5") == xl(raw, "M6") + xl(raw, "M7")
    for key, day, cell, row in (("a13_bad", "2026-09-10", "L5", "mismatch"),
                                ("a13_other", "2026-09-11", "M5", "passed")):
        out = await provenance(client, w.main, key)
        await assert_drilled(out, w.main, key, table="日客流", column="全日客流", role="measure",
                             pk={"日期": day}, sheet=SHEET, cell=cell, part_seq=2)
        assert [x.id for x in out.checks] == WIDE_IDS
        assert check_of(out, "R1") == RelatedCheck(id="R1", kind="relation_sum_eq", title=accepted.title,
                                                   part_status="mismatch", row_status=row, cell_status=None,
                                                   acceptance=accepted, detail=None)
        assert out.version == VersionView(source=src.name, snapshot_id=src.snapshot, mode="replace", union=False,
                                          tables=DAILY, manifest_view=True,
                                          parts=await d26_parts(src, has_row=True))
        assert xl(raw, cell) == result_value(w.main, key)


async def test_union_row_status_runs_on_the_union_rowid(world, fresh_engines, client):
    """2.7 的行级核对在累积并集上按**快照库**的 rowid 跑（cell_source.rowid 是并集 rowid，不是该期的 part_rowid）：
    8 月 D00 + 9 月 D26（写理由接受）按期累积，9 月 10 日这一行不成立、9 月 11 日成立。按 part_rowid 跑会落到 8 月
    的同号行上（8 月逐日成立），把不成立的那天说成成立。接受理由挂在第 2 期的 R1 上，第 1 期没有接受。"""
    w = world
    src = w.sources["d26acc"]
    sep_raw, _ = src.raws[1]
    mid = await manifest_id_of(src.imports[1])
    accepted = manifest_acceptance(mid, "R1")
    for key, day, cell, row, part_rowid in (("union_bad", "2026-09-10", "L5", "mismatch", 10),
                                            ("union_other", "2026-09-11", "M5", "passed", 11)):
        out = await provenance(client, w.main, key)
        await assert_drilled(out, w.main, key, table="日客流", column="全日客流", role="measure",
                             pk={"日期": day}, sheet=SHEET, cell=cell, part_seq=2)
        assert out.cell_source.part_rowid == part_rowid and out.cell_source.rowid == 31 + part_rowid
        assert check_of(out, "R1") == RelatedCheck(id="R1", kind="relation_sum_eq", title=accepted.title,
                                                   part_status="mismatch", row_status=row, cell_status=None,
                                                   acceptance=accepted, detail=None)
        assert out.version == VersionView(
            source=src.name, snapshot_id=src.snapshot, mode="accumulate", union=True, tables=DAILY,
            manifest_view=True,
            parts=[await expected_part(src, 0, has_row=False, period=august(), region=f"{SHEET}!B4:AG30"),
                   await expected_part(src, 1, has_row=True, period=september(), region=f"{SHEET}!B4:AF30",
                                       acceptances=[accepted])])
        assert xl(sep_raw, cell) == result_value(w.main, key)


async def test_a14_revoked_acceptance_stays_visible_on_the_pinned_version(own, client):
    """A14：A13 之后经接口 revoke-acceptance 作废 D26 这一期的接受：每期替换模式下就是回滚到 D00 那一版，请求体带
    revoke_plan 给的 expected_target_snapshot_id。再打开固定在 D26 版本上的那次运行、点同一格：仍是 inferred（被运行
    引用的版本受保护，数据文件还在）；version.parts[0].revoked 非空（时间、理由、署名），接受理由照旧在。"""
    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as c:
        src = await first_import(c, *flow_workbook(AUG, 31, seed=85))
        await add_period(c, src, *flow_workbook(SEP, 30, seed=86, variant="D26"),
                         acceptances=[{"check_id": "R1", "reason": REASON_D26}])
    own.append(src)
    run = await run_report({"d26": src}, [Cite("a13_bad", "d26", wide_sql("2026-09-10"), "全日客流")])
    before = await provenance(client, run, "a13_bad")
    assert before.status == "inferred" and before.version.parts[0].revoked is None
    [record] = [r for r in (await client.get(f"{API}/{src.id}/imports")).json() if r["id"] == src.imports[1]]
    plan = record["revoke_plan"]
    assert plan["action"] == "rollback" and plan["target_snapshot_id"] == src.snapshots[0]
    reason = "合成理由：接受依据有误"
    r = await client.post(f"{API}/{src.id}/imports/{src.imports[1]}/revoke-acceptance",
                          json={"confirm": True, "expected_current_snapshot_id": src.snapshot,
                                "expected_target_snapshot_id": plan["target_snapshot_id"], "reason": reason,
                                "signed_by": SIGNER})
    assert r.status_code == 200, r.text
    out = await provenance(client, run, "a13_bad")
    await assert_drilled(out, run, "a13_bad", table="日客流", column="全日客流", role="measure",
                         pk={"日期": "2026-09-10"}, sheet=SHEET, cell="L5", part_seq=2)
    [part] = out.version.parts
    assert part.revoked is not None and (part.revoked.reason, part.revoked.signed_by) == (reason, SIGNER)
    assert isinstance(part.revoked.at, str) and part.revoked.at
    assert out.version.parts == await d26_parts(src, has_row=True, revoked=part.revoked)
    assert check_of(out, "R1").acceptance == part.acceptances[0]


async def test_a15_a_masked_column_is_not_drilled(world, fresh_engines, client):
    """A15：数据源设 mask_columns="分区甲"，引用分区甲那一格：table_only · masked（行标签原文、主键值都是格子里的
    内容，展示出来就绕过了遮罩，调整 4）；数据版本里不给「查看导入清单」（manifest_view false）。"""
    w = world
    src = w.sources["masked"]
    out = await provenance(client, w.main, "a15")
    assert_table_only(out, "masked")
    assert out.version == VersionView(
        source=src.name, snapshot_id=src.snapshot, mode="replace", union=False, tables=DAILY, manifest_view=False,
        parts=[await expected_part(src, 0, has_row=False, period=august(), region=f"{SHEET}!B4:AG30")])
    assert (await query_step(client, w.main, "a15"))["masked"] == ["分区甲"]


async def test_a15_a_masked_primary_key_recorded_in_the_snapshot_is_enough(world, fresh_engines, client):
    """A15 补充（列表源）：只遮主键「地区」、引用「金额」也不下钻——否则遮罩的主键值会经 cell_source.pk 和
    recheck.params 交给界面。列表的「地区」角色是 text（不是维度、派生、常量），附带的格只有 column 为空的列表头，
    拦它的只有 R10 的「主键列」这一项。遮罩在运行时设、运行之后从数据源上清掉：只剩查询快照记下的 mask_columns，
    也要拦住（遮罩 = 数据源现在设的加上快照里记的，R10）。"""
    w = world
    src = w.sources["list_masked"]
    snap = query_snapshot(w.main, "a15_pk")
    assert snap["mask_columns"] == ["地区"]
    schema = artifact_store.load(snap["schema_artifact"])
    assert schema["tables"]["月报"]["primary_key"] == ["地区"]
    receipt = artifact_store.load(schema["import_manifests"][0])["receipt"]
    assert {c["name"]: c["role"] for t in receipt["tables"] if t["name"] == "月报" for c in t["columns"]}["地区"] == "text"
    async with SessionLocal() as session:
        assert "mask_columns" not in ((await session.get(DataSource, src.id)).options or {})
    out = await provenance(client, w.main, "a15_pk")
    assert_table_only(out, "masked")
    assert out.version == VersionView(
        source=src.name, snapshot_id=src.snapshot, mode="replace", union=False,
        tables=[TableRef(name="月报", kind="data")], manifest_view=False,
        parts=[await expected_part(src, 0, has_row=False, period=None, region="月报!C5:F9")])
    assert (await query_step(client, w.main, "a15_pk"))["masked"] == ["地区"]


async def test_a16_missing_primary_key_columns(world, fresh_engines, client):
    """A16：`SELECT "全日客流" …` 没有带上主键列：table_only · pk_missing，原文里写出没带齐的主键列「日期」。"""
    out = await provenance(client, world.main, "a16")
    assert_table_only(out, "pk_missing", text=refuse("pk_missing", 列=["日期"]).text)
    assert "（日期）" in out.reason.text
    await flow_version(world, out, DAILY)


async def test_a17_a_human_entered_period(world, fresh_engines, client):
    """A17：统计期人工录入（noperiod，署名「录入员甲」）：年份取自人工录入的统计期 year.source human、signed_by
    录入员甲；数据版本这一期的 period.source human、没有格子。人工录入的统计期没有 C1，只有 C2。"""
    w = world
    src = w.sources["human"]
    raw, _ = src.raws[1]
    out = await provenance(client, w.main, "a17")
    await assert_drilled(out, w.main, "a17", table="日客流", column="全日客流", role="measure",
                         pk={"日期": "2026-09-04"}, sheet=SHEET, cell="F5", part_seq=2)
    assert out.cell_source.year == YearSource(source="human", cells=[], signed_by=SIGNER, mixed=False)
    assert froms(out) == [axis("F"), label("B5", "全日客流（人次）", None)]
    assert out.version.parts == [await expected_part(src, 1, has_row=True, period=september("human"),
                                                     region=f"{SHEET}!B4:AF30")]
    assert [x.id for x in out.checks] == ["R1", "R2", "C2"]
    assert xl(raw, "F5") == result_value(w.main, "a17") and xl(raw, "B2") is None


async def test_a18_a_dimension_cell_is_its_row_label(world, fresh_engines, client):
    """A18：引用长表的维度列（时段，D24w 全角标签的 9 月，累积并集）：格子就是行标签格 B11；canonical.raw 是原表的
    「８－９」、canonical 是库里的「8-9」；附带的格是日期表头 N4 和分段标题 B9（被引用的就是行标签，不再重复）。"""
    w = world
    sep_raw, _ = w.sources["acc"].raws[1]
    out = await provenance(client, w.main, "a18")
    await assert_drilled(out, w.main, "a18", table="时段客流", column="时段", role="dim",
                         pk={"日期": "2026-09-12", "时段": "8-9"}, sheet=SHEET, cell="B11", part_seq=2)
    assert out.cell_source.canonical == CanonicalView(raw="８－９", canonical="8-9")
    assert froms(out) == [axis("N"), title("B9")]
    assert xl(sep_raw, "B11") == "８－９" and result_value(w.main, "a18") == "8-9"
    await acc_version(w, out, 2)


async def test_a19_a_list_cell(world, fresh_engines, client):
    """A19：列表形态（lab.c01，工作表「月报」）里分区丙的金额：F8（openpyxl 对照）；附带列表头 F5（text 是表头原文，
    column 为 null）；被引用的格不在合并区域里，merged_fill false；列表没有年份。"""
    w = world
    src = w.sources["list"]
    raw, _ = src.raws[0]
    out = await provenance(client, w.main, "a19")
    await assert_drilled(out, w.main, "a19", table="月报", column="金额", role="measure",
                         pk={"地区": "分区丙"}, sheet="月报", cell="F8", part_seq=1)
    assert froms(out) == [("col_header", None, "月报", "F5", "金额", None)]
    assert out.cell_source.year is None and out.cell_source.merged_fill is False
    assert xl(raw, "F8", "月报") == result_value(w.main, "a19") and xl(raw, "F5", "月报") == "金额"
    assert out.version == VersionView(
        source=src.name, snapshot_id=src.snapshot, mode="replace", union=False,
        tables=[TableRef(name="月报", kind="data")], manifest_view=True,
        parts=[await expected_part(src, 0, has_row=True, period=None, region="月报!C5:F9")])


def lines_of(excerpts: dict[str, str], run: RunInfo, key: str) -> list[str]:
    return excerpt_new_lines(excerpts[run.aliases[key]])


async def report_excerpts(run: RunInfo, units: list[str] | None = None) -> Any:
    """报告节点的取法（nodes/report.py）：遮罩按目录里数据源现在的设置查，loader 用工件库。"""
    from app.engine import judge as judging

    catalog = run.doc["catalog"]
    masked = await judging.source_masks(catalog)
    return judging.prepare(run.doc, catalog, units=units, loader=artifact_store.load, masked=masked)


async def test_a20_judge_excerpts_carry_provenance_lines(world, fresh_engines):
    """A20：新文档、配方源的裁判摘录（judge.prepare，报告节点的取法）。

    ① 摘录在 SQL 与「列：」之间多出 3.2 的来历行：D26 那条查询有「来源：上传的表格」、文件名、区域、「口径：表 日客流
       的说明」、「已接受（用户填写的理由，未经系统核对）：」和理由原文；D28 那条有带限定语的「区域外文字（…）：」小标题
       和 B32 的全文（数字照录）及所属的导入和统计期；人工录入统计期那条写「人工录入，署名「录入员甲」，未认证」；
    ② 第 32 行隐藏后再导入的变体：没有「闸机故障」，只有「另有 1 处区域外文字未列出…」；
    ③ 给 D28、D26 两个数据源设遮罩后：没有 B32 的全文、没有接受理由原文，只写条数和「理由未列出」；
    ④ 同样的输入调两次，摘录逐字相同；
    ⑤ 估价（JudgeRequest.cost，模型 claude-sonnet-5）：只判一句、引用一条 D28 查询 ≤ 0.01 美元；15 句、引用同一
       数据版本 5 条查询 ≤ 0.05 美元。"""
    w = world
    run = w.main
    d26, d28 = w.sources["d26"], w.sources["d28"]
    request = await report_excerpts(run)
    ex = request.excerpts
    # ① D26：来源、口径、已接受
    new = lines_of(ex, run, "a13_bad")
    mid = await manifest_id_of(d26.imports[1])
    m = artifact_store.load(mid)
    assert new and new[0].startswith("来源：上传的表格，每期替换，第 2 次导入：统计期 2026-09-01 至 2026-09-30"), new
    assert f"文件「{d26.raws[1][1]}」（sha256 {m['file']['raw_sha256'][:12]}）" in new[0]
    assert f"区域 {SHEET}!B4:AF30" in new[0]
    schema = artifact_store.load(query_snapshot(run, "a13_bad")["schema_artifact"])
    assert f"口径：表 日客流 的说明：{schema['tables']['日客流']['comment']}" in new
    title = manifest_acceptance(mid, "R1").title
    assert ACCEPT_HEAD in new
    assert f"  第 2 次导入（统计期 2026-09-01 至 2026-09-30）的核对「{title}」不成立，理由：{REASON_D26}（未署名）" in new
    # ① D28：区域外文字的全文、所属的导入和统计期（数字照录）
    new = lines_of(ex, run, "d28")
    assert OUTSIDE_HEAD in new
    assert f"  第 1 次导入（统计期 2026-09-01 至 2026-09-30）{SHEET}!B32：「{D28_NOTE}」" in new
    # ① 人工录入的统计期
    new = lines_of(ex, run, "a17")
    assert new and "统计期 2026-09-01 至 2026-09-30（人工录入，署名「录入员甲」，未认证）" in new[0], new
    # ② 隐藏行里的区域外文字不送全文，只计条数
    new = lines_of(ex, run, "d28h")
    assert not any("闸机故障" in line for line in new), new
    assert "另有 1 处区域外文字未列出（所在行列被隐藏，或导入时未记录是否隐藏），见证据面板" in [x.strip() for x in new]
    # ④ 确定：同样的输入两次逐字相同
    assert (await report_excerpts(run)).excerpts == ex
    # ③ 设遮罩：理由、区域外文字一概不写原文
    try:
        await set_masks(d28.id, "分区甲")
        await set_masks(d26.id, "分区甲")
        masked = (await report_excerpts(run)).excerpts
    finally:
        await set_masks(d28.id, None)
        await set_masks(d26.id, None)
    new = lines_of(masked, run, "d28")
    assert not any(D28_NOTE in line or "客流汇总表" in line for line in new), new
    assert "区域外另有 2 处文字（数据源设置了遮罩，未列出），见证据面板" in new
    new = lines_of(masked, run, "a13_bad")
    assert not any(REASON_D26 in line for line in new), new
    assert any(line.endswith("不成立（数据源设置了遮罩，理由未列出）") for line in new), new
    # ⑤ 估价
    one = await report_excerpts(run, [run.units["d28"]])
    assert [c.unit for c in one.cands] == [run.units["d28"]] and lines_of(one.excerpts, run, "d28")
    one_usd = one.cost("claude-sonnet-5", one.cands)
    fifteen_units = [run.units[k] for k in run.units if k.startswith("cost_")]
    assert len(fifteen_units) == 15 and len({run.aliases[k] for k in run.units if k.startswith("cost_")}) == 5
    many = await report_excerpts(run, fifteen_units)
    assert len(many.cands) == 15 and all(lines_of(many.excerpts, run, k) for k in run.units if k.startswith("cost_"))
    many_usd = many.cost("claude-sonnet-5", many.cands)
    print(f"\n[A20] 估价：一句一条查询 {one_usd:.4f} 美元；15 句 5 条查询 {many_usd:.4f} 美元")
    assert one_usd is not None and one_usd <= 0.01
    assert many_usd is not None and many_usd <= 0.05


async def test_a21_on_demand_judge_reaches_the_chain_only_for_marked_documents(world, fresh_engines, engine_up,
                                                                               client, judge_spy):
    """A21：按需裁判的取证范围（探索运行）。带标记的文档：`_sealed_loader(sealed, chain=True)` 取得到封存范围内的
    表结构快照经哈希链列出的导入清单；一份真实存在、但不在这条链上的清单（D26 源被替换掉的 8 月那一期）取不到。
    不带标记（chain=False）：链上的清单也取不到，和期 4 之前相同。接口一侧：按需裁判送给裁判的摘录，带标记的文档
    有来源行，期 4 之前的文档没有。"""
    from app.api import evidence as api

    w = world
    sealed = await api._sealed(w.main.id)
    assert sealed.trusted
    chained: set[str] = set()
    for sid in sealed.schemas:
        schema = artifact_store.load(sid)
        if isinstance(schema, dict) and schema.get("import_mode") == "recipe":
            chained |= set(schema.get("import_manifests") or [])
            if schema.get("snapshot_manifest"):
                chained.add(schema["snapshot_manifest"])
    assert len(chained) >= 8
    stray = await manifest_id_of(w.sources["d26"].imports[0])
    assert stray not in chained and isinstance(artifact_store.load(stray), dict)
    with_chain = api._sealed_loader(sealed, chain=True)
    without = api._sealed_loader(sealed, chain=False)
    for mid in sorted(chained):
        assert with_chain(mid) == artifact_store.load(mid), mid
        assert without(mid) is None, mid
    assert with_chain(stray) is None and without(stray) is None

    r = await client.post(f"/api/runs/{w.legacy.id}/evidence/judge", json={"units": [w.legacy.units["a6"]]})
    assert r.status_code == 200, r.text
    assert excerpt_new_lines(judge_spy[-1][w.legacy.aliases["a6"]]) == []
    r = await client.post(f"/api/runs/{w.main.id}/evidence/judge", json={"units": [w.main.units["a13_bad"]]})
    assert r.status_code == 200, r.text
    new = excerpt_new_lines(judge_spy[-1][w.main.aliases["a13_bad"]])
    assert new and new[0].startswith("来源：上传的表格") and any(REASON_D26 in line for line in new), new


async def test_a22_a_run_waiting_for_approval_is_drilled_and_marked_unsealed(world, fresh_engines, client):
    """A22：报告之后接一个人工审批节点，运行停在审批卡上（未封存）：照常推断（审批人正需要核对报告里的数），
    sealed false，没有 alert（不会被误判成清单无法读取）；查询步骤带 provenance: true。"""
    w = world
    async with SessionLocal() as session:
        assert (await session.get(Run, w.gate.id)).status == "interrupted"
    out = await provenance(client, w.gate, "a22")
    await assert_drilled(out, w.gate, "a22", table="日客流", column="全日客流", role="measure",
                         pk={"日期": "2026-08-05"}, sheet=SHEET, cell="G5", part_seq=1, sealed=False)
    step = await query_step(client, w.gate, "a22")
    assert step["provenance"] is True and step["sealed"] is False


async def test_a23_a_broken_seal_is_drilled_and_marked_unsealed(world, fresh_engines, client):
    """A23：A4 的运行封存之后，直接改库里一条封存范围内的事件（不重算封存哈希）：封存核对不通过，sealed false；
    照常推断，没有 alert（不是 manifest_unreadable）。"""
    w = world
    async with SessionLocal() as session:
        started = (await session.execute(select(RunEvent).where(
            RunEvent.run_id == w.broken.id, RunEvent.type == "node.started", RunEvent.node_id == "start"))).scalar_one()
        started.data = {**(started.data or {}), "note": "事后改过"}
        await session.commit()
    graph = (await client.get(f"/api/runs/{w.broken.id}/evidence")).json()
    assert graph["seal"]["sealed"] is True and graph["seal"]["ok"] is False
    out = await provenance(client, w.broken, "a23")
    await assert_drilled(out, w.broken, "a23", table="日客流", column="全日客流", role="measure",
                         pk={"日期": "2026-08-05"}, sheet=SHEET, cell="G5", part_seq=1, sealed=False)


@pytest.mark.parametrize("variant", sorted(A24_SQL))
async def test_a24_hidden_compound_queries_are_refused(world, fresh_engines, client, monkeypatch, variant):
    """A24：藏起来的集合运算（实测 B7 的 `ESCAPE '\\'`、`AS [t']` 两种写法，加上 WP-A 补的记号开头的 U+FEFF）：
    结果列名是「全日客流」，r0 的值实际取自分区甲；夹具让 2026-08-03 两列的值恰好相等，只比值的回查分不出来。
    必须由 R7 认出：table_only · multi_table，细分 compound；不给格子。

    R15 的兜底（EXPLAIN 里 UNION ALL 有两个 ResultRow）给出的结论和 R7 一字不差，只看接口结论分不出是哪一道
    拦下的：R7 失效时 R15 照样给 compound，这条用例照样通过（评审实测：R7 整个拿掉，单元测试 23 条失败、这里
    全过）。所以把 provenance_db.compile_facts 换成放行的桩（只读了本表、ResultRow 1、没有 Yield）——这时能给出
    compound 的只剩 R7，R7 一漏就会一路下钻到 R17，给出分区甲那一格却说是全日客流（回查的值恰好相等，也拦不住）。
    另外两条断言：桩一次也没被调用（R7 在打开数据文件之前就拒绝了）；对照组 A4 在同一个桩下照常推断、桩被调用过，
    证明桩确实接在接口用的那个名字上，不是换了个没人调的函数。"""
    from app.data import provenance_db

    w = world
    key = f"a24_{variant}"
    raw, _ = w.sources["flow"].raws[0]
    snap = query_snapshot(w.main, key)
    assert snap["columns"] == ["日期", "全日客流"] and snap["rows"] == [["2026-08-03", xl(raw, "E6")]]
    assert xl(raw, "E5") == xl(raw, "E6") and xl(raw, "E7") == 0            # 陷阱：两列的值恰好相等

    calls: list[str] = []

    async def lenient(view: Any, sql: str) -> CompileFacts:
        calls.append(sql)
        return CompileFacts(tables={"日客流"}, result_rows=1, yields=0)

    patch_everywhere(monkeypatch, provenance_db, "compile_facts", lenient)
    out = await provenance(client, w.main, key)
    assert_table_only(out, "multi_table", detail="compound")
    assert calls == []                                                     # R7 拒绝在先，没走到 R15
    await flow_version(w, out, DAILY)

    control = await provenance(client, w.main, "a4")
    assert control.status == "inferred" and control.cell_source is not None
    assert calls and set(calls) == {query_snapshot(w.main, "a4")["sql"]}


def patch_everywhere(monkeypatch, module: Any, name: str, fake: Any) -> None:
    """把 module.name 换成 fake；接口模块里按名字导入了同一个函数对象的，一并换掉（不依赖 WP-C 的导入写法）。"""
    from app.api import evidence as api

    real = getattr(module, name)
    monkeypatch.setattr(module, name, fake)
    for attr, obj in list(vars(api).items()):
        if obj is real:
            monkeypatch.setattr(api, attr, fake)


async def test_a25_a_recheck_mismatch_forces_a_full_rehash(own, client, monkeypatch):
    """A25：回查不一致时对数据文件强制重算一次整份哈希（R16，不走 EngineCache 的指纹捷径）。用桩让
    provenance_db.recheck 返回与查询快照不同的值：
    ① 数据文件没动：哈希相等 → table_only · recheck_mismatch，不标红（文件确实没变，是 SQL 的含义和判据认定的不同）；
    ② 数据文件原地改过（同样大小）、再用 os.utime 还原修改时间：指纹 (dev, ino, size, mtime_ns) 不变，EngineCache
       不会重新核对；强制重算发现哈希不等 → table_only · db_tampered，标红。"""
    from app.data import provenance_db

    src, run = await own_flow(own, seed=87)
    assert (await provenance(client, run, "a4")).status == "inferred"
    real = provenance_db.recheck

    async def lying(view: Any, *, table: str, column: str, pk: dict[str, Any]) -> Any:
        rows, sql, params = await real(view, table=table, column=column, pk=pk)
        return [(rowid, value + 1 if isinstance(value, int) else value) for rowid, value in rows], sql, params

    patch_everywhere(monkeypatch, provenance_db, "recheck", lying)
    out = await provenance(client, run, "a4")
    assert_table_only(out, "recheck_mismatch")
    # ②：同样大小的合法库原地写回，还原修改时间
    path = Path((await snapshot_row(src.snapshot)).db_path)
    st = os.stat(path)
    other = path.with_name(path.stem + ".edit.db")
    shutil.copyfile(path, other)
    os.chmod(other, 0o644)
    with closing(sqlite3.connect(other)) as conn:
        conn.execute('UPDATE "日客流" SET "分区甲" = "分区甲" + 1 WHERE "日期" = \'2026-08-21\'')
        conn.commit()
    data = other.read_bytes()
    other.unlink()
    assert len(data) == st.st_size and sha256(data) != sha256(path.read_bytes())
    os.chmod(path, 0o644)
    with open(path, "r+b") as f:
        f.write(data)
        f.truncate()
    os.chmod(path, 0o444)
    os.utime(path, ns=(st.st_atime_ns, st.st_mtime_ns))
    after = os.stat(path)
    assert (after.st_dev, after.st_ino, after.st_size, after.st_mtime_ns) == \
        (st.st_dev, st.st_ino, st.st_size, st.st_mtime_ns)
    out = await provenance(client, run, "a4")
    assert_table_only(out, "db_tampered", alert=True)


async def test_a26_a_pre_p4_document_never_touches_a_broken_manifest(own, client, judge_spy):
    """A26：期 4 之前形状的文档 + 坏清单：把表结构快照列出的导入清单文件改坏一个字节，再走按需裁判（接口），送给
    裁判的摘录与清单完好时逐字相同：老文档的 `_sealed_loader` chain 为假，根本不去取清单。"""
    src, run = await own_flow(own, seed=88, legacy=True)
    assert "provenance" not in run.doc
    unit, alias = run.units["a4"], run.aliases["a4"]
    r = await client.post(f"/api/runs/{run.id}/evidence/judge", json={"units": [unit]})
    assert r.status_code == 200, r.text
    intact = judge_spy[-1]
    assert alias in intact and excerpt_new_lines(intact[alias]) == []
    schema = artifact_store.load(query_snapshot(run, "a4")["schema_artifact"])
    [manifest] = schema["import_manifests"]
    path = artifact_store._path_of(manifest)
    raw = path.read_bytes()
    at = raw.index(b'"seq":')
    path.write_bytes(raw[:at] + b'"seQ":' + raw[at + 6:])
    with pytest.raises(ValueError):
        artifact_store.load(manifest)
    r = await client.post(f"/api/runs/{run.id}/evidence/judge", json={"units": [unit]})
    assert r.status_code == 200, r.text
    assert len(judge_spy) == 2 and judge_spy[-1] == intact


# ==========================================================================
# 6.5 性能（RECIPE_PERF=1 时才跑）：端到端耗时，经 AsyncClient 调推断来源接口
# ==========================================================================

perf = pytest.mark.skipif(os.environ.get("RECIPE_PERF") != "1",
                          reason="性能测试默认跳过：RECIPE_PERF=1 时运行（12 期累积、20 万行列表要造一会儿）")
#: 热路径取几次的中位数
PERF_SAMPLES = 7


async def _median_get(c: AsyncClient, url: str) -> float:
    took = []
    for _ in range(PERF_SAMPLES):
        t0 = time.perf_counter()
        r = await c.get(url)
        took.append(time.perf_counter() - t0)
        assert r.status_code == 200, r.text
    return sorted(took)[len(took) // 2]


def _median_chain(run: RunInfo, key: str) -> float:
    """清单取回的耗时（2.9 的开销表）：每次一个新的工件备忘，resolve_chain 逐份取回清单、复验哈希、核对链。
    接口里同样是每次请求一个新的备忘，所以这就是热路径里清单取回那一段。"""
    from app.api.evidence import _Artifacts
    from app.data.provenance import resolve_chain

    query = query_snapshot(run, key)
    schema = artifact_store.load(query["schema_artifact"])
    took = []
    for _ in range(PERF_SAMPLES):
        t0 = time.perf_counter()
        resolve_chain(query, schema, _Artifacts().load)
        took.append(time.perf_counter() - t0)
    return sorted(took)[len(took) // 2]


async def _measure(c: AsyncClient, run: RunInfo, key: str, src: Source, label: str, target: float) -> None:
    """首次请求（引擎缓存清空、整份文件哈希未校验）单独记、不设断言；热路径取中位数，超过目标的 1.5 倍算不通过。
    同一格的片段接口耗时作为固定开销（_sealed、verify_manifest、报告文档复验）、清单取回的耗时一并打印，拆出占比。"""
    seg = run.segs[key]["id"]
    url = f"/api/runs/{run.id}/evidence/segments/{seg}/provenance"
    await engines.invalidate(src.id)
    t0 = time.perf_counter()
    first = await c.get(url)
    cold = time.perf_counter() - t0
    assert first.status_code == 200, first.text
    out = from_json(first.json())
    assert contract_problems(out) == [] and out.status == "inferred", (out.status, out.reason)
    hot = await _median_get(c, url)
    fixed = await _median_get(c, f"/api/runs/{run.id}/evidence/segments/{seg}")
    chain = _median_chain(run, key)
    print(f"\n[perf] {label}：首次 {cold * 1000:.0f} ms；热路径 {hot * 1000:.0f} ms（目标 {target * 1000:.0f} ms）；"
          f"同一格的片段接口 {fixed * 1000:.0f} ms（固定开销占 {fixed / hot:.0%}）；"
          f"清单取回 {chain * 1000:.1f} ms（占 {chain / hot:.0%}）")
    assert hot <= target * FAIL_FACTOR, f"{label} 热路径 {hot * 1000:.0f} ms 超过目标的 {FAIL_FACTOR} 倍"


@perf
async def test_perf_d00_single_period(own, client):
    src, run = await own_flow(own, seed=91)
    await _measure(client, run, "a4", src, "D00 单期", 0.150)


@perf
async def test_perf_two_period_union(own, client):
    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as c:
        src = await first_import(c, *flow_workbook(AUG, 31, seed=92), mode="accumulate")
        await add_period(c, src, *flow_workbook(SEP, 30, seed=93))
    own.append(src)
    run = await run_report({"acc": src}, [Cite("a7", "acc", long_sql("2026-09-12", "8-9"), "客流")])
    await _measure(client, run, "a7", src, "8 月加 9 月并集", 0.200)


@perf
async def test_perf_twelve_period_accumulation(own, client):
    """12 期按期累积（flow_workbook 按月造 12 份，逐期导入）：清单取回随期数线性增长，这一条盯它。"""
    import calendar
    import datetime as dt

    months = [(dt.date(2025, m, 1), calendar.monthrange(2025, m)[1]) for m in range(1, 13)]
    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as c:
        src = await first_import(c, *flow_workbook(months[0][0], months[0][1], seed=100), mode="accumulate")
        for i, (start, days) in enumerate(months[1:], start=1):
            await add_period(c, src, *flow_workbook(start, days, seed=100 + i))
    own.append(src)
    run = await run_report({"acc": src}, [Cite("p12", "acc", wide_sql("2025-12-15"), "全日客流")])
    out = await provenance(client, run, "p12")
    assert len(out.version.parts) == 12 and out.version.parts[-1].has_row
    await _measure(client, run, "p12", src, "12 期按期累积", 0.300)


def _big_list(rows: int = 200_000) -> bytes:
    """20 万行的单期列表（造法同 test_recipe_perf.big_list，多一列唯一的「编号」作主键：没有主键就只给表级来历，
    量不到完整的下钻）。只写值、没有公式。"""
    import random

    from openpyxl import Workbook

    rnd = random.Random(7)
    wb = Workbook(write_only=True)
    ws = wb.create_sheet("明细")
    ws.append(["编号", "地区", "渠道", "销量", "金额"])
    for i in range(rows):
        ws.append([f"N{i:06d}", f"区{i % 7}", "线上" if i % 2 else "线下", rnd.randint(1, 999), rnd.randint(1, 99999)])
    buf = io.BytesIO()
    wb.save(buf)
    return excel_saved(buf.getvalue(), "明细", {})


@perf
async def test_perf_two_hundred_thousand_row_list(own, client):
    recipe = {"recipe_format": "agentlab-recipe/2",
              "sheets": [{"id": "s1", "match": {"name": "明细"}, "blocks": [{
                  "id": "明细", "layout": "list", "table": "销售明细",
                  "columns": [{"header": h, "name": h, "type": t} for h, t in
                              (("编号", "TEXT"), ("地区", "TEXT"), ("渠道", "TEXT"), ("销量", "INTEGER"),
                               ("金额", "INTEGER"))]}]}],
              "tables": [{"name": "销售明细", "grain": ["编号"]}]}
    raw, fn = _big_list(), "明细_20万行.xlsx"
    name = unique("perf")
    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test", timeout=600) as c:
        r = await c.post(f"{API}/imports/stage", files={"file": (fn, raw, XLSX)}, data={"name": name})
        assert r.status_code == 201, r.text
        st = r.json()
        r = await c.put(f"{API}/imports/{st['id']}/recipe", json={"recipe": recipe})
        assert r.status_code == 200 and r.json()["recipe_problems"] == [], r.text
        r = await c.post(f"{API}/imports/{st['id']}/trial", json={})
        st = r.json()
        assert st["trial"]["status"] == "passed", st["trial"]["problems"][:3]
        assert {t["name"]: t["rows"] for t in st["trial"]["receipt"]["tables"]} == {"销售明细": 200_000}
        out = await commit(c, st)
    src = Source(name=name, id=out["source"]["id"], snapshot=out["snapshot_id"], raws=[(raw, fn)],
                 imports=[out["import_id"]])
    own.append(src)
    run = await run_report({"big": src}, [Cite(
        "big", "big", 'SELECT "编号", "金额" FROM "销售明细" WHERE "编号" = \'N123456\'', "金额")])
    await _measure(client, run, "big", src, "20 万行列表单期", 0.300)
