"""期 3 集成验收（WP-8，P3-SPEC 11.2、11.4、12.9）：真管线（default_pipeline）经接口走完整流程。

不用假 Pipeline：暂存、规则起草、回答问题、修复按钮、框选、试运行、累积计划、物化、确认项、提交、版本页的启用 /
移除 / 作废 / 清除，全部是各工作包合并后的真代码；接口一律经 AsyncClient(transport=ASGITransport(app=app))，从暂存
开始，不直接调内部函数拼结果。只有接口不提供的观察点（快照记录、导入记录、构建记录、运行记录、文件）才直接读库、
看文件；造「数据目录被人改过」这类状态时才直接改库或改文件（11.4 的数据变异）。

公共做法（11.2）：D00 8 月暂存，回答问题（系统发现一律登记、占位符存为空值），按期 2 验收的写法把宽表改名「日客流」，
q_mode 选「按期累积」，提交得到 S1。之后各例按「上传新一期」重放。

夹具全部合成、假名（tests/fixtures/xlsx：分区甲、分区乙、分区丙……，数字随机）；不调用任何模型（期 3 没有新的模型
调用，规则起草不完整时的 AI 可用性查询换成「不可用」的桩）。断言照规格写，不放宽、不标 xfail。
"""
from __future__ import annotations

import json
import os
import re
import sqlite3
import subprocess
import sys
import time
import uuid
from contextlib import closing
from pathlib import Path
from typing import Any

import pytest
from httpx import ASGITransport, AsyncClient
from openpyxl import Workbook
from sqlalchemy import select, update

from app.core import artifact_store
from app.core.config import settings
from app.data import raw_store, table_versions
from app.db.base import SessionLocal
from app.db.models import (
    DataSource, ImportStaging, Run, SnapshotActivation, SourceSnapshot, TableBuild, TableImport, TableRecipe,
)
from app.main import app
from tests.fixtures.xlsx import excel_saved, lab, save
from tests.fixtures.xlsx.drift import AUG, DRIFT_CASES, OCT, SEP
from tests.fixtures.xlsx.flow import SHEET, flow_workbook

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
API = "/api/datasources"
TABLES = ("日客流", "时段客流", "时段客流_表内合计")
#: 2.12 的验收数字：8 月 31 天、9 月 30 天，并集 61 / 1037 / 183
AUG_ROWS = {"日客流": 31, "时段客流": 527, "时段客流_表内合计": 93}
SEP_ROWS = {"日客流": 30, "时段客流": 510, "时段客流_表内合计": 90}
UNION_ROWS = {"日客流": 61, "时段客流": 1037, "时段客流_表内合计": 183}
REASON = "合成理由：验收用例"


# ==========================================================================
# 夹具
# ==========================================================================


@pytest.fixture(autouse=True)
def store(monkeypatch, tmp_path):
    """每个测试一个独立的数据目录：原件存档、构建库、试运行库、并集、隔离区、工件互不串。"""
    data = tmp_path / "data"
    (data / "uploads").mkdir(parents=True)
    monkeypatch.setattr(settings, "data_dir", data)
    yield data


@pytest.fixture(autouse=True)
def real_pipeline_no_model(monkeypatch):
    """验收必须跑真管线：recipe_imports.pipeline() 接的是各工作包的真函数（含期 3 新增的累积、物化、修复、框选），
    不是别的测试留下的假实现。规则起草不完整时 StagingOut 会查 AI 可用性：换成「不可用」的桩，连模型对象都不建。"""
    from app.data import (
        recipe_accumulate, recipe_ai, recipe_engine, recipe_fixes, recipe_imports, recipe_notes, recipe_select,
        recipe_suggest,
    )
    from app.data.recipe_types import AiAvailability

    p = recipe_imports.pipeline()
    assert p.execute is recipe_engine.execute and p.draft is recipe_suggest.draft
    assert p.ai_availability is recipe_ai.ai_availability
    assert p.plan_accumulate is recipe_accumulate.plan_accumulate
    assert p.materialize_union is recipe_accumulate.materialize_union
    assert p.build_union_notes is recipe_notes.build_union_notes
    assert p.propose_fixes is recipe_fixes.propose_fixes and p.selection_edit is recipe_select.selection_edit

    async def no_model(session):
        return None, "", AiAvailability(False, reason="验收测试不接模型")

    monkeypatch.setattr(recipe_ai, "resolve_draft_model", no_model)


@pytest.fixture
async def client():
    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as c:
        yield c


# ==========================================================================
# 小工具（经 HTTP）
# ==========================================================================


def unique(prefix: str = "wp8") -> str:
    return f"{prefix}_{uuid.uuid4().hex[:10]}"


async def db_get(model, key):
    async with SessionLocal() as session:
        return await session.get(model, key)


async def stage(client, raw: bytes, filename: str, name: str | None = None) -> dict[str, Any]:
    resp = await client.post(f"{API}/imports/stage", files={"file": (filename, raw, XLSX)},
                             data={"name": name or unique()})
    assert resp.status_code == 201, resp.text
    return resp.json()


async def reupload(client, source_id: str, raw: bytes, filename: str) -> dict[str, Any]:
    resp = await client.post(f"{API}/{source_id}/reupload", files={"file": (filename, raw, XLSX)})
    assert resp.status_code == 201, resp.text
    return resp.json()


async def trial(client, sid: str, **body: Any) -> dict[str, Any]:
    resp = await client.post(f"{API}/imports/{sid}/trial", json=body)
    assert resp.status_code == 200, resp.text
    return resp.json()


async def answer(client, sid: str, answers: dict[str, dict[str, Any]]) -> dict[str, Any]:
    resp = await client.post(f"{API}/imports/{sid}/answers", json={"answers": answers})
    assert resp.status_code == 200, resp.text
    return resp.json()


async def put_recipe(client, sid: str, rec: dict[str, Any], **extra: Any) -> dict[str, Any]:
    resp = await client.put(f"{API}/imports/{sid}/recipe", json={"recipe": rec, **extra})
    assert resp.status_code == 200, resp.text
    return resp.json()


def confirm_ids(st: dict[str, Any]) -> list[str]:
    return [i["id"] for i in (st.get("trial") or {}).get("confirm_items") or []]


async def commit(client, st: dict[str, Any], *, confirmations: list[str] | None = None,
                 acceptances: list[dict[str, str]] | None = None):
    """提交；返回原始响应（调用方自己判状态码）。confirmations 缺省 = 勾上全部确认项。"""
    body: dict[str, Any] = {"trial_id": st["trial"]["trial_id"],
                            "confirmations": confirm_ids(st) if confirmations is None else confirmations}
    if acceptances is not None:
        body["acceptances"] = acceptances
    return await client.post(f"{API}/imports/{st['id']}/commit", json=body)


def rows_of(st: dict[str, Any]) -> dict[str, int]:
    return {t["name"]: t["rows"] for t in st["trial"]["receipt"]["tables"]}


def problems(st: dict[str, Any], code: str | None = None) -> list[dict[str, Any]]:
    return [p for p in (st.get("trial") or {}).get("problems") or [] if code is None or p["code"] == code]


def check(st: dict[str, Any], cid: str) -> dict[str, Any] | None:
    return next((c for c in st["trial"]["checks"] if c["id"] == cid), None)


def all_comments(notes: dict[str, Any] | None) -> str:
    out = []
    for n in (notes or {}).values():
        out.append(n.get("comment") or "")
        out += list((n.get("columns") or {}).values())
    return "\n".join(out)


def rename_wide(rec: dict[str, Any], old: str, new: str) -> dict[str, Any]:
    """把规则起草的宽表名改成 new：表名、写入它的分段（id 一起改）、关系里引用的表名（期 2 验收 9.7 第 1 步）。"""
    rec = json.loads(json.dumps(rec, ensure_ascii=False))
    for sheet in rec["sheets"]:
        for block in sheet["blocks"]:
            for seg in block.get("segments", []):
                if seg.get("table") == old:
                    seg["table"] = new
                if seg.get("id") == old:
                    seg["id"] = new
    for t in rec["tables"]:
        if t["name"] == old:
            t["name"] = new
    for r in rec["relations"]:
        if r.get("table") == old:
            r["table"] = new
        for side in ("a", "b"):
            if isinstance(r.get(side), dict) and r[side].get("table") == old:
                r[side]["table"] = new
    return rec


def answers_for(questions: list[dict[str, Any]], mode: str) -> dict[str, dict[str, Any]]:
    """系统发现一律「登记」，q_mode 按 mode，其余能存为空值的选「存为空值」（参考配方的写法）。"""
    out = {}
    for q in questions:
        values = [o["value"] for o in q["options"]]
        if q["id"].startswith("q_relation"):
            out[q["id"]] = {"value": "register"}
        elif q["id"] == "q_mode":
            out[q["id"]] = {"value": mode}
        elif "null" in values:
            out[q["id"]] = {"value": "null"}
        else:
            out[q["id"]] = {"value": values[0]}
    return out


async def first_import(client, *, mode: str = "accumulate", seed: int = 0, start=AUG, days: int = 31,
                       variant: str = "D00", name: str | None = None, raw: bytes | None = None,
                       edit: Any = None) -> dict[str, Any]:
    """公共做法：暂存 → 回答问题（q_mode=mode）→ 宽表改名「日客流」→ 试运行 → 勾全部确认项提交，得到 S1。
    raw 给了就用它（文件名照 flow_workbook 的写法）；edit 给了就在改名之后再改一次配方（A18 的列改名）。"""
    name = name or unique()
    made, fn = flow_workbook(start, days, seed=seed, variant=variant)
    raw = raw if raw is not None else made
    st = await stage(client, raw, fn, name)
    assert st["questions"][-1]["id"] == "q_mode", "q_mode 固定排在问题列表最后（2.2）"
    st = await answer(client, st["id"], answers_for(st["questions"], mode))
    wide = next(t["name"] for t in st["recipe"]["tables"] if t["name"] not in TABLES)
    rec = rename_wide(st["recipe"], wide, "日客流")
    st = await put_recipe(client, st["id"], edit(rec) if edit is not None else rec)
    st = await trial(client, st["id"])
    assert st["trial"]["status"] == "passed", st["trial"]["problems"]
    resp = await commit(client, st)
    assert resp.status_code == 201, resp.text
    out = resp.json()
    return {"name": name, "source_id": out["source"]["id"], "s1": out["snapshot_id"], "import_id": out["import_id"],
            "recipe_id": out["recipe_id"], "commit": out, "trial": st, "raw": raw, "file_name": fn}


async def add_period(client, source_id: str, raw: bytes, fn: str, *,
                     acceptances: list[dict[str, str]] | None = None) -> dict[str, Any]:
    """上传新一期 → 试运行通过 → 勾全部确认项提交。返回提交响应。"""
    st = await reupload(client, source_id, raw, fn)
    assert st["trial"]["status"] in ("passed", "needs_decision"), st["trial"]["problems"]
    resp = await commit(client, st, acceptances=acceptances)
    assert resp.status_code == 201, resp.text
    return resp.json()


async def august_september(client, *, seed: int = 0, name: str | None = None) -> dict[str, Any]:
    """S1=[8 月]、S2=[8 月, 9 月]，同一个配方（按期累积）。"""
    base = await first_import(client, seed=seed, name=name)
    raw, fn = flow_workbook(SEP, 30, seed=seed + 1)
    out = await add_period(client, base["source_id"], raw, fn)
    return {**base, "s2": out["snapshot_id"], "aug": base["import_id"], "sep": out["import_id"], "sep_commit": out,
            "sep_raw": raw, "sep_file": fn}


async def tool_rows(client, source_name: str, sql: str) -> list[list[Any]]:
    """经工具接口查当前版本（模型看到的那一份）。"""
    resp = await client.post(f"/api/tools/db_query__{source_name}/run", json={"args": {"sql": sql}})
    assert resp.status_code == 200, resp.text
    out = resp.json()
    assert out["ok"], out
    return json.loads(out["result"])["rows"]


async def count(client, name: str, table: str, where: str = "") -> int:
    return int((await tool_rows(client, name, f'SELECT COUNT(*) FROM "{table}"{where}'))[0][0])


async def by_month(client, name: str, table: str, where: str = "") -> list[list[Any]]:
    return await tool_rows(client, name, f'SELECT substr("日期", 1, 7), COUNT(*) FROM "{table}"{where} '
                                         "GROUP BY 1 ORDER BY 1")


async def current_id(source_id: str) -> str | None:
    return (await db_get(DataSource, source_id)).current_snapshot_id


async def current_periods(source_id: str) -> list[str]:
    snap = await db_get(SourceSnapshot, await current_id(source_id))
    return [(await db_get(TableImport, i)).period_start for i in snap.imports]


async def snapshots(client, source_id: str) -> list[dict[str, Any]]:
    resp = await client.get(f"{API}/{source_id}/snapshots")
    assert resp.status_code == 200, resp.text
    return resp.json()


async def imports_api(client, source_id: str) -> list[dict[str, Any]]:
    resp = await client.get(f"{API}/{source_id}/imports")
    assert resp.status_code == 200, resp.text
    return resp.json()


def date_of(y: int, m: int, d: int):
    import datetime as dt

    return dt.date(y, m, d)


def sql_file(path: str | Path, query: str, params: tuple[Any, ...] = ()) -> list[tuple[Any, ...]]:
    """只读打开库文件查一句（不改文件，库哈希不受影响）。"""
    with closing(sqlite3.connect(f"file:{path}?mode=ro", uri=True)) as conn:
        return list(conn.execute(query, params))


def columns_of(snap: SourceSnapshot, table: str) -> list[str]:
    return [c["name"] for c in snap.schema_cache["tables"][table]["columns"]]


# --- 改造夹具（保留公式的缓存值） ---


def edit_cells(raw: bytes, edits: dict[str, Any], sheet: str = SHEET) -> bytes:
    """在合成夹具上改几格（值为 None 即清空），其余照旧。openpyxl 存盘会丢掉公式的缓存值，这里先按原文件读出缓存值，
    存盘后再用 excel_saved 填回去：表内合计的公式照样有缓存，核对不受影响。"""
    import io

    from openpyxl import load_workbook

    book = load_workbook(io.BytesIO(raw))
    values = load_workbook(io.BytesIO(raw), data_only=True)
    ws, wv = book[sheet], values[sheet]
    cached = {c.coordinate: wv[c.coordinate].value for row in ws.iter_rows() for c in row
              if isinstance(c.value, str) and c.value.startswith("=")}
    for ref, value in edits.items():
        ws[ref] = value
    return excel_saved(save(book), sheet, cached)


def without_zone_b(raw: bytes) -> bytes:
    """这一期的文件不再有「分区乙（人次）」那一行（第 7 行整行清空）：原表不再登记分区乙。"""
    from openpyxl.utils import get_column_letter

    return edit_cells(raw, {f"{get_column_letter(c)}7": None for c in range(2, 40)})


# --- 修复按钮、框选 ---


def fix_of(st: dict[str, Any], kind: str) -> dict[str, Any]:
    found = [f for f in st["fixes"] if f["kind"] == kind]
    assert found, (kind, [f["kind"] for f in st["fixes"]])
    return found[0]


async def preview(client, sid: str, body: dict[str, Any]):
    return await client.post(f"{API}/imports/{sid}/edits/preview", json=body)


async def apply(client, sid: str, body: dict[str, Any], expected: str | None, **extra: Any):
    return await client.post(f"{API}/imports/{sid}/edits/apply", json={**body, "expected_sha256": expected, **extra})


async def fixed(client, st: dict[str, Any], kind: str, option: str | None = None,
                reason: str | None = None) -> tuple[dict[str, Any], dict[str, Any]]:
    """按提议预览再应用（3.1：预览 → 应用，客户端只发 fix_id、选项值和理由）。返回 (应用后的 StagingOut, 预览)。"""
    fx = fix_of(st, kind)
    opt = next(o for o in fx["options"] if option is None or o["value"] == option)
    if opt["needs_reason"] and reason is None:
        reason = REASON
    body = {"fix": {"id": fx["id"], "option": opt["value"], **({"reason": reason} if reason else {})}}
    resp = await preview(client, st["id"], body)
    assert resp.status_code == 200, resp.text
    pv = resp.json()
    assert pv["ok"], pv["problems"]
    resp = await apply(client, st["id"], body, pv["recipe_sha256_after"])
    assert resp.status_code == 200, resp.text
    return resp.json(), pv


# ==========================================================================
# A1–A3：按期累积、重叠拒收、替换该期
# ==========================================================================


async def test_a1_august_plus_september_accumulates_to_61_1037_183(client):
    """A1：reupload D01 → 试运行 passed，accumulate.action=append，union.rows={61, 1037, 183}，U1–U3 passed →
    提交 201、parts=2 → 当前快照 imports=[8 月, 9 月]、库在 snapshots/ 下，三表 61 / 1037 / 183，按月 31 / 30。"""
    base = await first_import(client, seed=11)
    sid, name = base["source_id"], base["name"]
    raw, fn = flow_workbook(SEP, 30, seed=12)
    st = await reupload(client, sid, raw, fn)
    t = st["trial"]
    assert st["kind"] == "reupload" and t["status"] == "passed", t["problems"]
    assert t["accumulate"]["action"] == "append" and t["accumulate"]["mode"] == "accumulate"
    assert t["accumulate"]["union"]["rows"] == UNION_ROWS
    assert rows_of(st) == SEP_ROWS, "本期的回执是 9 月这一期"
    assert [(c["id"], c["kind"], c["status"]) for c in t["union_checks"]] == [
        ("U1", "union_rows", "passed"), ("U2", "union_pk", "passed"), ("U3", "union_period", "passed")]
    resp = await commit(client, st)
    assert resp.status_code == 201, resp.text
    out = resp.json()
    assert out["parts"] == 2 and out["unchanged"] is False
    snap = await db_get(SourceSnapshot, out["snapshot_id"])
    imps = [await db_get(TableImport, i) for i in snap.imports]
    assert snap.imports == [base["import_id"], out["import_id"]]
    assert [(i.period_start, i.period_end) for i in imps] == [("2026-08-01", "2026-08-31"), ("2026-09-01", "2026-09-30")]
    assert Path(snap.db_path).parent.name == "snapshots" and snap.mode == "accumulate"
    for table, n in UNION_ROWS.items():
        assert await count(client, name, table) == n
        assert sql_file(snap.db_path, f'SELECT COUNT(*) FROM "{table}"')[0][0] == n
    assert await by_month(client, name, "日客流") == [["2026-08", 31], ["2026-09", 30]]


async def test_a2_partially_overlapping_period_is_rejected(client):
    """A2：S2 之后 reupload flow_workbook(2026-08-15, 31) → 试运行 rejected，problems 有 period_overlap，消息里有两段
    已有的统计期；commit 409 trial_not_passed；当前快照不变。"""
    v = await august_september(client, seed=21)
    raw, fn = flow_workbook(AUG.replace(day=15), 31, seed=23)
    st = await reupload(client, v["source_id"], raw, fn)
    t = st["trial"]
    assert t["status"] == "rejected"
    [overlap] = problems(st, "period_overlap")
    assert overlap["category"] == "structure" and overlap["fix"] is None
    assert "2026-08-15 至 2026-09-14" in overlap["message"]
    assert "2026-08-01 至 2026-08-31" in overlap["message"] and "2026-09-01 至 2026-09-30" in overlap["message"]
    resp = await commit(client, st)
    assert resp.status_code == 409 and resp.json()["code"] == "trial_not_passed", resp.text
    assert await current_id(v["source_id"]) == v["s2"]


async def test_a3_replacing_a_period_needs_confirmation(client):
    """A3：reupload 9 月 seed 不同（数字不同、统计期相同）→ passed、action=replace_period；不勾
    period_replace:2026-09-01~2026-09-30 提交 422 confirm_required（detail 里有这个 id）；勾上 201 → 新快照
    imports=[8 月, 9 月新]，行数仍是 61 / 1037 / 183，9 月的值与原来的不同。"""
    v = await august_september(client, seed=31)
    sid, name = v["source_id"], v["name"]
    probe = 'SELECT "全日客流" FROM "日客流" WHERE "日期" = \'2026-09-05\''
    old = (await tool_rows(client, name, probe))[0][0]
    raw, fn = flow_workbook(SEP, 30, seed=99)
    st = await reupload(client, sid, raw, fn)
    t = st["trial"]
    assert t["status"] == "passed", t["problems"]
    assert t["accumulate"]["action"] == "replace_period"
    item = "period_replace:2026-09-01~2026-09-30"
    assert item in confirm_ids(st)
    resp = await commit(client, st, confirmations=[i for i in confirm_ids(st) if i != item])
    assert resp.status_code == 422 and resp.json()["code"] == "confirm_required", resp.text
    assert item in resp.json()["detail"]
    assert await current_id(sid) == v["s2"]
    resp = await commit(client, st)
    assert resp.status_code == 201, resp.text
    out = resp.json()
    snap = await db_get(SourceSnapshot, out["snapshot_id"])
    assert snap.imports == [v["aug"], out["import_id"]] and out["import_id"] != v["sep"]
    for table, n in UNION_ROWS.items():
        assert await count(client, name, table) == n
    assert (await tool_rows(client, name, probe))[0][0] != old
    assert (await db_get(TableImport, v["sep"])).status == "superseded"
    # 回到替换之前的版本：被替换掉的旧 9 月重新是 active，新 9 月置 superseded，数字回到原来的
    resp = await client.post(f"{API}/{sid}/snapshots/{v['s2']}/activate",
                             json={"confirm": True, "expected_current_snapshot_id": out["snapshot_id"]})
    assert resp.status_code == 200, resp.text
    assert (await db_get(TableImport, v["sep"])).status == "active"
    assert (await db_get(TableImport, out["import_id"])).status == "superseded"
    assert (await tool_rows(client, name, probe))[0][0] == old


# ==========================================================================
# A4：固定在 8 月快照上的运行不受影响；被运行引用的并集快照不回收
# ==========================================================================


async def keep_run(source_id: str, name: str, snapshot: str) -> str:
    """登记一条引用这个快照的中断运行（写法同 test_run_data_versions.keep）。返回运行 id。"""
    async with SessionLocal() as session:
        run = Run(workflow_name="验收：保留旧版本", status="interrupted",
                  data_versions={source_id: {"snapshot": snapshot, "name": name}})
        session.add(run)
        await session.commit()
        return run.id


async def tools_for(names: list[str], *, data_versions: Any = None) -> dict[str, Any]:
    from app.tools.datasource import build_datasource_tools
    from app.tools.registry import ToolContext, Versions

    async with SessionLocal() as session:
        built = await build_datasource_tools(
            names, ToolContext(run_id="", node_id="n",
                               data_versions=Versions.CURRENT if data_versions is None else data_versions), session)
    return {t.name: t for t in built}


async def restart_with_unit_change(client, source_id: str, start=OCT, days: int = 31, seed: int = 0) -> dict[str, Any]:
    """这一期的文件把「分区甲（人次）」写成「分区甲（人）」，在这一期的暂存区里用 PUT 把配方的标签和单位一起改掉
    （单位由标签写死，unit_label_conflict 不许单改配方里的单位）：单位变了是不兼容变化 → restart。返回暂存区。"""
    raw, fn = flow_workbook(start, days, seed=seed)
    st = await reupload(client, source_id, edit_cells(raw, {"B6": "分区甲（人）"}), fn)
    assert st["trial"]["status"] == "rejected", "现行配方认不出改了单位的标签"
    rec = json.loads(json.dumps(st["recipe"], ensure_ascii=False))
    seg = next(s for s in rec["sheets"][0]["blocks"][0]["segments"] if s["id"] == "日客流")
    seg["labels"]["expect"] = ["分区甲（人）" if x == "分区甲（人次）" else x for x in seg["labels"]["expect"]]
    seg["measures"] = {("分区甲（人）" if k == "分区甲（人次）" else k): v for k, v in seg["measures"].items()}
    next(t for t in rec["tables"] if t["name"] == "日客流")["units"]["分区甲"] = "人"
    st = await put_recipe(client, st["id"], rec)
    assert st["recipe_problems"] == [], st["recipe_problems"]
    return await trial(client, st["id"])


async def test_a4_runs_pinned_to_august_keep_reading_august(client):
    """A4 前半：在 S1 上建运行记录（data_versions 固定 S1），之后提交 9 月 → 固定 S1 的查询工具查日客流是 31 行，
    冻结的表结构与 S1 的一致；不固定的调用查到 61。"""
    from app.data.engine import engines

    base = await first_import(client, seed=41)
    sid, name, s1 = base["source_id"], base["name"], base["s1"]
    await keep_run(sid, name, s1)
    raw, fn = flow_workbook(SEP, 30, seed=42)
    out = await add_period(client, sid, raw, fn)
    assert out["snapshot_id"] != s1 and out["parts"] == 2
    pins = {sid: {"snapshot": s1, "name": name}}
    try:
        query = (await tools_for([f"db_query__{name}"], data_versions=pins))[f"db_query__{name}"]
        text = await query.coroutine(sql='SELECT COUNT(*) FROM "日客流"')
        payload = json.loads(text)
        assert payload["rows"][0][0] == 31
        assert artifact_store.load(payload["artifact"])["data_version"] == s1
        frozen = artifact_store.load(payload["schema_artifact"])
        s1_row = await db_get(SourceSnapshot, s1)
        assert set(frozen["tables"]) == set(s1_row.schema_cache["tables"]) == set(TABLES)
        assert frozen["tables"]["日客流"]["columns"] == s1_row.schema_cache["tables"]["日客流"]["columns"]
        current = (await tools_for([f"db_query__{name}"]))[f"db_query__{name}"]
        text = await current.coroutine(sql='SELECT COUNT(*) FROM "日客流"')
        assert json.loads(text)["rows"][0][0] == 61
        assert artifact_store.load(json.loads(text)["artifact"])["data_version"] == out["snapshot_id"]
    finally:
        await engines.invalidate(sid)


async def test_a4_a_union_snapshot_referenced_by_a_run_survives_gc(client, monkeypatch):
    """A4 后半（保护规则单独验证，评审一-m8；1.5：KEEP 一律用 0）：在 S2（并集）上建运行记录，再走 restart 得到 S3；
    回收之后 S2 的并集文件仍在、记录没有 retired；删掉那条运行记录再回收，S2 被回收、并集文件删掉。"""
    monkeypatch.setattr(table_versions, "SNAPSHOT_KEEP", 0)
    v = await august_september(client, seed=43)
    sid = v["source_id"]
    run_id = await keep_run(sid, v["name"], v["s2"])
    st = await restart_with_unit_change(client, sid, seed=44)
    assert st["trial"]["accumulate"]["action"] == "restart"
    resp = await commit(client, st)
    assert resp.status_code == 201, resp.text
    s3 = resp.json()["snapshot_id"]
    s2 = await db_get(SourceSnapshot, v["s2"])
    union_file = Path(s2.db_path)
    async with SessionLocal() as session:
        await table_versions.gc(session)
    assert union_file.is_file() and (await db_get(SourceSnapshot, v["s2"])).retired_at is None
    assert (await db_get(SourceSnapshot, v["s1"])).retired_at is not None, "KEEP=0：没人引用的 S1 照期 1 回收"
    async with SessionLocal() as session:
        await session.delete(await session.get(Run, run_id))
        await session.commit()
        await table_versions.gc(session)
    assert (await db_get(SourceSnapshot, v["s2"])).retired_at is not None
    assert not union_file.exists()
    assert await current_id(sid) == s3


# ==========================================================================
# A5：全角或 en-dash 的标签两个月都能查到
# ==========================================================================


@pytest.mark.parametrize("variant", ["D24", "D24w"])
async def test_a5_dash_and_fullwidth_labels_are_queryable_across_months(client, variant):
    """A5：9 月用 en-dash（D24）或全角数字与全角减号（D24w）写时段标签：试运行差异卡有 label_writing、需确认；提交后
    WHERE "时段"='8-9' 是 61，按月 31 / 30（维度值按规范写法存储）。"""
    base = await first_import(client, seed=51)
    sid, name = base["source_id"], base["name"]
    raw, fn = flow_workbook(SEP, 30, seed=52, variant=variant)
    st = await reupload(client, sid, raw, fn)
    t = st["trial"]
    assert t["status"] == "passed", t["problems"]
    writing = [d for d in t["diff"] if d["kind"] == "label_writing"]
    assert writing and all(d["requires_confirm"] for d in writing), t["diff"]
    assert {"diff:label_writing:日间", "diff:label_writing:夜间"} <= set(confirm_ids(st))
    assert t["accumulate"]["action"] == "append" and t["accumulate"]["union"]["rows"] == UNION_ROWS
    resp = await commit(client, st)
    assert resp.status_code == 201, resp.text
    where = ' WHERE "时段" = \'8-9\''
    assert await count(client, name, "时段客流", where) == 61
    assert await by_month(client, name, "时段客流", where) == [["2026-08", 31], ["2026-09", 30]]
    labels = {r[0] for r in await tool_rows(client, name, 'SELECT DISTINCT "时段" FROM "时段客流"')}
    assert all(re.fullmatch(r"\d{1,2}-\d{1,2}", x) for x in labels), labels


# ==========================================================================
# A6、A6b、A6c、A15：拒收的漂移走修复按钮后能试运行并启用（3.4）
# ==========================================================================

#: 3.4 表与「其他两例」：拒收的问题、fix、坐标（None = 无坐标）、修复步骤 [(kind, 选项)]、修复后要勾的确认项、
#: 累积后的行数
FIX_CASES: dict[str, dict[str, Any]] = {
    "D04": {"code": "label_missing", "fix": "remove_label", "cell": f"{SHEET}!B10",
            "steps": [("remove_label", "remove")], "confirms": {"fix:remove_label:日间:7-8"},
            "rows": {"日客流": 61, "时段客流": 1007, "时段客流_表内合计": 183}},
    "D09": {"code": "title_not_found", "fix": "rename_title", "cell": None,
            "steps": [("rename_title", "|keep")], "confirms": {"fix:rename_title:日间"}, "rows": UNION_ROWS},
    "D13": {"code": "label_unexpected", "fix": "add_label", "cell": f"{SHEET}!B8",
            "steps": [("add_label", "add"), ("edit_members", "update")],
            "confirms": {"unit:日客流.分区丙", "relation:R1"}, "fix_items": 2, "rows": UNION_ROWS},
    "D18": {"code": "label_unparsed", "fix": "declare_total", "cell": f"{SHEET}!B21",
            "steps": [("declare_total", "keep")], "confirms": {"fix:declare_total:日间", "derived:日间合计"},
            "accept": "K1", "rows": {"日客流": 61, "时段客流": 1037, "时段客流_表内合计": 213}},
    "D16": {"code": "axis_extra_cells", "fix": "ignore_cells", "cell": None,
            "steps": [("ignore_cells", None)], "confirms": set(), "rows": UNION_ROWS},
    "D21": {"code": "row_unclaimed", "fix": "ignore_cells", "cell": f"{SHEET}!B32",
            "steps": [("ignore_cells", None)], "confirms": set(), "rows": UNION_ROWS},
}


async def _fix_case(client, code: str, seed: int) -> dict[str, Any]:
    case = FIX_CASES[code]
    base = await first_import(client, seed=seed)
    raw, fn = DRIFT_CASES[code].build(seed + 1)
    st = await reupload(client, base["source_id"], raw, fn)
    assert st["trial"]["status"] == "rejected"
    probs = problems(st, case["code"])
    assert probs and {p["fix"] for p in probs} == {case["fix"]}, [(p["code"], p["fix"]) for p in problems(st)]
    if case["cell"] is not None:
        assert case["cell"] in {c for p in probs for c in p["cells"]}, [p["cells"] for p in probs]
    # 问题旁边的按钮：提议挂在对应的那条问题上（fix_ids，评审三-B2）
    fx = fix_of(st, case["fix"])
    assert any(fx["id"] in p["fix_ids"] for p in probs), (fx["id"], [p["fix_ids"] for p in probs])
    for kind, option in case["steps"]:
        fx = fix_of(st, kind)
        value = next(o["value"] for o in fx["options"] if option is None or o["value"].endswith(option))
        st, pv = await fixed(client, st, kind, value)
        assert pv["recipe_sha256_after"] == st["recipe_sha256"]
    assert st["recipe_problems"] == [], st["recipe_problems"]
    st = await trial(client, st["id"])
    return {"base": base, "st": st, "case": case}


@pytest.mark.parametrize("code", ["D04", "D09", "D13", "D18"])
async def test_a6_drift_fixed_with_the_fix_button_can_be_enabled(client, code):
    """A6：reupload → rejected，问题的 fix 与 3.4 表一致 → StagingOut.fixes 有对应提议 → edits/preview → edits/apply
    （D13 接着应用 edit_members）→ trial（D18 写理由接受 K1）→ 勾全部确认项（含 fix:*）提交 201 → 并集行数与 3.4 一致。"""
    done = await _fix_case(client, code, seed=60 + list(FIX_CASES).index(code))
    st, case, name = done["st"], done["case"], done["base"]["name"]
    t = st["trial"]
    acceptances = None
    if case.get("accept"):
        assert t["status"] == "needs_decision" and t["acceptable"] == [case["accept"]], (t["status"], t["acceptable"])
        k1 = check(st, case["accept"])
        # 3.4：K1（新的日间合计）unverifiable，原因 null_detail 30 格（7-8 行整行是占位符，明细缺数）
        assert k1["status"] == "unverifiable" and k1["reasons"] == {"null_detail": 30} and k1["acceptable"] is True
        assert {c["id"]: c["status"] for c in t["checks"] if c["id"] in ("K2", "G1", "G2")} == {
            "K2": "passed", "G1": "passed", "G2": "passed"}
        acceptances = [{"check_id": case["accept"], "reason": REASON}]
    else:
        assert t["status"] == "passed", (t["problems"], [c for c in t["checks"] if c["status"] != "passed"])
    assert t["accumulate"]["action"] == "append" and t["accumulate"]["union"]["rows"] == case["rows"]
    ids = confirm_ids(st)
    fixes = [i for i in ids if i.startswith("fix:")]
    assert len(fixes) == case.get("fix_items", 1), fixes
    assert case["confirms"] <= set(ids), (case["confirms"], ids)
    assert not [i for i in ids if i.startswith(("breaking:", "accumulate_restart"))], ids
    resp = await commit(client, st, confirmations=[i for i in ids if not i.startswith("fix:")], acceptances=acceptances)
    assert resp.status_code == 422 and resp.json()["code"] == "confirm_required", "fix:* 是必勾项"
    resp = await commit(client, st, acceptances=acceptances)
    assert resp.status_code == 201, resp.text
    assert resp.json()["parts"] == 2
    for table, n in case["rows"].items():
        assert await count(client, name, table) == n, table
    if code == "D09":
        kinds = await tool_rows(client, name, 'SELECT DISTINCT "时段类别" FROM "时段客流" ORDER BY 1')
        assert sorted(k[0] for k in kinds) == ["夜间", "日间"]
    if code == "D13":
        assert await count(client, name, "日客流", ' WHERE "分区丙" IS NULL') == 31
        assert await by_month(client, name, "日客流", ' WHERE "分区丙" IS NULL') == [["2026-08", 31]]
        snap = await db_get(SourceSnapshot, resp.json()["snapshot_id"])
        assert "各期登记的分项列不同" in snap.schema_cache["tables"]["日客流"]["comment"]


@pytest.mark.parametrize("code", ["D16", "D21"])
async def test_a6b_ignore_columns_and_rows_then_enable(client, code):
    """A6b：D16 按表头忽略右侧的「合计」列（⑥ columns），D21 按行标签忽略「补录（人次）」（⑥ rows，要理由）；修复后
    能启用，行数与 D01 相同（30 / 510 / 90，并集 61 / 1037 / 183）；D21 的回执 rows_excluded 有 ignored_rows（第 32 行）。"""
    done = await _fix_case(client, code, seed=70 + (code == "D21"))
    st, name = done["st"], done["base"]["name"]
    t = st["trial"]
    assert t["status"] == "passed", (t["problems"], [c for c in t["checks"] if c["status"] != "passed"])
    assert rows_of(st) == SEP_ROWS
    excluded = t["receipt"]["rows_excluded"]
    if code == "D21":
        ignored = [x for x in excluded if x["reason"] == "ignored_rows"]
        assert ignored and ignored[0]["rows"] == [[32, 32]] and ignored[0]["anchor"] == "补录（人次）", excluded
        assert ignored[0]["cells"] == 3, "标签格加两格数（B32:D32）"
    resp = await commit(client, st)
    assert resp.status_code == 201, resp.text
    for table, n in UNION_ROWS.items():
        assert await count(client, name, table) == n
    if code == "D21":
        imp = await db_get(TableImport, resp.json()["import_id"])
        manifest = artifact_store.load(imp.manifest_artifact)
        assert [x["reason"] for x in manifest["receipt"]["rows_excluded"]] == [x["reason"] for x in excluded]
    # 后续每期按锚点忽略：锚点照旧的下一期直接通过；换了一个没见过的标签（列头）照样拒收，不会被这条规则吞掉
    from openpyxl.utils import get_column_letter

    raw, fn = flow_workbook(OCT, 31, seed=72, variant=code)
    st = await reupload(client, done["base"]["source_id"], raw, fn)
    assert st["trial"]["status"] == "passed", st["trial"]["problems"]
    await client.delete(f"{API}/imports/{st['id']}")
    other = {"B32": "备用分区（人次）"} if code == "D21" else {f"{get_column_letter(3 + 31)}4": "备注"}
    st = await reupload(client, done["base"]["source_id"], edit_cells(raw, other), fn)
    want = "row_unclaimed" if code == "D21" else "axis_extra_cells"
    assert st["trial"]["status"] == "rejected" and problems(st, want), st["trial"]["problems"]


async def test_a6c_d04_with_a_failing_r1_day_is_fixed_then_accepted(client):
    """A6c（评审一-M3）：D04 上再改一格分区甲，使 R1 有一天不成立 → 应用 ① 后 recipe_problems 为空（R1 的认领不查）→
    试运行 R1 mismatch、needs_decision → 写理由接受后提交 201；说明用 sum_eq_union_mixed、不含「已逐日核对」。"""
    base = await first_import(client, seed=81)
    raw, fn = DRIFT_CASES["D04"].build(82)
    raw = edit_cells(raw, {"L6": 1})          # 第 10 个日期（L 列）的分区甲改成 1：全日 ≠ 甲 + 乙
    st = await reupload(client, base["source_id"], raw, fn)
    assert st["trial"]["status"] == "rejected" and problems(st, "label_missing")
    st, _ = await fixed(client, st, "remove_label", "remove")
    assert st["recipe_problems"] == [], "改的是 dimension 段：R1 的认领按重放处理，不查（3.3 第 11 条）"
    st = await trial(client, st["id"])
    t = st["trial"]
    r1 = check(st, "R1")
    assert t["status"] == "needs_decision" and r1["status"] == "mismatch" and t["acceptable"] == ["R1"], t["status"]
    resp = await commit(client, st)
    assert resp.status_code == 422 and resp.json()["code"] == "acceptance_required"
    resp = await commit(client, st, acceptances=[{"check_id": "R1", "reason": REASON}])
    assert resp.status_code == 201, resp.text
    comment = (await db_get(SourceSnapshot, resp.json()["snapshot_id"])).schema_cache["tables"]["日客流"]["comment"]
    assert "在部分期的个别" in comment and "不成立或未能核对" in comment and "已逐日核对" not in comment, comment


async def test_a15_after_a_fix_only_changed_recipe_items_are_listed(client):
    """A15：D04 修复之后的确认项里，配方类只有 fix:*（source=edit），没有 placeholder:·、unit:*、relation:* 等没有变化
    的项（6.2：只列签名与现行配方不同的项）。"""
    done = await _fix_case(client, "D04", seed=91)
    st = done["st"]
    items = {i["id"]: i for i in st["trial"]["confirm_items"]}
    assert items["fix:remove_label:日间:7-8"]["source"] == "edit"
    recipe_items = sorted(i for i, v in items.items() if v.get("source", "recipe") == "recipe")
    assert recipe_items == [], recipe_items
    for stale in ("placeholder:·", "unit:日客流.全日客流", "relation:R1", "derived:夜间合计", "mode", "year_from:交叉表"):
        assert stale not in items, stale


# ==========================================================================
# A7：c01 框选回放一致（4.6 的三条）
# ==========================================================================


async def test_a7_c01_selection_replays_consistently(client):
    """4.6：① 同一份文件：月报!C5:F8 框为列表 → ok，锚点是四个表头，重放与框一致，换算后的配方与规则草稿逐字相同；
    ② 整张表挪到 E8（shift=(3, 2)）用同一份配方试运行：passed、T1 passed，列表表和合计表的 table_hashes 与未挪动的
    逐一相同（配方里没有坐标）；③ 泛化不了时明确说明：C5:F7 → selection_bottom_not_anchorable（建议扩大到第 8 行），
    同一框改选 bottom=auto → ok、data=C6:F8、重放一致；C6:F8 → selection_header_not_text；C5:G8 → selection_header_blank。"""
    raw, fn = lab.c01("literal")
    st = await stage(client, raw, fn)
    blocks = [b for s in st["recipe"]["sheets"] for b in s["blocks"]]
    assert [b["layout"] for b in blocks] == ["list"]
    sel = {"sheet": "月报", "ref": "C5:F8", "as": "list", "options": {"header_rows": 1, "bottom": "box"}}
    resp = await preview(client, st["id"], {"selection": sel})
    assert resp.status_code == 200, resp.text
    pv = resp.json()
    assert pv["ok"] is True, pv["problems"]
    assert [a["text"] for a in pv["anchors"] if a["kind"] == "header"] == ["地区", "产品", "销量", "金额"]
    box = {"header": "C5:F5", "data": "C6:F8", "total": "C9:F9"}
    assert pv["replay"]["match"] is True and pv["replay"]["expected"] == box and pv["replay"]["actual"] == box
    assert pv["replay"]["diffs"] == []
    assert pv["recipe_sha256_after"] == st["recipe_sha256"], "框选得到的规则与起草器完全一样"
    resp = await apply(client, st["id"], {"selection": sel}, pv["recipe_sha256_after"])
    assert resp.status_code == 200, resp.text
    st = await trial(client, resp.json()["id"])
    assert st["trial"]["status"] == "passed", st["trial"]["problems"]
    assert check(st, "T1")["status"] == "passed"
    hashes = st["trial"]["receipt"]["table_hashes"]
    list_table = blocks[0]["table"]
    total_table = blocks[0]["rows"]["total_row"]["keep_as"]
    assert {list_table, total_table} <= set(hashes)

    # ② 挪了位置也回放一致：用第 1 条得到的配方（PUT）
    raw2, fn2 = lab.c01("literal", shift=(3, 2))
    moved = await stage(client, raw2, fn2)
    moved = await put_recipe(client, moved["id"], st["recipe"])
    assert moved["recipe_problems"] == [], moved["recipe_problems"]
    moved = await trial(client, moved["id"])
    assert moved["trial"]["status"] == "passed", moved["trial"]["problems"]
    assert check(moved, "T1")["status"] == "passed"
    got = moved["trial"]["receipt"]["table_hashes"]
    assert got[list_table] == hashes[list_table] and got[total_table] == hashes[total_table]

    # ③ 泛化不了时明确说明
    bad = (await preview(client, st["id"], {"selection": {**sel, "ref": "C5:F7"}})).json()
    assert bad["ok"] is False and [p["code"] for p in bad["problems"]] == ["selection_bottom_not_anchorable"]
    assert "第 8 行" in bad["problems"][0]["message"], bad["problems"][0]["message"]
    auto = (await preview(client, st["id"], {"selection": {**sel, "ref": "C5:F7",
                                                           "options": {"header_rows": 1, "bottom": "auto"}}})).json()
    assert auto["ok"] is True, auto["problems"]
    assert auto["expected"]["data"] == "C6:F8" and auto["replay"]["match"] is True
    bad = (await preview(client, st["id"], {"selection": {**sel, "ref": "C6:F8"}})).json()
    assert bad["ok"] is False and "selection_header_not_text" in [p["code"] for p in bad["problems"]]
    bad = (await preview(client, st["id"], {"selection": {**sel, "ref": "C5:G8"}})).json()
    assert bad["ok"] is False and "selection_header_blank" in [p["code"] for p in bad["problems"]]


# ==========================================================================
# A8：回滚（启用旧版本）
# ==========================================================================


async def test_a8_rollback_to_an_earlier_version(client, monkeypatch):
    """A8：在 S2（累积）上 activate S1（带 expected_current=S2）→ 当前快照是 S1，current_recipe_id 是 S1 的配方，导入记录
    状态互换，写了 snapshot_activations；在 S2 上发起的固定运行仍查 61 行；confirm 缺省 422；expected_current 不对 409
    base_changed；遮罩列在目标版本里没有同名列时不带 ack_mask_lost 409 mask_lost，带上成功；activate 已回收的快照
    409 snapshot_unavailable（KEEP=0 造出已回收的快照，1.5）。

    9 月用 D13 走修复按钮（加「分区丙」），S2 的配方 R2 与 S1 的 R0 不是同一条：启用前现行配方是 R2，启用 S1 之后是 R0，
    两条配方的状态互换。两期同一个配方时「现行配方跟着快照走」这一条测不出来（WP-8 评审意见 3）。"""
    from app.data.engine import engines

    done = await _fix_case(client, "D13", seed=101)
    assert done["st"]["trial"]["status"] == "passed", done["st"]["trial"]["problems"]
    resp = await commit(client, done["st"])
    assert resp.status_code == 201, resp.text
    base, sep = done["base"], resp.json()
    v = {**base, "s2": sep["snapshot_id"], "aug": base["import_id"], "sep": sep["import_id"]}
    sid, name = v["source_id"], v["name"]
    r0, r2 = base["recipe_id"], sep["recipe_id"]
    assert r0 != r2 and (await db_get(SourceSnapshot, v["s1"])).recipe_id == r0
    assert (await db_get(SourceSnapshot, v["s2"])).recipe_id == r2
    assert (await db_get(DataSource, sid)).current_recipe_id == r2
    assert ((await db_get(TableRecipe, r0)).status, (await db_get(TableRecipe, r2)).status) == ("superseded", "active")
    await keep_run(sid, name, v["s2"])
    url = f"{API}/{sid}/snapshots/{v['s1']}/activate"
    resp = await client.post(url, json={"expected_current_snapshot_id": v["s2"]})
    assert resp.status_code == 422 and resp.json()["code"] == "confirm_required"
    resp = await client.post(url, json={"confirm": True})
    assert resp.status_code == 422 and resp.json()["code"] == "expected_current_required"
    resp = await client.post(url, json={"confirm": True, "expected_current_snapshot_id": v["s1"]})
    assert resp.status_code == 409 and resp.json()["code"] == "base_changed"
    async with SessionLocal() as session:
        await session.execute(update(DataSource).where(DataSource.id == sid).values(options={"mask_columns": "分区丁"}))
        await session.commit()
    listed = {s["id"]: s for s in await snapshots(client, sid)}
    assert listed[v["s1"]]["mask_lost"] == ["分区丁"] and listed[v["s1"]]["activatable"] is True
    resp = await client.post(url, json={"confirm": True, "expected_current_snapshot_id": v["s2"]})
    assert resp.status_code == 409 and resp.json()["code"] == "mask_lost"
    resp = await client.post(url, json={"confirm": True, "expected_current_snapshot_id": v["s2"], "reason": REASON,
                                        "ack_mask_lost": ["分区丁"], "signed_by": "合成署名甲"})
    assert resp.status_code == 200, resp.text
    out = resp.json()
    assert out["snapshot_id"] == v["s1"] and out["previous_snapshot_id"] == v["s2"]
    s1 = await db_get(SourceSnapshot, v["s1"])
    src = await db_get(DataSource, sid)
    assert src.current_snapshot_id == v["s1"] and src.current_recipe_id == s1.recipe_id == out["recipe_id"] == r0
    assert ((await db_get(TableRecipe, r0)).status, (await db_get(TableRecipe, r2)).status) == ("active", "superseded")
    assert (await db_get(TableImport, v["aug"])).status == "active"
    assert (await db_get(TableImport, v["sep"])).status == "superseded"
    async with SessionLocal() as session:
        last = (await session.execute(select(SnapshotActivation).where(SnapshotActivation.source_id == sid)
                                      .order_by(SnapshotActivation.created_at.desc()))).scalars().first()
    assert (last.kind, last.snapshot_id, last.previous_snapshot_id, last.signed_by) == (
        "activate", v["s1"], v["s2"], "合成署名甲")
    assert await count(client, name, "日客流") == 31
    try:
        pinned = (await tools_for([f"db_query__{name}"], data_versions={sid: {"snapshot": v["s2"], "name": name}}))
        text = await pinned[f"db_query__{name}"].coroutine(sql='SELECT COUNT(*) FROM "日客流"')
        assert json.loads(text)["rows"][0][0] == 61, "在 S2 上发起的固定运行仍查 S2"
    finally:
        await engines.invalidate(sid)
    # 已回收的快照不能启用：KEEP=0 下再提交 10 月，S1（不是当前、没人引用）被回收
    monkeypatch.setattr(table_versions, "SNAPSHOT_KEEP", 0)
    raw, fn = flow_workbook(OCT, 31, seed=102)
    s3 = (await add_period(client, sid, raw, fn))["snapshot_id"]
    assert (await db_get(SourceSnapshot, v["s1"])).retired_at is not None
    resp = await client.post(url, json={"confirm": True, "expected_current_snapshot_id": s3,
                                        "ack_mask_lost": ["分区丁"]})
    assert resp.status_code == 409 and resp.json()["code"] == "snapshot_unavailable"


# ==========================================================================
# A9：作废接受（累积模式 = 移除该期）
# ==========================================================================


async def test_a9_revoking_an_acceptance_removes_the_period_and_keeps_recipe_and_mode(client):
    """A9：D26（9 月 R1 不成立、写理由接受）累积提交 → 说明含 sum_eq_union_mixed 的片段、不含「已逐日核对」→
    revoke-acceptance → action=remove_period，新快照只有 8 月，快照 id 等于 S1（复用），current_recipe_id 和模式都没变，
    该导入 revoked 非空；再 activate 含已作废导入的快照 409 contains_revoked；remove 唯一一期 409 last_period。"""
    base = await first_import(client, seed=111)
    sid = base["source_id"]
    raw, fn = DRIFT_CASES["D26"].build(112)
    st = await reupload(client, sid, raw, fn)
    assert st["trial"]["status"] == "needs_decision" and st["trial"]["acceptable"] == ["R1"]
    resp = await commit(client, st, acceptances=[{"check_id": "R1", "reason": REASON}])
    assert resp.status_code == 201, resp.text
    s2, sep = resp.json()["snapshot_id"], resp.json()["import_id"]
    comment = (await db_get(SourceSnapshot, s2)).schema_cache["tables"]["日客流"]["comment"]
    assert "在部分期的个别" in comment and "不成立或未能核对" in comment and "已逐日核对" not in comment, comment
    recipe_before = (await db_get(DataSource, sid)).current_recipe_id
    resp = await client.post(f"{API}/{sid}/imports/{sep}/revoke-acceptance",
                             json={"confirm": True, "expected_current_snapshot_id": s2, "reason": REASON})
    assert resp.status_code == 200, resp.text
    out = resp.json()
    assert out["action"] == "remove_period" and out["snapshot_id"] == base["s1"]
    snap = await db_get(SourceSnapshot, out["snapshot_id"])
    assert snap.imports == [base["import_id"]] and snap.mode == "accumulate"
    assert (await db_get(DataSource, sid)).current_recipe_id == recipe_before
    imp = await db_get(TableImport, sep)
    assert imp.revoked and imp.revoked["reason"] == REASON and imp.status == "superseded"
    resp = await client.post(f"{API}/{sid}/snapshots/{s2}/activate",
                             json={"confirm": True, "expected_current_snapshot_id": base["s1"]})
    assert resp.status_code == 409 and resp.json()["code"] == "contains_revoked"
    resp = await client.post(f"{API}/{sid}/imports/{base['import_id']}/remove",
                             json={"confirm": True, "expected_current_snapshot_id": base["s1"], "reason": REASON})
    assert resp.status_code == 409 and resp.json()["code"] == "last_period"


# ==========================================================================
# A10：不兼容变更要重新开始累积
# ==========================================================================


async def test_a10_an_incompatible_change_restarts_the_accumulation(client):
    """A10：S2 之后 reupload 10 月，在 10 月的暂存区里用 PUT 把一列的单位改掉 → 试运行 action=restart，确认项有
    accumulate_restart；提交后当前快照只有 10 月，说明含 single_after_drop；S2 仍可启用（在 SNAPSHOT_KEEP 之内）。"""
    v = await august_september(client, seed=121)
    sid = v["source_id"]
    # 不兼容与部分重叠同时出现：restart 优先（结果只有本期，不涉及重叠），重叠的各期只作提示（2.4，评审一-m12）
    st = await restart_with_unit_change(client, sid, start=date_of(2026, 8, 15), days=31, seed=120)
    assert st["trial"]["status"] == "passed" and st["trial"]["accumulate"]["action"] == "restart", st["trial"]
    assert [p["start"] for p in st["trial"]["accumulate"]["overlaps"]] == ["2026-08-01", "2026-09-01"]
    await client.delete(f"{API}/imports/{st['id']}")
    st = await restart_with_unit_change(client, sid, seed=122)
    t = st["trial"]
    assert t["status"] == "passed", t["problems"]
    plan = t["accumulate"]
    assert plan["action"] == "restart" and plan["change"] == "semantic"
    assert [p["start"] for p in plan["dropped"]] == ["2026-08-01", "2026-09-01"]
    assert any("人次" in x and "人" in x for x in plan["semantic"]), plan["semantic"]
    assert "accumulate_restart" in confirm_ids(st)
    resp = await commit(client, st, confirmations=[i for i in confirm_ids(st) if i != "accumulate_restart"])
    assert resp.status_code == 422 and resp.json()["code"] == "confirm_required"
    resp = await commit(client, st)
    assert resp.status_code == 201, resp.text
    out = resp.json()
    assert out["parts"] == 1
    assert await current_periods(sid) == ["2026-10-01"]
    comment = (await db_get(SourceSnapshot, out["snapshot_id"])).schema_cache["tables"]["日客流"]["comment"]
    assert "当前版本只含最近一期的数据" in comment, comment
    listed = {s["id"]: s for s in await snapshots(client, sid)}
    assert listed[v["s2"]]["activatable"] is True and listed[v["s2"]]["available"] is True


# ==========================================================================
# A11：退役列保留
# ==========================================================================


async def drop_zone_b(client, st: dict[str, Any]) -> dict[str, Any]:
    """这一期的文件不再有分区乙：修复按钮去掉标签（①，连带去掉列和单位），R1（全日 = 甲 + 乙）删掉（③ 的「删除这条
    关系」），R2 按本期的系统发现重新认领（本期只剩口径不同那一条，编号是 F1；改认领只能在配方面板里 PUT）。"""
    st, _ = await fixed(client, st, "remove_label", "remove")
    st, _ = await fixed(client, st, "edit_members", "remove")
    rec = json.loads(json.dumps(st["recipe"], ensure_ascii=False))
    for r in rec["relations"]:
        if r["id"] == "R2":
            r["claims"] = "F1"
    st = await put_recipe(client, st["id"], rec)
    assert st["recipe_problems"] == [], st["recipe_problems"]
    assert [r["id"] for r in st["recipe"]["relations"]] == ["R2"]
    return await trial(client, st["id"])


async def test_a11_a_retired_column_is_kept_and_not_retired_again(client):
    """A11：去掉 measures 的分区乙（R1 不再登记）后累积 9 月 → 确认项有 retire:日客流.分区乙；并集里 8 月有值、9 月为空值；
    列说明含「自某一期起不再导入」，分区乙的单位仍是人次；日客流的说明用 sum_eq_union_partial，不含「已逐日核对」。
    再累积 10 月（评审二-M3）：确认项里没有 retire:*，也没有 breaking:日客流；计划的 retired_existing 有分区乙。"""
    base = await first_import(client, seed=131)
    sid, name = base["source_id"], base["name"]
    raw, fn = flow_workbook(SEP, 30, seed=132)
    st = await reupload(client, sid, without_zone_b(raw), fn)
    assert st["trial"]["status"] == "rejected" and problems(st, "label_missing")
    st = await drop_zone_b(client, st)
    t = st["trial"]
    assert t["status"] == "passed", t["problems"]
    assert t["accumulate"]["action"] == "append" and t["accumulate"]["change"] == "retire"
    assert t["accumulate"]["retired_new"] == [{"table": "日客流", "column": "分区乙"}]
    assert "retire:日客流.分区乙" in confirm_ids(st)
    assert not [i for i in confirm_ids(st) if i.startswith("breaking:日客流")], "退役与 breaking 同表去重（2.5）"
    resp = await commit(client, st)
    assert resp.status_code == 201, resp.text
    snap = await db_get(SourceSnapshot, resp.json()["snapshot_id"])
    assert await by_month(client, name, "日客流", ' WHERE "分区乙" IS NOT NULL') == [["2026-08", 31]]
    assert await by_month(client, name, "日客流", ' WHERE "分区乙" IS NULL') == [["2026-09", 30]]
    table = snap.schema_cache["tables"]["日客流"]
    col = next(c for c in table["columns"] if c["name"] == "分区乙")
    assert "单位：人次" in col["comment"] and "自某一期起不再导入" in col["comment"], col
    assert "只在部分期登记并核对" in table["comment"] and "已逐日核对" not in table["comment"], table["comment"]
    manifest = artifact_store.load(snap.manifest_artifact)
    assert manifest["columns"]["日客流"]["分区乙"]["status"] == "retired"

    raw, fn = flow_workbook(OCT, 31, seed=133)
    st = await reupload(client, sid, without_zone_b(raw), fn)
    t = st["trial"]
    assert t["status"] == "passed", t["problems"]
    assert t["accumulate"]["retired_new"] == []
    assert {"table": "日客流", "column": "分区乙"} in t["accumulate"]["retired_existing"]
    assert not [i for i in confirm_ids(st) if i.startswith(("retire", "breaking:日客流"))], confirm_ids(st)
    resp = await commit(client, st)
    assert resp.status_code == 201, resp.text
    assert await by_month(client, name, "日客流", ' WHERE "分区乙" IS NULL') == [["2026-09", 30], ["2026-10", 31]]


# ==========================================================================
# A12：模式切换
# ==========================================================================


async def test_a12_switching_modes_both_ways(client):
    """A12：替换模式的源上传新一期并回答 q_mode=accumulate → mode_switch:replace->accumulate，当前那一期作为第一期；
    反向切换出 mode_switch:accumulate->replace，提交后只剩一期。"""
    base = await first_import(client, seed=141, mode="replace")
    sid, name = base["source_id"], base["name"]
    assert (await db_get(SourceSnapshot, base["s1"])).mode == "replace"
    raw, fn = flow_workbook(SEP, 30, seed=142)
    st = await reupload(client, sid, raw, fn)
    assert "q_mode" in {q["id"] for q in st["questions"]}
    st = await answer(client, st["id"], {"q_mode": {"value": "accumulate"}})
    st = await trial(client, st["id"])
    t = st["trial"]
    assert t["status"] == "passed", t["problems"]
    assert t["accumulate"]["mode_switch"] == "replace->accumulate" and t["accumulate"]["action"] == "append"
    assert [p["start"] for p in t["accumulate"]["parts"]] == ["2026-08-01", "2026-09-01"]
    assert "mode_switch:replace->accumulate" in confirm_ids(st)
    resp = await commit(client, st)
    assert resp.status_code == 201, resp.text
    assert resp.json()["parts"] == 2 and await count(client, name, "日客流") == 61

    raw, fn = flow_workbook(OCT, 31, seed=143)
    st = await reupload(client, sid, raw, fn)
    st = await answer(client, st["id"], {"q_mode": {"value": "replace"}})
    st = await trial(client, st["id"])
    t = st["trial"]
    assert t["status"] == "passed", t["problems"]
    assert t["accumulate"]["mode_switch"] == "accumulate->replace" and t["accumulate"]["action"] == "replace"
    assert "mode_switch:accumulate->replace" in confirm_ids(st)
    resp = await commit(client, st)
    assert resp.status_code == 201, resp.text
    assert resp.json()["parts"] == 1
    snap = await db_get(SourceSnapshot, resp.json()["snapshot_id"])
    assert snap.mode == "replace" and len(snap.imports) == 1
    assert await count(client, name, "日客流") == 31


# ==========================================================================
# A13：快照清单与证据链
# ==========================================================================


async def test_a13_query_snapshot_commits_to_the_snapshot_and_import_manifests(client):
    """A13：在累积快照上跑一次查询工具：query_snapshot → schema_snapshot 里有 snapshot_manifest 和两个 import_manifests →
    快照清单的 union.db_sha256 等于快照的 db_sha256；各期 rows 区间相接，覆盖 1..61。"""
    from app.data.engine import engines

    v = await august_september(client, seed=151)
    sid, name = v["source_id"], v["name"]
    try:
        query = (await tools_for([f"db_query__{name}"]))[f"db_query__{name}"]
        payload = json.loads(await query.coroutine(sql='SELECT COUNT(*) FROM "日客流"'))
    finally:
        await engines.invalidate(sid)
    assert payload["rows"][0][0] == 61
    qsnap = artifact_store.load(payload["artifact"])
    assert qsnap["data_version"] == v["s2"]
    schema = artifact_store.load(payload["schema_artifact"])
    snap = await db_get(SourceSnapshot, v["s2"])
    assert schema["snapshot_manifest"] == snap.manifest_artifact
    imps = [await db_get(TableImport, i) for i in snap.imports]
    assert schema["import_manifests"] == [i.manifest_artifact for i in imps] and len(imps) == 2
    manifest = artifact_store.load(schema["snapshot_manifest"])
    assert manifest["format"] == "agentlab-snapshot-manifest/1" and manifest["snapshot_id"] == v["s2"]
    assert manifest["union"]["db_sha256"] == snap.db_sha256 == raw_store.sha256_file(Path(snap.db_path))
    assert [p["import_id"] for p in manifest["parts"]] == snap.imports
    for table, total in UNION_ROWS.items():
        spans = [p["rows"][table]["union"] for p in manifest["parts"]]
        assert spans[0][0] == 1 and spans[-1][1] == total, (table, spans)
        assert all(a[1] + 1 == b[0] for a, b in zip(spans, spans[1:])), (table, spans)
        assert [p["rows"][table]["part"][0] for p in manifest["parts"]] == [1, 1]
    for imp, part in zip(imps, manifest["parts"]):
        doc = artifact_store.load(part["manifest_artifact"])
        assert doc["import_id"] == imp.id and part["manifest_artifact"] == imp.manifest_artifact


# ==========================================================================
# A14：确定性与复用
# ==========================================================================


async def test_a14_same_september_file_under_another_name_reuses_the_union(client):
    """A14：9 月同一份文件换个名字再传（替换该期）→ 新导入记录（文件名不同，不是期 2 的「原样重传」），union_id 与 S2
    相同（并集 id 只取决于各期构建和目标表结构，不含导入记录 id），并集文件复用（新快照的库就是 S2 那个文件），
    db_sha256 相同。对照：同名原样再传不新建记录（unchanged）。"""
    v = await august_september(client, seed=161)
    sid = v["source_id"]
    s2_row = await db_get(SourceSnapshot, v["s2"])
    st = await reupload(client, sid, v["sep_raw"], "九月客流（再传）.xlsx")
    t = st["trial"]
    assert t["status"] == "passed" and t["accumulate"]["action"] == "replace_period", t["problems"]
    assert t["same_as_import"] == {"id": v["sep"], "seq": 2}, "同一个构建"
    assert t["accumulate"]["union"]["union_id"] == Path(s2_row.db_path).stem
    assert t["accumulate"]["union"]["db_sha256"] == s2_row.db_sha256
    resp = await commit(client, st)
    assert resp.status_code == 201, resp.text
    out = resp.json()
    assert out["unchanged"] is False and out["build_reused"] is True
    assert out["import_id"] != v["sep"] and out["snapshot_id"] != v["s2"]
    s3 = await db_get(SourceSnapshot, out["snapshot_id"])
    assert s3.db_path == s2_row.db_path and s3.db_sha256 == s2_row.db_sha256
    assert raw_store.sha256_file(Path(s3.db_path)) == s2_row.db_sha256
    assert (await db_get(TableImport, out["import_id"])).file_name == "九月客流（再传）.xlsx"
    async with SessionLocal() as session:
        unions = (await session.execute(select(TableBuild).where(
            TableBuild.source_id == sid, TableBuild.raw_sha256 == ""))).scalars().all()
    assert [u.id for u in unions] == [Path(s2_row.db_path).stem], "并集只登记一次"
    # 同名原样再传：不新建记录，当前版本不动
    st = await reupload(client, sid, v["sep_raw"], "九月客流（再传）.xlsx")
    resp = await commit(client, st)
    assert resp.status_code == 201 and resp.json()["unchanged"] is True, resp.text
    assert resp.json()["snapshot_id"] == out["snapshot_id"]


# ==========================================================================
# A16：并发：移除与提交交错
# ==========================================================================


async def test_a16_remove_with_a_stale_version_after_a_commit_is_base_changed(client):
    """A16 前半：S2=[8, 9] 上先取 snapshots 得到当前 id；另一个请求提交 10 月得 S3；再用旧的 expected_current 移除 9 月
    → 409 base_changed，当前快照仍是 S3，10 月仍在。"""
    v = await august_september(client, seed=171)
    sid = v["source_id"]
    seen = (await snapshots(client, sid))[0]["id"]
    assert seen == v["s2"]
    raw, fn = flow_workbook(OCT, 31, seed=172)
    s3 = (await add_period(client, sid, raw, fn))["snapshot_id"]
    resp = await client.post(f"{API}/{sid}/imports/{v['sep']}/remove",
                             json={"confirm": True, "expected_current_snapshot_id": seen, "reason": REASON})
    assert resp.status_code == 409 and resp.json()["code"] == "base_changed", resp.text
    assert await current_id(sid) == s3
    assert await current_periods(sid) == ["2026-08-01", "2026-09-01", "2026-10-01"]


async def test_a16_interleaved_publishes_the_later_one_is_refused(client, monkeypatch, store):
    """A16 后半：两次发布卡在锁前交错。提交 10 月的请求进了 publish_snapshot、还没进锁时，移除 9 月（按 S2 确认）先完成；
    提交随后在锁内重读指针，与暂存区的 base（S2）对不上 → 后到的一方 409 base_changed，当前版本是移除的结果，10 月没进去，
    暂存区退回 drafting、没留下临时文件。反过来（移除停在物化里、提交先完成）由 test_source_versions_api 覆盖。"""
    v = await august_september(client, seed=173)
    sid = v["source_id"]
    raw, fn = flow_workbook(OCT, 31, seed=174)
    oct_st = await reupload(client, sid, raw, fn)
    assert oct_st["trial"]["status"] == "passed"
    orig = table_versions.publish_snapshot
    removed: dict[str, Any] = {}

    async def interleaved(session, source, **kw):
        if (kw.get("activation") or {}).get("kind") == "commit" and not removed:
            removed["resp"] = await client.post(
                f"{API}/{sid}/imports/{v['sep']}/remove",
                json={"confirm": True, "expected_current_snapshot_id": v["s2"], "reason": REASON})
        return await orig(session, source, **kw)

    monkeypatch.setattr(table_versions, "publish_snapshot", interleaved)
    resp = await commit(client, oct_st)
    assert removed["resp"].status_code == 200, removed["resp"].text
    assert removed["resp"].json()["snapshot_id"] == v["s1"]
    assert resp.status_code == 409 and resp.json()["code"] == "base_changed", resp.text
    assert await current_id(sid) == v["s1"] and await current_periods(sid) == ["2026-08-01"]
    assert (await db_get(ImportStaging, oct_st["id"])).status == "drafting"
    assert list((store / "uploads").rglob("*.tmp-*")) == []


# ==========================================================================
# A17：移除时配方、模式不退回；已回收的快照复活
# ==========================================================================


async def test_a17_1_removing_after_a_mode_switch_keeps_the_target_recipe_and_mode(client):
    """A17①：替换模式 S1=[8]（配方 r1）→ 9 月 q_mode=accumulate 得 S2=[8, 9]（目标 r2）→ 移除 9 月：新快照 mode=accumulate、
    recipe=r2，id 不等于 S1；8 月的配方（r1）哈希不等于 r2，快照库是物化的单期并集（在 snapshots/ 下），说明含
    single_after_drop、不含「部分期」。"""
    base = await first_import(client, seed=181, mode="replace")
    sid, r1 = base["source_id"], base["recipe_id"]
    raw, fn = flow_workbook(SEP, 30, seed=182)
    st = await reupload(client, sid, raw, fn)
    st = await answer(client, st["id"], {"q_mode": {"value": "accumulate"}})
    st = await trial(client, st["id"])
    assert st["trial"]["status"] == "passed", st["trial"]["problems"]
    resp = await commit(client, st)
    assert resp.status_code == 201, resp.text
    s2, r2, sep = resp.json()["snapshot_id"], resp.json()["recipe_id"], resp.json()["import_id"]
    assert r2 != r1
    assert (await db_get(TableRecipe, r1)).recipe_sha256 != (await db_get(TableRecipe, r2)).recipe_sha256
    resp = await client.post(f"{API}/{sid}/imports/{sep}/remove",
                             json={"confirm": True, "expected_current_snapshot_id": s2, "reason": REASON})
    assert resp.status_code == 200, resp.text
    out = resp.json()
    assert out["snapshot_id"] not in (base["s1"], s2) and out["reused"] is False
    snap = await db_get(SourceSnapshot, out["snapshot_id"])
    assert snap.mode == "accumulate" and snap.recipe_id == r2 and snap.imports == [base["import_id"]]
    assert Path(snap.db_path).parent.name == "snapshots", "一期、但配方不是目标配方：按目标配方物化成单期并集"
    assert (await db_get(DataSource, sid)).current_recipe_id == r2
    comment = snap.schema_cache["tables"]["日客流"]["comment"]
    assert "当前版本只含最近一期的数据" in comment and "部分期" not in comment, comment


async def test_a17_2_removing_the_period_that_added_a_column_keeps_the_target_recipe(client):
    """A17②：[8, 9, 10] 中 10 月加了分区丙（修复按钮，目标 r3）→ 移除 10 月：目标仍是 r3，并集里分区丙为空值；下一期
    按 r3 重放（不用再修）。"""
    v = await august_september(client, seed=183)
    sid, name = v["source_id"], v["name"]
    raw, fn = flow_workbook(OCT, 31, seed=184, variant="D13")
    st = await reupload(client, sid, raw, fn)
    st, _ = await fixed(client, st, "add_label", "add")
    st, _ = await fixed(client, st, "edit_members", "update")
    st = await trial(client, st["id"])
    assert st["trial"]["status"] == "passed", st["trial"]["problems"]
    resp = await commit(client, st)
    assert resp.status_code == 201, resp.text
    s3, r3, oct_id = resp.json()["snapshot_id"], resp.json()["recipe_id"], resp.json()["import_id"]
    resp = await client.post(f"{API}/{sid}/imports/{oct_id}/remove",
                             json={"confirm": True, "expected_current_snapshot_id": s3, "reason": REASON})
    assert resp.status_code == 200, resp.text
    snap = await db_get(SourceSnapshot, resp.json()["snapshot_id"])
    assert snap.recipe_id == r3 and snap.imports == [v["aug"], v["sep"]] and snap.id != v["s2"]
    assert (await db_get(DataSource, sid)).current_recipe_id == r3
    assert "分区丙" in columns_of(snap, "日客流")
    assert await count(client, name, "日客流", ' WHERE "分区丙" IS NULL') == 61
    raw, fn = flow_workbook(date_of(2026, 11, 1), 30, seed=185, variant="D13")
    st = await reupload(client, sid, raw, fn)
    assert st["trial"]["status"] == "passed", st["trial"]["problems"]
    assert st["recipe_sha256"] == (await db_get(TableRecipe, r3)).recipe_sha256


async def test_a17_3_removing_revives_a_retired_snapshot_whose_union_file_was_deleted(client, monkeypatch):
    """A17③（1.5：KEEP=0）：S2=[8, 9]（并集）在提交 10 月之后被回收、并集文件被删；移除 10 月按同一组导入、同一个目标配方
    得到 S2 的 id → 复活（重新物化、link 回去），文件哈希等于原登记值。"""
    monkeypatch.setattr(table_versions, "SNAPSHOT_KEEP", 0)
    v = await august_september(client, seed=186)
    sid = v["source_id"]
    s2 = await db_get(SourceSnapshot, v["s2"])
    registered, union_file = s2.db_sha256, Path(s2.db_path)
    raw, fn = flow_workbook(OCT, 31, seed=187)
    out = await add_period(client, sid, raw, fn)
    assert (await db_get(SourceSnapshot, v["s2"])).retired_at is not None and not union_file.exists()
    resp = await client.post(f"{API}/{sid}/imports/{out['import_id']}/remove",
                             json={"confirm": True, "expected_current_snapshot_id": out["snapshot_id"],
                                   "reason": REASON})
    assert resp.status_code == 200, resp.text
    assert resp.json()["snapshot_id"] == v["s2"] and resp.json()["reused"] is True
    revived = await db_get(SourceSnapshot, v["s2"])
    assert revived.retired_at is None and revived.db_sha256 == registered
    assert union_file.is_file() and raw_store.sha256_file(union_file) == registered
    assert (await db_get(TableBuild, union_file.stem)).retired_at is None
    assert await count(client, v["name"], "日客流") == 61


# ==========================================================================
# A18：退役名撞名、单位风险
# ==========================================================================


def rename_column(rec: dict[str, Any], old: str, new: str) -> dict[str, Any]:
    """把日客流的一列改名：measures 映射、单位、关系的成员一起改（改名在配方里分不出来，按「旧名退役、新名新增」报）。"""
    rec = json.loads(json.dumps(rec, ensure_ascii=False))
    for sheet in rec["sheets"]:
        for block in sheet["blocks"]:
            for seg in block.get("segments", []):
                if isinstance(seg.get("measures"), dict):
                    seg["measures"] = {k: (new if v == old else v) for k, v in seg["measures"].items()}
    for t in rec["tables"]:
        if t["name"] == "日客流" and old in t.get("units", {}):
            t["units"] = {(new if k == old else k): v for k, v in t["units"].items()}
    for r in rec["relations"]:
        if r.get("table") == "日客流":
            r["parts"] = [new if x == old else x for x in r.get("parts", [])]
            if r.get("total") == old:
                r["total"] = new
    return rec


async def test_a18_1_a_new_column_colliding_with_a_retired_one_restarts(client):
    """A18①：当前版本有退役列「PM2_5」时，新增「pm2_5」→ restart（SQLite 列名不分大小写，并集建表会报重名，评审一-m6）。"""
    base = await first_import(client, seed=191, edit=lambda rec: rename_column(rec, "分区乙", "PM2_5"))
    sid = base["source_id"]
    raw, fn = flow_workbook(SEP, 30, seed=192)
    st = await reupload(client, sid, raw, fn)
    st = await put_recipe(client, st["id"], rename_column(st["recipe"], "PM2_5", "分区乙"))
    st = await trial(client, st["id"])
    t = st["trial"]
    assert t["status"] == "passed", t["problems"]
    assert t["accumulate"]["action"] == "append"
    assert {"table": "日客流", "column": "PM2_5"} in t["accumulate"]["retired_new"]
    assert "retire:日客流.PM2_5" in confirm_ids(st)
    resp = await commit(client, st)
    assert resp.status_code == 201, resp.text
    raw, fn = flow_workbook(OCT, 31, seed=193)
    st = await reupload(client, sid, raw, fn)
    st = await put_recipe(client, st["id"], rename_column(st["recipe"], "分区乙", "pm2_5"))
    assert st["recipe_problems"] == [], st["recipe_problems"]
    st = await trial(client, st["id"])
    t = st["trial"]
    assert t["status"] == "passed", t["problems"]
    assert t["accumulate"]["action"] == "restart"
    assert any("pm2_5" in x and "PM2_5" in x for x in t["accumulate"]["semantic"]), t["accumulate"]["semantic"]
    assert "accumulate_restart" in confirm_ids(st)


async def test_a18_2_unit_wording_outside_the_region_changing_needs_confirmation(client):
    """A18②：累积模式下 9 月区域外的「单位：人次」改成「单位：万人次」→ 确认项有 accumulate_unit_risk（单位写在配方之外，
    classify_changes 看不到，评审一-M10）。"""
    raw8, _ = flow_workbook(AUG, 31, seed=194)
    base = await first_import(client, seed=194, raw=edit_cells(raw8, {"B32": "注：单位：人次"}))
    raw9, fn9 = flow_workbook(SEP, 30, seed=195)
    st = await reupload(client, base["source_id"], edit_cells(raw9, {"B32": "注：单位：万人次"}), fn9)
    t = st["trial"]
    assert t["status"] == "passed", t["problems"]
    assert t["accumulate"]["action"] == "append"
    assert "accumulate_unit_risk" in confirm_ids(st), confirm_ids(st)
    item = next(i for i in t["confirm_items"] if i["id"] == "accumulate_unit_risk")
    assert "万人次" in item["label"] + (item.get("detail") or "")
    # 对照：区域外文字不变时不出
    raw9b, fn9b = flow_workbook(SEP, 30, seed=196)
    await client.delete(f"{API}/imports/{st['id']}")
    st = await reupload(client, base["source_id"], edit_cells(raw9b, {"B32": "注：单位：人次"}), fn9b)
    assert st["trial"]["status"] == "passed" and "accumulate_unit_risk" not in confirm_ids(st)


# ==========================================================================
# 版本页：清除原件（7.6）
# ==========================================================================


async def test_purging_raw_files_from_the_version_page(client, store):
    """7.6：清除原件要理由（422 reason_required）；当前版本里的一期可以清，不在当前版本里的导入（被替换掉的那一期）也能清
    （purge-raw 不看导入状态，1.5）；清除后原件文件删掉、导入记录 raw_state=purged、署名照记；同一份内容在别的导入里的
    引用一并清除（also_purged），引用它的未完成导入一并放弃；那几期构建的隔离文件一并删掉；清除不影响累积：下一期照常
    追加（期 3 不重放历史，2.5）。"""
    v = await august_september(client, seed=301)
    sid = v["source_id"]
    aug = await db_get(TableImport, v["aug"])
    # 8 月的构建曾被改过、已恢复：隔离区里有它的一份
    build = await db_get(TableBuild, aug.build_id)
    q = store / "uploads" / "quarantine" / sid
    q.mkdir(parents=True, exist_ok=True)
    stale = q / f"{aug.build_id}-20260101T000000Z.db"
    stale.write_bytes(b"quarantined copy")
    # 引用同一份 8 月原件的未完成导入（原样再传、还没提交）
    open_st = await reupload(client, sid, v["raw"], v["file_name"])
    listed = {r["id"]: r for r in await imports_api(client, sid)}
    assert listed[v["aug"]]["raw_open_stagings"] == 1
    url = f"{API}/{sid}/imports/{v['aug']}/purge-raw"
    resp = await client.post(url, json={"confirm": True, "reason": "  "})
    assert resp.status_code == 422 and resp.json()["code"] == "reason_required"
    resp = await client.post(url, json={"confirm": True, "reason": REASON, "signed_by": "合成署名丙"})
    assert resp.status_code == 200, resp.text
    out = resp.json()
    assert out["file_deleted"] is True and out["discarded_stagings"] == [open_st["id"]]
    assert out["import"]["raw_state"] == "purged" and out["import"]["purged"]["signed_by"] == "合成署名丙"
    assert not raw_store.raw_exists(aug.raw_sha256) and not stale.exists()
    assert Path(build.db_path).is_file(), "构建库（数据）照旧，只清原件"
    assert (await db_get(ImportStaging, open_st["id"])).status == "discarded"
    resp = await client.post(url, json={"confirm": True, "reason": REASON})
    assert resp.status_code == 409 and resp.json()["code"] == "raw_already_purged"
    # 不在当前版本里的导入：替换掉 9 月之后，旧 9 月的原件照样能清
    raw, fn = flow_workbook(SEP, 30, seed=304)
    replaced = await add_period(client, sid, raw, fn)
    assert (await db_get(TableImport, v["sep"])).status == "superseded"
    resp = await client.post(f"{API}/{sid}/imports/{v['sep']}/purge-raw", json={"confirm": True, "reason": REASON})
    assert resp.status_code == 200, resp.text
    # 清除不影响累积
    raw, fn = flow_workbook(OCT, 31, seed=303)
    out = await add_period(client, sid, raw, fn)
    assert out["parts"] == 3 and out["snapshot_id"] != replaced["snapshot_id"]
    assert await count(client, v["name"], "日客流") == 92


# ==========================================================================
# 遗留项 L1：被改过的构建文件（能重建就重建，否则明确报错）
# ==========================================================================


def simple_sheet(seed: int) -> bytes:
    """一份最简单的表格（期 1 的简单上传用）：表头加两行，数字随机。"""
    import io
    import random

    rnd = random.Random(seed)
    book = Workbook()
    ws = book.active
    ws.title = "甲表"
    ws.append(["分区", "数量"])
    ws.append(["分区甲", rnd.randint(1, 999)])
    ws.append(["分区乙", rnd.randint(1, 999)])
    buf = io.BytesIO()
    book.save(buf)
    return buf.getvalue()


def corrupt(path: Path, data: bytes = b"") -> None:
    os.chmod(path, 0o644)
    path.write_bytes(data)


def quarantined(store: Path, source_id: str) -> list[Path]:
    return sorted((store / "uploads" / "quarantine" / source_id).rglob("*.db"))


async def test_l1_upload_restores_a_tampered_build_and_cleans_the_quarantine(client, store):
    """L1（/upload）：同一文件重传、构建文件被改成 0 字节 → 201，build_restored=true，quarantine/<源 id>/ 里有原文件，
    快照的文件哈希恢复；构建记录的哈希也被改掉时 → 409 build_conflict，指针不动；隔离文件超过 7 天被回收删除；删除
    数据源后隔离目录被删。"""
    name, raw = unique(), simple_sheet(201)
    resp = await client.post(f"{API}/upload", files={"file": ("甲表.xlsx", raw, XLSX)}, data={"name": name})
    assert resp.status_code == 201 and resp.json()["build_restored"] is False, resp.text
    first = resp.json()
    sid = first["source"]["id"]
    snap = await db_get(SourceSnapshot, first["snapshot_id"])
    corrupt(Path(snap.db_path))
    resp = await client.post(f"{API}/upload", files={"file": ("甲表.xlsx", raw, XLSX)}, data={"name": name})
    assert resp.status_code == 201, resp.text
    assert resp.json()["build_restored"] is True
    [moved] = quarantined(store, sid)
    assert moved.stat().st_size == 0, "隔离的是被改过的那份"
    assert raw_store.sha256_file(Path(snap.db_path)) == snap.db_sha256, "快照的文件哈希恢复"
    # 记录的哈希也被改掉：无法用本次上传恢复 → 409 build_conflict，指针不动
    imp = await db_get(TableImport, first["import_id"])
    pointer = await current_id(sid)
    corrupt(Path(snap.db_path), b"tampered again")
    async with SessionLocal() as session:
        await session.execute(update(TableBuild).where(TableBuild.id == imp.build_id).values(db_sha256="0" * 64))
        await session.commit()
    resp = await client.post(f"{API}/upload", files={"file": ("甲表.xlsx", raw, XLSX)}, data={"name": name})
    assert resp.status_code == 409 and resp.json()["code"] == "build_conflict", resp.text
    assert await current_id(sid) == pointer
    # 隔离文件超过 QUARANTINE_DAYS 天：回收删掉
    old = time.time() - (table_versions.QUARANTINE_DAYS + 1) * 86400
    os.utime(moved, (old, old))
    async with SessionLocal() as session:
        await table_versions.gc(session)
    assert not moved.exists()
    # 删除数据源：隔离目录一并删掉
    corrupt(Path(snap.db_path), b"")
    async with SessionLocal() as session:
        await session.execute(update(TableBuild).where(TableBuild.id == imp.build_id)
                              .values(db_sha256=snap.db_sha256))
        await session.commit()
    resp = await client.post(f"{API}/upload", files={"file": ("甲表.xlsx", raw, XLSX)}, data={"name": name})
    assert resp.status_code == 201 and resp.json()["build_restored"] is True, resp.text
    assert quarantined(store, sid)
    resp = await client.delete(f"{API}/{sid}")
    assert resp.status_code == 204, resp.text
    assert not (store / "uploads" / "quarantine" / sid).exists()


async def test_l1_recipe_commit_restores_conflicts_and_quarantines_an_orphan(client, store):
    """L1（配方提交）：8 月（每期替换）→ 9 月 → 8 月原文件再传（构建与第一次相同），第一次的构建文件被改成 0 字节 →
    201、build_restored=true；构建记录的哈希也被改掉 → 409 build_conflict、指针不动；目标位置有一个没人登记、内容不同的
    文件（孤儿）→ 隔离后 link，发布成功。"""
    base = await first_import(client, seed=211, mode="replace")
    sid = base["source_id"]
    raw9, fn9 = flow_workbook(SEP, 30, seed=212)
    await add_period(client, sid, raw9, fn9)
    build = await db_get(TableBuild, (await db_get(TableImport, base["import_id"])).build_id)
    corrupt(Path(build.db_path))
    st = await reupload(client, sid, base["raw"], base["file_name"])
    assert st["trial"]["status"] == "passed", st["trial"]["problems"]
    resp = await commit(client, st)
    assert resp.status_code == 201, resp.text
    assert resp.json()["build_restored"] is True and resp.json()["build_id"] == build.id
    assert raw_store.sha256_file(Path(build.db_path)) == build.db_sha256
    assert len(quarantined(store, sid)) == 1

    # 记录的哈希也被改掉：9 月原文件再传，那份构建无法恢复
    sep_build = await db_get(TableBuild, (await db_get(TableImport, (await snapshots(client, sid))[1]["parts"][0]
                                                          ["import_id"])).build_id)
    corrupt(Path(sep_build.db_path))
    async with SessionLocal() as session:
        await session.execute(update(TableBuild).where(TableBuild.id == sep_build.id).values(db_sha256="0" * 64))
        await session.commit()
    pointer = await current_id(sid)
    st = await reupload(client, sid, raw9, fn9)
    resp = await commit(client, st)
    assert resp.status_code == 409 and resp.json()["code"] == "build_conflict", resp.text
    assert await current_id(sid) == pointer

    # 孤儿：10 月的构建位置上已经有一个没人登记的文件
    raw10, fn10 = flow_workbook(OCT, 31, seed=213)
    st = await reupload(client, sid, raw10, fn10)
    build_id = (await db_get(ImportStaging, st["id"])).trial["build_id"]
    orphan = table_versions.build_path(sid, build_id)
    orphan.parent.mkdir(parents=True, exist_ok=True)
    orphan.write_bytes(b"orphan nobody registered")
    resp = await commit(client, st)
    assert resp.status_code == 201, resp.text
    assert resp.json()["build_restored"] is False
    assert raw_store.sha256_file(orphan) == (await db_get(TableBuild, build_id)).db_sha256
    assert any(p.read_bytes() == b"orphan nobody registered" for p in quarantined(store, sid))


# ==========================================================================
# 遗留项 L2：改配方后回答延续
# ==========================================================================


async def test_l2_answers_survive_a_put_that_embodies_them(client):
    """L2：先回答 q_placeholder:·，再 PUT 一份体现了这个回答的配方 → answers 里仍有它，answers_dropped 为空；PUT 一份
    与回答矛盾的配方（去掉占位符）→ answers_dropped 有它（带问题文字）。"""
    raw, fn = flow_workbook(AUG, 31, seed=221)
    st = await stage(client, raw, fn)
    qid = "q_placeholder:·"
    st = await answer(client, st["id"], {qid: {"value": "null"}})
    st = await put_recipe(client, st["id"], st["recipe"])
    assert st["answers"].get(qid, {}).get("value") == "null" and st["answers_dropped"] == []
    rec = json.loads(json.dumps(st["recipe"], ensure_ascii=False))
    for sheet in rec["sheets"]:
        for block in sheet["blocks"]:
            if isinstance(block.get("values"), dict):
                block["values"]["placeholders"] = []
    st = await put_recipe(client, st["id"], rec)
    dropped = {d["id"]: d for d in st["answers_dropped"]}
    assert qid in dropped and dropped[qid]["text"] and dropped[qid]["reason"] in ("changed", "gone"), dropped
    assert qid not in st["answers"]


# ==========================================================================
# 遗留项 L3：沿用已确认的常量（3.3），经接口
# ==========================================================================


async def test_l3_keeping_a_confirmed_constant_follows_the_three_rules(client):
    """L3（3.3）：D09 改了日间的分段标题。(b) 沿用：base 里同一分段的 pick「日间」仍是新标题的子串 → 静态校验通过；
    反例：换成不在候选词里、base 里也没有的「时段」→ const_not_candidate；含数字的 pick → period_literal；(a) 候选词
    「日间分」→ 通过（但 breaking）；首次导入没有 base 时沿用不成立 → const_not_candidate；(c) 提交后下一期按重放校验
    （origin=replay，没有 base），子串成立 → 通过。"""
    base = await first_import(client, seed=231)
    sid = base["source_id"]
    raw, fn = DRIFT_CASES["D09"].build(232)
    st = await reupload(client, sid, raw, fn)

    def with_title(rec: dict[str, Any], pick: str) -> dict[str, Any]:
        rec = json.loads(json.dumps(rec, ensure_ascii=False))
        seg = next(s for s in rec["sheets"][0]["blocks"][0]["segments"] if s["id"] == "日间")
        seg["locate"]["title"] = "日间分时段客流（人次）"
        seg["const"]["时段类别"]["pick"] = pick
        return rec

    codes = lambda out: {p["code"] for p in out["recipe_problems"]}  # noqa: E731
    out = await put_recipe(client, st["id"], with_title(st["recipe"], "时段"))
    assert "const_not_candidate" in codes(out), out["recipe_problems"]
    out = await put_recipe(client, st["id"], with_title(st["recipe"], "9月"))
    assert "period_literal" in codes(out), out["recipe_problems"]
    out = await put_recipe(client, st["id"], with_title(st["recipe"], "日间分"))
    assert out["recipe_problems"] == [], "(a) 候选词照旧放行"
    kept = with_title(st["recipe"], "日间")
    out = await put_recipe(client, st["id"], kept)
    assert out["recipe_problems"] == [], "(b) 沿用：base 里同一分段的 pick，仍是新标题的子串"
    st = await trial(client, out["id"])
    assert st["trial"]["status"] == "passed", st["trial"]["problems"]
    assert not [i for i in confirm_ids(st) if i.startswith("breaking:")], confirm_ids(st)
    resp = await commit(client, st)
    assert resp.status_code == 201, resp.text
    # 首次导入没有 base：沿用不成立
    other = await stage(client, raw, fn)
    out = await put_recipe(client, other["id"], rename_wide(kept, "日客流", "日客流"))
    assert "const_not_candidate" in codes(out), out["recipe_problems"]
    # (c) 下一期按重放校验：子串成立就放行
    raw10, fn10 = flow_workbook(OCT, 31, seed=233, variant="D09")
    st = await reupload(client, sid, raw10, fn10)
    assert st["recipe_problems"] == [] and st["trial"]["status"] == "passed", (st["recipe_problems"],
                                                                                st["trial"]["problems"])


# ==========================================================================
# 遗留项 L4：排除的行（rows_excluded）；列表只剩表头必须确认（empty_block）
# ==========================================================================


def list_book(rows: list[tuple[Any, ...]], *, sheet: str = "明细", extra: list[tuple[Any, ...]] | None = None
              ) -> tuple[bytes, str]:
    """一张列表：表头「地区、产品、金额」加 rows（None 行即空行），再接 extra。"""
    book = Workbook()
    ws = book.active
    ws.title = sheet
    for row in [("地区", "产品", "金额"), *rows, *(extra or [])]:
        ws.append(list(row))
    return excel_saved(save(book), sheet, {}), f"{sheet}.xlsx"


def excluded_reasons(doc: dict[str, Any]) -> set[str]:
    return {x["reason"] for x in doc["rows_excluded"]}


async def _commit_and_manifest(client, st: dict[str, Any]) -> dict[str, Any]:
    resp = await commit(client, st)
    assert resp.status_code == 201, resp.text
    imp = await db_get(TableImport, resp.json()["import_id"])
    return artifact_store.load(imp.manifest_artifact)


async def test_l4_rows_excluded_reach_the_receipt_and_the_import_manifest(client):
    """L4：c06 exclude → hidden_excluded；c08 gap skip → blank_skipped；D21 忽略 → ignored_rows（A6b 已提交过一次，这里
    再看清单）；列表停止后的文字行 → after_stop。trial.receipt.rows_excluded 经接口给出，导入清单的 receipt 里也有。"""
    # c06：隐藏行不导入
    raw, fn = lab.c06()
    st = await stage(client, raw, fn)
    q = next(q for q in st["questions"] if q["id"].startswith("q_hidden:"))
    st = await answer(client, st["id"], {q["id"]: {"value": "exclude"}})
    st = await trial(client, st["id"])
    assert st["trial"]["status"] == "passed", st["trial"]["problems"]
    assert "hidden_excluded" in excluded_reasons(st["trial"]["receipt"])
    assert "hidden_excluded" in excluded_reasons((await _commit_and_manifest(client, st))["receipt"])
    # c08 gap：中间空行跳过
    raw, fn = lab.c08("gap")
    st = await trial(client, (await stage(client, raw, fn))["id"])
    assert st["trial"]["status"] == "passed"
    gap = [x for x in st["trial"]["receipt"]["rows_excluded"] if x["reason"] == "blank_skipped"]
    assert gap and gap[0]["rows"] == [[4, 4]], st["trial"]["receipt"]["rows_excluded"]
    assert "blank_skipped" in excluded_reasons((await _commit_and_manifest(client, st))["receipt"])
    # 列表停止后的文字行：空行之后还有两格文字（像数据，没导入）→ rows_after_stop 需确认，记 after_stop
    raw, fn = list_book([("分区甲", "产品甲", 10), ("分区乙", "产品乙", 20), (None, None, None),
                         ("补充说明", "见附件甲", None)])
    st = await stage(client, raw, fn)
    rec = json.loads(json.dumps(st["recipe"], ensure_ascii=False))
    rec["sheets"][0]["blocks"][0]["rows"]["blank_rows"] = "stop"
    st = await trial(client, (await put_recipe(client, st["id"], rec))["id"])
    assert st["trial"]["status"] == "passed", st["trial"]["problems"]
    assert problems(st, "rows_after_stop") and "after_stop" in excluded_reasons(st["trial"]["receipt"])
    assert "after_stop" in excluded_reasons((await _commit_and_manifest(client, st))["receipt"])
    # D21：按行标签忽略
    done = await _fix_case(client, "D21", seed=241)
    assert "ignored_rows" in excluded_reasons(done["st"]["trial"]["receipt"])
    assert "ignored_rows" in excluded_reasons((await _commit_and_manifest(client, done["st"]))["receipt"])


async def test_l4_a_list_that_drops_to_only_its_header_must_be_confirmed(client):
    """遗留项 4 原意（期 2 AU-1 修复③）：上一期有数据行、本期只剩表头的列表重传，必须出 empty_block:<块> 必勾项，
    不勾不能启用；勾上之后这张表是 0 行。"""
    raw, fn = list_book([("分区甲", "产品甲", 10), ("分区乙", "产品乙", 20), ("分区丙", "产品丙", 30)])
    st = await trial(client, (await stage(client, raw, fn))["id"])
    assert st["trial"]["status"] == "passed", st["trial"]["problems"]
    resp = await commit(client, st)
    assert resp.status_code == 201, resp.text
    sid, name = resp.json()["source"]["id"], resp.json()["source"]["name"]
    table = st["recipe"]["sheets"][0]["blocks"][0]["table"]
    block = st["recipe"]["sheets"][0]["blocks"][0]["id"]
    raw, fn = list_book([])
    st = await reupload(client, sid, raw, fn)
    t = st["trial"]
    assert t["status"] == "passed", t["problems"]
    assert problems(st, "list_empty"), t["problems"]
    item = f"empty_block:{block}"
    items = {i["id"]: i for i in t["confirm_items"]}
    assert item in items and items[item]["required"] is True, list(items)
    resp = await commit(client, st, confirmations=[i for i in confirm_ids(st) if i != item])
    assert resp.status_code == 422 and resp.json()["code"] == "confirm_required" and item in resp.json()["detail"]
    resp = await commit(client, st)
    assert resp.status_code == 201, resp.text
    assert await count(client, name, table) == 0


# ==========================================================================
# 遗留项 L5：另一个进程持有存储守卫
# ==========================================================================

_HOLD = """
import fcntl, os, sys
fd = os.open(sys.argv[1], os.O_RDWR | os.O_CREAT, 0o600)
fcntl.flock(fd, fcntl.LOCK_EX)
print("locked", flush=True)
sys.stdin.read()
"""


async def test_l5_another_process_holding_the_store_guard(client, store):
    """L5：子进程先占住 .store.lock → 本进程 acquire_store_guard() 返回 False；本进程调 startup：事先放好的临时文件和孤儿
    构建库一个都没删；/upload、stage、trial、commit、activate 返回 503 store_unavailable，GET 接口照常；过期处理不抛；
    子进程退出后重新调用返回 True，startup 照常清理。"""
    v = await august_september(client, seed=251)
    sid = v["source_id"]
    raw10, fn10 = flow_workbook(OCT, 31, seed=252)
    st = await reupload(client, sid, raw10, fn10)
    half = table_versions.build_path(sid, "1" * 64).with_name(f"{'1' * 64}.db{raw_store.TMP_MARK}zzz")
    half.write_bytes(b"half written")
    orphan = table_versions.build_path(sid, "2" * 64)
    orphan.write_bytes(b"orphan build")
    proc = subprocess.Popen([sys.executable, "-c", _HOLD, str(settings.uploads_dir / ".store.lock")],
                            stdin=subprocess.PIPE, stdout=subprocess.PIPE, text=True)
    try:
        assert proc.stdout.readline().strip() == "locked"
        assert table_versions.acquire_store_guard() is False
        ran: list[str] = []
        real_gc = table_versions.gc

        async def spy_gc(session):
            ran.append("gc")

        table_versions.gc = spy_gc
        try:
            async with SessionLocal() as session:
                await table_versions.startup(session)
        finally:
            table_versions.gc = real_gc
        assert ran == [] and half.exists() and orphan.exists(), "拿不到守卫：startup 一个文件都不删，不回收"

        def refused(resp) -> None:
            assert resp.status_code == 503, resp.text
            assert resp.json()["code"] == "store_unavailable"

        refused(await client.post(f"{API}/upload", files={"file": ("甲表.xlsx", simple_sheet(253), XLSX)},
                                  data={"name": unique()}))
        refused(await client.post(f"{API}/imports/stage", files={"file": (fn10, raw10, XLSX)},
                                  data={"name": unique()}))
        refused(await client.post(f"{API}/imports/{st['id']}/trial", json={}))
        refused(await commit(client, st))
        refused(await client.post(f"{API}/{sid}/snapshots/{v['s1']}/activate",
                                  json={"confirm": True, "expected_current_snapshot_id": v["s2"]}))
        assert (await client.get(API)).status_code == 200
        assert (await client.get(f"{API}/{sid}/snapshots")).status_code == 200
        assert (await client.get(f"{API}/{sid}/imports")).status_code == 200
        assert (await client.get(f"{API}/imports/{st['id']}")).status_code == 200
        assert await count(client, v["name"], "日客流") == 61
        async with SessionLocal() as session:
            assert await table_versions.expire_stagings(session) == []
        assert await current_id(sid) == v["s2"]
    finally:
        proc.stdin.close()
        proc.wait(timeout=10)
    try:
        assert table_versions.acquire_store_guard() is True
        async with SessionLocal() as session:
            await table_versions.startup(session)
        assert not half.exists() and not orphan.exists()
        resp = await commit(client, st)
        assert resp.status_code == 201 and resp.json()["parts"] == 3, resp.text
    finally:
        table_versions._release_guard()


# ==========================================================================
# 11.4 数据变异
# ==========================================================================


def shifted_period(raw: bytes, start, days: int) -> bytes:
    """把 9 月 D01 的统计期（B2）和日期表头（第 4 行）一起改成从 start 起 days 天：数字照旧，只挪统计期。"""
    import datetime as dt

    from openpyxl.utils import get_column_letter

    end = start + dt.timedelta(days=days - 1)
    edits: dict[str, Any] = {"B2": f"统计时间范围：{start.year}年{start.month}月{start.day}日至"
                                   f"{end.year}年{end.month}月{end.day}日"}
    for i in range(days):
        d = start + dt.timedelta(days=i)
        edits[f"{get_column_letter(3 + i)}4"] = f"{d.month}月{d.day}日"
    return edit_cells(raw, edits)


async def test_mutation_september_period_overlapping_august_by_one_day(client):
    """11.4 第 1 行：9 月 D01 的 B2 改成 08-31 至 09-29（与 8 月重叠一天）→ period_overlap，当前快照不动。日期表头跟着挪
    （否则先撞上「日期不在统计期内」的结构问题，同样拒收，只是原因不同，见下半）。"""
    base = await first_import(client, seed=261)
    raw, fn = flow_workbook(SEP, 30, seed=262)
    st = await reupload(client, base["source_id"], shifted_period(raw, date_of(2026, 8, 31), 30), fn)
    assert st["trial"]["status"] == "rejected"
    [overlap] = problems(st, "period_overlap")
    assert "2026-08-31 至 2026-09-29" in overlap["message"] and "2026-08-01 至 2026-08-31" in overlap["message"]
    resp = await commit(client, st)
    assert resp.status_code == 409 and resp.json()["code"] == "trial_not_passed"
    await client.delete(f"{API}/imports/{st['id']}")
    # 只改 B2、表头不动：9 月 30 日落在统计期之外，执行器先拒收（axis_year_missing），同样不动当前快照
    bad = edit_cells(raw, {"B2": "统计时间范围：2026年8月31日至2026年9月29日"})
    st = await reupload(client, base["source_id"], bad, fn)
    assert st["trial"]["status"] == "rejected" and problems(st, "axis_year_missing")
    assert await current_id(base["source_id"]) == base["s1"]


async def test_mutation_trial_files_tampered_or_deleted_before_commit(client):
    """11.4 第 2、3 行：本期构建的试运行库在提交前被改字节 → 409 trial_tampered；并集试运行库被改 → 409 union_tampered；
    并集试运行库被删 → 409 trial_required，暂存区退回 drafting。当前快照都不动。"""
    base = await first_import(client, seed=271)
    sid = base["source_id"]
    raw, fn = flow_workbook(SEP, 30, seed=272)
    st = await reupload(client, sid, raw, fn)
    row = await db_get(ImportStaging, st["id"])
    main = Path(row.trial_path)
    with open(main, "ab") as f:
        f.write(b"\0")
    resp = await commit(client, st)
    assert resp.status_code == 409 and resp.json()["code"] == "trial_tampered", resp.text
    st = await trial(client, st["id"])
    union = table_versions.union_trial_path(Path((await db_get(ImportStaging, st["id"])).trial_path))
    with open(union, "ab") as f:
        f.write(b"\0")
    resp = await commit(client, st)
    assert resp.status_code == 409 and resp.json()["code"] == "union_tampered", resp.text
    st = await trial(client, st["id"])
    union = table_versions.union_trial_path(Path((await db_get(ImportStaging, st["id"])).trial_path))
    union.unlink()
    resp = await commit(client, st)
    assert resp.status_code == 409 and resp.json()["code"] == "trial_required", resp.text
    assert (await db_get(ImportStaging, st["id"])).status == "drafting"
    assert await current_id(sid) == base["s1"]


async def test_mutation_published_part_or_its_period_column_tampered(client):
    """11.4 第 4、5 行：已发布的某一期构建库被改字节，再上传下一期 → 试运行 rejected、part_tampered；某一期导入记录的
    period_start 库列被改 → rejected、part_tampered（统计期取自导入清单）。"""
    v = await august_september(client, seed=281)
    sid = v["source_id"]
    build = await db_get(TableBuild, (await db_get(TableImport, v["aug"])).build_id)
    path = Path(build.db_path)
    original = path.read_bytes()
    corrupt(path, original[:-1] + bytes([original[-1] ^ 1]))
    raw, fn = flow_workbook(OCT, 31, seed=282)
    st = await reupload(client, sid, raw, fn)
    assert st["trial"]["status"] == "rejected" and problems(st, "part_tampered"), st["trial"]["problems"]
    corrupt(path, original)
    os.chmod(path, 0o444)
    async with SessionLocal() as session:
        await session.execute(update(TableImport).where(TableImport.id == v["sep"]).values(period_start="2026-09-02"))
        await session.commit()
    st = await trial(client, st["id"])
    assert st["trial"]["status"] == "rejected"
    tampered = problems(st, "part_tampered")
    assert tampered and "统计期" in tampered[0]["message"], st["trial"]["problems"]
    assert await current_id(sid) == v["s2"]


async def test_mutation_snapshot_hash_column_rewritten_to_match_a_swapped_file(client):
    """11.4 第 6 行：并集快照的库文件被换掉，source_snapshots.db_sha256 库列也改成与新文件一致的假值（文件哈希核对照样
    通过）→ 查询时拿快照清单交叉核对，报 snapshot_tampered（版本清单不一致），不会查到被换掉的数据。"""
    v = await august_september(client, seed=291)
    snap = await db_get(SourceSnapshot, v["s2"])
    path = Path(snap.db_path)
    os.chmod(path, 0o644)
    with closing(sqlite3.connect(path)) as conn:
        conn.execute('UPDATE "日客流" SET "全日客流" = 0')
        conn.commit()
    table_versions.settle_journal(path)
    fake = raw_store.sha256_file(path)
    async with SessionLocal() as session:
        await session.execute(update(SourceSnapshot).where(SourceSnapshot.id == v["s2"]).values(db_sha256=fake))
        await session.commit()
    from app.data.engine import engines

    await engines.invalidate(v["source_id"])
    resp = await client.post(f"/api/tools/db_query__{v['name']}/run",
                             json={"args": {"sql": 'SELECT SUM("全日客流") FROM "日客流"'}})
    assert resp.status_code == 200, resp.text
    out = resp.json()
    # 工具把数据不可用写成给模型看的结果（不抛），内容是拒绝查询的原因，没有数
    assert "查询失败（数据不可用）" in out["result"] and "版本清单不一致" in out["result"], out
    with pytest.raises(table_versions.SnapshotManifestMismatch) as err:
        async with SessionLocal() as session:
            await table_versions.resolve_source(session, await session.get(DataSource, v["source_id"]))
    assert err.value.code == "snapshot_tampered"


# ==========================================================================
# WP-8 修补 3：跳过空行的列表，表下隔着空行的说明不是数据（经真管线）
# ==========================================================================


async def test_fix3_note_under_a_list_with_a_middle_gap_stays_outside(client):
    """列表中间有一个空行（起草成 blank_rows=skip），表下再隔一个空行有「注：……」：试运行 4 行，「注：……」进区域外文字，
    跳过的空行只有中间那一行；提交后表里没有首列是「注：……」的记录。框选整块（含中间空行）重放与框一致。曾静默多读进
    一行首列是说明、其余为空值的记录（p3w1a/fix-2.md「需要你决定的事项」第 5 条）。"""
    raw, fn = list_book([("分区甲", "产品甲", 10), ("分区甲", "产品乙", 20), (None, None, None),
                         ("分区乙", "产品甲", 30), ("分区乙", "产品乙", 40)],
                        extra=[(None, None, None), ("注：以上为初步统计，以终稿为准。", None, None)])
    st = await stage(client, raw, fn)
    block = st["recipe"]["sheets"][0]["blocks"][0]
    assert block["rows"]["blank_rows"] == "skip"
    sel = {"sheet": "明细", "ref": "A1:C6", "as": "list", "options": {"header_rows": 1, "bottom": "box"}}
    pv = (await preview(client, st["id"], {"selection": sel})).json()
    assert pv["ok"] is True and pv["replay"]["match"] is True, (pv["problems"], pv["replay"])
    st = await trial(client, st["id"])
    t = st["trial"]
    assert t["status"] == "passed", t["problems"]
    assert rows_of(st)[block["table"]] == 4
    assert ("明细!A8", "注：以上为初步统计，以终稿为准。") in [(o["cell"], o["text"]) for o in t["receipt"]["outside_text"]]
    # 列表停在哪一行、因为什么也看得出来（评审意见 4）：说明那一行记进排除的行（after_stop，1 格，锚点是说明文字）；
    # 像说明，不出 rows_after_stop
    note = "注：以上为初步统计，以终稿为准。"
    assert [(x["reason"], x["rows"], x["cells"], x["anchor"]) for x in t["receipt"]["rows_excluded"]] == [
        ("blank_skipped", [[4, 4]], 0, None), ("after_stop", [[8, 8]], 1, note)]
    assert t["receipt"]["blank_rows_skipped"] == 1
    assert not problems(st, "rows_after_stop") and not [i for i in confirm_ids(st) if i.startswith("rows_after_stop")]
    assert "空行之后" not in t["notes"][block["table"]]["comment"]
    resp = await commit(client, st)
    assert resp.status_code == 201, resp.text
    name = resp.json()["source"]["name"]
    doc = artifact_store.load((await db_get(TableImport, resp.json()["import_id"])).manifest_artifact)
    assert ("after_stop", [[8, 8]], note) in [(x["reason"], x["rows"], x["anchor"]) for x in doc["receipt"]["rows_excluded"]]
    assert await count(client, name, block["table"]) == 4
    assert await count(client, name, block["table"], ' WHERE "地区" LIKE \'注%\'') == 0


def _lone_tail_book(seed: int) -> tuple[bytes, str]:
    """起草成 blank_rows=skip 的列表：中间一个空行，表尾隔一个空行是一条只填了首列的「分区丁」（评审意见 4 的探针）。
    金额随 seed 变，地区和产品不变。"""
    import random

    rnd = random.Random(seed)
    return list_book([("分区甲", "产品甲", rnd.randint(10, 99)), ("分区乙", "产品甲", rnd.randint(10, 99)),
                      (None, None, None), ("分区丙", "产品甲", rnd.randint(10, 99)), (None, None, None),
                      ("分区丁", None, None)])


async def test_fix3_a_lone_first_column_row_that_is_not_a_note_must_be_ticked_every_time(client):
    """评审意见 4：表尾只填了首列的「分区丁」照旧当表的下边界（与起草器一致，不导入），但它不像说明：出 confirm 类的
    rows_after_stop（「第 7 行只有首列文字，已当作表下说明，没有导入」），必勾项 rows_after_stop:<块>，不勾不能提交；
    排除的行记 after_stop（锚点「分区丁」）。下一期重放同样的表尾，照样要勾。曾只出现在区域外文字里：确认项只有
    blank_skip 和 mode，排除的行里也没有它。"""
    raw, fn = _lone_tail_book(711)
    st = await stage(client, raw, fn)
    block = st["recipe"]["sheets"][0]["blocks"][0]
    assert block["rows"]["blank_rows"] == "skip"
    st = await trial(client, st["id"])
    t = st["trial"]
    assert t["status"] == "passed", t["problems"]
    assert rows_of(st)[block["table"]] == 3
    found = problems(st, "rows_after_stop")
    assert len(found) == 1 and found[0]["category"] == "confirm" and found[0]["cells"] == ["明细!A7"], found
    assert "第 7 行只有首列文字" in found[0]["message"] and "已当作表下说明，没有导入" in found[0]["message"]
    assert ("after_stop", [[7, 7]], 1, "分区丁") in [(x["reason"], x["rows"], x["cells"], x["anchor"])
                                                   for x in t["receipt"]["rows_excluded"]]
    item = f"rows_after_stop:{block['id']}"
    items = {i["id"]: i for i in t["confirm_items"]}
    assert item in items and items[item]["required"] is True, list(items)
    assert "第 7 行只有首列文字" in items[item]["label"]
    assert "原表在空行之后另有未导入的文字行" in t["notes"][block["table"]]["comment"]
    resp = await commit(client, st, confirmations=[i for i in confirm_ids(st) if i != item])
    assert resp.status_code == 422 and resp.json()["code"] == "confirm_required" and item in resp.json()["detail"]
    resp = await commit(client, st)
    assert resp.status_code == 201, resp.text
    sid, name = resp.json()["source"]["id"], resp.json()["source"]["name"]
    assert await count(client, name, block["table"]) == 3
    assert await count(client, name, block["table"], ' WHERE "地区" = \'分区丁\'') == 0
    # 下一期（数字不同、表尾照旧）：同样出这一项，不勾照样不能提交
    raw, fn = _lone_tail_book(712)
    st = await reupload(client, sid, raw, fn)
    t = st["trial"]
    assert t["status"] == "passed", t["problems"]
    assert item in confirm_ids(st), confirm_ids(st)
    resp = await commit(client, st, confirmations=[i for i in confirm_ids(st) if i != item])
    assert resp.status_code == 422 and resp.json()["code"] == "confirm_required" and item in resp.json()["detail"]
    resp = await commit(client, st)
    assert resp.status_code == 201, resp.text
    assert await count(client, name, block["table"]) == 3


async def test_fix3_a_build_from_the_old_engine_is_not_reused_for_the_same_file(client, monkeypatch):
    """评审意见 1：修补 3 改了「列表 + 跳过空行」的执行结果，RECIPE_ENGINE_VER 没升（升了交叉表配方的构建全部换 id），
    改为这类配方的构建 id 带执行器语义标记（table_versions.recipe_engine_semantics）。先用修补前的执行器（_note_after_blank
    恒 False、构建 id 不带标记）发布一次：5 行，其中一条首列是「注：……」。恢复后把同一文件换个名字重传：试运行 4 行，
    提交不复用旧构建（换了 id），发布出去的库就是这次试运行的库，表里 4 行。曾复用旧构建：核对的 4 行、发布的 5 行。"""
    from app.data import recipe_engine

    note = "注：以上为初步统计，以终稿为准。"
    raw, fn = list_book([("分区甲", "产品甲", 31), ("分区甲", "产品乙", 47), (None, None, None),
                         ("分区乙", "产品甲", 52), ("分区乙", "产品乙", 68)],
                        extra=[(None, None, None), (note, None, None)])
    with monkeypatch.context() as old:
        old.setattr(recipe_engine._ListRun, "_note_after_blank", lambda self, inb: False)
        old.setattr(table_versions, "recipe_engine_semantics", lambda recipe: {})
        st = await stage(client, raw, fn)
        block = st["recipe"]["sheets"][0]["blocks"][0]
        assert block["rows"]["blank_rows"] == "skip"
        st = await trial(client, st["id"])
        assert st["trial"]["status"] == "passed", st["trial"]["problems"]
        assert rows_of(st)[block["table"]] == 5
        resp = await commit(client, st)
        assert resp.status_code == 201, resp.text
    first = resp.json()
    sid, name, table = first["source"]["id"], first["source"]["name"], block["table"]
    assert await count(client, name, table, ' WHERE "地区" LIKE \'注%\'') == 1
    old_build = await db_get(TableBuild, first["build_id"])
    assert "engine_semantics" not in old_build.options

    st = await reupload(client, sid, raw, "另一个名字.xlsx")
    t = st["trial"]
    assert t["status"] == "passed", t["problems"]
    assert rows_of(st)[table] == 4
    resp = await commit(client, st)
    assert resp.status_code == 201, resp.text
    out = resp.json()
    assert out["build_reused"] is False and out["build_id"] != first["build_id"]
    build = await db_get(TableBuild, out["build_id"])
    assert build.db_sha256 == t["receipt"]["db_sha256"], "发布出去的库就是这次试运行、用户核对过的那一份"
    assert build.options["engine_semantics"] == {table_versions.LIST_NOTE_AFTER_BLANK: 1}
    assert build.options_sha256 == table_versions.sha_json(build.options)
    assert await count(client, name, table) == 4
    assert await count(client, name, table, ' WHERE "地区" LIKE \'注%\'') == 0
