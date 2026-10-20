"""交叉表里按文字导入的指标（MeasuresSegment.text_measures）：契约、静态校验、执行、规则起草、修复按钮、确认项、说明。

起因：一份客流导出的顶部指标块里夹着「高峰时段」一行，每格写的是当天客流最多的时段（「18-19」），不是数字。
以前起草逐格报「不是数字」（一行 30 条），提示去加占位符；照做的话「18-19」整格静默存成空值、干跑照样通过；
从分段里去掉这一行再「忽略这些行」又不生效（按标签定位的分段认领连续的数据行）。

夹具在这里用 openpyxl 现造，名字全是假名（分区甲 / 分区乙）：B2 统计期、B3 标题、第 4 行日期、第 5–9 行无标题的
指标块（全日、分区甲、分区乙、高峰时段、时段人数）、空行、日间和夜间两段时段客流，没有合计行。第 4 天的时段全是「·」，
那天的高峰时段、时段人数也是「·」；第 2 天的高峰时段用 en dash 写（「18–19」）。
"""
from __future__ import annotations

import copy
import datetime as dt
import io
import random
import sqlite3
from typing import Any, Callable

import pytest
from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from app.data import recipe_engine, xlsx_cells, xlsx_scan
from app.data import recipe_fixes as F
from app.data import recipe_suggest as S
from app.data.recipe import apply_patch, canonical_recipe, validate_recipe
from app.data.recipe_confirm import ConfirmContext, confirm_items, table_changes
from app.data.recipe_types import Recipe, derive_tables

SHEET = "客流汇总"
START = dt.date(2026, 7, 1)
DAYS = 8                          # 至少 tabular.DATE_HEADER_MIN（7）格日期才按交叉表起草
DAY_HOURS = (8, 9, 10, 11)
NIGHT_HOURS = (18, 19, 20)
GAP = 3                           # 第 4 天：时段全是「·」
PEAK, PEAK_N = "高峰时段", "时段人数（人次）"


def peak_book(edit: Callable[[Any], None] | None = None) -> tuple[bytes, str]:
    rnd = random.Random(7)
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET
    end = START + dt.timedelta(days=DAYS - 1)
    last = get_column_letter(2 + DAYS)
    ws["B2"] = f"统计时间范围：{START.year}年{START.month}月{START.day}日至{end.year}年{end.month}月{end.day}日"
    ws.merge_cells(f"B2:{last}2")
    ws["B3"] = "某地区客流汇总表"
    ws.merge_cells(f"B3:{last}3")
    for j in range(DAYS):
        d = START + dt.timedelta(days=j)
        ws.cell(4, 3 + j, f"{d.month}月{d.day}日")
    hours: dict[int, list[Any]] = {h: [rnd.randint(100, 900) for _ in range(DAYS)] for h in DAY_HOURS + NIGHT_HOURS}
    zone_a = [rnd.randint(3000, 9000) for _ in range(DAYS)]
    zone_b = [rnd.randint(3000, 9000) for _ in range(DAYS)]
    peak: list[Any] = []
    peak_n: list[Any] = []
    for j in range(DAYS):
        if j == GAP:
            for h in hours:
                hours[h][j] = "·"
            peak.append("·")
            peak_n.append("·")
            continue
        h = max(hours, key=lambda k: hours[k][j])
        peak.append(f"{h}–{h + 1}" if j == 1 else f"{h}-{h + 1}")
        peak_n.append(hours[h][j])
    rows: list[tuple[str, list[Any]]] = [
        ("全日客流（人次）", [a + b for a, b in zip(zone_a, zone_b)]), ("分区甲（人次）", zone_a),
        ("分区乙（人次）", zone_b), (PEAK, peak), (PEAK_N, peak_n)]
    r = 5
    for label, values in rows:
        ws.cell(r, 2, label)
        for j, v in enumerate(values):
            ws.cell(r, 3 + j, v)
        r += 1
    ws.merge_cells(f"B{r}:{last}{r}")             # 第 10 行空行
    r += 1
    for title, hs in (("日间时段客流（人次）", DAY_HOURS), ("夜间时段客流（人次）", NIGHT_HOURS)):
        ws.cell(r, 2, title)
        ws.merge_cells(f"B{r}:{last}{r}")
        r += 1
        for h in hs:
            ws.cell(r, 2, f"{h}-{h + 1}")
            for j, v in enumerate(hours[h]):
                ws.cell(r, 3 + j, v)
            r += 1
    if edit is not None:
        edit(ws)
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue(), f"月报导出_{START.isoformat()}_{end.isoformat()}.xlsx"


class Book:
    def __init__(self, edit: Callable[[Any], None] | None = None):
        self.raw, self.filename = peak_book(edit)
        self.scan = xlsx_scan.scan(self.raw)
        self.grids = {s.name: xlsx_cells.read_grid(self.raw, self.scan, s.name, max_rows=S.DRAFT_MAX_ROWS)
                      for s in self.scan.sheets if s.state == "visible" and s.bounds is not None and s.nonempty}

    def run(self, recipe: dict[str, Any] | Recipe, db: str | None = None):
        parsed = recipe if isinstance(recipe, Recipe) else Recipe.model_validate(recipe)
        return recipe_engine.execute(parsed, self.raw, self.filename, self.scan, db)

    def draft(self):
        return S.draft(self.scan, self.grids, self.filename,
                       validate=lambda r, f, o: validate_recipe(r, facts=f, origin=o), dry_run=self.run)


@pytest.fixture(scope="module")
def drafted():
    book = Book()
    d = book.draft()
    return book, d


def measures_seg(recipe: dict[str, Any]) -> dict[str, Any]:
    return next(s for s in recipe["sheets"][0]["blocks"][0]["segments"] if s["role"] == "measures")


def full(recipe: dict[str, Any]) -> dict[str, Any]:
    return Recipe.model_validate(recipe).model_dump(mode="json")


def codes(problems: list[Any]) -> list[str]:
    return [p.code for p in problems]


# ==========================================================================
# 规则起草
# ==========================================================================


def test_rule_draft_turns_the_text_row_into_a_text_measure(drafted):
    _book, d = drafted
    assert d.complete, d.failures
    assert d.failures == []
    seg = measures_seg(d.recipe)
    assert seg["text_measures"] == {PEAK: {"parser": "hour_range"}}
    assert PEAK in seg["labels"]["expect"] and PEAK_N in seg["labels"]["expect"]
    # 文字列不写单位；「时段人数（人次）」是数字指标，单位照常拆
    units = next(t["units"] for t in d.recipe["tables"] if t["name"] == seg["table"])
    assert PEAK not in units and units.get("时段人数") == "人次"


def test_rule_draft_asks_to_confirm_the_text_measure(drafted):
    _book, d = drafted
    seg = measures_seg(d.recipe)
    card = next(c for c in d.cards if c.id == f"text_measure:{seg['id']}:{PEAK}")
    q = next(q for q in d.questions if q.id == card.question)
    assert card.title == f"第 8 行「{PEAK}」按时段导入"
    assert f"{SHEET}!B8" in card.cells and f"{SHEET}!C8" in card.cells
    assert [o["value"] for o in q.options] == ["text", "reject"]
    assert q.effects["text"] == []
    # 答「不是」：去掉这一条，这一行回到按数字导入，干跑拒收（不会静默存成空值）
    rejected = apply_patch(full(d.recipe), q.effects["reject"])
    assert PEAK not in measures_seg(rejected)["text_measures"]
    book = Book()
    ex = book.run(rejected)
    assert not ex.ok and "value_not_number" in codes(ex.problems)


def test_rule_draft_merges_bad_cells_of_one_row_into_one_failure():
    # 时段人数一行混进两格文字：标签写着单位，是数字指标，照旧起草失败，但一行只报一条，给模型的那条没有格子里的字
    d = Book(lambda ws: (ws.cell(9, 3, "约3000"), ws.cell(9, 4, "约4000"))).draft()
    assert not d.complete
    lines = [f for f in d.failures if f.startswith(f"工作表「{SHEET}」：") and ("第 9 行" in f or "C9" in f)]
    assert lines == [f"工作表「{SHEET}」：第 9 行「{PEAK_N}」有 2 个单元格不是数字，也不像表示无数据的符号"
                     "（如 C9「约3000」、D9「约4000」）"]
    assert not any("约3000" in f for f in d.failures_for_model)
    # 干跑另有一条按分段合并的，下一步不再叫人把「约3000」加成占位符
    dry = next(f for f in d.failures if "的分段" in f and "C9" in f)
    assert "添加占位符" not in dry and "请检查原表" in dry
    assert PEAK not in str(d.failures)                       # 高峰时段照样按文字指标起草，不报错
    single = Book(lambda ws: ws.cell(9, 3, "约3000")).draft()
    assert f"工作表「{SHEET}」：C9 的「约3000」不是数字，也不像表示无数据的符号" in single.failures


def test_text_measure_lands_as_a_text_column(drafted, tmp_path):
    book, d = drafted
    db = str(tmp_path / "t.db")
    ex = book.run(d.recipe, db)
    assert ex.ok, [(p.code, p.message) for p in ex.problems]
    seg = measures_seg(d.recipe)
    table = next(t for t in ex.tables if t.name == seg["table"])
    types = {c.name: c.type for c in table.columns}
    assert types["高峰时段"] == "TEXT" and types["时段人数"] == "INTEGER" and types["全日客流"] == "INTEGER"
    con = sqlite3.connect(db)
    try:
        decl = {row[1]: row[2] for row in con.execute(f'PRAGMA table_info("{seg["table"]}")')}
        rows = con.execute(f'SELECT "日期", "高峰时段", "时段人数" FROM "{seg["table"]}" ORDER BY "日期"').fetchall()
        hour_rows = con.execute('SELECT COUNT(*) FROM "时段客流"').fetchone()[0]
    finally:
        con.close()
    assert decl["高峰时段"] == "TEXT"
    assert len(rows) == DAYS and hour_rows == DAYS * (len(DAY_HOURS) + len(NIGHT_HOURS))
    # en dash 写的存规范写法；整天没有时段数的那天两列都是空值
    assert rows[1][1] is not None and "–" not in rows[1][1] and rows[1][1].count("-") == 1
    assert rows[GAP][1:] == (None, None)
    assert all(isinstance(r[2], int) for i, r in enumerate(rows) if i != GAP)
    # 高峰时段能和时段客流按规范写法对上
    con = sqlite3.connect(db)
    try:
        hit = con.execute(f'SELECT COUNT(*) FROM "{seg["table"]}" d JOIN "时段客流" h ON h."日期" = d."日期" '
                          'AND h."时段" = d."高峰时段" AND h."客流" = d."时段人数"').fetchone()[0]
    finally:
        con.close()
    assert hit == DAYS - 1
    assert any(c.column == "高峰时段" and "–" in c.raw for c in ex.canonicalized)


# ==========================================================================
# 执行器：文字指标的取值、报错
# ==========================================================================


def test_number_in_a_text_measure_is_value_not_text(drafted):
    _, d = drafted
    book = Book(lambda ws: ws.cell(8, 3, 18))
    ex = book.run(d.recipe)
    p = next(p for p in ex.problems if p.code == "value_not_text")
    assert p.category == "structure" and f"{SHEET}!C8" in p.cells
    assert "「指标类型」改为「数字」" in p.message
    assert "18" not in p.model_message.replace(SHEET, "").replace("C8", "")


def test_unparsable_hour_range_and_plain_text_parser(drafted):
    _, d = drafted
    book = Book(lambda ws: ws.cell(8, 3, "傍晚"))
    ex = book.run(d.recipe)
    p = next(p for p in ex.problems if p.code == "value_not_hour_range")
    assert "「指标类型」改为「文字」" in p.message and "傍晚" not in p.model_message
    as_text = full(d.recipe)
    measures_seg(as_text)["text_measures"][PEAK] = {"parser": "text"}
    ex2 = book.run(as_text)
    assert ex2.ok, [(p.code, p.message) for p in ex2.problems]


# ==========================================================================
# 没声明成文字指标时：报错的下一步和修复按钮
# ==========================================================================


def _without_text_measure(d: Any) -> dict[str, Any]:
    r = full(d.recipe)
    measures_seg(r)["text_measures"] = {}
    return r


def test_not_number_on_a_text_row_points_at_measure_kind_not_placeholder(drafted):
    book, d = drafted
    ex = book.run(_without_text_measure(d))
    p = next(p for p in ex.problems if p.code == "value_not_number")
    assert f"「{PEAK}」这一行写的都是文字" in p.message
    assert "「指标类型」改为「文字」或「时段」" in p.message
    assert "添加占位符" not in p.message
    seg = measures_seg(d.recipe)
    assert p.fix_args["texts"] == []
    assert p.fix_args["text_rows"] == [{"segment": seg["id"], "label": PEAK, "cell": f"{SHEET}!B8",
                                        "parsers": ["hour_range", "text"]}]


def test_not_number_tail_depends_on_what_the_cells_are(drafted):
    _, d = drafted
    # 不含数字的短符号：照旧提示加占位符
    ex = Book(lambda ws: ws.cell(5, 3, "—")).run(d.recipe)
    p = next(p for p in ex.problems if p.code == "value_not_number")
    assert p.message.endswith("如果它表示无数据，请在配方面板中添加占位符")
    # 含数字、又不是整行文字：不再提示加占位符
    ex = Book(lambda ws: ws.cell(5, 3, "约3000")).run(d.recipe)
    p = next(p for p in ex.problems if p.code == "value_not_number")
    assert "占位符" not in p.message.split("。", 1)[1].replace("表示无数据的符号", "")
    assert "请检查原表" in p.message and "text_rows" not in p.fix_args
    # 标签写着单位的指标整行是文字：是数字指标，不提示改指标类型
    ex = Book(lambda ws: [ws.cell(9, c, "18-19") for c in range(3, 3 + DAYS)]).run(d.recipe)
    p = next(p for p in ex.problems if p.code == "value_not_number")
    assert "指标类型" not in p.message and "text_rows" not in p.fix_args


def test_fix_button_declares_the_text_measure(drafted):
    book, d = drafted
    work = _without_text_measure(d)
    ex = book.run(work)
    problems = [p.__dict__ for p in ex.problems]
    props = F.propose_fixes(work, problems, [], d.facts)
    p = next(x for x in props if x.kind == "declare_text_measure")
    assert [o.value for o in p.options] == ["hour_range", "text"]
    assert all(o.breaking for o in p.options)
    assert p.anchor.index == next(i for i, x in enumerate(problems) if x["code"] == "value_not_number")
    res = F.fix_ops(work, p, "hour_range", None)
    assert res.ok and "按时段导入" in res.title
    new = apply_patch(work, res.ops)
    assert measures_seg(new)["text_measures"] == {PEAK: {"parser": "hour_range"}}
    parsed, probs = validate_recipe(new, facts=d.facts, origin="manual")
    assert probs == [] and parsed is not None
    assert book.run(parsed).ok
    # 已经按文字导入的不再提议
    assert not [x for x in F.propose_fixes(new, problems, [], d.facts) if x.kind == "declare_text_measure"]


def test_remove_label_also_drops_the_text_measure(drafted):
    _, d = drafted
    work = full(d.recipe)
    seg = measures_seg(work)
    prob = {"code": "label_missing", "fix": "remove_label", "cells": [],
            "fix_args": {"segment": seg["id"], "labels": [PEAK], "role": "measures"}}
    p = next(x for x in F.propose_fixes(work, [prob], [], d.facts) if x.kind == "remove_label")
    new = apply_patch(work, F.fix_ops(work, p, p.options[0].value, None).ops)
    assert PEAK not in measures_seg(new)["text_measures"]
    _parsed, probs = validate_recipe(new, facts=d.facts, origin="manual")
    assert "text_measure_unknown" not in codes(probs)


# ==========================================================================
# 静态校验、契约
# ==========================================================================


def test_static_checks_on_text_measures(drafted):
    _, d = drafted
    r = full(d.recipe)
    seg = measures_seg(r)
    seg["text_measures"]["不存在的指标"] = {"parser": "text"}
    seg["text_measures"][PEAK_N] = {"parser": "text"}
    table = next(t for t in r["tables"] if t["name"] == seg["table"])
    table["units"]["高峰时段"] = "人次"
    _parsed, probs = validate_recipe(r, facts=d.facts, origin="manual")
    got = {(p.code, p.path) for p in probs}
    sp = "/sheets/0/blocks/0/segments/0"
    assert ("text_measure_unknown", f"{sp}/text_measures/不存在的指标") in got
    assert ("text_measure_unit", f"{sp}/text_measures/{PEAK_N}") in got
    assert any(code == "unit_on_text" and path.endswith("/units/高峰时段") for code, path in got)


def test_placeholder_with_digits_is_rejected_except_on_replay_or_when_carried(drafted):
    _, d = drafted
    r = full(d.recipe)
    measures_seg(r)["text_measures"] = {}
    r["sheets"][0]["blocks"][0]["values"]["placeholders"].append({"text": "18-19", "meaning": "无数据"})
    _parsed, probs = validate_recipe(r, facts=d.facts, origin="manual")
    p = next(p for p in probs if p.code == "placeholder_digits")
    assert p.path == "/sheets/0/blocks/0/values/placeholders/1/text"
    assert "「指标类型」改为「文字」或「时段」" in p.message
    assert "placeholder_digits" not in codes(validate_recipe(r, origin="replay")[1])
    base = Recipe.model_validate(copy.deepcopy(r))
    assert "placeholder_digits" not in codes(validate_recipe(r, facts=d.facts, origin="manual", base=base)[1])


def test_contract_text_measures_are_text_columns_and_hash_compatible(drafted):
    _, d = drafted
    parsed = Recipe.model_validate(d.recipe)
    cols, problems = derive_tables(parsed)
    seg = measures_seg(d.recipe)
    assert problems == []
    assert {c.name: c.type for c in cols[seg["table"]]}["高峰时段"] == "TEXT"
    # 默认值（没有文字指标）不进规范形式：已存配方的哈希和构建 id 不变
    plain = full(d.recipe)
    measures_seg(plain)["text_measures"] = {}
    assert "text_measures" not in str(canonical_recipe(Recipe.model_validate(plain)))


# ==========================================================================
# 确认项、前后对照、说明
# ==========================================================================


def test_confirm_item_and_breaking_changes(drafted):
    book, d = drafted
    ex = book.run(d.recipe)
    seg = measures_seg(d.recipe)
    items = {i.id: i for i in confirm_items(ConfirmContext(kind="first", recipe=d.recipe, extraction=ex, checks=[]))}
    item = items[f"text_measure:{seg['id']}:{PEAK}"]
    assert "按时段导入" in item.label and "文字列" in item.label
    as_text = full(d.recipe)
    measures_seg(as_text)["text_measures"][PEAK] = {"parser": "text"}
    kinds = {(c.kind, c.column) for c in table_changes(d.recipe, as_text)}
    assert ("store", "高峰时段") in kinds
    plain = full(d.recipe)
    measures_seg(plain)["text_measures"] = {}
    assert ("type", "高峰时段") in {(c.kind, c.column) for c in table_changes(d.recipe, plain)}


def test_column_note_explains_nulls_of_the_text_measure(drafted, tmp_path):
    from app.data.recipe_checks import run_checks
    from app.data.recipe_notes import build_notes

    book, d = drafted
    db = str(tmp_path / "n.db")
    ex = book.run(d.recipe, db)
    parsed = Recipe.model_validate(d.recipe)
    notes = build_notes(parsed, ex, run_checks(db, ex, parsed), [])
    seg = measures_seg(d.recipe)
    assert "空值" in notes.tables[seg["table"]].columns.get("高峰时段", "")
