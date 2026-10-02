"""证据下钻的清单与格子（WP-B，P4-SPEC 2.2、2.5、2.7、3.2、6.3）。

清单用真管线造（模块级夹具跑一遍）：经接口（AsyncClient + ASGITransport）导入几份合成的表格，从数据库取出冻结的
表结构和导入清单，作为纯函数的输入。之后的用例只用内存里的工件和一个按 id 取的 loader，不再走管线：
- rep：8 月 D00 每期替换（S1），再上传 9 月 D26 写理由接受 R1（S2，单期、seq 2）；
- acc：8 月 D00 按期累积（S1，单期），再上传 9 月 D24w（全角时段标签，S2 是并集）；
- rm：8 月加 9 月按期累积之后移除 8 月，只剩 9 月一期（清单里的 snapshot_id 是当初的并集，不是现在这个）；
- lst：lab.c01 列表（带合计行，另存成表）；
- d28：9 月 D28（表下有含数字的说明 B32），首次导入；
- c04：lab.c04 列表（两行表头，上层横向合并、左两列纵向合并，合并的数据按左上格填充）。

坐标一律和 openpyxl 读到的合成夹具对照：被引用的格保存的值等于快照库里那一格的值。夹具全部合成、假名，
不调用任何模型。
"""
from __future__ import annotations

import asyncio
import copy
import hashlib
import io
import json
import sqlite3
import uuid
from contextlib import closing
from dataclasses import dataclass, field
from typing import Any

import pytest
from httpx import ASGITransport, AsyncClient
from openpyxl import load_workbook

from app.core import artifact_store
from app.core.config import settings
from app.data import provenance as P
from app.data import provenance_types as T
from app.db.base import SessionLocal
from app.db.base import engine as db_engine
from app.db.models import SourceSnapshot
from app.main import app
from app.tools.datasource import _FROZEN_SCHEMA_KEYS
from tests.fixtures.xlsx import lab
from tests.fixtures.xlsx.drift import AUG, SEP
from tests.fixtures.xlsx.flow import SHEET, flow_workbook

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
API = "/api/datasources"
TABLES = ("日客流", "时段客流", "时段客流_表内合计")
REASON = "合成理由：分区合计口径调整"
DAY_TITLE = "日间时段客流（人次）"


# ==========================================================================
# 模块级夹具：真管线跑一遍
# ==========================================================================


@dataclass
class Snap:
    id: str
    name: str
    source_id: str
    db_path: str
    db_sha256: str
    #: 冻结的表结构快照（同 datasource._store_schema 存进 schema_snapshot 的内容）
    schema: dict[str, Any]


@dataclass
class World:
    #: 工件 id → 内容：导入清单、快照清单（测试里的 loader 只认它）
    store: dict[str, Any]
    snaps: dict[str, Snap]
    #: 名字 → [每一期的原件字节]（按统计期）
    raws: dict[str, list[bytes]]
    imports: dict[str, list[str]] = field(default_factory=dict)


def _answers(questions: list[dict[str, Any]], mode: str, prefer: str | None = None) -> dict[str, dict[str, Any]]:
    """prefer：选项的 effects 里把某个字段设成这个值的优先选（lab.c04 合并单元格的问题选 fill）。"""
    out = {}
    for q in questions:
        values = [o["value"] for o in q["options"]]
        effects = q.get("effects") or {}
        wanted = [v for v in values if prefer and prefer in json.dumps(effects.get(v, []), ensure_ascii=False)]
        if wanted:
            out[q["id"]] = {"value": wanted[0]}
        elif q["id"].startswith("q_relation"):
            out[q["id"]] = {"value": "register"}
        elif q["id"] == "q_mode":
            out[q["id"]] = {"value": mode}
        elif "null" in values:
            out[q["id"]] = {"value": "null"}
        else:
            out[q["id"]] = {"value": values[0]}
    return out


def _rename_wide(rec: dict[str, Any], old: str, new: str) -> dict[str, Any]:
    rec = json.loads(json.dumps(rec, ensure_ascii=False))
    for sheet in rec["sheets"]:
        for block in sheet["blocks"]:
            for seg in block.get("segments", []):
                if seg.get("table") == old:
                    seg["table"] = new
                if seg.get("id") == old:
                    seg["id"] = new
    for t in rec["tables"]:
        if t["name"] == old:
            t["name"] = new
    for r in rec["relations"]:
        if r.get("table") == old:
            r["table"] = new
        for side in ("a", "b"):
            if isinstance(r.get(side), dict) and r[side].get("table") == old:
                r[side]["table"] = new
    return rec


async def _commit(c: AsyncClient, st: dict[str, Any], acceptances: list[dict[str, str]] | None = None) -> dict:
    body: dict[str, Any] = {"trial_id": st["trial"]["trial_id"],
                            "confirmations": [i["id"] for i in st["trial"]["confirm_items"]]}
    if acceptances:
        body["acceptances"] = acceptances
    r = await c.post(f"{API}/imports/{st['id']}/commit", json=body)
    assert r.status_code == 201, r.text
    return r.json()


async def _first(c: AsyncClient, raw: bytes, fn: str, mode: str, *, flow: bool = True,
                 prefer: str | None = None) -> dict[str, Any]:
    name = f"prov_{uuid.uuid4().hex[:8]}"
    r = await c.post(f"{API}/imports/stage", files={"file": (fn, raw, XLSX)}, data={"name": name})
    assert r.status_code == 201, r.text
    st = r.json()
    if st.get("questions"):
        r = await c.post(f"{API}/imports/{st['id']}/answers", json={"answers": _answers(st["questions"], mode, prefer)})
        assert r.status_code == 200, r.text
        st = r.json()
    if flow:
        wide = next(t["name"] for t in st["recipe"]["tables"] if t["name"] not in TABLES)
        r = await c.put(f"{API}/imports/{st['id']}/recipe",
                        json={"recipe": _rename_wide(st["recipe"], wide, "日客流")})
        assert r.status_code == 200, r.text
        st = r.json()
    r = await c.post(f"{API}/imports/{st['id']}/trial", json={})
    assert r.status_code == 200, r.text
    st = r.json()
    assert st["trial"]["status"] == "passed", st["trial"]["problems"]
    out = await _commit(c, st)
    return {"name": name, "source_id": out["source"]["id"], "snapshot": out["snapshot_id"],
            "import": out["import_id"]}


async def _add(c: AsyncClient, source_id: str, raw: bytes, fn: str,
               acceptances: list[dict[str, str]] | None = None) -> dict[str, Any]:
    r = await c.post(f"{API}/{source_id}/reupload", files={"file": (fn, raw, XLSX)})
    assert r.status_code == 201, r.text
    st = r.json()
    assert st["trial"]["status"] in ("passed", "needs_decision"), st["trial"]["problems"]
    return await _commit(c, st, acceptances)


async def _snap(name: str, snapshot_id: str, store: dict[str, Any]) -> Snap:
    async with SessionLocal() as s:
        row = await s.get(SourceSnapshot, snapshot_id)
    cache = row.schema_cache
    schema = {"source": name, "tables": cache["tables"], **{k: cache[k] for k in _FROZEN_SCHEMA_KEYS if k in cache}}
    for mid in cache.get("import_manifests") or []:
        store[mid] = artifact_store.load(mid)
    if cache.get("snapshot_manifest"):
        store[cache["snapshot_manifest"]] = artifact_store.load(cache["snapshot_manifest"])
    return Snap(snapshot_id, name, row.source_id, row.db_path, row.db_sha256, json.loads(json.dumps(schema)))


async def _build() -> World:
    store: dict[str, Any] = {}
    snaps: dict[str, Snap] = {}
    raws: dict[str, list[bytes]] = {}
    imports: dict[str, list[str]] = {}
    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as c:
        aug, aug_fn = flow_workbook(AUG, 31, seed=401)
        rep = await _first(c, aug, aug_fn, "replace")
        snaps["rep1"] = await _snap(rep["name"], rep["snapshot"], store)
        d26, d26_fn = flow_workbook(SEP, 30, seed=402, variant="D26")
        out = await _add(c, rep["source_id"], d26, d26_fn, acceptances=[{"check_id": "R1", "reason": REASON}])
        snaps["rep2"] = await _snap(rep["name"], out["snapshot_id"], store)
        raws["rep1"], raws["rep2"] = [aug], [d26]

        acc = await _first(c, aug, aug_fn, "accumulate")
        snaps["acc1"] = await _snap(acc["name"], acc["snapshot"], store)
        sepw, sepw_fn = flow_workbook(SEP, 30, seed=403, variant="D24w")
        out = await _add(c, acc["source_id"], sepw, sepw_fn)
        snaps["acc2"] = await _snap(acc["name"], out["snapshot_id"], store)
        raws["acc1"], raws["acc2"] = [aug], [aug, sepw]
        imports["acc2"] = [acc["import"], out["import_id"]]

        rm = await _first(c, aug, aug_fn, "accumulate")
        sep, sep_fn = flow_workbook(SEP, 30, seed=404)
        out = await _add(c, rm["source_id"], sep, sep_fn)
        r = await c.post(f"{API}/{rm['source_id']}/imports/{rm['import']}/remove",
                         json={"confirm": True, "expected_current_snapshot_id": out["snapshot_id"], "reason": REASON})
        assert r.status_code == 200, r.text
        snaps["rm_union"] = await _snap(rm["name"], out["snapshot_id"], store)
        snaps["rm"] = await _snap(rm["name"], r.json()["snapshot_id"], store)
        raws["rm"] = [sep]

        c01, c01_fn = lab.c01()
        lst = await _first(c, c01, c01_fn, "replace", flow=False)
        snaps["lst"] = await _snap(lst["name"], lst["snapshot"], store)
        raws["lst"] = [c01]

        d28, d28_fn = flow_workbook(SEP, 30, seed=405, variant="D28")
        one = await _first(c, d28, d28_fn, "replace")
        snaps["d28"] = await _snap(one["name"], one["snapshot"], store)
        raws["d28"] = [d28]

        c04, c04_fn = lab.c04()
        mrg = await _first(c, c04, c04_fn, "replace", flow=False, prefer="fill")
        snaps["c04"] = await _snap(mrg["name"], mrg["snapshot"], store)
        raws["c04"] = [c04]
    # 这个事件循环里建的连接留在连接池里，换了事件循环的用例取到会出错：交还之前全部关掉
    await db_engine.dispose()
    return World(store, snaps, raws, imports)


@pytest.fixture(scope="module")
def world(tmp_path_factory) -> World:
    from app.data import recipe_ai
    from app.data.recipe_types import AiAvailability

    async def no_model(session):
        return None, "", AiAvailability(False, reason="测试不接模型")

    data = tmp_path_factory.mktemp("provenance") / "data"
    (data / "uploads").mkdir(parents=True)
    mp = pytest.MonkeyPatch()
    mp.setattr(settings, "data_dir", data)
    mp.setattr(recipe_ai, "resolve_draft_model", no_model)
    loop = asyncio.new_event_loop()
    try:
        return loop.run_until_complete(_build())
    finally:
        loop.close()
        mp.undo()


# ==========================================================================
# 小工具
# ==========================================================================


def loader(store: dict[str, Any], *, broken: frozenset[str] = frozenset(), gone: frozenset[str] = frozenset()):
    """语义同 artifact_store.load：取不到返回 None，哈希不符抛 ValueError。每次给一份拷贝，用例改了不串。"""
    def load(aid: str) -> Any:
        if aid in broken:
            raise ValueError(f"工件 {aid[:12]}… 内容与哈希不符")
        if aid in gone or aid not in store:
            return None
        return copy.deepcopy(store[aid])
    return load


def query_of(snap: Snap, columns: list[str] | None = None, rows: list[list[Any]] | None = None) -> dict[str, Any]:
    """查询快照里推断要用的几项（source、data_version、columns）。"""
    return {"source": snap.name, "data_version": snap.id, "schema_artifact": "schema", "sql": "SELECT 1",
            "columns": columns or [], "rows": rows or []}


def chain_of(world: World, key: str, store: dict[str, Any] | None = None) -> T.Chain:
    snap = world.snaps[key]
    got = P.resolve_chain(query_of(snap), snap.schema, loader(store or world.store))
    assert isinstance(got, T.Chain), got
    return got


def db_row(snap: Snap, table: str, where: dict[str, Any], column: str) -> tuple[int, Any]:
    """按主键在快照库上取 (rowid, 值)：只读、immutable，不改文件。"""
    cond = " AND ".join(f'"{k}" = ?' for k in where)
    with closing(sqlite3.connect(f"file:{snap.db_path}?mode=ro&immutable=1", uri=True)) as conn:
        got = conn.execute(f'SELECT rowid, "{column}" FROM "{table}" WHERE {cond}', list(where.values())).fetchall()
    assert len(got) == 1, (table, where, got)
    return got[0]


def xl(raw: bytes, sheet: str, ref: str) -> Any:
    return load_workbook(io.BytesIO(raw), data_only=True)[sheet][ref].value


def located(world: World, key: str, table: str, where: dict[str, Any], column: str,
            store: dict[str, Any] | None = None) -> tuple[T.Located, Any]:
    snap = world.snaps[key]
    rowid, value = db_row(snap, table, where, column)
    got = P.locate(chain_of(world, key, store), table=table, column=column, rowid=rowid)
    assert isinstance(got, T.Located), got
    return got, value


def froms(cs: T.CellSource) -> list[tuple[str, str | None, str, str | None, str | None]]:
    return [(f.role, f.column, f.cell, f.text, f.locate_title) for f in cs.from_]


def edited(world: World, key: str, fn) -> tuple[dict[str, Any], str]:
    """把某个快照的（唯一一份）导入清单拷贝一份、交给 fn 改，换一个新 id 放进新的 store；返回 (store, 新 id)。
    loader 是测试自己的，不复验哈希，所以 id 随便取；表结构快照也同步改成列这份新清单。"""
    snap = world.snaps[key]
    store = copy.deepcopy(world.store)
    mid = snap.schema["import_manifests"][-1]
    doc = copy.deepcopy(store[mid])
    fn(doc)
    new = "edited" + mid[:20]
    store[new] = doc
    return store, new


def with_manifest(world: World, key: str, fn) -> tuple[T.Chain, dict[str, Any]]:
    """单期快照的清单被 fn 改过之后的链（和 store）。"""
    store, new = edited(world, key, fn)
    snap = world.snaps[key]
    schema = {**copy.deepcopy(snap.schema), "import_manifests": [new]}
    got = P.resolve_chain(query_of(snap), schema, loader(store))
    assert isinstance(got, T.Chain), got
    return got, store


def recipe_blocks(doc: dict[str, Any]) -> list[dict[str, Any]]:
    return [b for s in doc["recipe"]["canonical"]["sheets"] for b in s["blocks"]]


# ==========================================================================
# resolve_chain（2.2）
# ==========================================================================


def test_single_period_chain_takes_the_manifest_hash_and_seq(world):
    snap = world.snaps["rep1"]
    chain = chain_of(world, "rep1")
    assert (chain.source, chain.snapshot_id, chain.mode, chain.union) == (snap.name, snap.id, "replace", False)
    assert chain.expected_db_sha256 == snap.db_sha256 and chain.source_id == snap.source_id
    assert chain.snapshot_manifest is None and len(chain.parts) == 1
    part = chain.parts[0]
    assert (part.seq, part.manifest_id, part.union_rows) == (1, snap.schema["import_manifests"][0], None)
    assert part.manifest["import_id"] == part.import_id
    # 每期替换的第二期只有一项，seq 是导入记录的 2，不是在 parts 里的位置
    second = chain_of(world, "rep2")
    assert second.parts[0].seq == 2 and second.expected_db_sha256 == world.snaps["rep2"].db_sha256
    # 按期累积的第一期：单期，mode 取配方的 accumulate
    first = chain_of(world, "acc1")
    assert (first.union, first.mode, len(first.parts)) == (False, "accumulate", 1)


def test_union_chain_maps_parts_to_their_manifests(world):
    snap = world.snaps["acc2"]
    chain = chain_of(world, "acc2")
    sm = world.store[snap.schema["snapshot_manifest"]]
    assert (chain.union, chain.mode, chain.snapshot_id) == (True, "accumulate", snap.id)
    assert chain.expected_db_sha256 == sm["union"]["db_sha256"] == snap.db_sha256
    assert chain.snapshot_manifest == sm and chain.source_id == snap.source_id
    assert [p.manifest_id for p in chain.parts] == snap.schema["import_manifests"]
    assert [p.import_id for p in chain.parts] == world.imports["acc2"] and [p.seq for p in chain.parts] == [1, 2]
    for p, sp in zip(chain.parts, sm["parts"]):
        for table, rows in sp["rows"].items():
            assert p.union_rows[table] == [*rows["union"], *rows["part"]]
    assert chain.parts[1].union_rows["日客流"] == [32, 61, 1, 30]


def test_single_period_left_after_removal_does_not_compare_snapshot_ids(world):
    """移除 8 月之后只剩 9 月：冻结的 import_manifests 沿用 9 月原来的清单，清单里的 snapshot_id 是当初的并集。
    单期不比快照 id（改成比，这一例要失败），期望的库哈希取这份清单的 db_sha256。"""
    snap = world.snaps["rm"]
    assert "snapshot_manifest" not in snap.schema and len(snap.schema["import_manifests"]) == 1
    manifest = world.store[snap.schema["import_manifests"][0]]
    assert manifest["snapshot_id"] == world.snaps["rm_union"].id != snap.id
    chain = chain_of(world, "rm")
    assert (chain.union, chain.snapshot_id, chain.parts[0].seq) == (False, snap.id, 2)
    assert chain.expected_db_sha256 == manifest["db_sha256"] == snap.db_sha256


def test_chain_problems(world):
    rep, acc = world.snaps["rep1"], world.snaps["acc2"]
    load = loader(world.store)
    # 单期 import_manifests 不是恰好 1 份：链本身对不上
    for mids in ([], [rep.schema["import_manifests"][0]] * 2):
        got = P.resolve_chain(query_of(rep), {**rep.schema, "import_manifests": mids}, load)
        assert (got.code, got.detail) == ("chain_mismatch", "import_manifests_count")
    # 并集：快照清单的 snapshot_id 不等于查询快照的 data_version
    got = P.resolve_chain({**query_of(acc), "data_version": world.snaps["acc1"].id}, acc.schema, load)
    assert (got.code, got.detail) == ("chain_mismatch", "snapshot_id")
    # 并集：parts 与 import_manifests 逐个相等（顺序也算）
    swapped = {**acc.schema, "import_manifests": list(reversed(acc.schema["import_manifests"]))}
    got = P.resolve_chain(query_of(acc), swapped, load)
    assert (got.code, got.detail) == ("chain_mismatch", "parts_vs_import_manifests")
    dropped = {**acc.schema, "import_manifests": acc.schema["import_manifests"][:1]}
    assert P.resolve_chain(query_of(acc), dropped, load).code == "chain_mismatch"
    # 取回抛 ValueError（哈希不符）、返回 None（取不回）：manifest_unreadable，不往外抛
    for kw in ({"broken": frozenset(rep.schema["import_manifests"])}, {"gone": frozenset(rep.schema["import_manifests"])}):
        got = P.resolve_chain(query_of(rep), rep.schema, loader(world.store, **kw))
        assert got.code == "manifest_unreadable", got
    sm_id = acc.schema["snapshot_manifest"]
    assert P.resolve_chain(query_of(acc), acc.schema, loader(world.store, broken=frozenset({sm_id}))).code \
        == "manifest_unreadable"
    part2 = acc.schema["import_manifests"][1]
    assert P.resolve_chain(query_of(acc), acc.schema, loader(world.store, gone=frozenset({part2}))).code \
        == "manifest_unreadable"

    def boom(aid):
        raise OSError("磁盘错误")
    assert P.resolve_chain(query_of(rep), rep.schema, boom).code == "manifest_unreadable"


def test_union_chain_cross_checks_each_part(world):
    acc = world.snaps["acc2"]
    sm_id = acc.schema["snapshot_manifest"]

    def chain_with(fn):
        store = copy.deepcopy(world.store)
        fn(store[sm_id])
        return P.resolve_chain(query_of(acc), acc.schema, loader(store))

    def other_import(sm):
        sm["parts"][1]["import_id"] = "别的导入"
    assert chain_with(other_import).detail == "part_identity"

    def other_source(sm):
        sm["source_id"] = "别的数据源"
    assert chain_with(other_source).detail == "source_id"

    def bad_rows(sm):
        sm["parts"][0]["rows"]["日客流"] = {"union": [1, 31]}
    assert chain_with(bad_rows).detail == "part_rows"


# ==========================================================================
# locate（2.5）：五种形态逐个对坐标
# ==========================================================================


def test_wide_table_measure_cell(world):
    got, value = located(world, "rep1", "日客流", {"日期": "2026-08-05"}, "全日客流")
    cs = got.cell
    assert (cs.sheet, cs.cell, cs.column_role, cs.kind, cs.part_seq) == (SHEET, "G5", "measure", "data", 1)
    assert cs.part_rowid == cs.rowid == got.part_rowid == 5
    assert xl(world.raws["rep1"][0], SHEET, "G5") == value
    assert froms(cs) == [("axis_header", "日期", "G4", None, None),
                         ("row_label", None, "B5", "全日客流（人次）", None)]
    assert (cs.header, cs.unit) == ("全日客流（人次）", "人次")
    assert cs.year == T.YearSource(source="period", cells=[f"{SHEET}!B2"], signed_by=None, mixed=False)
    assert cs.canonical is None and cs.merged_fill is False and not cs.pk and cs.recheck is None
    assert xl(world.raws["rep1"][0], SHEET, "B5") == "全日客流（人次）"


def test_anchor_is_the_cited_column_itself(world):
    """被引用的列有溯源段就取它自己的格（不是同一行第一个有溯源段的列）。"""
    raw = world.raws["rep1"][0]
    for column, ref in (("分区甲", "AG6"), ("分区乙", "AG7")):
        got, value = located(world, "rep1", "日客流", {"日期": "2026-08-31"}, column)
        assert got.cell.cell == ref and xl(raw, SHEET, ref) == value
        assert froms(got.cell)[1] == ("row_label", None, f"B{ref[2:]}", xl(raw, SHEET, f"B{ref[2:]}"), None)


def test_long_table_value_cell_from_section_title(world):
    got, value = located(world, "rep1", "时段客流", {"日期": "2026-08-20", "时段": "13-14"}, "客流")
    cs = got.cell
    raw = world.raws["rep1"][0]
    assert (cs.cell, cs.column_role) == ("V16", "value") and xl(raw, SHEET, "V16") == value
    assert froms(cs) == [("axis_header", "日期", "V4", None, None),
                         ("row_label", "时段", "B16", "13-14", None),
                         ("section_title", "时段类别", "B9", None, DAY_TITLE)]
    assert xl(raw, SHEET, "B9") == DAY_TITLE and cs.canonical is None


def test_reported_total_cell_gets_the_total_label_and_canonical(world):
    got, value = located(world, "rep1", "时段客流_表内合计", {"日期": "2026-08-10", "合计项": "18-22时合计"}, "客流")
    cs = got.cell
    raw = world.raws["rep1"][0]
    assert (cs.cell, cs.kind) == ("L28", "reported_total") and xl(raw, SHEET, "L28") == value
    # derived_label → total_label；原文取 DerivedItem.label_raw，规范写法（库里合计项的值）取 label
    assert froms(cs) == [("axis_header", "日期", "L4", None, None),
                         ("total_label", "合计项", "B28", "18-22 时合计", None)]
    assert xl(raw, SHEET, "B28") == "18-22 时合计"
    assert cs.canonical == T.CanonicalView(raw="18-22 时合计", canonical="18-22时合计")


def test_list_cell_gets_its_column_header(world):
    got, value = located(world, "lst", "月报", {"地区": lab.ZONES[2]}, "金额")
    cs = got.cell
    raw = world.raws["lst"][0]
    assert (cs.sheet, cs.cell, cs.column_role) == ("月报", "F8", "measure") and xl(raw, "月报", "F8") == value
    assert froms(cs) == [("col_header", None, "F5", "金额", None)] and xl(raw, "月报", "F5") == "金额"
    assert cs.year is None and cs.canonical is None and cs.merged_fill is False
    # 列表的文本列同样有溯源段：被引用的就是它自己那一格
    got, value = located(world, "lst", "月报", {"地区": lab.ZONES[2]}, "地区")
    assert (got.cell.cell, got.cell.column_role) == ("C8", "text") and xl(raw, "月报", "C8") == value


def test_list_total_row(world):
    got, value = located(world, "lst", "月报_表内合计", {"合计项": "合计"}, "金额")
    cs = got.cell
    raw = world.raws["lst"][0]
    assert (cs.cell, cs.kind) == ("F9", "reported_total") and xl(raw, "月报", "F9") == value
    assert froms(cs) == [("col_header", None, "F5", "金额", None), ("total_label", "合计项", "C9", "合计", None)]
    # 合计项列本身：取合计标签那一格
    got, value = located(world, "lst", "月报_表内合计", {"合计项": "合计"}, "合计项")
    assert got.cell.cell == "C9" and xl(raw, "月报", "C9") == value
    assert froms(got.cell) == []


def test_union_cell_lands_in_the_second_period(world):
    snap = world.snaps["acc2"]
    got, value = located(world, "acc2", "时段客流", {"日期": "2026-09-12", "时段": "8-9"}, "客流")
    cs = got.cell
    sm = world.store[snap.schema["snapshot_manifest"]]
    start = sm["parts"][1]["rows"]["时段客流"]["union"][0]
    assert (cs.part_seq, got.part.seq) == (2, 2) and cs.part_rowid == cs.rowid - start + 1
    assert cs.cell == "N11" and xl(world.raws["acc2"][1], SHEET, "N11") == value
    assert froms(cs) == [("axis_header", "日期", "N4", None, None),
                         ("row_label", "时段", "B11", "８－９", None),
                         ("section_title", "时段类别", "B9", None, DAY_TITLE)]
    assert cs.canonical == T.CanonicalView(raw="８－９", canonical="8-9")
    # 8 月那一期的行在第 1 期
    got, value = located(world, "acc2", "时段客流", {"日期": "2026-08-07", "时段": "19-20"}, "客流")
    assert (got.cell.part_seq, got.cell.cell) == (1, "I23") and got.part_rowid == got.cell.rowid
    assert xl(world.raws["acc2"][0], SHEET, "I23") == value


def test_key_columns_use_the_axis_row_the_row_label_or_the_section_title(world):
    raw = world.raws["acc2"][1]
    where = {"日期": "2026-09-12", "时段": "8-9"}
    got, value = located(world, "acc2", "时段客流", where, "时段")              # dim：行标签格
    assert (got.cell.cell, got.cell.column_role) == ("B11", "dim") and value == "8-9"
    assert xl(raw, SHEET, "B11") == "８－９"
    assert froms(got.cell) == [("axis_header", "日期", "N4", None, None),
                               ("section_title", "时段类别", "B9", None, DAY_TITLE)]
    assert got.cell.canonical == T.CanonicalView(raw="８－９", canonical="8-9")
    got, _ = located(world, "acc2", "时段客流", where, "日期")                 # axis：轴行
    assert got.cell.cell == "N4" and [f.role for f in got.cell.from_] == ["row_label", "section_title"]
    got, _ = located(world, "acc2", "时段客流", where, "时段类别")             # const：分段标题
    assert got.cell.cell == "B9" and [f.role for f in got.cell.from_] == ["axis_header", "row_label"]
    got, _ = located(world, "acc2", "时段客流", where, "起始小时")             # derive：由行标签解析
    assert got.cell.cell == "B11" and got.cell.canonical is None
    got, _ = located(world, "rep1", "日客流", {"日期": "2026-08-05"}, "日期")  # 宽表的日期：只有表头格
    assert got.cell.cell == "G4" and got.cell.from_ == []


def test_measure_column_without_lineage_does_not_fall_back(world):
    """measure / value / text 没有溯源段一律 no_lineage，不回退到同一行别的列（回退取到的是另一列的格子）。"""
    def drop(doc):
        del doc["receipt"]["lineage"]["日客流"]["全日客流"]
    chain, store = with_manifest(world, "rep1", drop)
    rowid, _ = db_row(world.snaps["rep1"], "日客流", {"日期": "2026-08-05"}, "全日客流")
    got = P.locate(chain, table="日客流", column="全日客流", rowid=rowid)
    assert got == T.LocateProblem("no_lineage", "no_lineage_for_row")
    # 键列照常回退到第一个有溯源段的列（分区甲）
    got = P.locate(chain, table="日客流", column="日期", rowid=rowid)
    assert isinstance(got, T.Located) and got.cell.cell == "G4"


def test_const_column_needs_a_section_title(world):
    def by_labels(doc):
        for b in recipe_blocks(doc):
            for seg in b.get("segments", []):
                if seg["id"] == "日间":
                    seg["locate"] = {"by": "labels"}
    chain, _ = with_manifest(world, "rep1", by_labels)
    rowid, _ = db_row(world.snaps["rep1"], "时段客流", {"日期": "2026-08-20", "时段": "13-14"}, "客流")
    for column in ("时段类别", "客流"):
        assert P.locate(chain, table="时段客流", column=column, rowid=rowid) == \
            T.LocateProblem("no_lineage", "section_title")

    def no_titles(doc):
        doc["receipt"]["regions"] = [m for m in doc["receipt"]["regions"] if m["role"] != "section_title"]
    chain, _ = with_manifest(world, "rep1", no_titles)
    assert P.locate(chain, table="时段客流", column="时段类别", rowid=rowid).code == "no_lineage"


def test_section_title_is_the_nearest_one_above_the_segment(world):
    """夜间段的分段标题是 B21，不是更上面的 B9。"""
    got, _ = located(world, "rep1", "时段客流", {"日期": "2026-08-07", "时段": "19-20"}, "时段类别")
    assert got.cell.cell == "B21"


def test_phase2_regions_without_block_keys(world):
    """期 2 生成的清单 regions 没有 block 键：同一工作表上配方恰好只有一块时认它；两块时不猜。"""
    def strip(doc):
        for m in doc["receipt"]["regions"]:
            m.pop("block", None)
    chain, _ = with_manifest(world, "rep1", strip)
    rowid, _ = db_row(world.snaps["rep1"], "时段客流", {"日期": "2026-08-20", "时段": "13-14"}, "客流")
    got = P.locate(chain, table="时段客流", column="客流", rowid=rowid)
    assert isinstance(got, T.Located) and got.cell.cell == "V16"
    assert [f.cell for f in got.cell.from_] == ["V4", "B16", "B9"]

    def two_blocks(doc):
        strip(doc)
        sheet = doc["recipe"]["canonical"]["sheets"][0]
        extra = {"id": "另一块", "layout": "list", "table": "别的表",
                 "columns": [{"header": "甲", "name": "甲", "type": "TEXT"}]}
        sheet["blocks"].append(extra)
        doc["recipe"]["canonical"]["tables"].append({"name": "别的表"})
    chain, _ = with_manifest(world, "rep1", two_blocks)
    assert P.locate(chain, table="时段客流", column="客流", rowid=rowid) == \
        T.LocateProblem("no_lineage", "anchor_block")


def test_anchor_outside_every_value_region_or_two_axes(world):
    def no_values(doc):
        doc["receipt"]["regions"] = [m for m in doc["receipt"]["regions"] if m["role"] != "value"]
    chain, _ = with_manifest(world, "rep1", no_values)
    rowid, _ = db_row(world.snaps["rep1"], "日客流", {"日期": "2026-08-05"}, "全日客流")
    assert P.locate(chain, table="日客流", column="全日客流", rowid=rowid).detail == "anchor_block"

    def two_axes(doc):
        doc["receipt"]["axes"].append(dict(doc["receipt"]["axes"][0]))
    chain, _ = with_manifest(world, "rep1", two_axes)
    assert P.locate(chain, table="日客流", column="全日客流", rowid=rowid).detail == "axis"


@pytest.mark.parametrize("form, expect", [("text", (False, True)), ("mixed", (True, True)), ("date", None)])
def test_year_rule_by_axis_form(world, form, expect):
    def set_form(doc):
        doc["receipt"]["axes"][0]["form"] = form
    chain, _ = with_manifest(world, "rep1", set_form)
    rowid, _ = db_row(world.snaps["rep1"], "日客流", {"日期": "2026-08-05"}, "全日客流")
    year = P.locate(chain, table="日客流", column="全日客流", rowid=rowid).cell.year
    if expect is None:
        assert year is None                      # 日期格自带年份
    else:
        assert year.mixed is expect[0] and year.source == "period" and year.cells == [f"{SHEET}!B2"]


def test_year_from_a_human_period_and_without_year_from(world):
    def human(doc):
        doc["period"] = {"start": "2026-08-01", "end": "2026-08-31", "source": "human", "cells": [],
                         "signed_by": "录入员甲"}
    chain, _ = with_manifest(world, "rep1", human)
    rowid, _ = db_row(world.snaps["rep1"], "日客流", {"日期": "2026-08-05"}, "全日客流")
    year = P.locate(chain, table="日客流", column="全日客流", rowid=rowid).cell.year
    assert year == T.YearSource(source="human", cells=[], signed_by="录入员甲", mixed=False)

    def no_year_from(doc):
        recipe_blocks(doc)[0]["axis"] = {"year_from": None}
    chain, _ = with_manifest(world, "rep1", no_year_from)
    assert P.locate(chain, table="日客流", column="全日客流", rowid=rowid).cell.year is None


def test_merged_fill_follows_the_list_block_recipe(world):
    def fill(doc):
        recipe_blocks(doc)[0]["merged_data"] = "fill"
    chain, _ = with_manifest(world, "lst", fill)
    rowid, _ = db_row(world.snaps["lst"], "月报", {"地区": lab.ZONES[0]}, "地区")
    assert P.locate(chain, table="月报", column="地区", rowid=rowid).cell.merged_fill is True


def test_multi_row_header_lists_only_cells_that_hold_text(world):
    """多行表头（lab.c04）：回执的列表头区域是 A3:E3、C4:F4——段把中间的空格一并覆盖，D3 是合并区 C3:D3 里的空格，
    F3 在合并区 E3:F3 里、不在任何段内。from 只列段的端点（必定有字），段中间的格不列：逐行逐列扫一遍，from 里的
    每一格在原表里都写着字；被引用的格保存的值等于库里的值，合并区里的空格只在 merged_fill 时出现。"""
    snap = world.snaps["c04"]
    raw = world.raws["c04"][0]
    chain = chain_of(world, "c04")
    receipt = chain.parts[0].manifest["receipt"]
    assert [(m["role"], m["ref"]) for m in receipt["regions"] if m["role"] == "col_header"] == [
        ("col_header", "A3:E3"), ("col_header", "C4:F4")]
    table = "分地区销售"
    columns = [c["name"] for c in receipt["tables"][0]["columns"]]
    with closing(sqlite3.connect(f"file:{snap.db_path}?mode=ro&immutable=1", uri=True)) as conn:
        rows = conn.execute(f'SELECT rowid, * FROM "{table}" ORDER BY rowid').fetchall()
    assert len(rows) == 6
    seen: dict[str, set[tuple[str, ...]]] = {c: set() for c in columns}
    book = load_workbook(io.BytesIO(raw), data_only=True)
    for rowid, *values in rows:
        for column, value in zip(columns, values):
            got = P.locate(chain, table=table, column=column, rowid=rowid)
            assert isinstance(got, T.Located), (column, rowid, got)
            cs = got.cell
            assert cs.merged_fill is True and cs.year is None
            original = book[cs.sheet][cs.cell].value
            assert original == value or (original is None and column == "地区"), (column, cs.cell, original, value)
            for f in cs.from_:
                assert f.role == "col_header" and f.text is None, f
                assert book[f.sheet][f.cell].value not in (None, ""), (column, f.cell)
            seen[column].add(tuple(f.cell for f in cs.from_))
    # 每一列各行给的表头格相同；纵向合并的「地区」取上层 A3，「产品」B3、上半年的金额 D4 落在段中间，不列
    assert {c: sorted(v) for c, v in seen.items()} == {
        "地区": [("A3",)], "产品": [()], "c_2026年上半年_销量": [("C4",)], "c_2026年上半年_金额": [()],
        "c_2026年下半年_销量": [("E3",)], "c_2026年下半年_金额": [("F4",)]}


def test_union_rowid_out_of_every_period_is_a_chain_mismatch(world):
    chain = chain_of(world, "acc2")
    last = chain.parts[-1].union_rows["日客流"][1]
    for rowid in (last + 1, 10 ** 6):
        assert P.locate(chain, table="日客流", column="全日客流", rowid=rowid) == \
            T.LocateProblem("chain_mismatch", "rowid_out_of_parts")
    assert P.locate(chain, table="不存在的表", column="全日客流", rowid=1).code == "chain_mismatch"
    # 偏移：第 2 期的第一行 = 并集的第 32 行
    got = P.locate(chain, table="日客流", column="全日客流", rowid=32)
    assert (got.cell.part_seq, got.part_rowid, got.cell.cell) == (2, 1, "C5")


def test_rows_unknown_to_the_receipt(world):
    chain = chain_of(world, "rep1")
    assert P.locate(chain, table="日客流", column="全日客流", rowid=999).detail == "no_lineage_for_row"
    assert P.locate(chain, table="日客流", column="不存在", rowid=1).detail == "column_not_in_receipt"
    assert P.locate(chain, table="日客流", column="全日客流", rowid=0).detail == "rowid"


# ==========================================================================
# version_view（2.8.1 的 version）
# ==========================================================================


def test_version_view_lists_every_period_from_the_manifests(world):
    snap = world.snaps["acc2"]
    view = P.version_view(chain_of(world, "acc2"), tables=["时段客流", "日客流", "时段客流", "不存在的表",
                                                          "时段客流_表内合计"])
    assert (view.source, view.snapshot_id, view.mode, view.union, view.manifest_view) == (
        snap.name, snap.id, "accumulate", True, True)
    assert view.tables == [T.TableRef(name="时段客流", kind="data"), T.TableRef(name="日客流", kind="data"),
                           T.TableRef(name="时段客流_表内合计", kind="reported_total")]
    assert [p.seq for p in view.parts] == [1, 2] and [p.manifest for p in view.parts] == snap.schema["import_manifests"]
    for part, mid, raw in zip(view.parts, snap.schema["import_manifests"], world.raws["acc2"]):
        m = world.store[mid]
        assert (part.import_id, part.file_name, part.raw_sha256) == (m["import_id"], m["file"]["name"],
                                                                     m["file"]["raw_sha256"])
        assert part.raw_sha256 == hashlib.sha256(raw).hexdigest()
        assert (part.recipe_sha256, part.committed_at, part.signed_by) == (m["recipe"]["sha256"], m["created_at"], None)
        assert part.period == T.PeriodView(start=m["period"]["start"], end=m["period"]["end"], source="cells",
                                           cells=[f"{SHEET}!B2"], signed_by=None)
        # 当前状态留给接口一侧补
        assert (part.recipe_seq, part.raw_state, part.purged, part.revoked, part.has_row) == (None,) * 4 + (False,)
    assert [p.region for p in view.parts] == [f"{SHEET}!B4:AG30", f"{SHEET}!B4:AF30"]
    assert [p.excluded_rows for p in view.parts] == [0, 0]
    # 名字只对 ASCII 不区分大小写（name_key），中文逐字比
    assert P.version_view(chain_of(world, "lst"), tables=["月报"]).tables == [T.TableRef(name="月报", kind="data")]


def test_version_view_acceptances_and_old_receipts(world):
    view = P.version_view(chain_of(world, "rep2"), tables=["日客流"])
    m = world.store[world.snaps["rep2"].schema["import_manifests"][0]]
    acc = m["acceptances"]["overrides"][0]
    assert view.parts[0].acceptances == [T.AcceptanceView(
        check_id="R1", title="「日客流」：全日客流 = 分区甲 + 分区乙", kind="override", reason=REASON, signed_by=None,
        at=acc["at"])]
    assert view.mode == "replace" and view.parts[0].seq == 2

    def old(doc):
        # 期 3 之前的回执没有 rows_excluded（未记录，不是 0）；期 2 的 regions 没有 block 键（按块内角色认）
        del doc["receipt"]["rows_excluded"]
        for mk in doc["receipt"]["regions"]:
            mk.pop("block", None)
        doc["acceptances"]["waivers"] = [{"check_id": "K1", "reason": "合成理由：公式未缓存", "signed_by": "录入员甲"}]
        doc["period"] = {"start": "2026-09-01", "end": "2026-09-30", "source": "human", "cells": [],
                         "signed_by": "录入员甲"}
    chain, _ = with_manifest(world, "rep2", old)
    part = P.version_view(chain, tables=[]).parts[0]
    assert part.excluded_rows is None and part.region == f"{SHEET}!B4:AF30"
    assert [(a.check_id, a.kind, a.title, a.signed_by, a.at) for a in part.acceptances] == [
        ("R1", "override", "「日客流」：全日客流 = 分区甲 + 分区乙", None, acc["at"]),
        ("K1", "waiver", "合计按明细重算（分段「夜间合计」）", "录入员甲", None)]
    assert part.period == T.PeriodView(start="2026-09-01", end="2026-09-30", source="human", cells=[],
                                       signed_by="录入员甲")


def test_region_and_excluded_rows_helpers():
    receipt = {"regions": [
        {"sheet": "甲", "role": "value", "ref": "C5:D7", "block": "b1"},
        {"sheet": "甲", "role": "outside_text", "ref": "B40", "block": None},
        {"sheet": "甲", "role": "row_label", "ref": "B5:B9", "block": "b1"},
        {"sheet": "乙", "role": "col_header", "ref": "A1:E1", "block": "b2"},
    ], "rows_excluded": [{"rows": [[3, 4], [9, 9]]}, {"rows": [[20, 22]]}]}
    assert P.region_of(receipt) == "甲!B5:D9、乙!A1:E1"
    assert P.excluded_rows_of(receipt) == 6
    assert P.region_of({"regions": [{"sheet": "甲", "role": "context", "ref": "B2", "block": None}]}) is None


# ==========================================================================
# related_checks（2.7）
# ==========================================================================


def plans(world: World, key: str, table: str, where: dict[str, Any], column: str,
          store: dict[str, Any] | None = None) -> list[T.RelatedCheckPlan]:
    got, _ = located(world, key, table, where, column, store)
    return P.related_checks(chain_of(world, key, store), got, table=table, column=column)


def test_relation_checks_for_the_cited_column(world):
    got = plans(world, "rep1", "日客流", {"日期": "2026-08-05"}, "全日客流")
    assert [(p.check.id, p.check.kind, p.check.part_status) for p in got] == [
        ("R1", "relation_sum_eq", "passed"), ("R2", "relation_not_comparable", "info"),
        ("C1", "context_agree", "passed"), ("C2", "filename_period", "passed")]
    assert got[0].rule == T.SumEqRule(table="日客流", total="全日客流", parts=["分区甲", "分区乙"],
                                      types={"全日客流": "INTEGER", "分区甲": "INTEGER", "分区乙": "INTEGER"})
    assert all(p.rule is None for p in got[1:])
    assert all(p.check.row_status is None and p.check.cell_status is None and p.check.acceptance is None for p in got)
    # 被引用的是 parts 之一同样列 R1；日期不是 total 也不是 parts
    assert [p.check.id for p in plans(world, "rep1", "日客流", {"日期": "2026-08-05"}, "分区甲")][0] == "R1"
    assert [p.check.id for p in plans(world, "rep1", "日客流", {"日期": "2026-08-05"}, "日期")] == ["R2", "C1", "C2"]
    # 长表只有口径不同 R2（a 是时段客流）
    assert [p.check.id for p in plans(world, "rep1", "时段客流", {"日期": "2026-08-20", "时段": "13-14"}, "客流")] \
        == ["R2", "C1", "C2"]


def test_accepted_mismatch_carries_its_reason(world):
    got = plans(world, "rep2", "日客流", {"日期": "2026-09-10"}, "全日客流")
    r1 = got[0].check
    assert (r1.id, r1.part_status, r1.row_status) == ("R1", "mismatch", None)
    assert r1.acceptance is not None and (r1.acceptance.reason, r1.acceptance.kind) == (REASON, "override")
    assert got[0].rule is not None and got[1].check.acceptance is None


def test_total_cell_lists_k_and_g_of_its_segment(world):
    got = plans(world, "rep1", "时段客流_表内合计", {"日期": "2026-08-10", "合计项": "18-22时合计"}, "客流")
    assert [(p.check.id, p.check.cell_status, p.rule) for p in got] == [
        ("K1", "ok", None), ("G1", "ok", None), ("C1", None, None), ("C2", None, None)]
    # 合计项列不是合计格：不列 K、G
    got = plans(world, "rep1", "时段客流_表内合计", {"日期": "2026-08-10", "合计项": "18-22时合计"}, "合计项")
    assert [p.check.id for p in got] == ["C1", "C2"]
    # 列表的合计行是 T；列表没有统计期，不列 C
    got = plans(world, "lst", "月报_表内合计", {"合计项": "合计"}, "金额")
    assert [(p.check.id, p.check.kind, p.check.cell_status) for p in got] == [("T1", "column_sum", "ok")]
    assert plans(world, "lst", "月报", {"地区": lab.ZONES[0]}, "金额") == []


def test_g_on_a_literal_total_is_not_formula_and_k_details(world):
    where = {"日期": "2026-08-10", "合计项": "18-22时合计"}

    def literal(doc):
        for d in doc["receipt"]["derived"]:
            if d["cell"] == f"{SHEET}!L28":
                d["is_formula"], d["formula"] = False, None
        k1 = next(c for c in doc["checks"] if c["id"] == "K1")
        k1.update(status="unverifiable", unverifiable=2, cells=[f"{SHEET}!L28", f"{SHEET}!M28"],
                  details=["2 格无法核对：文件未经计算保存"])
    store, new = edited(world, "rep1", literal)
    snap = world.snaps["rep1"]
    schema = {**copy.deepcopy(snap.schema), "import_manifests": [new]}
    chain = P.resolve_chain(query_of(snap), schema, loader(store))
    rowid, _ = db_row(snap, "时段客流_表内合计", where, "客流")
    loc = P.locate(chain, table="时段客流_表内合计", column="客流", rowid=rowid)
    got = {p.check.id: p.check for p in P.related_checks(chain, loc, table="时段客流_表内合计", column="客流")}
    # 段里别的格还是公式，plan_checks 照样编 G1；这一格写死，不能说「这一格一致」
    assert got["G1"].cell_status == "not_formula" and got["G1"].part_status == "passed"
    assert (got["K1"].cell_status, got["K1"].part_status) == ("unverifiable", "unverifiable")
    assert got["K1"].detail == "2 格无法核对：文件未经计算保存" and got["G1"].detail is None


@pytest.mark.parametrize("check, expect", [
    ({"status": "passed"}, "ok"),
    ({"status": "unverifiable", "unverifiable": 2, "cells": ["甲!L28", "甲!M28"]}, "unverifiable"),
    ({"status": "unverifiable", "unverifiable": 2, "cells": ["甲!M28", "甲!N28"]}, "ok"),
    ({"status": "unverifiable", "unverifiable": 25, "cells": [f"甲!X{i}" for i in range(20)]}, "unknown"),
    ({"status": "mismatch", "failed": 1, "unverifiable": 1, "cells": ["甲!L28", "甲!M28"]}, None),
    ({"status": "mismatch", "failed": 1, "unverifiable": 1, "cells": ["甲!M28", "甲!L28"]}, "unverifiable"),
    ({"status": "mismatch", "failed": 1, "unverifiable": 0, "cells": ["甲!M28"]}, "ok"),
    ({"status": "info"}, None),
])
def test_cell_status_branches(check, expect):
    assert P.cell_status(check, "甲!L28", is_formula=True, formula_check=False) == expect
    assert P.cell_status(check, "甲!L28", is_formula=True, formula_check=True) == expect
    assert P.cell_status(check, "甲!L28", is_formula=False, formula_check=True) == "not_formula"
    assert P.cell_status(check, "甲!L28", is_formula=False, formula_check=False) == expect


def test_context_checks_only_when_the_year_comes_from_the_period(world):
    def human(doc):
        doc["period"] = {"start": "2026-08-01", "end": "2026-08-31", "source": "human", "cells": [],
                         "signed_by": "录入员甲"}
        doc["checks"] = [c for c in doc["checks"] if c["id"] != "C1"]     # 人工录入没有 C1
    chain, _ = with_manifest(world, "rep1", human)
    rowid, _ = db_row(world.snaps["rep1"], "日客流", {"日期": "2026-08-05"}, "全日客流")
    loc = P.locate(chain, table="日客流", column="全日客流", rowid=rowid)
    assert [p.check.id for p in P.related_checks(chain, loc, table="日客流", column="全日客流")] == ["R1", "R2", "C2"]

    def date_cells(doc):
        doc["receipt"]["axes"][0]["form"] = "date"
    chain, _ = with_manifest(world, "rep1", date_cells)
    loc = P.locate(chain, table="日客流", column="全日客流", rowid=rowid)
    assert [p.check.id for p in P.related_checks(chain, loc, table="日客流", column="全日客流")] == ["R1", "R2"]


# ==========================================================================
# chain_artifacts（2.8.3）
# ==========================================================================


def test_chain_artifacts_lists_the_chain_and_never_raises(world):
    acc, rep = world.snaps["acc2"], world.snaps["rep1"]
    sm_id = acc.schema["snapshot_manifest"]
    assert P.chain_artifacts(acc.schema, loader(world.store)) == {sm_id, *acc.schema["import_manifests"]}
    assert P.chain_artifacts(rep.schema, loader(world.store)) == set(rep.schema["import_manifests"])
    # 只列在快照清单里的各期清单：快照清单取得回才算
    only_sm = {"snapshot_manifest": sm_id, "import_manifests": []}
    assert P.chain_artifacts(only_sm, loader(world.store)) == {sm_id, *acc.schema["import_manifests"]}
    assert P.chain_artifacts(only_sm, loader(world.store, broken=frozenset({sm_id}))) == set()
    assert P.chain_artifacts(only_sm, loader(world.store, gone=frozenset({sm_id}))) == set()

    def boom(aid):
        raise RuntimeError("取回出错")
    assert P.chain_artifacts(only_sm, boom) == set()
    assert P.chain_artifacts(None, boom) == set()                                  # type: ignore[arg-type]
    assert P.chain_artifacts({"import_manifests": "不是列表"}, boom) == set()
    assert P.chain_artifacts({}, loader(world.store)) == set()

    class Odd(dict):
        def get(self, key, default=None):
            raise KeyError(key)
    assert P.chain_artifacts(Odd(), loader(world.store)) == set()


# ==========================================================================
# judge_lines（3.2）
# ==========================================================================

QUAL_ACCEPT = "已接受（用户填写的理由，未经系统核对）："
QUAL_OUTSIDE = ("区域外文字（原表区域外的文字原文，未经系统核对；其中的数字不是查询结果，不能用来支持或否定报告中的"
                "数字）：")
CUT = "…另有部分来历信息未列出，见证据面板"


def lines_of(world: World, key: str, columns: list[str], tables: list[str], *, hidden: set[str] = frozenset(),
             store: dict[str, Any] | None = None, schema: dict[str, Any] | None = None,
             query: dict[str, Any] | None = None) -> list[str]:
    snap = world.snaps[key]
    return P.judge_lines(query or query_of(snap, columns), schema or snap.schema, loader(store or world.store),
                         tables=tables, hidden=set(hidden))


def part_line(world: World, key: str, k: int) -> str:
    """第 k 份清单（按 import_manifests 的顺序）应有的分期描述，按 3.2 的写法逐字拼。"""
    m = world.store[world.snaps[key].schema["import_manifests"][k]]
    p = m["period"]
    return (f"第 {m['seq']} 次导入：统计期 {p['start']} 至 {p['end']}（取自 {SHEET}!B2），"
            f"文件「{m['file']['name']}」（sha256 {m['file']['raw_sha256'][:12]}），区域 {P.region_of(m['receipt'])}")


def test_judge_lines_for_a_union(world):
    snap = world.snaps["acc2"]
    tables = snap.schema["tables"]
    got = lines_of(world, "acc2", ["日期", "全日客流"], ["日客流"])
    assert got == [
        "来源：上传的表格，按期累积，共 2 期（2026-08-01 至 2026-09-30）",
        "  " + part_line(world, "acc2", 0),
        "  " + part_line(world, "acc2", 1),
        f"口径：表 日客流 的说明：{tables['日客流']['comment']}",
        "口径：列 日期 的说明：格式 YYYY-MM-DD，年份按统计期补全",
        "口径：列 全日客流 的说明：单位：人次",
        QUAL_OUTSIDE,
        "  第 2 次导入（统计期 2026-09-01 至 2026-09-30）客流汇总!B3：「客流汇总表」",
        "  第 1 次导入（统计期 2026-08-01 至 2026-08-31）客流汇总!B3：「客流汇总表」",
    ]
    assert part_line(world, "acc2", 0).endswith("区域 客流汇总!B4:AG30")
    # 同样的工件调两次逐字相同（摘录进断点指纹）
    assert lines_of(world, "acc2", ["日期", "全日客流"], ["日客流"]) == got


def test_judge_lines_single_period_with_an_acceptance(world):
    got = lines_of(world, "rep2", ["日期", "全日客流"], ["日客流"])
    assert got[0] == "来源：上传的表格，每期替换，" + part_line(world, "rep2", 0)
    i = got.index(QUAL_ACCEPT)
    assert got[i + 1] == ("  第 2 次导入（统计期 2026-09-01 至 2026-09-30）的核对「「日客流」：全日客流 = 分区甲 + 分区乙」"
                          f"不成立，理由：{REASON}（未署名）")
    assert got.index(QUAL_ACCEPT) < got.index(QUAL_OUTSIDE)


def test_judge_lines_outside_text_is_sent_in_full_with_its_period(world):
    got = lines_of(world, "d28", ["日期", "全日客流"], ["日客流"])
    i = got.index(QUAL_OUTSIDE)
    # 含数字的排在前面，全文照录，每行写所属的导入和统计期
    assert got[i + 1:] == [
        "  第 1 次导入（统计期 2026-09-01 至 2026-09-30）客流汇总!B32：「注：9月15日闸机故障，当日客流为估算值」",
        "  第 1 次导入（统计期 2026-09-01 至 2026-09-30）客流汇总!B3：「客流汇总表」",
    ]
    m = world.store[world.snaps["d28"].schema["import_manifests"][0]]
    assert {o["cell"]: o["hidden"] for o in m["receipt"]["outside_text"]} == {
        f"{SHEET}!B3": False, f"{SHEET}!B32": False}


@pytest.mark.parametrize("hidden_value", [True, None, "missing"])
def test_outside_text_not_known_to_be_visible_is_only_counted(world, hidden_value):
    """hidden 为 True、为 None、或者老清单没有这个键：不送全文，只计入「另有 n 处」（改成 not hidden 让 None 也送，
    这一例要失败）。"""
    def mark(doc):
        for o in doc["receipt"]["outside_text"]:
            if o["cell"].endswith("!B32"):
                if hidden_value == "missing":
                    del o["hidden"]
                else:
                    o["hidden"] = hidden_value
    store, new = edited(world, "d28", mark)
    schema = {**copy.deepcopy(world.snaps["d28"].schema), "import_manifests": [new]}
    got = lines_of(world, "d28", ["日期", "全日客流"], ["日客流"], store=store, schema=schema)
    assert not any("闸机故障" in x for x in got)
    i = got.index(QUAL_OUTSIDE)
    assert got[i + 1:] == ["  第 1 次导入（统计期 2026-09-01 至 2026-09-30）客流汇总!B3：「客流汇总表」",
                           "  另有 1 处区域外文字未列出（所在行列被隐藏，或导入时未记录是否隐藏），见证据面板"]

    def all_hidden(doc):
        for o in doc["receipt"]["outside_text"]:
            o.pop("hidden", None)
    store, new = edited(world, "d28", all_hidden)
    schema = {**copy.deepcopy(world.snaps["d28"].schema), "import_manifests": [new]}
    got = lines_of(world, "d28", ["日期", "全日客流"], ["日客流"], store=store, schema=schema)
    assert QUAL_OUTSIDE not in got
    assert got[-1] == "另有 2 处区域外文字未列出（所在行列被隐藏，或导入时未记录是否隐藏），见证据面板"


def test_masks_drop_reasons_outside_text_and_masked_column_notes(world):
    got = lines_of(world, "rep2", ["日期", "全日客流"], ["日客流"], hidden={"全日客流"})
    text = "\n".join(got)
    assert REASON not in text and "口径：列 全日客流" not in text and "口径：列 日期" in text
    assert ("  第 2 次导入（统计期 2026-09-01 至 2026-09-30）的核对「「日客流」：全日客流 = 分区甲 + 分区乙」不成立"
            "（数据源设置了遮罩，理由未列出）") in got
    assert QUAL_OUTSIDE not in got and got[-1] == "区域外另有 1 处文字（数据源设置了遮罩，未列出），见证据面板"
    got = lines_of(world, "d28", ["日期", "全日客流"], ["日客流"], hidden={"分区甲"})
    assert not any("闸机故障" in x for x in got)
    assert got[-1] == "区域外另有 2 处文字（数据源设置了遮罩，未列出），见证据面板"
    # 遮罩列按小写比（同 judge._hidden 的口径）
    upper = lines_of(world, "lst", ["地区", "金额"], ["月报"], hidden={"金额"})
    assert QUAL_OUTSIDE not in upper


def test_judge_lines_scope(world):
    snap = world.snaps["rep1"]
    load = loader(world.store)
    # 手工源（没有 data_version）、没有表结构快照：一行都不加
    assert P.judge_lines({**query_of(snap), "data_version": None}, snap.schema, load, tables=["日客流"],
                         hidden=set()) == []
    assert P.judge_lines(query_of(snap), None, load, tables=["日客流"], hidden=set()) == []
    # 简单导入（不是按配方导入）：只有口径
    simple = {k: v for k, v in snap.schema.items() if k not in ("import_mode", "import_manifests")}
    got = P.judge_lines(query_of(snap, ["日期"]), simple, load, tables=["日客流"], hidden=set())
    assert got and all(x.startswith("口径：") for x in got)
    # 未规整的表：说明就是 UNSHAPED_NOTE
    from app.data.tabular import UNSHAPED_NOTE

    raw = copy.deepcopy(simple)
    raw["tables"]["日客流"]["comment"] = UNSHAPED_NOTE
    assert P.judge_lines(query_of(snap), raw, load, tables=["日客流"], hidden=set()) == [
        f"口径：表 日客流 的说明：{UNSHAPED_NOTE}"]
    # 链解析失败：写一行说明，口径照写，不中断
    def boom(aid):
        raise ValueError("哈希不符")
    got = P.judge_lines(query_of(snap, ["日期"]), snap.schema, boom, tables=["日客流"], hidden=set())
    assert got[0] == "来源：导入清单无法读取或与登记不一致，未列出来历"
    assert got[1].startswith("口径：表 日客流") and not any(x.startswith(("已接受", "区域外")) for x in got)
    # SQL 里没有的表不写说明
    assert not any(x.startswith("口径：表 时段客流") for x in got)


# ---- 上限与截断：在真清单上拼出多期、长文字的变体（loader 是测试自己的，id 随便取）


def synthetic(world: World, n: int, *, accepts: int = 0, outside: int = 0, outside_len: int = 20,
              comment_len: int | None = None, columns: int = 0, column_len: int = 10
              ) -> tuple[dict[str, Any], dict[str, Any], dict[str, Any]]:
    """以 acc2 为模板拼一个 n 期的并集：(query, schema, store)。每期 accepts 条接受、outside 格可见的区域外文字
    （单数格含数字）；comment_len 给了就把表的说明换成这么长（0 = 没有说明）；columns 给了就换成这么多列、每列都有
    说明。"""
    acc = world.snaps["acc2"]
    store = copy.deepcopy(world.store)
    sm = copy.deepcopy(store[acc.schema["snapshot_manifest"]])
    template = store[acc.schema["import_manifests"][0]]
    mids, parts = [], []
    for k in range(1, n + 1):
        m = copy.deepcopy(template)
        m.update(import_id=f"导入{k}", seq=k, period={"start": f"2026-{k:02d}-01", "end": f"2026-{k:02d}-28",
                                                     "source": "cells", "cells": [f"{SHEET}!B2"]})
        m["file"] = {**m["file"], "name": f"月报_{k:02d}.xlsx"}
        m["acceptances"] = {"overrides": [{"check_id": "R1", "reason": f"第{k}期理由{j}", "signed_by": None}
                                          for j in range(accepts)], "waivers": []}
        m["receipt"]["outside_text"] = [
            {"sheet": SHEET, "cell": f"{SHEET}!B{40 + j}", "text": f"第{k}期注{j}" + "字" * outside_len,
             "kind": "text_digits" if j % 2 else "text", "period_source": False, "hidden": False}
            for j in range(outside)]
        mid = f"清单{k}"
        store[mid] = m
        mids.append(mid)
        parts.append({**sm["parts"][0], "import_id": f"导入{k}", "seq": k, "manifest_artifact": mid})
    sm.update(snapshot_id="合成并集", parts=parts)
    store["合成快照清单"] = sm
    schema = {**copy.deepcopy(acc.schema), "import_manifests": mids, "snapshot_manifest": "合成快照清单"}
    names = [f"列{j}" for j in range(columns)]
    meta = schema["tables"]["日客流"]
    if comment_len is not None:
        meta["comment"] = "说" * comment_len
    if columns:
        meta["columns"] = [{"name": c, "type": "INTEGER", "comment": f"{c}的说明" + "明" * column_len} for c in names]
    query = {"source": acc.name, "data_version": "合成并集", "columns": names or ["日期", "全日客流"], "rows": []}
    return query, schema, store


def test_judge_lines_caps(world):
    def run(**kw):
        query, schema, store = synthetic(world, comment_len=0, **kw)
        got = P.judge_lines(query, schema, loader(store), tables=["日客流"], hidden=set())
        assert CUT not in got, len("\n".join(got))
        return got, _sections(got)
    # 分期行最多 PROV_PARTS 期，取最近的；总述行写全部期数和起止
    got, sec = run(n=5)
    assert got[0] == "来源：上传的表格，按期累积，共 5 期（2026-01-01 至 2026-05-28）"
    assert [x.split("：")[0] for x in sec["parts"]] == ["  第 3 次导入", "  第 4 次导入", "  第 5 次导入"]
    assert got[len(sec["parts"]) + 1] == "  另有 2 期，见证据面板"
    # 列的说明最多 PROV_COLUMN_NOTES 行，按结果列的顺序
    got, sec = run(n=1, columns=8)
    assert [x.split(" 的说明")[0] for x in sec["columns"]] == [f"口径：列 列{j}" for j in range(T.PROV_COLUMN_NOTES)]
    # 已接受最多 PROV_ACCEPT 条（取最近几期的），按期次排序
    got, sec = run(n=3, accepts=2)
    assert [x.split("（")[0] for x in sec["accepts"]] == ["  第 2 次导入"] * 2 + ["  第 3 次导入"] * 2
    assert got[-1] == "  另有 2 条，见证据面板"
    # 区域外文字最多 PROV_OUTSIDE 格：先含数字的（B41），再按期次从新到旧
    got, sec = run(n=3, outside=2)
    assert [(x.split("（")[0], x.split("）")[1].split("：")[0]) for x in sec["outside"]] == [
        ("  第 3 次导入", f"{SHEET}!B41"), ("  第 2 次导入", f"{SHEET}!B41"), ("  第 1 次导入", f"{SHEET}!B41"),
        ("  第 3 次导入", f"{SHEET}!B40")]
    assert got[-1] == "  另有 2 处区域外文字未列出（超出条数上限），见证据面板"


def test_outside_text_longer_than_the_cap_is_cut(world):
    query, schema, store = synthetic(world, 1, outside=1, outside_len=300)
    got = P.judge_lines(query, schema, loader(store), tables=["日客流"], hidden=set())
    line = next(x for x in got if f"{SHEET}!B40" in x)
    assert line.endswith("…」（已截断）")
    body = line.split("：「", 1)[1].removesuffix("…」（已截断）")
    assert len(body) == T.PROV_OUTSIDE_CHARS and body.startswith("第1期注0字")
    query, schema, store = synthetic(world, 1, outside=1, outside_len=T.PROV_OUTSIDE_CHARS - 5)
    got = P.judge_lines(query, schema, loader(store), tables=["日客流"], hidden=set())
    assert not any("已截断" in x for x in got)


def test_judge_lines_wording_for_human_periods_signatures_and_waivers(world):
    """3.2 写死的几处措辞逐字钉住：统计期人工录入（有署名、没署名）、清单里没有统计期、接受理由的署名一律标
    「未认证」（0.1）、waivers 写「无法核对」。"""
    query, schema, store = synthetic(world, 3, outside=1, comment_len=0)
    m1, m2, m3 = (store[f"清单{k}"] for k in (1, 2, 3))
    m1["period"] = {"start": "2026-01-01", "end": "2026-01-28", "source": "human", "cells": [],
                    "signed_by": "录入员甲"}
    m2["period"] = {"start": "2026-02-01", "end": "2026-02-28", "source": "human", "cells": [], "signed_by": None}
    del m3["period"]
    m3["receipt"].pop("period", None)
    m1["acceptances"] = {"overrides": [{"check_id": "R1", "reason": "合成理由甲", "signed_by": "录入员甲"}],
                         "waivers": []}
    waived = next(c for c in m3["checks"] if c["id"] != "R1")
    m3["acceptances"] = {"overrides": [], "waivers": [
        {"check_id": waived["id"], "reason": "合成理由乙", "signed_by": "录入员乙"},
        {"check_id": waived["id"], "reason": "合成理由丙", "signed_by": None}]}
    got = P.judge_lines(query, schema, loader(store), tables=["日客流"], hidden=set())

    def tail(m: dict[str, Any]) -> str:
        return (f"文件「{m['file']['name']}」（sha256 {m['file']['raw_sha256'][:12]}），"
                f"区域 {P.region_of(m['receipt'])}")
    r1 = "「日客流」：全日客流 = 分区甲 + 分区乙"
    assert got == [
        "来源：上传的表格，按期累积，共 3 期（2026-01-01 至 2026-02-28）",
        f"  第 1 次导入：统计期 2026-01-01 至 2026-01-28（人工录入，署名「录入员甲」，未认证），{tail(m1)}",
        f"  第 2 次导入：统计期 2026-02-01 至 2026-02-28（人工录入，未署名），{tail(m2)}",
        f"  第 3 次导入：统计期未记录，{tail(m3)}",
        "口径：列 日期 的说明：格式 YYYY-MM-DD，年份按统计期补全",
        "口径：列 全日客流 的说明：单位：人次",
        QUAL_ACCEPT,
        f"  第 1 次导入（统计期 2026-01-01 至 2026-01-28）的核对「{r1}」不成立，理由：合成理由甲（署名「录入员甲」，未认证）",
        f"  第 3 次导入（统计期未记录）的核对「{waived['title']}」无法核对，理由：合成理由乙（署名「录入员乙」，未认证）",
        f"  第 3 次导入（统计期未记录）的核对「{waived['title']}」无法核对，理由：合成理由丙（未署名）",
        QUAL_OUTSIDE,
        f"  第 3 次导入（统计期未记录）{SHEET}!B40：「第3期注0{'字' * 20}」",
        f"  第 2 次导入（统计期 2026-02-01 至 2026-02-28）{SHEET}!B40：「第2期注0{'字' * 20}」",
        f"  第 1 次导入（统计期 2026-01-01 至 2026-01-28）{SHEET}!B40：「第1期注0{'字' * 20}」",
    ]


def _sections(got: list[str]) -> dict[str, list[str]]:
    sec: dict[str, list[str]] = {"parts": [], "columns": [], "accepts": [], "outside": [], "tables": []}
    mode = None
    for x in got:
        if x.startswith("口径：表 "):
            sec["tables"].append(x)
        elif x.startswith("口径：列 "):
            sec["columns"].append(x)
        elif x == QUAL_ACCEPT:
            mode = "accepts"
        elif x == QUAL_OUTSIDE:
            mode = "outside"
        elif x.startswith("  第 ") and "次导入：" in x:
            sec["parts"].append(x)
        elif x.startswith("  第 ") and mode:
            sec[mode].append(x)
    return sec


#: 表的说明越来越长，五步依次被用到（最短的不删，最长的截到表的说明）
_LENGTHS = tuple(range(0, 2400, 60))


def _truncated(world: World, comment_len: int) -> list[str]:
    query, schema, store = synthetic(world, 4, accepts=1, outside=1, columns=3, comment_len=comment_len)
    return P.judge_lines(query, schema, loader(store), tables=["日客流"], hidden=set())


def test_judge_lines_truncation_order(world):
    """超过 PROV_CHARS 时按固定顺序删：分期行（旧的先删，总述行保留）→ 列的说明（从最后一列起）→ 已接受（旧的先删）
    → 区域外文字（旧的先删）→ 表的说明截断加「…」；删过任何一项，末尾写一行说明。"""
    order = ["parts", "columns", "accepts", "outside", "tables"]
    full = {"parts": 3, "columns": 3, "accepts": 4, "outside": 4}
    states = set()
    for comment_len in _LENGTHS:
        got = _truncated(world, comment_len)
        assert len("\n".join(got)) <= T.PROV_CHARS, comment_len
        assert got[0] == "来源：上传的表格，按期累积，共 4 期（2026-01-01 至 2026-04-28）"
        sec = _sections(got)
        lost = {k: len(sec[k]) < v for k, v in full.items()}
        lost["tables"] = comment_len > 0 and (not sec["tables"] or sec["tables"][0].endswith("…"))
        for i, k in enumerate(order):
            if lost[k]:
                # 后面的类别动过了，前面的必须已经删光
                assert all(not sec[j] for j in order[:i]), (comment_len, k, sec)
        assert (got[-1] == CUT) == any(lost.values()), comment_len
        # 各类内部：分期行、已接受、区域外文字留下最近的；列的说明留下前几列
        assert [int(x.split()[1]) for x in sec["parts"]] == [2, 3, 4][3 - len(sec["parts"]):]
        assert [x.split(" 的说明")[0] for x in sec["columns"]] == [f"口径：列 列{j}" for j in range(len(sec["columns"]))]
        assert [int(x.split()[1]) for x in sec["accepts"]] == [1, 2, 3, 4][4 - len(sec["accepts"]):]
        assert sorted(int(x.split()[1]) for x in sec["outside"]) == [1, 2, 3, 4][4 - len(sec["outside"]):]
        # 同样的输入调两次逐字相同
        assert _truncated(world, comment_len) == got
        states.add(next((k for k in reversed(order) if lost[k]), None))
    assert states == {None, *order}, states
