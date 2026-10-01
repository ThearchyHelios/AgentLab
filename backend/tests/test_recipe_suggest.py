"""WP-4 确定性起草（recipe_suggest）的测试。

夹具一律在这里用 openpyxl 现造（假名：分区甲 / 分区乙，表名「客流汇总」），网格由 grid_from_ws 从 openpyxl 普通模式
直接造出来，不依赖读取层（WP-2）；validate / dry_run 用桩，不依赖静态校验（WP-1）和执行器（WP-2）。
"""
from __future__ import annotations

import datetime as dt
import io
import json
import random
from pathlib import Path
from typing import Any

import pytest
from openpyxl import Workbook
from openpyxl.utils import column_index_from_string, get_column_letter

from app.data import recipe_suggest as S
from app.data.recipe_types import (
    REASON_SLOT,
    AxisOut,
    Extraction,
    Grid,
    GridCell,
    OutsideText,
    PeriodOut,
    Problem,
    Recipe,
    RecipeProblem,
)
from app.data.xlsx_scan import scan as xlsx_scan

FLOW = json.loads((Path(__file__).parent / "fixtures" / "recipes" / "flow_recipe.json").read_text(encoding="utf-8"))


# ==========================================================================
# 夹具工具（test_recipe_ai.py 也用）
# ==========================================================================


def grid_from_ws(ws: Any, cache: dict[str, Any] | None = None) -> Grid:
    """openpyxl 普通模式的工作表 → Grid。公式格的保存值由 cache（坐标 → 值）给出（openpyxl 写公式不带缓存）。"""
    cache = cache or {}
    cells: dict[tuple[int, int], GridCell] = {}
    for row in ws.iter_rows():
        for cell in row:
            v = cell.value
            if v is None:
                continue
            if isinstance(v, str) and v.startswith("="):
                cells[(cell.row, cell.column)] = GridCell(cache.get(cell.coordinate), v)
            elif isinstance(v, str) and not v.strip():
                continue
            else:
                cells[(cell.row, cell.column)] = GridCell(v)
    merges = [(m.min_row, m.min_col, m.max_row, m.max_col) for m in ws.merged_cells.ranges]
    hidden_rows = {r for r, d in ws.row_dimensions.items() if d.hidden}
    hidden_cols: set[int] = set()
    for key, d in ws.column_dimensions.items():
        if d.hidden:
            lo = d.min or column_index_from_string(key)
            hi = d.max or lo
            hidden_cols.update(range(lo, hi + 1))
    bounds = None
    if cells:
        rs = [r for r, _ in cells]
        cs = [c for _, c in cells]
        bounds = (min(rs), min(cs), max(rs), max(cs))
    return Grid(ws.title, bounds, cells, merges, hidden_rows, hidden_cols)


def build(wb: Workbook, cache: dict[str, dict[str, Any]] | None = None):
    """保存成字节、跑期 1 的扫描、按可见工作表造网格。返回 (scan, grids)。"""
    bio = io.BytesIO()
    wb.save(bio)
    sc = xlsx_scan(bio.getvalue())
    grids = {ws.title: grid_from_ws(ws, (cache or {}).get(ws.title))
             for ws in wb.worksheets if ws.sheet_state == "visible"}
    return sc, grids


def flow_book(start: dt.date = dt.date(2026, 8, 1), days: int = 31, *, seed: int = 0,
              big: tuple[int, int] = (3000, 9000)) -> tuple[Workbook, dict[str, dict[str, Any]], dict[str, Any]]:
    """按 P2-SPEC 9.1 的坐标造「结构仿照客流表」：返回 (工作簿, 公式缓存, 填进去的数)。"""
    rnd = random.Random(seed)
    wb = Workbook()
    ws = wb.active
    ws.title = "客流汇总"
    last = 2 + days
    end = start + dt.timedelta(days=days - 1)
    L = get_column_letter
    ws["B2"] = f"统计时间范围：{start.year}年{start.month}月{start.day}日至{end.year}年{end.month}月{end.day}日"
    ws.merge_cells(f"B2:{L(last)}2")
    ws["B3"] = "客流汇总表"
    ws.merge_cells(f"B3:{L(last)}3")
    for i in range(days):
        d = start + dt.timedelta(days=i)
        ws.cell(4, 3 + i, f"{d.month}月{d.day}日")
    labels = ["全日客流（人次）", "分区甲（人次）", "分区乙（人次）"]
    for r, lab in zip((5, 6, 7), labels):
        ws.cell(r, 2, lab)
    nums: dict[str, Any] = {"jia": [], "yi": []}
    ws.merge_cells(f"B8:{L(last)}8")
    ws["B9"] = "日间时段客流（人次）"
    ws.merge_cells(f"B9:{L(last)}9")
    for k, h in enumerate(range(7, 18)):
        ws.cell(10 + k, 2, f"{h}-{h + 1}")
    ws["B21"] = "夜间时段客流（人次）"
    ws.merge_cells(f"B21:{L(last)}21")
    for k, h in enumerate(range(18, 24)):
        ws.cell(22 + k, 2, f"{h}-{h + 1}")
    for r, lab in zip((28, 29, 30), ["18-22 时合计", "22-24 时合计", "18-24 时合计"]):
        ws.cell(r, 2, lab)
    cache: dict[str, Any] = {}
    for i in range(days):
        c = 3 + i
        jia, yi = rnd.randint(*big), rnd.randint(2000, 8000)
        nums["jia"].append(jia)
        nums["yi"].append(yi)
        ws.cell(5, c, jia + yi)
        ws.cell(6, c, jia)
        ws.cell(7, c, yi)
        ws.cell(10, c, "·")
        slots = {}
        for r in list(range(11, 21)) + list(range(22, 28)):
            slots[r] = rnd.randint(50, 900)
        if sum(slots.values()) == jia + yi:
            slots[11] += 1
        for r, v in slots.items():
            ws.cell(r, c, v)
        col = L(c)
        ws.cell(28, c, f"=SUM({col}22:{col}25)")
        ws.cell(29, c, f"=SUM({col}26:{col}27)")
        ws.cell(30, c, f"=SUM({col}22:{col}27)")
        cache[f"{col}28"] = sum(slots[r] for r in range(22, 26))
        cache[f"{col}29"] = sum(slots[r] for r in range(26, 28))
        cache[f"{col}30"] = sum(slots[r] for r in range(22, 28))
    return wb, {"客流汇总": cache}, nums


def ok_validate(calls: list | None = None):
    """静态校验桩：pydantic 收就算通过。calls 记下 (配方, facts, origin)。"""
    def _v(data: dict[str, Any], facts: Any, origin: str):
        if calls is not None:
            calls.append((data, facts, origin))
        return Recipe.model_validate(data), []
    return _v


def canonical(d: dict[str, Any]) -> dict[str, Any]:
    return Recipe.model_validate(d).model_dump(mode="json", exclude_defaults=True)


def rename_wide(recipe: dict[str, Any], old: str, new: str) -> dict[str, Any]:
    """把宽表名、该段 id 和关系里引用的表名一并换掉（P2-SPEC 第 10 节 WP-4 测试的做法）。"""
    text = json.dumps(recipe, ensure_ascii=False)
    return json.loads(text.replace(json.dumps(old, ensure_ascii=False), json.dumps(new, ensure_ascii=False)))


# ==========================================================================
# 参考布局
# ==========================================================================


@pytest.fixture(scope="module")
def flow():
    wb, cache, nums = flow_book()
    sc, grids = build(wb, cache)
    return sc, grids, nums


def test_reference_layout_draft_matches_reference_recipe(flow):
    sc, grids, _ = flow
    calls: list = []
    d = S.draft(sc, grids, "月报导出_2026-08-01_2026-08-31.xlsx", validate=ok_validate(calls))
    assert d.complete, d.failures
    assert d.origin == "rules" and d.failures == [] and d.failures_for_model == []
    assert calls and calls[0][2] == "rules" and calls[0][1] is d.facts
    wide = d.recipe["tables"][0]["name"]
    assert wide == "客流汇总_按日"
    assert canonical(rename_wide(d.recipe, wide, "日客流")) == canonical(FLOW)


def apply_ops(doc: dict[str, Any], ops: list[dict[str, Any]], reason: str = "巧合") -> dict[str, Any]:
    """RFC 6902 的 add / replace / remove 子集（测试自用；生产里是 WP-1 的 recipe.apply_patch）。"""
    doc = json.loads(json.dumps(doc, ensure_ascii=False))

    def fill(v: Any) -> Any:
        if v == REASON_SLOT:
            return reason
        if isinstance(v, dict):
            return {k: fill(x) for k, x in v.items()}
        if isinstance(v, list):
            return [fill(x) for x in v]
        return v

    for op in ops:
        parts = [p.replace("~1", "/").replace("~0", "~") for p in op["path"].split("/")[1:]]
        node: Any = doc
        for p in parts[:-1]:
            node = node[int(p)] if isinstance(node, list) else node[p]
        last = parts[-1]
        if isinstance(node, list):
            if op["op"] == "remove":
                node.pop(int(last))
            elif last == "-":
                node.append(fill(op["value"]))
            elif op["op"] == "add":
                node.insert(int(last), fill(op["value"]))
            else:
                node[int(last)] = fill(op["value"])
        else:
            if op["op"] == "remove":
                del node[last]
            else:
                if op["op"] == "replace":
                    assert last in node, op
                node[last] = fill(op["value"])
    return doc


def test_reference_layout_facts(flow):
    sc, grids, _ = flow
    facts = S.detect_facts(grids, sc)
    assert [(f.id, f.kind) for f in facts.facts] == [("F1", "sum_eq"), ("F2", "not_equal_sum")]
    f1, f2 = facts.facts
    assert f1.text == "第 5 行 = 第 6 行 + 第 7 行（31 列中 31 列成立）"
    assert f1.detail == {"segment": "客流汇总_按日", "total": "全日客流（人次）",
                         "parts": ["分区甲（人次）", "分区乙（人次）"], "rows": [5, 6, 7]}
    assert f2.text == "日间、夜间各行之和 与 第 5 行：31 列中 0 列相等"
    assert f2.detail == {"a": ["日间", "夜间"], "b_segment": "客流汇总_按日", "b": "全日客流（人次）",
                         "equal": 0, "checked": 31}
    assert facts.candidates["日间时段客流（人次）"][0] == "日间"
    assert facts.nonnumeric == {"·": 31}
    # 只有关系不给数值：事实的文字里没有任何一个填进去的数
    _, _, nums = flow
    blob = json.dumps([f.__dict__ for f in facts.facts], ensure_ascii=False)
    assert not any(str(v) in blob for v in nums["jia"] + nums["yi"] if v > 100)


def test_reference_layout_cards_and_questions(flow):
    sc, grids, _ = flow
    d = S.draft(sc, grids, "x.xlsx", validate=ok_validate())
    assert [c.id for c in d.cards] == [
        "axis:交叉表", "period:s1", "titles:交叉表", "fact:F1", "merge:时段客流", "placeholder:交叉表:·",
        "derived:夜间合计", "fact:F2", "outside:s1", "mode"]
    titles = [c.title for c in d.cards]
    assert titles[0] == "第 4 行是日期表头，按交叉表导入"
    assert titles[1] == "统计期取自 B2"
    assert titles[2] == "分段标题 B9、B21"
    assert titles[3].startswith("第 5 行 = 第 6 行 + 第 7 行")
    assert titles[4] == "日间、夜间合进一张表「时段客流」"
    assert titles[5] == "「·」表示无数据吗"
    assert titles[6] == "第 28–30 行是合计，改作核对并另存"
    assert titles[7] == "各时段之和与全日客流 31 天中 0 天相等"
    assert titles[8] == "B3 是区域外文字"
    assert titles[9] == "导入模式：每期替换"
    assert "按期累积将在后续版本提供" in d.cards[9].reason
    assert d.cards[1].cells == ["客流汇总!B2"]
    assert d.cards[8].cells == ["客流汇总!B3"]
    by_id = {c.id: c for c in d.cards}
    assert by_id["placeholder:交叉表:·"].question == "q_placeholder:·"
    assert by_id["fact:F2"].question == "q_relation:F2"

    qs = {q.id: q for q in d.questions}
    assert [q.id for q in d.questions] == ["q_relation:F1", "q_placeholder:·", "q_relation:F2"]
    assert all(q.default is None for q in d.questions)
    for qid in ("q_relation:F1", "q_relation:F2"):
        q = qs[qid]
        opts = {o["value"]: o for o in q.options}
        assert opts["register"]["needs_reason"] is False
        assert opts["dismiss"]["needs_reason"] is True
        assert REASON_SLOT in json.dumps(q.effects["dismiss"], ensure_ascii=False)
    ph = qs["q_placeholder:·"]
    assert [o["value"] for o in ph.options] == ["null", "reject"]


def test_question_effects_apply_and_are_state_independent(flow):
    sc, grids, _ = flow
    d = S.draft(sc, grids, "x.xlsx", validate=ok_validate())
    qs = {q.id: q for q in d.questions}
    base = d.recipe
    dismissed = apply_ops(base, qs["q_relation:F1"].effects["dismiss"])
    rel = dismissed["relations"][0]
    assert rel == {"id": "R1", "kind": "dismissed", "claims": "F1", "reason": "巧合"}
    Recipe.model_validate(dismissed)
    # 先选「不登记」再改选「登记」，等于直接选「登记」（整条 replace，与状态无关）
    again = apply_ops(dismissed, qs["q_relation:F1"].effects["register"])
    assert canonical(again) == canonical(base)
    # 在改过的配方上重算问题，effects 仍然把它写回同一条
    _, qs2 = S.questions_for(dismissed, d.facts, grids)
    q2 = next(q for q in qs2 if q.id == "q_relation:F1")
    assert canonical(apply_ops(dismissed, q2.effects["register"])) == canonical(base)
    rejected = apply_ops(base, qs["q_placeholder:·"].effects["reject"])
    assert rejected["sheets"][0]["blocks"][0]["values"]["placeholders"] == []
    assert apply_ops(base, qs["q_placeholder:·"].effects["null"]) == base


def test_detect_facts_remaps_to_renamed_recipe(flow):
    """用户把宽表和该段 id 改名成「日客流」：给了 recipe 时，事实里的分段 id 跟着换，认领检查才对得上。"""
    sc, grids, _ = flow
    renamed = rename_wide(S.draft(sc, grids, "x.xlsx", validate=ok_validate()).recipe, "客流汇总_按日", "日客流")
    facts = S.detect_facts(grids, sc, recipe=renamed)
    assert facts.facts[0].detail["segment"] == "日客流"
    assert facts.facts[1].detail["b_segment"] == "日客流"
    assert facts.facts[1].detail["a"] == ["日间", "夜间"]
    # 不给 recipe 时是规则起草的取法
    assert S.detect_facts(grids, sc).facts[0].detail["segment"] == "客流汇总_按日"
    # 标签写法只按 match_key 对：换成配方里 expect 的写法
    flow2 = json.loads(json.dumps(FLOW))
    seg = flow2["sheets"][0]["blocks"][0]["segments"][0]
    seg["labels"]["expect"][1] = "分区甲 （人次）"
    seg["measures"] = {"全日客流（人次）": "全日客流", "分区甲 （人次）": "分区甲", "分区乙（人次）": "分区乙"}
    f = S.remap_facts(S.detect_facts(grids, sc), flow2).facts[0]
    assert f.detail["parts"] == ["分区甲 （人次）", "分区乙（人次）"]


def test_draft_with_dry_run_structure_problem_is_incomplete(flow):
    sc, grids, _ = flow
    seen: list = []

    def dry(recipe):
        seen.append(recipe)
        return Extraction(ok=False, problems=[
            Problem("cell_unclaimed", "structure", "第 31 行有 2 个单元格没有去处：987654", ["客流汇总!C31"],
                    model_message="第 31 行有 2 个单元格没有去处"),
            Problem("period_missing", "input", "未能从表格中解析出统计期"),
            Problem("outside_digits", "confirm", "区域外有含数字的文字"),
        ])
    d = S.draft(sc, grids, "x.xlsx", validate=ok_validate(), dry_run=dry)
    assert seen and isinstance(seen[0], Recipe)
    assert not d.complete
    assert d.failures == ["第 31 行有 2 个单元格没有去处：987654"]
    assert d.failures_for_model == ["cell_unclaimed：第 31 行有 2 个单元格没有去处（客流汇总!C31）"]
    assert any(c.id == "unclaimed:s1" for c in d.cards)
    # input、confirm 类不算起草不完整
    d2 = S.draft(sc, grids, "x.xlsx", validate=ok_validate(),
                 dry_run=lambda r: Extraction(ok=False, problems=[Problem("period_missing", "input", "x")]))
    assert d2.complete


def test_draft_with_static_problems_is_incomplete(flow):
    sc, grids, _ = flow

    def bad(data, facts, origin):
        return None, [RecipeProblem("/tables/0/units/全日客流", "unit_unknown", "表「日客流」：单位不在单位词表中")]
    d = S.draft(sc, grids, "x.xlsx", validate=bad, dry_run=lambda r: pytest.fail("静态校验没过不该干跑"))
    assert not d.complete and d.recipe is not None
    assert d.failures == ["表「日客流」：单位不在单位词表中"]
    assert d.failures_for_model == ["/tables/0/units/全日客流：表「日客流」：单位不在单位词表中"]

    def boom(data, facts, origin):
        raise RuntimeError("x")
    d2 = S.draft(sc, grids, "x.xlsx", validate=boom)
    assert not d2.complete and d2.failures


def test_questions_for_region_cards_come_from_extraction(flow):
    sc, grids, _ = flow
    d = S.draft(sc, grids, "x.xlsx", validate=ok_validate())
    ext = Extraction(
        ok=True,
        period=PeriodOut("2026-08-01", "2026-08-31", "cells", cells=["客流汇总!B40"]),
        outside_text=[OutsideText("客流汇总", "B41", "备注说明", "text"),
                      OutsideText("客流汇总", "B42", "注：9月15日闸机故障", "text_digits")],
        axes=[AxisOut("交叉表", 4, "8月1日", "8月31日", 31, "text")],
        problems=[Problem("row_unclaimed", "structure", "第 44 行没有分段认领", ["客流汇总!B44"])],
        placeholders={"·": 29},
    )
    cards, _ = S.questions_for(d.recipe, d.facts, grids, extraction=ext)
    by_id = {c.id: c for c in cards}
    assert by_id["period:s1"].cells == ["客流汇总!B40"]
    assert by_id["period:s1"].title == "统计期取自 B40"
    assert by_id["outside:s1"].cells == ["客流汇总!B41", "客流汇总!B42"]
    assert "提交时逐条确认" in by_id["outside:s1"].reason
    assert by_id["unclaimed:s1"].cells == ["客流汇总!B44"]
    assert "29 格" in by_id["placeholder:交叉表:·"].reason
    # 干跑说统计期是人工录入的：统计期卡片换成「每期人工录入」
    ext2 = Extraction(ok=True, period=PeriodOut("2026-08-01", "2026-08-31", "human"))
    cards2, qs2 = S.questions_for(d.recipe, d.facts, grids, extraction=ext2)
    assert next(c for c in cards2 if c.id == "period:s1").question == "q_period"
    assert any(q.id == "q_period" for q in qs2)


def test_questions_for_invalid_recipe_does_not_raise(flow):
    sc, grids, _ = flow
    cards, qs = S.questions_for({"recipe_format": "x"}, S.detect_facts(grids, sc), grids)
    assert [c.id for c in cards] == ["mode"] and qs == []


# ==========================================================================
# 交叉表的变体
# ==========================================================================


def _crosstab(ws, labels_rows: list[tuple[int, str]], days: int = 8, start_col: int = 3, axis_row: int = 4,
              value=lambda r, c: (r * 7 + c * 3) % 50 + 10):
    for i in range(days):
        ws.cell(axis_row, start_col + i, f"8月{i + 1}日")
    for r, lab in labels_rows:
        ws.cell(r, start_col - 1, lab)
        if lab.endswith("时段客流（人次）") or lab.endswith("（标题）"):
            continue
        for i in range(days):
            ws.cell(r, start_col + i, value(r, i))


def test_overlapping_hour_segments_are_not_merged():
    wb = Workbook()
    ws = wb.active
    ws.title = "时段"
    ws["B2"] = "统计期：2026年8月1日至2026年8月8日"
    rows = [(5, "工作日时段客流（人次）"), (6, "7-8"), (7, "8-9"), (9, "周末时段客流（人次）"), (10, "7-8"), (11, "8-9")]
    _crosstab(ws, rows)
    sc, grids = build(wb)
    d = S.draft(sc, grids, "x.xlsx", validate=ok_validate())
    assert d.complete, d.failures
    segs = d.recipe["sheets"][0]["blocks"][0]["segments"]
    assert [s["table"] for s in segs] == ["工作日时段客流", "周末时段客流"]
    assert all(s["const"] == {} for s in segs)
    assert [s["id"] for s in segs] == ["工作日时段客流", "周末时段客流"]
    assert [t["grain"] for t in d.recipe["tables"]] == [["日期", "时段"], ["日期", "时段"]]
    assert not any(c.id.startswith("merge:") for c in d.cards)
    assert d.recipe["sheets"][0]["context"][0]["prefer_prefix"] == "统计期"


def test_multiple_sheets_get_unique_segment_and_block_ids():
    wb = Workbook()
    for k, name in enumerate(["一号", "二号"]):
        ws = wb.active if k == 0 else wb.create_sheet(name)
        ws.title = name
        ws["B2"] = "2026年8月1日至2026年8月8日"
        rows = [(5, "合计（人次）"), (6, "甲（人次）"), (7, "乙（人次）"),
                (9, "日间时段客流（人次）"), (10, "7-8"), (11, "8-9"),
                (13, "夜间时段客流（人次）"), (14, "18-19"), (15, "19-20")]
        _crosstab(ws, rows)
    hidden = wb.create_sheet("底稿")
    hidden["A1"] = "不导入"
    hidden.sheet_state = "hidden"
    sc, grids = build(wb)
    d = S.draft(sc, grids, "x.xlsx", validate=ok_validate())
    assert d.complete, d.failures
    sheets = d.recipe["sheets"]
    assert [s["id"] for s in sheets] == ["s1", "s2"]
    block_ids = [b["id"] for s in sheets for b in s["blocks"]]
    seg_ids = [g["id"] for s in sheets for b in s["blocks"] for g in b["segments"]]
    assert block_ids == ["交叉表", "交叉表_s2"]
    assert len(seg_ids) == len(set(seg_ids))
    assert "日间" in seg_ids and "日间_s2" in seg_ids
    names = [t["name"] for t in d.recipe["tables"]]
    assert len(names) == len(set(names))
    assert names[:2] == ["一号_按日", "时段客流"] and "时段客流_2" in names
    assert any(c.id == "hidden_sheet:底稿" for c in d.cards)
    assert d.cards[-1].id == "mode"


def test_unparseable_total_row_fails_with_reason():
    wb = Workbook()
    ws = wb.active
    ws.title = "表"
    ws["B2"] = "2026年8月1日至2026年8月8日"
    rows = [(5, "日间时段客流（人次）"), (6, "7-8"), (7, "8-9"), (8, "上午合计")]
    _crosstab(ws, rows)
    sc, grids = build(wb)
    d = S.draft(sc, grids, "x.xlsx", validate=ok_validate())
    assert not d.complete
    assert any("第 8 行像合计行，但标签无法确定区间" in f for f in d.failures)


def test_missing_period_asks_for_human_input():
    wb = Workbook()
    ws = wb.active
    ws.title = "表"
    _crosstab(ws, [(5, "甲（人次）"), (6, "乙（人次）")])
    sc, grids = build(wb)
    d = S.draft(sc, grids, "x.xlsx", validate=ok_validate())
    assert d.complete
    assert d.recipe["sheets"][0]["context"][0]["prefer_prefix"] is None
    q = next(q for q in d.questions if q.id == "q_period")
    assert [o["value"] for o in q.options] == ["human"] and q.effects == {"human": []}
    assert next(c for c in d.cards if c.id == "period:s1").title == "未找到统计期：每期人工录入"


def test_value_area_text_fails_but_model_version_has_no_value():
    wb = Workbook()
    ws = wb.active
    ws.title = "表"
    ws["B2"] = "2026年8月1日至2026年8月8日"
    _crosstab(ws, [(5, "甲（人次）"), (6, "乙（人次）")])
    ws["E6"] = "约987654"
    ws["F6"] = "暂缺"
    sc, grids = build(wb)
    d = S.draft(sc, grids, "x.xlsx", validate=ok_validate())
    assert not d.complete
    assert any("约987654" in f for f in d.failures)
    assert any("E6" in f for f in d.failures_for_model)
    assert not any("987654" in f or "暂缺" in f for f in d.failures_for_model)


def test_failures_for_model_scrubs_cell_values_from_dry_run(flow):
    """model_message 为空时只发 code 和坐标；即使 model_message 写进了格子里的数，也换成 <值>。"""
    sc, grids, _ = flow
    grids = dict(grids)
    g = grids["客流汇总"]
    cells = dict(g.cells)
    cells[(31, 3)] = GridCell(987654)
    cells[(31, 4)] = GridCell("123,457")
    grids["客流汇总"] = Grid(g.sheet, g.bounds, cells, g.merges, g.hidden_rows, g.hidden_cols)
    probs = [Problem("outside_number", "structure", "C31 的 987654 在数据区外", ["客流汇总!C31"]),
             Problem("value_not_number", "structure", "D31「123,457」", ["客流汇总!D31"],
                     model_message="D31 写着 123,457，不是数")]
    lines = S.problem_lines_for_model(probs, grids)
    assert lines[0] == "outside_number（客流汇总!C31）"
    assert lines[1] == "value_not_number：D31 写着 <值>，不是数（客流汇总!D31）"
    assert not any("987654" in x or "123,457" in x for x in lines)


def test_blank_and_formula_cells_in_value_area_become_questions():
    wb = Workbook()
    ws = wb.active
    ws.title = "表"
    ws["B2"] = "2026年8月1日至2026年8月8日"
    _crosstab(ws, [(5, "甲（人次）"), (6, "乙（人次）")])
    ws["D5"] = None
    ws["E6"] = "=E5*2"
    sc, grids = build(wb, {"表": {"E6": 40}})
    d = S.draft(sc, grids, "x.xlsx", validate=ok_validate())
    values = d.recipe["sheets"][0]["blocks"][0]["values"]
    assert values["blank"] == "reject" and values["formula"] == "reject"
    qs = {q.id: q for q in d.questions}
    assert qs["q_blank:交叉表"].default is None and qs["q_formula:交叉表"].default is None
    after = apply_ops(d.recipe, qs["q_blank:交叉表"].effects["null"] + qs["q_formula:交叉表"].effects["accept_cached"])
    v = after["sheets"][0]["blocks"][0]["values"]
    assert v["blank"] == "null" and v["formula"] == "accept_cached"
    # 紧凑形式（去掉默认值，values 整个不在）上重算的 effects 也能应用
    compact = canonical(d.recipe)
    assert "values" not in compact["sheets"][0]["blocks"][0]
    _, qs2 = S.questions_for(compact, d.facts, grids)
    q = next(x for x in qs2 if x.id == "q_blank:交叉表")
    assert apply_ops(compact, q.effects["null"])["sheets"][0]["blocks"][0]["values"]["blank"] == "null"


def test_hidden_rows_in_region_ask_without_default():
    wb = Workbook()
    ws = wb.active
    ws.title = "表"
    ws["B2"] = "2026年8月1日至2026年8月8日"
    _crosstab(ws, [(5, "甲（人次）"), (6, "乙（人次）"), (7, "丙（人次）")])
    ws.row_dimensions[6].hidden = True
    sc, grids = build(wb)
    d = S.draft(sc, grids, "x.xlsx", validate=ok_validate())
    q = next(q for q in d.questions if q.id == "q_hidden:表")
    assert q.default is None
    assert [o["value"] for o in q.options] == ["reject_if_any", "include", "exclude"]
    after = apply_ops(d.recipe, q.effects["exclude"])
    assert after["sheets"][0]["hidden"] == {"rows": "exclude", "cols": "reject_if_any"}


def test_non_text_row_label_is_a_failure():
    wb = Workbook()
    ws = wb.active
    ws.title = "表"
    ws["B2"] = "2026年8月1日至2026年8月8日"
    _crosstab(ws, [(5, "甲（人次）"), (6, "乙（人次）")])
    ws["B6"] = 2025
    sc, grids = build(wb)
    d = S.draft(sc, grids, "x.xlsx", validate=ok_validate())
    assert not d.complete
    assert any("B6" in f and "不是文字" in f for f in d.failures)
    assert "2025" not in json.dumps(d.recipe, ensure_ascii=False)


def test_unknown_unit_keeps_raw_name_and_gets_a_card():
    wb = Workbook()
    ws = wb.active
    ws.title = "表"
    ws["B2"] = "2026年8月1日至2026年8月8日"
    _crosstab(ws, [(5, "销量（箱）"), (6, "金额（万元）")])
    sc, grids = build(wb)
    d = S.draft(sc, grids, "x.xlsx", validate=ok_validate())
    seg = d.recipe["sheets"][0]["blocks"][0]["segments"][0]
    assert seg["measures"] == {"销量（箱）": "销量_箱", "金额（万元）": "金额"}
    assert d.recipe["tables"][0]["units"] == {"金额": "万元"}
    assert any(c.title == "「销量（箱）」的单位「箱」不在单位词表中" for c in d.cards)


# ==========================================================================
# 列表（P2-SPEC 6.1 第 11 条）
# ==========================================================================


def test_list_header_types_and_grain():
    wb = Workbook()
    ws = wb.active
    ws.title = "销售明细"
    ws.append(["订单号", "日期", "数量", "金额（元）", "备注"])
    rows = [["A001", dt.datetime(2026, 8, 1), 3, 12.5, "首单"],
            ["A002", dt.datetime(2026, 8, 2), 1, 7.0, None],
            ["A003", dt.datetime(2026, 8, 3), 2, 9.25, "加急"]]
    for r in rows:
        ws.append(r)
    sc, grids = build(wb)
    d = S.draft(sc, grids, "x.xlsx", validate=ok_validate())
    assert d.complete, d.failures
    block = d.recipe["sheets"][0]["blocks"][0]
    assert block["id"] == "列表1" and block["layout"] == "list" and block["table"] == "销售明细"
    assert [(c["header"], c["name"], c["type"]) for c in block["columns"]] == [
        ("订单号", "订单号", "TEXT"), ("日期", "日期", "DATE"), ("数量", "数量", "INTEGER"),
        ("金额（元）", "金额", "REAL"), ("备注", "备注", "TEXT")]
    assert d.recipe["tables"] == [{"name": "销售明细", "grain": ["订单号"], "kind": "data",
                                   "units": {"金额": "元"}, "note": ""}]
    assert d.recipe["sheets"][0]["context"] == []
    assert d.cards[0].id == "list:列表1"


def test_list_two_row_header_with_merged_upper_row():
    wb = Workbook()
    ws = wb.active
    ws.title = "汇总"
    ws["A1"] = "地区"
    ws.merge_cells("A1:A2")
    ws["B1"] = "销售"
    ws.merge_cells("B1:C1")
    ws["B2"] = "数量"
    ws["C2"] = "金额"
    for r, row in enumerate([["华东", 10, 100], ["华北", 20, 120]], start=3):
        for c, v in enumerate(row, start=1):
            ws.cell(r, c, v)
    sc, grids = build(wb)
    d = S.draft(sc, grids, "x.xlsx", validate=ok_validate())
    block = d.recipe["sheets"][0]["blocks"][0]
    assert block["header_rows"] == 2
    assert [c["header"] for c in block["columns"]] == ["地区", "销售_数量", "销售_金额"]
    assert d.complete, d.failures


def test_list_blank_row_in_middle_is_skipped_with_card():
    wb = Workbook()
    ws = wb.active
    ws.title = "明细"
    for row in [["地区", "产品", "金额"], ["华东", "甲", 10], ["华东", "乙", 20], [None, None, None],
                ["华北", "甲", 30], ["华北", "乙", 40]]:
        ws.append(row)
    sc, grids = build(wb)
    d = S.draft(sc, grids, "x.xlsx", validate=ok_validate())
    block = d.recipe["sheets"][0]["blocks"][0]
    assert block["rows"]["blank_rows"] == "skip"
    assert len(d.recipe["sheets"][0]["blocks"]) == 1
    assert any(c.id == "blank_skip:列表1" for c in d.cards)


def test_list_stacked_tables_become_two_blocks():
    wb = Workbook()
    ws = wb.active
    ws.title = "汇总"
    ws["B1"] = "一、销售"
    for r, row in enumerate([["地区", "销量", "金额"], ["华东", 10, 100], ["华北", 20, 120], ["华南", 15, 80]], start=2):
        for c, v in enumerate(row, start=2):
            ws.cell(r, c, v)
    ws["B7"] = "二、费用"
    for r, row in enumerate([["部门", "费用"], ["甲部", 5], ["乙部", 7], ["丙部", 8]], start=8):
        for c, v in enumerate(row, start=2):
            ws.cell(r, c, v)
    sc, grids = build(wb)
    d = S.draft(sc, grids, "x.xlsx", validate=ok_validate())
    blocks = d.recipe["sheets"][0]["blocks"]
    assert [(b["id"], b["table"]) for b in blocks] == [("列表1", "销售"), ("列表2", "费用")]
    assert [c["header"] for c in blocks[1]["columns"]] == ["部门", "费用"]
    assert d.complete, d.failures
    out = next(c for c in d.cards if c.id == "outside:s1")
    assert out.cells == ["汇总!B1", "汇总!B7"]


def test_list_side_by_side_tables_split_on_empty_column():
    wb = Workbook()
    ws = wb.active
    ws.title = "并排"
    for r, row in enumerate([["地区", "金额"], ["华东", 1], ["华北", 2]], start=1):
        ws.cell(r, 1, row[0])
        ws.cell(r, 2, row[1])
    for r, row in enumerate([["部门", "费用"], ["甲部", 5], ["乙部", 7]], start=1):
        ws.cell(r, 4, row[0])
        ws.cell(r, 5, row[1])
    sc, grids = build(wb)
    d = S.draft(sc, grids, "x.xlsx", validate=ok_validate())
    blocks = d.recipe["sheets"][0]["blocks"]
    assert [[c["header"] for c in b["columns"]] for b in blocks] == [["地区", "金额"], ["部门", "费用"]]
    assert [b["table"] for b in blocks] == ["并排", "并排_2"]


def test_list_total_row_is_verified_and_kept():
    wb = Workbook()
    ws = wb.active
    ws.title = "销售"
    for row in [["地区", "销量", "金额"], ["华东", 10, 100], ["华北", 20, 120], ["合计", 30, 220]]:
        ws.append(row)
    sc, grids = build(wb)
    d = S.draft(sc, grids, "x.xlsx", validate=ok_validate())
    block = d.recipe["sheets"][0]["blocks"][0]
    assert block["rows"]["total_row"] == {"label_column": "地区", "pick": "合计", "keep_as": "销售_表内合计"}
    assert [t["name"] for t in d.recipe["tables"]] == ["销售", "销售_表内合计"]
    assert d.recipe["tables"][1]["grain"] == ["合计项"] and d.recipe["tables"][1]["kind"] == "reported_total"
    assert any(c.id == "total_row:列表1" and "第 4 行" in c.title for c in d.cards)


def test_list_vertical_merge_in_text_column_fills_with_question():
    wb = Workbook()
    ws = wb.active
    ws.title = "分组"
    for row in [["地区", "产品", "金额"], ["华东", "甲", 10], [None, "乙", 20], ["华北", "甲", 30], [None, "乙", 40]]:
        ws.append(row)
    ws.merge_cells("A2:A3")
    ws.merge_cells("A4:A5")
    sc, grids = build(wb)
    d = S.draft(sc, grids, "x.xlsx", validate=ok_validate())
    block = d.recipe["sheets"][0]["blocks"][0]
    assert block["merged_data"] == "fill"
    q = next(q for q in d.questions if q.id == "q_merged:列表1")
    assert q.default is None
    assert apply_ops(d.recipe, q.effects["reject"])["sheets"][0]["blocks"][0]["merged_data"] == "reject"


def test_list_header_with_year_gets_a_card():
    wb = Workbook()
    ws = wb.active
    ws.title = "年度"
    for row in [["地区", "2026年上半年_销量"], ["华东", 10], ["华北", 20]]:
        ws.append(row)
    sc, grids = build(wb)
    d = S.draft(sc, grids, "x.xlsx", validate=ok_validate())
    assert any(c.id == "header_year:列表1" for c in d.cards)


def test_list_period_cell_outside_becomes_context():
    wb = Workbook()
    ws = wb.active
    ws.title = "销售"
    ws["A1"] = "统计时间范围：2026年8月1日至2026年8月31日"
    for r, row in enumerate([["地区", "金额"], ["华东", 1], ["华北", 2]], start=3):
        ws.cell(r, 1, row[0])
        ws.cell(r, 2, row[1])
    sc, grids = build(wb)
    d = S.draft(sc, grids, "x.xlsx", validate=ok_validate())
    assert d.recipe["sheets"][0]["context"][0]["prefer_prefix"] == "统计时间范围"
    assert any(c.id == "period:s1" for c in d.cards)


def test_unrecognized_key_value_form():
    wb = Workbook()
    ws = wb.active
    ws.title = "登记表"
    ws["A1"] = "项目登记表"
    for r, (k, v) in enumerate([("项目名称：", "某工程"), ("负责人：", "张三"), ("预算（万元）：", 120),
                                ("开始日期：", dt.datetime(2026, 8, 1)), ("备注：", "无")], start=3):
        ws.cell(r, 1, k)
        ws.cell(r, 2, v)
    sc, grids = build(wb)
    d = S.draft(sc, grids, "x.xlsx", validate=lambda *a: pytest.fail("认不出时不该静态校验"))
    assert d.recipe is None and not d.complete
    assert any("认不出版式" in f for f in d.failures)
    assert d.failures_for_model and not any("张三" in f or "120" in f for f in d.failures_for_model)

    # 不带冒号的键值表单：表头像，但下面没有一列是统一的数字或日期
    wb2 = Workbook()
    ws2 = wb2.active
    ws2.title = "登记表"
    for k, v in [("项目名称", "某工程"), ("负责人", "张三"), ("预算", 120), ("部门", "工程部")]:
        ws2.append([k, v])
    sc2, grids2 = build(wb2)
    d2 = S.draft(sc2, grids2, "x.xlsx", validate=ok_validate())
    assert not d2.complete
    assert any("不是列表" in f for f in d2.failures)


def test_draft_never_raises_on_odd_input():
    wb = Workbook()
    ws = wb.active
    ws.title = "空"
    sc, grids = build(wb)
    d = S.draft(sc, grids, "x.xlsx", validate=ok_validate())
    assert d.recipe is None and not d.complete and d.failures


def test_layout_hints():
    wb, cache, _ = flow_book()
    sc, grids = build(wb, cache)
    h = S.layout_hints(grids["客流汇总"])
    assert (h.axis_row, h.label_col) == (4, 2)
    wb2 = Workbook()
    ws = wb2.active
    ws.title = "表"
    ws["A1"] = "标题"
    for row in [["地区", "金额"], ["华东", 1], ["华北", 2]]:
        ws.append(row)
    _, grids2 = build(wb2)
    h2 = S.layout_hints(grids2["表"])
    assert h2.axis_row is None and h2.header_rows == {2} and h2.top_header_row == 2


def test_constants_match_period1():
    from app.data import tabular

    assert S.DATE_HEADER_MIN == tabular.DATE_HEADER_MIN == 7
    assert S.DRAFT_MAX_ROWS == 500


def test_fact_search_stays_fast_on_large_measures_segment():
    """几十行的指标段不能让组合数爆炸（暂存时起草要在一秒量级内完成）。"""
    import time

    wb = Workbook()
    ws = wb.active
    ws.title = "大段"
    ws["B2"] = "2026年8月1日至2026年8月31日"
    rows = [(5 + i, f"指标{chr(0x4e00 + i)}（人次）") for i in range(45)]
    _crosstab(ws, rows, days=31, value=lambda r, c: (r * 131 + c * 17) % 997 + 3)
    sc, grids = build(wb)
    t0 = time.perf_counter()
    S.detect_facts(grids, sc)
    assert time.perf_counter() - t0 < 2.0


def test_question_ids_are_unique_across_blocks():
    wb = Workbook()
    for k, name in enumerate(["一号", "二号"]):
        ws = wb.active if k == 0 else wb.create_sheet(name)
        ws.title = name
        ws["B2"] = "2026年8月1日至2026年8月8日"
        _crosstab(ws, [(5, "甲（人次）"), (6, "乙（人次）")])
        ws["D6"] = "·"
    sc, grids = build(wb)
    d = S.draft(sc, grids, "x.xlsx", validate=ok_validate())
    ids = [q.id for q in d.questions]
    assert len(ids) == len(set(ids))
    assert "q_placeholder:·" in ids and "q_placeholder:·@交叉表_s2" in ids
    assert {c.question for c in d.cards if c.question} <= set(ids)


# ==========================================================================
# 列表：数据行不当表头、合计字样不在第一列或在中间
# ==========================================================================


def _roster(blank_before_wang: bool = False) -> Workbook:
    """名单：第 4 行（王五）的金额空着。这一行全是文字、互不相同、下面有数，单看「像表头」。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "名单"
    ws.append(["姓名", "电话", "金额", "备注"])
    people = [("张三", "138-1234-5678", 10, "老客户"), ("李四", "139-8765-4321", 20, "新客户"),
              ("王五", "137-2222-3333", None, "转介绍"), ("赵六", "136-2222-4444", 40, "老客户"),
              ("钱七", "135-4444-5555", 50, "新客户")]
    for i, p in enumerate(people):
        if blank_before_wang and i == 2:
            ws.append([None, None, None, None])
        ws.append(list(p))
    return wb


def test_list_row_with_blank_number_is_not_a_header():
    sc, grids = build(_roster())
    d = S.draft(sc, grids, "x.xlsx", validate=ok_validate())
    assert d.complete, d.failures
    blocks = d.recipe["sheets"][0]["blocks"]
    assert len(blocks) == 1
    assert [c["name"] for c in blocks[0]["columns"]] == ["姓名", "电话", "金额", "备注"]
    assert blocks[0]["columns"][2]["type"] == "INTEGER"
    assert S.layout_hints(grids["名单"]).header_rows == {1}
    blob = json.dumps(d.recipe, ensure_ascii=False)
    assert "王五" not in blob and "137" not in blob
    assert any(c.id == "list:列表1" and "第 2–6 行是数据" in c.reason for c in d.cards)
    # 空行之后紧跟着这种行：和上方同类型，跳过空行继续，不当成第二张表
    sc2, grids2 = build(_roster(blank_before_wang=True))
    d2 = S.draft(sc2, grids2, "x.xlsx", validate=ok_validate())
    assert d2.complete, d2.failures
    assert len(d2.recipe["sheets"][0]["blocks"]) == 1
    assert d2.recipe["sheets"][0]["blocks"][0]["rows"]["blank_rows"] == "skip"
    assert S.layout_hints(grids2["名单"]).header_rows == {1}


def test_list_stacked_with_title_but_no_blank_row_splits():
    """「小标题 + 表头」也算两张表的分界（没有空行时）。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "汇总"
    for row in [["地区", "销量", "金额"], ["分区甲", 10, 100], ["分区乙", 20, 120], ["二、费用", None, None],
                ["部门", "费用", None], ["部门甲", 5, None], ["部门乙", 7, None]]:
        ws.append(row)
    sc, grids = build(wb)
    d = S.draft(sc, grids, "x.xlsx", validate=ok_validate())
    blocks = d.recipe["sheets"][0]["blocks"]
    assert [[c["header"] for c in b["columns"]] for b in blocks] == [["地区", "销量", "金额"], ["部门", "费用"]]
    assert blocks[1]["table"] == "费用"
    assert d.complete, d.failures


def test_list_header_glued_to_data_is_a_failure_not_a_split():
    """另一张表的表头紧挨着上方的数据（没有空行或标题）：不切开、不把它当表头，起草失败写明原因。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "汇总"
    for row in [["地区", "销量"], ["分区甲", 10], ["分区乙", 20], ["部门", "费用"], ["部门甲", 5], ["部门乙", 7]]:
        ws.append(row)
    sc, grids = build(wb)
    d = S.draft(sc, grids, "x.xlsx", validate=ok_validate())
    assert not d.complete
    assert len(d.recipe["sheets"][0]["blocks"]) == 1
    assert any("第 4 行像另一张表的表头，B 列应是数字或日期" in f and "没有空行或标题" in f for f in d.failures)
    assert S.layout_hints(grids["汇总"]).header_rows == {1}


def test_list_total_word_outside_first_column_fails():
    wb = Workbook()
    ws = wb.active
    ws.title = "销售"
    ws.append(["日期", "地区", "金额"])
    for i, (zone, v) in enumerate([("分区甲", 100), ("分区乙", 120), ("分区丙", 130), ("分区丁", 140),
                                   ("分区戊", 160)]):
        ws.append([dt.datetime(2026, 8, 1 + i), zone, v])
    ws.append([None, "合计", 650])
    sc, grids = build(wb)
    d = S.draft(sc, grids, "x.xlsx", validate=ok_validate())
    assert not d.complete
    assert "工作表「销售」：第 7 行像合计行（B 列「合计」），但合计字样不在第一列，无法按合计行核对" in d.failures
    assert any("第 7 行像合计行（B 列以「合计」开头）" in f for f in d.failures_for_model)
    assert not any("650" in f for f in d.failures + d.failures_for_model)
    assert d.recipe["sheets"][0]["blocks"][0]["rows"]["total_row"] is None


def test_list_subtotal_rows_in_middle_fail():
    # 小计写在第二列、夹在中间
    wb = Workbook()
    ws = wb.active
    ws.title = "门店"
    for row in [["地区", "门店", "金额"], ["分区甲", "门店一", 100], ["分区甲", "门店二", 200], ["分区甲", "小计", 300],
                ["分区乙", "门店三", 150], ["分区乙", "门店四", 150], ["分区乙", "小计", 300]]:
        ws.append(row)
    sc, grids = build(wb)
    d = S.draft(sc, grids, "x.xlsx", validate=ok_validate())
    assert not d.complete
    assert [f for f in d.failures if "像合计行" in f] == [
        "工作表「门店」：第 4 行像合计行（B 列「小计」），但合计字样不在第一列，无法按合计行核对",
        "工作表「门店」：第 7 行像合计行（B 列「小计」），但合计字样不在第一列，无法按合计行核对"]
    # 小计写在第一列、但下面还有明细：不是末尾的合计行，整块照样收进来，不另报「数据区外的数字」
    wb2 = Workbook()
    ws2 = wb2.active
    ws2.title = "门店"
    for row in [["门店", "金额"], ["门店一", 100], ["门店二", 200], ["小计", 300], ["门店三", 150], ["门店四", 150],
                ["合计", 600]]:
        ws2.append(row)
    sc2, grids2 = build(wb2)
    d2 = S.draft(sc2, grids2, "x.xlsx", validate=ok_validate())
    assert not d2.complete
    assert d2.failures == ["工作表「门店」：第 4 行像合计行（A 列「小计」），但不在列表末尾，无法确定它汇总的是哪几行"]
    block = d2.recipe["sheets"][0]["blocks"][0]
    assert block["rows"]["total_row"]["pick"] == "合计"
    # 末尾第一列的合计行照旧改作核对，不算失败（test_list_total_row_is_verified_and_kept）


# ==========================================================================
# 系统发现按配方改写：只在事实所在的工作表里找
# ==========================================================================


def _two_flow_sheets(names=("客流汇总", "客流汇总二")) -> Workbook:
    wb = Workbook()
    for k, name in enumerate(names):
        ws = wb.active if k == 0 else wb.create_sheet(name)
        ws.title = name
        ws["B2"] = "2026年8月1日至2026年8月8日"
        rows = [(5, "全日（人次）"), (6, "甲（人次）"), (7, "乙（人次）"), (9, "日间时段客流（人次）"), (10, "7-8"),
                (11, "8-9")]

        def val(r, c, k=k):
            jia, yi = 10 + c + k, 20 + 2 * c
            return {5: jia + yi, 6: jia, 7: yi}.get(r, (r * 7 + c * 3) % 50 + 10)
        _crosstab(ws, rows, value=val)
    return wb


def test_remap_facts_stays_on_the_fact_sheet():
    sc, grids = build(_two_flow_sheets())
    d = S.draft(sc, grids, "x.xlsx", validate=ok_validate())
    assert d.complete, d.failures
    by_id = {f.id: f for f in d.facts.facts}
    assert by_id["F3"].sheet == "客流汇总二" and by_id["F3"].detail["segment"] == "客流汇总二_按日"
    # 规则草稿原样按配方改写：一个字没改，事实也不能被改指到第一张表的分段
    remapped = {f.id: f.detail for f in S.detect_facts(grids, sc, recipe=d.recipe).facts}
    assert remapped == {f.id: f.detail for f in d.facts.facts}
    assert remapped["F4"]["b_segment"] == "客流汇总二_按日"
    # 第二张表的宽表和明细分段改了名：按名字认得上的同表分段，明细分段按标题认
    renamed = rename_wide(d.recipe, "客流汇总二_按日", "日客流二")
    renamed = json.loads(json.dumps(renamed, ensure_ascii=False).replace('"id": "日间时段客流_2"', '"id": "日间二"'))
    seg_ids = [g["id"] for b in renamed["sheets"][1]["blocks"] for g in b["segments"]]
    assert seg_ids == ["日客流二", "日间二"]
    remapped = {f.id: f.detail for f in S.detect_facts(grids, sc, recipe=renamed).facts}
    assert remapped["F1"]["segment"] == "客流汇总_按日" and remapped["F2"]["a"] == ["日间时段客流"]
    assert remapped["F3"]["segment"] == "日客流二"
    assert remapped["F4"]["b_segment"] == "日客流二" and remapped["F4"]["a"] == ["日间二"]


def test_remap_facts_two_long_tables_on_one_sheet():
    """一张工作表上两张长表（工作日、周末的时段相交，不合表）：改过 id 的明细分段按标题认，不把两张表拼在一起。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "时段"
    ws["B2"] = "统计期：2026年8月1日至2026年8月8日"
    rows = [(5, "全日（人次）"), (7, "工作日时段客流（人次）"), (8, "7-8"), (9, "8-9"),
            (11, "周末时段客流（人次）"), (12, "7-8"), (13, "8-9")]
    _crosstab(ws, rows)
    sc, grids = build(wb)
    d = S.draft(sc, grids, "x.xlsx", validate=ok_validate())
    ne = [f for f in d.facts.facts if f.kind == "not_equal_sum"]
    assert [f.detail["a"] for f in ne] == [["工作日时段客流"], ["周末时段客流"]]
    renamed = json.loads(json.dumps(d.recipe, ensure_ascii=False)
                         .replace('"id": "工作日时段客流"', '"id": "工作日"')
                         .replace('"id": "周末时段客流"', '"id": "周末"'))
    got = [f.detail["a"] for f in S.detect_facts(grids, sc, recipe=renamed).facts if f.kind == "not_equal_sum"]
    assert got == [["工作日"], ["周末"]]
    # 不给标题（直接调 remap_facts）时认不全：两张长表，不取「整张工作表的全部分段」，原样留着
    plain = S.remap_facts(d.facts, renamed)
    assert [f.detail["a"] for f in plain.facts if f.kind == "not_equal_sum"] == [["工作日时段客流"], ["周末时段客流"]]


def test_remap_facts_follows_renamed_sheet_only_when_unambiguous():
    sc, grids = build(_two_flow_sheets(names=("客流汇总",)))
    d = S.draft(sc, grids, "x.xlsx", validate=ok_validate())
    moved = json.loads(json.dumps(d.recipe, ensure_ascii=False))
    moved["sheets"][0]["match"]["name"] = "上月的表名"
    moved = rename_wide(moved, "客流汇总_按日", "日客流")
    # 文件里只有一张可见工作表、配方里一张可回退：按回退认
    assert S.detect_facts(grids, sc, recipe=moved).facts[0].detail["segment"] == "日客流"
    # 不告诉有几张可见工作表时不回退
    assert S.remap_facts(d.facts, moved).facts[0].detail["segment"] == "客流汇总_按日"


# ==========================================================================
# 卡片、命名、失败原因
# ==========================================================================


def test_hidden_sheet_card_is_not_duplicated_with_dry_run():
    from app.data.recipe_types import SheetsOut

    wb, cache, _ = flow_book()
    spare = wb.create_sheet("备用")
    spare["A1"] = "不导入"
    spare.sheet_state = "hidden"
    sc, grids = build(wb, cache)
    d = S.draft(sc, grids, "x.xlsx", validate=ok_validate(),
                dry_run=lambda r: Extraction(ok=True, sheets=SheetsOut(skipped_hidden=[{"sheet": "备用"}])))
    ids = [c.id for c in d.cards]
    assert len(ids) == len(set(ids))
    assert ids[-2:] == ["hidden_sheet:备用", "mode"]
    # 没有干跑时由 draft 按扫描结果补上
    d2 = S.draft(sc, grids, "x.xlsx", validate=ok_validate())
    assert [c.id for c in d2.cards][-2:] == ["hidden_sheet:备用", "mode"]


def test_deduped_long_table_keeps_a_meaningful_value_column():
    sc, grids = build(_two_flow_sheets())
    d = S.draft(sc, grids, "x.xlsx", validate=ok_validate())
    segs = [g for s in d.recipe["sheets"] for b in s["blocks"] for g in b["segments"] if g["role"] == "dimension"]
    assert [(g["table"], g["value"]) for g in segs] == [("日间时段客流", "客流"), ("日间时段客流_2", "客流")]
    ne = [r for r in d.recipe["relations"] if r["kind"] == "not_comparable"]
    assert [r["a"] for r in ne] == [{"table": "日间时段客流", "value": "客流"},
                                    {"table": "日间时段客流_2", "value": "客流"}]


def test_internal_errors_do_not_expose_exception_names(flow, monkeypatch, caplog):
    sc, grids, _ = flow

    def boom(*a, **k):
        raise KeyError("张三")
    monkeypatch.setattr(S, "_analyze", boom)
    with caplog.at_level("WARNING", logger="app.data.recipe_suggest"):
        d = S.draft(sc, grids, "x.xlsx", validate=ok_validate())
    assert not d.complete and d.recipe is None
    assert d.failures == ["规则起草未能完成：分析表格版式时出错。请在配方面板中粘贴一份完整的配方"]
    assert "KeyError" in caplog.text and "张三" not in caplog.text
    monkeypatch.undo()

    def bad_validate(data, facts, origin):
        raise RuntimeError("x")
    d2 = S.draft(sc, grids, "x.xlsx", validate=bad_validate)
    assert d2.failures == ["配方检查未能完成：检查时出错。请在配方面板中修改配方后重试"]

    def bad_dry(recipe):
        raise ValueError("x")
    d3 = S.draft(sc, grids, "x.xlsx", validate=ok_validate(), dry_run=bad_dry)
    assert d3.failures == ["按草稿检查表格未能完成：检查时出错。请在配方面板中修改配方后重试"]
    for f in [*d.failures, *d2.failures, *d3.failures,
              *d.failures_for_model, *d2.failures_for_model, *d3.failures_for_model]:
        assert not any(name in f for name in ("Error", "Exception"))


def test_long_list_with_many_blank_numbers_stays_one_block_and_fast():
    """几百行、隔行空着金额的名单：每一行都「像表头」，既不能切成几百块，也不能逐行从头推类型（行数的平方）。"""
    import time

    wb = Workbook()
    ws = wb.active
    ws.title = "名单"
    ws.append(["姓名", "部门", "金额", "备注"])
    for i in range(499):
        ws.append([f"人员{i:03d}", f"部门{'甲乙丙'[i % 3]}", None if i % 2 else i + 1, f"说明{i}"])
    sc, grids = build(wb)
    t0 = time.perf_counter()
    d = S.draft(sc, grids, "x.xlsx", validate=ok_validate())
    S.layout_hints(grids["名单"])
    assert time.perf_counter() - t0 < 1.0
    assert d.complete, d.failures
    assert len(d.recipe["sheets"][0]["blocks"]) == 1
