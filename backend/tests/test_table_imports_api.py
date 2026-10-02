"""按配方导入的接口（期 2，WP-5c）：真管线的端到端，加上用假 Pipeline 测的接口形状与错误码。

真管线（default_pipeline）跑「结构仿照客流表」的合成夹具（tests/fixtures/xlsx/flow.py，假名、随机数）：
- 暂存 → 规则起草 → 回答问题 → 改表名（PUT）→ 试运行 → 提交；
- 同版式的下月文件「上传新一期」→ 回执与差异卡 → 启用；
- 同一文件再传 → same_as_import → 提交不新建记录；
- 简单导入的源切换为按配方导入。
其余（错误码、AI 起草的门槛、并发、路由不被截走）用 test_recipe_imports 里的假 Pipeline。不调用任何真实模型。
"""
from __future__ import annotations

import datetime as dt
import json
import uuid
from pathlib import Path
from typing import Any

import pytest
from httpx import ASGITransport, AsyncClient
from sqlalchemy import select

from app.core import artifact_store
from app.core.config import settings
from app.data import recipe, table_versions
from app.data.engine import engines, run_query
from app.db.base import SessionLocal
from app.db.models import DataSource, ImportStaging, SourceSnapshot, TableImport, TableRecipe
from app.main import app
from tests.fixtures.xlsx.flow import FLOW_EXPECT, flow_workbook
from tests.test_recipe_imports import (  # noqa: F401 - fake 是夹具，导入即可用
    FLOW_FULL, api_stage, fake, raw_bytes, recipe_source, staged,
)

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
FLOW_RECIPE = json.loads((Path(__file__).parent / "fixtures" / "recipes" / "flow_recipe.json").read_text("utf-8"))
AUG, SEP = dt.date(2026, 8, 1), dt.date(2026, 9, 1)


@pytest.fixture(autouse=True)
def store(monkeypatch, tmp_path):
    """每个测试一个独立的数据目录：上传目录、原件存档、工件都在里面。"""
    data = tmp_path / "data"
    (data / "uploads").mkdir(parents=True)
    monkeypatch.setattr(settings, "data_dir", data)
    yield data


@pytest.fixture
async def client():
    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as c:
        yield c


def unique(prefix: str = "wp5c") -> str:
    return f"{prefix}_{uuid.uuid4().hex[:10]}"


async def stage(client, name: str, raw: bytes, filename: str, **form: str):
    return await client.post("/api/datasources/imports/stage", files={"file": (filename, raw, XLSX)},
                             data={"name": name, **form})


async def scalar(source_id: str, sql: str):
    async with SessionLocal() as session:
        row = await session.get(DataSource, source_id)
        view = await table_versions.resolve_source(session, row)
    try:
        return (await run_query(view, sql)).rows[0][0]
    finally:
        await engines.invalidate(view.id)


async def get(model, key):
    async with SessionLocal() as session:
        return await session.get(model, key)


def rename_wide(rec: dict[str, Any], old: str, new: str) -> dict[str, Any]:
    """把规则起草的宽表名改成 new：表名、写入它的分段（id 一起改）、关系里引用的表名。"""
    rec = json.loads(json.dumps(rec, ensure_ascii=False))
    for sheet in rec["sheets"]:
        for block in sheet["blocks"]:
            for seg in block.get("segments", []):
                if seg.get("table") == old:
                    seg["table"] = new
                if seg.get("id") == old:
                    seg["id"] = new
    for t in rec["tables"]:
        if t["name"] == old:
            t["name"] = new
    for r in rec["relations"]:
        if r.get("table") == old:
            r["table"] = new
        for side in ("a", "b"):
            if isinstance(r.get(side), dict) and r[side].get("table") == old:
                r[side]["table"] = new
    return rec


def all_confirms(staging: dict[str, Any]) -> list[str]:
    return [i["id"] for i in staging["trial"]["confirm_items"]]


# --------------------------------------------------------------------------
# 真管线：首次导入
# --------------------------------------------------------------------------


async def first_import(client, name: str, seed: int = 1) -> dict[str, Any]:
    """D00 走完首次导入，返回提交的响应。"""
    raw, fn = flow_workbook(AUG, 31, seed=seed)
    resp = await stage(client, name, raw, fn)
    assert resp.status_code == 201, resp.text
    st = resp.json()
    sid = st["id"]
    answers = {q["id"]: {"value": "register" if q["id"].startswith("q_relation") else "null"}
               for q in st["questions"]}
    resp = await client.post(f"/api/datasources/imports/{sid}/answers", json={"answers": answers})
    assert resp.status_code == 200, resp.text
    wide = next(t["name"] for t in st["recipe"]["tables"] if t["name"].endswith("_按日"))
    resp = await client.put(f"/api/datasources/imports/{sid}/recipe",
                            json={"recipe": rename_wide(resp.json()["recipe"], wide, "日客流")})
    assert resp.status_code == 200, resp.text
    resp = await client.post(f"/api/datasources/imports/{sid}/trial", json={})
    assert resp.status_code == 200, resp.text
    st = resp.json()
    assert st["trial"]["status"] == "passed", st["trial"]["problems"]
    resp = await client.post(f"/api/datasources/imports/{sid}/commit",
                             json={"trial_id": st["trial"]["trial_id"], "confirmations": all_confirms(st)})
    assert resp.status_code == 201, resp.text
    return {"staging": st, "commit": resp.json()}


async def test_real_pipeline_first_import_stage_draft_answer_trial_commit(client):
    name = unique()
    raw, fn = flow_workbook(AUG, 31, seed=3)
    resp = await stage(client, name, raw, fn)
    assert resp.status_code == 201, resp.text
    st = resp.json()
    assert st["kind"] == "first" and st["status"] == "drafting"
    assert st["source"]["exists"] is False and st["source"]["import_mode"] is None
    assert st["draft"]["complete"] is True and "failures_for_model" not in st["draft"]
    assert st["ai"]["offered"] is False
    assert len(st["cards"]) == 10
    assert st["recipe_problems"] == [] and st["recipe_origin"] == "rules"
    assert st["grids"][0]["sheet"] == "客流汇总" and st["grids"][0]["bounds"] == "B2:AG30"
    assert st["sheets"][0]["nonempty"] == FLOW_EXPECT["nonempty"]
    assert st["marks"], "干跑的区域标记"
    qids = [q["id"] for q in st["questions"]]
    assert {"q_relation:F1", "q_relation:F2"} <= set(qids)
    sid = st["id"]
    # 新源在提交前不存在
    async with SessionLocal() as session:
        assert (await session.execute(select(DataSource).where(DataSource.name == name))).first() is None

    # 回答问题：先选「不登记」再改选「登记」，等于直接选「登记」
    first = await client.post(f"/api/datasources/imports/{sid}/answers", json={"answers": {
        "q_relation:F1": {"value": "dismiss", "reason": "合成理由：判断为巧合"}}})
    assert first.status_code == 200, first.text
    assert next(r for r in first.json()["recipe"]["relations"] if r["id"] == "R1")["kind"] == "dismissed"
    answers = {q: {"value": "register" if q.startswith("q_relation") else "null"} for q in qids}
    resp = await client.post(f"/api/datasources/imports/{sid}/answers", json={"answers": answers})
    assert resp.status_code == 200, resp.text
    st = resp.json()
    assert st["answers"]["q_relation:F1"] == {"value": "register", "reason": None}
    assert [q["id"] for q in st["questions"]] == qids, "问题集合按起点配方生成，回答不改变它"

    # 改表名（PUT）：起点换成它，回答清空；与参考配方逐字相同
    resp = await client.put(f"/api/datasources/imports/{sid}/recipe",
                            json={"recipe": rename_wide(st["recipe"], "客流汇总_按日", "日客流")})
    assert resp.status_code == 200, resp.text
    st = resp.json()
    assert st["answers"] == {} and st["recipe_origin"] == "manual" and st["recipe_problems"] == []
    ref, _ = recipe.validate_recipe(FLOW_RECIPE, origin="manual")
    mine, _ = recipe.validate_recipe(st["recipe"], origin="manual")
    assert recipe.recipe_sha256(mine) == recipe.recipe_sha256(ref)

    resp = await client.post(f"/api/datasources/imports/{sid}/trial", json={})
    assert resp.status_code == 200, resp.text
    st = resp.json()
    trial = st["trial"]
    assert st["status"] == "trialed" and trial["status"] == "passed", trial["problems"]
    assert set(all_confirms(st)) == set(FLOW_EXPECT["confirms"])
    assert {c["id"]: c["status"] for c in trial["checks"]} == FLOW_EXPECT["checks"]
    assert trial["diff"] is None and trial["same_as_import"] is None
    assert "extraction" not in trial and "null_counts" not in trial
    assert trial["receipt"]["canonicalized_total"] == 0 and trial["receipt"]["db_sha256"]
    assert {t["name"]: t["rows"] for t in trial["receipt"]["tables"]} == FLOW_EXPECT["rows"]
    assert "已逐日核对" in trial["notes"]["日客流"]["comment"]
    assert trial["notes"]["日客流"]["columns"]["全日客流"] == "单位：人次"

    # 少勾一项 → 422，detail 列出缺的那一项
    some = all_confirms(st)
    resp = await client.post(f"/api/datasources/imports/{sid}/commit",
                             json={"trial_id": trial["trial_id"], "confirmations": some[1:]})
    assert resp.status_code == 422 and resp.json()["code"] == "confirm_required"
    assert some[0] in resp.json()["detail"]

    resp = await client.post(f"/api/datasources/imports/{sid}/commit",
                             json={"trial_id": trial["trial_id"], "confirmations": some, "signed_by": "测试员"})
    assert resp.status_code == 201, resp.text
    out = resp.json()
    assert out["unchanged"] is False and out["build_reused"] is False
    src = out["source"]
    assert src["name"] == name and src["origin"] == "upload" and src["import_mode"] == "recipe"
    assert src["current_recipe"]["id"] == out["recipe_id"] and src["current_recipe"]["seq"] == 1
    assert src["current_recipe"]["origin"] == "manual"
    assert src["open_staging"] is None
    assert await scalar(src["id"], 'SELECT COUNT(*) FROM "时段客流"') == FLOW_EXPECT["rows"]["时段客流"]

    # 表结构里的说明；schema_cache 承诺了导入清单
    resp = await client.get(f"/api/datasources/{src['id']}/schema", params={"table": "日客流"})
    assert resp.status_code == 200, resp.text
    schema = resp.json()
    assert "已逐日核对" in json.dumps(schema, ensure_ascii=False)
    snap = await get(SourceSnapshot, out["snapshot_id"])
    imp = await get(TableImport, out["import_id"])
    assert snap.schema_cache["import_mode"] == "recipe"
    assert snap.schema_cache["import_manifests"] == [imp.manifest_artifact]
    manifest = artifact_store.load(imp.manifest_artifact)
    assert manifest["format"] == "agentlab-import-manifest/1"
    assert manifest["db_sha256"] == snap.db_sha256 == manifest["trial_db_sha256"]
    for key in ("axes", "tables", "ledger", "period", "labels", "outside_text", "lineage", "canonicalized"):
        assert key in manifest["receipt"], key
    assert "problems" not in manifest["receipt"]
    assert manifest["recipe"]["id"] == out["recipe_id"] and manifest["signed_by"] == {"name": "测试员", "verified": False}
    assert {c["id"] for c in manifest["checks"]} >= {"C1", "C2"}
    assert {c["id"] for c in imp.checks} >= {"C1", "C2"}
    assert imp.period_start == "2026-08-01" and imp.period_end == "2026-08-31" and imp.signed_by == "测试员"
    assert imp.recipe_id == out["recipe_id"] and imp.staging_id == sid
    assert {c["id"] for c in imp.confirmations} == set(some)
    rec = await get(TableRecipe, out["recipe_id"])
    assert rec.status == "active" and rec.seq == 1 and rec.staging_id == sid
    # 暂存区结束即瘦身
    resp = await client.get(f"/api/datasources/imports/{sid}")
    st = resp.json()
    assert st["status"] == "committed" and st["grids"] == [] and st["trial"] is None
    assert st["committed_import_id"] == out["import_id"]
    row = await get(ImportStaging, sid)
    assert row.scan is None and row.grid_preview is None and row.trial_path is None and row.recipe_sha256
    # 已结束的暂存区不能再写
    resp = await client.post(f"/api/datasources/imports/{sid}/trial", json={})
    assert resp.status_code == 409 and resp.json()["code"] == "staging_closed"
    # 按配方导入的源：简单上传 409、当前配方、导入记录的新字段
    resp = await client.post("/api/datasources/upload", files={"file": (fn, raw, XLSX)}, data={"name": name})
    assert resp.status_code == 409 and resp.json()["code"] == "recipe_source"
    resp = await client.get(f"/api/datasources/{src['id']}/recipe")
    assert resp.status_code == 200, resp.text
    cur = resp.json()
    assert cur["recipe_id"] == out["recipe_id"] and cur["recipe"]["mode"] == "replace"
    resp = await client.get(f"/api/datasources/{src['id']}/imports")
    item = resp.json()[0]
    assert item["recipe_id"] == out["recipe_id"] and item["period_start"] == "2026-08-01"
    assert item["overrides"] == 0 and item["waivers"] == 0 and item["signed_by"] == "测试员"


# --------------------------------------------------------------------------
# 真管线：上传新一期
# --------------------------------------------------------------------------


async def test_real_pipeline_reupload_next_month_receipt_diff_and_activate(client):
    name = unique()
    done = await first_import(client, name)
    source_id = done["commit"]["source"]["id"]
    first_snapshot = done["commit"]["snapshot_id"]
    raw, fn = flow_workbook(SEP, 30, seed=7)
    resp = await client.post(f"/api/datasources/{source_id}/reupload", files={"file": (fn, raw, XLSX)})
    assert resp.status_code == 201, resp.text
    st = resp.json()
    assert st["kind"] == "reupload" and st["status"] == "trialed" and st["draft"] is None
    assert st["ai"]["offered"] is False and st["source"]["import_mode"] == "recipe"
    trial = st["trial"]
    assert trial["status"] == "passed", trial["problems"]
    assert {d["kind"] for d in trial["diff"]} == {"period", "rows"}
    assert not any(d["requires_confirm"] for d in trial["diff"])
    # 配方没变：配方类确认项都不出，只剩本期类（这里没有）
    assert all_confirms(st) == []
    assert trial["base_snapshot_id"] == first_snapshot
    assert {t["name"]: t["rows"] for t in trial["receipt"]["tables"]} == {"日客流": 30, "时段客流": 510,
                                                                         "时段客流_表内合计": 90}
    # 数据源列表里看得到这个未结束的导入
    resp = await client.get("/api/datasources")
    row = next(s for s in resp.json() if s["id"] == source_id)
    assert row["open_staging"]["id"] == st["id"] and row["open_staging"]["kind"] == "reupload"

    resp = await client.post(f"/api/datasources/imports/{st['id']}/commit",
                             json={"trial_id": trial["trial_id"], "confirmations": []})
    assert resp.status_code == 201, resp.text
    out = resp.json()
    assert out["unchanged"] is False and out["recipe_id"] == done["commit"]["recipe_id"]
    assert out["source"]["current_snapshot"]["id"] == out["snapshot_id"] != first_snapshot
    assert await scalar(source_id, 'SELECT COUNT(*) FROM "日客流"') == 30
    async with SessionLocal() as session:
        recipes = list((await session.execute(
            select(TableRecipe).where(TableRecipe.source_id == source_id))).scalars())
        imps = list((await session.execute(
            select(TableImport).where(TableImport.source_id == source_id).order_by(TableImport.seq))).scalars())
    assert len(recipes) == 1, "配方没变：不新建配方"
    # 被替换的那一期：发布后的回收会把没有运行引用的版本标成 retired（原件按配方源最近 12 次保留）
    assert imps[0].status in ("superseded", "retired") and imps[1].status == "active"
    assert imps[0].raw_state == "kept"
    assert imps[1].recipe_id == recipes[0].id and imps[1].period_start == "2026-09-01"


async def test_real_pipeline_same_file_again_is_unchanged(client):
    name = unique()
    raw, fn = flow_workbook(AUG, 31, seed=11)
    done = await first_import(client, name, seed=11)
    source_id = done["commit"]["source"]["id"]
    resp = await client.post(f"/api/datasources/{source_id}/reupload", files={"file": (fn, raw, XLSX)})
    assert resp.status_code == 201, resp.text
    st = resp.json()
    trial = st["trial"]
    assert trial["same_as_import"]["id"] == done["commit"]["import_id"]
    assert trial["diff"] == []
    resp = await client.post(f"/api/datasources/imports/{st['id']}/commit",
                             json={"trial_id": trial["trial_id"], "confirmations": all_confirms(st)})
    assert resp.status_code == 201, resp.text
    out = resp.json()
    assert out["unchanged"] is True and out["import_id"] == done["commit"]["import_id"]
    assert out["snapshot_id"] == done["commit"]["snapshot_id"]
    async with SessionLocal() as session:
        n = len(list((await session.execute(
            select(TableImport).where(TableImport.source_id == source_id))).scalars()))
    assert n == 1
    row = await get(ImportStaging, st["id"])
    assert row.status == "committed" and row.trial_path is None


async def test_real_pipeline_changed_context_input_invalidates_earlier_trial_id(client):
    """trial_key 含人工录入（WP-5c：改配方、改录入后失效）：同一配方、同一统计期换了署名再试运行，trial_id
    必须换；拿前一次（另一份录入）的 trial_id 提交返回 409 trial_required，不能发布成最新那次的结果。"""
    raw, fn = flow_workbook(AUG, 31, seed=1201, variant="noperiod")
    resp = await stage(client, unique(), raw, fn)
    assert resp.status_code == 201, resp.text
    sid = resp.json()["id"]
    resp = await client.put(f"/api/datasources/imports/{sid}/recipe", json={"recipe": FLOW_RECIPE})
    assert resp.status_code == 200, resp.text
    resp = await client.post(f"/api/datasources/imports/{sid}/trial", json={})
    assert resp.status_code == 200 and resp.json()["trial"]["status"] == "needs_input", resp.text
    period = {"统计期": {"start": "2026-08-01", "end": "2026-08-31"}}
    resp = await client.post(f"/api/datasources/imports/{sid}/trial",
                             json={"context_inputs": period, "signed_by": "录入员甲"})
    first = resp.json()
    assert first["trial"]["status"] == "passed", first["trial"]["problems"]
    resp = await client.post(f"/api/datasources/imports/{sid}/trial",
                             json={"context_inputs": period, "signed_by": "录入员乙"})
    second = resp.json()["trial"]
    assert second["status"] == "passed" and second["trial_id"] != first["trial"]["trial_id"]
    resp = await client.post(f"/api/datasources/imports/{sid}/commit", json={
        "trial_id": first["trial"]["trial_id"], "confirmations": all_confirms(first)})
    assert resp.status_code == 409 and resp.json()["code"] == "trial_required", resp.text
    assert (await get(ImportStaging, sid)).status != "committed"


# --------------------------------------------------------------------------
# 真管线：AI 起草（假模型）
# --------------------------------------------------------------------------

KV_RECIPE = {
    "recipe_format": "agentlab-recipe/2",
    "sheets": [{"id": "s1", "match": {"name": "基本情况"}, "blocks": [{
        "id": "列表1", "layout": "list", "table": "基本情况", "after_title": "基本情况登记表",
        "columns": [{"header": "单位名称", "name": "项目", "type": "TEXT"},
                    {"header": "示例单位甲", "name": "取值", "type": "TEXT"}]}]}],
    "tables": [{"name": "基本情况"}],
}


async def test_real_pipeline_ai_draft_sends_exactly_the_previewed_text(client, monkeypatch):
    """规则起草认不出的表单（lab.kv_form）→ 查看要发送的全文 → 同意 → 假模型起草 → 试运行 → 提交。

    模型第一次收到的系统提示 + 提示词，与同意框里展示的全文逐字相同；不调用任何真实模型。
    """
    from langchain_core.messages import AIMessage

    from app.data import recipe_ai
    from app.data.recipe_types import AiAvailability
    from tests.fixtures.xlsx.lab import kv_form

    calls: list[list[Any]] = []

    class Model:
        async def ainvoke(self, messages):
            calls.append(list(messages))
            return AIMessage(content=json.dumps(KV_RECIPE, ensure_ascii=False),
                             usage_metadata={"input_tokens": 3000, "output_tokens": 400, "total_tokens": 3400})

    async def resolve(session):
        return Model(), "fake-model", AiAvailability(True, model="fake-model", provider="假接入")

    monkeypatch.setattr(recipe_ai, "resolve_draft_model", resolve)
    raw, fn = kv_form()
    resp = await stage(client, unique(), raw, fn)
    assert resp.status_code == 201, resp.text
    st = resp.json()
    assert st["draft"]["complete"] is False and st["ai"]["offered"] is True and st["ai"]["available"] is True
    sid = st["id"]
    resp = await client.get(f"/api/datasources/imports/{sid}/draft-ai/preview")
    assert resp.status_code == 200, resp.text
    pv = resp.json()
    assert pv["model"] == "fake-model" and "# 配方 JSON Schema" in pv["text"] and "# 单位词表" in pv["text"]
    assert "示例单位甲" not in pv["text"].split("# 工作表压缩表示")[0], "系统提示里没有表格内容"
    assert "35600" not in pv["text"] and "860" not in pv["text"], "数字格的值不外发"
    resp = await client.post(f"/api/datasources/imports/{sid}/draft-ai",
                             json={"consent": True, "preview_sha256": pv["sha256"], "signed_by": "测试员"})
    assert resp.status_code == 200, resp.text
    st = resp.json()
    first = calls[0]
    assert first[0].content + "\n\n" + first[1].content == pv["text"]
    assert st["recipe_origin"] == "ai" and st["recipe_problems"] == [] and st["cards"][0]["id"] == "ai_origin"
    assert st["ai_consents"][0]["preview_sha256"] == pv["sha256"] and st["ai_draft"]["total_tokens"] == 3400

    st = (await client.post(f"/api/datasources/imports/{sid}/trial", json={})).json()
    assert st["trial"]["status"] == "passed", st["trial"]["problems"]
    ids = all_confirms(st)
    assert {"ai_origin", "rename:基本情况.项目", "columns:基本情况"} <= set(ids)
    resp = await client.post(f"/api/datasources/imports/{sid}/commit",
                             json={"trial_id": st["trial"]["trial_id"],
                                   "confirmations": [i for i in ids if i != "ai_origin"]})
    assert resp.status_code == 422 and "ai_origin" in resp.json()["detail"]
    resp = await client.post(f"/api/datasources/imports/{sid}/commit",
                             json={"trial_id": st["trial"]["trial_id"], "confirmations": ids})
    assert resp.status_code == 201, resp.text
    out = resp.json()
    rec = await get(TableRecipe, out["recipe_id"])
    assert rec.origin == "ai" and rec.ai_usage["calls"][0]["input_tokens"] == 3000
    imp = await get(TableImport, out["import_id"])
    manifest = artifact_store.load(imp.manifest_artifact)
    assert manifest["ai"]["consents"][0]["preview_sha256"] == pv["sha256"]
    assert await scalar(out["source"]["id"], 'SELECT COUNT(*) FROM "基本情况"') > 0


async def test_real_pipeline_model_switch_after_preview_is_stale_and_model_resolved_once(client, monkeypatch):
    """AU-7 / SE-4：看完预览后助手的模型从本机换成外部服务 → 409 ai_preview_stale，不发送；重新查看再同意时，
    同意记录、实际调用、用量记的是同一个模型，POST 里只解析一次模型。"""
    from langchain_core.messages import AIMessage

    from app.data import recipe_ai
    from app.data.recipe_types import AiAvailability
    from tests.fixtures.xlsx.lab import kv_form

    sent: list[str] = []
    current = {"name": "local-model-A", "provider": "本机接入"}
    resolves: list[str] = []

    class Model:
        def __init__(self, name: str) -> None:
            self.name = name

        async def ainvoke(self, messages):
            sent.append(self.name)
            resolves.append("send")
            return AIMessage(content=json.dumps(KV_RECIPE, ensure_ascii=False),
                             usage_metadata={"input_tokens": 10, "output_tokens": 5, "total_tokens": 15})

    async def resolve(session):
        resolves.append(current["name"])
        return (Model(current["name"]), current["name"],
                AiAvailability(True, model=current["name"], provider=current["provider"]))

    monkeypatch.setattr(recipe_ai, "resolve_draft_model", resolve)
    raw, fn = kv_form()
    sid = (await stage(client, unique(), raw, fn)).json()["id"]
    pv = (await client.get(f"/api/datasources/imports/{sid}/draft-ai/preview")).json()
    assert pv["model"] == "local-model-A"
    current.update(name="cloud-model-B", provider="外部接入")
    resp = await client.post(f"/api/datasources/imports/{sid}/draft-ai",
                             json={"consent": True, "preview_sha256": pv["sha256"]})
    assert resp.status_code == 409 and resp.json()["code"] == "ai_preview_stale", resp.text
    assert sent == []
    pv = (await client.get(f"/api/datasources/imports/{sid}/draft-ai/preview")).json()
    assert pv["model"] == "cloud-model-B"
    resolves.clear()
    resp = await client.post(f"/api/datasources/imports/{sid}/draft-ai",
                             json={"consent": True, "preview_sha256": pv["sha256"]})
    assert resp.status_code == 200, resp.text
    st = resp.json()
    # 发送之前只解析了一次（同意的门槛里那次）；之后的是响应里查可用性，与发送无关
    assert resolves[:resolves.index("send")] == ["cloud-model-B"], resolves
    assert sent == ["cloud-model-B"]
    assert st["ai_consents"][-1]["model"] == "cloud-model-B" and st["ai_consents"][-1]["provider"] == "外部接入"
    assert [u["model"] for u in st["ai_draft"]["usage"]] == ["cloud-model-B"]


# --------------------------------------------------------------------------
# 假管线：接口形状与错误码
# --------------------------------------------------------------------------


async def test_stage_errors(client, fake, monkeypatch):
    resp = await api_stage(client, "Bad-Name")
    assert resp.status_code == 400 and resp.json()["code"] == "name_invalid"
    resp = await client.post("/api/datasources/imports/stage", data={"name": unique()})
    assert resp.status_code == 422 and resp.json()["code"] == "body_invalid"
    resp = await api_stage(client, unique(), raw=b"bad file")
    assert resp.status_code == 400 and resp.json()["code"] == "file_unreadable"
    assert "无法读取" in resp.json()["detail"]
    # 手工源同名 → name_taken；已按配方导入 → recipe_source_exists
    manual = unique()
    async with SessionLocal() as session:
        session.add(DataSource(name=manual, kind="sqlite", database="/tmp/合成.db"))
        await session.commit()
    resp = await api_stage(client, manual)
    assert resp.status_code == 409 and resp.json()["code"] == "name_taken"
    done = await recipe_source(client)
    resp = await api_stage(client, done["source"]["name"])
    assert resp.status_code == 409 and resp.json()["code"] == "recipe_source_exists"
    # 超限 → 413（两条上传路径共用 read_upload）
    monkeypatch.setattr(settings, "max_upload_mb", 0)
    resp = await api_stage(client, unique())
    assert resp.status_code == 413 and resp.json()["code"] == "file_too_large"
    resp = await client.post("/api/datasources/upload", files={"file": ("x.xlsx", b"PK", XLSX)},
                             data={"name": unique()})
    assert resp.status_code == 413 and resp.json()["code"] == "file_too_large"


async def test_not_found_and_body_invalid(client, fake):
    for method, path in (("get", "/api/datasources/imports/nope"), ("delete", "/api/datasources/imports/nope")):
        resp = await getattr(client, method)(path)
        assert resp.status_code == 404 and resp.json()["code"] == "staging_not_found", path
    resp = await client.post("/api/datasources/imports/nope/trial", json={})
    assert resp.status_code == 404 and resp.json()["code"] == "staging_not_found"
    st = await staged(client)
    sid = st["id"]
    for method, path, kw in (
        ("post", "answers", {"content": b"[1, 2]"}), ("put", "recipe", {"content": b"not json"}),
        ("put", "recipe", {"json": {"recipe": [1]}}), ("post", "answers", {"json": {"answers": ["x"]}}),
        ("post", "commit", {"json": {"trial_id": "x", "confirmations": "mode"}}),
        ("post", "commit", {"json": {"trial_id": "x", "acceptances": {"R1": "理由"}}}),
    ):
        headers = {"content-type": "application/json"} if "content" in kw else None
        resp = await getattr(client, method)(f"/api/datasources/imports/{sid}/{path}", headers=headers, **kw)
        assert resp.status_code == 422 and resp.json()["code"] == "body_invalid", (path, kw)
    resp = await client.post(f"/api/datasources/imports/{sid}/commit", json={})
    assert resp.status_code == 409 and resp.json()["code"] == "trial_required"


async def test_put_recipe_with_schema_problems_is_stored_but_cannot_trial(client, fake):
    st = await staged(client)
    sid = st["id"]
    broken = {"recipe_format": "agentlab-recipe/2", "tables": []}
    resp = await client.put(f"/api/datasources/imports/{sid}/recipe", json={"recipe": broken})
    assert resp.status_code == 200, resp.text
    out = resp.json()
    assert out["recipe"] == broken and out["recipe_problems"][0]["code"] == "schema"
    assert out["recipe_problems"][0]["path"]
    assert fake.calls["detect_facts"][-1] == broken
    resp = await client.post(f"/api/datasources/imports/{sid}/trial", json={})
    assert resp.status_code == 409 and resp.json()["code"] == "recipe_invalid"
    # 改回合法的配方：问题集合按新的起点重算
    fake.calls["questions_for"].clear()
    resp = await client.put(f"/api/datasources/imports/{sid}/recipe", json={"recipe": FLOW_FULL})
    assert resp.json()["recipe_problems"] == [] and fake.calls["questions_for"][-1]["recipe"] == FLOW_FULL


async def test_signed_by_from_header_and_body(client, fake):
    st = await staged(client)
    from urllib.parse import quote

    resp = await client.post(f"/api/datasources/imports/{st['id']}/trial", json={},
                             headers={"X-Actor": quote("页眉署名")})
    assert resp.status_code == 200
    assert (await get(ImportStaging, st["id"])).signed_by == "页眉署名"
    resp = await client.post(f"/api/datasources/imports/{st['id']}/trial", json={"signed_by": "请求体署名"},
                             headers={"X-Actor": quote("页眉署名")})
    assert (await get(ImportStaging, st["id"])).signed_by == "请求体署名"


async def test_datasource_out_new_fields(client, fake):
    manual = unique()
    async with SessionLocal() as session:
        session.add(DataSource(name=manual, kind="sqlite", database="/tmp/合成.db"))
        await session.commit()
    done = await recipe_source(client)
    rows = {s["name"]: s for s in (await client.get("/api/datasources")).json()}
    assert rows[manual]["import_mode"] is None and rows[manual]["current_recipe"] is None
    assert rows[manual]["open_staging"] is None
    mine = rows[done["source"]["name"]]
    assert mine["import_mode"] == "recipe"
    assert set(mine["current_recipe"]) == {"id", "seq", "origin", "activated_at", "signed_by"}
    re_ = (await client.post(f"/api/datasources/{mine['id']}/reupload",
                             files={"file": ("x.xlsx", raw_bytes(), XLSX)})).json()
    rows = {s["name"]: s for s in (await client.get("/api/datasources")).json()}
    assert rows[done["source"]["name"]]["open_staging"]["id"] == re_["id"]
    assert set(rows[done["source"]["name"]]["open_staging"]) == {"id", "kind", "status", "created_at"}


async def test_routes_are_not_taken_by_datasource_wildcards(client, fake):
    """/imports/… 和 /{source_id}/reupload|redraft|recipe 不被 datasources 的通配路由截走：每个打一次，看 code。"""
    resp = await client.get("/api/datasources/imports/imports")
    assert resp.json().get("code") == "staging_not_found"
    resp = await client.post("/api/datasources/nope/reupload", files={"file": ("x.xlsx", b"PK", XLSX)})
    assert resp.status_code == 404 and resp.json().get("code") == "source_not_found"
    resp = await client.get("/api/datasources/nope/recipe")
    assert resp.status_code == 404 and resp.json().get("code") == "source_not_found"
    resp = await client.delete("/api/datasources/imports/nope")
    assert resp.json().get("code") == "staging_not_found"
    resp = await client.get("/api/datasources/imports/nope/draft-ai/preview")
    assert resp.json().get("code") == "staging_not_found"


async def test_switch_from_simple_reports_mask_lost_and_compress_masks_case_insensitively(client, monkeypatch):
    """AU-6 / SE-9：简单导入遮罩了「工资_元」，切换后规则起草的列名是「工资」：确认项里必勾 mask_lost:工资_元。
    遮罩写成大小写不同的「PHONE」时，AI 压缩表示同样按遮罩列处理（与证据面板一致，不区分大小写）。"""
    import io

    from openpyxl import Workbook

    from app.data import recipe_imports

    wb = Workbook()
    ws = wb.active
    ws.title = "员工"
    ws.append(["工号", "姓名", "Phone", "工资（元）"])
    for i in range(12):
        ws.append([f"E{i:03d}", f"员工{i}", f"1380000{i:04d}", 5000 + i * 137])
    buf = io.BytesIO()
    wb.save(buf)
    raw = buf.getvalue()
    name = unique()
    resp = await client.post("/api/datasources/upload", files={"file": ("员工.xlsx", raw, XLSX)}, data={"name": name})
    assert resp.status_code == 201, resp.text
    src = resp.json()["source"]["id"]
    resp = await client.patch(f"/api/datasources/{src}", json={"options": {"mask_columns": ["工资_元", "PHONE"]}})
    assert resp.status_code == 200, resp.text
    st = (await stage(client, name, raw, "员工.xlsx")).json()
    assert st["kind"] == "switch"
    st = (await client.post(f"/api/datasources/imports/{st['id']}/trial", json={})).json()
    items = {i["id"]: i for i in st["trial"]["confirm_items"]}
    assert "mask_lost:工资_元" in items, list(items)
    assert "mask_lost:PHONE" not in items          # 新结构里还有 Phone（不区分大小写）
    # SE-9：压缩表示的遮罩表头不区分大小写
    from app.db.base import SessionLocal
    from app.db.models import ImportStaging

    async with SessionLocal() as session:
        staging = await session.get(ImportStaging, st["id"])
        heads = await recipe_imports._mask_headers(session, staging)
    assert "Phone" in heads and "工资（元）" in heads


async def test_bad_number_cell_is_file_unreadable_with_fixed_text(client):
    """SE-7：数字格的 <v> 里写了非数字：stage 回 400 file_unreadable（带 code、固定中文文案），不回英文原文和格子内容。"""
    from openpyxl import Workbook

    from tests.fixtures.xlsx import patch_parts, save

    wb = Workbook()
    ws = wb.active
    ws.title = "S"
    ws["A1"], ws["B1"] = "项目", "数量"
    for r in range(2, 8):
        ws.cell(r, 1, f"甲{r}")
        ws.cell(r, 2, r)
    raw = patch_parts(save(wb), {"xl/worksheets/sheet1.xml": lambda x: x.replace("<v>3</v>", "<v>机密值abc</v>", 1)})
    resp = await stage(client, unique(), raw, "bad.xlsx")
    assert resp.status_code == 400 and resp.json()["code"] == "file_unreadable", resp.text
    assert "机密值abc" not in resp.text and "invalid literal" not in resp.text


async def test_json_body_size_and_depth_limits(client, fake):
    """SE-5：超过 1 MB 的请求体回 413 body_too_large；嵌套超过 32 层回 422 body_invalid，不入库，之后 GET 照常。"""
    from app.api import table_imports as api_mod

    sid = (await stage(client, unique(), raw_bytes(), "a.xlsx")).json()["id"]
    url = f"/api/datasources/imports/{sid}/recipe"
    big = {"recipe": {"recipe_format": "agentlab-recipe/2", "junk": "x" * (api_mod.BODY_MAX_BYTES + 10)}}
    resp = await client.put(url, json=big)
    assert resp.status_code == 413 and resp.json()["code"] == "body_too_large"
    deep: Any = "x"
    for _ in range(300):
        deep = [deep]
    resp = await client.put(url, json={"recipe": {"recipe_format": "agentlab-recipe/2", "junk": deep}})
    assert resp.status_code == 422 and resp.json()["code"] == "body_invalid"
    resp = await client.get(f"/api/datasources/imports/{sid}")
    assert resp.status_code == 200, resp.text
    # 字符串里的括号不算嵌套
    ok = {"recipe": {"recipe_format": "agentlab-recipe/2", "note": "[" * 100}}
    resp = await client.put(url, json=ok)
    assert resp.status_code == 200, resp.text


async def test_recipe_problems_are_capped(client, monkeypatch):
    """SE-5：几千个问题只存、只回前 200 条，另附一条总数。"""
    from app.data import recipe_imports
    from tests.fixtures.xlsx.lab import kv_form

    raw, fn = kv_form()
    sid = (await stage(client, unique(), raw, fn)).json()["id"]
    rec = json.loads(json.dumps(KV_RECIPE, ensure_ascii=False))
    rec["tables"][0]["units"] = {f"不存在的列{i}": "元" for i in range(1000)}
    resp = await client.put(f"/api/datasources/imports/{sid}/recipe", json={"recipe": rec})
    assert resp.status_code == 200, resp.text
    probs = resp.json()["recipe_problems"]
    assert len(probs) == recipe_imports.RECIPE_PROBLEMS_MAX + 1
    assert probs[-1]["code"] == "too_many_problems" and "共 1000 个" in probs[-1]["message"]
