"""上传表格的版本底座：构建、导入记录、快照，发布、迁移、回收、原件存档。

要守住的几件事：
- 发布失败（解析拒收、文件对不上、写记录出错）时，指针不动、旧版本完整可查、不留新文件；
- 同一个源重传同一个文件是同一个构建、新的导入记录；两个源传同一个文件是两个构建；
- 运行钉住的快照不被回收，钉住的快照没了就明确报错，绝不退回当前版本；
- 原件按哈希只存一份、只读、没有下载口子，清除后记录还在。

夹具一律现造（openpyxl），标签用假名。解析器（tabular.load_into）的新接口由 WP-B 实现：
它还没合进来时，用这里的最小替身（只认简单的一行表头 + 数据）；合进来以后自动换成真的。
"""
from __future__ import annotations

import hashlib
import inspect
import io
import os
import sqlite3
import stat
import uuid
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import pytest
from sqlalchemy import func, select

from app.core.config import settings
from app.data import raw_store, table_versions, tabular
from app.data.engine import engines, run_query
from app.data.table_versions import SnapshotMissing
from app.db.base import SessionLocal
from app.db.models import DataSource, Run, SourceSnapshot, TableBuild, TableImport

# --------------------------------------------------------------------------
# 解析器替身（WP-B 合并前）
# --------------------------------------------------------------------------

REAL_PARSER = "raw_mode" in inspect.signature(tabular.load_into).parameters


@dataclass
class StubTable:
    name: str
    sheet: str
    columns: list[dict[str, Any]]
    rows: int
    region: str
    blank_rows_skipped: int = 0
    columns_trimmed: list[str] = field(default_factory=list)
    unshaped: bool = False


@dataclass
class StubReport:
    tables: list[StubTable] = field(default_factory=list)
    skipped_sheets: list[dict[str, Any]] = field(default_factory=list)
    conversions: list[dict[str, Any]] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)


class StubNeedsDecision(tabular.UnsupportedTable):
    """tabular.NeedsDecision 的形状：UnsupportedTable 的子类，带 kind 和 details。"""

    def __init__(self, message: str, *, kind: str, details: dict[str, Any]) -> None:
        super().__init__(message)
        self.kind = kind
        self.details = details


def stub_load_into(db_path: str, raw: bytes, filename: str, *, header_row: int = 1,
                   mixed: str = "reject", raw_mode: bool = False, scan: Any = None) -> StubReport:
    """一行表头 + 数据的最小解析：整列是整数就 INTEGER，否则 TEXT；STRICT 表。"""
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter

    if not filename.lower().endswith((".xlsx", ".xlsm")):
        raise tabular.UnsupportedTable(f"无法读取 {filename}：仅支持 Excel（.xlsx）")
    try:
        book = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception as e:  # noqa: BLE001
        raise tabular.UnsupportedTable("无法打开该 Excel 文件") from e
    report = StubReport()
    conn = sqlite3.connect(db_path)
    try:
        for ws in book.worksheets:
            rows = [list(r) for r in ws.iter_rows(min_row=header_row, values_only=True)]
            if not rows:
                continue
            headers = [str(h) for h in rows[0]]
            body = rows[1:]
            types = []
            for i in range(len(headers)):
                vals = [r[i] for r in body if r[i] is not None]
                types.append("INTEGER" if vals and all(isinstance(v, int) for v in vals) else "TEXT")
            cols = ", ".join(f'"{h}" {t}' for h, t in zip(headers, types))
            conn.execute(f'CREATE TABLE "{ws.title}" ({cols}) STRICT')
            conn.executemany(
                f'INSERT INTO "{ws.title}" VALUES ({", ".join("?" * len(headers))})',
                [[None if v is None else (v if t == "INTEGER" else str(v)) for v, t in zip(r, types)]
                 for r in body],
            )
            report.tables.append(StubTable(
                name=ws.title, sheet=ws.title,
                columns=[{"name": h, "type": t, "header": h} for h, t in zip(headers, types)],
                rows=len(body),
                region=f"A{header_row}:{get_column_letter(len(headers))}{header_row + len(body)}",
                unshaped=raw_mode,
            ))
        conn.commit()
    finally:
        conn.close()
        book.close()
    return report


@pytest.fixture(autouse=True)
def parser(monkeypatch):
    if not REAL_PARSER:
        monkeypatch.setattr(tabular, "load_into", stub_load_into)


@pytest.fixture(autouse=True)
def store(monkeypatch, tmp_path):
    """每个测试一个独立的数据目录：上传目录、原件存档都在里面，回收和启动清理只看得到这里。"""
    data = tmp_path / "data"
    (data / "uploads").mkdir(parents=True)
    monkeypatch.setattr(settings, "data_dir", data)
    yield data
    # 测试目录随后会被删，权限 0444 的文件挡不住删除（目录可写即可）


# --------------------------------------------------------------------------
# 工具
# --------------------------------------------------------------------------


def xlsx(sheets: dict[str, list[list]]) -> bytes:
    from openpyxl import Workbook

    book = Workbook()
    book.remove(book.active)
    for title, rows in sheets.items():
        sheet = book.create_sheet(title)
        for row in rows:
            sheet.append(row)
    buf = io.BytesIO()
    book.save(buf)
    return buf.getvalue()


V1 = {"甲表": [["区域", "数量"], ["分区甲", 1], ["分区乙", 2]],
      "乙表": [["时段", "客流"], ["早", 10], ["晚", 20]]}
V2 = {"甲表": [["区域", "数量"], ["分区甲", 100]],
      "乙表": [["时段", "客流"], ["早", 1000]]}


def unique(prefix: str = "wpc") -> str:
    return f"{prefix}_{uuid.uuid4().hex[:10]}"


def fresh(sheets: dict) -> dict:
    """内容独一份的工作簿。原件按内容共用：别的测试传过同样字节的文件，原件就不会随回收删除。"""
    return {**sheets, "丙表": [["标记"], [uuid.uuid4().hex]]}


async def publish(name: str, sheets: dict | None = None, *, raw: bytes | None = None,
                  filename: str = "客流.xlsx", **kw: Any):
    """建（或更新）一个上传源并发布。返回 (源 id, ImportResult)。"""
    raw = raw if raw is not None else xlsx(sheets or V1)
    async with SessionLocal() as session:
        row = (await session.execute(select(DataSource).where(DataSource.name == name))).scalar_one_or_none()
        if row is None:
            row = DataSource(id=uuid.uuid4().hex, name=name, kind="sqlite", readonly=True, origin="upload")
        result = await table_versions.publish_upload(session, row, raw, filename, **kw)
        return row.id, result


async def load(source_id: str) -> DataSource:
    async with SessionLocal() as session:
        return await session.get(DataSource, source_id)


async def resolve(source_id: str, pinned: str | None = None):
    async with SessionLocal() as session:
        row = await session.get(DataSource, source_id)
        return await table_versions.resolve_source(session, row, pinned)


async def scalar(source, sql: str):
    try:
        return (await run_query(source, sql)).rows[0][0]
    finally:
        await engines.invalidate(source.id)


async def imports_of(source_id: str) -> list[TableImport]:
    async with SessionLocal() as session:
        return list((await session.execute(
            select(TableImport).where(TableImport.source_id == source_id).order_by(TableImport.seq)
        )).scalars())


async def get(model, key):
    async with SessionLocal() as session:
        return await session.get(model, key)


def files_under(root: Path) -> list[Path]:
    return sorted(p for p in root.rglob("*") if p.is_file()) if root.exists() else []


# --------------------------------------------------------------------------
# id
# --------------------------------------------------------------------------


def test_build_id_covers_source_raw_options_and_engine():
    opts = table_versions.parse_options("a.xlsx", header_row=1, mixed="reject", raw_mode=False)
    base = table_versions.build_id("s1", "a" * 64, opts)
    assert len(base) == 64
    assert table_versions.build_id("s2", "a" * 64, opts) != base
    assert table_versions.build_id("s1", "b" * 64, opts) != base
    other = table_versions.parse_options("a.xlsx", header_row=2, mixed="reject", raw_mode=False)
    assert table_versions.build_id("s1", "a" * 64, other) != base
    # Excel 换个文件名还是同一个构建；CSV 的表名取自文件名，所以文件名算进去
    renamed = table_versions.parse_options("b.xlsx", header_row=1, mixed="reject", raw_mode=False)
    assert table_versions.build_id("s1", "a" * 64, renamed) == base
    csv_a = table_versions.parse_options("a.csv", header_row=1, mixed="reject", raw_mode=False)
    csv_b = table_versions.parse_options("b.csv", header_row=1, mixed="reject", raw_mode=False)
    assert table_versions.build_id("s1", "a" * 64, csv_a) != table_versions.build_id("s1", "a" * 64, csv_b)


def test_build_path_refuses_odd_source_ids():
    with pytest.raises(ValueError):
        table_versions.build_path("../etc", "x")


# --------------------------------------------------------------------------
# 原件存档
# --------------------------------------------------------------------------


def test_put_raw_stores_once_by_hash_and_read_only(store):
    sha, path = raw_store.put_raw(b"hello")
    assert sha == hashlib.sha256(b"hello").hexdigest()
    assert path == store / "uploads" / "raw" / sha[:2] / sha
    assert path.read_bytes() == b"hello"
    assert stat.S_IMODE(path.stat().st_mode) == 0o444
    again, path2 = raw_store.put_raw(b"hello")
    assert (again, path2) == (sha, path)
    assert raw_store.raw_exists(sha)
    # 没有临时文件残留
    assert [p.name for p in path.parent.iterdir()] == [sha]


def test_put_raw_replaces_a_corrupted_copy(store):
    sha, path = raw_store.put_raw(b"payload")
    os.chmod(path, 0o644)
    path.write_bytes(b"")
    raw_store.put_raw(b"payload")
    assert path.read_bytes() == b"payload"


def test_purge_raw_and_bad_hashes(store):
    sha, path = raw_store.put_raw(b"to be purged")
    assert raw_store.purge_raw(sha) is True
    assert not path.exists() and not raw_store.raw_exists(sha)
    assert raw_store.purge_raw(sha) is False
    assert raw_store.raw_exists("../../etc/passwd") is False
    with pytest.raises(ValueError):
        raw_store.raw_path("../../x")


def test_raw_store_has_no_reader():
    """原件不提供下载：模块里没有读回内容的函数，接口层也没有对应的路由。"""
    public = {n for n in dir(raw_store) if not n.startswith("_")}
    assert not {n for n in public if n.startswith(("get", "read", "open", "load"))}
    from app.main import app

    paths = [getattr(r, "path", "") for r in app.routes]
    assert not [p for p in paths if "raw" in p and "purge" not in p]


# --------------------------------------------------------------------------
# 发布
# --------------------------------------------------------------------------


async def test_publish_creates_build_import_snapshot_and_switches_pointer(store):
    name = unique()
    sid, result = await publish(name)
    row = await load(sid)
    assert row.origin == "upload" and row.readonly is True
    assert row.current_snapshot_id == result.snapshot_id

    build = await get(TableBuild, result.build_id)
    snap = await get(SourceSnapshot, result.snapshot_id)
    [imp] = await imports_of(sid)
    assert imp.id == result.import_id and imp.seq == 1 and imp.status == "active"
    assert imp.raw_state == "kept" and raw_store.raw_exists(imp.raw_sha256)
    assert snap.imports == [imp.id] and snap.db_path == build.db_path == row.database
    assert Path(build.db_path).parent == store / "uploads" / "tables" / sid / "builds"
    assert build.db_sha256 == raw_store.sha256_file(Path(build.db_path)) == snap.db_sha256
    assert {t["name"] for t in result.report["tables"]} == {"甲表", "乙表"}
    # 库文件只读：任何代码路径都写不进去
    assert stat.S_IMODE(Path(build.db_path).stat().st_mode) == 0o444
    with pytest.raises(sqlite3.OperationalError):
        conn = sqlite3.connect(build.db_path)
        try:
            conn.execute('INSERT INTO "甲表" VALUES (\'x\', 1)')
            conn.commit()
        finally:
            conn.close()
    # 结构冻结在快照里，和数据源上的一致
    assert set(snap.schema_cache["tables"]) == {"甲表", "乙表"}
    assert row.schema_cache == snap.schema_cache
    # 没有临时文件
    assert not [p for p in files_under(store / "uploads") if raw_store.TMP_MARK in p.name]

    view = await resolve(sid)
    assert view.snapshot_id == snap.id and view.immutable is True and view.readonly is True
    assert view.expected_sha256 == build.db_sha256 and view.database == snap.db_path
    assert await scalar(view, 'SELECT SUM("数量") FROM "甲表"') == 3


async def test_same_file_twice_is_one_build_and_two_imports(store):
    name = unique()
    raw = xlsx(V1)
    sid, first = await publish(name, raw=raw)
    _, second = await publish(name, raw=raw, filename="改了名字.xlsx")
    assert second.build_id == first.build_id and second.build_reused is True
    assert first.build_reused is False
    assert second.import_id != first.import_id and second.snapshot_id != first.snapshot_id
    imps = await imports_of(sid)
    assert [i.seq for i in imps] == [1, 2]
    assert imps[1].status == "active" and imps[0].status != "active"
    assert imps[1].file_name == "改了名字.xlsx"
    # 库文件只有一个，仍在、仍是当前版本
    builds = files_under(store / "uploads" / "tables" / sid)
    assert [p.stem for p in builds] == [first.build_id]
    assert (await load(sid)).current_snapshot_id == second.snapshot_id


async def test_two_sources_same_file_get_separate_builds(store):
    raw = xlsx(V1)
    a, ra = await publish(unique(), raw=raw)
    b, rb = await publish(unique(), raw=raw)
    assert ra.build_id != rb.build_id
    ba, bb = await get(TableBuild, ra.build_id), await get(TableBuild, rb.build_id)
    assert ba.db_path != bb.db_path and ba.source_id == a and bb.source_id == b
    # 删掉 a 并回收：b 的库和共用的原件都还在
    async with SessionLocal() as session:
        await table_versions.retire_source(session, a)
        await session.delete(await session.get(DataSource, a))
        await session.commit()
        await table_versions.gc(session)
    assert not Path(ba.db_path).exists()
    assert Path(bb.db_path).exists()
    assert raw_store.raw_exists(ba.raw_sha256)
    assert {i.status for i in await imports_of(a)} == {"retired"}
    assert await scalar(await resolve(b), 'SELECT COUNT(*) FROM "乙表"') == 2


async def test_concurrent_reuploads_are_serialized(store):
    """同一个源同时来两次重传：序号不重、最后只有一条 active，指针指向其中一次。"""
    import asyncio

    name = unique()
    sid, _ = await publish(name)
    results = await asyncio.gather(publish(name, fresh(V2)), publish(name, fresh(V1)))
    imps = await imports_of(sid)
    assert [i.seq for i in imps] == [1, 2, 3]
    assert [i.status for i in imps].count("active") == 1
    current = (await load(sid)).current_snapshot_id
    assert current in {r.snapshot_id for _, r in results}
    active = next(i for i in imps if i.status == "active")
    assert (await get(SourceSnapshot, current)).imports == [active.id]


async def test_parses_run_at_most_two_at_a_time(store, monkeypatch):
    """解析在线程里做、内存和上传数成正比：同时来五个上传，最多两个在解析，其余排队。"""
    import asyncio
    import threading
    import time

    lock = threading.Lock()
    state = {"now": 0, "peak": 0}

    def slow(final, raw, filename, header_row, mixed, raw_mode):
        with lock:
            state["now"] += 1
            state["peak"] = max(state["peak"], state["now"])
        time.sleep(0.05)
        with lock:
            state["now"] -= 1
        raise tabular.UnsupportedTable("测试用的拒收")

    monkeypatch.setattr(table_versions, "_parse_to_tmp", slow)
    got = await asyncio.gather(*(publish(unique()) for _ in range(5)), return_exceptions=True)
    assert all(isinstance(e, tabular.UnsupportedTable) for e in got)
    assert state["peak"] == table_versions.MAX_PARALLEL_PARSES == 2


async def test_rejected_parse_keeps_old_version_and_leaves_nothing(store, monkeypatch):
    name = unique()
    sid, first = await publish(name)
    before = files_under(store / "uploads")

    def refuse(*_a, **_k):
        raise StubNeedsDecision("需要先确认", kind="shape", details={"reasons": ["示例"]})

    monkeypatch.setattr(tabular, "load_into", refuse)
    with pytest.raises(tabular.UnsupportedTable) as info:
        await publish(name, V2)
    assert getattr(info.value, "kind", None) == "shape"
    assert files_under(store / "uploads") == before          # 没留原件、没留临时库
    assert (await load(sid)).current_snapshot_id == first.snapshot_id
    assert len(await imports_of(sid)) == 1
    view = await resolve(sid)
    assert await scalar(view, 'SELECT SUM("数量") FROM "甲表"') == 3
    assert await scalar(view, 'SELECT SUM("客流") FROM "乙表"') == 30


async def test_unexpected_parse_error_is_wrapped_for_the_user(store, monkeypatch):
    def boom(*_a, **_k):
        raise KeyError("x")

    monkeypatch.setattr(tabular, "load_into", boom)
    with pytest.raises(table_versions.ParseFailed, match="请确认文件是未加密的 Excel 或 CSV"):
        await publish(unique())
    assert not files_under(store / "uploads")


async def test_failure_after_link_rolls_back_and_removes_new_files(store, monkeypatch):
    """写记录这一步失败：指针不动、旧版完整可查，这次新建的库文件和原件都删掉。"""
    name = unique()
    sid, first = await publish(name)
    before = files_under(store / "uploads")

    async def broken(view, report):
        raise RuntimeError("模拟探查失败")

    monkeypatch.setattr(table_versions, "_introspect_view", broken)
    with pytest.raises(RuntimeError, match="模拟探查失败"):
        await publish(name, V2)
    assert files_under(store / "uploads") == before
    row = await load(sid)
    assert row.current_snapshot_id == first.snapshot_id
    assert [i.status for i in await imports_of(sid)] == ["active"]
    view = await resolve(sid)
    assert await scalar(view, 'SELECT SUM("数量") FROM "甲表"') == 3


async def test_new_source_is_not_created_when_publishing_fails(store, monkeypatch):
    async def broken(view, report):
        raise RuntimeError("模拟失败")

    monkeypatch.setattr(table_versions, "_introspect_view", broken)
    name = unique()
    with pytest.raises(RuntimeError):
        await publish(name)
    async with SessionLocal() as session:
        assert (await session.execute(select(DataSource).where(DataSource.name == name))).first() is None
    assert not files_under(store / "uploads" / "raw")


def quarantined(store: Path, source_id: str) -> list[Path]:
    return files_under(store / "uploads" / "quarantine" / source_id)


async def test_existing_target_with_wrong_hash_is_restored_and_quarantined(store):
    """同一个构建的目标文件已存在、哈希和登记的不一样（0 字节残骸、被改过）：这次的临时库哈希等于登记值，
    就把已有文件挪进隔离区、用临时库恢复（期 3 遗留项 1；期 1 这里一律失败，同一个文件永远传不上去）。"""
    name = unique()
    raw = xlsx(V1)
    sid, first = await publish(name, raw=raw)
    build = await get(TableBuild, first.build_id)
    path = Path(build.db_path)
    os.chmod(path, 0o644)
    path.write_bytes(b"")
    _, second = await publish(name, raw=raw)
    assert second.build_reused is True and second.build_restored is True
    assert second.build_id == first.build_id
    assert (await load(sid)).current_snapshot_id == second.snapshot_id
    # 恢复的是登记的那一份：旧快照、新快照指着同一个文件，哈希都对得上
    assert raw_store.sha256_file(path) == build.db_sha256 == (await get(TableBuild, first.build_id)).db_sha256
    assert stat.S_IMODE(path.stat().st_mode) == 0o444
    [moved] = quarantined(store, sid)
    assert moved.name.startswith(f"{first.build_id}-") and moved.stat().st_size == 0
    assert stat.S_IMODE(moved.stat().st_mode) == 0o444
    assert await scalar(await resolve(sid), 'SELECT SUM("数量") FROM "甲表"') == 3


async def test_existing_target_and_record_both_wrong_is_a_build_conflict(store):
    """文件和构建记录的哈希都被改了，本次上传也恢复不了：build_conflict，指针不动，什么都不挪。"""
    name = unique()
    raw = xlsx(V1)
    sid, first = await publish(name, raw=raw)
    async with SessionLocal() as session:
        row = await session.get(TableBuild, first.build_id)
        row.db_sha256 = "f" * 64
        await session.commit()
    path = Path((await get(TableBuild, first.build_id)).db_path)
    os.chmod(path, 0o644)
    path.write_bytes(b"")
    with pytest.raises(table_versions.PublishError, match="无法用本次上传恢复") as err:
        await publish(name, raw=raw)
    assert err.value.code == "build_conflict"
    assert (await load(sid)).current_snapshot_id == first.snapshot_id
    assert len(await imports_of(sid)) == 1
    assert path.exists() and path.stat().st_size == 0      # 不是我们建的，不替它删
    assert not quarantined(store, sid)
    assert raw_store.raw_exists((await get(TableBuild, first.build_id)).raw_sha256)


async def test_unregistered_file_at_target_is_quarantined_and_publish_succeeds(store):
    """目标位置有一个没有记录、内容不同的文件（没人登记的孤儿）：挪进隔离区再 link，发布成功。"""
    sid = uuid.uuid4().hex
    raw = xlsx(V1)
    opts = table_versions.parse_options("客流.xlsx", header_row=1, mixed="reject", raw_mode=False)
    target = table_versions.build_path(sid, table_versions.build_id(sid, hashlib.sha256(raw).hexdigest(), opts))
    target.parent.mkdir(parents=True)
    target.write_bytes(b"")
    name = unique()
    async with SessionLocal() as session:
        row = DataSource(id=sid, name=name, kind="sqlite", readonly=True, origin="upload")
        result = await table_versions.publish_upload(session, row, raw, "客流.xlsx")
    assert result.build_reused is False and result.build_restored is False
    assert (await load(sid)).current_snapshot_id == result.snapshot_id
    assert raw_store.sha256_file(target) == (await get(TableBuild, result.build_id)).db_sha256
    [moved] = quarantined(store, sid)
    assert moved.stat().st_size == 0
    assert await scalar(await resolve(sid), 'SELECT SUM("数量") FROM "甲表"') == 3


async def test_reupload_probes_the_new_file_even_when_the_old_engine_is_cached(store):
    """旧版本的引擎还在缓存里时重传：冻结进新快照的必须是新库的结构，不是缓存引擎指着的旧库。

    引擎缓存的键若只有源 id（WP-A 之前），探查拿到的就是旧库的引擎。探查改用一次性的键，
    发布时再核对探查到的表和解析回执是否一致。
    """
    name = unique()
    sid, _ = await publish(name, {"旧表": [["区域", "数量"], ["分区甲", 1]]})
    # 查一次当前版本，不清缓存：引擎留在缓存里，指着旧库
    assert (await run_query(await resolve(sid), 'SELECT COUNT(*) FROM "旧表"')).rows[0][0] == 1
    try:
        _, second = await publish(name, {"新表": [["时段", "客流"], ["早", 10]]})
        snap = await get(SourceSnapshot, second.snapshot_id)
        assert list(snap.schema_cache["tables"]) == ["新表"]
        assert [t["name"] for t in second.report["tables"]] == ["新表"]
        assert list((await load(sid)).schema_cache["tables"]) == ["新表"]
        assert await scalar(await resolve(sid), 'SELECT SUM("客流") FROM "新表"') == 10
    finally:
        await engines.invalidate(sid)


async def test_publish_refuses_a_probe_that_does_not_match_the_report(store, monkeypatch):
    """探查到的表和解析回执对不上、或者探查本身失败：不冻结进快照，指针不动，新文件删掉。"""
    from app.data import introspect as introspect_mod

    name = unique()
    sid, first = await publish(name)
    before = files_under(store / "uploads")
    real = introspect_mod.introspect

    async def stale(source, **kw):
        cache = await real(source, **kw)
        cache["tables"] = {"别的表": {"qualified": "别的表", "columns": []}}
        return cache

    async def failed(source, **kw):
        return {"tables": {}, "truncated": False, "total": 0, "failed": True, "error": "模拟探查失败"}

    for fake, message in ((stale, "对不上"), (failed, "模拟探查失败")):
        monkeypatch.setattr(introspect_mod, "introspect", fake)
        with pytest.raises(table_versions.PublishError, match=message):
            await publish(name, fresh(V2))
        assert files_under(store / "uploads") == before
        assert (await load(sid)).current_snapshot_id == first.snapshot_id
        assert [i.status for i in await imports_of(sid)] == ["active"]
    monkeypatch.setattr(introspect_mod, "introspect", real)
    assert await scalar(await resolve(sid), 'SELECT SUM("数量") FROM "甲表"') == 3


async def test_a_schema_option_on_an_upload_source_does_not_break_publishing(store):
    """上传库只有 main：数据源 options 里留着的 schema（迁移来的老源、早先手工改过的）不参与探查。"""
    name = unique()
    sid, _ = await publish(name)
    async with SessionLocal() as session:
        row = await session.get(DataSource, sid)
        row.options = {"schema": "nope", "mask_columns": "区域"}
        await session.commit()
    _, second = await publish(name, fresh(V2))
    snap = await get(SourceSnapshot, second.snapshot_id)
    assert {"甲表", "乙表"} <= set(snap.schema_cache["tables"])
    assert not snap.schema_cache.get("failed")
    row = await load(sid)
    # 发布顺手把用不上的 schema 摘掉，其余配置保留
    assert row.options == {"mask_columns": "区域"}
    view = await resolve(sid)
    assert "schema" not in view.options and view.options["mask_columns"] == "区域"
    assert await scalar(view, 'SELECT SUM("数量") FROM "甲表"') == 100


def _wal_parser(keep_open: list):
    """替身解析器：把临时库切成 WAL 再写。keep_open 非 None 时留一个连接开着，-wal 合并不掉。"""

    def parse(db_path, raw, filename, **kw):
        conn = sqlite3.connect(db_path)
        conn.execute("PRAGMA journal_mode=WAL")
        conn.close()
        report = stub_load_into(db_path, raw, filename, **kw)
        if keep_open is not None:
            holder = sqlite3.connect(db_path, check_same_thread=False)
            holder.execute("SELECT COUNT(*) FROM sqlite_master").fetchall()
            keep_open.append(holder)
        return report

    return parse


async def test_published_files_are_in_rollback_journal_mode(store, monkeypatch):
    """解析器用了 WAL 也要切回回滚日志：库以 immutable 打开时 SQLite 不看日志，留在 WAL 里会少数据。"""
    monkeypatch.setattr(tabular, "load_into", _wal_parser(None))
    sid, result = await publish(unique(), fresh(V1))
    path = Path((await get(TableBuild, result.build_id)).db_path)
    header = path.read_bytes()[:100]
    assert header[18] == 1 and header[19] == 1       # 1 = 回滚日志，2 = WAL
    assert not Path(f"{path}-wal").exists() and not Path(f"{path}-shm").exists()
    assert await scalar(await resolve(sid), 'SELECT SUM("数量") FROM "甲表"') == 3


async def test_publish_stops_when_the_wal_cannot_be_merged(store, monkeypatch):
    name = unique()
    sid, first = await publish(name)
    before = files_under(store / "uploads")
    holders: list = []
    monkeypatch.setattr(tabular, "load_into", _wal_parser(holders))
    try:
        with pytest.raises(table_versions.PublishError, match="日志"):
            await publish(name, fresh(V2))
    finally:
        for conn in holders:
            conn.close()
    assert files_under(store / "uploads") == before
    assert (await load(sid)).current_snapshot_id == first.snapshot_id


async def test_unshaped_tables_get_a_note_without_numbers(store, monkeypatch):
    def unshaped(db_path, raw, filename, **kw):
        report = stub_load_into(db_path, raw, filename, **kw)
        report.tables[0].unshaped = True
        report.tables[1].unshaped = False
        return report

    monkeypatch.setattr(tabular, "load_into", unshaped)
    sid, result = await publish(unique(), raw_mode=True)
    snap = await get(SourceSnapshot, result.snapshot_id)
    tables = snap.schema_cache["tables"]
    assert tables["甲表"]["comment"] == table_versions.UNSHAPED_NOTE
    assert not tables["乙表"].get("comment")
    assert not any(ch.isdigit() for ch in table_versions.UNSHAPED_NOTE)
    from app.data import introspect

    assert "未经规整" in introspect.describe_table(await resolve(sid), "甲表")


# --------------------------------------------------------------------------
# 快照解析
# --------------------------------------------------------------------------


async def test_manual_sources_resolve_to_themselves(tmp_path):
    async with SessionLocal() as session:
        row = DataSource(name=unique(), kind="sqlite", database=str(tmp_path / "m.db"))
        session.add(row)
        await session.commit()
        assert await table_versions.resolve_source(session, row) is row


async def test_pinned_snapshot_must_belong_to_the_source(store):
    a, ra = await publish(unique())
    b, _ = await publish(unique())
    with pytest.raises(SnapshotMissing, match="不会改用当前版本"):
        await resolve(b, ra.snapshot_id)
    with pytest.raises(SnapshotMissing):
        await resolve(b, "f" * 64)


async def test_bind_snapshot_uses_the_snapshot_not_the_source(store):
    sid, first = await publish(unique())
    async with SessionLocal() as session:
        row = await session.get(DataSource, sid)
        snap = await session.get(SourceSnapshot, first.snapshot_id)
        row.schema_cache = {"tables": {}}       # 数据源上的缓存被改了，视图不受影响
        row.database = "/elsewhere.db"
        view = table_versions.bind_snapshot(row, snap)
        name, snap_path = row.name, snap.db_path
        await session.rollback()
    assert set(view.schema_cache["tables"]) == {"甲表", "乙表"}
    assert view.database == snap_path and view.readonly and view.immutable
    assert view.id == sid and view.name == name and view.kind == "sqlite"
    # 视图里的结构是一份拷贝：改它不会改到快照
    view.schema_cache["tables"].clear()
    assert set((await get(SourceSnapshot, first.snapshot_id)).schema_cache["tables"]) == {"甲表", "乙表"}


def _rewrite_snapshot(path: str, sql: str) -> None:
    """改快照库里的数据（模拟被改过的文件），改完切回回滚日志模式，不留 -wal / -journal。"""
    os.chmod(path, 0o644)
    conn = sqlite3.connect(path)
    conn.execute(sql)
    conn.commit()
    conn.execute("PRAGMA journal_mode=DELETE")
    conn.close()


async def test_snapshot_with_an_empty_hash_is_refused(store):
    """快照登记的哈希是空的（记录损坏、被手工改过）：不能退回「不核对照样打开」，读到被改过的数据。"""
    sid, first = await publish(unique())
    snap = await get(SourceSnapshot, first.snapshot_id)
    _rewrite_snapshot(snap.db_path, 'UPDATE "甲表" SET "数量" = 1000 WHERE "区域" = \'分区甲\'')
    with pytest.raises(Exception, match="可能被修改过"):      # 对照：哈希还在时，被改过的文件被拒
        await scalar(await resolve(sid), 'SELECT SUM("数量") FROM "甲表"')
    for blank in ("", "   "):
        async with SessionLocal() as session:
            (await session.get(SourceSnapshot, first.snapshot_id)).db_sha256 = blank
            await session.commit()
        with pytest.raises(SnapshotMissing, match="缺少数据文件的哈希登记"):
            await resolve(sid)
        with pytest.raises(SnapshotMissing, match="缺少数据文件的哈希登记"):
            await resolve(sid, first.snapshot_id)


async def test_reusing_a_build_with_an_empty_hash_fills_it_in(store):
    """构建记录的哈希是空的，重传同一个文件复用了它：新快照不能继承空哈希，构建记录也补上。"""
    name, raw = unique(), xlsx(fresh(V1))
    sid, first = await publish(name, raw=raw)
    async with SessionLocal() as session:
        (await session.get(TableBuild, first.build_id)).db_sha256 = ""
        await session.commit()
    _, second = await publish(name, raw=raw)
    assert second.build_reused and second.build_id == first.build_id
    build = await get(TableBuild, first.build_id)
    snap = await get(SourceSnapshot, second.snapshot_id)
    assert build.db_sha256 and snap.db_sha256 == build.db_sha256 == raw_store.sha256_file(Path(build.db_path))
    assert await scalar(await resolve(sid), 'SELECT SUM("数量") FROM "甲表"') == 3


# --------------------------------------------------------------------------
# 回收
# --------------------------------------------------------------------------


async def test_old_versions_are_reclaimed_when_nothing_references_them(store, monkeypatch):
    # 期 3 起回收在当前快照、运行引用的之外另外保留每个源最近 SNAPSHOT_KEEP 个快照；调成 0（期 1 的规则）
    # 才测得出「没人引用就回收」
    monkeypatch.setattr(table_versions, "SNAPSHOT_KEEP", 0)
    name = unique()
    sid, first = await publish(name, fresh(V1))
    old = await get(TableBuild, first.build_id)
    _, second = await publish(name, fresh(V2))
    assert not Path(old.db_path).exists()
    assert (await get(TableBuild, first.build_id)).retired_at is not None
    assert (await get(SourceSnapshot, first.snapshot_id)).retired_at is not None
    imps = await imports_of(sid)
    assert imps[0].status == "retired" and imps[0].raw_state == "purged"
    assert imps[0].purged and imps[0].purged.get("auto") is True
    assert not raw_store.raw_exists(imps[0].raw_sha256)
    assert raw_store.raw_exists(imps[1].raw_sha256)
    with pytest.raises(SnapshotMissing, match="不会改用当前版本"):
        await resolve(sid, first.snapshot_id)


async def test_gc_keeps_snapshots_referenced_by_interrupted_runs(store, monkeypatch):
    # 同上：SNAPSHOT_KEEP 调成 0，保护只来自运行引用，运行记录删掉以后才会被回收
    monkeypatch.setattr(table_versions, "SNAPSHOT_KEEP", 0)
    name = unique()
    sid, first = await publish(name, fresh(V1))
    async with SessionLocal() as session:
        run = Run(workflow_name="wpc-gc", status="interrupted",
                  data_versions={sid: {"snapshot": first.snapshot_id, "name": name}})
        session.add(run)
        await session.commit()
        run_id = run.id
    try:
        _, second = await publish(name, fresh(V2))
        old = await get(TableBuild, first.build_id)
        assert Path(old.db_path).exists() and old.retired_at is None
        assert (await get(SourceSnapshot, first.snapshot_id)).retired_at is None
        imps = await imports_of(sid)
        assert imps[0].status == "superseded" and raw_store.raw_exists(imps[0].raw_sha256)
        # 钉住的旧版仍可查，当前版本是新的
        assert await scalar(await resolve(sid, first.snapshot_id), 'SELECT SUM("数量") FROM "甲表"') == 3
        assert await scalar(await resolve(sid), 'SELECT SUM("数量") FROM "甲表"') == 100
    finally:
        async with SessionLocal() as session:
            await session.delete(await session.get(Run, run_id))
            await session.commit()
    # 运行没了，再回收就删掉
    async with SessionLocal() as session:
        await table_versions.gc(session)
    assert not Path((await get(TableBuild, first.build_id)).db_path).exists()
    assert not raw_store.raw_exists((await imports_of(sid))[0].raw_sha256)
    with pytest.raises(SnapshotMissing):
        await resolve(sid, first.snapshot_id)


async def test_gc_never_deletes_outside_the_uploads_dir(store, tmp_path):
    outside = tmp_path / "outside.db"
    outside.write_bytes(b"keep me")
    async with SessionLocal() as session:
        session.add(TableBuild(id=hashlib.sha256(os.urandom(8)).hexdigest(), source_id="nosuchsource",
                               db_path=str(outside)))
        await session.commit()
        await table_versions.gc(session)
    assert outside.read_bytes() == b"keep me"


# --------------------------------------------------------------------------
# 启动：孤儿清理、迁移
# --------------------------------------------------------------------------


async def test_startup_removes_orphans_left_by_a_crash(store):
    """模拟「link 之后、提交之前」进程没了：库文件和原件都在，却没有记录。启动时清掉。"""
    sid, kept = await publish(unique())
    kept_build = await get(TableBuild, kept.build_id)

    orphan_src = uuid.uuid4().hex
    orphan = table_versions.build_path(orphan_src, "0" * 64)
    orphan.parent.mkdir(parents=True)
    tmp = orphan.parent / f"{orphan.name}{raw_store.TMP_MARK}abc"
    conn = sqlite3.connect(tmp)
    conn.execute("CREATE TABLE t (a INTEGER)")
    conn.commit()
    conn.close()
    table_versions._publish_file(tmp, orphan, None)
    assert orphan.exists() and not tmp.exists()
    half = orphan.parent / f"{'1' * 64}.db{raw_store.TMP_MARK}zzz"
    half.write_bytes(b"half written")
    orphan_sha, orphan_raw = raw_store.put_raw(b"raw nobody registered")

    async with SessionLocal() as session:
        await table_versions.startup(session)
    assert not orphan.exists() and not half.exists() and not orphan_raw.exists()
    assert Path(kept_build.db_path).exists() and raw_store.raw_exists(kept_build.raw_sha256)
    assert await scalar(await resolve(sid), 'SELECT COUNT(*) FROM "甲表"') == 2


def _legacy_db(path: Path) -> None:
    conn = sqlite3.connect(path)
    conn.execute('CREATE TABLE "明细" ("金额" INTEGER)')
    conn.executemany('INSERT INTO "明细" VALUES (?)', [(5,), (7,)])
    conn.commit()
    conn.close()


async def test_startup_migrates_legacy_uploads_idempotently(store):
    tables = table_versions.tables_root()
    tables.mkdir(parents=True)
    a_path, b_path = tables / "old_a.db", tables / "old_b.db"
    _legacy_db(a_path)
    _legacy_db(b_path)
    assert a_path.read_bytes() == b_path.read_bytes()       # 内容相同，v0 的 id 也不能撞
    async with SessionLocal() as session:
        a = DataSource(name=unique("old"), kind="sqlite", database=str(a_path), readonly=True)
        b = DataSource(name=unique("old"), kind="sqlite", database=str(b_path), readonly=True,
                       schema_cache={"tables": {"明细": {"qualified": "明细", "columns": []}}})
        session.add_all([a, b])
        await session.commit()
        ids = (a.id, b.id)

    async with SessionLocal() as session:
        await table_versions.startup(session)
    rows = [await load(i) for i in ids]
    assert [r.origin for r in rows] == ["upload", "upload"]
    snaps = [await get(SourceSnapshot, r.current_snapshot_id) for r in rows]
    builds = [await get(TableBuild, (await imports_of(r.id))[0].build_id) for r in rows]
    assert snaps[0].id != snaps[1].id and builds[0].id != builds[1].id
    sha = raw_store.sha256_file(a_path)
    assert builds[0].id == hashlib.sha256(f"legacy{ids[0]}{sha}".encode()).hexdigest()
    assert builds[0].engine_ver == table_versions.LEGACY_ENGINE_VER
    for r, snap, path in zip(rows, snaps, (a_path, b_path)):
        [imp] = await imports_of(r.id)
        assert imp.raw_state == "absent" and imp.status == "active"
        assert snap.db_path == os.path.realpath(path) and snap.db_sha256 == sha
        assert stat.S_IMODE(path.stat().st_mode) == 0o444
    # 没有结构缓存的补探一次；有的原样冻结
    assert "明细" in snaps[0].schema_cache["tables"]
    assert snaps[1].schema_cache == {"tables": {"明细": {"qualified": "明细", "columns": []}}}
    assert await scalar(await resolve(ids[0]), 'SELECT SUM("金额") FROM "明细"') == 12

    async with SessionLocal() as session:
        await table_versions.startup(session)
    for r in rows:
        assert len(await imports_of(r.id)) == 1
        assert (await load(r.id)).current_snapshot_id == r.current_snapshot_id
    async with SessionLocal() as session:
        count = (await session.execute(
            select(func.count()).select_from(TableBuild).where(TableBuild.source_id.in_(ids))
        )).scalar()
    assert count == 2


async def test_one_failing_legacy_source_does_not_block_the_rest_of_startup(store, monkeypatch):
    """一个老上传迁不过去（文件属主不对、chmod 报 EPERM 之类）：别的源照常迁，孤儿清理和回收照常做。"""
    tables = table_versions.tables_root()
    tables.mkdir(parents=True)
    paths = [tables / "aaa_fail.db", tables / "bbb_ok.db"]
    for path in paths:
        _legacy_db(path)
    async with SessionLocal() as session:
        rows = [DataSource(name=unique("old"), kind="sqlite", database=str(p), readonly=True) for p in paths]
        session.add_all(rows)
        await session.commit()
        ids = [r.id for r in rows]
    # 一个没人登记的孤儿库、一个没人引用的旧构建：startup 走完时都该被收拾掉
    orphan = table_versions.build_path(uuid.uuid4().hex, "0" * 64)
    orphan.parent.mkdir(parents=True)
    orphan.write_bytes(b"orphan")
    stale_sid, stale = await publish(unique(), fresh(V1))
    stale_build = await get(TableBuild, stale.build_id)
    async with SessionLocal() as session:
        # 删了源、还没来得及回收（比如回收那一步失败了）
        await table_versions.retire_source(session, stale_sid)
        await session.delete(await session.get(DataSource, stale_sid))
        await session.commit()

    real = table_versions._prepare_legacy
    calls: list[str] = []

    def flaky(path):
        calls.append(path)
        if len(calls) == 1:
            raise PermissionError(1, "Operation not permitted", path)
        return real(path)

    monkeypatch.setattr(table_versions, "_prepare_legacy", flaky)
    async with SessionLocal() as session:
        await table_versions.startup(session)       # 不抛

    assert len(calls) == 2
    failed_path = calls[0]
    by_path = {str(os.path.realpath(p)): i for p, i in zip(paths, ids)}
    failed_id = by_path[failed_path]
    ok_id = next(i for i in ids if i != failed_id)
    assert (await load(failed_id)).origin == "manual"
    assert (await load(failed_id)).current_snapshot_id is None
    migrated = await load(ok_id)
    assert migrated.origin == "upload" and migrated.current_snapshot_id
    assert await scalar(await resolve(ok_id), 'SELECT SUM("金额") FROM "明细"') == 12
    assert not orphan.exists()
    assert (await get(TableBuild, stale.build_id)).retired_at is not None
    assert not Path(stale_build.db_path).exists()

    # 下次启动把上次失败的那个补上
    monkeypatch.setattr(table_versions, "_prepare_legacy", real)
    async with SessionLocal() as session:
        await table_versions.startup(session)
    assert (await load(failed_id)).origin == "upload"


async def test_startup_leaves_other_sqlite_sources_alone(store, tmp_path):
    """不是老上传的 SQLite 源不迁移；但指向上传目录的（迁移认不出来的更深一层、带连接参数的别名）
    强制改成只读：那里的文件归版本底座管，升级前登记的可写手工源不能接着往里写（方案 3.7）。"""
    elsewhere = tmp_path / "manual.db"
    _legacy_db(elsewhere)
    nested = table_versions.tables_root() / "deeper"
    nested.mkdir(parents=True)
    _legacy_db(nested / "x.db")
    _legacy_db(nested / "y.db")
    async with SessionLocal() as session:
        m = DataSource(name=unique("man"), kind="sqlite", database=str(elsewhere), readonly=False)
        n = DataSource(name=unique("man"), kind="sqlite", database=str(nested / "x.db"), readonly=False)
        q = DataSource(name=unique("man"), kind="sqlite", database=f"{nested / 'y.db'}?timeout=5",
                       readonly=False)
        session.add_all([m, n, q])
        await session.commit()
        ids = (m.id, n.id, q.id)
        await table_versions.startup(session)
    for i in ids:
        row = await load(i)
        assert row.origin == "manual" and row.current_snapshot_id is None
    assert stat.S_IMODE(elsewhere.stat().st_mode) != 0o444
    assert (await load(m.id)).readonly is False             # 上传目录外的手工库照旧可写
    assert (await load(n.id)).readonly is True and (await load(q.id)).readonly is True
    async with SessionLocal() as session:
        with pytest.raises(Exception, match="readonly|只读|拒绝"):
            await run_query(await session.get(DataSource, n.id), 'INSERT INTO "明细" VALUES (9)')
    assert sqlite3.connect(nested / "x.db").execute('SELECT COUNT(*) FROM "明细"').fetchone()[0] == 2


async def test_legacy_upload_can_be_updated_by_a_new_upload(store, monkeypatch):
    # 同上：SNAPSHOT_KEEP 调成 0，迁移出的 v0 版本没人引用，随回收删除
    monkeypatch.setattr(table_versions, "SNAPSHOT_KEEP", 0)
    tables = table_versions.tables_root()
    tables.mkdir(parents=True)
    legacy = tables / "old_c.db"
    _legacy_db(legacy)
    name = unique("old")
    async with SessionLocal() as session:
        row = DataSource(name=name, kind="sqlite", database=str(legacy), readonly=True)
        session.add(row)
        await session.commit()
        await table_versions.startup(session)
    sid, result = await publish(name, V1)
    assert (await load(sid)).current_snapshot_id == result.snapshot_id
    assert not legacy.exists()          # 老库没有运行引用，随回收删除
    assert [i.status for i in await imports_of(sid)] == ["retired", "active"]
