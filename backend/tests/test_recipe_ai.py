"""WP-4 AI 兜底起草（recipe_ai）的单元测试（P2-SPEC 9.6 的 WP-4 部分）。

不调用任何真实模型：ai_draft 注入假模型（async ainvoke(messages) -> AIMessage）；可用性的用例用独立的内存库建接入，
resolve_draft_model 走到 build_chat_model 时 monkeypatch 成假模型。夹具一律合成、假名。
"""
from __future__ import annotations

import asyncio
import datetime as dt
import json
from typing import Any

import pytest
from langchain_core.messages import AIMessage, HumanMessage, SystemMessage
from openpyxl import Workbook

from app.data import recipe_ai as A
from app.data import recipe_suggest as S
from app.data.recipe_types import Draft, DraftFacts, Extraction, Fact, Grid, GridCell, Problem, Recipe, RecipeProblem
from tests.test_recipe_suggest import FLOW, build, flow_book, ok_validate

# ==========================================================================
# 假模型与桩
# ==========================================================================


def reply(content: Any, *, inp: int = 100, out: int = 50, meta: dict[str, Any] | None = None) -> AIMessage:
    return AIMessage(content=content, usage_metadata={"input_tokens": inp, "output_tokens": out,
                                                      "total_tokens": inp + out},
                     response_metadata=meta or {})


class FakeModel:
    """按预设顺序返回；元素是异常就抛，是协程函数就 await 它。记下每次收到的消息列表。"""

    def __init__(self, replies: list[Any], name: str = "fake-model") -> None:
        self.replies = list(replies)
        self.calls: list[list[Any]] = []
        self.model_name = name

    async def ainvoke(self, messages: list[Any]) -> Any:
        self.calls.append(list(messages))
        r = self.replies.pop(0)
        if isinstance(r, BaseException):
            raise r
        if callable(r):
            return await r()
        return r


def flow_json(**patch: Any) -> str:
    data = json.loads(json.dumps(FLOW))
    data.update(patch)
    return json.dumps(data, ensure_ascii=False)


def claims_validate(calls: list | None = None):
    """静态校验桩：pydantic、AI 的 note 必须为空（ai_note_forbidden）、每条事实恰好被认领一次（fact_unclaimed）。"""
    def _v(data: dict[str, Any], facts: DraftFacts | None, origin: str):
        if calls is not None:
            calls.append((data, facts, origin))
        try:
            r = Recipe.model_validate(data)
        except Exception as e:  # noqa: BLE001
            return None, [RecipeProblem("/", "schema", str(e)[:200])]
        probs = []
        if origin == "ai":
            for i, t in enumerate(r.tables):
                if t.note:
                    probs.append(RecipeProblem(f"/tables/{i}/note", "ai_note_forbidden", f"表「{t.name}」：AI 起草的配方说明必须为空"))
        claimed = [x.claims for x in r.relations if x.claims]
        for f in (facts.facts if facts else []):
            if claimed.count(f.id) != 1:
                probs.append(RecipeProblem("/relations", "fact_unclaimed", f"系统发现的关系「{f.text}」没有被认领"))
        return (r if not probs else r), probs
    return _v


FACTS = DraftFacts(facts=[
    Fact("F1", "sum_eq", "客流汇总", "第 5 行 = 第 6 行 + 第 7 行（31 列中 31 列成立）",
         {"segment": "日客流", "total": "全日客流（人次）", "parts": ["分区甲（人次）", "分区乙（人次）"], "rows": [5, 6, 7]}),
    Fact("F2", "not_equal_sum", "客流汇总", "日间、夜间各行之和 与 第 5 行：31 列中 0 列相等",
         {"a": ["日间", "夜间"], "b_segment": "日客流", "b": "全日客流（人次）", "equal": 0, "checked": 31}),
], candidates={"日间时段客流（人次）": ["日间", "日间时段客流"], "夜间时段客流（人次）": ["夜间", "夜间时段客流"]})


@pytest.fixture(scope="module")
def flow():
    wb, cache, nums = flow_book()
    sc, grids = build(wb, cache)
    return sc, grids, nums


async def run(flow, model, *, validate=None, dry_run=None, rules: Draft | None = None, facts=FACTS, **kw):
    sc, grids, _ = flow
    return await A.ai_draft(None, grids=grids, scan=sc, facts=facts, rules_draft=rules, filename="x.xlsx",  # type: ignore[arg-type]
                            validate=validate or claims_validate(), dry_run=dry_run, model=model, **kw)


# ==========================================================================
# 常量
# ==========================================================================


def test_constants_match_copilot():
    from app.api import copilot

    assert A.COPILOT_SETTING_KEY == copilot.COPILOT_SETTING_KEY
    assert A.AI_MAX_TOKENS == copilot.COPILOT_MAX_TOKENS
    assert A.AI_FALLBACK_MAX_TOKENS == copilot._FALLBACK_MAX_TOKENS
    assert A._BUDGET_WORDS == copilot._BUDGET_WORDS
    assert A.AI_MAX_ATTEMPTS == 3


# ==========================================================================
# 可用性（独立的内存库）
# ==========================================================================


class _MemDB:
    """独立的内存库：不碰测试会话共用的库，也不受别的测试建的接入影响。"""

    async def __aenter__(self):
        from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker, create_async_engine
        from sqlalchemy.pool import StaticPool

        import app.db.models  # noqa: F401 - 建表前要先登记全部模型
        from app.db.base import Base

        self.engine = create_async_engine("sqlite+aiosqlite:///:memory:", poolclass=StaticPool)
        async with self.engine.begin() as conn:
            await conn.run_sync(Base.metadata.create_all)
        self.session = async_sessionmaker(self.engine, expire_on_commit=False, class_=AsyncSession)()
        return self.session

    async def __aexit__(self, *exc):
        await self.session.close()
        await self.engine.dispose()


@pytest.fixture
async def mem_session():
    async with _MemDB() as session:
        yield session


async def _setup(session, providers: list[dict[str, Any]], saved: dict[str, str] | None = None) -> None:
    from app.core.crypto import encrypt
    from app.db.models import Provider, Setting

    for p in providers:
        p = dict(p)
        if p.get("api_key"):
            p["api_key"] = encrypt(p["api_key"])
        session.add(Provider(**p))
    if saved is not None:
        session.add(Setting(key=A.COPILOT_SETTING_KEY, value=saved))
    await session.commit()


@pytest.fixture
def no_env_keys(monkeypatch):
    for k in ("ANTHROPIC_API_KEY", "ANTHROPIC_AUTH_TOKEN", "OPENAI_API_KEY", "ANTHROPIC_BASE_URL"):
        monkeypatch.delenv(k, raising=False)


CASES = {
    "none": ([], None),
    "mock_only": ([{"name": "演示", "kind": "mock", "models": [{"id": "mock-fast"}], "default_model": "mock-fast"}], None),
    "name_disabled": ([{"name": "甲接入", "kind": "anthropic", "api_key": "k", "models": [{"id": "claude-x"}],
                        "default_model": "claude-x", "enabled": False}], {"provider": "甲接入", "model": ""}),
    "model_disabled": ([{"name": "乙接入", "kind": "anthropic", "api_key": "k", "models": [{"id": "claude-y"}],
                         "default_model": "claude-y", "enabled": False}], {"provider": "", "model": "claude-y"}),
    "missing_key": ([{"name": "丙接入", "kind": "anthropic", "api_key": None, "models": [{"id": "claude-z"}],
                      "default_model": "claude-z"}], None),
}


@pytest.mark.asyncio
async def test_unavailable_cases(no_env_keys, monkeypatch, flow):
    """没有接入、只有演示接入、存的接入名已停用、存的模型名所属接入已停用、缺 key：都不可用，原因各不相同，
    假模型一次都没被调。"""
    from app.providers import factory

    built: list[FakeModel] = []
    real = factory.build_chat_model

    def spy(provider, spec):
        real(provider, spec)                  # 缺 key 时这里就抛 ProviderNotConfigured
        fake = FakeModel([])
        built.append(fake)
        return fake
    monkeypatch.setattr(factory, "build_chat_model", spy)
    sc, grids, _ = flow
    reasons: dict[str, str] = {}
    for case, (providers, saved) in CASES.items():
        async with _MemDB() as session:
            await _setup(session, providers, saved)
            avail = await A.ai_availability(session)
            assert not avail.available and avail.reason, case
            with pytest.raises(A.AiUnavailable) as e:
                await A.ai_draft(session, grids=grids, scan=sc, facts=FACTS, rules_draft=None, filename="x.xlsx",
                                 validate=claims_validate(), dry_run=None)
            assert str(e.value) == avail.reason
            reasons[case] = avail.reason
    assert all(not m.calls for m in built)
    assert len(set(reasons.values())) == len(CASES), reasons
    assert "已停用" in reasons["name_disabled"] and "甲接入" in reasons["name_disabled"]
    assert "已停用" in reasons["model_disabled"] and "乙接入" in reasons["model_disabled"]
    assert "API Key" in reasons["missing_key"]
    assert "演示" in reasons["mock_only"]
    assert "未配置模型接入" in reasons["none"]


@pytest.mark.asyncio
async def test_available_builds_model_without_request(mem_session, no_env_keys, monkeypatch):
    from app.providers import factory

    await _setup(mem_session, [{"name": "丁接入", "kind": "anthropic", "api_key": "sk-test",
                                "models": [{"id": "claude-w"}], "default_model": "claude-w"}],
                 {"provider": "丁接入", "model": "claude-w"})
    specs: list[Any] = []

    def fake_build(provider, spec):
        specs.append(spec)
        return FakeModel([], name=spec.model)
    monkeypatch.setattr(factory, "build_chat_model", fake_build)
    model, model_id, avail = await A.resolve_draft_model(mem_session)
    assert avail.available and avail.model == "claude-w" and avail.provider == "丁接入" and model_id == "claude-w"
    assert [s.max_tokens for s in specs] == [A.AI_MAX_TOKENS, A.AI_FALLBACK_MAX_TOKENS]
    assert all(s.thinking == "off" and s.temperature == 0 for s in specs)
    assert not model.model.calls                          # 只建对象，不发请求


@pytest.mark.asyncio
async def test_budget_rejection_falls_back_to_small_budget(mem_session, no_env_keys, monkeypatch, flow):
    """服务方以 400「max_tokens too large」拒绝 → 退回 8192 重来，之后的调用都用 8192。"""
    from app.providers import factory

    await _setup(mem_session, [{"name": "戊接入", "kind": "anthropic", "api_key": "sk-test",
                                "models": [{"id": "claude-v"}], "default_model": "claude-v"}])

    class BudgetError(Exception):
        status_code = 400

    models: dict[int, FakeModel] = {}

    def fake_build(provider, spec):
        if spec.max_tokens == A.AI_MAX_TOKENS:
            m = FakeModel([BudgetError("max_tokens: 32768 > 8192, which is the maximum allowed")])
        else:
            m = FakeModel([reply("不是 JSON"), reply(flow_json())])
        models[spec.max_tokens] = m
        return m
    monkeypatch.setattr(factory, "build_chat_model", fake_build)
    sc, grids, _ = flow
    res = await A.ai_draft(mem_session, grids=grids, scan=sc, facts=FACTS, rules_draft=None, filename="x.xlsx",
                           validate=claims_validate(), dry_run=None)
    assert res.draft is not None and res.attempts == 2
    assert len(models[A.AI_MAX_TOKENS].calls) == 1
    assert len(models[A.AI_FALLBACK_MAX_TOKENS].calls) == 2


# ==========================================================================
# 压缩表示不外泄
# ==========================================================================


def _leaky_book():
    """一张交叉表 + 一张列表：特征数、像数的文字、带常量的公式、姓名电话、隐藏行列、遮罩列。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "客流汇总"
    ws["B2"] = "统计时间范围：2026年8月1日至2026年8月8日"
    for i in range(8):
        ws.cell(4, 3 + i, f"8月{i + 1}日")
    ws["B5"] = "全日客流（人次）"
    ws["B6"] = "分区甲（人次）"
    ws["B7"] = "分区乙（人次）"
    ws["B8"] = "秘密指标（人次）"
    for i in range(8):
        ws.cell(5, 3 + i, 987654 + i)
        ws.cell(6, 3 + i, 123457 + i)
        ws.cell(7, 3 + i, 5555.75)
        ws.cell(8, 3 + i, 424242)
    ws["C9"] = "1,234"
    ws["B9"] = "补录（人次）"
    ws["D9"] = "=3000+2500"
    ws["E9"] = "=C5*1.0675"
    ws["F9"] = '="内部"&B5'
    ws["G9"] = "='附表'!C5+1"
    ws.row_dimensions[8].hidden = True
    ws["L5"] = "隐藏列里的话"
    ws.column_dimensions["L"].hidden = True
    lst = wb.create_sheet("名单")
    lst.append(["姓名", "电话", "金额", "工号"])
    people = [("张三", "138-1234-5678", 10, "X-保密-1"), ("李四", "139-8765-4321", 20, "X-保密-2"),
              ("王五", "137-0000-1111", 30, "X-保密-3"), ("赵六", "136-2222-3333", 40, "X-保密-4"),
              ("钱七", "135-4444-5555", 50, "X-保密-5")]
    for p in people:
        lst.append(list(p))
    cache = {"客流汇总": {"D9": 5500, "E9": 1054320.6, "F9": "内部987654", "G9": 99}}
    return wb, cache


def test_compress_does_not_leak_values():
    wb, cache = _leaky_book()
    sc, grids = build(wb, cache)
    facts = S.detect_facts(grids, sc)
    c = A.compress(grids, sc, facts, mask_headers={"工号"})
    text = c.text
    for g in grids.values():
        for cell in g.cells.values():
            # 特征数逐个断言（小数字会和坐标、统计期里的年份撞上，不在这里查）
            if isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool) and abs(cell.value) >= 5000:
                assert str(cell.value) not in text, cell.value
    for secret in ("987654", "123457", "5555.75", "424242", "1,234", "3000", "2500", "1.0675", "内部", "1054320",
                   "张三", "李四", "138-1234-5678", "139-8765-4321", "秘密指标", "隐藏列里的话", "保密"):
        assert secret not in text, secret
    assert "<文本数字>" in text
    assert "=<数>+<数>" in text and "=C5*<数>" in text and "=<文本>&B5" in text
    assert "G9 ='附表'!<格>+<数>" in text           # 跨工作表引用只留表名
    assert "隐藏行：第 8 行（文字 1 格、数字 8 格）" in text
    assert "隐藏列：L 列（文字 1 格、数字 0 格）" in text
    assert "A 列「姓名」是数据列（只给画像）：5 格，不同取值 5 个" in text
    assert "B 列「电话」是数据列" in text
    assert "D 列「工号」是遮罩列（只给画像）" in text
    assert "C9「<文本数字>」" in text
    # 结构文字照发
    for shown in ("全日客流（人次）", "统计时间范围：2026年8月1日至2026年8月8日", "序列「{}月{}日」8 格", "姓名", "金额"):
        assert shown in text, shown
    # 同一份网格两次压缩，sha256 相同
    again = A.compress(grids, sc, facts, mask_headers={"工号"})
    assert again.sha256 == c.sha256 and again.text == text and c.chars == len(text)


def test_sanitize_formula():
    assert A.sanitize_formula("=3000+2500") == "=<数>+<数>"
    assert A.sanitize_formula("=C5*1.0675") == "=C5*<数>"
    assert A.sanitize_formula('=IF(B5>50000,"高","低")') == "=IF(B5><数>,<文本>,<文本>)"
    assert A.sanitize_formula("='内部'!C5+SUM(C1:C4)") == "='内部'!<格>+SUM(C1:C4)"


def test_compress_flow_shape(flow):
    sc, grids, nums = flow
    c = A.compress(grids, sc, S.detect_facts(grids, sc))
    assert "## 工作表「客流汇总」 已用区域 B2:AG30，非空格 771" in c.text
    assert "C28:AG28 =SUM(C22:C25) 逐列平移（31 格，均有保存值）" in c.text
    assert "C10:AG10「·」×31" in c.text
    assert "C5:AG7 整数" in c.text
    assert "F1（等式" in c.text and "分段 id「客流汇总_按日」" in c.text
    assert "「日间时段客流（人次）」：日间、日间时段客流" in c.text
    assert not any(str(v) in c.text for v in nums["jia"] + nums["yi"])
    # 压缩表示只有规格里的四个字段（接口层照此存库）
    assert set(c.__dataclass_fields__) == {"text", "chars", "truncated", "sha256"}


def test_compress_folds_large_sheet():
    wb = Workbook()
    ws = wb.active
    ws.title = "大表"
    ws.append(["编号", "金额"])
    for i in range(400):
        ws.append([f"N{i:04d}", 1000 + i])
    sc, grids = build(wb)
    c = A.compress(grids, sc, DraftFacts())
    assert c.truncated
    assert "第 201–381 行（共 181 行）折叠为列画像：" in c.text      # 前 200 行、后 20 行照常
    assert "N0300" not in c.text


@pytest.mark.asyncio
async def test_too_large_raises_before_calling_model(flow, monkeypatch):
    monkeypatch.setattr(A, "COMPRESS_MAX_CHARS_PER_SHEET", 100)
    model = FakeModel([reply(flow_json())])
    with pytest.raises(A.AiTooLarge):
        await run(flow, model)
    assert model.calls == []
    monkeypatch.setattr(A, "COMPRESS_MAX_CHARS_PER_SHEET", 12_000)
    monkeypatch.setattr(A, "COMPRESS_MAX_CHARS", 100)
    with pytest.raises(A.AiTooLarge):
        await run(flow, model)
    assert model.calls == []


# ==========================================================================
# 起草、修订、失败
# ==========================================================================


@pytest.mark.asyncio
async def test_invalid_json_then_valid(flow):
    model = FakeModel([reply("我先想一想……"), reply(flow_json(), inp=200, out=80)])
    res = await run(flow, model)
    assert res.draft is not None and res.draft.origin == "ai" and res.error is None
    assert res.attempts == 2 and len(res.usage) == 2
    assert [u.ok for u in res.usage] == [True, True]
    assert res.usage[1].input_tokens == 200 and res.usage[1].output_tokens == 80
    assert all(u.model == "fake-model" and u.attempt == i + 1 for i, u in enumerate(res.usage))
    assert res.draft.cards[0].id == "ai_origin" and res.draft.cards[-1].id == "mode"
    assert res.draft.complete
    assert Recipe.model_validate(res.draft.recipe).model_dump(mode="json", exclude_defaults=True) == \
        Recipe.model_validate(FLOW).model_dump(mode="json", exclude_defaults=True)
    assert res.compressed_chars > 0 and len(res.compressed_sha256) == 64
    # 修订是同一轮对话：system + 首次 prompt + 上一轮回复（原样）+ repair
    second = model.calls[1]
    assert len(second) == 4
    assert isinstance(second[0], SystemMessage) and isinstance(second[1], HumanMessage)
    assert second[0].content == A.RECIPE_DRAFT_SYSTEM
    assert second[1].content == model.calls[0][1].content
    assert isinstance(second[2], AIMessage) and second[2].content == "我先想一想……"
    assert isinstance(second[3], HumanMessage) and "不是合法的 JSON" in second[3].content


@pytest.mark.asyncio
async def test_block_list_content_is_extracted(flow):
    content = [{"type": "thinking", "thinking": "先写交叉表 {\"x\": 1}", "signature": "s"},
               {"type": "text", "text": "```json\n" + flow_json() + "\n```"}]
    model = FakeModel([reply(content)])
    res = await run(flow, model)
    assert res.draft is not None and res.attempts == 1


@pytest.mark.asyncio
async def test_truncated_reply_is_not_repaired(flow):
    model = FakeModel([reply(flow_json()[:200], meta={"stop_reason": "max_tokens"}),
                       reply(flow_json())])
    res = await run(flow, model)
    assert res.draft is not None and res.attempts == 2
    assert [u.ok for u in res.usage] == [False, True]
    assert len(model.calls[1]) == 2                  # 没有 repair：同样的消息列表重来
    assert [m.content for m in model.calls[1]] == [m.content for m in model.calls[0]]
    # 三次都截断：error 是截断的原因，不是「不是 JSON」
    cut = reply("{", meta={"finish_reason": "length"})
    res2 = await run(flow, FakeModel([cut, cut, cut]))
    assert res2.draft is None and "输出被截断" in (res2.error or "")
    assert [u.ok for u in res2.usage] == [False, False, False]


@pytest.mark.asyncio
async def test_timeouts_fail_without_raising(flow, monkeypatch):
    monkeypatch.setattr(A, "AI_TIMEOUT_S", 0.05)

    async def slow():
        await asyncio.sleep(1)
        return reply(flow_json())
    model = FakeModel([slow, slow, slow])
    res = await run(flow, model)
    assert res.draft is None and res.attempts == 3
    assert "调用超时" in (res.error or "")
    assert len(res.usage) == 3 and not any(u.ok for u in res.usage)
    assert all(len(c) == 2 for c in model.calls)


@pytest.mark.asyncio
async def test_call_exception_then_success(flow):
    model = FakeModel([RuntimeError("connection reset"), reply(flow_json())])
    res = await run(flow, model)
    assert res.draft is not None and res.attempts == 2
    assert [u.ok for u in res.usage] == [False, True]
    res2 = await run(flow, FakeModel([RuntimeError("boom")] * 3))
    assert res2.draft is None and (res2.error or "").startswith("调用失败：")


@pytest.mark.asyncio
async def test_note_and_unclaimed_fact_trigger_repair(flow):
    with_note = json.loads(flow_json())
    with_note["tables"][0]["note"] = "这是AI写的说明"
    unclaimed = json.loads(flow_json())
    unclaimed["relations"] = unclaimed["relations"][:1]
    calls: list = []
    model = FakeModel([reply(json.dumps(with_note, ensure_ascii=False)),
                       reply(json.dumps(unclaimed, ensure_ascii=False)), reply(flow_json())])
    res = await run(flow, model, validate=claims_validate(calls))
    assert res.draft is not None and res.attempts == 3
    assert [c[2] for c in calls] == ["ai", "ai", "ai"]
    assert "/tables/0/note：" in model.calls[1][-1].content and "ai_note_forbidden" not in model.calls[1][-1].content
    assert "没有被认领" in model.calls[2][-1].content


@pytest.mark.asyncio
async def test_dry_run_problems_are_fed_back_with_model_message_only(flow):
    sc, grids, _ = flow
    grids = dict(grids)
    g = grids["客流汇总"]
    cells = dict(g.cells)
    cells[(31, 3)] = GridCell(987654)
    cells[(31, 4)] = GridCell("约7777")
    grids["客流汇总"] = Grid(g.sheet, g.bounds, cells, g.merges, g.hidden_rows, g.hidden_cols)
    n = {"k": 0}

    def dry(recipe):
        n["k"] += 1
        if n["k"] == 1:
            return Extraction(ok=False, problems=[
                Problem("outside_number", "structure", "第 31 行 C31 的 987654 在数据区外", ["客流汇总!C31"]),
                Problem("value_not_number", "structure", "D31 的「约7777」不是数字", ["客流汇总!D31"],
                        model_message="D31 不是数字"),
                Problem("outside_digits", "confirm", "区域外有含数字的文字"),
            ])
        return Extraction(ok=True)
    model = FakeModel([reply(flow_json()), reply(flow_json())])
    res = await A.ai_draft(None, grids=grids, scan=sc, facts=FACTS, rules_draft=None, filename="x.xlsx",  # type: ignore[arg-type]
                           validate=claims_validate(), dry_run=dry, model=model)
    assert res.draft is not None and res.draft.complete and res.attempts == 2
    fed = model.calls[1][-1].content
    assert "outside_number（客流汇总!C31）" in fed
    assert "value_not_number：D31 不是数字（客流汇总!D31）" in fed
    assert "987654" not in fed and "7777" not in fed and "outside_digits" not in fed


@pytest.mark.asyncio
async def test_cpu_work_runs_off_the_event_loop_and_async_dry_run_is_awaited(flow, monkeypatch):
    """ai_draft 跑在事件循环上：静态校验、干跑、卡片问题都不能在循环线程上同步执行（一次干跑可能要几秒，
    会堵住同进程的所有请求）。协程函数的 dry_run 直接 await（接口层那份自己占解析名额、进线程池）；
    普通函数的 dry_run 放进线程池。"""
    import threading

    loop_thread = threading.get_ident()
    where: dict[str, list[int]] = {"validate": [], "questions_for": [], "sync_dry": []}
    inner = claims_validate()

    def validate(data, facts, origin):
        where["validate"].append(threading.get_ident())
        return inner(data, facts, origin)

    real_questions = S.questions_for

    def questions_for(*a, **kw):
        where["questions_for"].append(threading.get_ident())
        return real_questions(*a, **kw)
    monkeypatch.setattr(S, "questions_for", questions_for)

    awaited: list[str] = []

    async def dry_async(recipe):
        awaited.append(type(recipe).__name__)
        await asyncio.sleep(0)
        return Extraction(ok=True)

    res = await run(flow, FakeModel([reply(flow_json())]), validate=validate, dry_run=dry_async)
    assert res.draft is not None and res.draft.complete and awaited == ["Recipe"]
    assert where["validate"] and all(t != loop_thread for t in where["validate"])
    assert where["questions_for"] and all(t != loop_thread for t in where["questions_for"])

    def dry_sync(recipe):
        where["sync_dry"].append(threading.get_ident())
        return Extraction(ok=True)

    res = await run(flow, FakeModel([reply(flow_json())]), dry_run=dry_sync)
    assert res.draft is not None and where["sync_dry"] and all(t != loop_thread for t in where["sync_dry"])

    # 协程版干跑出错：同普通版，变成一条起草失败原因，不让起草整个失败
    async def boom(recipe):
        raise RuntimeError("内部细节")
    res = await run(flow, FakeModel([reply(flow_json())] * 3), dry_run=boom)
    assert res.draft is not None and not res.draft.complete
    assert res.draft.failures == [S._FAIL_TEXT["dry_run"]] and "内部细节" not in json.dumps(res.draft.failures)


@pytest.mark.asyncio
async def test_last_dry_run_problems_still_return_draft(flow):
    """次数用完时，静态校验通过但干跑仍有结构问题的配方照样返回（问题随草稿交给用户）。"""
    prob = Problem("cell_unclaimed", "structure", "第 31 行没有去处", ["客流汇总!B31"], model_message="第 31 行没有去处")
    model = FakeModel([reply(flow_json())] * 3)
    res = await run(flow, model, dry_run=lambda r: Extraction(ok=False, problems=[prob]))
    assert res.attempts == 3 and res.draft is not None and not res.draft.complete
    assert res.draft.failures == ["第 31 行没有去处"] and res.error is None
    # 后面一次退化成不合法的 JSON：仍返回之前那份静态校验通过的配方
    model2 = FakeModel([reply(flow_json()), reply("坏了"), reply("还是坏的")])
    res2 = await run(flow, model2, dry_run=lambda r: Extraction(ok=False, problems=[prob]))
    assert res2.draft is not None and res2.attempts == 3


@pytest.mark.asyncio
async def test_never_valid_returns_none_with_usage(flow):
    model = FakeModel([reply("x"), reply("{\"recipe_format\": 1}"), reply("y")])
    res = await run(flow, model)
    assert res.draft is None and res.error == "AI 未能起草出合法的配方"
    assert len(res.usage) == 3 and res.attempts == 3


@pytest.mark.asyncio
async def test_prompt_uses_failures_for_model_not_failures(flow):
    rules = Draft(recipe=None, complete=False, origin="rules",
                  failures=["C12 的「约987654」不是数字"], failures_for_model=["工作表「客流汇总」：C12 不是数字"])
    model = FakeModel([reply(flow_json())])
    await run(flow, model, rules=rules)
    prompt = model.calls[0][1].content
    assert "987654" not in prompt
    assert "工作表「客流汇总」：C12 不是数字" in prompt
    assert "# 配方 JSON Schema" in prompt and '"recipe_format"' in prompt
    assert "人次、人、元" in prompt
    # 半成品配方也带上（紧凑形式）
    sc, grids, _ = flow
    rd = S.draft(sc, grids, "x.xlsx", validate=ok_validate())
    model2 = FakeModel([reply(flow_json())])
    await run(flow, model2, rules=rd)
    assert "半成品配方" in model2.calls[0][1].content and "客流汇总_按日" in model2.calls[0][1].content


@pytest.mark.asyncio
async def test_given_compressed_is_what_is_sent(flow):
    sc, grids, _ = flow
    c = A.compress(grids, sc, FACTS)
    model = FakeModel([reply(flow_json())])
    res = await run(flow, model, compressed=c)
    assert c.text in model.calls[0][1].content
    assert res.compressed_sha256 == c.sha256


@pytest.mark.asyncio
async def test_rules_draft_section_hides_hidden_and_numeric_labels():
    wb = Workbook()
    ws = wb.active
    ws.title = "表"
    ws["B2"] = "2026年8月1日至2026年8月8日"
    for i in range(8):
        ws.cell(4, 3 + i, f"8月{i + 1}日")
    for r, lab in ((5, "甲（人次）"), (6, "藏起来的指标"), (7, "2025")):
        ws.cell(r, 2, lab)
        for i in range(8):
            ws.cell(r, 3 + i, 10 + r + i)
    ws.row_dimensions[6].hidden = True
    sc, grids = build(wb)
    rd = S.draft(sc, grids, "x.xlsx", validate=ok_validate())
    assert rd.recipe is not None and "藏起来的指标" in json.dumps(rd.recipe, ensure_ascii=False)
    model = FakeModel([reply(flow_json())])
    await A.ai_draft(None, grids=grids, scan=sc, facts=rd.facts, rules_draft=rd, filename="x.xlsx",  # type: ignore[arg-type]
                     validate=ok_validate(), dry_run=None, model=model)
    prompt = model.calls[0][1].content
    assert "藏起来的指标" not in prompt
    assert '"2025"' not in prompt


# ==========================================================================
# 评审修复：数据行不当表头外发、压缩表示只靠 text、日期轴不发值
# ==========================================================================


@pytest.mark.asyncio
async def test_list_row_with_blank_number_does_not_leak():
    """名单第 4 行的金额空着：这一行不能被当成表头，姓名、电话、备注都不出现在压缩表示和提示词里。"""
    from tests.test_recipe_suggest import _roster

    for blank in (False, True):
        sc, grids = build(_roster(blank_before_wang=blank))
        rd = S.draft(sc, grids, "x.xlsx", validate=ok_validate())
        c = A.compress(grids, sc, rd.facts)
        model = FakeModel([reply(flow_json())])
        await A.ai_draft(None, grids=grids, scan=sc, facts=rd.facts, rules_draft=rd, filename="x.xlsx",  # type: ignore[arg-type]
                         validate=ok_validate(), dry_run=None, model=model)
        prompt = model.calls[0][1].content
        for secret in ("张三", "王五", "钱七", "138-1234-5678", "137-2222-3333", "137_2222_3333", "转介绍", "老客户"):
            assert secret not in c.text, (blank, secret)
            assert secret not in prompt, (blank, secret)
        assert "A 列「姓名」是数据列（只给画像）：5 格" in c.text
        assert "B 列「电话」是数据列" in c.text and "D 列「备注」是数据列" in c.text


@pytest.mark.asyncio
async def test_rules_section_scrubs_column_names_made_from_hidden_text():
    """万一起草把数据行认成了表头：配方里的列名是 to_sql_name 变过的写法，同样要遮掉。"""
    from tests.test_recipe_suggest import _roster

    sc, grids = build(_roster())
    rd = S.draft(sc, grids, "x.xlsx", validate=ok_validate())
    bad = json.loads(json.dumps(rd.recipe, ensure_ascii=False))
    cols = bad["sheets"][0]["blocks"][0]["columns"]
    cols[0].update(header="王五", name="王五")
    cols[1].update(header="137-2222-3333", name="c_137_2222_3333")
    bad["tables"][0]["grain"] = ["王五"]
    leaky = Draft(recipe=bad, complete=False, origin="rules", failures=[], failures_for_model=[])
    model = FakeModel([reply(flow_json())])
    await A.ai_draft(None, grids=grids, scan=sc, facts=rd.facts, rules_draft=leaky, filename="x.xlsx",  # type: ignore[arg-type]
                     validate=ok_validate(), dry_run=None, model=model)
    prompt = model.calls[0][1].content
    assert "半成品配方" in prompt
    assert "王五" not in prompt and "137" not in prompt


@pytest.mark.asyncio
async def test_compressed_with_only_the_four_fields_is_sent_once(flow):
    """接口层把压缩表示按四个字段存库再取回：提示词的前三段就是 text，不重复、不重算。"""
    sc, grids, _ = flow
    facts = S.detect_facts(grids, sc)
    c = A.compress(grids, sc, facts)
    stored = A.Compressed(text=c.text, chars=c.chars, truncated=c.truncated, sha256=c.sha256)
    m1, m2 = FakeModel([reply(flow_json())]), FakeModel([reply(flow_json())])
    await run(flow, m1, facts=facts, compressed=stored)
    await run(flow, m2, facts=facts)
    prompt = m1.calls[0][1].content
    for head in ("# 工作表压缩表示", "# 系统发现", "# 分段标题候选词", "# 单位词表", "# 规则起草的半成品与失败原因",
                 "# 配方 JSON Schema"):
        assert prompt.count(head) == 1, head
    assert prompt.startswith(c.text + "\n\n# 单位词表\n")
    assert prompt == m2.calls[0][1].content
    # 与 RECIPE_DRAFT_PROMPT 的写法逐字相同（前三段取压缩表示）
    assert A.RECIPE_DRAFT_PROMPT == A.RECIPE_SECTIONS_PROMPT + "\n\n" + A.RECIPE_REQUEST_PROMPT


def test_date_axis_cells_send_type_only():
    """日期轴上是日期格（不是「8月1日」这样的文字）时只发类型签名：表格里没有统计期格时，轴上的年月日就是数据。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "日报"
    for i in range(8):
        ws.cell(4, 3 + i, dt.datetime(2026, 8, 1 + i))
    for r, lab in ((5, "全日客流（人次）"), (6, "分区甲（人次）")):
        ws.cell(r, 2, lab)
        for i in range(8):
            ws.cell(r, 3 + i, 1000 + r * 10 + i)
    sc, grids = build(wb)
    assert S.layout_hints(grids["日报"]).axis_row == 4
    c = A.compress(grids, sc, S.detect_facts(grids, sc))
    assert "C4:J4 日期" in c.text
    for leak in ("2026-08", "2026年", "8月1日", "08-01"):
        assert leak not in c.text, leak
    assert "B5「全日客流（人次）」" in c.text


# ==========================================================================
# 整体审查（AU / SE）修复
# ==========================================================================


def _hidden_unit_book() -> Workbook:
    """交叉表：第 7 行隐藏，标签带单位后缀（「绝密分区（人次）」），并且是 F1（5 = 6 + 7）的成员。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "表"
    ws["B2"] = "统计时间范围：2026年8月1日至2026年8月8日"
    for i in range(8):
        ws.cell(4, 3 + i, f"8月{i + 1}日")
    labels = ((5, "全日客流（人次）"), (6, "分区甲（人次）"), (7, "绝密分区（人次）"))
    for r, lab in labels:
        ws.cell(r, 2, lab)
    for i in range(8):
        a, b = 100 + i, 30 + i
        ws.cell(6, 3 + i, a)
        ws.cell(7, 3 + i, b)
        ws.cell(5, 3 + i, a + b)
    ws.row_dimensions[7].hidden = True
    return wb


@pytest.mark.asyncio
async def test_hidden_label_with_unit_suffix_never_leaves_in_any_call():
    """SE-1、SE-2：隐藏行的标签（含去掉单位后缀的「绝密分区」）不出现在首轮全文、也不出现在任何一轮修订消息里。"""
    sc, grids = build(_hidden_unit_book())
    rd = S.draft(sc, grids, "x.xlsx", validate=ok_validate())
    facts = S.detect_facts(grids, sc)
    assert rd.recipe is not None and "绝密分区" in json.dumps(rd.recipe, ensure_ascii=False)
    assert any("绝密分区" in json.dumps(f.detail, ensure_ascii=False) for f in facts.facts)
    c = A.compress(grids, sc, facts)
    first = A.first_call_text(c, grids=grids, scan=sc, rules_draft=rd)
    assert "绝密" not in first
    assert "半成品配方" in first and "人次" in first        # 半成品照发，只遮隐藏的部分

    # 模型把首轮看到的半成品原样抄回来（<隐藏> 对不上真实标签）：静态校验的消息会引用系统发现的标签原文
    echoed = json.dumps(A._Outbound(A._secret_texts(grids, sc), grids).json(rd.recipe), ensure_ascii=False)
    from app.data import recipe as R

    def validate(data, f, origin):
        return R.validate_recipe(data, facts=f, origin=origin)
    model = FakeModel([reply(echoed), reply(echoed), reply(echoed)])
    res = await A.ai_draft(None, grids=grids, scan=sc, facts=facts, rules_draft=rd, filename="x.xlsx",  # type: ignore[arg-type]
                           validate=validate, dry_run=None, model=model, compressed=c)
    assert res.attempts == 3 and len(model.calls) == 3
    repairs = [m.content for call in model.calls for m in call if isinstance(m, HumanMessage)]
    assert len(repairs) >= 3 and any("上一次的配方有以下问题" in x for x in repairs)
    for text in repairs:
        assert "绝密" not in text


def test_hidden_column_header_with_unit_suffix_is_scrubbed():
    """SE-2：隐藏列的表头「机密奖金（元）」去掉单位后成了列名「机密奖金」和 units 的键，同样不能外发。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "明细"
    for j, h in enumerate(["日期", "门店", "机密奖金（元）", "数量（件）"]):
        ws.cell(1, j + 1, h)
    for i in range(10):
        ws.cell(2 + i, 1, dt.datetime(2026, 8, 1 + i))
        ws.cell(2 + i, 2, ["店甲", "店乙"][i % 2])
        ws.cell(2 + i, 3, 100 + i)
        ws.cell(2 + i, 4, 5 + i)
    ws.column_dimensions["C"].hidden = True
    sc, grids = build(wb)
    rd = S.draft(sc, grids, "x.xlsx", validate=ok_validate())
    assert rd.recipe is not None and "机密奖金" in json.dumps(rd.recipe, ensure_ascii=False)
    text = A.first_call_text(A.compress(grids, sc, rd.facts), grids=grids, scan=sc,
                             rules_draft=Draft(recipe=rd.recipe, complete=False, origin="rules"))
    assert "机密" not in text and "半成品配方" in text and "数量" in text


def test_outbound_line_scrubs_quoted_numeric_text_and_hidden_variants():
    sc, grids = build(_hidden_unit_book())
    ob = A._Outbound(A._secret_texts(grids, sc), grids)
    line = ob.line("应为「全日客流（人次）」=「分区甲（人次）」+「绝密分区（人次）」；列「绝密分区」；表「绝密分区_人次」")
    assert "绝密" not in line and "全日客流（人次）" in line and "分区甲（人次）" in line


def test_total_word_in_middle_of_data_column_is_profiled_not_sent():
    """SE-3：备注列中间一格以「累计」开头：它是数据，只进画像，不当合计标签原样外发；最左一列的「合计」照发。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "名单"
    ws.append(["姓名", "产品", "数量", "金额（元）", "备注"])
    rows = [("张一", "甲品", 3, 30, "老客户来电"), ("张二", "甲品", 2, 20, "累计消费12000元"),
            ("张三", "甲品", 1, 10, "新客户到店"), ("张四", "甲品", 5, 50, "转介绍来的"),
            ("张五", "甲品", 4, 40, "门店活动价"), ("张六", "甲品", 6, 60, "线上下单的")]
    for r in rows:
        ws.append(list(r))
    ws.append(["合计", None, 21, 210, None])
    sc, grids = build(wb)
    c = A.compress(grids, sc, S.detect_facts(grids, sc))
    assert "累计消费" not in c.text and "12000" not in c.text
    assert "E 列「备注」是数据列（只给画像）：6 格" in c.text
    assert "A8「合计」" in c.text


def test_extract_json_deep_nesting_is_value_error():
    """SE-6：深层嵌套（json.loads 会抛 RecursionError）一律按「不是合法的 JSON」处理。"""
    deep = '{"x":' + "[" * 12000 + "}"
    with pytest.raises(ValueError):
        A._extract_json(deep)
    closed = '{"x":' + "[" * 100 + "]" * 100 + "}"
    with pytest.raises(ValueError):
        A._extract_json(closed)
    assert A._extract_json('{"a": "[[[[{{{{", "b": [[1]]}') == {"a": "[[[[{{{{", "b": [[1]]}


@pytest.mark.asyncio
async def test_deeply_nested_reply_returns_none_with_usage(flow):
    """SE-6：回复是约 1.2 万层不闭合的「[」时 ai_draft 不抛异常，返回 draft=None、每次调用的用量都在。"""
    deep = '{"x":' + "[" * 12000 + "}"
    model = FakeModel([reply(deep), reply(deep), reply(deep)])
    res = await run(flow, model)
    assert res.draft is None and res.error and len(res.usage) == 3 == res.attempts


@pytest.mark.asyncio
async def test_exception_after_reply_is_returned_not_raised(flow, monkeypatch):
    """SE-6：拿到回复之后的处理（这里是 check_recipe 本身）抛了异常：不逃出 ai_draft，用量照记。"""
    def boom(*a, **kw):
        raise RecursionError("maximum recursion depth exceeded")
    monkeypatch.setattr(S, "check_recipe", boom)
    model = FakeModel([reply(flow_json())])
    res = await run(flow, model)
    assert res.draft is None and len(res.usage) == 1 and res.error


@pytest.mark.asyncio
async def test_given_model_id_is_recorded_and_budget_wrapper_has_a_name(flow):
    """AU-7 / SE-4：接口层把同意时解析出的模型和它的 id 交进来，用量记的就是这个 id；_BudgetFallback 的
    model_name 是 id，不是底层模型对象的 repr。"""
    model = FakeModel([reply(flow_json())], name="ignored")
    res = await run(flow, model, model_id="local-model-A")
    assert [u.model for u in res.usage] == ["local-model-A"]
    wrapped = A._BudgetFallback(object(), object(), "cloud-model-B")
    assert wrapped.model_name == "cloud-model-B"
