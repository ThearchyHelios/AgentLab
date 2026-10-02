"""读取层（xlsx_cells）、扫描的逐行游程与序列化（xlsx_scan.row_runs / scan_to_json）、tabular 的公开名。

夹具都在测试里用 openpyxl 现造，标签用假名；需要 openpyxl 写不出来的形状（共享公式、行号乱序、陈旧的
<dimension>、公式缓存值）时在 zip 层改 XML。
"""
from __future__ import annotations

import datetime as dt
import io
import json
import re
import zipfile

import pytest
from openpyxl import Workbook

from app.data import tabular, xlsx_cells, xlsx_scan
from app.data.xlsx_scan import UnsupportedTable


def _save(wb: Workbook) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _patch(raw: bytes, part: str, fn) -> bytes:
    zin = zipfile.ZipFile(io.BytesIO(raw))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        for it in zin.infolist():
            data = zin.read(it.filename)
            if it.filename == part:
                data = fn(data.decode("utf-8")).encode("utf-8")
            zout.writestr(it, data)
    return out.getvalue()


def _shared_formula_book() -> bytes:
    """B5 是 A5 的共享公式（<f t="shared" si="0"/>），两格都有缓存值。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "表甲"
    for r in range(1, 5):
        ws.cell(r, 1, r)
        ws.cell(r, 2, r * 10)
    ws["A5"] = "=SUM(A1:A4)"
    ws["B5"] = "=SUM(B1:B4)"

    def fix(t: str) -> str:
        t = re.sub(r'<c r="A5"([^>]*)><f>SUM\(A1:A4\)</f>(?:<v\s*/>|<v></v>)?',
                   r'<c r="A5"\1><f t="shared" ref="A5:B5" si="0">SUM(A1:A4)</f><v>10</v>', t)
        t = re.sub(r'<c r="B5"([^>]*)><f>SUM\(B1:B4\)</f>(?:<v\s*/>|<v></v>)?',
                   r'<c r="B5"\1><f t="shared" si="0"/><v>100</v>', t)
        assert 'si="0"/>' in t
        return t

    return _patch(_save(wb), "xl/worksheets/sheet1.xml", fix)


def _mixed_book() -> bytes:
    """各种取值：文本、只含空白的文本、0、False、日期、错误值、无缓存的公式；一张隐藏工作表。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "表甲"
    ws["B2"] = "分区甲"
    ws["C2"] = "   "                   # 只含空白：扫描、读取都算空
    ws["D2"] = 0                       # 0 不算空
    ws["E2"] = False                   # False 不算空
    ws["B3"] = dt.datetime(2026, 8, 1)
    ws["C3"] = "#DIV/0!"
    ws["D3"] = "=B9*2"                 # openpyxl 写公式不带缓存值：读出来值是空、公式照给
    ws["B4"] = 12.5
    ws.merge_cells("B6:D6")
    ws["B6"] = "合并标题"
    ws.row_dimensions[4].hidden = True
    ws.column_dimensions["E"].hidden = True
    hidden = wb.create_sheet("隐藏表")
    hidden["A1"] = "口令"
    hidden.sheet_state = "hidden"
    return _save(wb)


# --------------------------------------------------------------------------
# tabular 的公开名
# --------------------------------------------------------------------------


def test_tabular_public_names_keep_old_aliases():
    assert tabular._open_book is tabular.open_book
    assert tabular._worksheet_for is tabular.worksheet_for


@pytest.mark.parametrize("value,text", [
    (dt.datetime(2026, 8, 1), "2026-08-01"),                 # 期 1 的 #23：零点日期不能写成「… 00:00:00」
    (dt.datetime(2026, 8, 1, 9, 30), "2026-08-01 09:30:00"),
    (dt.date(2026, 8, 1), "2026-08-01"),
    (dt.time(7, 5), "07:05:00"),
    (3.0, "3"), (2.5, "2.5"), (12, "12"), (True, "TRUE"),
    ("  分区甲 ", "分区甲"), ("   ", ""), (None, ""),
])
def test_as_text_matches_simple_import(value, text):
    assert tabular.as_text(value) == text


# --------------------------------------------------------------------------
# 扫描：逐行游程、序列化
# --------------------------------------------------------------------------


def test_row_runs_merge_consecutive_rows_with_same_count():
    wb = Workbook()
    ws = wb.active
    for r in (1, 2, 3):
        ws.cell(r, 1, "甲")
        ws.cell(r, 2, r)
    ws.cell(5, 1, "乙")                 # 第 4 行空、第 5 行 1 格
    ws.cell(6, 1, "丙")
    ws.cell(7, 1, "丁")
    ws.cell(7, 3, 1)
    sc = xlsx_scan.scan(_save(wb))
    assert sc.sheets[0].row_runs == [(1, 3, 2), (5, 6, 1), (7, 7, 2)]
    assert sum((b - a + 1) * n for a, b, n in sc.sheets[0].row_runs) == sc.sheets[0].nonempty


def test_row_runs_fall_back_when_xml_rows_out_of_order():
    wb = Workbook()
    ws = wb.active
    for r in (1, 2, 3):
        ws.cell(r, 1, f"标签{r}")

    def swap(t: str) -> str:
        rows = re.findall(r"<row [^>]*>.*?</row>", t)
        assert len(rows) == 3
        return t.replace(rows[0] + rows[1] + rows[2], rows[0] + rows[2] + rows[1])

    raw = _patch(_save(wb), "xl/worksheets/sheet1.xml", swap)
    sc = xlsx_scan.scan(raw)
    assert sc.sheets[0].row_runs == [(1, 3, 1)]


def test_row_runs_add_up_duplicate_rows():
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "甲"
    ws["A2"] = "乙"

    def dup(t: str) -> str:
        row = re.search(r'<row r="2"[^>]*>.*?</row>', t).group(0)
        return t.replace(row, row + row.replace('r="A2"', 'r="B2"'))

    sc = xlsx_scan.scan(_patch(_save(wb), "xl/worksheets/sheet1.xml", dup))
    assert sc.sheets[0].row_runs == [(1, 1, 1), (2, 2, 2)]


def test_scan_json_round_trip():
    sc = xlsx_scan.scan(_mixed_book())
    data = xlsx_scan.scan_to_json(sc)
    again = xlsx_scan.scan_from_json(json.loads(json.dumps(data, ensure_ascii=False)))
    assert again == sc
    assert data == json.loads(json.dumps(data))      # 返回值本身就是 JSON 形状（元组已写成列表）


def test_scan_from_json_tolerates_old_and_newer_shapes():
    sc = xlsx_scan.scan(_mixed_book())
    data = xlsx_scan.scan_to_json(sc)
    for sheet in data["sheets"]:
        del sheet["row_runs"]                         # 老暂存区没有这个字段
        sheet["future_field"] = 1                     # 更新的版本多写的字段
    data["future_top"] = True
    again = xlsx_scan.scan_from_json(data)
    assert [s.row_runs for s in again.sheets] == [[] for _ in again.sheets]
    assert [s.nonempty for s in again.sheets] == [s.nonempty for s in sc.sheets]
    assert again.sheets[0].hidden_rows == sc.sheets[0].hidden_rows


# --------------------------------------------------------------------------
# 网格
# --------------------------------------------------------------------------


def test_read_grid_values_and_nonempty_rule_match_scan():
    raw = _mixed_book()
    sc = xlsx_scan.scan(raw)
    grid = xlsx_cells.read_grid(raw, sc, "表甲")
    s = sc.sheet("表甲")
    assert len(grid.cells) == s.nonempty
    per_row = {r: len(grid.row(r)) for r in grid.row_numbers()}
    expected = {r: n for a, b, n in s.row_runs for r in range(a, b + 1)}
    assert per_row == expected
    assert (2, 3) not in grid.cells                   # 只含空白的文本
    assert grid.get(2, 4).value == 0 and grid.get(2, 5).value is False
    assert grid.get(3, 2).value == dt.datetime(2026, 8, 1)
    assert grid.get(3, 3).value == "#DIV/0!"
    uncached = grid.get(3, 4)
    assert uncached.value is None and uncached.formula == "=B9*2" and not uncached.has_cache
    assert grid.bounds == (2, 2, 6, 5)
    assert grid.merges == [(6, 2, 6, 4)]
    assert grid.hidden_rows == {4} and grid.hidden_cols == {5}
    assert grid.truncated is False


def test_read_grid_refuses_hidden_and_unknown_sheets():
    raw = _mixed_book()
    sc = xlsx_scan.scan(raw)
    with pytest.raises(UnsupportedTable):
        xlsx_cells.read_grid(raw, sc, "隐藏表")
    with pytest.raises(UnsupportedTable):
        xlsx_cells.read_grid(raw, sc, "没有这张表")


def test_shared_formula_reads_expanded_text():
    raw = _shared_formula_book()
    sc = xlsx_scan.scan(raw)
    grid = xlsx_cells.read_grid(raw, sc, "表甲")
    assert grid.get(5, 1).formula == "=SUM(A1:A4)" and grid.get(5, 1).value == 10
    assert grid.get(5, 2).formula == "=SUM(B1:B4)" and grid.get(5, 2).value == 100
    rows = list(xlsx_cells.iter_rows(raw, sc, "表甲", min_row=5, max_row=5, min_col=1, max_col=2))
    assert [(r, [(c, g.formula) for c, g in cells]) for r, cells in rows] == [
        (5, [(1, "=SUM(A1:A4)"), (2, "=SUM(B1:B4)")])]


def test_iter_rows_gives_only_nonempty_cells_in_column_order():
    raw = _mixed_book()
    sc = xlsx_scan.scan(raw)
    rows = list(xlsx_cells.iter_rows(raw, sc, "表甲", min_row=1, max_row=6, min_col=1, max_col=5))
    assert [r for r, _ in rows] == [2, 3, 4, 6]
    assert [c for c, _ in rows[0][1]] == [2, 4, 5]
    assert all(cells for _, cells in rows)


def test_grid_too_large_and_max_rows(monkeypatch):
    wb = Workbook()
    ws = wb.active
    ws.title = "表甲"
    for r in range(1, 21):
        ws.cell(r, 1, f"项{r}")
        ws.cell(r, 2, r)
    raw = _save(wb)
    sc = xlsx_scan.scan(raw)
    monkeypatch.setattr(xlsx_cells, "GRID_MAX_CELLS", 30)
    with pytest.raises(xlsx_cells.GridTooLarge):
        xlsx_cells.read_grid(raw, sc, "表甲")
    grid = xlsx_cells.read_grid(raw, sc, "表甲", max_rows=5)
    assert grid.truncated and grid.row_numbers() == [1, 2, 3, 4, 5]
    assert grid.bounds == (1, 1, 20, 2)                # 边界仍是整张表的
    # 给了 max_rows 也不会超过格数上限：前 50 行就超了，读到上一行为止
    capped = xlsx_cells.read_grid(raw, sc, "表甲", max_rows=50)
    assert capped.truncated and len(capped.cells) == 30 and capped.row_numbers() == list(range(1, 16))
    monkeypatch.setattr(xlsx_cells, "GRID_MAX_CELLS", 200_000)
    whole = xlsx_cells.read_grid(raw, sc, "表甲", max_rows=50)
    assert not whole.truncated and len(whole.cells) == 40


def test_stale_dimension_does_not_change_the_grid():
    raw = _mixed_book()
    stale = _patch(raw, "xl/worksheets/sheet1.xml", lambda t: re.sub(r'<dimension ref="[^"]*"/>',
                                                                     '<dimension ref="B2:B2"/>', t))
    assert stale != raw
    g1 = xlsx_cells.read_grid(raw, xlsx_scan.scan(raw), "表甲")
    g2 = xlsx_cells.read_grid(stale, xlsx_scan.scan(stale), "表甲")
    assert g1.cells == g2.cells


# --------------------------------------------------------------------------
# 预览
# --------------------------------------------------------------------------


def test_preview_shape_and_kinds():
    raw = _mixed_book()
    sc = xlsx_scan.scan(raw)
    p = xlsx_cells.preview(xlsx_cells.read_grid(raw, sc, "表甲"))
    assert set(p) == {"sheet", "bounds", "total_rows", "total_cols", "truncated", "rows", "cols", "cells",
                      "formulas", "merges", "hidden_rows", "hidden_cols"}
    assert p["bounds"] == "B2:E6" and p["total_rows"] == 5 and p["total_cols"] == 4 and not p["truncated"]
    kinds = {(r, c): (text, kind) for r, c, text, kind in p["cells"]}
    assert kinds[(2, 2)] == ("分区甲", "text")
    assert kinds[(2, 4)] == ("0", "number")
    assert kinds[(2, 5)] == ("FALSE", "bool")
    assert kinds[(3, 2)] == ("2026-08-01", "date")
    assert kinds[(3, 3)] == ("#DIV/0!", "error")
    assert kinds[(3, 4)] == ("", "formula_uncached")
    assert kinds[(4, 2)] == ("12.5", "number")
    assert p["formulas"] == {"D3": "=B9*2"}
    assert p["merges"] == ["B6:D6"] and p["hidden_rows"] == [4] and p["hidden_cols"] == [5]
    json.dumps(p, ensure_ascii=False)


def test_preview_window_truncates():
    raw = _shared_formula_book()
    sc = xlsx_scan.scan(raw)
    p = xlsx_cells.preview(xlsx_cells.read_grid(raw, sc, "表甲"), max_rows=2, max_cols=1)
    assert p["truncated"] and p["rows"] == [1, 2] and p["cols"] == [1]
    assert {(r, c) for r, c, _, _ in p["cells"]} == {(1, 1), (2, 1)}
    assert p["formulas"] == {}
    formula = xlsx_cells.preview(xlsx_cells.read_grid(raw, sc, "表甲"))
    assert [k for _, _, _, k in formula["cells"] if k == "formula"] == ["formula", "formula"]
    assert formula["formulas"] == {"A5": "=SUM(A1:A4)", "B5": "=SUM(B1:B4)"}
