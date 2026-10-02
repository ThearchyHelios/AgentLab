"""框选转锚点（recipe_select，WP-2，P3-SPEC 第 4 节）：八种 as 各一例成功，泛化不了的每个 code 各一例；c01 的规则
等价（框选结果与草稿 canonical 相同）；列表复用起草器（占位符、合并填充、跳过空行）；after_title 跳过含数字的格；
客户端表名的校验和撞名（含退役名）；重放比对。

网格用读取层（xlsx_cells.read_grid）从合成夹具读出；干跑用期 2 的执行器。期 2 的执行器不给 RegionMark 填块 id，
recipe_select 在没有块 id 时退回用起草器的定位函数近似，这里两种都测（块 id 的那一路用手写的 Extraction）。
"""
from __future__ import annotations

import copy
import datetime as dt
import io
import json
from pathlib import Path
from typing import Any

import pytest
from openpyxl import Workbook, load_workbook

from app.data import recipe_engine, xlsx_cells, xlsx_scan
from app.data import recipe_suggest as S
from app.data.names import collide_key
from app.data.recipe import apply_patch, recipe_sha256, validate_recipe
from app.data.recipe_fixes import EditRequestError
from app.data.recipe_select import replay_compare, selection_edit
from app.data.recipe_types import (
    EditResult,
    Extraction,
    GridCell,
    Recipe,
    RegionMark,
    Selection,
)
from tests.fixtures.xlsx import lab, save
from tests.fixtures.xlsx.flow import flow_workbook

FLOW = Recipe.model_validate(json.loads((Path(__file__).parent / "fixtures" / "recipes" / "flow_recipe.json")
                                        .read_text(encoding="utf-8"))).model_dump(mode="json")
SEP = dt.date(2026, 9, 1)
SHEET = "客流汇总"


class Book:
    """一份合成文件：原件、扫描、网格、规则草稿（完整形式）和按某份配方的干跑。"""

    def __init__(self, raw: bytes, fn: str, *, max_rows: int = 500):
        self.raw, self.fn = raw, fn
        self.scan = xlsx_scan.scan(raw)
        self.grids = {s.name: xlsx_cells.read_grid(raw, self.scan, s.name, max_rows=max_rows)
                      for s in self.scan.sheets if s.state == "visible" and s.nonempty}

    @property
    def grid(self):
        return next(iter(self.grids.values()))

    def draft(self):
        return S.draft(self.scan, self.grids, self.fn,
                       validate=lambda r, f, o: validate_recipe(r, facts=f, origin=o), dry_run=self.dry)

    def dry(self, recipe: Any) -> Extraction:
        r = recipe if isinstance(recipe, Recipe) else Recipe.model_validate(recipe)
        return recipe_engine.execute(r, self.raw, self.fn, self.scan, None)


def full(recipe: dict[str, Any]) -> dict[str, Any]:
    return Recipe.model_validate(recipe).model_dump(mode="json")


def sha(recipe: dict[str, Any]) -> str:
    return recipe_sha256(Recipe.model_validate(recipe))


def sel(ref: str, as_: str, sheet: str = SHEET, **options) -> Selection:
    return Selection(sheet, ref, as_, options)


def codes(res: EditResult) -> list[str]:
    return [p.code for p in res.problems]


@pytest.fixture(scope="module")
def c01():
    b = Book(*lab.c01("literal"))
    d = b.draft()
    assert d.complete, d.failures
    work = full(d.recipe)
    return b, work, b.dry(work)


@pytest.fixture(scope="module")
def flow_d01():
    return Book(*flow_workbook(SEP, 30, variant="D01"))


def flow_book(variant: str) -> tuple[Book, Extraction]:
    b = Book(*flow_workbook(SEP, 30, variant=variant))
    return b, b.dry(FLOW)


# ==========================================================================
# 列表：c01（4.6 第 1、3 条）
# ==========================================================================


def test_c01_box_gives_the_same_recipe_as_the_rule_draft(c01):
    b, work, ext = c01
    res = selection_edit(work, b.grid, sel("C5:F8", "list", "月报"), extraction=ext)
    assert res.ok, res.problems
    assert res.kind == "selection" and res.key == "list:列表1" and res.block == "列表1"
    # 框换算出的块替换规则草稿的同一个块，工作配方的哈希不变：框选得到的规则和起草器完全一样
    assert res.recipe_sha256_after == sha(work)
    assert res.expected == {"header": "C5:F5", "data": "C6:F8", "total": "C9:F9"}
    assert [(a.kind, a.text) for a in res.anchors] == [
        ("header", "地区"), ("header", "产品"), ("header", "销量"), ("header", "金额"), ("total_word", "合计")]
    new = apply_patch(work, res.ops)
    rc = replay_compare(res, b.dry(new))
    assert rc.match and rc.diffs == [] and rc.actual == rc.expected


def test_c01_shifted_file_replays_with_the_same_table_hashes(c01):
    """4.6 第 2 条：整张表挪了位置，按框选得到的配方照样认出，表哈希逐一相同（配方里没有坐标）。"""
    b, work, ext = c01
    res = selection_edit(work, b.grid, sel("C5:F8", "list", "月报"), extraction=ext)
    new = apply_patch(work, res.ops)
    moved = Book(*lab.c01("literal", shift=(3, 2)))
    ext2 = moved.dry(new)
    assert ext2.ok and [p.code for p in ext2.problems if p.category == "structure"] == []
    assert ext2.table_hashes == b.dry(new).table_hashes


def test_c01_box_too_short_is_not_anchorable(c01):
    b, work, ext = c01
    res = selection_edit(work, b.grid, sel("C5:F7", "list", "月报"), extraction=ext)
    assert not res.ok and res.ops == [] and codes(res) == ["selection_bottom_not_anchorable"]
    msg = res.problems[0].message
    assert "第 8 行还有数据" in msg and "第 8 行）" in msg and "下边界按规则推断" in msg
    assert res.problems[0].category == "recipe"


def test_c01_bottom_auto_infers_the_end(c01):
    b, work, ext = c01
    res = selection_edit(work, b.grid, sel("C5:F7", "list", "月报", bottom="auto"), extraction=ext)
    assert res.ok and res.expected == {"header": "C5:F5", "data": "C6:F8", "total": "C9:F9"}
    assert res.recipe_sha256_after == sha(work)
    assert any("重放的数据区到第 8 行" in s for s in res.summary)
    rc = replay_compare(res, b.dry(apply_patch(work, res.ops)))
    assert rc.match


@pytest.mark.parametrize("ref", ["C5:F9", "C5:F10"])
def test_c01_box_including_the_total_row_reads_it_as_total(c01, ref):
    """把合计行（和它下面的空行）也框进来：合计行照起草器认作合计行，不当数据行导入（H2：当数据行的话
    SUM(金额) 翻倍，干跑、核对、重放比对都看不出来）。框底多带空行时与框到合计行相同，配方与规则草稿逐字相同。"""
    b, work, ext = c01
    res = selection_edit(work, b.grid, sel(ref, "list", "月报"), extraction=ext)
    assert res.ok, res.problems
    assert res.expected == {"header": "C5:F5", "data": "C6:F8", "total": "C9:F9"}
    assert res.recipe_sha256_after == sha(work)
    assert ("total_word", "合计") in [(a.kind, a.text) for a in res.anchors]
    # 配方为空时同样：框选新建的配方与规则草稿相同
    fresh = selection_edit(None, b.grid, sel(ref, "list", "月报"), extraction=None)
    assert fresh.ok and fresh.recipe_sha256_after == sha(work)
    new = apply_patch(work, res.ops)
    assert replay_compare(res, b.dry(new)).match


def test_c01_box_down_to_the_notes_reports_the_total_in_the_middle(c01):
    """框一直拉到下方的备注、制表人行：合计行夹在框中间，无法确定它汇总的是哪几行，报错并建议把下边移到合计行，
    不把合计、备注当数据导入。"""
    b, work, ext = c01
    res = selection_edit(work, b.grid, sel("C5:F12", "list", "月报"), extraction=ext)
    assert not res.ok and codes(res) == ["selection_bottom_not_anchorable"]
    msg = res.problems[0].message
    assert "第 9 行像合计行" in msg and "把框的下边移到第 9 行" in msg
    assert res.problems[0].cells == ["月报!C9"]


def test_box_with_trailing_blank_rows_ends_at_the_last_filled_row():
    """框底多带了空行、空行下面是备注：下边界看框内最后一个非空行的下一行（空行，遇到空行停止），与框到最后
    一个数据行相同。看框的下一行的话会撞上备注，提示「把框扩大到第 5 行」，等于让人把备注也框进来。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "名单"
    for r, row in enumerate([("地区", "销量"), ("分区甲", 3), ("分区乙", 5), (None, None), ("注：合成备注", None)],
                            start=1):
        for c, v in enumerate(row, start=1):
            if v is not None:
                ws.cell(r, c, v)
    b = Book(save(wb), "tail.xlsx")
    tight = selection_edit(None, b.grid, sel("A1:B3", "list", "名单"), extraction=None)
    loose = selection_edit(None, b.grid, sel("A1:B4", "list", "名单"), extraction=None)
    assert tight.ok and loose.ok, loose.problems
    assert loose.expected == tight.expected == {"header": "A1:B1", "data": "A2:B3", "total": None}
    assert loose.recipe_sha256_after == tight.recipe_sha256_after
    assert apply_patch({}, loose.ops)["sheets"][0]["blocks"][0]["rows"] == {"blank_rows": "stop", "total_row": None}


def _subtotals() -> Book:
    """列表中间夹着小计：分区甲两行、小计、分区乙两行、合计。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "明细"
    rows = [("地区", "产品", "金额"), ("分区甲", "产品甲", 10), ("分区甲", "产品乙", 20), ("小计", None, 30),
            ("分区乙", "产品甲", 5), ("分区乙", "产品乙", 6), ("合计", None, 41)]
    for r, row in enumerate(rows, start=1):
        for c, v in enumerate(row, start=1):
            if v is not None:
                ws.cell(r, c, v)
    return Book(save(wb), "subtotal.xlsx")


@pytest.mark.parametrize("ref,bottom,where", [
    ("A1:C7", "box", "第 4 行像合计行"),       # 小计夹在框中间
    ("A1:C4", "box", "第 4 行像合计行"),       # 框到小计为止：它下面紧接着还有明细，是夹在中间的小计
    ("A1:C3", "box", "第 4 行像合计行"),       # 框到小计上一行：下一行是小计，同上
    ("A1:C3", "auto", "第 4 行像合计行"),      # 按规则推断：起草器在这里记失败，框选同样不收
])
def test_subtotal_in_the_middle_is_never_imported_as_data(ref, bottom, where):
    b = _subtotals()
    res = selection_edit(None, b.grid, sel(ref, "list", "明细", bottom=bottom), extraction=None)
    assert not res.ok and set(codes(res)) == {"selection_bottom_not_anchorable"}, res.problems
    assert where in res.problems[0].message and "明细!A4" in res.problems[0].cells


def test_total_word_in_another_column_is_not_imported_as_data():
    """合计字样写在第二列（第一列空着）：起草器记失败，框选同样报错，不当明细导入。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "明细"
    for r, row in enumerate([("地区", "产品", "金额"), ("分区甲", "产品甲", 10), ("分区乙", "产品乙", 20),
                             (None, "合计", 30)], start=1):
        for c, v in enumerate(row, start=1):
            if v is not None:
                ws.cell(r, c, v)
    b = Book(save(wb), "total_b.xlsx")
    res = selection_edit(None, b.grid, sel("A1:C4", "list", "明细"), extraction=None)
    assert codes(res) == ["selection_bottom_not_anchorable"]
    assert "B 列是合计字样" in res.problems[0].message and res.problems[0].cells == ["明细!B4"]


def test_c01_box_on_data_row_header_not_text(c01):
    b, work, ext = c01
    res = selection_edit(work, b.grid, sel("C6:F8", "list", "月报"), extraction=ext)
    assert codes(res) == ["selection_header_not_text"]


def test_c01_box_with_blank_header_column(c01):
    b, work, ext = c01
    res = selection_edit(work, b.grid, sel("C5:G8", "list", "月报"), extraction=ext)
    assert codes(res) == ["selection_header_blank"] and res.problems[0].cells == ["月报!G5"]


def test_list_without_recipe_creates_one_by_root_replace(c01):
    """配方为空（规则起草失败）时用根替换新建配方：框选也是起草失败时不用 AI 的兜底（4.2 第 7 条）。"""
    b, work, _ext = c01
    res = selection_edit(None, b.grid, sel("C5:F8", "list", "月报"), extraction=None)
    assert res.ok and res.ops[0]["op"] == "replace" and res.ops[0]["path"] == ""
    assert res.recipe_sha256_after == sha(work)


def test_header_duplicate():
    wb = Workbook()
    ws = wb.active
    ws.title = "明细"
    for r, row in enumerate([("地区", "金额", "金 额"), ("分区甲", 1, 2), ("分区乙", 3, 4)], start=1):
        for c, v in enumerate(row, start=1):
            ws.cell(r, c, v)
    b = Book(save(wb), "dup.xlsx")
    res = selection_edit(None, b.grid, sel("A1:C3", "list", "明细"), extraction=None)
    assert codes(res) == ["selection_header_duplicate"]


def test_extra_columns_next_to_the_box(c01):
    b, work, ext = c01
    res = selection_edit(work, b.grid, sel("C5:E8", "list", "月报"), extraction=ext)
    assert codes(res) == ["selection_extra_columns"]
    assert "「金额」" in res.problems[0].message and "C5:F8" in res.problems[0].message
    assert res.problems[0].cells == ["月报!F5"]
    # 被替换的块上已经按表头忽略了这一列：不算多出的列
    with_ignore = copy.deepcopy(work)
    with_ignore["sheets"][0]["blocks"][0]["ignore_columns"] = [{"header": "金额", "reason": "合成理由"}]
    ok = selection_edit(with_ignore, b.grid, sel("C5:E8", "list", "月报"), extraction=ext)
    assert ok.ok, ok.problems


def test_box_bottom_on_blank_row_stops_there():
    b = _stacked("补充明细")
    res = selection_edit(None, b.grid, sel("A2:B4", "list", "汇总"), extraction=None)
    assert res.ok, res.problems
    blk = apply_patch({}, res.ops)["sheets"][0]["blocks"][0]
    assert blk["rows"] == {"blank_rows": "stop", "total_row": None}
    assert res.expected == {"header": "A2:B2", "data": "A3:B4", "total": None}


def test_list_options_are_validated(c01):
    b, work, ext = c01
    with pytest.raises(EditRequestError) as e:
        selection_edit(work, b.grid, sel("C5:F8", "list", "月报", header_rows=4), extraction=ext)
    assert e.value.code == "edit_invalid"
    with pytest.raises(EditRequestError):
        selection_edit(work, b.grid, sel("C5:F8", "list", "月报", bottom="下面"), extraction=ext)


def test_out_of_grid_on_truncated_grid():
    b = Book(*lab.c01("literal"), max_rows=6)
    assert b.grid.truncated
    res = selection_edit(None, b.grid, sel("C5:F8", "list", "月报"), extraction=None)
    assert codes(res) == ["selection_out_of_grid"]
    auto = selection_edit(None, b.grid, sel("C5:F6", "list", "月报", bottom="auto"), extraction=None)
    assert auto.ok and "数据延续到已读入范围之外，完整范围在试运行时核对" in auto.notes


# ==========================================================================
# 列表复用起草器：合并表头、占位符、合并填充、跳过空行
# ==========================================================================


@pytest.mark.parametrize("build,sheet,ref,n", [
    (lambda: lab.c08("gap"), "明细", "A1:C6", 1),        # 中间一行空白：blank_rows=skip
    (lambda: lab.c04(), "合并", "A3:F10", 2),            # 两行表头、地区纵向合并：merged_data=fill
])
def test_list_selection_keeps_draft_settings(build, sheet, ref, n):
    b = Book(*build())
    d = b.draft()
    work = full(d.recipe)
    res = selection_edit(work, b.grid, sel(ref, "list", sheet, header_rows=n), extraction=b.dry(work))
    assert res.ok, res.problems
    assert res.recipe_sha256_after == sha(work), "框选的列表规则与起草器相同"
    blk = apply_patch(work, res.ops)["sheets"][0]["blocks"][0]
    if sheet == "明细":
        assert blk["rows"]["blank_rows"] == "skip"
        assert any("中间的空行跳过" in s for s in res.summary)
    else:
        assert blk["merged_data"] == "fill" and blk["header_rows"] == 2


def test_list_selection_keeps_placeholders():
    wb = Workbook()
    ws = wb.active
    ws.title = "名单"
    for r, row in enumerate([("地区", "销量"), ("分区甲", 3), ("分区乙", "-"), ("分区丙", 7)], start=1):
        for c, v in enumerate(row, start=1):
            ws.cell(r, c, v)
    b = Book(save(wb), "ph.xlsx")
    res = selection_edit(None, b.grid, sel("A1:B4", "list", "名单"), extraction=None)
    assert res.ok
    blk = apply_patch({}, res.ops)["sheets"][0]["blocks"][0]
    assert blk["values"]["placeholders"] == [{"text": "-", "meaning": "无数据"}]
    assert blk["columns"][1]["type"] == "INTEGER"


# ==========================================================================
# after_title、表名
# ==========================================================================


def _stacked(title2: str | None) -> Book:
    """上下两张表头相同的列表：上面一张有标题「销售汇总」，下面一张上方是含年月的说明（和可选的第二行标题）。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "汇总"
    ws["A1"] = "销售汇总"
    rows = [(2, ("地区", "销量")), (3, ("分区甲", 1)), (4, ("分区乙", 2)),
            (8, ("地区", "销量")), (9, ("分区丙", 3)), (10, ("分区丁", 4))]
    for r, row in rows:
        for c, v in enumerate(row, start=1):
            ws.cell(r, c, v)
    ws["A7"] = "2026年9月 补充"
    if title2:
        ws["A6"] = title2
    return Book(save(wb), "stacked.xlsx")


def test_after_title_skips_cells_with_digits():
    b = _stacked("补充明细")
    res = selection_edit(None, b.grid, sel("A8:B10", "list", "汇总"), extraction=None)
    assert res.ok, res.problems
    blk = apply_patch({}, res.ops)["sheets"][0]["blocks"][0]
    assert blk["after_title"] == "补充明细"
    assert ("after_title", "补充明细") in [(a.kind, a.text) for a in res.anchors]
    assert validate_recipe(apply_patch({}, res.ops))[1] == []


def test_after_title_ambiguous_when_only_titles_with_digits():
    b = _stacked(None)
    res = selection_edit(None, b.grid, sel("A8:B10", "list", "汇总"), extraction=None)
    assert codes(res) == ["selection_anchor_ambiguous"]
    assert "含年月或数字" in res.problems[0].message


def test_client_table_name_is_validated_and_checked_for_collisions(c01):
    b, work, ext = c01
    bad = selection_edit(None, b.grid, sel("C5:F8", "list", "月报", table="1号表"), extraction=None)
    assert codes(bad) == ["selection_name_invalid"]
    taken = selection_edit(None, b.grid, sel("C5:F8", "list", "月报", table="旧销售"), extraction=None,
                           retired_names={"旧销售": ["销量"]})
    assert codes(taken) == ["selection_name_taken"]
    taken2 = selection_edit(None, b.grid, sel("C5:F8", "list", "月报", table="ＳＡＬＥＳ"), extraction=None,
                            retired_names={"sales": []})
    assert codes(taken2) in (["selection_name_invalid"], ["selection_name_taken"])
    ok = selection_edit(work, b.grid, sel("C5:F8", "list", "月报", table="销售月报"), extraction=ext)
    assert ok.ok
    new = apply_patch(work, ok.ops)
    assert [t["name"] for t in new["tables"]] == ["销售月报", "销售月报_表内合计"]
    assert validate_recipe(new)[1] == []


def test_new_column_name_follows_retired_spelling(c01):
    b, work, ext = c01
    res = selection_edit(work, b.grid, sel("C5:F8", "list", "月报"), extraction=ext,
                         retired_names={"月报": ["销量"]})
    assert res.ok and res.notes == []
    wb = Workbook()
    ws = wb.active
    ws.title = "表"
    for r, row in enumerate([("地区", "pm2_5"), ("分区甲", 1), ("分区乙", 2)], start=1):
        for c, v in enumerate(row, start=1):
            ws.cell(r, c, v)
    b2 = Book(save(wb), "pm.xlsx")
    first = selection_edit(None, b2.grid, sel("A1:B3", "list", "表", table="监测"), extraction=None)
    rec = full(apply_patch({}, first.ops))
    assert rec["sheets"][0]["blocks"][0]["columns"][1]["name"] == "pm2_5"
    # 当前版本里「监测」表的「PM2_5」已退役：重新框选这张表时，新列沿用退役列的写法（SQLite 列名不分大小写）
    res2 = selection_edit(rec, b2.grid, sel("A1:B3", "list", "表"), extraction=b2.dry(rec),
                          retired_names={"监测": ["PM2_5"]})
    assert res2.ok and apply_patch(rec, res2.ops)["sheets"][0]["blocks"][0]["columns"][1]["name"] == "PM2_5"
    assert any("PM2_5" in n for n in res2.notes)


# ==========================================================================
# 重叠
# ==========================================================================


def test_overlap_with_two_blocks():
    b = Book(*lab.c08("side"))
    d = b.draft()
    work = full(d.recipe)
    assert len(work["sheets"][0]["blocks"]) == 2
    res = selection_edit(work, b.grid, sel("A1:E3", "list", "并排"), extraction=b.dry(work))
    assert codes(res) == ["selection_overlap"]
    # 只框一张：替换它，沿用块 id
    one = selection_edit(work, b.grid, sel("D1:E3", "list", "并排"), extraction=b.dry(work))
    assert one.ok and one.key == "list:列表2" and one.recipe_sha256_after == sha(work)


def test_non_overlapping_box_adds_a_block(c01):
    b, work, ext = c01
    wb = Workbook()
    ws = wb.active
    ws.title = "月报"
    for r, row in enumerate([("地区", "销量"), ("分区甲", 1), ("分区乙", 2)], start=1):
        for c, v in enumerate(row, start=1):
            ws.cell(r, c, v)
    for r, row in enumerate([("部门", "费用"), ("部门甲", 5), ("部门乙", 7)], start=6):
        for c, v in enumerate(row, start=1):
            ws.cell(r, c, v)
    bk = Book(save(wb), "two.xlsx")
    first = selection_edit(None, bk.grid, sel("A1:B3", "list", "月报"), extraction=None)
    rec = full(apply_patch({}, first.ops))
    second = selection_edit(rec, bk.grid, sel("A6:B8", "list", "月报"), extraction=bk.dry(rec))
    assert second.ok and second.ops[0] == {"op": "add", "path": "/sheets/0/blocks/-", "value": second.ops[0]["value"]}
    new = apply_patch(rec, second.ops)
    assert [blk["id"] for blk in new["sheets"][0]["blocks"]] == ["列表1", "列表2"]
    assert validate_recipe(new)[1] == []


# ==========================================================================
# 交叉表（整块）
# ==========================================================================


def test_crosstab_from_rule_draft(flow_d01):
    b = flow_d01
    res = selection_edit(None, b.grid, sel("C5:H12", "crosstab"), extraction=None, draft_fn=b.draft)
    assert res.ok, res.problems
    assert res.expected == {"axis": "C4:AF4", "labels": "B5:B30", "values": "C5:AF30"}
    assert res.anchors[0].kind == "axis"
    new = apply_patch({}, res.ops)
    assert validate_recipe(new)[1] == []
    rc = replay_compare(res, b.dry(new))
    assert rc.match, rc.diffs
    # 已有交叉表：框在它上面替换它。起草器的名字按来源对齐到现行配方，版式没变时就是原来的配方
    rep = selection_edit(FLOW, b.grid, sel("C5:H12", "crosstab"), extraction=b.dry(FLOW), draft_fn=b.draft)
    assert rep.ok and rep.key == "crosstab:交叉表"
    assert rep.recipe_sha256_after == sha(FLOW)
    # 现行配方里另有一条关系占了 R2：起草的关系与它重号时顺延
    extra = copy.deepcopy(FLOW)
    extra["relations"][1]["id"] = "R5"
    rep2 = selection_edit(extra, b.grid, sel("C5:H12", "crosstab"), extraction=b.dry(extra), draft_fn=b.draft)
    new2 = apply_patch(extra, rep2.ops)
    assert validate_recipe(new2)[1] == [] and len(new2["sheets"][0]["blocks"]) == 1
    assert sorted(r["id"] for r in new2["relations"]) == ["R1", "R5"]


def test_crosstab_on_new_sheet_renumbers_clashing_relations(flow_d01):
    b = flow_d01
    other = copy.deepcopy(FLOW)
    other["sheets"][0]["match"]["name"] = "另一张表"
    other["sheets"][0]["id"] = "s9"
    res = selection_edit(other, b.grid, sel("C5:H12", "crosstab"), extraction=None, draft_fn=b.draft)
    assert res.ok, res.problems
    new = apply_patch(other, res.ops)
    assert [s["id"] for s in new["sheets"]] == ["s9", "s2"]
    assert [r["id"] for r in new["relations"]] == ["R1", "R2", "R3", "R4"]
    names = [t["name"] for t in new["tables"]]
    assert len({collide_key(n) for n in names}) == len(names), names
    assert validate_recipe(new)[1] == []


def test_selection_options_text_never_enters_the_patch(c01):
    b, work, ext = c01
    res = selection_edit(work, b.grid, sel("C5:F8", "list", "月报", note="随便写的说明", segment="恶意"),
                         extraction=ext)
    assert res.ok and "随便写的说明" not in json.dumps(res.ops, ensure_ascii=False)


def test_crosstab_needs_enough_overlap(flow_d01):
    b = flow_d01
    res = selection_edit(None, b.grid, sel("AG40:AH45", "crosstab"), extraction=None, draft_fn=b.draft)
    assert codes(res) == ["selection_no_crosstab"]
    lst = Book(*lab.c01("literal"))
    res2 = selection_edit(None, lst.grid, sel("C5:F8", "crosstab", "月报"), extraction=None, draft_fn=lst.draft)
    assert codes(res2) == ["selection_no_crosstab"]


# ==========================================================================
# 分段、合计行、分段标题
# ==========================================================================


def test_segment_measures_on_unclaimed_rows():
    b, ext = flow_book("D21")
    res = selection_edit(FLOW, b.grid, sel("B32:AF32", "segment", role="measures", table="补录"), extraction=ext)
    assert res.ok, res.problems
    assert res.key == "segment:补录" and res.expected["labels"] == "B32"
    new = apply_patch(FLOW, res.ops)
    seg = new["sheets"][0]["blocks"][0]["segments"][-1]
    assert seg["role"] == "measures" and seg["locate"]["by"] == "labels" and seg["labels"]["expect"] == ["补录（人次）"]
    assert seg["measures"] == {"补录（人次）": "补录"}
    assert new["tables"][-1] == {"name": "补录", "grain": ["日期"], "kind": "data", "units": {"补录": "人次"}, "note": ""}
    assert validate_recipe(new)[1] == []
    rc = replay_compare(res, b.dry(new))
    assert rc.match, rc.diffs


def test_segment_problems(flow_d01):
    b, ext = flow_book("D21")
    with pytest.raises(EditRequestError):
        selection_edit(FLOW, b.grid, sel("B32:AF32", "segment"), extraction=ext)
    assert codes(selection_edit(FLOW, b.grid, sel("C32:AF32", "segment", role="measures"), extraction=ext)) == [
        "selection_not_in_block"]
    assert codes(selection_edit(FLOW, b.grid, sel("B31:AF32", "segment", role="measures"), extraction=ext)) == [
        "selection_label_blank"]
    assert codes(selection_edit(FLOW, b.grid, sel("B6:AF7", "segment", role="measures"), extraction=ext)) == [
        "selection_overlap"]
    taken = selection_edit(FLOW, b.grid, sel("B32:AF32", "segment", role="measures", table="日客流"), extraction=ext)
    assert codes(taken) == ["selection_name_taken"]


def _early_grid(b: Book, title: str = "早间时段客流（人次）", labels: tuple[str, ...] = ("5-6", "6-7")):
    """D21 的网格上，第 32 行改成分段标题、下面两行是新的时段（框选为维度分段用）。"""
    grid = copy.deepcopy(b.grid)
    grid.cells[(32, 2)] = GridCell(title)
    for c in range(3, 5):
        grid.cells.pop((32, c), None)
    for k, h in enumerate(labels):
        grid.cells[(33 + k, 2)] = GridCell(h)
        for c in range(3, 33):
            grid.cells[(33 + k, c)] = GridCell(10 + c)
    grid.__dict__.pop("_rows", None)
    return grid


def test_segment_dimension_with_section_title_joins_the_matching_table():
    """起草器第 5 条：解析器相同（hour_range）、标签不相交、单位相同的维度分段写进同一张表「时段客流」，常量列
    「时段类别」照第 8b 条取分段标题的候选词「早间」，与「日间」「夜间」不同。不新建表。"""
    b, ext = flow_book("D21")
    grid = _early_grid(b)
    res = selection_edit(FLOW, grid, sel("B33:AF34", "segment", role="dimension"), extraction=ext)
    assert res.ok, res.problems
    assert len(res.ops) == 1, "并入已有的表，不新增表定义"
    seg = res.ops[0]["value"]
    assert seg == {"id": "早间", "role": "dimension", "table": "时段客流",
                   "locate": {"by": "section_title", "title": "早间时段客流（人次）", "segment": None},
                   "labels": {"expect": ["5-6", "6-7"]},
                   "dim": {"name": "时段", "parser": "hour_range", "derive": {"起始小时": "start", "结束小时": "end"}},
                   "value": "客流", "const": {"时段类别": {"pick": "早间"}}, "stop_parser": None}
    assert res.ops[0]["path"] == "/sheets/0/blocks/0/segments/4", "插在已有分段之后"
    assert res.expected["title"] == "B32" and res.notes == []
    assert any("并入表「时段客流」" in x for x in res.summary) and any("「时段类别」取「早间」" in x for x in res.summary)
    assert validate_recipe(apply_patch(FLOW, res.ops))[1] == []


def test_segment_dimension_merged_replays_on_the_file():
    """并入「时段客流」后按真文件干跑：新时段的 60 行写进同一张表，没有问题，重放比对一致。"""
    raw, fn = flow_workbook(SEP, 30, variant="D21")
    wb = load_workbook(io.BytesIO(raw), data_only=True)
    ws = wb.active
    ws.cell(32, 2).value = "早间时段客流（人次）"
    for c in range(3, 33):
        ws.cell(32, c).value = None
    for k, h in enumerate(("5-6", "6-7")):
        ws.cell(33 + k, 2).value = h
        for c in range(3, 33):
            ws.cell(33 + k, c).value = 10 + c + k
    b = Book(save(wb), fn)
    res = selection_edit(FLOW, b.grid, sel("B33:AF34", "segment", role="dimension"), extraction=b.dry(FLOW))
    assert res.ok and res.ops[0]["value"]["table"] == "时段客流"
    new = apply_patch(FLOW, res.ops)
    ext = b.dry(new)
    assert ext.ok and ext.problems == []
    assert ext.expected_rows["时段客流"] == b.dry(FLOW).expected_rows["时段客流"] + 60
    assert replay_compare(res, ext).match


def test_segment_dimension_new_table_says_why_it_was_not_merged():
    """有同类的表却并不进去时新建表，并在说明里写明没有并入、为什么：客户端指定了新表名；那张表的分段没有常量列
    （并入要给现有分段加列）；分段标题里选不出与已有取值不同的词。标签和已有分段重叠的不算同类，不提。"""
    b, ext = flow_book("D21")
    named = selection_edit(FLOW, _early_grid(b), sel("B33:AF34", "segment", role="dimension", table="早间客流"),
                           extraction=ext)
    assert named.ok and named.ops[1]["value"]["name"] == "早间客流"
    assert named.notes == ["按指定的表名新建了表，没有并入同类的表「时段客流」"]
    # 那张表只有一个分段、没有常量列
    single = copy.deepcopy(FLOW)
    segs = single["sheets"][0]["blocks"][0]["segments"]
    segs[1]["const"] = {}
    del segs[2:]
    alone = selection_edit(single, _early_grid(b), sel("B33:AF34", "segment", role="dimension"), extraction=ext)
    assert alone.ok and alone.ops[0]["value"]["table"] == "早间时段客流"
    assert "没有并入同类的表「时段客流」" in alone.notes[0] and "没有用来区分分段的常量列" in alone.notes[0]
    # 标题里能区分的词与已有取值相同
    same = selection_edit(FLOW, _early_grid(b, title="日间时段客流（人次）"),
                          sel("B33:AF34", "segment", role="dimension"), extraction=ext)
    assert same.ok and same.ops[0]["value"]["table"] != "时段客流"
    assert "分段标题里没有能与已有分段区分的词" in same.notes[0]
    # 标签与日间重叠：不是同类，新建表、不提并入
    overlap = selection_edit(FLOW, _early_grid(b, labels=("7-8", "8-9")), sel("B33:AF34", "segment", role="dimension"),
                             extraction=ext)
    assert overlap.ok and overlap.ops[0]["value"]["table"] == "早间时段客流" and overlap.notes == []


def test_segment_dimension_text_label_too_long():
    b, ext = flow_book("D21")
    grid = copy.deepcopy(b.grid)
    grid.cells[(32, 2)] = GridCell("补" * 41)
    grid.__dict__.pop("_rows", None)
    res = selection_edit(FLOW, grid, sel("B32:AF32", "segment", role="dimension"), extraction=ext)
    assert codes(res) == ["selection_label_unparsed"]


def test_derived_d18():
    b, ext = flow_book("D18")
    res = selection_edit(FLOW, b.grid, sel("B21:AF21", "derived", keep=True), extraction=ext)
    assert res.ok and res.key == "derived:日间合计"
    new = apply_patch(FLOW, res.ops)
    assert validate_recipe(new, base=FLOW)[1] == []
    ext2 = b.dry(new)
    assert ext2.ok
    assert replay_compare(res, ext2).match


def test_derived_problems():
    b, ext = flow_book("D18")
    assert codes(selection_edit(FLOW, b.grid, sel("B11:AF11", "derived"), extraction=ext)) == ["selection_label_unparsed"]
    b2, ext2 = flow_book("D21")
    res = selection_edit(FLOW, b2.grid, sel("B32:AF32", "derived"), extraction=ext2)
    assert codes(res) == ["selection_label_unparsed"]
    grid = copy.deepcopy(b2.grid)
    grid.cells[(32, 2)] = GridCell("7-12 时合计")
    grid.__dict__.pop("_rows", None)
    assert codes(selection_edit(FLOW, grid, sel("B32:AF32", "derived"), extraction=ext2)) == [
        "selection_not_after_segment"]


def test_section_title_d09_keeps_const():
    b, ext = flow_book("D09")
    res = selection_edit(FLOW, b.grid, sel("B9", "section_title", segment="日间"), extraction=ext)
    assert res.ok and res.key == "section_title:日间" and res.breaking is False
    assert res.ops == [{"op": "replace", "path": "/sheets/0/blocks/0/segments/1/locate",
                        "value": {"by": "section_title", "title": "日间分时段客流（人次）", "segment": None}}]
    new = apply_patch(FLOW, res.ops)
    assert validate_recipe(new, base=FLOW)[1] == []
    assert b.dry(new).ok


def test_section_title_does_not_guess_a_new_word_when_old_pick_is_gone():
    """新标题里没有原来的常量取值：常量列整列换成哪个词只能由人选（4.1「不猜」）。框选不替人挑第一个候选词，
    报 selection_label_unparsed，指到修复按钮（它把候选词逐个列成选项）和配方面板。"""
    b, ext = flow_book("D09")
    grid = copy.deepcopy(b.grid)
    grid.cells[(9, 2)] = GridCell("白天时段客流（人次）")
    grid.__dict__.pop("_rows", None)
    res = selection_edit(FLOW, grid, sel("B9", "section_title", segment="日间"), extraction=ext)
    assert not res.ok and res.ops == [] and codes(res) == ["selection_label_unparsed"]
    msg = res.problems[0].message
    assert "常量列「时段类别」现在取「日间」" in msg and "修复按钮" in msg and "配方面板" in msg
    assert res.problems[0].cells == [f"{SHEET}!B9"]
    with pytest.raises(EditRequestError):
        selection_edit(FLOW, grid, sel("B9", "section_title", segment="夜间合计"), extraction=ext)
    assert codes(selection_edit(FLOW, grid, sel("B9:B10", "section_title", segment="日间"), extraction=ext)) == [
        "selection_not_in_block"]


# ==========================================================================
# 忽略
# ==========================================================================


def test_ignore_rows_d21():
    b, ext = flow_book("D21")
    res = selection_edit(FLOW, b.grid, sel("B32:D32", "ignore_rows", reason="补录行不导入"), extraction=ext)
    assert res.ok and res.key == "ignore_rows:交叉表"
    assert res.ops == [{"op": "add", "path": "/sheets/0/blocks/0/ignore_rows/-",
                        "value": {"label": "补录（人次）", "reason": "补录行不导入"}}]
    assert validate_recipe(apply_patch(FLOW, res.ops))[1] == []
    with pytest.raises(EditRequestError) as e:
        selection_edit(FLOW, b.grid, sel("B32:D32", "ignore_rows"), extraction=ext)
    assert e.value.code == "reason_required"


def test_ignore_rows_with_digits():
    b, ext = flow_book("D21")
    grid = copy.deepcopy(b.grid)
    grid.cells[(32, 2)] = GridCell("注：9月补录")
    grid.__dict__.pop("_rows", None)
    res = selection_edit(FLOW, grid, sel("B32:D32", "ignore_rows", reason="x"), extraction=ext)
    assert codes(res) == ["selection_anchor_has_digits"]


def test_ignore_columns_d16():
    b, ext = flow_book("D16")
    res = selection_edit(FLOW, b.grid, sel("AG4:AG30", "ignore_columns", reason="右侧合计列"), extraction=ext)
    assert res.ok and res.key == "ignore_columns:交叉表"
    assert res.ops == [{"op": "add", "path": "/sheets/0/blocks/0/ignore_columns/-",
                        "value": {"header": "合计", "reason": "右侧合计列"}}]
    assert validate_recipe(apply_patch(FLOW, res.ops))[1] == []
    assert codes(selection_edit(FLOW, b.grid, sel("AE4:AE9", "ignore_columns", reason="x"), extraction=ext)) == [
        "selection_not_in_block"]


def test_ignore_columns_header_problems():
    b, ext = flow_book("D16")
    grid = copy.deepcopy(b.grid)
    grid.cells[(4, 34)] = GridCell(123)
    grid.cells[(4, 35)] = GridCell("第2组合计")
    grid.__dict__.pop("_rows", None)
    assert codes(selection_edit(FLOW, grid, sel("AH4", "ignore_columns", reason="x"), extraction=ext)) == [
        "selection_header_not_text"]
    assert codes(selection_edit(FLOW, grid, sel("AI4", "ignore_columns", reason="x"), extraction=ext)) == [
        "selection_anchor_has_digits"]
    assert codes(selection_edit(FLOW, grid, sel("AJ4", "ignore_columns", reason="x"), extraction=ext)) == [
        "selection_header_blank"]


def test_ignore_outside():
    b, ext = flow_book("D01")
    grid = copy.deepcopy(b.grid)
    grid.cells[(33, 2)] = GridCell("补录说明")
    grid.cells[(33, 3)] = GridCell(123)
    grid.cells[(34, 3)] = GridCell(456)
    grid.cells[(35, 2)] = GridCell("9月补录")
    grid.cells[(35, 3)] = GridCell(789)
    grid.__dict__.pop("_rows", None)
    res = selection_edit(FLOW, grid, sel("B33:C33", "ignore_outside", reason="说明行"), extraction=ext)
    assert res.ok and res.key == "ignore_outside:s1"
    assert res.ops == [{"op": "add", "path": "/sheets/0/ignore_outside/-",
                        "value": {"anchor": "补录说明", "reason": "说明行"}}]
    assert validate_recipe(apply_patch(FLOW, res.ops))[1] == []
    assert codes(selection_edit(FLOW, grid, sel("B34:C34", "ignore_outside", reason="x"), extraction=ext)) == [
        "selection_no_anchor"]
    assert codes(selection_edit(FLOW, grid, sel("B35:C35", "ignore_outside", reason="x"), extraction=ext)) == [
        "selection_anchor_has_digits"]


def test_selection_on_sheet_without_blocks():
    b, _ext = flow_book("D21")
    res = selection_edit(None, b.grid, sel("B32:D32", "ignore_rows", reason="x"), extraction=None)
    assert codes(res) == ["selection_not_in_block"]


# ==========================================================================
# 重放比对（4.5）
# ==========================================================================


def _edit(key: str, expected: dict, block: str | None = "列表1") -> EditResult:
    return EditResult(ok=True, kind="selection", key=key, title="", summary=[], ops=[], anchors=[], notes=[],
                      problems=[], expected=expected, block=block)


def test_replay_compare_reports_extra_rows_in_words():
    ext = Extraction(ok=True, regions=[
        RegionMark("月报", "col_header", "C5:F5", "列表1"), RegionMark("月报", "value", "C6:F10", "列表1"),
        RegionMark("月报", "value", "H6:I20", "列表2"), RegionMark("月报", "outside_text", "C1:C3", None)])
    rc = replay_compare(_edit("list:列表1", {"header": "C5:F5", "data": "C6:F8", "total": None}), ext)
    assert not rc.match and rc.actual == {"header": "C5:F5", "data": "C6:F10", "total": None}
    assert rc.diffs == ["重放的数据区到第 10 行，比框多 2 行（第 9–10 行）"]


def test_replay_compare_missing_total_and_window():
    ext = Extraction(ok=True, regions=[
        RegionMark("月报", "col_header", "C5:F5", "列表1"), RegionMark("月报", "value", "C6:F900", "列表1")])
    rc = replay_compare(_edit("list:列表1", {"header": "C5:F5", "data": "C6:F900", "total": "C901:F901"}), ext,
                        window_rows=500)
    # 部分干跑：期望和实际都裁到前 500 行再比，窗口外的合计行不算缺
    assert rc.match and rc.window_rows == 500 and rc.expected["data"] == "C6:F500"
    rc2 = replay_compare(_edit("list:列表1", {"header": "C5:F5", "data": "C6:F900", "total": "C901:F901"}), ext)
    assert not rc2.match and "重放时没有认出合计行（框里是 C901:F901）" in rc2.diffs


def test_replay_compare_partial_kinds_check_coverage():
    ext = Extraction(ok=True, regions=[RegionMark(SHEET, "ignored", "B32:D32", "交叉表"),
                                       RegionMark(SHEET, "row_label", "B5:B30", "交叉表")])
    ok = replay_compare(_edit("ignore_rows:交叉表", {"ignored": "B32"}, "交叉表"), ext)
    assert ok.match and ok.actual == {"ignored": "B32:D32"}
    bad = replay_compare(_edit("ignore_rows:交叉表", {"ignored": "B33"}, "交叉表"), ext)
    assert not bad.match and bad.diffs


def test_edit_result_keys_do_not_leak_coordinates_into_recipe(c01):
    """坐标不进配方：换算出的补丁里没有任何 A1 写法（表头上方的标题、锚点都是文字）。"""
    import re
    b, work, ext = c01
    res = selection_edit(work, b.grid, sel("C5:F8", "list", "月报"), extraction=ext)
    blob = json.dumps(res.ops, ensure_ascii=False)
    assert not re.search(r"(?<![A-Za-z])[A-Z]{1,3}[0-9]{1,5}(?![0-9])", blob.replace("agentlab-recipe/2", ""))
    assert collide_key("月报") in {collide_key(t["name"]) for t in apply_patch(work, res.ops)["tables"]}
