"""配方执行器：交叉表（P2-SPEC 4.2、4.3、4.5、4.6、4.7）。

夹具是测试里现造的「结构仿照客流表」（坐标同 9.1，名字全是假名、数字随机），按需要改几格造出各种漂移；
配方用契约里的参考配方 flow_recipe.json。不依赖 WP-1 的夹具。
"""
from __future__ import annotations

import copy
import datetime as dt
import hashlib
import io
import json
import os
import random
import re
import sqlite3
import zipfile
from collections.abc import Callable
from pathlib import Path
from typing import Any

import pytest
from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from app.data import recipe_engine, xlsx_scan
from app.data.recipe_types import Extraction, PeriodInput, Recipe

FLOW = json.loads((Path(__file__).parent / "fixtures" / "recipes" / "flow_recipe.json").read_text(encoding="utf-8"))
SHEET = "客流汇总"


# --------------------------------------------------------------------------
# 夹具
# --------------------------------------------------------------------------


def inject_cache(raw: bytes, values: dict[str, Any], *, keep_full_calc: bool = False,
                 sheet_part: str = "xl/worksheets/sheet1.xml") -> bytes:
    """openpyxl 写公式不带缓存值：在 zip 层给指定格补上 <v>；默认同时去掉 workbook.xml 的 fullCalcOnLoad。"""
    zin = zipfile.ZipFile(io.BytesIO(raw))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        for it in zin.infolist():
            data = zin.read(it.filename)
            if it.filename == sheet_part and values:
                t = data.decode("utf-8")
                for ref, v in values.items():
                    t, n = re.subn(rf'(<c r="{ref}"[^>]*>)<f>(.*?)</f>(?:<v\s*/>|<v></v>)?',
                                   lambda m, v=v: f"{m.group(1)}<f>{m.group(2)}</f><v>{v}</v>", t)
                    assert n == 1, ref
                data = t.encode("utf-8")
            if it.filename == "xl/workbook.xml" and not keep_full_calc:
                data = re.sub(rb'\s*fullCalcOnLoad="1"', b"", data)
            zout.writestr(it, data)
    return out.getvalue()


def patch_part(raw: bytes, part: str, fn: Callable[[str], str]) -> bytes:
    zin = zipfile.ZipFile(io.BytesIO(raw))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        for it in zin.infolist():
            data = zin.read(it.filename)
            if it.filename == part:
                data = fn(data.decode("utf-8")).encode("utf-8")
            zout.writestr(it, data)
    return out.getvalue()


def flow_book(start: dt.date = dt.date(2026, 8, 1), days: int = 31, seed: int = 0, *,
              day_hours: tuple[int, ...] = tuple(range(7, 18)), night_first: bool = False, blank_row: bool = True,
              edit: Callable[[Any, dict[str, Any]], None] | None = None, uncached: bool = False,
              keep_full_calc: bool = False, filename: str | None = None, period: str | None = None
              ) -> tuple[bytes, str]:
    """「结构仿照客流表」：B2 统计期、B3 标题、第 4 行日期、B5:B7 指标、空行、日间段（第一行整行「·」）、
    夜间段和三行合计公式（补了缓存值）。edit(ws, info) 在保存前改格子；info["cache"] 是公式缓存值，
    info["rows"] 记各段所在的行。"""
    rnd = random.Random(seed)
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET
    end = start + dt.timedelta(days=days - 1)
    last = get_column_letter(2 + days)
    ws["B2"] = period if period is not None else (
        f"统计时间范围：{start.year}年{start.month}月{start.day}日至{end.year}年{end.month}月{end.day}日")
    ws.merge_cells(f"B2:{last}2")
    ws["B3"] = "客流汇总表"
    ws.merge_cells(f"B3:{last}3")
    for j in range(days):
        d = start + dt.timedelta(days=j)
        ws.cell(4, 3 + j, f"{d.month}月{d.day}日")
    info: dict[str, Any] = {"cache": {}, "rows": {}, "last": last, "days": days}

    def measures(r: int) -> int:
        for k, label in enumerate(("全日客流（人次）", "分区甲（人次）", "分区乙（人次）")):
            ws.cell(r + k, 2, label)
        for j in range(days):
            a, b = rnd.randint(3000, 9000), rnd.randint(2000, 8000)
            ws.cell(r, 3 + j, a + b)
            ws.cell(r + 1, 3 + j, a)
            ws.cell(r + 2, 3 + j, b)
        info["rows"]["日客流"] = [r, r + 1, r + 2]
        return r + 3

    def day(r: int) -> int:
        ws.cell(r, 2, "日间时段客流（人次）")
        ws.merge_cells(f"B{r}:{last}{r}")
        info["rows"]["日间标题"] = r
        rows = []
        for i, h in enumerate(day_hours):
            rr = r + 1 + i
            ws.cell(rr, 2, f"{h}-{h + 1}")
            for j in range(days):
                ws.cell(rr, 3 + j, "·" if h == 7 else rnd.randint(50, 900))
            rows.append(rr)
        info["rows"]["日间"] = rows
        return r + 1 + len(day_hours)

    def night(r: int) -> int:
        ws.cell(r, 2, "夜间时段客流（人次）")
        ws.merge_cells(f"B{r}:{last}{r}")
        info["rows"]["夜间标题"] = r
        vals: dict[tuple[int, int], int] = {}
        rows = []
        for i, h in enumerate(range(18, 24)):
            rr = r + 1 + i
            ws.cell(rr, 2, f"{h}-{h + 1}")
            for j in range(days):
                v = rnd.randint(50, 900)
                ws.cell(rr, 3 + j, v)
                vals[(rr, j)] = v
            rows.append(rr)
        info["rows"]["夜间"] = rows
        totals = []
        for k, (label, a, b) in enumerate((("18-22 时合计", 0, 3), ("22-24 时合计", 4, 5), ("18-24 时合计", 0, 5))):
            rr = r + 7 + k
            ws.cell(rr, 2, label)
            for j in range(days):
                col = get_column_letter(3 + j)
                ws.cell(rr, 3 + j, f"=SUM({col}{rows[a]}:{col}{rows[b]})")
                info["cache"][f"{col}{rr}"] = sum(vals[(rows[x], j)] for x in range(a, b + 1))
            totals.append(rr)
        info["rows"]["合计"] = totals
        return r + 10

    r = measures(5)
    if blank_row:
        ws.merge_cells(f"B{r}:{last}{r}")
        r += 1
    for block in ((night, day) if night_first else (day, night)):
        r = block(r)
    info["bottom"] = r - 1
    if edit is not None:
        edit(ws, info)
    buf = io.BytesIO()
    wb.save(buf)
    raw = inject_cache(buf.getvalue(), {} if uncached else info["cache"], keep_full_calc=keep_full_calc)
    return raw, filename or f"月报导出_{start.isoformat()}_{end.isoformat()}.xlsx"


def recipe_dict(**patch: Any) -> dict[str, Any]:
    data = copy.deepcopy(FLOW)
    for path, value in patch.items():
        cur: Any = data
        keys = path.split("__")
        for k in keys[:-1]:
            cur = cur[int(k)] if k.isdigit() else cur[k]
        last = keys[-1]
        if isinstance(cur, list):
            cur[int(last)] = value
        else:
            cur[last] = value
    return data


def run(book: tuple[bytes, str], recipe: dict[str, Any] | None = None, db: str | None = None,
        scan: xlsx_scan.WorkbookScan | None = None, **kw: Any) -> Extraction:
    raw, filename = book
    return recipe_engine.execute(Recipe.model_validate(recipe or FLOW), raw, filename,
                                 scan or xlsx_scan.scan(raw), db, **kw)


def codes(ex: Extraction) -> set[str]:
    return {p.code for p in ex.problems}


def problem(ex: Extraction, code: str):
    found = [p for p in ex.problems if p.code == code]
    assert found, f"没有 {code}：{[(p.code, p.message) for p in ex.problems]}"
    return found[0]


def set_value(ws: Any, ref: str, value: Any) -> None:
    ws[ref] = value


# --------------------------------------------------------------------------
# 参考布局（9.1 的期望）
# --------------------------------------------------------------------------


def test_reference_layout_matches_receipt(tmp_path):
    db = str(tmp_path / "flow.db")
    ex = run(flow_book(), db=db)
    assert ex.ok and not ex.problems and not ex.partial
    (ledger,) = ex.ledger
    assert ledger.nonempty_scan == ledger.nonempty_read == 771 and ledger.unclaimed == 0
    assert ledger.roles == {"value": 620, "derived_value": 93, "derived_label": 3, "col_header": 31,
                            "row_label": 20, "section_title": 2, "context": 1, "outside_text": 1}
    assert [(t.name, t.rows, t.kind, t.grain) for t in ex.tables] == [
        ("日客流", 31, "data", ["日期"]), ("时段客流", 527, "data", ["日期", "时段"]),
        ("时段客流_表内合计", 93, "reported_total", ["日期", "合计项"])]
    assert ex.expected_rows == {"日客流": 31, "时段客流": 527, "时段客流_表内合计": 93}
    assert ex.placeholders == {"·": 31}
    assert [(c.id, c.status) for c in ex.context_checks] == [("C1", "passed"), ("C2", "passed")]
    assert ex.canonicalized == []
    assert [(o.cell, o.kind, o.period_source) for o in ex.outside_text] == [(f"{SHEET}!B3", "text", False)]
    assert ex.period.source == "cells" and ex.period.texts == {
        f"{SHEET}!B2": "统计时间范围：2026年8月1日至2026年8月31日"} and ex.period.annotated == {}
    assert ex.block_order == {SHEET: ["日客流", "日间", "夜间", "夜间合计"]}
    assert ex.derived_form == {"夜间合计": "formula"}
    assert [(a.row, a.first, a.last, a.count, a.form) for a in ex.axes] == [(4, "2026-08-01", "2026-08-31", 31, "text")]
    assert ex.sheets.matched == {"s1": SHEET} and ex.sheets.renamed == {} and ex.sheets.other_visible == []
    assert len(ex.derived) == 93 and all(d.ref_problem is None and d.func == "SUM" for d in ex.derived)
    assert ex.lineage["日客流"]["全日客流"] == [[1, SHEET, "C5", 31, "right"]]
    assert ex.lineage["时段客流"]["客流"][:2] == [[1, SHEET, "C10", 31, "right"], [32, SHEET, "C11", 31, "right"]]
    assert set(ex.table_hashes) == {"日客流", "时段客流", "时段客流_表内合计"}
    columns = {t.name: [(c.name, c.type, c.role, c.unit) for c in t.columns] for t in ex.tables}
    assert columns["时段客流"] == [("日期", "TEXT", "axis", None), ("时段", "TEXT", "dim", None),
                                ("时段类别", "TEXT", "const", None), ("起始小时", "INTEGER", "derive", None),
                                ("结束小时", "INTEGER", "derive", None), ("客流", "INTEGER", "value", "人次")]
    assert ex.tables[0].sources == ["s1/交叉表/日客流"]
    assert ex.tables[1].sources == ["s1/交叉表/日间", "s1/交叉表/夜间"]
    conn = sqlite3.connect(db)
    try:
        assert conn.execute("PRAGMA journal_mode").fetchone()[0] == "delete"
        ddl = conn.execute("SELECT sql FROM sqlite_master WHERE name='时段客流'").fetchone()[0]
        assert ddl.endswith("STRICT") and 'PRIMARY KEY ("日期", "时段")' in ddl
        assert conn.execute("SELECT COUNT(*), COUNT(客流) FROM 时段客流 WHERE 时段='7-8'").fetchone() == (31, 0)
        assert conn.execute("SELECT 时段类别, 起始小时, 结束小时 FROM 时段客流 WHERE 时段='18-19' LIMIT 1"
                            ).fetchone() == ("夜间", 18, 19)
        assert conn.execute("SELECT COUNT(*) FROM 日客流 WHERE 全日客流 <> 分区甲 + 分区乙").fetchone() == (0,)
        assert conn.execute("SELECT DISTINCT 合计项 FROM 时段客流_表内合计 ORDER BY rowid").fetchall() == [
            ("18-22时合计",), ("22-24时合计",), ("18-24时合计",)]
    finally:
        conn.close()
    assert not os.path.exists(db + "-wal") and not os.path.exists(db + "-journal")


def test_regions_cover_every_claimed_cell():
    raw, fn = flow_book(days=3)
    ex = run((raw, fn))
    from app.data.xlsx_cells import parse_range
    rects = [(m.role, parse_range(m.ref)) for m in ex.regions]
    assert {m.role for m in ex.regions} == {"context", "outside_text", "col_header", "row_label", "value",
                                            "section_title", "derived_label", "derived_value"}
    def covered(role: str, r: int, c: int) -> bool:
        return any(rl == role and s[0] <= r <= s[2] and s[1] <= c <= s[3] for rl, s in rects)
    assert covered("context", 2, 2) and covered("outside_text", 3, 2) and covered("value", 5, 3)
    assert covered("derived_value", 30, 5) and covered("section_title", 21, 2)


def test_regions_do_not_split_on_blank_edge_days():
    # 首日、末日是空格（blank=null）的行，值的区域仍按整段日期列出，不另起矩形（4.9）
    base = [(m.role, m.ref) for m in run(flow_book(days=3)).regions]
    def blank_edges(ws, info):
        for r in info["rows"]["日间"][1::2]:
            ws.cell(r, 5).value = None
        ws.cell(info["rows"]["日间"][3], 3).value = None
    ex = run(flow_book(days=3, edit=blank_edges), recipe_dict(sheets__0__blocks__0__values__blank="null"))
    assert ex.ok and [(m.role, m.ref) for m in ex.regions] == base
    assert ("value", "C10:E20") in base


# --------------------------------------------------------------------------
# 评审后新增的行为
# --------------------------------------------------------------------------


def test_null_in_a_primary_key_column_is_reported_not_dropped():
    # grain 含指标列、blank=null 时那一格存空值：主键为空的行曾被静默丢掉（ok=True、少一行、N 核对照样通过）
    data = recipe_dict(sheets__0__blocks__0__values__blank="null")
    data["tables"][0]["grain"] = ["日期", "分区甲"]
    ex = run(flow_book(days=3, edit=lambda ws, info: set_value(ws, "D6", None)), data)
    assert not ex.ok and codes(ex) == {"value_blank"}
    p = problem(ex, "value_blank")
    assert p.cells == [f"{SHEET}!D6"] and "主键列「分区甲」" in p.message and "主键" in p.model_message
    # 合计另存表的主键含值列、合计格没有保存值：同样报出来，不静默少存
    data = recipe_dict()
    data["tables"][2]["grain"] = ["日期", "合计项", "客流"]
    ex = run(flow_book(days=3, uncached=True), data)
    assert not ex.ok and codes(ex) == {"value_blank"}
    p = problem(ex, "value_blank")
    assert len(p.cells) == 9 and p.cells[0] == f"{SHEET}!C28" and "有 9 处为空值" in p.message
    # 转换失败的格（已经报过）不再另报主键为空
    data = recipe_dict()
    data["tables"][0]["grain"] = ["日期", "分区甲"]
    ex = run(flow_book(days=3, edit=lambda ws, info: set_value(ws, "D6", "若干")), data)
    assert codes(ex) == {"value_not_number"}


def test_merged_blank_row_is_not_merged_in_values_but_merge_over_values_is():
    ex = run(flow_book(days=3))
    assert "merged_in_values" not in codes(ex)            # B8:E8 空行合并（乙-5）
    def merge_values(ws, info):
        r = info["rows"]["日间"][2]
        ws.merge_cells(f"C{r}:D{r}")
    p = problem(run(flow_book(days=3, edit=merge_values)), "merged_in_values")
    assert p.category == "structure" and p.cells == [f"{SHEET}!C12"]


def _blank_day_row(ws, info):
    r = info["rows"]["日间"][0]
    for j in range(info["days"]):
        ws.cell(r, 3 + j).value = None


def test_all_blank_period_row_is_a_data_row(tmp_path):
    ex = run(flow_book(days=3, edit=_blank_day_row))
    assert codes(ex) == {"value_blank"}                   # 不是 title_not_found / label_missing / row_unclaimed
    assert problem(ex, "value_blank").cells == [f"{SHEET}!C10", f"{SHEET}!D10", f"{SHEET}!E10"]
    db = str(tmp_path / "null.db")
    ex = run(flow_book(days=3, edit=_blank_day_row),
             recipe_dict(sheets__0__blocks__0__values__blank="null"), db=db)
    assert ex.ok
    conn = sqlite3.connect(db)
    assert conn.execute("SELECT COUNT(*), COUNT(客流) FROM 时段客流 WHERE 时段='7-8'").fetchone() == (3, 0)
    conn.close()


def test_annotated_period_cells_go_to_outside_text_pure_ones_do_not():
    ex = run(flow_book(days=3, edit=lambda ws, info: set_value(ws, "B3", "客流汇总表（2026年8月）")))
    assert ex.ok and ex.period.annotated == {f"{SHEET}!B3": "客流汇总表"}
    assert ex.period.cells == [f"{SHEET}!B2", f"{SHEET}!B3"]
    assert [(o.cell, o.kind, o.period_source) for o in ex.outside_text] == [(f"{SHEET}!B3", "text_digits", True)]
    assert ex.context_checks[0].status == "passed" and ex.context_checks[0].checked == 2
    assert "outside_digits" not in codes(ex)               # 免不免确认由确认项按上一期判断（7.5），不在这里报
    assert ex.ledger[0].roles["context"] == 2 and "outside_text" not in ex.ledger[0].roles
    # D28b：脚注里写了年月，同样是统计期来源、同样记进 outside_text
    ex = run(flow_book(days=3, edit=lambda ws, info: set_value(ws, "B32", "注：2026年8月数据为初步统计，待修订")))
    assert ex.period.annotated == {f"{SHEET}!B32": "注数据为初步统计待修订"}
    assert (f"{SHEET}!B32", "text_digits", True) in [(o.cell, o.kind, o.period_source) for o in ex.outside_text]


def test_outside_text_with_digits_needs_confirmation():
    ex = run(flow_book(days=3, edit=lambda ws, info: set_value(ws, "B32", "注：8月15日闸机故障，当日客流为估算值")))
    assert ex.ok
    p = problem(ex, "outside_digits")
    assert p.category == "confirm" and p.cells == [f"{SHEET}!B32"]
    assert (f"{SHEET}!B32", "text_digits", False) in [(o.cell, o.kind, o.period_source) for o in ex.outside_text]


def test_outside_number_is_structure_and_merges_per_row():
    def edit(ws, info):
        ws["B1"] = 987654                  # 日期表头以上：区域外
        ws["C1"] = "1,234"
    ex = run(flow_book(days=3, edit=edit))
    ps = [p for p in ex.problems if p.code == "outside_number"]
    assert len(ps) == 1 and ps[0].fix == "ignore_cells" and ps[0].cells == [f"{SHEET}!B1", f"{SHEET}!C1"]
    assert not ex.ok


def test_subtotal_formula_and_hidden_rows_make_refs_hidden():
    def subtotal(ws, info):
        ws[f"C{info['rows']['合计'][0]}"] = f"=SUBTOTAL(9,C{info['rows']['夜间'][0]}:C{info['rows']['夜间'][3]})"
        info["cache"].pop("C28")
    ex = run(flow_book(days=3, edit=subtotal))
    item = next(d for d in ex.derived if d.cell == f"{SHEET}!C28")
    assert item.func == "SUBTOTAL" and item.ref_problem == "hidden" and item.value is None  # 改了公式、没有缓存
    def hide(ws, info):
        ws.row_dimensions[info["rows"]["夜间"][1]].hidden = True
    include = recipe_dict(sheets__0__hidden={"rows": "include", "cols": "reject_if_any"})
    ex = run(flow_book(days=3, edit=hide), include)
    assert ex.ok and ex.hidden == {SHEET: {"rows": [23], "cols": [], "policy_rows": "include",
                                           "policy_cols": "reject_if_any"}}
    item = next(d for d in ex.derived if d.cell == f"{SHEET}!C28")
    assert item.hidden_in_range and item.ref_problem is None and len(item.ref_rowids) == 4
    assert not next(d for d in ex.derived if d.cell == f"{SHEET}!C29").hidden_in_range
    exclude = recipe_dict(sheets__0__hidden={"rows": "exclude", "cols": "reject_if_any"})
    ex = run(flow_book(days=3, edit=hide), exclude)
    assert ex.ok and ex.expected_rows["时段客流"] == 16 * 3
    item = next(d for d in ex.derived if d.cell == f"{SHEET}!C28")
    assert item.ref_problem == "hidden" and item.ref_rowids is None and item.hidden_in_range
    assert ex.ledger[0].roles["hidden_excluded"] == 4


def test_excluded_hidden_total_row_is_neither_checked_nor_kept():
    def hide(ws, info):
        ws.row_dimensions[info["rows"]["合计"][0]].hidden = True
    exclude = recipe_dict(sheets__0__hidden={"rows": "exclude", "cols": "reject_if_any"})
    ex = run(flow_book(days=3, edit=hide), exclude)
    assert ex.ok and len(ex.derived) == 6 and ex.expected_rows["时段客流_表内合计"] == 6
    assert {d.label for d in ex.derived} == {"22-24时合计", "18-24时合计"}


def test_hidden_rows_and_cols_rejected_by_default():
    def hide(ws, info):
        ws.row_dimensions[info["rows"]["夜间"][1]].hidden = True
        ws.column_dimensions["D"].hidden = True
        ws.auto_filter.ref = "B4:E30"
    ex = run(flow_book(days=3, edit=hide))
    rows, cols = problem(ex, "hidden_rows"), problem(ex, "hidden_cols")
    assert "第 23 行" in rows.message and "筛选" in rows.message and "D 列" in cols.message
    assert not ex.ok


def test_derived_item_key_carries_the_const_column(tmp_path):
    db = str(tmp_path / "k.db")
    ex = run(flow_book(days=3), db=db)
    item = next(d for d in ex.derived if d.cell == f"{SHEET}!D29")
    assert item.key == {"日期": "2026-08-02", "时段类别": "夜间"}
    assert (item.label, item.label_raw, item.start, item.end) == ("22-24时合计", "22-24 时合计", 22, 24)
    assert (item.base_table, item.base_value, item.start_col, item.end_col) == ("时段客流", "客流", "起始小时", "结束小时")
    assert item.keep_table == "时段客流_表内合计" and item.is_formula and item.formula == "=SUM(D26:D27)"
    conn = sqlite3.connect(db)
    rowids = [r for (r,) in conn.execute(
        "SELECT rowid FROM 时段客流 WHERE 日期='2026-08-02' AND 时段类别='夜间' AND 起始小时>=22 AND 结束小时<=24")]
    conn.close()
    assert item.ref_rowids == sorted(rowids)


def test_literal_totals_and_unrecognized_formulas():
    def literal(ws, info):
        for ref, v in info["cache"].items():
            ws[ref] = v
        info["cache"].clear()
    ex = run(flow_book(days=3, edit=literal))
    assert ex.ok and ex.derived_form == {"夜间合计": "literal"}
    assert all(not d.is_formula and d.ref_rowids is None and d.func is None for d in ex.derived)
    def odd(ws, info):
        ws["C28"] = "=C22+C23+C24+C25"
        ws["D28"] = "=SUM(D22:D24)*1"
        ws["E28"] = "=SUM(E22:E26)"
        info["cache"].update({"C28": 1, "D28": 1, "E28": 1})
    ex = run(flow_book(days=3, edit=odd))
    by = {d.cell.split("!")[1]: d for d in ex.derived}
    assert by["C28"].func == "plus" and by["C28"].ref_problem is None and len(by["C28"].ref_rowids) == 4
    assert by["D28"].ref_problem == "unrecognized" and by["D28"].func is None
    assert by["E28"].ref_problem is None and len(by["E28"].ref_rowids) == 5       # 引用对不对由核对模块判
    assert ex.derived_form == {"夜间合计": "formula"}
    def mixed(ws, info):
        ws["C28"] = info["cache"].pop("C28")
    assert run(flow_book(days=3, edit=mixed)).derived_form == {"夜间合计": "mixed"}


def test_uncached_totals_are_none_not_problems():
    ex = run(flow_book(days=3, uncached=True))
    assert ex.ok and all(d.value is None and d.is_formula for d in ex.derived)


# --------------------------------------------------------------------------
# 问题收集策略
# --------------------------------------------------------------------------


def test_one_failed_segment_does_not_stop_the_others():
    ex = run(flow_book(days=3, edit=lambda ws, info: set_value(ws, "B9", "日间分时段客流（人次）")))   # D09
    assert codes(ex) == {"title_not_found"}            # 日间的 11 行不再逐行报 row_unclaimed
    p = problem(ex, "title_not_found")
    assert p.fix == "rename_title" and p.category == "structure"
    assert ex.expected_rows == {"日客流": 3, "时段客流": 6 * 3, "时段客流_表内合计": 9}
    assert len(ex.derived) == 9


def test_label_duplicate_and_pk_duplicate_reported_together():
    ex = run(flow_book(days=3, edit=lambda ws, info: set_value(ws, "B12", "8－9")))
    assert {"label_duplicate", "label_missing", "pk_duplicate"} <= codes(ex)
    dup = problem(ex, "pk_duplicate")
    assert dup.cells[:2] == [f"{SHEET}!C11", f"{SHEET}!C12"] and "C11 和 C12" in dup.message


# --------------------------------------------------------------------------
# 标签集合、分段定位
# --------------------------------------------------------------------------


def test_label_set_problems():
    ex = run(flow_book(days=3, day_hours=tuple(range(8, 18))))                 # D04：日间少了 7-8
    assert codes(ex) == {"label_missing"} and problem(ex, "label_missing").fix == "remove_label"
    def add_zone(ws, info):                                                     # D13
        ws.unmerge_cells("B8:E8")
        ws["B8"] = "分区丙（人次）"
        for j in range(3):
            ws.cell(8, 3 + j, 5)
    p = problem(run(flow_book(days=3, edit=add_zone)), "label_unexpected")
    assert p.fix == "add_label" and p.cells == [f"{SHEET}!B8"]
    def subtotal_row(ws, info):                                                 # D18
        ws.insert_rows(21)
        ws["B21"] = "7-18 时合计"
        for j in range(3):
            ws.cell(21, 3 + j, 1)
    raw, fn = flow_book(days=3, edit=subtotal_row, uncached=True)
    p = problem(run((raw, fn)), "label_unparsed")
    assert p.fix == "declare_total" and p.cells == [f"{SHEET}!B21"]


def test_label_order_and_writing_changes_are_not_problems():
    def reorder(ws, info):                                                       # D19
        for j in range(3 + 1):
            c = 2 + j
            a, b, t = ws.cell(5, c).value, ws.cell(6, c).value, ws.cell(7, c).value
            ws.cell(5, c).value, ws.cell(6, c).value, ws.cell(7, c).value = b, t, a
    ex = run(flow_book(days=3, edit=reorder))
    assert ex.ok
    assert ex.labels[0].raw == ["分区甲（人次）", "分区乙（人次）", "全日客流（人次）"]
    def writing(ws, info):                                                       # D11 / D24
        ws["B11"] = "08:00-09:00"
        ws["B12"] = "9–10"
    ex = run(flow_book(days=3, edit=writing))
    assert ex.ok and ex.canonicalized == []
    day = next(lab for lab in ex.labels if lab.segment == "日间")
    assert day.raw[1:3] == ["08:00-09:00", "9–10"] and day.canonical[1:3] == ["8-9", "9-10"]


def test_title_and_segment_codes():
    def twice(ws, info):
        ws.unmerge_cells("B8:E8")
        ws["B8"] = "日间时段客流（人次）"
    p = problem(run(flow_book(days=3, edit=twice)), "title_ambiguous")
    assert p.cells == [f"{SHEET}!B8", f"{SHEET}!B9"]
    bad_pick = recipe_dict()
    bad_pick["sheets"][0]["blocks"][0]["segments"][2]["const"] = {"时段类别": {"pick": "晚上"}}
    p = problem(run(flow_book(days=3), bad_pick), "const_missing")
    assert p.cells == [f"{SHEET}!B21"] and "晚上" in p.message
    gone = recipe_dict()
    seg = gone["sheets"][0]["blocks"][0]["segments"][0]
    seg["labels"]["expect"] = ["甲（人次）", "乙（人次）"]
    seg["measures"] = {"甲（人次）": "甲", "乙（人次）": "乙"}
    gone["tables"][0]["units"] = {}
    gone["relations"] = []
    assert "segment_not_found" in codes(run(flow_book(days=3), gone))
    overlap = recipe_dict()
    overlap["sheets"][0]["blocks"][0]["segments"].append({
        "id": "日间按标签", "role": "dimension", "table": "另表", "locate": {"by": "labels"},
        "labels": {"expect": ["8-9", "9-10"]}, "dim": {"name": "时段", "parser": "hour_range"}, "value": "客流"})
    overlap["tables"].append({"name": "另表", "grain": ["日期", "时段"]})
    assert "segment_overlap" in codes(run(flow_book(days=3), overlap))


def test_row_codes():
    def extra_row(ws, info):                                                     # D21
        ws["B32"] = "补录（人次）"
        ws["C32"] = "1,234"
        ws["D32"] = "2,345"
    ex = run(flow_book(days=3, edit=extra_row))
    assert codes(ex) == {"row_unclaimed"}
    p = problem(ex, "row_unclaimed")
    assert p.fix == "add_label" and p.cells[0] == f"{SHEET}!B32"
    ex = run(flow_book(days=3, edit=lambda ws, info: set_value(ws, "D32", 987654)))
    assert codes(ex) == {"row_without_label"}


def test_cells_right_of_the_axis_on_data_rows_have_nowhere_to_go():
    def edit(ws, info):
        ws["G11"] = "备注"
    ex = run(flow_book(days=3, edit=edit))
    p = problem(ex, "cell_unclaimed")
    assert p.cells == [f"{SHEET}!G11"] and ex.ledger[0].unclaimed == 1


# --------------------------------------------------------------------------
# 日期轴
# --------------------------------------------------------------------------


def test_axis_location_codes():
    def no_dates(ws, info):
        for j in range(3):
            ws.cell(4, 3 + j).value = f"第{j}列"
    ex = run(flow_book(days=3, edit=no_dates))
    assert codes(ex) == {"axis_not_found"}          # 块级失败：别的格不再逐个报
    def twice(ws, info):
        for j in range(3):
            ws.cell(40, 3 + j).value = ws.cell(4, 3 + j).value
    assert "axis_ambiguous" in codes(run(flow_book(days=3, edit=twice)))
    p = problem(run(flow_book(days=5, edit=lambda ws, info: set_value(ws, "E4", "备注"))), "axis_gap")
    assert p.cells == [f"{SHEET}!E4"]
    def extra(ws, info):                                                          # D16 / P25
        ws["F4"] = "合计"
        for r in range(5, 31):
            ws.cell(r, 6).value = 999
    ex = run(flow_book(days=3, edit=extra))
    assert codes(ex) == {"axis_extra_cells"}       # F 列的 999 不再报成区域外数字
    assert problem(ex, "axis_extra_cells").fix == "ignore_cells"


def test_axis_year_and_assertions():
    sep = flow_book(days=3, period="统计时间范围：2026年9月1日至2026年9月3日",
                    filename="月报导出_2026-09-01_2026-09-03.xlsx")
    assert "axis_year_missing" in codes(run(sep))
    long = flow_book(start=dt.date(2026, 1, 5), days=3,
                     period="统计时间范围：2026年1月1日至2027年1月31日", filename="x.xlsx")
    assert "axis_year_ambiguous" in codes(run(long))
    ex = run(flow_book(days=5, edit=lambda ws, info: set_value(ws, "E4", "8月2日")))   # P6
    assert codes(ex) == {"axis_duplicate"}
    def swap(ws, info):
        ws["D4"], ws["E4"] = "8月3日", "8月2日"
    ex = run(flow_book(days=5, edit=swap))
    assert codes(ex) == {"axis_not_contiguous"}
    relaxed = recipe_dict()
    relaxed["sheets"][0]["blocks"][0]["axis"]["checks"] = ["covers_context"]
    assert run(flow_book(days=5, edit=swap), relaxed).ok
    p = problem(run(flow_book(days=5, period="统计时间范围：2026年8月1日至2026年8月6日")), "axis_coverage")
    assert "共 5 格" in p.message and "共 6 天" in p.message


def test_date_cells_axis_form():
    def as_dates(ws, info):
        for j in range(3):
            ws.cell(4, 3 + j).value = dt.datetime(2026, 8, 1 + j)
    ex = run(flow_book(days=3, edit=as_dates))
    assert ex.ok and ex.axes[0].form == "date"
    ex = run(flow_book(days=3, edit=lambda ws, info: set_value(ws, "C4", dt.date(2026, 8, 1))))
    assert ex.ok and ex.axes[0].form == "mixed"


# --------------------------------------------------------------------------
# 值的转换（4.5）
# --------------------------------------------------------------------------


@pytest.mark.parametrize("value,policy,expected", [
    (123, {}, 123),
    (123.0, {}, 123),
    (1.5, {}, "value_not_integer"),
    (7, {"type": "REAL"}, 7.0),
    (2.25, {"type": "REAL"}, 2.25),
    ("·", {}, None),
    (" · ", {}, None),
    ("1,234", {"text_number": "parse_thousands"}, 1234),
    ("1,234", {}, "value_text_number"),
    (None, {"blank": "null"}, None),
    (None, {}, "value_blank"),
    ("#N/A", {}, "value_error"),
    ("约5000", {}, "value_not_number"),
    ("1234", {}, "value_not_number"),
    (True, {}, "value_not_number"),
    (dt.datetime(2026, 8, 1), {}, "value_not_number"),
])
def test_value_conversion(tmp_path, value, policy, expected):
    data = recipe_dict()
    data["sheets"][0]["blocks"][0]["values"].update(policy)
    book = flow_book(days=3, edit=lambda ws, info: set_value(ws, "C11", value))
    db = str(tmp_path / "v.db")
    ex = run(book, data, db=db)
    if isinstance(expected, str):
        assert codes(ex) == {expected} and not ex.ok and not os.path.exists(db)
        return
    assert ex.ok, [(p.code, p.message) for p in ex.problems]
    conn = sqlite3.connect(db)
    got = conn.execute("SELECT 客流 FROM 时段客流 WHERE 日期='2026-08-01' AND 时段='8-9'").fetchone()[0]
    conn.close()
    assert got == expected and type(got) is type(expected)


def test_formula_policies_in_the_data_area():
    def formula(ws, info):
        ws["C11"] = "=C12+1"
        info["cache"]["C11"] = 42
    ex = run(flow_book(days=3, edit=formula))
    assert codes(ex) == {"value_formula"}
    accept = recipe_dict(sheets__0__blocks__0__values__formula="accept_cached")
    ex = run(flow_book(days=3, edit=formula), accept)
    assert ex.ok and ex.formula_cells_accepted == 1
    def uncached(ws, info):
        ws["C11"] = "=C12+1"
    assert codes(run(flow_book(days=3, edit=uncached), accept)) == {"formula_uncached"}
    ex = run(flow_book(days=3, edit=formula, keep_full_calc=True), accept)
    assert codes(ex) == {"formula_full_calc"} and ex.full_calc_on_load
    # 合计格是公式不受这两条限制：交给核对模块按标签区间重算
    ex = run(flow_book(days=3, keep_full_calc=True))
    assert ex.ok and ex.full_calc_on_load


# --------------------------------------------------------------------------
# 统计期
# --------------------------------------------------------------------------


def test_period_conflict_prefers_prefix():
    ex = run(flow_book(days=3, edit=lambda ws, info: set_value(ws, "B3", "客流汇总表（2026年8月2日至2026年8月4日）")))
    c1 = ex.context_checks[0]
    assert (c1.id, c1.status, c1.category, c1.acceptable, c1.failed) == ("C1", "mismatch", "data_quality", True, 1)
    assert ex.period.start == "2026-08-01"                 # prefer_prefix「统计时间范围」那格
    assert ex.ok                                           # 数据质量类：不阻断执行，由核对结果决定
    ex = run(flow_book(days=3, edit=lambda ws, info: set_value(ws, "B3", "客流汇总表（2026年9月）")))
    assert ex.context_checks[0].status == "mismatch"


def test_period_missing_and_human_input():
    no_period = flow_book(days=3, period="客流汇总")
    ex = run(no_period)
    p = problem(ex, "period_missing")
    assert p.category == "input" and "2026-08-01 至 2026-08-03" in p.message
    assert codes(ex) == {"period_missing"} and not ex.ok
    entered = {"统计期": PeriodInput(dt.date(2026, 8, 1), dt.date(2026, 8, 3), signed_by="张三")}
    ex = run(no_period, context_inputs=entered)
    assert ex.ok and ex.period.source == "human" and ex.period.signed_by == "张三"
    assert [(c.id, c.status) for c in ex.context_checks] == [("C2", "passed")]
    ex = run(flow_book(days=3, filename="9月客流.xlsx"))
    assert [(c.id, c.status) for c in ex.context_checks] == [("C1", "passed"), ("C2", "info")]
    ex = run(flow_book(days=3, filename="月报导出_2026-07-01_2026-07-03.xlsx"))
    c2 = ex.context_checks[1]
    assert (c2.status, c2.acceptable, c2.category) == ("mismatch", True, "data_quality") and ex.ok
    off = recipe_dict()
    off["sheets"][0]["context"][0]["cross_check"] = "none"
    assert [c.id for c in run(flow_book(days=3), off).context_checks] == ["C1"]


# --------------------------------------------------------------------------
# 工作表、写库、哈希、部分干跑、两遍对账
# --------------------------------------------------------------------------


def test_sheet_matching(tmp_path):
    def rename(wb_raw: bytes) -> bytes:
        return patch_part(wb_raw, "xl/workbook.xml", lambda t: t.replace('name="客流汇总"', 'name="客流汇总九月"'))
    raw, fn = flow_book(days=3)
    ex = run((rename(raw), fn))
    assert ex.ok and ex.sheets.renamed == {"客流汇总": "客流汇总九月"} and ex.sheets.matched == {"s1": "客流汇总九月"}
    strict = recipe_dict()
    strict["sheets"][0]["match"]["fallback"] = "none"
    assert codes(run((rename(raw), fn), strict)) == {"sheet_missing"}

    def with_extra(ws, info):
        wb = ws.parent
        wb.create_sheet("说明")["A1"] = "填表说明"
        hidden = wb.create_sheet("配置")
        hidden["A1"] = 987654
        hidden.sheet_state = "veryHidden"
    book = flow_book(days=3, edit=with_extra)
    ex = run(book)
    assert ex.ok and ex.sheets.other_visible == ["说明"]
    assert ex.sheets.skipped_hidden == [{"sheet": "配置", "state": "veryHidden"}]
    reject = recipe_dict(other_visible_sheets="reject")
    assert codes(run(book, reject)) == {"sheet_unexpected"}


def test_db_only_when_clean(tmp_path):
    db = str(tmp_path / "bad.db")
    ex = run(flow_book(days=3, day_hours=tuple(range(8, 18))), db=db)
    assert not ex.ok and not os.path.exists(db) and ex.table_hashes == {}
    ex = run(flow_book(days=3), db=None)
    assert ex.ok and ex.table_hashes == {} and ex.expected_rows["时段客流"] == 51
    with pytest.raises(ValueError):
        run(flow_book(days=3), db=str(tmp_path / "x.db"), max_rows=10)


def test_exception_removes_the_db(tmp_path, monkeypatch):
    db = str(tmp_path / "boom.db")
    def boom(self):
        raise RuntimeError("中途出错")
    monkeypatch.setattr(recipe_engine._Engine, "_hashes", boom)
    with pytest.raises(RuntimeError):
        run(flow_book(days=3), db=db)
    assert not os.path.exists(db)


def test_recipe_shape_problem_blocks_everything(tmp_path):
    data = recipe_dict()
    seg = data["sheets"][0]["blocks"][0]["segments"][0]
    seg["measures"] = {"全日客流(人次)": "全日客流", "分区甲（人次）": "分区甲", "分区乙（人次）": "分区乙"}
    db = str(tmp_path / "shape.db")
    ex = run(flow_book(days=3), data, db=db)
    assert [(p.code, p.category) for p in ex.problems] == [("measures_keys_mismatch", "recipe")]
    assert not ex.ok and not os.path.exists(db)


def _sqlite_order_key(row: tuple[Any, ...]) -> tuple[Any, ...]:
    # SQLite 的排序：NULL < 数 < 文本
    return tuple((0, 0) if v is None else (1, v) if isinstance(v, (int, float)) else (2, v) for v in row)


def test_table_hash_without_primary_key_matches_python_sort(tmp_path):
    data = recipe_dict()
    data["tables"][1]["grain"] = []
    db = str(tmp_path / "nopk.db")
    ex = run(flow_book(days=3), data, db=db)
    assert ex.ok
    conn = sqlite3.connect(db)
    ddl = conn.execute("SELECT sql FROM sqlite_master WHERE name='时段客流'").fetchone()[0]
    rows = conn.execute("SELECT * FROM 时段客流").fetchall()
    conn.close()
    assert "PRIMARY KEY" not in ddl
    h = hashlib.sha256()
    for row in sorted(rows, key=_sqlite_order_key):
        h.update(json.dumps(list(row), ensure_ascii=False, separators=(",", ":")).encode("utf-8") + b"\n")
    assert ex.table_hashes["时段客流"] == h.hexdigest()


def test_partial_dry_run_skips_whole_sheet_checks():
    ex = run(flow_book(days=3), max_rows=12)                 # 第 2–13 行
    assert ex.partial and ex.ok and not ex.problems
    assert ex.expected_rows == {"日客流": 3, "时段客流": 4 * 3, "时段客流_表内合计": 0}
    assert ex.ledger[0].nonempty_read < ex.ledger[0].nonempty_scan
    ex = run(flow_book(days=3, edit=lambda ws, info: set_value(ws, "B32", "补录")), max_rows=25)
    assert ex.partial and not ex.problems                    # 窗口之外的行不报


def test_pass_mismatch_when_row_runs_disagree(tmp_path):
    raw, fn = flow_book(days=3)
    sc = xlsx_scan.scan(raw)
    a, b, n = sc.sheets[0].row_runs[0]
    sc.sheets[0].row_runs[0] = (a, b, n - 1)
    db = str(tmp_path / "m.db")
    ex = run((raw, fn), scan=sc, db=db)
    p = problem(ex, "pass_mismatch")
    assert "第 2 行：扫描到 0 个非空格，读取到 1 个" in p.message and not ex.ok and not os.path.exists(db)


def test_pass_mismatch_by_total_when_row_runs_missing(tmp_path):
    """老暂存区的扫描经 scan_from_json 读回来时没有 row_runs：逐行比不了，只剩合计数把关（4.1：两遍对账的
    合计数还要等于 SheetScan.nonempty）。少读一格也必须报 pass_mismatch、不建库。"""
    raw, fn = flow_book(days=3)
    sc = xlsx_scan.scan(raw)
    sc.sheets[0].row_runs = []
    sc.sheets[0].nonempty += 1
    db = str(tmp_path / "t.db")
    ex = run((raw, fn), scan=sc, db=db)
    problem(ex, "pass_mismatch")
    assert not ex.ok and not os.path.exists(db)


def test_outside_formula_cells_are_numbers_whatever_the_cache(tmp_path):
    """区域外的公式格一律按 number 处理（4.2 第 6 步）：缓存值是文字（="备"&"注"）、或者没有缓存值，
    也按区域外数字拒收，不能当成文字只记录。"""
    raw, fn = flow_book(days=3)
    row = ('<row r="32"><c r="B32" t="str"><f>"备"&amp;"注"</f><v>备注</v></c></row>'
           '<row r="34"><c r="B34"><f>LEN(B32)</f></c></row>')
    raw = patch_part(raw, "xl/worksheets/sheet1.xml", lambda t: t.replace("</sheetData>", row + "</sheetData>", 1))
    ex = run((raw, fn), db=str(tmp_path / "f.db"))
    ps = [p for p in ex.problems if p.code == "outside_number"]          # 按行各一条
    assert {p.category for p in ps} == {"structure"}
    assert [c for p in ps for c in p.cells] == [f"{SHEET}!B32", f"{SHEET}!B34"], [p.cells for p in ps]
    assert {o.cell for o in ex.outside_text} == {f"{SHEET}!B3"}           # 只有原本的标题是区域外文字
    assert not ex.ok


def test_stale_dimension_gives_identical_hashes(tmp_path):
    raw, fn = flow_book(days=3)
    a = run((raw, fn), db=str(tmp_path / "a.db"))
    for ref in ("B2:E27", "B2:D30"):                         # P23 截行、P24 截列
        stale = patch_part(raw, "xl/worksheets/sheet1.xml",
                           lambda t, ref=ref: re.sub(r'<dimension ref="[^"]*"/>', f'<dimension ref="{ref}"/>', t))
        b = run((stale, fn), db=str(tmp_path / f"{ref.replace(':', '')}.db"))
        assert b.ok and b.table_hashes == a.table_hashes


def test_measures_moved_below_with_a_new_title():
    def move(ws, info):                                                           # D06
        ws.unmerge_cells("B8:E8")
        ws["B31"] = "日客流（人次）"
        for k in range(3):
            for c in range(2, 6):
                ws.cell(32 + k, c).value = ws.cell(5 + k, c).value
                ws.cell(5 + k, c).value = None
    ex = run(flow_book(days=3, edit=move))
    assert ex.ok
    assert ex.block_order == {SHEET: ["日间", "夜间", "夜间合计", "日客流"]}
    assert ex.lineage["日客流"]["全日客流"] == [[1, SHEET, "C32", 3, "right"]]
    assert (f"{SHEET}!B31", "text") in [(o.cell, o.kind) for o in ex.outside_text]


def test_block_order_follows_the_sheet():
    ex = run(flow_book(days=3, night_first=True))            # D05
    assert ex.ok and ex.block_order == {SHEET: ["日客流", "夜间", "夜间合计", "日间"]}
    # 长表按分段在工作表里的先后写：夜间在前，rowid 从夜间第一行起
    assert ex.lineage["时段客流"]["客流"][0] == [1, SHEET, "C10", 3, "right"]
    first_total = next(d for d in ex.derived if d.cell == f"{SHEET}!C16")
    assert first_total.ref_rowids == [1, 4, 7, 10]
    ex = run(flow_book(days=3, blank_row=False))             # D14：没有空行
    assert ex.ok
