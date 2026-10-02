"""证据下钻（期 4）接口一侧：文档标记、查询步骤的提示字段、按需裁判 loader 的链扩展、推断来源接口的路由和错误码、
规则表 R0–R17（P4-SPEC 2.3）与 R14 之后的异常映射。

三种写法：
- 单元（R0–R4、loader、提示字段）：在测试里现造查询快照、表结构快照（内容寻址存进临时数据目录），拼一份内存里的
  `_Sealed`，直接调 `segment_provenance`、`_chain`、`_sealed_loader`；
- 真导入（R5–R17）：经接口把合成的客流表（tests/fixtures/xlsx/flow.py，假名、随机数）按配方导入成上传源，经工具接口
  真跑查询得到查询快照，再拼内存里的 `_Sealed` 调 `segment_provenance`。要构造的反例（重复主键、篡改的清单、
  改过的登记哈希……）在这份真数据上改出来；
- 接口：真 runner 加脚本化的模型，跑「工具节点 db_query__<源> → 报告节点」的图，经 AsyncClient(ASGITransport)
  取推断来源。手工 SQLite 源和上传源各一。

每份结论都断言 contract_problems(out) == []。端到端验收（A1–A26）由 WP-E 覆盖。不调用任何模型。
"""
from __future__ import annotations

import asyncio
import copy
import hashlib
import io
import json
import os
import sqlite3
import sys
import threading
import types
import uuid
from dataclasses import asdict, dataclass, field, fields
from datetime import datetime, timezone
from pathlib import Path

import pytest
from httpx import ASGITransport, AsyncClient
from langchain_core.messages import AIMessage
from openpyxl import load_workbook
from sqlalchemy import select

from app.api import evidence as api
from app.core import artifact_store
from app.core.config import settings
from app.data.engine import SnapshotTampered, engines
from app.data.provenance_types import (
    DOC_PROVENANCE,
    SCHEMA,
    AcceptanceView,
    Alert,
    CanonicalView,
    CellRef,
    CellSource,
    DirectSelect,
    FromCell,
    LocateProblem,
    PartView,
    PeriodView,
    ProvenanceOut,
    Reason,
    Recheck,
    RelatedCheck,
    ReportRef,
    StateNote,
    TableRef,
    VersionView,
    YearSource,
    contract_problems,
)
from app.data.table_versions import SnapshotManifestMismatch, SnapshotMissing
from app.db.base import SessionLocal
from app.db.models import DataSource, Run
from app.engine.evidence import iter_segments
from app.engine.runner import run_manager
from app.main import app
from tests.fixtures.xlsx.flow import SHEET, flow_workbook

SOURCE = "flow_demo"
SQL = 'SELECT "日期", "全日客流" FROM "日客流" WHERE "日期" = \'2026-08-05\''


# --------------------------------------------------------------------------
# 单元夹具：现造的工件、内存里的封存范围
# --------------------------------------------------------------------------


async def put(obj, kind="query_snapshot") -> str:
    return await artifact_store.put_json(obj, kind=kind)


def _path(artifact: str):
    return settings.data_dir / "artifacts" / artifact[:2] / f"{artifact}.json"


def tamper(artifact: str) -> None:
    """内容改一个字：取回时哈希复验不通过（load 抛 ValueError）。"""
    path = _path(artifact)
    path.write_text(path.read_text("utf-8").replace("2026", "2027", 1), encoding="utf-8")


async def schema_snapshot(**extra) -> str:
    return await put({"schema": {"tables": [{"name": "日客流", "primary_key": ["日期"], "columns": [
        {"name": "日期", "type": "TEXT"}, {"name": "全日客流", "type": "INTEGER"}]}]},
        "synced_at": "2026-10-01T00:00:00+00:00", **extra}, kind="schema_snapshot")


async def query_snapshot(schema_id: str | None, *, data_version: str | None = "ab" * 32, **extra) -> str:
    snap = {"columns": ["日期", "全日客流"], "rows": [["2026-08-05", 8754]], "row_count": 1, "truncated": False,
            "elapsed_ms": 1, "sql": SQL, "column_types": {"日期": "date", "全日客流": "number"}, "source": SOURCE,
            **({"schema_artifact": schema_id} if schema_id else {}),
            **({"data_version": data_version} if data_version else {}), **extra}
    return await put(snap)


def sealed_of(*, queries=(), schemas=(), trusted=True, ledger=()) -> api._Sealed:
    seal = {"sealed": trusted, "ok": True if trusted else None, "manifest_seq": 9 if trusted else None,
            "legacy": False}
    return api._Sealed(run_id="run-p4c", graph={}, seal=seal, events=[],
                       queries={q: {"node_id": "fetch", "tool": f"db_query__{SOURCE}"} for q in queries},
                       schemas={s: {"node_id": "fetch"} for s in schemas}, ledger=set(ledger))


_TOOL = object()


def doc_of(query_id: str, *, marked=True, tool=_TOOL, kind="query", extra_catalog=None, source=SOURCE) -> dict:
    tool = f"db_query__{source}" if tool is _TOOL else tool
    doc = {"schema": "agentlab.report/1", "entity_syntax": 2, "run_id": "run-p4c", "node_id": "write",
           "catalog": {"Q1": {"kind": kind, "artifact": query_id, "tool": tool, "node_id": "fetch", "source": source},
                       **(extra_catalog or {})}, "blocks": []}
    if marked:
        doc["provenance"] = DOC_PROVENANCE
    return doc


def cell_seg(*, column="全日客流", row=0, **cite) -> dict:
    return {"id": "s4", "kind": "value", "state": "deterministic", "text": "8,754",
            "cite": {"kind": "cell", "status": "resolved", "alias": "Q1", "locator": {"row": row, "column": column},
                     **cite}}


def report_of(doc: dict) -> api._Report:
    return api._Report(node_id="write", doc_artifact="d0" * 32, ok=True, repairs=0, stats=None, doc=doc,
                       hash_ok=True, fields=[])


def ok(out: ProvenanceOut) -> ProvenanceOut:
    assert contract_problems(out) == [], contract_problems(out)
    return out


async def provenance(doc, seg, sealed) -> ProvenanceOut:
    return ok(await api.segment_provenance(report_of(doc), seg, sealed, masks=api._Masks()))


@pytest.fixture
def drill(monkeypatch):
    """第二部分（R5 起）的替身：记下 R4 交过去的东西，回一个占位的结论。"""
    calls: list[dict] = []

    async def fake(base, **kw):
        calls.append({"base": base, **kw})
        return ProvenanceOut(status="none", reason=api.refuse("manifest_unreadable"),
                             alert=api.alert("manifest_unreadable"), **base)

    monkeypatch.setattr(api, "_drill", fake)
    return calls


# --------------------------------------------------------------------------
# R0–R4
# --------------------------------------------------------------------------


async def test_r0_legacy_doc_still_names_the_cell(drill):
    """期 4 之前的文档（没有标记）一律 legacy_doc；cell 只看片段本身是不是单元格引用，与 status 无关。"""
    sid = await schema_snapshot(import_mode="recipe")
    qid = await query_snapshot(sid)
    out = await provenance(doc_of(qid, marked=False), cell_seg(), sealed_of(queries=[qid], schemas=[sid]))
    assert out.status == "none" and out.reason.code == "legacy_doc" and out.alert is None
    assert out.version is None and out.cell_source is None and out.checks == []
    assert out.cell == CellRef(alias="Q1", row=0, column="全日客流", artifact=qid)
    assert out.report == ReportRef(node_id="write", doc_artifact="d0" * 32) and out.segment == "s4"
    assert out.sealed is True and out.schema == SCHEMA
    # 标记的版本号不认识（以后的文法）也不按 1 的规则判
    out = await provenance({**doc_of(qid), "provenance": 2}, cell_seg(), sealed_of(queries=[qid], schemas=[sid]))
    assert out.reason.code == "legacy_doc"
    assert drill == []


@pytest.mark.parametrize("case", ["metric", "unresolved", "not_deterministic", "entry_not_query", "agent_field",
                                  "code_node"])
async def test_r1_only_direct_query_cells(case, drill):
    sid = await schema_snapshot(import_mode="recipe")
    qid = await query_snapshot(sid)
    seg, doc = cell_seg(), doc_of(qid)
    if case == "metric":
        seg = {"id": "s9", "kind": "value", "state": "deterministic", "text": "8,754",
               "cite": {"kind": "metric", "status": "resolved", "alias": "M1", "locator": {"metric": "total"}}}
    elif case == "unresolved":
        seg["cite"]["status"] = "unresolved"
    elif case == "not_deterministic":
        seg["state"] = "flagged"
    elif case == "entry_not_query":
        doc = doc_of(qid, kind="retrieval")
    elif case == "agent_field":
        doc = doc_of(qid, tool="agent")
    else:
        doc = doc_of(qid, tool=None)
    out = await provenance(doc, seg, sealed_of(queries=[qid], schemas=[sid]))
    assert out.status == "none" and out.reason.code == "not_cell" and out.version is None
    # 单元格引用没解析成功也照样给 cell（2.8.1：与 status 无关）；指标片段没有 cell
    assert (out.cell is None) == (case == "metric")
    assert drill == []


async def test_r2_query_snapshot_must_be_returned_in_this_run(drill):
    sid = await schema_snapshot(import_mode="recipe")
    qid = await query_snapshot(sid)
    # 不在本次运行的事件里
    out = await provenance(doc_of(qid), cell_seg(), sealed_of(queries=[], schemas=[sid]))
    assert (out.status, out.reason.code, out.alert) == ("none", "not_sealed", None)
    # 在事件里，但取回时哈希复验不通过
    bad = await query_snapshot(sid, elapsed_ms=2)
    tamper(bad)
    out = await provenance(doc_of(bad), cell_seg(), sealed_of(queries=[bad], schemas=[sid]))
    assert out.reason.code == "not_sealed"
    # 在事件里，文件不在了
    gone = await query_snapshot(sid, elapsed_ms=3)
    _path(gone).unlink()
    out = await provenance(doc_of(gone), cell_seg(), sealed_of(queries=[gone], schemas=[sid]))
    assert out.reason.code == "not_sealed"
    assert drill == []


async def test_r3_manual_sources_are_not_uploads(drill):
    sid = await schema_snapshot()
    qid = await query_snapshot(sid, data_version=None)
    out = await provenance(doc_of(qid), cell_seg(), sealed_of(queries=[qid], schemas=[sid]))
    assert (out.status, out.reason.code, out.version) == ("none", "not_upload", None)
    assert drill == []


async def test_r4_schema_snapshot_and_import_mode(drill):
    recipe = await schema_snapshot(import_mode="recipe")
    # 查询快照没记表结构快照
    qid = await query_snapshot(None)
    assert (await provenance(doc_of(qid), cell_seg(), sealed_of(queries=[qid]))).reason.code == "not_sealed"
    # 表结构快照不在本次运行的事件里
    qid = await query_snapshot(recipe)
    assert (await provenance(doc_of(qid), cell_seg(), sealed_of(queries=[qid]))).reason.code == "not_sealed"
    # 在事件里，哈希不符
    bad = await schema_snapshot(import_mode="recipe", total=1)
    tamper(bad)
    qid = await query_snapshot(bad)
    out = await provenance(doc_of(qid), cell_seg(), sealed_of(queries=[qid], schemas=[bad]))
    assert out.reason.code == "not_sealed" and out.alert is None
    # 简单导入（没有 import_mode）、v0 迁移出的版本：只说一句，不标红
    simple = await schema_snapshot()
    qid = await query_snapshot(simple)
    out = await provenance(doc_of(qid), cell_seg(), sealed_of(queries=[qid], schemas=[simple]))
    assert (out.status, out.reason.code, out.alert, out.version) == ("none", "simple_upload", None, None)
    assert drill == []


@pytest.mark.parametrize("trusted", [True, False])
async def test_r4_passed_hands_over_to_part_two_regardless_of_seal(drill, trusted):
    """配方导入的上传源：R0–R4 都满足，交给第二部分。未封存（停在审批上）、封存核对不通过的运行照常往下判，
    `sealed` 照实给，不会被误判成清单无法读取（推断来源接口不走 _sealed_loader）。"""
    sid = await schema_snapshot(import_mode="recipe", import_manifests=["c" * 64])
    qid = await query_snapshot(sid)
    sealed = sealed_of(queries=[qid], schemas=[sid], trusted=trusted)
    doc = doc_of(qid)
    out = await api.segment_provenance(report_of(doc), cell_seg(), sealed, masks=api._Masks())
    assert len(drill) == 1
    call = drill[0]
    assert call["query"]["data_version"] == "ab" * 32 and call["schema"]["import_mode"] == "recipe"
    assert call["entry"] is doc["catalog"]["Q1"] and call["sealed"] is sealed
    assert call["base"]["sealed"] is trusted and out.sealed is trusted
    # 前几条取过的工件留在同一份备忘里，第二部分不必再取
    assert call["artifacts"]._memo[qid][1] is True and call["artifacts"]._memo[sid][1] is True


async def test_artifacts_memo_fetches_once_in_a_thread(monkeypatch):
    qid = await query_snapshot(None, elapsed_ms=11)
    bad = await query_snapshot(None, elapsed_ms=12)
    tamper(bad)
    calls: list[tuple[str, bool]] = []
    real = api._fetch

    def spy(artifact):
        calls.append((artifact, threading.current_thread() is threading.main_thread()))
        return real(artifact)

    monkeypatch.setattr(api, "_fetch", spy)
    memo = api._Artifacts()
    first = await memo.get(qid)
    assert first[1] is True and (await memo.get(qid)) is first and memo.load(qid) == first[0]
    assert calls == [(qid, False)]                    # 只取一次，在线程里
    assert (await memo.get(bad))[1] is False
    with pytest.raises(ValueError):
        memo.load(bad)                                # 同步 loader 的语义同 artifact_store.load
    assert memo.load("f" * 64) is None and memo.load("") is None
    assert await memo.get("") == (None, None)


# --------------------------------------------------------------------------
# R14 之后的异常映射
# --------------------------------------------------------------------------


def test_snapshot_failure_order():
    """SnapshotManifestMismatch 同时是 SnapshotMissing 和 SnapshotTampered，必须最先认（顺序是变异守卫点）。"""
    mismatch = SnapshotManifestMismatch("清单不一致")
    assert isinstance(mismatch, SnapshotMissing) and isinstance(mismatch, SnapshotTampered)
    reason, red = api._snapshot_failure(mismatch)
    assert (reason.code, red) == ("chain_mismatch", Alert(code="chain_mismatch", text=reason.text))
    reason, red = api._snapshot_failure(SnapshotTampered("文件被改"))
    assert (reason.code, red.code) == ("db_tampered", "db_tampered") and red.text == reason.text
    reason, red = api._snapshot_failure(SnapshotMissing("文件不在"))
    assert (reason.code, red) == ("snapshot_gone", None)
    assert all(isinstance(e, api._SNAPSHOT_ERRORS)
               for e in (mismatch, SnapshotTampered("x"), SnapshotMissing("x")))
    with pytest.raises(KeyError):
        api._snapshot_failure(KeyError("不是数据文件的异常，原样抛出"))
    # 三种结论都是 table_only 的合法组合
    for exc in (mismatch, SnapshotTampered("x"), SnapshotMissing("x")):
        reason, red = api._snapshot_failure(exc)
        assert isinstance(reason, Reason)
        out = ProvenanceOut(report=ReportRef(node_id="w", doc_artifact="d"), segment="s1", status="table_only",
                            reason=reason, alert=red, sealed=True,
                            version=VersionView(source=SOURCE, snapshot_id="ab" * 32, mode="replace", union=False))
        assert contract_problems(out) == []


# --------------------------------------------------------------------------
# 查询步骤的提示字段（2.8.2）
# --------------------------------------------------------------------------


@pytest.mark.parametrize("marked,upload", [(True, True), (True, False), (False, True), (False, False)])
async def test_query_step_hint_needs_marker_and_upload(marked, upload):
    sid = await schema_snapshot(import_mode="recipe")
    qid = await query_snapshot(sid, data_version="ab" * 32 if upload else None)
    sealed = sealed_of(queries=[qid], schemas=[sid])
    plain = await api._chain(cell_seg(), doc_of(qid, marked=False), sealed, "write")
    steps = await api._chain(cell_seg(), doc_of(qid, marked=marked), sealed, "write")
    assert len(steps) == 1 and steps[0]["step"] == "query"
    if marked and upload:
        assert steps[0]["provenance"] is True
        # 只多这一个键，其余一字不差
        assert {k: v for k, v in steps[0].items() if k != "provenance"} == plain[0]
    else:
        assert "provenance" not in steps[0] and steps[0] == plain[0]


async def test_query_step_hint_ignores_seal_but_needs_the_snapshot():
    sid = await schema_snapshot(import_mode="recipe")
    qid = await query_snapshot(sid)
    # 未封存的运行也给提示：推断来源照常推断
    steps = await api._chain(cell_seg(), doc_of(qid), sealed_of(queries=[qid], schemas=[sid], trusted=False), "write")
    assert steps[0]["provenance"] is True and steps[0]["sealed"] is False
    # 快照不是本次运行交回过的：一行都不给，也不给提示
    steps = await api._chain(cell_seg(), doc_of(qid), sealed_of(queries=[], schemas=[sid]), "write")
    assert steps[0]["note"] == api.UNSEALED_QUERY and "provenance" not in steps[0]


async def test_metric_chain_query_steps_never_carry_the_hint():
    """同一份带标记的文档、同一份上传源的快照：指标链里的查询步骤不带提示（指标片段只会回 not_cell）。"""
    sid = await schema_snapshot(import_mode="recipe")
    qid = await query_snapshot(sid)
    card = await put({"caliber": "周报口径", "caliber_version": "v1", "metrics": [{
        "id": "total", "name": "全日客流", "value": 8754, "unit": "人次", "decimals": 0,
        "inputs": [{"via": "tool_cell", "artifact": qid, "node_id": "fetch",
                    "locator": {"row": 0, "column": "全日客流"}}]}]}, kind="metric_set")
    doc = doc_of(qid, extra_catalog={"M1": {"kind": "metric", "artifact": card, "node_id": "card",
                                            "locator": {"metric": "total"}, "name": "全日客流"}})
    seg = {"id": "s9", "kind": "value", "state": "deterministic", "text": "8,754",
           "cite": {"kind": "metric", "status": "resolved", "alias": "M1", "locator": {"metric": "total"}}}
    steps = await api._chain(seg, doc, sealed_of(queries=[qid], schemas=[sid], ledger=[("card", card)]), "write")
    queries = [s for s in steps if s.get("step") == "query"]
    assert len(queries) == 1 and queries[0]["artifact"] == qid and queries[0]["rows"]
    assert "provenance" not in queries[0]
    # 同一份文档里直接引用同一格的片段照常带
    assert (await api._chain(cell_seg(), doc, sealed_of(queries=[qid], schemas=[sid]), "write"))[0]["provenance"]


# --------------------------------------------------------------------------
# 按需裁判的 loader：链扩展（2.8.3）
# --------------------------------------------------------------------------


@pytest.fixture
async def chain_world(monkeypatch):
    """封存范围里一份查询快照、一份表结构快照；表结构快照经哈希链列出一份导入清单（由 chain_artifacts 的替身给出）。"""
    manifest = await put({"source_id": "s" * 32, "seq": 1, "period": {"start": "2026-08-01"}}, kind="import_manifest")
    stray = await put({"source_id": "s" * 32, "seq": 9, "period": {"start": "2026-09-01"}}, kind="import_manifest")
    sid = await schema_snapshot(import_mode="recipe", import_manifests=[manifest])
    qid = await query_snapshot(sid)
    calls: list[tuple[dict, object]] = []

    def fake_chain(schema, loader):
        calls.append((schema, loader))
        return set(schema.get("import_manifests") or [])

    monkeypatch.setattr(api, "_chain_artifacts", fake_chain)
    return types.SimpleNamespace(manifest=manifest, stray=stray, schema=sid, query=qid, calls=calls)


async def test_sealed_loader_without_chain_is_unchanged(chain_world):
    w = chain_world
    loader = api._sealed_loader(sealed_of(queries=[w.query], schemas=[w.schema]))
    assert loader(w.query)["data_version"] == "ab" * 32 and loader(w.schema)["import_mode"] == "recipe"
    assert loader(w.manifest) is None and loader(w.stray) is None
    assert w.calls == []                                  # 老文档根本不去碰清单


async def test_sealed_loader_chain_is_lazy_and_bounded(chain_world):
    w = chain_world
    sealed = sealed_of(queries=[w.query], schemas=[w.schema])
    loader = api._sealed_loader(sealed, chain=True)
    assert loader(w.query) is not None and w.calls == []  # 基本范围内的不触发扩展
    assert loader(w.manifest)["seq"] == 1                 # 链上的清单取得到
    assert len(w.calls) == 1 and w.calls[0][0]["import_manifests"] == [w.manifest]
    assert loader(w.stray) is None                        # 不在链上的取不到
    assert loader("f" * 64) is None
    assert len(w.calls) == 1                              # 同一个 loader 只算一次
    api._sealed_loader(sealed, chain=True)(w.manifest)
    assert len(w.calls) == 2                              # 新的一次请求重新算


async def test_sealed_loader_chain_still_needs_a_trusted_seal(chain_world):
    w = chain_world
    loader = api._sealed_loader(sealed_of(queries=[w.query], schemas=[w.schema], trusted=False), chain=True)
    assert loader(w.query) is None and loader(w.manifest) is None and w.calls == []


async def test_sealed_loader_chain_skips_unreadable_schema_snapshots(chain_world):
    w = chain_world
    bad = await schema_snapshot(import_mode="recipe", import_manifests=[w.stray], total=3)
    tamper(bad)
    loader = api._sealed_loader(sealed_of(queries=[w.query], schemas=[bad, w.schema, "e" * 64]), chain=True)
    assert loader(w.manifest)["seq"] == 1 and loader(w.stray) is None
    assert [c[0]["import_manifests"] for c in w.calls] == [[w.manifest]]


def test_chain_artifacts_delegates_to_the_provenance_module(monkeypatch):
    seen = []
    fake = types.ModuleType("app.data.provenance")
    fake.chain_artifacts = lambda schema, load: seen.append((schema, load)) or {"m1", "m2"}
    monkeypatch.setitem(sys.modules, "app.data.provenance", fake)
    loader = object()
    assert api._chain_artifacts({"import_manifests": ["m1"]}, loader) == {"m1", "m2"}
    assert seen == [({"import_manifests": ["m1"]}, loader)]


def test_chain_artifacts_uses_the_real_chain():
    """不经替身：表结构快照列出的导入清单在链上（provenance.chain_artifacts），快照清单取不回、哈希不符时只是少列、
    不抛（按需裁判不因一份坏清单报错）。"""
    def broken(_aid):
        raise ValueError("哈希不符")

    assert api._chain_artifacts({"import_manifests": ["m1", "m2"]}, lambda a: None) == {"m1", "m2"}
    assert api._chain_artifacts({"import_manifests": ["m1"], "snapshot_manifest": "sm"}, lambda a: None) == {"m1"}
    assert api._chain_artifacts({"import_manifests": ["m1"], "snapshot_manifest": "sm"}, broken) == {"m1"}
    sm = {"parts": [{"manifest_artifact": "m1"}, {"manifest_artifact": "m0"}]}
    assert api._chain_artifacts({"import_manifests": ["m1"], "snapshot_manifest": "sm"},
                                lambda a: sm if a == "sm" else None) == {"m1", "m0", "sm"}


@pytest.mark.parametrize("marked", [True, False])
async def test_judge_now_extends_the_loader_only_for_marked_docs(monkeypatch, marked):
    from app.engine import judge as judging

    chains: list[bool] = []
    monkeypatch.setattr(api, "_sealed_loader", lambda sealed, *, chain=False: chains.append(chain) or (lambda a: None))
    monkeypatch.setattr(judging, "prepare", lambda *a, **kw: types.SimpleNamespace(skipped={}, cands=[]))
    doc = {"blocks": [], "catalog": {}, **({"provenance": DOC_PROVENANCE} if marked else {})}
    out = await api._judge_now(sealed_of(), report_of(doc), ["u1"])
    assert chains == [marked] and out["calls"] == 0


# --------------------------------------------------------------------------
# 接口：真运行
# --------------------------------------------------------------------------

TEXT = "2026-08-05 的全日客流为 [[v:Q1.r0.全日客流]] 人次，口径卡记为 [[m:total]]。"


@pytest.fixture
async def engine_up():
    await run_manager.setup()
    yield
    await run_manager.shutdown()


@pytest.fixture
async def client():
    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as c:
        yield c


@pytest.fixture
async def flow_source(tmp_path, engine_up):
    """手工 SQLite 源 flow_demo（不是上传表格：查询快照没有 data_version）。"""
    path = tmp_path / "flow.db"
    db = sqlite3.connect(path)
    db.execute('CREATE TABLE "日客流" ("日期" TEXT PRIMARY KEY, "全日客流" INTEGER)')
    db.executemany('INSERT INTO "日客流" VALUES (?, ?)', [(f"2026-08-{d:02d}", 1000 + 37 * d) for d in range(1, 11)])
    db.commit()
    db.close()
    async with SessionLocal() as session:
        row = (await session.execute(select(DataSource).where(DataSource.name == SOURCE))).scalar_one_or_none()
        if row is None:
            row = DataSource(name=SOURCE, kind="sqlite", database=str(path), readonly=True, enabled=True)
            session.add(row)
        row.database, row.kind, row.enabled, row.readonly, row.options = str(path), "sqlite", True, True, {}
        await session.commit()
        source_id = row.id
    await engines.invalidate(source_id)
    yield path
    await engines.invalidate(source_id)


def node(nid, ntype, **config):
    return {"id": nid, "type": ntype, "position": {"x": 0, "y": 0}, "data": {"label": nid, "config": config}}


def chain(*nodes):
    return {"nodes": list(nodes), "edges": [{"source": a["id"], "target": b["id"]} for a, b in zip(nodes, nodes[1:])]}


def graph(*, report=True):
    nodes = [node("start", "input"), node("fetch", "tool", tool=f"db_query__{SOURCE}", args={"sql": SQL}),
             node("card", "metrics", caliber="周报口径", caliber_version="v1", metrics=[
                 {"id": "total", "name": "全日客流", "unit": "人次", "decimals": 0,
                  "expression": "cell(nodes.fetch, 0, '全日客流')"}])]
    if report:
        nodes.append(node("write", "report", instructions="写日报", on_violation="flag"))
        nodes.append(node("out", "output", fields=[{"name": "日报", "value": "{{ nodes.write.text }}"}]))
    else:
        nodes.append(node("out", "output", fields=[{"name": "日报", "value": "无"}]))
    return chain(*nodes)


async def run(monkeypatch, *, report=True) -> Run:
    from app.providers import mock_model

    monkeypatch.setattr(mock_model.MockChatModel, "_decide", lambda self, messages: AIMessage(content=TEXT))
    run_id = (await run_manager.start(graph=graph(report=report), input_payload={})).id
    for _ in range(400):
        await asyncio.sleep(0.05)
        async with SessionLocal() as session:
            row = await session.get(Run, run_id)
        if row.status in ("succeeded", "failed") and row.manifest_seq is not None:
            assert row.status == "succeeded", row.error
            return row
    raise AssertionError(f"run {run_id} 没有结束：{row.status}")


def doc_id(row: Run) -> str:
    return row.output["_evidence"]["doc_artifact"]


def load_doc(row: Run) -> dict:
    return json.loads(_path(doc_id(row)).read_text("utf-8"))


def seg_of(doc: dict, kind: str) -> dict:
    return next(s for s in iter_segments(doc) if (s.get("cite") or {}).get("kind") == kind)


def _opt(cls, value):
    return cls(**value) if value is not None else None


def from_json(body: dict) -> ProvenanceOut:
    """响应 JSON → ProvenanceOut（逐层还原嵌套的类型），给 contract_problems 用。还原后再 asdict 必须逐字等于原 JSON：
    键名（含 cell_source 的「from」）、顺序、取值都没有在路由里走样。"""
    assert list(body) == [f.name for f in fields(ProvenanceOut)]
    v, cs = body["version"], body["cell_source"]
    version = None if v is None else VersionView(**{
        **v, "tables": [TableRef(**t) for t in v["tables"]],
        "parts": [PartView(**{**p, "period": _opt(PeriodView, p["period"]),
                              "acceptances": [AcceptanceView(**a) for a in p["acceptances"]],
                              "purged": _opt(StateNote, p["purged"]), "revoked": _opt(StateNote, p["revoked"])})
                  for p in v["parts"]]})
    cell_source = None if cs is None else CellSource(**{
        **cs, "from": [FromCell(**f) for f in cs["from"]], "year": _opt(YearSource, cs["year"]),
        "canonical": _opt(CanonicalView, cs["canonical"]), "recheck": _opt(Recheck, cs["recheck"])})
    out = ProvenanceOut(**{**body, "report": ReportRef(**body["report"]), "cell": _opt(CellRef, body["cell"]),
                           "reason": _opt(Reason, body["reason"]), "alert": _opt(Alert, body["alert"]),
                           "version": version, "cell_source": cell_source,
                           "checks": [RelatedCheck(**{**c, "acceptance": _opt(AcceptanceView, c["acceptance"])})
                                      for c in body["checks"]]})
    assert asdict(out) == body
    return out


async def test_new_documents_carry_the_marker(monkeypatch, flow_source, client):
    row = await run(monkeypatch)
    doc = load_doc(row)
    # 工件按规范 JSON 存（键排序），标记在顶层、参与内容哈希
    assert doc["provenance"] == DOC_PROVENANCE == 1
    assert artifact_store.content_hash(artifact_store.canonical_json(doc)) == doc_id(row)
    assert doc["stats"]["violations"] == 0


async def test_route_answers_with_status_not_errors(monkeypatch, flow_source, client):
    """手工源的单元格回 not_upload，指标片段回 not_cell：推断不出来不是错误，都是 200。片段接口的查询步骤
    （手工源）不带提示。"""
    row = await run(monkeypatch)
    doc = load_doc(row)
    cell, metric = seg_of(doc, "cell"), seg_of(doc, "metric")
    resp = await client.get(f"/api/runs/{row.id}/evidence/segments/{cell['id']}/provenance")
    assert resp.status_code == 200, resp.text
    body = resp.json()
    out = ok(from_json(body))
    assert (out.status, out.reason.code, out.sealed) == ("none", "not_upload", True)
    assert out.report == ReportRef(node_id="write", doc_artifact=doc_id(row)) and out.segment == cell["id"]
    assert out.cell.alias == "Q1" and out.cell.row == 0 and out.cell.column == "全日客流"
    assert out.cell.artifact == doc["catalog"]["Q1"]["artifact"]
    # 指定报告节点也一样
    again = await client.get(f"/api/runs/{row.id}/evidence/segments/{cell['id']}/provenance?report=write")
    assert again.json() == body

    resp = await client.get(f"/api/runs/{row.id}/evidence/segments/{metric['id']}/provenance")
    out = ok(from_json(resp.json()))
    assert (out.status, out.reason.code, out.cell) == ("none", "not_cell", None)

    step = (await client.get(f"/api/runs/{row.id}/evidence/segments/{cell['id']}")).json()["chain"][0]
    assert step["step"] == "query" and step["rows"] and "provenance" not in step


async def test_route_error_codes(monkeypatch, flow_source, client):
    row = await run(monkeypatch)
    doc = load_doc(row)
    cell = seg_of(doc, "cell")
    base = f"/api/runs/{row.id}/evidence/segments"

    async def code(url):
        resp = await client.get(url)
        return resp.status_code, resp.json().get("code")

    assert await code(f"/api/runs/nosuchrun/evidence/segments/{cell['id']}/provenance") == (404, "run_not_found")
    assert await code(f"{base}/{cell['id']}/provenance?report=nosuchnode") == (404, "evidence_report_not_found")
    assert await code(f"{base}/s999/provenance") == (404, "evidence_segment_not_found")
    path = _path(doc_id(row))
    original = path.read_text("utf-8")
    path.write_text(original.replace('"provenance":1', '"provenance":2'), encoding="utf-8")
    assert await code(f"{base}/{cell['id']}/provenance") == (409, "evidence_doc_tampered")
    path.unlink()
    assert await code(f"{base}/{cell['id']}/provenance") == (404, "evidence_doc_missing")
    path.write_text(original, encoding="utf-8")
    assert (await client.get(f"{base}/{cell['id']}/provenance")).status_code == 200


async def test_route_without_a_report(monkeypatch, flow_source, client):
    row = await run(monkeypatch, report=False)
    resp = await client.get(f"/api/runs/{row.id}/evidence/segments/s1/provenance")
    assert (resp.status_code, resp.json().get("code")) == (404, "evidence_report_not_found")


async def test_route_on_a_pre_p4_document(monkeypatch, flow_source, client):
    """期 4 之前组装的文档（报告节点存下的就是没有标记的老形状）：legacy_doc，cell 照样给。"""
    from app.engine.nodes import report as report_node

    real = report_node.compose_doc

    def legacy(*a, **kw):
        doc = real(*a, **kw)
        doc.pop("provenance", None)
        return doc

    monkeypatch.setattr(report_node, "compose_doc", legacy)
    row = await run(monkeypatch)
    doc = load_doc(row)
    assert "provenance" not in doc
    cell = seg_of(doc, "cell")
    out = ok(from_json((await client.get(f"/api/runs/{row.id}/evidence/segments/{cell['id']}/provenance")).json()))
    assert (out.status, out.reason.code) == ("none", "legacy_doc") and out.cell is not None
    assert asdict(out)["cell"]["artifact"] == doc["catalog"]["Q1"]["artifact"]


# ==========================================================================
# 第二部分：真导入的上传源（R5–R17、异常映射、相关核对、当前状态）
# ==========================================================================

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
DS_API = "/api/datasources"
FLOW_TABLES = ("日客流", "时段客流", "时段客流_表内合计")
REASON = "合成理由：分区合计口径调整"
WIDE_SQL = 'SELECT "日期", "全日客流" FROM "日客流" WHERE "日期" = \'2026-08-05\''


@dataclass
class Upload:
    """经接口按配方导入的上传源：名字、id、当前快照、各期导入记录和原件（按导入顺序）。"""

    name: str
    source_id: str
    snapshot: str
    imports: list[str] = field(default_factory=list)
    raws: list[bytes] = field(default_factory=list)
    files: list[str] = field(default_factory=list)


@pytest.fixture
def upload_world(monkeypatch, tmp_path):
    """每个用例一个独立的数据目录（原件、构建库、快照库、工件互不串）。导入不接模型：规则起草不完整时会查 AI 可用性，
    换成「不可用」的桩，连模型对象都不建。"""
    from app.data import recipe_ai
    from app.data.recipe_types import AiAvailability

    data = tmp_path / "data"
    (data / "uploads").mkdir(parents=True)
    monkeypatch.setattr(settings, "data_dir", data)

    async def no_model(session):
        return None, "", AiAvailability(False, reason="测试不接模型")

    monkeypatch.setattr(recipe_ai, "resolve_draft_model", no_model)
    return data


def _answers(questions: list[dict], mode: str) -> dict[str, dict]:
    """系统发现一律「登记」，q_mode 按 mode，其余能存为空值的选「存为空值」（参考配方的写法）。"""
    out = {}
    for q in questions:
        values = [o["value"] for o in q["options"]]
        if q["id"].startswith("q_relation"):
            out[q["id"]] = {"value": "register"}
        elif q["id"] == "q_mode":
            out[q["id"]] = {"value": mode}
        elif "null" in values:
            out[q["id"]] = {"value": "null"}
        else:
            out[q["id"]] = {"value": values[0]}
    return out


def _rename_wide(rec: dict, old: str, new: str) -> dict:
    rec = json.loads(json.dumps(rec, ensure_ascii=False))
    for sheet in rec["sheets"]:
        for block in sheet["blocks"]:
            for seg in block.get("segments", []):
                if seg.get("table") == old:
                    seg["table"] = new
                if seg.get("id") == old:
                    seg["id"] = new
    for t in rec["tables"]:
        if t["name"] == old:
            t["name"] = new
    for r in rec["relations"]:
        if r.get("table") == old:
            r["table"] = new
        for side in ("a", "b"):
            if isinstance(r.get(side), dict) and r[side].get("table") == old:
                r[side]["table"] = new
    return rec


async def _commit(c: AsyncClient, st: dict, acceptances: list[dict] | None = None) -> dict:
    body: dict = {"trial_id": st["trial"]["trial_id"], "confirmations": [i["id"] for i in st["trial"]["confirm_items"]]}
    if acceptances:
        body["acceptances"] = acceptances
    r = await c.post(f"{DS_API}/imports/{st['id']}/commit", json=body)
    assert r.status_code == 201, r.text
    return r.json()


async def import_flow(c: AsyncClient, *, mode: str = "replace", seed: int = 4101, start=None, days: int = 31,
                      variant: str = "D00") -> Upload:
    """合成客流表：宽表改名「日客流」后导入。"""
    from tests.fixtures.xlsx.drift import AUG

    raw, fn = flow_workbook(start or AUG, days, seed=seed, variant=variant)
    return await import_workbook(c, raw, fn, mode=mode)


async def import_workbook(c: AsyncClient, raw: bytes, fn: str, *, mode: str = "replace",
                          rename: bool = True) -> Upload:
    """暂存 → 回答问题 → （客流表）宽表改名「日客流」→ 试运行 → 勾全部确认项提交（同期 3 验收的公共做法）。
    rename=False 给不是客流表的夹具（lab.c01 列表），表名照规则起草的。"""
    name = f"p4c2_{uuid.uuid4().hex[:8]}"
    r = await c.post(f"{DS_API}/imports/stage", files={"file": (fn, raw, XLSX)}, data={"name": name})
    assert r.status_code == 201, r.text
    st = r.json()
    if st.get("questions"):
        r = await c.post(f"{DS_API}/imports/{st['id']}/answers", json={"answers": _answers(st["questions"], mode)})
        assert r.status_code == 200, r.text
        st = r.json()
    if rename:
        wide = next(t["name"] for t in st["recipe"]["tables"] if t["name"] not in FLOW_TABLES)
        r = await c.put(f"{DS_API}/imports/{st['id']}/recipe",
                        json={"recipe": _rename_wide(st["recipe"], wide, "日客流")})
        assert r.status_code == 200, r.text
    r = await c.post(f"{DS_API}/imports/{st['id']}/trial", json={})
    assert r.status_code == 200, r.text
    st = r.json()
    assert st["trial"]["status"] == "passed", st["trial"]["problems"]
    out = await _commit(c, st)
    return Upload(name=name, source_id=out["source"]["id"], snapshot=out["snapshot_id"], imports=[out["import_id"]],
                  raws=[raw], files=[fn])


async def add_period(c: AsyncClient, up: Upload, *, seed: int, variant: str = "D00", days: int = 30,
                     acceptances: list[dict] | None = None) -> Upload:
    from tests.fixtures.xlsx.drift import SEP

    raw, fn = flow_workbook(SEP, days, seed=seed, variant=variant)
    r = await c.post(f"{DS_API}/{up.source_id}/reupload", files={"file": (fn, raw, XLSX)})
    assert r.status_code == 201, r.text
    st = r.json()
    assert st["trial"]["status"] in ("passed", "needs_decision"), st["trial"]["problems"]
    out = await _commit(c, st, acceptances)
    return Upload(name=up.name, source_id=up.source_id, snapshot=out["snapshot_id"],
                  imports=[*up.imports, out["import_id"]], raws=[*up.raws, raw], files=[*up.files, fn])


@pytest.fixture
async def rep(upload_world, client):
    """8 月 D00，每期替换（单期快照）。"""
    up = await import_flow(client)
    yield up
    await engines.invalidate(up.source_id)


async def run_sql(c: AsyncClient, up: Upload, sql: str) -> tuple[str, str]:
    """经工具接口真跑一次查询（上传源查当前版本）：(查询快照 id, 表结构快照 id)。存下的 SQL 是守卫规范过的那一条。"""
    r = await c.post(f"/api/tools/db_query__{up.name}/run", json={"args": {"sql": sql}})
    assert r.status_code == 200, r.text
    out = r.json()
    assert out["ok"], out
    payload = json.loads(out["result"])
    assert artifact_store.load(payload["artifact"])["data_version"] == up.snapshot
    return payload["artifact"], payload["schema_artifact"]


async def drill_of(up: Upload, qid: str, sid: str, *, column: str = "全日客流", row: int = 0,
                   trusted: bool = True) -> ProvenanceOut:
    """拼内存里的封存范围（这份查询快照、这份表结构快照），走 segment_provenance；断言契约、JSON 往返不走样。"""
    out = await provenance(doc_of(qid, source=up.name), cell_seg(column=column, row=row),
                           sealed_of(queries=[qid], schemas=[sid], trusted=trusted))
    from_json(json.loads(json.dumps(asdict(out), ensure_ascii=False)))
    return out


async def forged_query(qid: str, **changes) -> str:
    """查询快照改几项另存一份（新 id）：在真数据上造反例。"""
    return await put({**artifact_store.load(qid), **changes})


async def forged_schema(sid: str, edit) -> str:
    schema = copy.deepcopy(artifact_store.load(sid))
    edit(schema)
    return await put(schema, kind="schema_snapshot")


def cell_value(raw: bytes, sheet: str, ref: str):
    return load_workbook(io.BytesIO(raw), data_only=True)[sheet][ref].value


async def edit_row(model, key, **values) -> dict:
    """改数据库里的一行，返回原值（用例最后改回去）。"""
    async with SessionLocal() as session:
        row = await session.get(model, key)
        before = {k: getattr(row, k) for k in values}
        for k, v in values.items():
            setattr(row, k, v)
        await session.commit()
    return before


# --------------------------------------------------------------------------
# 满足全部条件：inferred
# --------------------------------------------------------------------------


async def test_inferred_wide_cell(rep, client):
    """宽表的一格（P4-SPEC 2.5 推演表第一行、6.2 A4）：格子、附带的格、年份、回查、相关核对、数据版本。"""
    qid, sid = await run_sql(client, rep, WIDE_SQL)
    out = await drill_of(rep, qid, sid)
    snap = artifact_store.load(qid)
    assert (out.status, out.reason, out.alert, out.sealed) == ("inferred", None, None, True)
    cs = out.cell_source
    assert (cs.table, cs.column, cs.column_role, cs.kind) == ("日客流", "全日客流", "measure", "data")
    assert (cs.sheet, cs.cell) == (SHEET, "G5") and cs.part_seq == 1 and cs.part_rowid == cs.rowid
    assert cell_value(rep.raws[0], SHEET, "G5") == snap["rows"][0][1]
    assert [(f.role, f.column, f.cell, f.text) for f in cs.from_] == [
        ("axis_header", "日期", "G4", None), ("row_label", None, "B5", "全日客流（人次）")]
    assert cs.year == YearSource(source="period", cells=[f"{SHEET}!B2"], signed_by=None, mixed=False)
    assert cs.pk == {"日期": "2026-08-05"} and cs.raw_purged is False and cs.merged_fill is False
    assert cs.recheck == Recheck(sql='SELECT rowid, "全日客流" FROM "日客流" WHERE "日期" = ?',
                                 params=["2026-08-05"], ok=True)
    r1 = next(c for c in out.checks if c.id == "R1")
    assert (r1.kind, r1.part_status, r1.row_status, r1.acceptance) == ("relation_sum_eq", "passed", "passed", None)
    v = out.version
    assert (v.source, v.snapshot_id, v.mode, v.union, v.manifest_view) == (rep.name, rep.snapshot, "replace", False, True)
    assert v.tables == [TableRef(name="日客流", kind="data")]
    assert len(v.parts) == 1
    part = v.parts[0]
    assert (part.import_id, part.seq, part.has_row, part.file_name) == (rep.imports[0], 1, True, rep.files[0])
    assert part.raw_sha256 == hashlib.sha256(rep.raws[0]).hexdigest()
    assert (part.raw_state, part.purged, part.revoked, part.recipe_seq) == ("kept", None, None, 1)
    assert part.region == f"{SHEET}!B4:AG30" and part.period.start == "2026-08-01"
    # 路由一律 asdict：cell_source 的「from」在 JSON 里叫 from
    assert "from" in asdict(out)["cell_source"] and "from_" not in asdict(out)["cell_source"]


async def test_unsealed_runs_still_infer(rep, client):
    """停在审批上、封存核对不通过的运行（trusted 为假）照常推断：sealed 照实给 false，不标红（1.3 调整 11）。"""
    qid, sid = await run_sql(client, rep, WIDE_SQL)
    out = await drill_of(rep, qid, sid, trusted=False)
    assert (out.status, out.sealed, out.alert) == ("inferred", False, None)
    assert out.cell_source.cell == "G5"


async def test_heavy_steps_run_in_threads(rep, client, monkeypatch):
    """清单取回（resolve_chain）、找回格子和相关核对（要解析配方、还原回执）在线程里跑，不卡事件循环（2.9）。"""
    seen: dict[str, bool] = {}

    def spy(name, real):
        def wrapped(*a, **kw):
            seen[name] = threading.current_thread() is threading.main_thread()
            return real(*a, **kw)
        return wrapped

    for name in ("resolve_chain", "locate", "related_checks"):
        monkeypatch.setattr(api.prov, name, spy(name, getattr(api.prov, name)))
    qid, sid = await run_sql(client, rep, WIDE_SQL)
    assert (await drill_of(rep, qid, sid)).status == "inferred"
    assert seen == {"resolve_chain": False, "locate": False, "related_checks": False}


# --------------------------------------------------------------------------
# R5、R6：清单链
# --------------------------------------------------------------------------


async def test_r6_chain_mismatch_and_r5_manifest_unreadable(rep, client):
    qid, sid = await run_sql(client, rep, WIDE_SQL)
    schema = artifact_store.load(sid)
    [mid] = schema["import_manifests"]

    async def with_schema(edit) -> ProvenanceOut:
        bad = await forged_schema(sid, edit)
        return await drill_of(rep, await forged_query(qid, schema_artifact=bad), bad)

    # R6：单期的 import_manifests 必须恰好一个（0 个、2 个都是链本身对不上）：none，version 为 null，标红
    other = await put({**artifact_store.load(mid), "created_at": "2026-10-01T00:00:00+00:00"}, kind="import_manifest")
    for manifests in ([mid, other], []):
        out = await with_schema(lambda s, m=manifests: s.update(import_manifests=m))
        assert (out.status, out.reason.code, out.alert.code, out.version) == ("none", "chain_mismatch",
                                                                             "chain_mismatch", None)
        assert out.cell_source is None and out.checks == []
    # R5：清单哈希不符、取不回：none，manifest_unreadable，标红
    path = _path(mid)
    original = path.read_text("utf-8")
    path.write_text(original.replace('"seq":1', '"seq":7', 1), encoding="utf-8")
    out = await drill_of(rep, qid, sid)
    assert (out.status, out.reason.code, out.alert.code, out.version) == ("none", "manifest_unreadable",
                                                                         "manifest_unreadable", None)
    path.unlink()
    out = await drill_of(rep, qid, sid)
    assert (out.reason.code, out.alert.code) == ("manifest_unreadable", "manifest_unreadable")
    path.write_text(original, encoding="utf-8")
    assert (await drill_of(rep, qid, sid)).status == "inferred"


# --------------------------------------------------------------------------
# R7–R11：只看查询快照和冻结表结构
# --------------------------------------------------------------------------


@pytest.mark.parametrize("sql,column,code,detail", [
    ('SELECT "日期", "分区甲" AS "分区乙" FROM "日客流" WHERE "日期" = \'2026-08-05\'', "分区乙", "alias", "alias"),
    ('SELECT "日期" "日子", "全日客流" FROM "日客流" WHERE "日期" = \'2026-08-05\'', "全日客流", "alias", "alias"),
    ('SELECT a."日期", b."全日客流" FROM "日客流" a JOIN "日客流" b ON a."日期" = b."日期" '
     'WHERE a."日期" = \'2026-08-05\'', "全日客流", "multi_table", "join"),
    ('SELECT a."日期", b."全日客流" FROM "日客流" a, "日客流" b WHERE a."日期" = b."日期" '
     'AND a."日期" = \'2026-08-05\'', "全日客流", "multi_table", "comma_join"),
    ('SELECT SUM("全日客流") AS "合计", COUNT(*) AS "天数" FROM "日客流"', "合计", "expression", "expression"),
    ('SELECT "日期", "全日客流" FROM "日客流" WHERE "日期" LIKE \'2026-08-0%\' ESCAPE \'\\\' '
     'UNION ALL SELECT "日期", "分区甲" FROM "日客流" WHERE "日期" = \'2026-08-31\'', "全日客流", "multi_table",
     "compound"),
])
async def test_r7_sql_refusals_keep_the_table_level_version(rep, client, sql, column, code, detail):
    """R7：受限文法之外的写法只给表级来历（version 照给），原因按 2.4 分组，细分给测试和日志。"""
    qid, sid = await run_sql(client, rep, sql)
    out = await drill_of(rep, qid, sid, column=column)
    assert (out.status, out.reason.code, out.reason.detail, out.alert) == ("table_only", code, detail, None)
    assert out.cell_source is None and out.checks == []
    assert out.version.tables == [TableRef(name="日客流", kind="data")]
    assert out.version.parts[0].has_row is False


async def test_r8_defensive_column_mapping(rep, client, monkeypatch):
    """R8：被引用的结果列对不上选取项（识别器已拒绝重名，只是防御）→ unparsed，不抛。"""
    qid, sid = await run_sql(client, rep, WIDE_SQL)
    monkeypatch.setattr(api.direct_select, "recognize", lambda sql, **kw: DirectSelect("日客流", None, ["日期"]))
    out = await drill_of(rep, qid, sid)
    assert (out.status, out.reason.code, out.reason.detail) == ("table_only", "unparsed", "")
    out = await drill_of(rep, qid, sid, column="不在结果里")
    assert out.reason.code == "unparsed"


async def test_r9_primary_key(rep, client):
    # 没带齐主键列：原文里写出缺的那一列
    qid, sid = await run_sql(client, rep, 'SELECT "全日客流" FROM "日客流" WHERE "日期" = \'2026-08-05\'')
    out = await drill_of(rep, qid, sid)
    assert (out.status, out.reason.code) == ("table_only", "pk_missing") and "（日期）" in out.reason.text
    # 主键重复（同一张表经集合运算拼出来的结果可能这样）：与解析无关的兜底，multi_table · duplicate_pk
    qid, sid = await run_sql(client, rep, WIDE_SQL)
    row = artifact_store.load(qid)["rows"][0]
    dup = await forged_query(qid, rows=[row, list(row)], row_count=2)
    out = await drill_of(rep, dup, sid)
    assert (out.reason.code, out.reason.detail) == ("multi_table", "duplicate_pk")
    # 冻结表结构里这张表没有主键
    nopk = await forged_schema(sid, lambda s: s["tables"]["日客流"].update(primary_key=[]))
    out = await drill_of(rep, await forged_query(qid, schema_artifact=nopk), nopk)
    assert (out.status, out.reason.code) == ("table_only", "no_pk")


async def test_r10_masks(rep, client):
    """R10：被引用的列、主键列、推断要用到的键列在遮罩里（数据源现在设的 ∪ 快照记下的）都不下钻；数据源设有任何
    遮罩时数据版本一节不出「查看导入清单」。"""
    from app.db.models import DataSource as Source

    wide, sid = await run_sql(client, rep, 'SELECT "日期", "全日客流", "分区甲" FROM "日客流" '
                                            'WHERE "日期" = \'2026-08-05\'')
    hourly, hsid = await run_sql(client, rep, 'SELECT "日期", "时段", "客流" FROM "时段客流" '
                                              'WHERE "日期" = \'2026-08-05\' AND "时段" = \'13-14\'')
    assert (await drill_of(rep, wide, sid)).version.manifest_view is True
    receipt = artifact_store.load(artifact_store.load(hsid)["import_manifests"][0])["receipt"]
    hourly_cols = next(t for t in receipt["tables"] if t["name"] == "时段客流")["columns"]
    const = next(c["name"] for c in hourly_cols if c["role"] == "const")
    derive = next(c["name"] for c in hourly_cols if c["role"] == "derive")

    async def masked_as(mask: str) -> None:
        await edit_row(Source, rep.source_id, options={"mask_columns": mask})

    try:
        await masked_as("分区甲")
        out = await drill_of(rep, wide, sid, column="分区甲")
        assert (out.status, out.reason.code, out.version.manifest_view) == ("table_only", "masked", False)
        # 别的列照常推断，但数据源设了遮罩：不在面板里出导入清单
        out = await drill_of(rep, wide, sid)
        assert (out.status, out.version.manifest_view) == ("inferred", False)
        # 主键列。宽表的「日期」同时是日期表头格写出的列，附带格的复核也会拦；只靠主键这一项拦的情形见
        # test_r10_a_masked_primary_key_is_enough
        await masked_as("日期")
        assert (await drill_of(rep, wide, sid)).reason.code == "masked"
        await masked_as(const)                                    # 常量列：分段标题的定位文字就是它的值
        out = await drill_of(rep, hourly, hsid, column="客流")
        assert (out.status, out.reason.code) == ("table_only", "masked")
        # 派生列（起始小时）不在查询里、也不在附带的格里，但行标签原文就是它的来源：按回执的键列判
        await masked_as(derive)
        out = await drill_of(rep, hourly, hsid, column="客流")
        assert (out.status, out.reason.code) == ("table_only", "masked")
        await masked_as("")
        assert (await drill_of(rep, hourly, hsid, column="客流")).status == "inferred"
        # 查询当时记下的遮罩（数据源事后改名、删掉、撤了遮罩也照遮）
        recorded = await forged_query(wide, mask_columns=["全日客流"])
        out = await drill_of(rep, recorded, sid)
        assert (out.reason.code, out.version.manifest_view) == ("masked", False)
    finally:
        await masked_as("")


async def test_r10_a_masked_primary_key_is_enough(upload_world, client):
    """R10 的「主键列」这一项单独拦得住：列表形态（lab.c01）的主键「地区」角色是 text（不是维度、派生、常量，回执
    键列不含它），附带的格只有 column 为空的列表头（附带格的复核也拦不住）。只遮「地区」、引用「金额」也不下钻，
    否则遮罩的主键值会经 cell_source.pk 和 recheck.params 交给界面。"""
    from app.db.models import DataSource as Source
    from tests.fixtures.xlsx import lab

    up = await import_workbook(client, *lab.c01(), rename=False)
    try:
        qid, sid = await run_sql(client, up, 'SELECT "地区", "金额" FROM "月报" WHERE "地区" = \'分区丙\'')
        schema = artifact_store.load(sid)
        assert schema["tables"]["月报"]["primary_key"] == ["地区"]
        receipt = artifact_store.load(schema["import_manifests"][0])["receipt"]
        roles = {c["name"]: c["role"] for t in receipt["tables"] if t["name"] == "月报" for c in t["columns"]}
        assert roles["地区"] == "text"
        # 前提：没遮罩时推断得出，主键值在 pk 和回查参数里，附带的格不写任何列的值
        clear = await drill_of(up, qid, sid, column="金额")
        assert clear.status == "inferred", clear.reason
        assert clear.cell_source.pk == {"地区": "分区丙"} and clear.cell_source.recheck.params == ["分区丙"]
        assert [(f.role, f.column) for f in clear.cell_source.from_] == [("col_header", None)]
        async with SessionLocal() as session:
            options = dict((await session.get(Source, up.source_id)).options or {})
        await edit_row(Source, up.source_id, options={**options, "mask_columns": "地区"})
        try:
            out = await drill_of(up, qid, sid, column="金额")
        finally:
            await edit_row(Source, up.source_id, options=options)
        assert (out.status, out.reason.code, out.reason.detail, out.alert) == ("table_only", "masked", "", None)
        assert out.cell_source is None and out.checks == [] and out.version.manifest_view is False
    finally:
        await engines.invalidate(up.source_id)


async def test_r11_null_value_and_null_pk(rep, client):
    # 夹具第 10 行（7-8 时）是占位符，存为空值：片段本身是 resolved（显示「—」），不推断来源
    qid, sid = await run_sql(client, rep, 'SELECT "日期", "时段", "客流" FROM "时段客流" '
                                          'WHERE "时段" = \'7-8\' AND "日期" = \'2026-08-03\'')
    assert artifact_store.load(qid)["rows"][0][2] is None
    out = await drill_of(rep, qid, sid, column="客流")
    assert (out.status, out.reason.code, out.alert) == ("table_only", "null_value", None)
    qid, sid = await run_sql(client, rep, WIDE_SQL)
    row = artifact_store.load(qid)["rows"][0]
    out = await drill_of(rep, await forged_query(qid, rows=[[None, row[1]]]), sid)
    assert out.reason.code == "null_pk"


# --------------------------------------------------------------------------
# R12、R13：数据文件记录
# --------------------------------------------------------------------------


async def test_r12_r13_data_file_record(rep, client, monkeypatch):
    from app.db.models import SourceSnapshot

    qid, sid = await run_sql(client, rep, WIDE_SQL)
    async with SessionLocal() as session:
        record = await session.get(SourceSnapshot, rep.snapshot)
        path, sha, source_id = Path(record.db_path), record.db_sha256, record.source_id

    def never(*_a, **_kw):
        raise AssertionError("R12 不满足时不该去绑定、打开数据文件")

    # R12：记录已回收、文件不在、记录不存在：只给表级来历，不标红；在 R13、R14 之前就判掉
    with monkeypatch.context() as m:
        m.setattr(api.provenance_db, "view_of", never)
        before = await edit_row(SourceSnapshot, rep.snapshot, retired_at=datetime.now(timezone.utc))
        out = await drill_of(rep, qid, sid)
        assert (out.status, out.reason.code, out.alert) == ("table_only", "snapshot_gone", None)
        assert out.version is not None and out.cell_source is None
        await edit_row(SourceSnapshot, rep.snapshot, **before)
        moved = path.with_suffix(".moved")
        path.rename(moved)
        assert (await drill_of(rep, qid, sid)).reason.code == "snapshot_gone"
        moved.rename(path)
        stray = await forged_query(qid, data_version="ab" * 32)   # 单期不比快照 id，链照样能解析
        out = await drill_of(rep, stray, sid)
        assert (out.status, out.reason.code) == ("table_only", "snapshot_gone")
    # R13：登记哈希和清单里期望的库哈希不等（含为空）、source_id 不一致：标红。不带细分（契约的 Reason.detail
    # 只装 DETAILS 和 ChainProblem / LocateProblem 的细分），和异常映射的 chain_mismatch 同一个形状（WP-E A5b）
    for change in ({"db_sha256": "0" * 64}, {"db_sha256": ""}, {"source_id": "f" * 32}):
        await edit_row(SourceSnapshot, rep.snapshot, **change)
        try:
            out = await drill_of(rep, qid, sid)
        finally:
            await edit_row(SourceSnapshot, rep.snapshot, db_sha256=sha, source_id=source_id)
        assert (out.status, out.reason.code, out.reason.detail) == ("table_only", "chain_mismatch", "")
        assert out.alert == Alert(code="chain_mismatch", text=out.reason.text)
        assert out.version is not None and out.cell_source is None and out.checks == []
    assert (await drill_of(rep, qid, sid)).status == "inferred"


# --------------------------------------------------------------------------
# R14 之后：异常映射、R15、R16、R17
# --------------------------------------------------------------------------

_FAILURES = {
    "chain_mismatch": lambda: SnapshotManifestMismatch("清单不一致"),
    "db_tampered": lambda: SnapshotTampered("文件被改"),
    "snapshot_gone": lambda: SnapshotMissing("文件不在"),
}


def _raise_at(m: pytest.MonkeyPatch, stage: str, exc: BaseException) -> None:
    async def boom(*_a, **_kw):
        raise exc

    def boom_sync(*_a, **_kw):
        raise exc

    if stage == "view_of":
        m.setattr(api.provenance_db, "view_of", boom_sync)
    elif stage == "engine":
        m.setattr(engines, "get", boom)
    else:
        m.setattr(api.provenance_db, stage, boom)


@pytest.mark.parametrize("stage", ["view_of", "engine", "compile_facts", "recheck", "row_status"])
async def test_snapshot_errors_after_r14_map_in_order(rep, client, monkeypatch, stage):
    """绑定、取引擎、编译核对、回查、行级核对，任何一步抛三类数据文件异常都按 2.3 的顺序映射：不给格子，已算出的
    行级状态一律丢掉，数据版本照给。SnapshotManifestMismatch 同时是另外两类，必须先认成 chain_mismatch。"""
    qid, sid = await run_sql(client, rep, WIDE_SQL)
    clean = await drill_of(rep, qid, sid)
    assert clean.status == "inferred" and any(c.row_status == "passed" for c in clean.checks)
    for code, make in _FAILURES.items():
        with monkeypatch.context() as m:
            _raise_at(m, stage, make())
            out = await drill_of(rep, qid, sid)
        assert (out.status, out.reason.code) == ("table_only", code), (stage, code)
        assert (out.alert.code if out.alert else None) == (None if code == "snapshot_gone" else code)
        assert out.cell_source is None and out.checks == [] and out.version is not None
        assert not any(p.has_row for p in out.version.parts)


async def test_other_errors_are_not_swallowed(rep, client, monkeypatch):
    """不是那三类的异常不当成「推断不出来」吞掉（会是 500，测试和日志看得见）。"""
    qid, sid = await run_sql(client, rep, WIDE_SQL)

    async def boom(*_a, **_kw):
        raise KeyError("意料之外")

    monkeypatch.setattr(api.provenance_db, "recheck", boom)
    with pytest.raises(KeyError):
        await drill_of(rep, qid, sid)


@pytest.mark.parametrize("sql,detail", [
    # 实测 B7：藏在 ESCAPE '\' 后面的 UNION ALL（ResultRow 有 2 个）
    ('SELECT "日期", "全日客流" FROM "日客流" WHERE "日期" LIKE \'2026-08-0%\' ESCAPE \'\\\' '
     'UNION ALL SELECT "日期", "分区甲" FROM "日客流" WHERE "日期" = \'2026-08-31\'', "compound"),
    # UNION（有 Yield）、INTERSECT（ResultRow 只有 1 个，靠 Yield 认出）
    ('SELECT "日期", "全日客流" FROM "日客流" WHERE "日期" = \'2026-08-05\' '
     'UNION SELECT "日期", "分区甲" FROM "日客流" WHERE "日期" = \'2026-08-31\'', "compound"),
    ('SELECT "日期", "全日客流" FROM "日客流" INTERSECT SELECT "日期", "全日客流" FROM "日客流"', "compound"),
    # 子查询读了另一张表：授权回调报出来
    ('SELECT "日期", "全日客流" FROM "日客流" WHERE "日期" IN (SELECT "日期" FROM "时段客流")', "authorizer"),
])
async def test_r15_compile_facts_catch_what_the_recognizer_let_through(rep, client, monkeypatch, sql, detail):
    """R15 是与解析无关的兜底：用桩让识别器放行，编译核对照样拦下（读到的表、ResultRow 恰好 1 个、没有 Yield）。"""
    qid, sid = await run_sql(client, rep, WIDE_SQL)
    monkeypatch.setattr(api.direct_select, "recognize",
                        lambda _sql, **kw: DirectSelect("日客流", None, ["日期", "全日客流"]))
    out = await drill_of(rep, await forged_query(qid, sql=sql), sid)
    assert (out.status, out.reason.code, out.reason.detail, out.alert) == ("table_only", "multi_table", detail, None)


async def test_r15_uncompilable_sql_is_unparsed(rep, client, monkeypatch):
    qid, sid = await run_sql(client, rep, WIDE_SQL)
    monkeypatch.setattr(api.direct_select, "recognize",
                        lambda _sql, **kw: DirectSelect("日客流", None, ["日期", "全日客流"]))
    bad = await forged_query(qid, sql='SELECT "日期", "全日客流" FROM "日客流" WHERE "没有这一列" = 1')
    out = await drill_of(rep, bad, sid)
    assert (out.status, out.reason.code, out.alert) == ("table_only", "unparsed", None)


async def test_r16_recheck(rep, client, monkeypatch):
    qid, sid = await run_sql(client, rep, WIDE_SQL)
    value = artifact_store.load(qid)["rows"][0][1]
    rehashed: list[bool] = []
    real_rehash, real_recheck = api.provenance_db.rehash, api.provenance_db.recheck

    async def spy_rehash(view):
        rehashed.append(await real_rehash(view))
        return rehashed[-1]

    monkeypatch.setattr(api.provenance_db, "rehash", spy_rehash)
    assert (await drill_of(rep, qid, sid)).status == "inferred" and rehashed == []    # 一致时不重算

    def answering(rows):
        async def fake(view, **kw):
            _, sql, params = await real_recheck(view, **kw)
            return rows, sql, params
        return fake

    with monkeypatch.context() as m:
        m.setattr(api.provenance_db, "recheck", answering([]))
        out = await drill_of(rep, qid, sid)
        assert (out.status, out.reason.code, out.reason.detail) == ("table_only", "recheck_missing", "")
        assert out.alert is None
        m.setattr(api.provenance_db, "recheck", answering([(5, value), (6, value)]))
        assert (await drill_of(rep, qid, sid)).reason.code == "recheck_multiple"
    # 类型严格相等：快照里记成 8754.0（float），回查得到 8754（int）不算相等。文件没变：不标红
    as_float = await forged_query(qid, rows=[["2026-08-05", float(value)]])
    out = await drill_of(rep, as_float, sid)
    assert (out.status, out.reason.code, out.alert, rehashed) == ("table_only", "recheck_mismatch", None, [True])
    # 文件原地改过一个字节、再还原修改时间（指纹不变，EngineCache 认不出）：值对不上时强制重算哈希 → 标红
    path = await _db_path(rep)
    _poke(path)
    with monkeypatch.context() as m:
        m.setattr(api.provenance_db, "recheck", answering([(5, value + 1)]))
        out = await drill_of(rep, qid, sid)
    assert (out.status, out.reason.code, out.alert.code) == ("table_only", "db_tampered", "db_tampered")
    assert rehashed == [True, False]


async def test_r16_recheck_errors(rep, client, monkeypatch):
    """回查本身出错：SQLite 报错只给表级来历（recheck_missing，不带契约外的细分、不标红）；连接池事件里抛出、被
    SQLAlchemy 包成 StatementError 的数据文件异常照原样映射（借连接时指纹不符 → db_tampered，标红）。"""
    from sqlalchemy.exc import StatementError

    qid, sid = await run_sql(client, rep, WIDE_SQL)

    def raising(exc: BaseException):
        async def fake(view, **kw):
            raise exc
        return fake

    with monkeypatch.context() as m:
        m.setattr(api.provenance_db, "recheck", raising(sqlite3.OperationalError("no such column")))
        out = await drill_of(rep, qid, sid)
    assert (out.status, out.reason.code, out.reason.detail, out.alert) == ("table_only", "recheck_missing", "", None)
    assert out.cell_source is None and out.checks == [] and out.version is not None
    wrapped = StatementError("借连接时核对失败", "SELECT 1", None, SnapshotTampered("指纹不符"))
    with monkeypatch.context() as m:
        m.setattr(api.provenance_db, "recheck", raising(wrapped))
        out = await drill_of(rep, qid, sid)
    assert (out.status, out.reason.code, out.reason.detail) == ("table_only", "db_tampered", "")
    assert out.alert == Alert(code="db_tampered", text=out.reason.text) and out.cell_source is None


async def _db_path(up: Upload) -> str:
    from app.db.models import SourceSnapshot

    async with SessionLocal() as session:
        return (await session.get(SourceSnapshot, up.snapshot)).db_path


def _poke(path: str) -> None:
    """原地改一个字节（大小不变），再用 os.utime 还原修改时间：EngineCache 的指纹认不出这种改动。"""
    st = os.stat(path)
    os.chmod(path, 0o644)
    with open(path, "r+b") as f:
        f.seek(-1, os.SEEK_END)
        last = f.read(1)
        f.seek(-1, os.SEEK_END)
        f.write(bytes([last[0] ^ 0x01]))
    os.chmod(path, 0o444)
    os.utime(path, ns=(st.st_atime_ns, st.st_mtime_ns))
    assert os.stat(path).st_mtime_ns == st.st_mtime_ns


@pytest.mark.parametrize("problem,code,red", [
    (LocateProblem("no_lineage", "anchor_block"), "no_lineage", False),
    (LocateProblem("chain_mismatch", "rowid_out_of_parts"), "chain_mismatch", True),
])
async def test_r17_locate_problems(rep, client, monkeypatch, problem, code, red):
    qid, sid = await run_sql(client, rep, WIDE_SQL)
    monkeypatch.setattr(api.prov, "locate", lambda *a, **kw: problem)
    out = await drill_of(rep, qid, sid)
    assert (out.status, out.reason.code, out.reason.detail) == ("table_only", code, problem.detail)
    assert (out.alert is not None) == red and out.cell_source is None and out.version is not None


# --------------------------------------------------------------------------
# 相关核对、累积并集、当前状态
# --------------------------------------------------------------------------


async def test_row_status_and_acceptance_on_an_accepted_period(upload_world, client):
    """每期替换：先 D00，再 D26 并写理由接受 R1。不成立的那一天这一行判 mismatch，其余 passed；接受理由挂在 R1 上，
    数据版本只有 D26 这一期（seq 2），它的接受理由列在那一期下面。"""
    up = await add_period(client, await import_flow(client, seed=4102), seed=4103, variant="D26",
                          acceptances=[{"check_id": "R1", "reason": REASON}])
    try:
        qid, sid = await run_sql(client, up, 'SELECT "日期", "全日客流", "分区甲", "分区乙" FROM "日客流"')
        rows = artifact_store.load(qid)["rows"]
        bad = [i for i, r in enumerate(rows) if None not in r[1:] and r[1] != r[2] + r[3]]
        assert len(bad) == 1
        good = next(i for i, r in enumerate(rows) if None not in r[1:] and r[1] == r[2] + r[3])
        outs = {}
        for i in (bad[0], good):
            out = await drill_of(up, qid, sid, row=i)
            assert out.status == "inferred", out.reason
            outs[i] = out
        for i, row_status in ((bad[0], "mismatch"), (good, "passed")):
            r1 = next(c for c in outs[i].checks if c.id == "R1")
            assert (r1.part_status, r1.row_status) == ("mismatch", row_status)
            assert (r1.acceptance.kind, r1.acceptance.reason, r1.acceptance.signed_by) == ("override", REASON, None)
            [part] = outs[i].version.parts
            assert (part.seq, part.import_id, part.has_row) == (2, up.imports[1], True)
            assert [a.check_id for a in part.acceptances] == ["R1"] and part.acceptances[0].reason == REASON
            assert outs[i].cell_source.part_seq == 2
    finally:
        await engines.invalidate(up.source_id)


async def test_union_cell_and_row_status_on_union_rowid(upload_world, client, monkeypatch):
    """累积并集（8 月 + 9 月）：9 月那一格落在第 2 期，part_rowid 按快照清单的区间换算；行级核对按**并集** rowid 跑。"""
    up = await add_period(client, await import_flow(client, mode="accumulate", seed=4104), seed=4105)
    try:
        qid, sid = await run_sql(client, up, 'SELECT "日期", "时段", "客流" FROM "时段客流" '
                                             'WHERE "日期" = \'2026-09-12\' AND "时段" = \'8-9\'')
        out = await drill_of(up, qid, sid, column="客流")
        assert out.status == "inferred", out.reason
        cs = out.cell_source
        schema = artifact_store.load(sid)
        rows = artifact_store.load(schema["snapshot_manifest"])["parts"][1]["rows"]["时段客流"]
        assert cs.part_seq == 2 and cs.part_rowid == cs.rowid - rows["union"][0] + 1 and cs.rowid > cs.part_rowid
        assert cell_value(up.raws[1], SHEET, cs.cell) == artifact_store.load(qid)["rows"][0][2]
        # 两列主键：pk 的键、回查 SQL 的 WHERE 和 params 都按冻结表结构 primary_key 的顺序（契约 Recheck）
        assert list(cs.pk) == schema["tables"]["时段客流"]["primary_key"] == ["日期", "时段"]
        assert cs.recheck == Recheck(sql='SELECT rowid, "客流" FROM "时段客流" WHERE "日期" = ? AND "时段" = ?',
                                     params=["2026-09-12", "8-9"], ok=True)
        v = out.version
        assert (v.mode, v.union) == ("accumulate", True)
        assert [(p.seq, p.has_row) for p in v.parts] == [(1, False), (2, True)]
        assert [p.import_id for p in v.parts] == up.imports

        seen: list[int] = []
        real = api.provenance_db.row_status

        async def spy(view, rule, snapshot_rowid):
            seen.append(snapshot_rowid)
            return await real(view, rule, snapshot_rowid)

        monkeypatch.setattr(api.provenance_db, "row_status", spy)
        qid, sid = await run_sql(client, up, 'SELECT "日期", "全日客流" FROM "日客流" WHERE "日期" = \'2026-09-12\'')
        out = await drill_of(up, qid, sid)
        assert out.status == "inferred" and out.cell_source.part_seq == 2
        assert seen == [out.cell_source.rowid] and out.cell_source.rowid != out.cell_source.part_rowid
        assert next(c for c in out.checks if c.id == "R1").row_status == "passed"
    finally:
        await engines.invalidate(up.source_id)


async def test_recheck_follows_the_primary_key_order_not_the_select_order(rep, client):
    """选取项的顺序和主键相反（「时段」在「日期」前面）：pk 的键、回查的 WHERE 和 params 仍按冻结表结构
    primary_key 的顺序，面板「技术细节」显示的回查 SQL 和表结构一致。"""
    qid, sid = await run_sql(client, rep, 'SELECT "客流", "时段", "日期" FROM "时段客流" '
                                          'WHERE "日期" = \'2026-08-05\' AND "时段" = \'13-14\'')
    out = await drill_of(rep, qid, sid, column="客流")
    assert out.status == "inferred", out.reason
    cs = out.cell_source
    assert list(cs.pk) == artifact_store.load(sid)["tables"]["时段客流"]["primary_key"] == ["日期", "时段"]
    assert cs.pk == {"日期": "2026-08-05", "时段": "13-14"}
    assert cs.recheck == Recheck(sql='SELECT rowid, "客流" FROM "时段客流" WHERE "日期" = ? AND "时段" = ?',
                                 params=["2026-08-05", "13-14"], ok=True)


async def test_current_state_from_the_import_record(rep, client):
    """原件已清除、接受已作废是导入记录上的当前状态：数据版本照实写（StateNote），格子照样给，标原件已清除。"""
    from app.db.models import TableImport

    qid, sid = await run_sql(client, rep, WIDE_SQL)
    purged = {"at": "2026-09-30T08:00:00+00:00", "reason": "合成理由：到期清除", "signed_by": "录入员甲",
              "signed_by_verified": False}
    revoked = {"at": "2026-09-30T09:00:00+00:00", "reason": "合成理由：接受有误", "signed_by": None}
    await edit_row(TableImport, rep.imports[0], raw_state="purged", purged=purged, revoked=revoked)
    out = await drill_of(rep, qid, sid)
    assert out.status == "inferred" and out.cell_source.raw_purged is True and out.cell_source.cell == "G5"
    part = out.version.parts[0]
    assert part.raw_state == "purged"
    assert part.purged == StateNote(at=purged["at"], signed_by="录入员甲", reason=purged["reason"])
    assert part.revoked == StateNote(at=revoked["at"], signed_by=None, reason=revoked["reason"])


# --------------------------------------------------------------------------
# 接口：上传源的单元格（真 runner）
# --------------------------------------------------------------------------


async def test_route_on_an_upload_cell_returns_200(upload_world, client, engine_up, monkeypatch):
    """配方源单元格：查询步骤带 provenance: true，推断来源接口回 200、inferred，契约不变式全过。"""
    from app.providers import mock_model

    up = await import_flow(client, seed=4106)
    text = "2026-08-05 的全日客流为 [[v:Q1.r0.全日客流]] 人次。"
    monkeypatch.setattr(mock_model.MockChatModel, "_decide", lambda self, messages: AIMessage(content=text))
    g = chain(node("start", "input"), node("fetch", "tool", tool=f"db_query__{up.name}", args={"sql": WIDE_SQL}),
              node("write", "report", instructions="写日报", on_violation="flag"),
              node("out", "output", fields=[{"name": "日报", "value": "{{ nodes.write.text }}"}]))
    run_id = (await run_manager.start(graph=g, input_payload={})).id
    for _ in range(400):
        await asyncio.sleep(0.05)
        async with SessionLocal() as session:
            row = await session.get(Run, run_id)
        if row.status in ("succeeded", "failed") and row.manifest_seq is not None:
            break
    assert row.status == "succeeded", row.error
    try:
        doc = load_doc(row)
        cell = seg_of(doc, "cell")
        step = (await client.get(f"/api/runs/{row.id}/evidence/segments/{cell['id']}")).json()["chain"][0]
        assert step["step"] == "query" and step["provenance"] is True
        resp = await client.get(f"/api/runs/{row.id}/evidence/segments/{cell['id']}/provenance")
        assert resp.status_code == 200, resp.text
        out = ok(from_json(resp.json()))
        assert (out.status, out.sealed) == ("inferred", True)
        assert (out.cell_source.sheet, out.cell_source.cell) == (SHEET, "G5")
        assert out.version.snapshot_id == up.snapshot and out.version.parts[0].has_row is True
        assert out.report == ReportRef(node_id="write", doc_artifact=doc_id(row))
    finally:
        await engines.invalidate(up.source_id)
