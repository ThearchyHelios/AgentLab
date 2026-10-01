"""合成 xlsx 夹具的自检（P2-SPEC 9.1–9.4）：坐标与规格的表一致、公式有缓存、变异只改到预定的格子。

夹具是执行器、核对、差异卡验收的输入；它自己错了，下游的验收全都失去意义，所以这里逐格对。
"""
from __future__ import annotations

import datetime as dt
import io
import re

import pytest
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from app.data import xlsx_scan
from tests.fixtures.xlsx import has_full_calc, read_part, sheet_part
from tests.fixtures.xlsx import lab
from tests.fixtures.xlsx.drift import AUG, DRIFT_CASES, GROUPS, OCT, SEP
from tests.fixtures.xlsx.flow import (
    DAY_TITLE,
    FLOW_EXPECT,
    NIGHT_TITLE,
    PLACEHOLDER,
    SHEET,
    VARIANTS,
    flow_workbook,
)
from tests.fixtures.xlsx.mutate import MUTATIONS, mutate

D00_RAW, D00_NAME = flow_workbook(AUG, 31)


def book(raw: bytes, *, data_only: bool = False):
    return load_workbook(io.BytesIO(raw), data_only=data_only)


def cells(raw: bytes, sheet: str = SHEET, *, data_only: bool = False) -> dict[str, object]:
    """非空格 → 值（data_only=False 时公式格是公式原文）。"""
    ws = book(raw, data_only=data_only)[sheet]
    out = {}
    for row in ws.iter_rows():
        for c in row:
            if c.value is not None:
                out[c.coordinate] = c.value
    return out


def col(i: int) -> str:
    """第 i 个日期（0 起）所在的列字母：C 起。"""
    return get_column_letter(3 + i)


# ==========================================================================
# 9.1 D00
# ==========================================================================


def test_d00_layout_matches_spec():
    assert D00_NAME == "月报导出_2026-08-01_2026-08-31.xlsx"
    wb = book(D00_RAW)
    assert wb.sheetnames == [SHEET]
    ws = wb[SHEET]
    assert ws.sheet_state == "visible"
    assert {str(m) for m in ws.merged_cells.ranges} == {"B2:AG2", "B3:AG3", "B8:AG8", "B9:AG9", "B21:AG21"}
    assert ws["B2"].value == "统计时间范围：2026年8月1日至2026年8月31日"
    assert ws["B3"].value == "客流汇总表"
    assert ws["B4"].value is None
    assert [ws[f"{col(i)}4"].value for i in range(31)] == [f"8月{d}日" for d in range(1, 32)]
    assert ws["AH4"].value is None
    assert [ws[f"B{r}"].value for r in (5, 6, 7)] == ["全日客流（人次）", "分区甲（人次）", "分区乙（人次）"]
    assert all(ws.cell(8, c).value is None for c in range(2, 34))
    assert ws["B9"].value == DAY_TITLE and ws["B21"].value == NIGHT_TITLE
    assert [ws[f"B{r}"].value for r in range(10, 21)] == [f"{h}-{h + 1}" for h in range(7, 18)]
    assert [ws[f"B{r}"].value for r in range(22, 28)] == [f"{h}-{h + 1}" for h in range(18, 24)]
    assert [ws[f"B{r}"].value for r in (28, 29, 30)] == ["18-22 时合计", "22-24 时合计", "18-24 时合计"]
    assert all(ws[f"{col(i)}10"].value == PLACEHOLDER for i in range(31))
    assert all(ws.cell(9, c).value is None and ws.cell(21, c).value is None for c in range(3, 34))
    assert ws.max_row == 30


def test_d00_values_and_formulas():
    ws = book(D00_RAW)[SHEET]
    cached = book(D00_RAW, data_only=True)[SHEET]
    for i in range(31):
        c = col(i)
        a, b, total = ws[f"{c}6"].value, ws[f"{c}7"].value, ws[f"{c}5"].value
        assert 3000 <= a <= 9000 and 2000 <= b <= 8000 and total == a + b
        hours = {r: ws[f"{c}{r}"].value for r in [*range(11, 21), *range(22, 28)]}
        assert all(isinstance(v, int) and 50 <= v <= 900 for v in hours.values())
        assert sum(hours.values()) != total          # 各时段之和 ≠ 全日客流
        assert ws[f"{c}28"].value == f"=SUM({c}22:{c}25)"
        assert ws[f"{c}29"].value == f"=SUM({c}26:{c}27)"
        assert ws[f"{c}30"].value == f"=SUM({c}22:{c}27)"
        assert cached[f"{c}28"].value == sum(hours[r] for r in range(22, 26))
        assert cached[f"{c}29"].value == sum(hours[r] for r in (26, 27))
        assert cached[f"{c}30"].value == sum(hours[r] for r in range(22, 28))
    assert not has_full_calc(D00_RAW)


def test_d00_scan_matches_expectations():
    scan = xlsx_scan.scan(D00_RAW)
    assert not scan.full_calc_on_load and [s.name for s in scan.sheets] == [SHEET]
    s = scan.sheets[0]
    assert s.bounds.a1() == FLOW_EXPECT["bounds"]
    assert s.nonempty == FLOW_EXPECT["nonempty"] == sum(FLOW_EXPECT["ledger"].values())
    assert len(s.merged) == FLOW_EXPECT["merged"]
    assert s.formulas == FLOW_EXPECT["formulas"] and s.formulas_uncached.total == 0
    assert s.hidden_rows == [] and s.hidden_cols == []


def test_flow_workbook_is_deterministic_and_seeded():
    assert flow_workbook(AUG, 31)[0] == D00_RAW
    assert flow_workbook(AUG, 31, seed=1)[0] != D00_RAW
    with pytest.raises(ValueError):
        flow_workbook(AUG, 31, variant="D99")


def test_r2_never_equal_across_seeds():
    for seed in range(20):
        ws = book(flow_workbook(SEP, 30, seed=seed)[0])[SHEET]
        for i in range(30):
            c = col(i)
            hours = [ws[f"{c}{r}"].value for r in [*range(11, 21), *range(22, 28)]]
            assert sum(hours) != ws[f"{c}5"].value


# ==========================================================================
# 9.2 漂移用例
# ==========================================================================


def test_drift_cases_catalogue():
    assert len(DRIFT_CASES) == 31 and "D28b" in DRIFT_CASES
    assert set(DRIFT_CASES) <= VARIANTS
    assert sorted(c for g in GROUPS.values() for c in g) == sorted(DRIFT_CASES)
    assert [len(GROUPS[k]) for k in GROUPS] == [12, 8, 3, 6, 2]
    for code, case in DRIFT_CASES.items():
        assert case.code == code
        assert case.expect in ("passed", "confirm", "rejected", "decision")
        if case.expect == "rejected":
            assert case.rows is None and case.diff_kinds is None and case.codes
        else:
            assert case.rows is not None and case.diff_kinds is not None
        if code != "D00" and case.diff_kinds is not None:
            assert "period" in case.diff_kinds
    assert {c for c, k in DRIFT_CASES.items() if k.expect == "passed"} == set(GROUPS["通过"])
    assert DRIFT_CASES["D26"].expect == "decision" and DRIFT_CASES["D08"].expect == "rejected"


def _sheet(raw: bytes, *, data_only: bool = False):
    return book(raw, data_only=data_only)[SHEET]


def _labels(ws) -> list[object]:
    return [ws.cell(r, 2).value for r in range(1, ws.max_row + 1)]


def _row_of(ws, label: str) -> int:
    return _labels(ws).index(label) + 1


#: 每例的「改到了」检查（ws 是 data_only=False 的工作表，cached 是 data_only=True 的）
def _check_variant(code: str, raw: bytes, name: str, start: dt.date, days: int) -> None:
    ws, cached = _sheet(raw), _sheet(raw, data_only=True)
    end = start + dt.timedelta(days=days - 1)
    last = col(days - 1)
    labels = _labels(ws)
    period = f"统计时间范围：{start.year}年{start.month}月{start.day}日至{end.year}年{end.month}月{end.day}日"
    expect_name = f"月报导出_{start.isoformat()}_{end.isoformat()}.xlsx"
    if code == "D23":
        assert name == "9月客流.xlsx"
    else:
        assert name == expect_name
    if code in ("D00", "D01", "D02", "D03", "D03b", "D23"):     # D23 只改文件名（上面已查）
        assert ws["B2"].value == period and ws[f"{last}4"].value == f"{end.month}月{end.day}日"
        assert ws.cell(4, 4 + days).value is None
    elif code == "D12":
        assert ws["B2"].value == f"统计时间范围：{start.year}年{start.month}月{start.day}日至{end.month}月{end.day}日"
    elif code == "D14":
        assert ws["B8"].value == DAY_TITLE and ws["B29"].value == "18-24 时合计" and ws.max_row == 29
        assert ws["C27"].value == "=SUM(C21:C24)"
    elif code == "D15":
        assert {str(m) for m in ws.merged_cells.ranges} == {f"B2:{last}2", f"B3:{last}3", f"B8:{last}8"}
    elif code == "D17":
        assert ws["B3"].value == "客流汇总表（2026年9月）"
    elif code == "D25":
        assert has_full_calc(raw) and cached["C28"].value is not None
    elif code == "D27":
        assert ws["B3"].value == "客流汇总表（2026年9月1日至2026年9月30日）"
    elif code == "D05":
        assert ws["B9"].value == NIGHT_TITLE and ws["B19"].value == DAY_TITLE
        assert ws["B16"].value == "18-22 时合计" and ws["C16"].value == "=SUM(C10:C13)"
    elif code == "D06":
        assert ws["B5"].value == DAY_TITLE and ws["B28"].value == "日客流（人次）"
        assert labels[28:31] == ["全日客流（人次）", "分区甲（人次）", "分区乙（人次）"] and ws["B27"].value is None
    elif code in ("D07", "D08"):
        for r, span in ((28, range(22, 26)), (29, (26, 27)), (30, range(22, 28))):
            for i in range(days):
                c = col(i)
                v = ws[f"{c}{r}"].value
                assert isinstance(v, int)
                want = sum(ws[f"{c}{h}"].value for h in span)
                assert v == want + (100 if (code == "D08" and r == 28 and i == 4) else 0)
        assert not any(isinstance(v, str) and v.startswith("=") for v in cells(raw).values())
    elif code == "D10":
        assert ws["C4"].value == dt.datetime(2026, 9, 1) and ws["C4"].number_format == 'm"月"d"日"'
    elif code == "D11":
        assert ws["B10"].value == "07:00-08:00" and ws["B27"].value == "23:00-24:00"
    elif code == "D19":
        assert [ws[f"B{r}"].value for r in (5, 6, 7)] == ["分区甲（人次）", "分区乙（人次）", "全日客流（人次）"]
    elif code == "D22":
        assert [ws[f"B{r}"].value for r in (28, 29, 30)] == ["18:00-22:00合计", "22:00-24:00合计", "18:00-24:00合计"]
    elif code == "D24":
        assert ws["B10"].value == "7–8" and ws["B28"].value == "18-22 时合计"
    elif code == "D20":
        wb = book(raw)
        assert wb.sheetnames == [SHEET, "说明"] and wb["说明"].sheet_state == "visible" and wb["说明"]["A1"].value
    elif code == "D28":
        assert ws["B31"].value is None and ws["B32"].value == "注：9月15日闸机故障，当日客流为估算值"
    elif code == "D28b":
        assert ws["B31"].value is None and ws["B32"].value == "注：2026年9月数据为初步统计，待修订"
    elif code == "D04":
        assert "7-8" not in labels and ws["B10"].value == "8-9" and ws["B20"].value == NIGHT_TITLE
    elif code == "D09":
        assert ws["B9"].value == "日间分时段客流（人次）"
    elif code == "D13":
        assert ws["B8"].value == "分区丙（人次）" and ws["B10"].value == DAY_TITLE
        assert all(ws[f"{col(i)}5"].value == sum(ws[f"{col(i)}{r}"].value for r in (6, 7, 8)) for i in range(days))
    elif code == "D16":
        extra = get_column_letter(4 + days - 1)
        assert ws[f"{extra}4"].value == "合计" and ws[f"{extra}9"].value is None
        for r in (5, 10, 28):
            assert ws[f"{extra}{r}"].value == f"=SUM(C{r}:{last}{r})"
            nums = [ws[f"{col(i)}{r}"].value for i in range(days)]
            want = sum(v for v in nums if isinstance(v, int)) if r != 28 else sum(
                cached[f"{col(i)}28"].value for i in range(days))
            assert cached[f"{extra}{r}"].value == want
    elif code == "D18":
        assert ws["B21"].value == "7-18 时合计" and ws["C21"].value == "=SUM(C10:C20)"
        assert cached["C21"].value == sum(ws[f"C{r}"].value for r in range(11, 21))
        assert ws["B22"].value == NIGHT_TITLE
    elif code == "D21":
        assert ws["B31"].value is None
        assert [ws[f"{c}32"].value for c in "BCD"] == ["补录（人次）", "1,234", "2,345"]
    elif code == "D26":
        c = col(9)
        assert ws[f"{c}5"].value == ws[f"{c}6"].value + ws[f"{c}7"].value + 7
        assert all(ws[f"{col(i)}5"].value == ws[f"{col(i)}6"].value + ws[f"{col(i)}7"].value
                   for i in range(days) if i != 9)
    else:  # pragma: no cover
        raise AssertionError(f"没有为 {code} 写检查")


_STARTS = {"D00": (AUG, 31), "D02": (OCT, 31), "D03": (OCT, 15), "D03b": (OCT, 7)}


@pytest.mark.parametrize("code", sorted(DRIFT_CASES))
def test_drift_case_builds_with_the_intended_change(code):
    case = DRIFT_CASES[code]
    start, days = _STARTS.get(code, (SEP, 30))
    for seed in (0, 1):
        raw, name = case.build(seed)
        _check_variant(code, raw, name, start, days)
        scan = xlsx_scan.scan(raw)           # 扫描层能读（不在扫描阶段被拒）
        assert scan.sheets[0].formulas_uncached.total == 0
        assert has_full_calc(raw) is (code == "D25")
    if case.rows is not None and code not in ("D13",):
        assert case.rows == (days, 17 * days, 3 * days)


def test_d24w_hour_labels_are_fullwidth():
    """期 3 的 D24w（P3-SPEC 11.1）：时段标签用全角数字和全角减号，规范写法仍是「7-8」；其余与 D01 相同。"""
    from app.data import recipe_parsers as P

    raw, name = flow_workbook(SEP, 30, variant="D24w")
    base, base_name = flow_workbook(SEP, 30)
    assert name == base_name
    ws, plain = _sheet(raw), _sheet(base)
    hours = [*range(10, 21), *range(22, 28)]
    assert [ws[f"B{r}"].value for r in (10, 13, 27)] == ["７－８", "１０－１１", "２３－２４"]
    for r in hours:
        label = ws[f"B{r}"].value
        assert all(not ("0" <= ch <= "9") and ch != "-" for ch in label), label
        assert "－" in label and P.hour_range(label).canonical == plain[f"B{r}"].value
    # 只改了时段标签：其余格（数字、合计标签、公式、统计期）与 D01 逐格相同
    changed = _diff(cells(base), cells(raw))
    assert changed == {f"B{r}" for r in hours}
    assert _diff(cells(base, data_only=True), cells(raw, data_only=True)) == changed
    assert ws["B28"].value == "18-22 时合计" and not has_full_calc(raw)
    assert flow_workbook(SEP, 30, variant="D24w")[0] == raw
    xlsx_scan.scan(raw)


def test_noperiod_variant_drops_b2_only():
    raw, _ = flow_workbook(SEP, 30, variant="noperiod")
    ws = _sheet(raw)
    assert ws["B2"].value is None and ws["B3"].value == "客流汇总表" and ws["C4"].value == "9月1日"
    assert "B2:AF2" not in {str(m) for m in ws.merged_cells.ranges}


# ==========================================================================
# 9.4 数据变异
# ==========================================================================


def _diff(a: dict[str, object], b: dict[str, object]) -> set[str]:
    return {k for k in a.keys() | b.keys() if a.get(k) != b.get(k)}


def _dimension(raw: bytes) -> str:
    return re.search(r'<dimension ref="([^"]+)"/>', read_part(raw, sheet_part(raw, SHEET))).group(1)


TOTAL_CELLS = {f"{col(i)}{r}" for i in range(31) for r in (28, 29, 30)}
AH_CELLS = {"AH4"} | {f"AH{r}" for r in (5, 6, 7, *range(10, 21), *range(22, 31))}
AG_CELLS = {"AG4"} | {f"AG{r}" for r in (5, 6, 7, *range(10, 21), *range(22, 31))}

#: 代号 → (公式视图里变了的格, 缓存值视图里变了的格, 变异后的 <dimension>)
MUTATION_TARGETS = {
    "P1": (TOTAL_CELLS, set(), "B2:AG30"),
    "P1b": (set(), TOTAL_CELLS, "B2:AG30"),
    "P2": ({"D15"}, {"D15"}, "B2:AG30"),
    "P6": ({"R4"}, {"R4"}, "B2:AG30"),
    "P9": ({"C28"}, {"C28"}, "B2:AG30"),
    "P20": ({"B32", "C32", "D32"}, {"B32", "C32", "D32"}, "B2:AG30"),
    "P23": (set(), set(), "B2:AG27"),
    "P24": (set(), set(), "B2:AF30"),
    "P25": (AH_CELLS, AH_CELLS, "B2:AG30"),
    "P26": (AG_CELLS, AG_CELLS, "B2:AF30"),
    "PKDUP": ({"B12"}, {"B12"}, "B2:AG30"),
}


def test_mutation_catalogue():
    assert set(MUTATIONS) == set(MUTATION_TARGETS)
    assert {"P1", "P1b", "P2", "P6", "P9", "P20", "P23", "P24", "P25", "P26"} <= set(MUTATIONS)
    with pytest.raises(ValueError):
        mutate(D00_RAW, "P99")
    with pytest.raises(ValueError):          # 只接受 D00 的版式
        mutate(flow_workbook(SEP, 30, variant="D21")[0], "P1")


@pytest.mark.parametrize("case", sorted(MUTATION_TARGETS))
def test_mutation_changes_only_the_intended_cells(case):
    raw = mutate(D00_RAW, case)
    formula_cells, cached_cells, dimension = MUTATION_TARGETS[case]
    assert _diff(cells(D00_RAW), cells(raw)) == formula_cells
    assert _diff(cells(D00_RAW, data_only=True), cells(raw, data_only=True)) == cached_cells
    assert _dimension(raw) == dimension
    assert has_full_calc(raw) is False
    xlsx_scan.scan(raw)


def test_mutation_values():
    before, before_v = cells(D00_RAW), cells(D00_RAW, data_only=True)
    after = cells(mutate(D00_RAW, "P1"), data_only=False)
    assert all(isinstance(after[c], int) and after[c] == before_v[c] for c in TOTAL_CELLS)
    after_v = cells(mutate(D00_RAW, "P1b"), data_only=True)
    assert all(c not in after_v for c in TOTAL_CELLS)
    assert "D15" not in cells(mutate(D00_RAW, "P2"))
    p6 = cells(mutate(D00_RAW, "P6"))
    assert p6["R4"] == p6["Q4"] == "8月15日"
    p9, p9v = cells(mutate(D00_RAW, "P9")), cells(mutate(D00_RAW, "P9"), data_only=True)
    assert p9["C28"] == "=SUM(C22:C24)" and p9v["C28"] == sum(before[f"C{r}"] for r in (22, 23, 24))
    p20 = cells(mutate(D00_RAW, "P20"))
    assert (p20["B32"], p20["C32"], p20["D32"]) == ("备用分区（人次）", 12, 15)
    p25 = cells(mutate(D00_RAW, "P25"))
    assert p25["AH4"] == "合计" and all(p25[c] == 999 for c in AH_CELLS - {"AH4"})
    assert cells(mutate(D00_RAW, "PKDUP"))["B12"] == "8－9"
    scan = xlsx_scan.scan(mutate(D00_RAW, "P1b")).sheets[0]
    assert scan.formulas_uncached.total == 93


# ==========================================================================
# 9.3 实验室用例
# ==========================================================================


def _all(raw: bytes) -> dict[tuple[str, str], object]:
    wb = book(raw)
    return {(ws.title, c.coordinate): c.value for ws in wb.worksheets for row in ws.iter_rows() for c in row
            if c.value is not None}


LAB_CASES = ([("c01", v, lab.c01(v)) for v in lab.C01_VARIANTS] + [("c04", "", lab.c04())]
             + [("c05", v, lab.c05(v)) for v in lab.C05_VARIANTS] + [("c06", "", lab.c06())]
             + [("c08", v, lab.c08(v)) for v in lab.C08_VARIANTS] + [("c10", m, lab.c10(m)) for m in lab.C10_MONTHS]
             + [("kv", "", lab.kv_form()), ("c01", "shift", lab.c01("literal", shift=(3, 2)))])


@pytest.mark.parametrize("name,variant,built", LAB_CASES, ids=[f"{n}-{v}" for n, v, _ in LAB_CASES])
def test_lab_cases_build_with_fake_names(name, variant, built):
    raw, filename = built
    assert filename.endswith(".xlsx")
    xlsx_scan.scan(raw)
    text = " ".join(str(v) for v in _all(raw).values())
    for real in ("华东", "华北", "华南", "张三", "secret", "口令"):
        assert real not in text


def test_lab_c01():
    for variant in lab.C01_VARIANTS:
        raw, _ = lab.c01(variant)
        ws, cached = book(raw)["月报"], book(raw, data_only=True)["月报"]
        assert [ws.cell(5, c).value for c in range(3, 7)] == ["地区", "产品", "销量", "金额"]
        assert ws["C9"].value == "合计" and {str(m) for m in ws.merged_cells.ranges} == {"C1:F1"}
        digits = [ref for ref in ("C1", "C2", "C3", "C11", "C12") if any(ch.isdigit() for ch in ws[ref].value)]
        assert digits == ["C1", "C3"]
        if variant == "literal":
            assert (ws["E9"].value, ws["F9"].value) == (45, 300)
        else:
            assert (ws["E9"].value, ws["F9"].value) == ("=SUM(E6:E8)", "=SUM(F6:F8)")
            want = (45, 300) if variant == "cached" else (None, None)
            assert (cached["E9"].value, cached["F9"].value) == want
        assert not has_full_calc(raw)
    with pytest.raises(ValueError):
        lab.c01("nope")


def test_lab_c01_shift():
    """期 3 的 c01 挪位（P3-SPEC 4.6 第 2 条、11.1）：整张表往下 3 行、往右 2 列，表头落在 E8，上方另加两行
    不含数字的说明；shift=(0, 0) 与不传参数逐字节相同。"""
    for variant in lab.C01_VARIANTS:
        assert lab.c01(variant, shift=(0, 0)) == lab.c01(variant)
    for variant in lab.C01_VARIANTS:
        raw, name = lab.c01(variant, shift=(3, 2))
        assert name == f"c01_{variant}_shifted.xlsx"
        ws, cached = book(raw)["月报"], book(raw, data_only=True)["月报"]
        assert [ws.cell(8, c).value for c in range(5, 9)] == ["地区", "产品", "销量", "金额"]
        assert [ws.cell(9, c).value for c in range(5, 9)] == ["分区甲", "产品甲", 10, 100]
        assert ws["E4"].value == "2026年8月 销售月报" and {str(m) for m in ws.merged_cells.ranges} == {"E4:H4"}
        assert (ws["E5"].value, ws["E12"].value, ws["E14"].value) == ("单位：万元", "合计", "注：数据来源于业务系统，金额含税。")
        assert [ws["E1"].value, ws["E2"].value] == list(lab.C01_SHIFT_NOTES) and ws["E3"].value is None
        # 原来的位置都空了：C 列、第 5 行左侧没有任何格
        assert all(ws.cell(r, c).value is None for r in range(1, 16) for c in (1, 2, 3, 4))
        texts = {ref: ws[ref].value for ref in ("E1", "E2", "E4", "E5", "E6", "E14", "E15")}
        assert [ref for ref, v in texts.items() if any(ch.isdigit() for ch in v)] == ["E4", "E6"]
        if variant == "literal":
            assert (ws["G12"].value, ws["H12"].value) == (45, 300)
        else:
            assert (ws["G12"].value, ws["H12"].value) == ("=SUM(G9:G11)", "=SUM(H9:H11)")
            want = (45, 300) if variant == "cached" else (None, None)
            assert (cached["G12"].value, cached["H12"].value) == want
        assert not has_full_calc(raw)
        assert lab.c01(variant, shift=(3, 2))[0] == raw
    # 挪位只改坐标：非空格的值（公式里的引用除外）与原表一一对应
    plain = sorted(str(v) for v in cells(lab.c01("literal")[0], "月报").values())
    moved = cells(lab.c01("literal", shift=(3, 2))[0], "月报")
    assert sorted(str(v) for v in moved.values() if v not in lab.C01_SHIFT_NOTES) == plain
    with pytest.raises(ValueError):
        lab.c01("literal", shift=(-1, 0))


def test_lab_c04_merges():
    ws = book(lab.c04()[0])["合并"]
    assert {str(m) for m in ws.merged_cells.ranges} == {"A1:F1", "A3:A4", "B3:B4", "C3:D3", "E3:F3", "A5:A7", "A8:A10"}
    assert (ws["C3"].value, ws["E3"].value) == ("2026年上半年", "2026年下半年")
    assert [ws.cell(4, c).value for c in range(3, 7)] == ["销量", "金额", "销量", "金额"]
    assert (ws["A5"].value, ws["A8"].value, ws["A6"].value) == ("分区甲", "分区乙", None)


def test_lab_c05_formula_variants():
    want = {"uncached": ([None] * 4, False), "zero_full_calc": ([0] * 4, True), "cached": ([10, 30, 10, 50], False)}
    for variant, (values, full) in want.items():
        raw, _ = lab.c05(variant)
        ws, cached = book(raw)["明细"], book(raw, data_only=True)["明细"]
        assert [ws[f"D{r}"].value for r in range(2, 6)] == ["=B2*C2", "=B3*C3", "=B4*C4", "=SUM(D2:D4)"]
        assert [cached[f"D{r}"].value for r in range(2, 6)] == values
        assert has_full_calc(raw) is full
    assert lab.c05("openpyxl")[0] == lab.c05("uncached")[0]


def test_lab_c06_hidden():
    raw, _ = lab.c06()
    wb = book(raw)
    assert [(ws.title, ws.sheet_state) for ws in wb.worksheets] == [
        ("明细", "visible"), ("草稿", "hidden"), ("配置", "veryHidden")]
    ws = wb["明细"]
    hidden = sorted(r for r, d in ws.row_dimensions.items() if d.hidden)
    assert hidden == [3, 4, 5, 7, 9, 11]
    assert ws.column_dimensions["C"].hidden and ws.auto_filter.ref == "A1:D11"
    s = xlsx_scan.scan(raw).sheets[0]
    assert s.hidden_rows == [(3, 5), (7, 7), (9, 9), (11, 11)] and s.hidden_cols == [(3, 3)]


def test_lab_c08_layouts():
    ws = book(lab.c08("stacked")[0])["汇总"]
    assert (ws["B1"].value, ws["B2"].value, ws["B7"].value, ws["B8"].value) == ("一、销售", "地区", "二、费用", "部门")
    ws = book(lab.c08("gap")[0])["明细"]
    assert [ws.cell(4, c).value for c in (1, 2, 3)] == [None, None, None] and ws["A5"].value == "分区乙"
    ws = book(lab.c08("side")[0])["并排"]
    assert (ws["A1"].value, ws["C1"].value, ws["D1"].value) == ("地区", None, "部门")


def test_lab_c10_months():
    want = {7: ("C5", ["地区", "产品", "销量", "金额"]), 8: ("C7", ["地区", "产品", "单价", "销量", "金额"]),
            9: ("B6", ["地区　", "产品", "销量\n", "金 额"]), 10: ("C5", ["地区", "产品", "销量", "销售额"])}
    for month, (anchor, headers) in want.items():
        raw, name = lab.c10(month)
        ws = book(raw)["月报"]
        cell = ws[anchor]
        assert [ws.cell(cell.row, cell.column + j).value for j in range(len(headers))] == headers
        assert ws.cell(cell.row + 4, cell.column).value == "合计"
        assert name == f"c10_2026-{month:02d}.xlsx"
        notes = [ws.cell(r, cell.column).value for r in range(1, cell.row)]
        assert not any(isinstance(v, str) and any(ch.isdigit() for ch in v) for v in notes)
    with pytest.raises(ValueError):
        lab.c10(11)


def test_lab_kv_form_has_no_header_row():
    ws = book(lab.kv_form()[0])["基本情况"]
    assert ws["A1"].value == "基本情况登记表" and ws["A3"].value == "单位名称" and ws.max_column == 2
