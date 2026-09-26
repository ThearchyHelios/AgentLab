"""把上传的文件变成文本。

在此之前 upload 只接受 UTF-8：PDF 传上去直接 415，而 mime 字段记了却从不
参与任何判断。知识库里最常见的恰恰是 PDF 和 Word。

PDF/DOCX 的解析库走可选依赖（pyproject 的 [docs]），和 [db]、[microvm] 一样
——它们加起来几十兆，而只用纯文本的人不该为此付出安装时间。没装时给的报错
要能照着做，不是一句 ImportError。

HTML 用已经在依赖里的 beautifulsoup4，不额外引入东西。
"""

from __future__ import annotations

import io
import re

#: 认得出的扩展名 → 处理方式。mime 不可靠（浏览器给 .md 报 application/octet-stream）
_PDF = {".pdf"}
_DOCX = {".docx"}
_PPTX = {".pptx"}
_HTML = {".html", ".htm"}

#: 按 UTF-8 文本读的常见格式。清单外的文件也会按文本试，这里列的是上传框里
#: 该让人选得到的那些——前端的 accept 从 GET /api/kb/formats 取，只在这一处定义
_TEXT = (".txt", ".md", ".markdown", ".json", ".log", ".yaml", ".yml", ".py", ".ts", ".js")

#: 表格数据。它们有更好的去处——数据源那条路会把它变成可以真查的表，
#: 而传进知识库只能当文本检索：数字是模型"读"出来的不是"算"出来的，
#: 而这个项目的地基恰恰是"所有算术下沉到 SQL 或口径卡"（见 engine/issuance.py）
_TABULAR = {".xlsx", ".xls", ".csv", ".tsv"}

#: 老的二进制 Office 格式。解析它们要装 libreoffice 之类的外部转换器，代价和
#: 收益不成比例——但认出来并说清楚"另存为新格式"，比让用户对着 ImportError
#: 或者"不是 UTF-8 文本"发愣强得多
_LEGACY_OFFICE = {".ppt": "PowerPoint", ".doc": "Word", ".xls": "Excel"}


class UnsupportedDocument(ValueError):
    """认不出或解析不了。message 直接给用户看。"""


def supported_formats() -> dict[str, list[str]]:
    """知识库收哪些格式、拒哪些。上传框的 accept 和后端的判断必须是同一份清单：
    以前前端手写了一份，漏了 .pptx（后端支持）、多了 .csv（后端必拒）。"""
    extensions = sorted(_PDF | _DOCX | _PPTX) + sorted(_HTML) + list(_TEXT)
    return {
        "extensions": extensions,
        "text": list(_TEXT),
        "tabular": sorted(_TABULAR),
        "legacy": sorted(_LEGACY_OFFICE),
    }


def _why(e: BaseException) -> str:
    """解析库的原话，去掉类名。它们常常是英文，但比"读不开"多一点线索。"""
    text = str(e).strip()
    return (text.splitlines()[0][:120] if text else "") or "没有更多说明"


def _ext(filename: str) -> str:
    idx = filename.rfind(".")
    return filename[idx:].lower() if idx >= 0 else ""


def _from_pdf(raw: bytes) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as e:
        raise UnsupportedDocument(
            "要读 PDF 得先装解析库：pip install 'agentlab-backend[docs]'"
        ) from e
    try:
        reader = PdfReader(io.BytesIO(raw))
    except Exception as e:  # noqa: BLE001
        raise UnsupportedDocument(
            f"这个 PDF 读不开：文件可能损坏、加了密码，或者其实不是 PDF（{_why(e)}）。"
            "用原软件打开另存一份再传"
        ) from e

    # 按页拼，页与页之间留空行——切块是按段落切的，页边界正好是天然的段落边界
    pages = []
    for i, page in enumerate(reader.pages):
        try:
            pages.append(page.extract_text() or "")
        except Exception:  # noqa: BLE001 - 单页坏了不该毁掉整份文档
            pages.append(f"（第 {i + 1} 页解析失败）")
    # PDF 抽出来的每行都拖着排版留下的填充空格，不清掉会混进切块和词频统计
    cleaned = []
    for page in pages:
        lines = [ln.rstrip() for ln in page.splitlines()]
        body = "\n".join(ln for ln in lines if ln.strip())
        if body:
            cleaned.append(body)
    text = "\n\n".join(cleaned)
    if not text.strip():
        raise UnsupportedDocument(
            "这个 PDF 里没有可提取的文字——多半是扫描件。"
            "扫描件要先过 OCR，这里不做图像识别。"
        )
    return text


def _from_docx(raw: bytes) -> str:
    try:
        import docx
    except ImportError as e:
        raise UnsupportedDocument(
            "要读 Word 得先装解析库：pip install 'agentlab-backend[docs]'"
        ) from e
    try:
        doc = docx.Document(io.BytesIO(raw))
    except Exception as e:  # noqa: BLE001
        raise UnsupportedDocument(
            f"这个 Word 文档读不开：文件可能损坏、加了密码，或者其实不是 Word（{_why(e)}）。"
            "用原软件打开另存一份再传"
        ) from e

    parts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    # 表格也要：合同、口径说明这类文档的关键信息常常整个在表里，
    # 只取段落等于把它们全丢了
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    text = "\n\n".join(parts)
    if not text.strip():
        raise UnsupportedDocument("这个 Word 文档里没有文字内容")
    return text


def _from_pptx(raw: bytes) -> str:
    try:
        from pptx import Presentation
    except ImportError as e:
        raise UnsupportedDocument(
            "要读 PowerPoint 得先装解析库：pip install 'agentlab-backend[docs]'"
        ) from e
    try:
        deck = Presentation(io.BytesIO(raw))
    except Exception as e:  # noqa: BLE001
        raise UnsupportedDocument(
            f"这个 PowerPoint 读不开：文件可能损坏、加了密码，或者其实不是 PowerPoint（{_why(e)}）。"
            "用原软件打开另存一份再传"
        ) from e

    slides: list[str] = []
    for page in deck.slides:
        parts: list[str] = []
        for shape in page.shapes:
            # 表格和文本框要分开取：一张幻灯片的关键信息常常整个在表里，
            # 而 has_text_frame 对表格是 False
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    cells = [c.text.strip() for c in row.cells if c.text.strip()]
                    if cells:
                        parts.append(" | ".join(cells))
            elif getattr(shape, "has_text_frame", False):
                text = shape.text_frame.text.strip()
                if text:
                    parts.append(text)
        if parts:
            slides.append("\n".join(parts))

    # 幻灯片之间留空行。和 PDF 按页拼是同一个道理：页/片的边界本来就是
    # 天然的段落边界，而切块是按段落切的
    text = "\n\n".join(slides)
    if not text.strip():
        raise UnsupportedDocument(
            "这个 PowerPoint 里没有可提取的文字——多半整页都是图片。"
            "图里的字要先过 OCR，这里不做图像识别。"
        )
    return text


def _from_html(raw: bytes) -> str:
    from bs4 import BeautifulSoup

    soup = BeautifulSoup(raw.decode("utf-8", errors="replace"), "html.parser")
    for tag in soup(["script", "style", "nav", "footer"]):
        tag.decompose()
    text = soup.get_text("\n")
    # 连续空行压成一个：HTML 转出来到处是空行，会把切块的段落判断搅乱
    return re.sub(r"\n{3,}", "\n\n", text).strip()


def extract(raw: bytes, filename: str, mime: str = "") -> str:
    """文件字节 → 文本。认不出就按 UTF-8 文本试。"""
    ext = _ext(filename)
    if ext in _PDF or mime == "application/pdf":
        return _from_pdf(raw)
    if ext in _DOCX or "wordprocessingml" in mime:
        return _from_docx(raw)
    if ext in _PPTX or "presentationml" in mime:
        return _from_pptx(raw)
    if ext in _HTML or mime == "text/html":
        return _from_html(raw)
    if ext in _TABULAR:
        raise UnsupportedDocument(
            f"「{filename}」是表格，知识库不收表格。请到「数据源」点「传表格」导入："
            "它会变成一张能用 SQL 查的表，数字是算出来的、查得到出自哪条查询。"
            "放进知识库只能按文字检索，数字靠模型去读，容易读错，也追不到来源。"
        )
    if ext in _LEGACY_OFFICE:
        # 认出来再拒，而不是让它掉进下面那个"不是 UTF-8"的兜底——
        # 后者对着一个 .ppt 说"不是文本"，用户根本不知道该怎么办
        raise UnsupportedDocument(
            f"读不了 {filename}：{_LEGACY_OFFICE[ext]} 的老二进制格式（{ext}）需要"
            f"外部转换器才能解析。请在 Office 里另存为 {ext}x 再传。"
        )
    try:
        return raw.decode("utf-8")
    except UnicodeDecodeError as e:
        raise UnsupportedDocument(
            f"读不了 {filename or '这个文件'}：它不是 UTF-8 文本，"
            f"也不是认得出的 PDF / Word / PowerPoint / HTML。"
        ) from e
