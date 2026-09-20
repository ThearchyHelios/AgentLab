"""把 Excel / CSV 变成一个可以用 SQL 查的 SQLite 库。

**为什么不走知识库那条路。** 这个项目有一条不可回退的决策：叙述层无算术权限，
所有算术下沉到 SQL 或口径卡，叙述只能引用（见 engine/issuance.py）。表格数据
传进知识库只能被切块检索，于是数字是模型从片段里"读"出来的——复核层
（engine/review.py）拦得住凭空多出的数字，拦不住"从一堆片段里读错了一格"。
所以表格必须变成表。

**为什么落成 SQLite 而不是新造一种数据源。** kind='sqlite' 早就支持，
engine.build_url 拿 database 当文件路径。落成 .db 文件再建一条 DataSource，
下游一行都不用改：SQL 守卫、结构探查、db_query__<name> 工具、查询快照进
工件库、Copilot 的数据源清单，全都白拿。

**建表必须用这里的直连，不能走 data.engine.run_query。** guard 的
_ALWAYS_DENIED 里有 create，那是对的——模型不该能建表。而导入是我们自己
发起的可信操作，两件事不该共用一条通道。
"""

from __future__ import annotations

import csv
import io
import re
import sqlite3
from dataclasses import dataclass, field
from typing import Any

#: 扫多少行来推断列类型。全表扫一遍对几十万行的文件太慢，而类型这种东西
#: 前几千行看不出来的，后面多半也看不出来——真看错了，SQLite 的动态类型
#: 也还能把值原样存下去，不会丢数据
_TYPE_SAMPLE = 5000

_CSV = {".csv", ".tsv"}
_EXCEL = {".xlsx", ".xlsm"}
_LEGACY_EXCEL = {".xls"}


class UnsupportedTable(ValueError):
    """认不出或读不了。message 直接给用户看。"""


@dataclass
class LoadedTable:
    name: str
    columns: list[str]
    types: list[str]
    rows: int
    #: 原始 sheet 名。清洗成表名之后，用户要能对上是哪一张
    source_name: str = ""


@dataclass
class LoadReport:
    tables: list[LoadedTable] = field(default_factory=list)

    @property
    def total_rows(self) -> int:
        return sum(t.rows for t in self.tables)


def _ext(filename: str) -> str:
    idx = filename.rfind(".")
    return filename[idx:].lower() if idx >= 0 else ""


# --------------------------------------------------------------------------
# 表名 / 列名
# --------------------------------------------------------------------------


def _table_name(raw: str, taken: set[str]) -> str:
    """sheet 名 → 表名。中文原样保留，只替换真正会断句的字符。

    一开始清洗成了纯 ASCII，结果「明细」「汇总」两张表都变成 sheet、sheet_2——
    模型和用户都认不出哪张是哪张。而 SQLite 的标识符本来就允许非 ASCII，
    `FROM 明细` 不加引号也解析得了，所以那道清洗纯属自找麻烦。

    真需要处理的只有空白和标点：`销售 明细` 不引号就断成两个 token。
    正则的词字符类按 Unicode 判定，汉字算词字符，正好留下。
    """
    slug = re.sub(r"[^\w]+", "_", (raw or "").strip(), flags=re.UNICODE).strip("_")
    if not slug:
        slug = "sheet"
    elif slug[0].isdigit():
        slug = f"t_{slug}"      # 标识符不能数字开头
    base, n = slug[:48], 2
    slug = base
    while slug in taken:
        slug = f"{base}_{n}"
        n += 1
    taken.add(slug)
    return slug


def _clean_headers(raw_row: list[Any]) -> list[str]:
    """表头单元格 → 列名。**保持原样**，只治两种病态。

    不转写、不 slug 化：Excel 表头几乎都是中文，转写是有损的，而 SQLite 认
    带引号的标识符，db_schema__<name> 又会把准确列名交给模型。
    """
    out: list[str] = []
    seen: dict[str, int] = {}
    for i, cell in enumerate(raw_row):
        name = str(cell).strip() if cell is not None else ""
        # 空表头给个位置名。合并单元格、末尾多出来的空列都会走到这里
        if not name:
            name = f"col_{i + 1}"
        # 双引号会把带引号的标识符打断，是唯一必须动的字符
        name = name.replace('"', "'")
        if name in seen:
            seen[name] += 1
            name = f"{name}_{seen[name]}"
        else:
            seen[name] = 1
        out.append(name)
    return out


# --------------------------------------------------------------------------
# 类型推断
# --------------------------------------------------------------------------

_INT_RE = re.compile(r"^[+-]?\d{1,18}$")
_FLOAT_RE = re.compile(r"^[+-]?(\d+\.?\d*|\.\d+)([eE][+-]?\d+)?$")


def _norm(value: Any) -> Any:
    """单元格 → 要存进去的值。空的一律 NULL。

    空字符串和 NULL 必须分清楚：存成 "" 的话，一个本该是数字的列就被一个
    空单元格拖成了 TEXT，然后 SUM() 静默地不再是你以为的那个 SUM()。
    """
    if value is None:
        return None
    if isinstance(value, str):
        text = value.strip()
        return text or None
    return value


def infer_type(values: list[Any]) -> str:
    """一列的 SQLite 类型。

    SQLite 是动态类型，但列**亲和性**决定比较和排序的行为：整数列存成 TEXT，
    ORDER BY 就变成字典序，'10' 排在 '9' 前面；而这正是"Excel 里的数字能不能
    真的被 SQL 算"的分界线。所以这一步不能省。

    日期统一成 TEXT——SQLite 没有日期类型，而 ISO 格式的字典序恰好等于时间序，
    存成 ISO 字符串既能排序也能比较。
    """
    import datetime as _dt

    seen = False
    all_int = all_float = True
    for value in values:
        norm = _norm(value)
        if norm is None:
            continue
        seen = True
        if isinstance(norm, bool):
            # bool 是 int 的子类，但"真/假"不是数量，不该被当成 INTEGER 去求和
            return "TEXT"
        if isinstance(norm, (_dt.datetime, _dt.date, _dt.time)):
            return "TEXT"
        if isinstance(norm, int):
            all_float = all_float and True
            continue
        if isinstance(norm, float):
            all_int = False
            continue
        text = str(norm)
        if all_int and not _INT_RE.match(text):
            all_int = False
        if all_float and not _FLOAT_RE.match(text):
            all_float = False
        if not all_float:
            return "TEXT"
    if not seen:
        return "TEXT"      # 整列都是空的，别猜
    return "INTEGER" if all_int else ("REAL" if all_float else "TEXT")


def _coerce(value: Any, sql_type: str) -> Any:
    """按推断出的类型把值转过去。转不动就原样存——宁可类型不齐，不可丢数据。"""
    import datetime as _dt

    norm = _norm(value)
    if norm is None:
        return None
    if isinstance(norm, (_dt.datetime, _dt.date, _dt.time)):
        return norm.isoformat()
    if sql_type == "INTEGER":
        try:
            return int(str(norm))
        except (TypeError, ValueError):
            return str(norm)
    if sql_type == "REAL":
        try:
            return float(str(norm))
        except (TypeError, ValueError):
            return str(norm)
    return norm if isinstance(norm, (int, float)) else str(norm)


# --------------------------------------------------------------------------
# 读文件
# --------------------------------------------------------------------------

#: CSV 的编码。中文环境里 GBK 系的 CSV 极其常见（Excel 另存为 CSV 的默认行为），
#: 一律按 UTF-8 读会得到一堆乱码列名——而乱码是不报错的，只是搜不到、看不懂。
#: 不引入 chardet：按真实场景的频次逐个试就够了，最后 latin-1 保证不抛异常。
_ENCODINGS = ("utf-8-sig", "utf-8", "gb18030", "latin-1")


def _decode_csv(raw: bytes) -> str:
    for enc in _ENCODINGS:
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("latin-1", errors="replace")


def _read_csv(raw: bytes, filename: str, *, header_row: int) -> list[tuple[str, list[list[Any]]]]:
    text = _decode_csv(raw)
    delimiter = "\t" if _ext(filename) == ".tsv" else None
    if delimiter is None:
        try:
            delimiter = csv.Sniffer().sniff(text[:8192], delimiters=",;\t|").delimiter
        except csv.Error:
            delimiter = ","      # 嗅不出来就按最常见的来
    rows = [list(r) for r in csv.reader(io.StringIO(text), delimiter=delimiter)]
    stem = filename.rsplit("/", 1)[-1].rsplit(".", 1)[0]
    return [(stem, rows[header_row - 1:])]


def _read_excel(raw: bytes, *, header_row: int) -> list[tuple[str, list[list[Any]]]]:
    try:
        from openpyxl import load_workbook
    except ImportError as e:
        raise UnsupportedTable(
            "要读 Excel 得先装解析库：pip install 'agentlab-backend[docs]'"
        ) from e
    try:
        # read_only 走流式，几十万行也不会把整个工作簿读进内存；
        # data_only 取公式算出来的值而不是公式本身——用户要的是数
        book = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception as e:  # noqa: BLE001
        raise UnsupportedTable(f"这个 Excel 读不开：{type(e).__name__}: {e}") from e

    out: list[tuple[str, list[list[Any]]]] = []
    for sheet in book.worksheets:
        rows = [list(r) for r in sheet.iter_rows(min_row=header_row, values_only=True)]
        # 整张空的 sheet 就跳过，不建空表。Excel 里常挂着几张没删的空白页
        if any(any(_norm(c) is not None for c in row) for row in rows):
            out.append((sheet.title, rows))
    book.close()
    if not out:
        raise UnsupportedTable("这个 Excel 里没有任何有内容的工作表")
    return out


def read_tables(
    raw: bytes, filename: str, *, header_row: int = 1
) -> list[tuple[str, list[list[Any]]]]:
    """文件字节 → [(sheet 名, 行)]。第一行是表头。"""
    if header_row < 1:
        raise UnsupportedTable("表头行号从 1 开始")
    ext = _ext(filename)
    if ext in _EXCEL:
        return _read_excel(raw, header_row=header_row)
    if ext in _CSV:
        return _read_csv(raw, filename, header_row=header_row)
    if ext in _LEGACY_EXCEL:
        raise UnsupportedTable(
            f"读不了 {filename}：.xls 是老的二进制格式，需要外部转换器。"
            "请在 Excel 里另存为 .xlsx 或 .csv 再传。"
        )
    raise UnsupportedTable(
        f"读不了 {filename}：只认 Excel（.xlsx）和 CSV / TSV。"
    )


# --------------------------------------------------------------------------
# 建库
# --------------------------------------------------------------------------


def load_into(
    db_path: str, raw: bytes, filename: str, *, header_row: int = 1
) -> LoadReport:
    """把文件里的每张表写进 db_path 指的 SQLite 库。

    同名表**整张替换**（DROP + CREATE）。重传的语义是"这份数据更新了"，
    追加会让行数悄悄翻倍，而没有任何地方会提示。

    用 sqlite3 直连而不是 data.engine：那条路上有 SQL 守卫，而守卫拒绝
    CREATE——那是对的，模型不该能建表。导入是我们自己发起的可信操作。
    """
    sheets = read_tables(raw, filename, header_row=header_row)
    report = LoadReport()
    taken: set[str] = set()

    conn = sqlite3.connect(db_path)
    try:
        for sheet_name, rows in sheets:
            if not rows:
                continue
            headers = _clean_headers(rows[0])
            body = rows[1:]
            if not headers:
                continue

            # 行长参差是常态：Excel 的末尾空列、CSV 的缺尾逗号。按表头补齐/截断
            width = len(headers)
            body = [(row + [None] * width)[:width] for row in body]

            sample = body[:_TYPE_SAMPLE]
            types = [infer_type([row[i] for row in sample]) for i in range(width)]

            table = _table_name(sheet_name, taken)
            cols = ", ".join(f'"{h}" {t}' for h, t in zip(headers, types))
            conn.execute(f'DROP TABLE IF EXISTS "{table}"')
            conn.execute(f'CREATE TABLE "{table}" ({cols})')
            if body:
                placeholders = ", ".join("?" * width)
                conn.executemany(
                    f'INSERT INTO "{table}" VALUES ({placeholders})',
                    ([_coerce(row[i], types[i]) for i in range(width)] for row in body),
                )
            report.tables.append(
                LoadedTable(
                    name=table, columns=headers, types=types,
                    rows=len(body), source_name=sheet_name,
                )
            )
        conn.commit()
    finally:
        conn.close()

    if not report.tables:
        raise UnsupportedTable("文件里没有读到任何数据行")
    return report
